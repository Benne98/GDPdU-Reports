import re
import sys
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.styles import Border, Side

PROJECT_ROOT = Path(__file__).resolve().parent
SCRIPTS_DIR = PROJECT_ROOT / "scripts"
BACKEND_DIR = PROJECT_ROOT / "backend"
if str(SCRIPTS_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPTS_DIR))
if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))

from gst_excel_theme import THEME, apply_zero_row_conditional_formatting  # noqa: E402
from databook_periods import ordered_reporting_columns_from_df, split_fy_and_ytd  # noqa: E402
from report_row_layout import build_lead_is_row_structure, l3_order_from_mapping  # noqa: E402

# ==================================================
# DESKTOP DEFAULTS
# ==================================================
from databook_workbook import MASTER_WORKBOOK_STR  # noqa: E402

DESKTOP_DIR = PROJECT_ROOT / "Desktop"

INPUT_FILE = MASTER_WORKBOOK_STR
MAPPING_FILE = str(DESKTOP_DIR / "PL_recon_Mapping.xlsx")
SHEET_MASTER = "Master_PL"
SHEET_OUT = "Lead_IS"
SHEET_RECON = "PL_Reconciliation"

PROJECT_NAME = "Desktop Test"
GROUP_NAME = "Group"
UNIT_LABEL = "kEUR"

SOURCE_COL_CANDIDATES = ("L5", "L6")
REPORTED_FILTER_VALUE = "reported"
L4_SORT_BASIS = "latest_fy"

PL_POSITION_RENAME_MAP = {
    "Other Expenses": "Other operating expenses",
    "Salaries & Bonuses": "Wages & salaries",
    "Reduction in Earnings": "Discounts",
    "Other income": "Other operating income",
    "Personnel Expenses": "Personnel expenses",
    "Social Security": "Social security",
    "Tax": "Taxes on income",
}

LEAD_SPACER_WIDTH = 6.43

TOP_BORDER_ONLY = {"Net sales", "Gross profit", "EBIT", "EBT"}
TOP_AND_BOTTOM = {"EBITDA", "Net result"}

PL_TOTALS_CONFIG_RAW = [
    {
        "label": "Total output",
        "components": ["Net sales", "Δ Finished goods & WIP", "Own work capitalised"],
        "insert_after": "Own work capitalised",
    },
    {
        "label": "Gross profit",
        "components": ["Total output", "Cost of goods sold"],
        "insert_after": "Cost of goods sold",
    },
    {
        "label": "Net operating expenses",
        "components": ["Other operating income", "Other operating expenses"],
        "insert_after": "Other operating expenses",
    },
    {
        "label": "EBITDA",
        "components": ["Gross profit", "Personnel expenses", "Net operating expenses"],
        "insert_after": "Net operating expenses",
    },
    {
        "label": "EBIT",
        "components": ["EBITDA", "D&A"],
        "insert_after": "D&A",
    },
    {
        "label": "EBT",
        "components": ["EBIT", "Financial result"],
        "insert_after": "Financial result",
    },
    {
        "label": "Net result",
        "components": ["EBT", "Taxes on income", "Other taxes"],
        "insert_after": "Other taxes",
    },
]

KPI_TITLE_TEXT = "KPIs in % total output"
KPI_LABEL_RENAME = {
    "Gross profit": "Gross margin",
    "Cost of goods sold": "Product margin",
    "EBITDA": "EBITDA margin",
    "EBIT": "EBIT margin",
}
KPI_BOLD_LABELS = {"Gross margin", "EBITDA margin", "EBIT margin"}

MAP_START_COL = 5
MAP_END_COL = 8

PROJECT_TITLE_ROW = 1
SUBTITLE_ROW = 2
TITLE_ROW = 6
HEADER_ROW_7 = 7
HEADER_ROW = 8
DATA_START_ROW = 9
RECON_BLOCK_TITLE_ROW = 7
RECON_HEADER_ROW = 8

# ==================================================
# STYLES (gst_excel_theme / PL recon work)
# ==================================================
FONT_BASE = THEME.font_base
FONT_BASE_BOLD = THEME.font_bold
FONT_HEADER = THEME.font_header
FONT_MAPPING = Font(name=THEME.font_name, size=THEME.font_size, color=THEME.text_header)
FONT_TITLE = Font(
    name=THEME.font_name,
    size=THEME.font_size,
    color=THEME.text_brand_title,
    bold=True,
)
FONT_PROJECT = Font(
    name=THEME.font_name,
    size=THEME.font_size_title,
    color=THEME.text_brand_title,
    bold=False,
)
FONT_SUBTITLE = Font(
    name=THEME.font_name,
    size=THEME.font_size_subtitle,
    color=THEME.text_brand_title,
    bold=False,
)

FONT_KPI_TITLE = Font(
    name=THEME.font_name, size=THEME.font_size, color="00A7B5", italic=True, bold=True,
)
FONT_KPI_LABEL = Font(
    name=THEME.font_name, size=THEME.font_size, color=THEME.text_kpi, italic=True, bold=False,
)
FONT_KPI_LABEL_BOLD = Font(
    name=THEME.font_name, size=THEME.font_size, color=THEME.text_kpi, italic=True, bold=True,
)
FONT_CAGR = Font(name=THEME.font_name, size=THEME.font_size, color=THEME.text_cagr, bold=False)
FONT_CAGR_BOLD = Font(name=THEME.font_name, size=THEME.font_size, color=THEME.text_cagr, bold=True)

FILL_GREY = THEME.fill_tech
FILL_WHITE = THEME.fill_white
FILL_HEADER = THEME.fill_header
FILL_HEADER7 = THEME.fill_header
FILL_KPI = THEME.fill_subtotal
FILL_YELLOW = PatternFill("solid", fgColor="FFFF00")

ALIGN_LEFT = Alignment(horizontal="left", vertical="center")
ALIGN_CENTER = Alignment(horizontal="center", vertical="center")
ALIGN_RIGHT = Alignment(horizontal="right", vertical="center")

THIN = Side(style="thin", color=THEME.border_color)
TOP_BORDER = Border(top=THIN)
TOP_BOTTOM_BORDER = Border(top=THIN, bottom=THIN)
BOTTOM_BORDER = Border(bottom=THIN)
NO_BORDER = Border()

NUM_FMT_INT = "#,##0;(#,##0);-"
NUM_FMT_RATE = "0.0;(0.0);-"

# ==================================================
# HELPERS
# ==================================================
def col_letter(i: int) -> str:
    return get_column_letter(i)


def detect_year_columns(df_: pd.DataFrame) -> list[str]:
    """FY columns only (for recon cross-checks)."""
    fy, _ = split_fy_and_ytd(ordered_reporting_columns_from_df(df_))
    if not fy:
        raise ValueError("Keine FY-Spalten gefunden (erwartet z. B. FY22A, FY23A).")
    return fy


def detect_reporting_period_columns(df_: pd.DataFrame) -> list[str]:
    periods = ordered_reporting_columns_from_df(df_)
    if not periods:
        raise ValueError("Keine FY/YTD-Spalten gefunden (erwartet z. B. FY22A, YTD25A).")
    return periods


def resolve_source_section_column(df_: pd.DataFrame) -> str:
    markers = {REPORTED_FILTER_VALUE, "adjusted", "consolidation"}

    def marker_count(col: str) -> int:
        if col not in df_.columns:
            return 0
        series = df_[col].astype(str).str.strip().str.lower()
        return int(series.isin(markers).sum())

    best_col = None
    best_count = 0
    for col in SOURCE_COL_CANDIDATES:
        count = marker_count(col)
        if count > best_count:
            best_col = col
            best_count = count
    if best_col and best_count > 0:
        return best_col
    for col in SOURCE_COL_CANDIDATES:
        if col in df_.columns:
            return col
    return SOURCE_COL_CANDIDATES[0]


def compute_lead_is_layout(n_years: int):
    pos1 = 10
    y1 = list(range(pos1 + 1, pos1 + 1 + n_years))
    cagr1 = y1[-1] + 1
    spacer = cagr1 + 1
    pos2 = spacer + 1
    y2 = list(range(pos2 + 1, pos2 + 1 + n_years))
    cagr2 = y2[-1] + 1
    return pos1, y1, cagr1, spacer, pos2, y2, cagr2


def insert_after_anchor(struct, anchor_label, new_row):
    for i, row in enumerate(struct):
        if row["label"] == anchor_label:
            struct.insert(i + 1, new_row)
            return True
    return False


def normalize_value(x, mapping):
    if x is None or (isinstance(x, float) and pd.isna(x)):
        return x
    s = str(x).strip()
    return mapping.get(s, s)


def norm_pl(x):
    return normalize_value(x, PL_POSITION_RENAME_MAP)


def year_short(y):
    return y.replace("FY", "")


def cagr_formula(first_col, last_col, excel_row, cagr_n: int) -> str:
    if cagr_n <= 0:
        return "n/a"
    return (
        f'=IFERROR((({col_letter(last_col)}{excel_row}/'
        f'{col_letter(first_col)}{excel_row})^(1/{cagr_n})-1)*100,"n/a")'
    )


def find_recon_consolidation_year_cols(ws_recon, years: list[str]) -> dict[str, int]:
    recon_con_start = None
    for cc in range(1, ws_recon.max_column + 1):
        v = ws_recon.cell(RECON_BLOCK_TITLE_ROW, cc).value
        if isinstance(v, str) and v.strip() == "Consolidation":
            recon_con_start = cc
            break
    if recon_con_start is None:
        raise RuntimeError("Could not find 'Consolidation' block in PL_Reconciliation (row 7).")

    recon_year_col: dict[str, int] = {}
    for cc in range(recon_con_start, ws_recon.max_column + 1):
        v = ws_recon.cell(RECON_HEADER_ROW, cc).value
        if isinstance(v, str):
            key = v.strip()
            if key in years:
                recon_year_col[key] = cc
        if len(recon_year_col) >= len(years):
            break

    missing = [y for y in years if y not in recon_year_col]
    if missing:
        raise RuntimeError(f"Consolidation block missing year headers: {missing}")
    return recon_year_col


TOP_BORDER_ONLY_N = {norm_pl(x) for x in TOP_BORDER_ONLY}
TOP_AND_BOTTOM_N = {norm_pl(x) for x in TOP_AND_BOTTOM}

PL_TOTALS_CONFIG = [
    {
        "label": norm_pl(t["label"]),
        "components": [norm_pl(c) for c in t["components"]],
        "insert_after": norm_pl(t["insert_after"]),
    }
    for t in PL_TOTALS_CONFIG_RAW
]

# ==================================================
# MAIN
# ==================================================
print("Lead_IS_work — Desktop defaults:")
print(f"  input:   {INPUT_FILE}")
print(f"  mapping: {MAPPING_FILE}")

wb = load_workbook(INPUT_FILE)
if SHEET_MASTER not in wb.sheetnames:
    raise RuntimeError(f"'{SHEET_MASTER}' nicht gefunden in {INPUT_FILE}.")

if SHEET_OUT in wb.sheetnames:
    del wb[SHEET_OUT]
ws = wb.create_sheet(SHEET_OUT)

df_master = pd.read_excel(INPUT_FILE, sheet_name=SHEET_MASTER, engine="openpyxl")
PERIODS = detect_reporting_period_columns(df_master)
FY_COLS, YTD_COLS = split_fy_and_ytd(PERIODS)
CAGR_N = len(FY_COLS) - 1

required_cols = {"L3", "L4"} | set(PERIODS)
missing = required_cols - set(df_master.columns)
if missing:
    raise ValueError(f"Master_PL fehlt Spalten: {missing}")

l5_col = resolve_source_section_column(df_master)
POS1_COL, Y1_COLS, CAGR1_COL, SPACER_COL, POS2_COL, Y2_COLS, CAGR2_COL = compute_lead_is_layout(len(PERIODS))

year_range_txt = f"{FY_COLS[0]} - {FY_COLS[-1]}"
TITLE_REPORTED = f"{GROUP_NAME} | Reported income statement {year_range_txt}"
TITLE_PROFORMA = f"{GROUP_NAME} | Pro forma income statement {year_range_txt}"

MASTER_START = 2
MASTER_END = len(df_master) + 1


def master_range(colname):
    L = get_column_letter(df_master.columns.get_loc(colname) + 1)
    return f"{SHEET_MASTER}!${L}${MASTER_START}:${L}${MASTER_END}"


l3_rng = master_range("L3")
l4_rng = master_range("L4")
l5_rng = master_range(l5_col)
year_rng = {y: master_range(y) for y in PERIODS}

map_df = pd.read_excel(MAPPING_FILE, sheet_name=0, engine="openpyxl")
if "L3" not in map_df.columns:
    raise ValueError("Mapping-Datei muss Spalte 'L3' enthalten.")

map_df = map_df.copy()
map_df["L3"] = map_df["L3"].astype(str).apply(norm_pl)

l3_order = l3_order_from_mapping(map_df)
row_structure = build_lead_is_row_structure(
    df_master,
    l3_order,
    {"l4_sort_basis": L4_SORT_BASIS},
    source_col=l5_col,
    label_fn=norm_pl,
)

existing_labels = {r["label"] for r in row_structure}

for t in PL_TOTALS_CONFIG:
    comps = [c for c in t["components"] if c in existing_labels]
    missing_comps = [c for c in t["components"] if c not in existing_labels]
    if missing_comps:
        print(f"Warnung: Total '{t['label']}' — fehlende Komponenten: {missing_comps}")
    new_row = {
        "type": "total",
        "label": t["label"],
        "components": comps,
        "L3": t["label"],
        "L4": "",
    }
    if not insert_after_anchor(row_structure, t["insert_after"], new_row):
        print(
            f"Warnung: Anchor '{t['insert_after']}' für Total '{t['label']}' "
            "nicht gefunden — Total übersprungen."
        )
        continue
    existing_labels.add(t["label"])

# Titles
pt = ws.cell(PROJECT_TITLE_ROW, POS1_COL, f"Project {PROJECT_NAME}")
pt.font = FONT_PROJECT
pt.alignment = ALIGN_LEFT
ws.row_dimensions[PROJECT_TITLE_ROW].height = 36

st = ws.cell(SUBTITLE_ROW, POS1_COL, f"{GROUP_NAME} Lead income statement")
st.font = FONT_SUBTITLE
st.alignment = ALIGN_LEFT

ws.merge_cells(start_row=TITLE_ROW, start_column=POS1_COL, end_row=TITLE_ROW, end_column=CAGR1_COL)
t1 = ws.cell(TITLE_ROW, POS1_COL, TITLE_REPORTED)
t1.font = FONT_TITLE
t1.alignment = ALIGN_LEFT

ws.merge_cells(start_row=TITLE_ROW, start_column=POS2_COL, end_row=TITLE_ROW, end_column=CAGR2_COL)
t2 = ws.cell(TITLE_ROW, POS2_COL, TITLE_PROFORMA)
t2.font = FONT_TITLE
t2.alignment = ALIGN_LEFT

for cc in range(POS1_COL, CAGR1_COL + 1):
    if cc != SPACER_COL:
        ws.cell(HEADER_ROW_7, cc).fill = FILL_HEADER7
        ws.cell(HEADER_ROW_7, cc).font = FONT_BASE_BOLD
for cc in range(POS2_COL, CAGR2_COL + 1):
    ws.cell(HEADER_ROW_7, cc).fill = FILL_HEADER7
    ws.cell(HEADER_ROW_7, cc).font = FONT_BASE_BOLD

cagr_range_txt = f"{year_short(FY_COLS[0])} - {year_short(FY_COLS[-1])}"
for cagr_col in (CAGR1_COL, CAGR2_COL):
    ws.cell(HEADER_ROW_7, cagr_col, "CAGR").alignment = ALIGN_RIGHT
    ws.cell(HEADER_ROW_7, cagr_col).font = FONT_BASE_BOLD
    ws.cell(HEADER_ROW, cagr_col, cagr_range_txt).alignment = ALIGN_RIGHT
    ws.cell(HEADER_ROW, cagr_col).font = FONT_BASE_BOLD
    ws.cell(HEADER_ROW, cagr_col).fill = FILL_HEADER

ws.column_dimensions[col_letter(CAGR1_COL)].width = 12
ws.column_dimensions[col_letter(CAGR2_COL)].width = 12

ws.cell(HEADER_ROW, MAP_START_COL + 0, "Reported").font = FONT_HEADER
ws.cell(HEADER_ROW, MAP_START_COL + 1, "").font = FONT_HEADER
ws.cell(HEADER_ROW, MAP_START_COL + 2, "L3").font = FONT_HEADER
ws.cell(HEADER_ROW, MAP_START_COL + 3, "L4").font = FONT_HEADER
for c in range(MAP_START_COL, MAP_END_COL + 1):
    ws.cell(HEADER_ROW, c).alignment = ALIGN_LEFT
    ws.cell(HEADER_ROW, c).fill = FILL_HEADER

ws.cell(HEADER_ROW, POS1_COL, UNIT_LABEL).font = FONT_HEADER
ws.cell(HEADER_ROW, POS1_COL).alignment = ALIGN_LEFT
ws.cell(HEADER_ROW, POS1_COL).fill = FILL_HEADER

for idx, y in enumerate(PERIODS):
    cc = Y1_COLS[idx]
    h = ws.cell(HEADER_ROW, cc, y)
    h.font = FONT_HEADER
    h.alignment = ALIGN_RIGHT
    h.fill = FILL_HEADER

ws.column_dimensions[col_letter(SPACER_COL)].width = LEAD_SPACER_WIDTH
ws.cell(HEADER_ROW, SPACER_COL).fill = FILL_WHITE
ws.cell(HEADER_ROW_7, SPACER_COL).fill = FILL_WHITE

ws.cell(HEADER_ROW, POS2_COL, UNIT_LABEL).font = FONT_HEADER
ws.cell(HEADER_ROW, POS2_COL).alignment = ALIGN_LEFT
ws.cell(HEADER_ROW, POS2_COL).fill = FILL_HEADER

for idx, y in enumerate(PERIODS):
    cc = Y2_COLS[idx]
    h = ws.cell(HEADER_ROW, cc, y)
    h.font = FONT_HEADER
    h.alignment = ALIGN_RIGHT
    h.fill = FILL_HEADER

ws.column_dimensions[col_letter(POS1_COL)].width = 32
ws.column_dimensions[col_letter(POS2_COL)].width = 32
for c in Y1_COLS + Y2_COLS:
    ws.column_dimensions[col_letter(c)].width = 7.86
for c in range(MAP_START_COL, MAP_END_COL + 1):
    ws.column_dimensions[col_letter(c)].width = 12

for cc in range(MAP_START_COL, MAP_END_COL + 1):
    ws.cell(HEADER_ROW, cc).border = BOTTOM_BORDER
for cc in range(POS1_COL, CAGR1_COL + 1):
    if cc != SPACER_COL:
        ws.cell(HEADER_ROW, cc).border = BOTTOM_BORDER
for cc in range(POS2_COL, CAGR2_COL + 1):
    ws.cell(HEADER_ROW, cc).border = BOTTOM_BORDER

row_index = {}
for i, r in enumerate(row_structure):
    excel_row = DATA_START_ROW + i
    r["_idx"] = i
    r["_excel_row"] = excel_row
    row_index[r["label"]] = excel_row

TOTAL_OUTPUT_LABEL = norm_pl("Total output")
if TOTAL_OUTPUT_LABEL not in row_index:
    fallback = norm_pl("Net sales")
    if fallback in row_index:
        print(f"Warnung: '{TOTAL_OUTPUT_LABEL}' fehlt — KPI-Basis: '{fallback}'.")
        TOTAL_OUTPUT_LABEL = fallback
    else:
        raise RuntimeError("Total output row not found in row_structure.")
TOTAL_OUTPUT_ROW = row_index[TOTAL_OUTPUT_LABEL]


def subtotal_span(idx, l3_label):
    k = idx - 1
    while k >= 0 and row_structure[k]["type"] == "detail" and row_structure[k]["L3"] == l3_label:
        k -= 1
    start = DATA_START_ROW + (k + 1)
    end = DATA_START_ROW + (idx - 1)
    return start, end


for i, r in enumerate(row_structure):
    excel_row = r["_excel_row"]

    ws.cell(excel_row, MAP_START_COL + 0, "Reported").font = FONT_MAPPING
    ws.cell(excel_row, MAP_START_COL + 1, "").font = FONT_MAPPING
    ws.cell(excel_row, MAP_START_COL + 2, r["L3"]).font = FONT_MAPPING
    ws.cell(excel_row, MAP_START_COL + 3, r["L4"]).font = FONT_MAPPING
    for c in range(MAP_START_COL, MAP_END_COL + 1):
        ws.cell(excel_row, c).alignment = ALIGN_LEFT

    is_real_subtotal = (
        r["type"] == "subtotal"
        and any(rr["type"] == "detail" and rr["L3"] == r["label"] for rr in row_structure)
    )
    is_total = r["type"] == "total"
    is_bold_main = is_total or is_real_subtotal

    p1 = ws.cell(excel_row, POS1_COL, r["label"])
    p1.alignment = ALIGN_LEFT
    p1.font = FONT_BASE_BOLD if is_bold_main else FONT_BASE

    p2 = ws.cell(excel_row, POS2_COL)
    p2.value = f"={col_letter(POS1_COL)}{excel_row}"
    p2.alignment = ALIGN_LEFT
    p2.font = FONT_BASE_BOLD if is_bold_main else FONT_BASE

    l3_crit = f"${col_letter(MAP_START_COL + 2)}${excel_row}"
    l4_crit = f"${col_letter(MAP_START_COL + 3)}${excel_row}"
    rep_crit = f"${col_letter(MAP_START_COL + 0)}${excel_row}"

    for y_idx, year in enumerate(PERIODS):
        c = ws.cell(excel_row, Y1_COLS[y_idx])
        c.number_format = NUM_FMT_INT
        c.alignment = ALIGN_RIGHT
        c.font = FONT_BASE_BOLD if is_bold_main else FONT_BASE

        if r["type"] == "detail":
            c.value = (
                f"=SUMIFS({year_rng[year]},"
                f"{l3_rng},{l3_crit},"
                f"{l4_rng},{l4_crit},"
                f"{l5_rng},{rep_crit}"
                f")/1000"
            )
        elif r["type"] == "subtotal":
            start, end = subtotal_span(i, r["label"])
            colL = col_letter(Y1_COLS[y_idx])
            if end >= start:
                c.value = f"=SUM({colL}{start}:{colL}{end})"
            else:
                c.value = (
                    f"=SUMIFS({year_rng[year]},"
                    f"{l3_rng},{l3_crit},"
                    f"{l4_rng},{l4_crit}"
                    f")/1000"
                )
        else:
            refs = []
            for comp in r.get("components", []):
                if comp in row_index:
                    refs.append(f"{col_letter(Y1_COLS[y_idx])}{row_index[comp]}")
            c.value = f"=SUM({','.join(refs)})" if refs else 0

    cagr1 = ws.cell(excel_row, CAGR1_COL)
    cagr1.alignment = ALIGN_RIGHT
    cagr1.number_format = NUM_FMT_RATE
    cagr1.font = FONT_CAGR_BOLD if is_total else FONT_CAGR
    cagr1.fill = FILL_KPI
    cagr1.border = NO_BORDER
    cagr1.value = cagr_formula(Y1_COLS[0], Y1_COLS[len(FY_COLS) - 1], excel_row, CAGR_N)

    for y_idx, year in enumerate(PERIODS):
        c = ws.cell(excel_row, Y2_COLS[y_idx])
        c.number_format = NUM_FMT_INT
        c.alignment = ALIGN_RIGHT
        c.font = FONT_BASE_BOLD if is_bold_main else FONT_BASE

        if r["type"] == "detail":
            c.value = (
                f"=SUMIFS({year_rng[year]},"
                f"{l3_rng},{l3_crit},"
                f"{l4_rng},{l4_crit}"
                f")/1000"
            )
        elif r["type"] == "subtotal":
            start, end = subtotal_span(i, r["label"])
            colL = col_letter(Y2_COLS[y_idx])
            if end >= start:
                c.value = f"=SUM({colL}{start}:{colL}{end})"
            else:
                c.value = (
                    f"=SUMIFS({year_rng[year]},"
                    f"{l3_rng},{l3_crit},"
                    f"{l4_rng},{l4_crit}"
                    f")/1000"
                )
        else:
            refs = []
            for comp in r.get("components", []):
                if comp in row_index:
                    refs.append(f"{col_letter(Y2_COLS[y_idx])}{row_index[comp]}")
            c.value = f"=SUM({','.join(refs)})" if refs else 0

    cagr2 = ws.cell(excel_row, CAGR2_COL)
    cagr2.alignment = ALIGN_RIGHT
    cagr2.number_format = NUM_FMT_RATE
    cagr2.font = FONT_CAGR_BOLD if is_total else FONT_CAGR
    cagr2.fill = FILL_KPI
    cagr2.border = NO_BORDER
    cagr2.value = cagr_formula(Y2_COLS[0], Y2_COLS[len(FY_COLS) - 1], excel_row, CAGR_N)

for r in row_structure:
    if r["type"] == "detail":
        rr = r["_excel_row"]
        ws.row_dimensions[rr].outlineLevel = 1
        ws.row_dimensions[rr].hidden = False

for r in row_structure:
    label = r["label"]
    if label not in TOP_BORDER_ONLY_N and label not in TOP_AND_BOTTOM_N:
        continue
    excel_row = r["_excel_row"]
    border_to_set = TOP_BOTTOM_BORDER if label in TOP_AND_BOTTOM_N else TOP_BORDER
    ws.cell(excel_row, POS1_COL).border = border_to_set
    for c in Y1_COLS:
        ws.cell(excel_row, c).border = border_to_set
    ws.cell(excel_row, POS2_COL).border = border_to_set
    for c in Y2_COLS:
        ws.cell(excel_row, c).border = border_to_set

LAST_TABLE_ROW = DATA_START_ROW + len(row_structure) - 1
KPI_TITLE_ROW = LAST_TABLE_ROW + 1
KPI_FIRST_ROW = KPI_TITLE_ROW + 1

for cc in range(POS1_COL, CAGR2_COL + 1):
    if cc == SPACER_COL:
        ws.cell(KPI_TITLE_ROW, cc).fill = FILL_WHITE
    else:
        ws.cell(KPI_TITLE_ROW, cc).fill = FILL_KPI
    ws.cell(KPI_TITLE_ROW, cc).border = NO_BORDER

for col in (POS1_COL, POS2_COL):
    cell = ws.cell(KPI_TITLE_ROW, col, KPI_TITLE_TEXT)
    cell.font = FONT_KPI_TITLE
    cell.alignment = ALIGN_LEFT

total_output_idx = next(
    (i for i, rr in enumerate(row_structure) if rr["label"] == TOTAL_OUTPUT_LABEL),
    0,
)
end_label = norm_pl("EBIT")
kpi_candidates = []
for rr in row_structure[total_output_idx + 1:]:
    kpi_candidates.append(rr)
    if rr["label"] == end_label:
        break

kpi_lines = [x for x in kpi_candidates if x["type"] in {"subtotal", "total"}]
kpi_row = KPI_FIRST_ROW
for base in kpi_lines:
    base_label = base["label"]
    base_row = row_index.get(base_label)
    if base_row is None:
        continue
    kpi_label = KPI_LABEL_RENAME.get(base_label, base_label)
    label_is_bold = kpi_label in KPI_BOLD_LABELS

    for cc in range(POS1_COL, CAGR2_COL + 1):
        if cc == SPACER_COL:
            ws.cell(kpi_row, cc).fill = FILL_WHITE
        else:
            ws.cell(kpi_row, cc).fill = FILL_KPI
        ws.cell(kpi_row, cc).border = NO_BORDER

    for col in (POS1_COL, POS2_COL):
        c = ws.cell(kpi_row, col, kpi_label)
        c.font = FONT_KPI_LABEL_BOLD if label_is_bold else FONT_KPI_LABEL
        c.alignment = ALIGN_LEFT

    for col in Y1_COLS:
        v = ws.cell(kpi_row, col)
        v.value = f'=IFERROR({col_letter(col)}{base_row}/{col_letter(col)}{TOTAL_OUTPUT_ROW}*100,"n/a")'
        v.number_format = NUM_FMT_RATE
        v.alignment = ALIGN_RIGHT
        v.font = FONT_KPI_LABEL_BOLD if label_is_bold else FONT_KPI_LABEL

    v = ws.cell(kpi_row, CAGR1_COL)
    v.value = None
    v.number_format = NUM_FMT_RATE
    v.alignment = ALIGN_RIGHT
    v.font = FONT_KPI_LABEL_BOLD if label_is_bold else FONT_KPI_LABEL

    for col in Y2_COLS:
        v = ws.cell(kpi_row, col)
        v.value = f'=IFERROR({col_letter(col)}{base_row}/{col_letter(col)}{TOTAL_OUTPUT_ROW}*100,"n/a")'
        v.number_format = NUM_FMT_RATE
        v.alignment = ALIGN_RIGHT
        v.font = FONT_KPI_LABEL_BOLD if label_is_bold else FONT_KPI_LABEL

    v = ws.cell(kpi_row, CAGR2_COL)
    v.value = None
    v.number_format = NUM_FMT_RATE
    v.alignment = ALIGN_RIGHT
    v.font = FONT_KPI_LABEL_BOLD if label_is_bold else FONT_KPI_LABEL

    kpi_row += 1
    if base_label == end_label:
        break

KPI_LAST_ROW = kpi_row - 1

CHECK_SRC_ROW = KPI_LAST_ROW + 2
CHECK_DELTA_ROW = KPI_LAST_ROW + 3
NET_RESULT_LABEL = norm_pl("Net result")
NET_RESULT_ROW = row_index[NET_RESULT_LABEL]

if SHEET_RECON not in wb.sheetnames:
    raise RuntimeError("Sheet 'PL_Reconciliation' not found for check.")
ws_recon = wb[SHEET_RECON]
recon_year_col = find_recon_consolidation_year_cols(ws_recon, FY_COLS)

recon_net_row = None
for rr in range(1, ws_recon.max_row + 1):
    v = ws_recon.cell(rr, POS1_COL).value
    if isinstance(v, str) and v.strip() == NET_RESULT_LABEL:
        recon_net_row = rr
        break
if recon_net_row is None:
    raise RuntimeError("Could not find 'Net result' row in PL_Reconciliation (col J).")

ws.cell(CHECK_SRC_ROW, POS1_COL, "Source - PL_Reconciliation").font = FONT_BASE
ws.cell(CHECK_SRC_ROW, POS1_COL).alignment = ALIGN_LEFT
ws.cell(CHECK_SRC_ROW, POS2_COL).value = f"={col_letter(POS1_COL)}{CHECK_SRC_ROW}"
ws.cell(CHECK_SRC_ROW, POS2_COL).font = FONT_BASE
ws.cell(CHECK_DELTA_ROW, POS1_COL, "Check").font = FONT_BASE
ws.cell(CHECK_DELTA_ROW, POS1_COL).alignment = ALIGN_LEFT
ws.cell(CHECK_DELTA_ROW, POS2_COL).value = f"={col_letter(POS1_COL)}{CHECK_DELTA_ROW}"
ws.cell(CHECK_DELTA_ROW, POS2_COL).font = FONT_BASE

DELTA_RED_FONT = Font(name=THEME.font_name, size=THEME.font_size, color=THEME.delta_negative)

for i, y in enumerate(FY_COLS):
    src1 = ws.cell(CHECK_SRC_ROW, Y1_COLS[i])
    src1.value = f"={SHEET_RECON}!{col_letter(recon_year_col[y])}{recon_net_row}"
    src1.number_format = NUM_FMT_INT
    src1.alignment = ALIGN_RIGHT
    src1.font = FONT_BASE

    src2 = ws.cell(CHECK_SRC_ROW, Y2_COLS[i])
    src2.value = f"={SHEET_RECON}!{col_letter(recon_year_col[y])}{recon_net_row}"
    src2.number_format = NUM_FMT_INT
    src2.alignment = ALIGN_RIGHT
    src2.font = FONT_BASE

    d1 = ws.cell(CHECK_DELTA_ROW, Y1_COLS[i])
    d1.value = f"={col_letter(Y1_COLS[i])}{NET_RESULT_ROW}-{col_letter(Y1_COLS[i])}{CHECK_SRC_ROW}"
    d1.number_format = NUM_FMT_INT
    d1.alignment = ALIGN_RIGHT
    d1.font = DELTA_RED_FONT

    d2 = ws.cell(CHECK_DELTA_ROW, Y2_COLS[i])
    d2.value = f"={col_letter(Y2_COLS[i])}{NET_RESULT_ROW}-{col_letter(Y2_COLS[i])}{CHECK_SRC_ROW}"
    d2.number_format = NUM_FMT_INT
    d2.alignment = ALIGN_RIGHT
    d2.font = DELTA_RED_FONT

for rr in (CHECK_SRC_ROW, CHECK_DELTA_ROW):
    for ccol in (CAGR1_COL, CAGR2_COL):
        ws.cell(rr, ccol).value = None
        ws.cell(rr, ccol).fill = FILL_WHITE
        ws.cell(rr, ccol).border = NO_BORDER
        ws.cell(rr, ccol).font = FONT_BASE

MAX_USED_ROW = CHECK_DELTA_ROW + 25
for rr in range(2, MAX_USED_ROW + 1):
    ws.row_dimensions[rr].height = 12

FILL_END_COL = CAGR2_COL + 10
for rr in range(1, MAX_USED_ROW + 1):
    for cc in range(1, POS1_COL):
        ws.cell(rr, cc).fill = FILL_GREY
    for cc in range(POS1_COL, FILL_END_COL + 1):
        ws.cell(rr, cc).fill = FILL_WHITE

for cc in range(POS1_COL, CAGR1_COL + 1):
    if cc != SPACER_COL:
        ws.cell(HEADER_ROW_7, cc).fill = FILL_HEADER7
        ws.cell(HEADER_ROW, cc).fill = FILL_HEADER
for cc in range(POS2_COL, CAGR2_COL + 1):
    ws.cell(HEADER_ROW_7, cc).fill = FILL_HEADER7
    ws.cell(HEADER_ROW, cc).fill = FILL_HEADER
for cc in range(MAP_START_COL, MAP_END_COL + 1):
    ws.cell(HEADER_ROW, cc).fill = FILL_HEADER

for rr in range(KPI_TITLE_ROW, KPI_LAST_ROW + 1):
    for cc in range(POS1_COL, CAGR2_COL + 1):
        if cc == SPACER_COL or cc in (CAGR1_COL, CAGR2_COL):
            ws.cell(rr, cc).fill = FILL_WHITE
            ws.cell(rr, cc).border = NO_BORDER
        else:
            ws.cell(rr, cc).fill = FILL_KPI
            ws.cell(rr, cc).border = NO_BORDER

for rr in range(DATA_START_ROW, LAST_TABLE_ROW + 1):
    for ccol in (CAGR1_COL, CAGR2_COL):
        ws.cell(rr, ccol).fill = FILL_KPI
        ws.cell(rr, ccol).border = NO_BORDER

for rr in range(1, MAX_USED_ROW + 1):
    ws.cell(rr, SPACER_COL).fill = FILL_WHITE

for ccol in (CAGR1_COL, CAGR2_COL):
    ws.cell(NET_RESULT_ROW, ccol).fill = FILL_KPI
    ws.cell(NET_RESULT_ROW, ccol).border = NO_BORDER
    ws.cell(NET_RESULT_ROW, ccol).font = FONT_CAGR_BOLD

for rr in range(NET_RESULT_ROW + 1, LAST_TABLE_ROW + 1):
    for ccol in (CAGR1_COL, CAGR2_COL):
        ws.cell(rr, ccol).fill = FILL_WHITE
        ws.cell(rr, ccol).border = NO_BORDER
        ws.cell(rr, ccol).font = FONT_CAGR_BOLD

for rr in range(KPI_TITLE_ROW, KPI_LAST_ROW + 1):
    for ccol in (CAGR1_COL, CAGR2_COL):
        ws.cell(rr, ccol).fill = FILL_WHITE
        ws.cell(rr, ccol).border = NO_BORDER

table_cols_source_yellow = [POS1_COL] + Y1_COLS + [POS2_COL] + Y2_COLS
for cc in table_cols_source_yellow:
    ws.cell(CHECK_SRC_ROW, cc).fill = FILL_YELLOW

for rr in (CHECK_SRC_ROW, CHECK_DELTA_ROW):
    ws.row_dimensions[rr].outlineLevel = 2
    ws.row_dimensions[rr].hidden = True

ws.column_dimensions.group("A", "I", hidden=True)
ws.column_dimensions["J"].collapsed = True
ws.sheet_properties.outlinePr.summaryBelow = True
ws.sheet_properties.outlinePr.summaryRight = True
ws.sheet_view.showOutlineSymbols = True

apply_zero_row_conditional_formatting(
    ws,
    first_row=DATA_START_ROW,
    last_row=LAST_TABLE_ROW,
    year_col_indices=Y1_COLS + Y2_COLS,
    style_start_col=POS1_COL,
    style_end_col=CAGR2_COL,
    exclude_cols={SPACER_COL},
)

wb.save(INPUT_FILE)
print(f"Fertig. Reiter '{SHEET_OUT}' gespeichert in: {INPUT_FILE}")
print(f"Master_PL Source-Spalte: {l5_col}")
print(f"Verwendete Perioden: {PERIODS} (FY={FY_COLS}, CAGR_N={CAGR_N})")
print("Check rows linked to PL_Reconciliation / Consolidation / Net result.")
