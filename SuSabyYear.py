"""
Build a master trial-balance pivot from Summen- und Saldenlisten (SuSa).

Driven by JSON config (first CLI argument) as produced by the FDD backend.
Column roles come from user-confirmed column_mapping (no heuristic detection in production).
"""

from __future__ import annotations

import json
import os
import re
import sys
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import pandas as pd

from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter
from openpyxl.utils.dataframe import dataframe_to_rows

from susa_column_mapping import (
    EN_MONTH_ABBR,
    apply_sign_indicator,
    build_period_blocks_from_profile,
    excel_col_to_idx,
    expand_period_blocks,
    find_header_row,
    filter_ap_ar_by_length,
    idx_to_excel_col,
    looks_like_account,
    parse_month_year_from_label,
    parse_number_de,
    profile_for_entity,
    resolve_header_row,
    signed_amount,
    col_series,
)

# GST-aligned Excel theme (shared with General Sales Table)
_SCRIPTS = Path(__file__).resolve().parent / "scripts"
if str(_SCRIPTS) not in sys.path:
    sys.path.insert(0, str(_SCRIPTS))
from gst_excel_theme import THEME  # noqa: E402

# ==============================================================
# Helpers
# ==============================================================


def parse_month_from_sheetname(sheet: str) -> int:
    s = sheet.lower()
    s = s.replace("ä", "ae").replace("ö", "oe").replace("ü", "ue").replace("ß", "ss")

    month_map = {
        1:  ["jan", "januar", "january"],
        2:  ["feb", "februar", "february"],
        3:  ["mar", "maerz", "mrz", "march"],
        4:  ["apr", "april"],
        5:  ["mai", "may"],
        6:  ["jun", "juni", "june"],
        7:  ["jul", "juli", "july"],
        8:  ["aug", "august"],
        9:  ["sep", "sept", "september"],
        10: ["okt", "oct", "oktober", "october"],
        11: ["nov", "november"],
        12: ["dez", "dec", "dezember", "december"],
    }

    for mm, keys in month_map.items():
        if any(k in s for k in keys):
            return mm

    m = re.search(r"\b(0?[1-9]|1[0-2])\b", s)
    if m:
        return int(m.group(1))

    raise ValueError(f"Month not identifiable from sheet name: {sheet}")


def parse_month_from_filename(filename: str) -> Optional[int]:
    try:
        return parse_month_from_sheetname(Path(filename).stem)
    except ValueError:
        return None


def make_period_str(year: int, month: int) -> str:
    return f"{EN_MONTH_ABBR[month - 1]}-{year}"


def apply_letter_mapping(df: pd.DataFrame, mapping: Dict[str, str]) -> pd.DataFrame:
    out = {}
    ncols = df.shape[1]
    for target, letter in mapping.items():
        if not letter:
            continue
        idx = excel_col_to_idx(letter)
        if idx >= ncols:
            raise ValueError(f"Configured column '{letter}' for '{target}' out of range.")
        out[target] = df.iloc[:, idx]
    return pd.DataFrame(out)


def assign_months_twelve_workbooks(paths: List[str]) -> List[Tuple[str, int]]:
    pairs: List[Tuple[str, int, str]] = []
    for p in paths:
        m = parse_month_from_filename(os.path.basename(p))
        pairs.append((p, m if m is not None else -1, os.path.basename(p).lower()))
    if all(m > 0 for _, m, _ in pairs):
        return sorted([(p, m) for p, m, _ in pairs], key=lambda x: x[1])
    ordered_paths = sorted(paths, key=lambda x: os.path.basename(x).lower())
    return [(p, i + 1) for i, p in enumerate(ordered_paths)]


def _has_column_mapping(config: dict) -> bool:
    cm = config.get("column_mapping")
    if not cm:
        return False
    if isinstance(cm, str):
        try:
            cm = json.loads(cm)
        except json.JSONDecodeError:
            return False
    return bool(cm.get("default"))


def finalize_accounts(df: pd.DataFrame, config: dict) -> pd.DataFrame:
    df["Account"] = df["Account"].astype(str).str.replace(".0", "", regex=False).str.strip()
    df = df[df["Account"].apply(looks_like_account)]
    length = config.get("ap_ar_remove_account_length")
    if length is not None:
        try:
            length = int(length)
        except (TypeError, ValueError):
            length = None
    # Only filter when AP/AR sub-ledgers are present (user answered "yes" and gave digit length).
    if length:
        df = filter_ap_ar_by_length(df, length)
    return df


def scale_balance(df: pd.DataFrame, config: dict) -> pd.DataFrame:
    df = df.copy()
    if config.get("scale_to_keur", True):
        df["Balance"] = pd.to_numeric(df["Balance"], errors="coerce") / 1000
    return df


def read_data_block_legacy(
    raw: pd.DataFrame,
    config: dict,
    year: int,
    month: int,
    entity_name: str,
) -> Optional[pd.DataFrame]:
    header_row = find_header_row(raw)
    df = raw.dropna(how="all").reset_index(drop=True)
    df = df.iloc[header_row + 1:].copy()
    if df.empty:
        return None
    cols_cfg = dict(config.get("columns") or {})
    df = apply_letter_mapping(df, cols_cfg)
    df = finalize_accounts(df, config)
    if df.empty:
        return None
    df["Balance"] = df["Balance"].apply(parse_number_de)
    mode = str(config.get("sign_mode") or "sh_column").strip()
    if mode == "sh_column" and "S" in df.columns:
        df["S"] = df["S"].astype(str).str.upper().str.strip()
        df.loc[df["S"] == "S", "Balance"] *= -1
    if config.get("scale_to_keur", True):
        df["Balance"] = pd.to_numeric(df["Balance"], errors="coerce") / 1000
    df["Year"] = year
    df["Month"] = month
    df["Period"] = make_period_str(year, month)
    df["Entity"] = entity_name
    return df


def _opening_balance_period(year: int) -> Tuple[int, str]:
    return 0, f"EB-{year}"


def read_opening_balance_mapped(
    raw: pd.DataFrame,
    config: dict,
    entity_name: str,
    year: int,
) -> Optional[pd.DataFrame]:
    """Read Eröffnungsbilanz (amount + optional S/H/X indicator column(s))."""
    profile = profile_for_entity(config, entity_name)
    cols = profile.get("columns") or profile
    ob_col = cols.get("opening_balance")
    if not ob_col:
        return None

    header_row = resolve_header_row(raw, profile)

    body = raw.dropna(how="all").reset_index(drop=True).iloc[header_row + 1:].copy()
    if body.empty:
        return None

    acc_col = cols.get("account") or cols.get("Account")
    desc_col = cols.get("account_description") or cols.get("Account description")
    if not acc_col:
        return None

    out = pd.DataFrame()
    out["Account"] = col_series(body, acc_col)
    out["Account description"] = col_series(body, desc_col) if desc_col else ""

    eb_sh = cols.get("eb_sh")
    eb_s = cols.get("eb_s") or cols.get("eb_sh_alt")
    eb_h = cols.get("eb_h")
    balances = []
    for i in range(len(body)):
        row = body.iloc[i]
        amt = row.iloc[excel_col_to_idx(ob_col)]
        if eb_sh:
            ind = row.iloc[excel_col_to_idx(eb_sh)]
            balances.append(apply_sign_indicator(amt, ind))
        elif eb_s or eb_h:
            ind_s = row.iloc[excel_col_to_idx(eb_s)] if eb_s else None
            ind_h = row.iloc[excel_col_to_idx(eb_h)] if eb_h else None
            ind = ind_s if pd.notna(ind_s) and str(ind_s).strip() else ind_h
            balances.append(apply_sign_indicator(amt, ind))
        else:
            balances.append(parse_number_de(amt))
    out["Balance"] = balances
    out = finalize_accounts(out, config)
    if out.empty:
        return None
    out = scale_balance(out, config)
    month, period = _opening_balance_period(year)
    out["Year"] = year
    out["Month"] = month
    out["Period"] = period
    out["Entity"] = entity_name
    return out


def _period_amount_from_row(
    row: pd.Series,
    sign_mode: str,
    block: dict,
) -> float:
    """Period layout: amount column first, indicator column(s) after."""
    amount_letter = block.get("amount") or block.get("amount_col")
    sh_letter = block.get("sh") or block.get("sh_col")
    soll_letter = block.get("soll") or block.get("soll_col")
    haben_letter = block.get("haben") or block.get("haben_col")

    sh_s_letter = block.get("sh_s") or soll_letter
    sh_h_letter = block.get("sh_h") or haben_letter
    if sign_mode == "sh_two_columns":
        amt_raw = row.iloc[excel_col_to_idx(amount_letter)] if amount_letter else None
        if amount_letter and pd.notna(amt_raw):
            ind_s = row.iloc[excel_col_to_idx(sh_s_letter)] if sh_s_letter else None
            ind_h = row.iloc[excel_col_to_idx(sh_h_letter)] if sh_h_letter else None
            ind = ind_s if pd.notna(ind_s) and str(ind_s).strip() else ind_h
            return apply_sign_indicator(amt_raw, ind)
        soll_raw = row.iloc[excel_col_to_idx(sh_s_letter)] if sh_s_letter else None
        haben_raw = row.iloc[excel_col_to_idx(sh_h_letter)] if sh_h_letter else None
        return signed_amount(None, sign_mode, soll_raw=soll_raw, haben_raw=haben_raw)

    if sign_mode == "sh_column" and sh_letter and amount_letter:
        return apply_sign_indicator(
            row.iloc[excel_col_to_idx(amount_letter)],
            row.iloc[excel_col_to_idx(sh_letter)],
        )
    return signed_amount(
        row.iloc[excel_col_to_idx(amount_letter)] if amount_letter else None,
        sign_mode,
    )


def read_period_row_mapped(
    raw: pd.DataFrame,
    config: dict,
    entity_name: str,
    year: int,
    month: int,
    period_block: Optional[dict] = None,
) -> Optional[pd.DataFrame]:
    profile = profile_for_entity(config, entity_name)
    header_row = resolve_header_row(raw, profile)

    df = raw.dropna(how="all").reset_index(drop=True)
    body = df.iloc[header_row + 1:].copy()
    if body.empty:
        return None

    cols = profile.get("columns") or profile
    acc_col = cols.get("account") or cols.get("Account")
    desc_col = cols.get("account_description") or cols.get("Account description")
    if not acc_col:
        raise ValueError("column_mapping must define account column letter.")

    out = pd.DataFrame()
    out["Account"] = col_series(body, acc_col)
    out["Account description"] = col_series(body, desc_col) if desc_col else ""

    sign_mode = str(config.get("sign_mode") or "sh_column").strip()

    if period_block:
        block = period_block
    else:
        blocks = expand_period_blocks(profile, body.shape[1], sign_mode, year)
        block = blocks[0] if blocks else {}

    balances = [_period_amount_from_row(body.iloc[i], sign_mode, block) for i in range(len(body))]
    out["Balance"] = balances
    out = finalize_accounts(out, config)
    if out.empty:
        return None

    out = scale_balance(out, config)
    # Use the FY being processed (cell year), not the mapper preview year in period_blocks.
    block_year = int(year)
    block_month = int(period_block.get("month") if period_block else month)
    out["Year"] = block_year
    out["Month"] = block_month
    out["Period"] = make_period_str(block_year, block_month)
    out["Entity"] = entity_name
    return out


def apply_bilanz_eb_plus_movements(long_df: pd.DataFrame, *, bs_only: bool = False) -> pd.DataFrame:
    """
    Balance sheet (Bilanz): period saldo = opening balance + cumulative movements.
    P&L accounts keep period movements unchanged when bs_only=True.

    Opening balance applies only once (first chronological SuSa). Later fiscal years
    continue from the prior period's running balance (e.g. Jan-25 follows Dec-24).
    """
    if long_df.empty or "Month" not in long_df.columns:
        return long_df

    balance_keys = ["Account", "Account description"]
    if "Entity" in long_df.columns:
        balance_keys = ["Entity", "Account", "Account description"]

    eb_rows = long_df[long_df["Month"] == 0].copy()
    mov_rows = long_df[long_df["Month"] > 0].copy()
    if eb_rows.empty or mov_rows.empty:
        return long_df

    # When movements are provided, PL should not include an opening-balance row.
    if bs_only and "Account Type" in eb_rows.columns:
        eb_rows = eb_rows[eb_rows["Account Type"] == "BS"].copy()

    if "Source FY" in eb_rows.columns and not eb_rows.empty:
        first_eb_fy = int(pd.to_numeric(eb_rows["Source FY"], errors="coerce").min())
        eb_rows = eb_rows[eb_rows["Source FY"] == first_eb_fy].copy()

    eb_map = eb_rows.set_index(balance_keys)["Balance"].to_dict()

    sort_cols = ["Year", "Month"]
    if "Source FY" in mov_rows.columns:
        sort_cols = ["Source FY", "Year", "Month"]
    mov_rows = mov_rows.sort_values(sort_cols)
    out_parts: List[pd.DataFrame] = [eb_rows]

    for keys, grp in mov_rows.groupby(balance_keys, sort=False):
        grp = grp.sort_values(sort_cols)
        if bs_only and "Account Type" in grp.columns and grp["Account Type"].iloc[0] != "BS":
            out_parts.append(grp)
            continue
        eb_val = eb_map.get(keys, 0.0)
        if pd.isna(eb_val):
            eb_val = 0.0
        cumulative = float(eb_val)
        adjusted = []
        for _, row in grp.iterrows():
            m = float(row["Balance"]) if pd.notna(row["Balance"]) else 0.0
            cumulative += m
            nr = row.copy()
            nr["Balance"] = cumulative
            adjusted.append(nr)
        if adjusted:
            out_parts.append(pd.DataFrame(adjusted))

    return pd.concat(out_parts, ignore_index=True)


def pl_movements_from_balances(
    df_long: pd.DataFrame,
    *,
    fiscal_start_month: int,
) -> pd.DataFrame:
    """
    When input value_type == 'balances', P&L period columns are cumulative movements.
    Convert PL balances to monthly movements: move_m = bal_m - bal_{m-1}, first month uses 0 as previous.
    """
    if df_long.empty:
        return df_long
    if not {"Month", "Balance", "Reporting FY"}.issubset(df_long.columns):
        return df_long
    if "Account Type" not in df_long.columns:
        return df_long

    df = df_long.copy()
    pl_mask = df["Account Type"] == "PL"
    if not pl_mask.any():
        return df_long

    # Only months (no EB) contribute to PL movements.
    pl = df.loc[pl_mask & (df["Month"] > 0)].copy()
    if pl.empty:
        return df_long

    idx_cols = ["Account"]
    if "Entity" in pl.columns:
        idx_cols = ["Entity", "Account"]

    pl["FY_INDEX"] = pd.to_numeric(pl["Reporting FY"], errors="coerce").fillna(0).astype(int)
    pl["FY_MONTH_ORDER"] = ((pl["Month"] - int(fiscal_start_month)) % 12).astype(int)

    sort_cols = idx_cols + ["FY_INDEX", "FY_MONTH_ORDER", "Month"]
    if "Calendar Year" in pl.columns:
        sort_cols.append("Calendar Year")
    pl = pl.sort_values(sort_cols, kind="stable")

    pl["__prev__"] = pl.groupby(idx_cols + ["FY_INDEX"], group_keys=False)["Balance"].shift(1)
    pl["__prev__"] = pl["__prev__"].fillna(0.0)
    pl["Balance"] = pl["Balance"].fillna(0.0) - pl["__prev__"]
    pl.drop(columns=["__prev__"], inplace=True)

    df.loc[pl.index, "Balance"] = pl["Balance"]
    return df


def read_data_block(
    raw: pd.DataFrame,
    config: dict,
    year: int,
    month: int,
    entity_name: str,
) -> Optional[pd.DataFrame]:
    if not _has_column_mapping(config):
        return read_data_block_legacy(raw, config, year, month, entity_name)
    return read_period_row_mapped(raw, config, entity_name, year, month)


def process_monthly_sheets_workbook(
    path: str, year: int, config: dict, entity_name: str, *, include_opening_balance: bool
) -> pd.DataFrame:
    xls = pd.ExcelFile(path, engine="openpyxl")
    all_sheets: List[pd.DataFrame] = []
    for sheet in xls.sheet_names:
        raw = pd.read_excel(path, sheet_name=sheet, header=None, engine="openpyxl")
        if include_opening_balance and _has_column_mapping(config):
            ob = read_opening_balance_mapped(raw, config, entity_name, year)
            if ob is not None and not ob.empty:
                all_sheets.append(ob)
                include_opening_balance = False
        try:
            month = parse_month_from_sheetname(sheet)
        except ValueError:
            continue
        block = read_data_block(raw, config, year, month, entity_name)
        if block is not None and not block.empty:
            all_sheets.append(block)
    if not all_sheets:
        raise ValueError(f"No readable monthly sheets in workbook: {path}")
    return pd.concat(all_sheets, ignore_index=True)


def process_monthly_workbooks_cell(
    paths: List[str], year: int, config: dict, entity_name: str, *, include_opening_balance: bool
) -> pd.DataFrame:
    if len(paths) != 12:
        raise ValueError(f"Expected 12 workbooks, got {len(paths)} for {entity_name} FY{year}")
    assigned = assign_months_twelve_workbooks(paths)
    parts: List[pd.DataFrame] = []
    for pth, month in assigned:
        xls = pd.ExcelFile(pth, engine="openpyxl")
        sheet = xls.sheet_names[0]
        raw = pd.read_excel(pth, sheet_name=sheet, header=None, engine="openpyxl")
        if include_opening_balance and _has_column_mapping(config):
            ob = read_opening_balance_mapped(raw, config, entity_name, year)
            if ob is not None and not ob.empty:
                parts.append(ob)
                include_opening_balance = False
        block = read_data_block(raw, config, year, month, entity_name)
        if block is not None and not block.empty:
            parts.append(block)
    if not parts:
        raise ValueError(f"No data read from monthly workbooks for {entity_name} FY{year}")
    return pd.concat(parts, ignore_index=True)


def process_single_sheet_workbook(
    path: str, year: int, config: dict, entity_name: str, *, include_opening_balance: bool
) -> pd.DataFrame:
    raw = pd.read_excel(path, sheet_name=0, header=None, engine="openpyxl")
    if not _has_column_mapping(config):
        header_row = find_header_row(raw)
        df = raw.dropna(how="all").reset_index(drop=True)
        df = df.iloc[header_row + 1:].copy()
        if df.empty:
            raise ValueError(f"Empty sheet in {path}")
        cols_cfg = dict(config.get("columns") or {})
        if "Month" not in cols_cfg:
            raise ValueError("single_sheet layout requires columns.Month (Excel letter) in config.")
        df = apply_letter_mapping(df, cols_cfg)
        df = finalize_accounts(df, config)
        df["Month"] = pd.to_numeric(df["Month"], errors="coerce").fillna(0).astype(int)
        df = df[(df["Month"] >= 1) & (df["Month"] <= 12)]
        df["Balance"] = df["Balance"].apply(parse_number_de)
        mode = str(config.get("sign_mode") or "sh_column").strip()
        if mode == "sh_column" and "S" in df.columns:
            df["S"] = df["S"].astype(str).str.upper().str.strip()
            df.loc[df["S"] == "S", "Balance"] *= -1
        df = scale_balance(df, config)
        df["Year"] = year
        df["Period"] = [make_period_str(year, int(m)) for m in df["Month"]]
        df["Entity"] = entity_name
        return df

    profile = profile_for_entity(config, entity_name)
    header_row = resolve_header_row(raw, profile)
    body = raw.dropna(how="all").reset_index(drop=True).iloc[header_row + 1:].copy()
    if body.empty:
        raise ValueError(f"Empty sheet in {path}")

    sign_mode = str(config.get("sign_mode") or "sh_column").strip()
    blocks = build_period_blocks_from_profile(raw, profile, sign_mode, year)
    if not blocks:
        blocks = expand_period_blocks(profile, body.shape[1], sign_mode, year)
    if not blocks:
        raise ValueError("single_sheet requires period_blocks or repeating pattern in column_mapping.")

    parts: List[pd.DataFrame] = []
    if include_opening_balance:
        ob = read_opening_balance_mapped(raw, config, entity_name, year)
        if ob is not None and not ob.empty:
            parts.append(ob)

    for block in blocks:
        month = int(block.get("month") or 0)
        if month < 1 or month > 12:
            continue
        block_for_read = {**block, "year": year}
        chunk = read_period_row_mapped(
            raw, config, entity_name, year, month, period_block=block_for_read
        )
        if chunk is not None and not chunk.empty:
            parts.append(chunk)

    if not parts:
        raise ValueError(f"No period data extracted from {path}")
    return pd.concat(parts, ignore_index=True)


# ==============================================================
# Long-format metadata + Master_BS / Master_PL pipeline
# ==============================================================

SHEET_BS = "Master_BS"
SHEET_PL = "Master_PL"
CHECK_ROW_LABEL = "Check"
CHECK_TOLERANCE = 0.001
CANVAS_EXTRA_COLS = 20
CANVAS_EXTRA_ROWS_BELOW_CHECK = 200

_META_COLS_BASE = [
    "Entity",
    "Account",
    "Account description",
    "L1 - BS/PL",
    "L2",
    "L3",
    "L4",
    "L5",
    "L6",
]
META_COLS_PL = _META_COLS_BASE + ["Comments", "Q&A", "Answer of Target"]
META_COLS_BS = _META_COLS_BASE + ["NA", "Comments", "Q&A", "Answer of Target"]

TEXT_OUTPUT_COLS_PL = len(_META_COLS_BASE)  # Entity … L6
TEXT_OUTPUT_COLS_BS = len(_META_COLS_BASE) + 1  # Entity … NA

INDEX_COLS_PL = _META_COLS_BASE
INDEX_COLS_BS = _META_COLS_BASE + ["NA"]

NET_INCOME_ACCOUNT = "2869"
NET_INCOME_DESCRIPTION = "Jahresüberschuss/Jahresfehlbetrag"


def enrich_extracted_block(
    df: pd.DataFrame,
    *,
    entity: str,
    source_fy: int,
    source_file: str,
    sheet_name: str = "",
    source_order: int = 0,
) -> pd.DataFrame:
    out = df.copy()
    out["Entity"] = entity
    out["Source FY"] = source_fy
    out["Reporting FY"] = source_fy
    out["Calendar Year"] = out["Year"]
    out["Source Group"] = Path(source_file).stem if source_file else ""
    out["Source File"] = source_file
    out["Sheet"] = sheet_name
    out["Source Order"] = source_order
    return out


def infer_calendar_years_from_sequence(months_in_order: List[int], source_fy: int) -> List[int]:
    if not months_in_order:
        return []
    wrap_idx = None
    for i in range(1, len(months_in_order)):
        if months_in_order[i] < months_in_order[i - 1]:
            wrap_idx = i
            break
    if wrap_idx is None:
        return [source_fy] * len(months_in_order)
    return [source_fy - 1 if i < wrap_idx else source_fy for i in range(len(months_in_order))]


def finalize_periods(df_long: pd.DataFrame, fiscal_start_month: int = 1) -> pd.DataFrame:
    df = df_long.copy()
    month_mask = df["Month"].notna() & df["Source FY"].notna()
    month_seq_map: Dict[Tuple[Any, ...], Dict[int, int]] = {}

    for (entity, source_fy, source_group), grp in df.loc[month_mask].groupby(
        ["Entity", "Source FY", "Source Group"]
    ):
        ordered = (
            grp[["Month", "Source Order"]]
            .dropna(subset=["Month"])
            .sort_values(["Source Order", "Month"], kind="stable")
            .drop_duplicates(subset=["Month"], keep="first")
        )
        months = [int(x) for x in ordered["Month"].tolist() if int(x) > 0]
        if not months:
            continue
        years = infer_calendar_years_from_sequence(months, int(source_fy))
        key = (entity, source_fy, source_group)
        month_seq_map[key] = {m: i + 1 for i, m in enumerate(months)}
        for m, cal_y in zip(months, years):
            mask = (
                (df["Entity"] == entity)
                & (df["Source FY"] == source_fy)
                & (df["Source Group"] == source_group)
                & (df["Month"] == m)
            )
            df.loc[mask, "Calendar Year"] = cal_y
            df.loc[mask, "Period"] = df.loc[mask].apply(
                lambda r: make_period_str(int(r["Calendar Year"]), int(r["Month"])),
                axis=1,
            )

    def _month_seq(row: pd.Series) -> Optional[int]:
        key = (row["Entity"], row["Source FY"], row["Source Group"])
        if key not in month_seq_map or pd.isna(row["Month"]) or int(row["Month"]) <= 0:
            return None
        return month_seq_map[key].get(int(row["Month"]))

    df["Month Seq"] = df.apply(_month_seq, axis=1)
    df["Reporting FY"] = df["Source FY"]
    return df


def _normalize_mapping_frame(raw: pd.DataFrame) -> pd.DataFrame:
    rename = {}
    for col in raw.columns:
        c = str(col).strip()
        if c.lower() in ("account description", "account desc", "konto", "kontobezeichnung", "beschreibung"):
            rename[col] = "Account description"
        elif c in ("L1 - BS/PL", "L1", "BS/PL"):
            rename[col] = "L1 - BS/PL"
        elif c.startswith("L2"):
            rename[col] = "L2"
        elif c.startswith("L3"):
            rename[col] = "L3"
        elif c.startswith("L4"):
            rename[col] = "L4"
        elif c.upper() == "NA" or c.strip().lower() in ("l6 na mapping", "na mapping"):
            rename[col] = "NA"
    df = raw.rename(columns=rename)
    # Mapping files come as BS_Kontenmapping / PL_Kontenmapping and typically contain
    # Account description + L1..L4. We map based on description (not account number).
    required = ["Account description", "L1 - BS/PL", "L2", "L3", "L4", "Account Type"]
    optional = ["NA"]
    for col in required + optional:
        if col not in df.columns:
            df[col] = None
    # If mapping uses plain L1 (BS/PL), align it to the expected columns.
    if df["L1 - BS/PL"].isna().all() and "L1" in raw.columns:
        df["L1 - BS/PL"] = raw["L1"]
    df["Account Type"] = df["Account Type"].where(df["Account Type"].notna(), df["L1 - BS/PL"])

    df["Account description"] = df["Account description"].astype(str).str.strip()
    out_cols = required + [c for c in optional if c in df.columns]
    return df[out_cols]


def _norm_account_description(s: object) -> str:
    txt = "" if s is None else str(s)
    txt = " ".join(txt.strip().lower().split())
    txt = txt.replace("\u00a0", " ")
    return " ".join(txt.split())


def _duplicate_bs_desc_norms(df_mapping: pd.DataFrame) -> set[str]:
    bs = df_mapping[df_mapping["Account Type"].astype(str).str.upper() == "BS"].copy()
    if bs.empty:
        return set()
    counts = bs.groupby("__desc_norm", dropna=False).size()
    return set(counts[counts > 1].index.astype(str))


def _latest_balance_by_entity_desc(df_long: pd.DataFrame) -> Dict[Tuple[str, str], float]:
    """Balance in the last available month (Month>0) per Entity × normalized description."""
    if df_long.empty:
        return {}
    work = df_long.copy()
    work["__desc_norm"] = work.get("Account description", "").astype(str).map(_norm_account_description)
    work["_month"] = pd.to_numeric(work.get("Month"), errors="coerce")
    work = work[work["_month"].fillna(0) > 0]
    if work.empty:
        return {}
    if "Calendar Year" in work.columns:
        work["_year"] = pd.to_numeric(work["Calendar Year"], errors="coerce")
    else:
        work["_year"] = pd.to_numeric(work.get("Year"), errors="coerce")
    work["_bal"] = pd.to_numeric(work.get("Balance"), errors="coerce").fillna(0.0)

    out: Dict[Tuple[str, str], float] = {}
    for (entity, desc_norm), grp in work.groupby(
        [work["Entity"].astype(str), work["__desc_norm"]], dropna=False
    ):
        latest = grp.sort_values(["_year", "_month"], kind="stable").iloc[-1]
        out[(str(entity), str(desc_norm))] = float(latest["_bal"])
    return out


def _resolve_sign_dependent_mapping(
    df_mapping_all: pd.DataFrame,
    df_long: pd.DataFrame,
) -> pd.DataFrame:
    """
    When BS account descriptions appear more than once in the mapping file, pick the row
    based on the sign of the balance in the last available month: negative → NA=ND,
    otherwise the alternate row (e.g. cash/TWC).
    """
    m = df_mapping_all.copy()
    m["__desc_key"] = m.get("Account description", "").astype(str).str.strip()
    m["__desc_norm"] = m["__desc_key"].map(_norm_account_description)

    base = df_long.copy()
    base["__desc_norm"] = base.get("Account description", "").astype(str).map(_norm_account_description)

    dup_norms = _duplicate_bs_desc_norms(m)
    sign_bal = _latest_balance_by_entity_desc(base)

    rows: List[dict] = []
    for desc_norm, grp in m.groupby("__desc_norm", dropna=False):
        desc_norm_s = str(desc_norm)
        if desc_norm_s not in dup_norms:
            row = grp.iloc[0].to_dict()
            row["Entity"] = None
            rows.append(row)
            continue

        if str(grp.iloc[0].get("Account Type", "")).upper() != "BS":
            row = grp.iloc[0].to_dict()
            row["Entity"] = None
            rows.append(row)
            continue

        na_series = grp["NA"].astype(str).str.strip().str.upper() if "NA" in grp.columns else pd.Series(dtype=str)
        nd_rows = grp[na_series == "ND"]
        pos_rows = grp[na_series != "ND"]
        if nd_rows.empty or pos_rows.empty:
            row = grp.iloc[0].to_dict()
            row["Entity"] = None
            rows.append(row)
            continue

        entities = base.loc[base["__desc_norm"] == desc_norm_s, "Entity"].dropna().astype(str).unique()
        if len(entities) == 0:
            row = pos_rows.iloc[0].to_dict()
            row["Entity"] = None
            rows.append(row)
            continue

        for ent in entities:
            bal = sign_bal.get((str(ent), desc_norm_s), 0.0)
            pick = nd_rows.iloc[0] if bal < 0 else pos_rows.iloc[0]
            row = pick.to_dict()
            row["Entity"] = str(ent)
            rows.append(row)

    return pd.DataFrame(rows)


def resolve_mapping_paths(config: dict) -> Tuple[Optional[str], Optional[str]]:
    from susa_mapping_paths import resolve_susa_kontenmapping_paths

    return resolve_susa_kontenmapping_paths(config)


def load_account_type_mapping(config: dict) -> pd.DataFrame:
    bs_path, pl_path = resolve_mapping_paths(config)
    frames: List[pd.DataFrame] = []
    for path, account_type in ((bs_path, "BS"), (pl_path, "PL")):
        if path and Path(path).is_file():
            try:
                raw = pd.read_excel(path, engine="openpyxl")
                frame = _normalize_mapping_frame(raw)
                frame["Account Type"] = account_type
                frames.append(frame)
            except Exception as exc:
                print(f"[WARN] Could not load mapping file {path}: {exc}")
    if not frames:
        print("[WARN] No BS/PL mapping files found; L1–L4 columns will be empty.")
        return pd.DataFrame(
            columns=["Account description", "L1 - BS/PL", "L2", "L3", "L4", "Account Type", "NA"]
        )
    return pd.concat(frames, ignore_index=True, sort=False)


def build_account_type_lookup(df_mapping: pd.DataFrame) -> Dict[str, str]:
    if df_mapping.empty:
        return {}
    return {
        str(row["Account description"]): str(row.get("Account Type") or "UNKNOWN")
        for _, row in df_mapping.iterrows()
        if pd.notna(row["Account description"])
    }


def _attach_mapping(df_long: pd.DataFrame, df_mapping_all: pd.DataFrame) -> pd.DataFrame:
    if df_mapping_all.empty:
        out = df_long.copy()
        out["L1 - BS/PL"] = None
        out["L2"] = None
        out["L3"] = None
        out["L4"] = None
        out["L5"] = ""
        out["L6"] = "Reported"
        out["Account Type"] = "UNKNOWN"
        return out

    base = df_long.copy()
    base["__desc_key"] = base.get("Account description", "").astype(str).str.strip()
    base["__desc_norm"] = base["__desc_key"].map(_norm_account_description)

    mapping_for_dup = df_mapping_all.copy()
    mapping_for_dup["__desc_norm"] = (
        mapping_for_dup.get("Account description", "").astype(str).str.strip().map(_norm_account_description)
    )
    dup_norms = _duplicate_bs_desc_norms(mapping_for_dup)

    m = _resolve_sign_dependent_mapping(df_mapping_all, base)
    m["__desc_key"] = m.get("__desc_key", m.get("Account description", "")).astype(str).str.strip()
    m["__desc_norm"] = m["__desc_key"].map(_norm_account_description)

    map_cols = ["L1 - BS/PL", "L2", "L3", "L4", "Account Type"]
    if "NA" in m.columns:
        map_cols.append("NA")

    m_global = m[m["Entity"].isna() | (m["Entity"].astype(str).str.strip() == "")].copy()
    m_entity = m[m["Entity"].notna() & (m["Entity"].astype(str).str.strip() != "")].copy()

    merged = base.merge(
        m_global[["__desc_key"] + map_cols],
        on="__desc_key",
        how="left",
        suffixes=("", "_map"),
    )

    if not m_entity.empty and dup_norms:
        dup_mask = merged["__desc_norm"].isin(dup_norms)
        if dup_mask.any():
            entity_merge = base.loc[dup_mask, ["Entity", "__desc_key"]].merge(
                m_entity[["Entity", "__desc_key"] + map_cols],
                on=["Entity", "__desc_key"],
                how="left",
            )
            for col in map_cols:
                merged.loc[dup_mask, col] = entity_merge[col].to_numpy()

    # Second pass: normalized match for still-unmapped rows
    need_second = merged["L1 - BS/PL"].isna()
    if need_second.any():
        m_norm_global = m_global.drop(columns=["__desc_key"]).rename(columns={"__desc_norm": "__desc_norm"})
        second = base.loc[need_second, ["__desc_norm"]].merge(
            m_norm_global[["__desc_norm"] + map_cols],
            on="__desc_norm",
            how="left",
        )
        for col in map_cols:
            merged.loc[need_second, col] = merged.loc[need_second, col].combine_first(second[col])

        still_unmapped = merged["L1 - BS/PL"].isna() & merged["__desc_norm"].isin(dup_norms)
        if still_unmapped.any() and not m_entity.empty:
            m_norm_entity = m_entity.drop(columns=["__desc_key"]).rename(columns={"__desc_norm": "__desc_norm"})
            second_entity = base.loc[still_unmapped, ["Entity", "__desc_norm"]].merge(
                m_norm_entity[["Entity", "__desc_norm"] + map_cols],
                on=["Entity", "__desc_norm"],
                how="left",
            )
            for col in map_cols:
                merged.loc[still_unmapped, col] = merged.loc[still_unmapped, col].combine_first(
                    second_entity[col]
                )

    merged["Account Type"] = merged["Account Type"].fillna("UNKNOWN")
    merged["L5"] = ""
    merged["L6"] = "Reported"
    if "NA" in merged.columns:
        merged.loc[merged["Account Type"] != "BS", "NA"] = None
    merged.drop(columns=["__desc_key", "__desc_norm"], inplace=True, errors="ignore")
    return merged


def _fy_end_periods_by_fy(
    df: pd.DataFrame,
    fy_values: List[int],
    *,
    fiscal_start_month: int,
    fy_end_month: int,
) -> Dict[int, str]:
    """Map each Reporting FY to the Period label at fiscal year-end (or last available month)."""
    work = df[df["Month"] > 0].copy()
    result: Dict[int, str] = {}
    for fy in fy_values:
        grp = work[work["Reporting FY"] == fy]
        if grp.empty:
            continue
        end_rows = grp[grp["Month"] == fy_end_month]
        if not end_rows.empty:
            period = (
                end_rows.dropna(subset=["Period"])
                .sort_values(["Calendar Year", "Month"], kind="stable")
                .iloc[-1]["Period"]
            )
        else:
            ordered = grp.assign(
                FY_MONTH_ORDER=lambda d: ((d["Month"] - fiscal_start_month) % 12),
            ).sort_values(["FY_MONTH_ORDER", "Month", "Calendar Year"], kind="stable")
            period = ordered.iloc[-1]["Period"]
        if pd.notna(period):
            result[int(fy)] = str(period)
    return result


def select_bs_fy_end_rows(
    bs_long: pd.DataFrame,
    index_cols: List[str],
    *,
    fiscal_start_month: int,
    fy_end_month: int,
) -> pd.DataFrame:
    """Year-end BS balance per account/FY (Month>0 only; prefers fy_end_month)."""
    mov = bs_long[bs_long["Month"] > 0].copy()
    if mov.empty:
        return mov.iloc[0:0]

    parts: List[pd.Series] = []
    for _, grp in mov.groupby(index_cols + ["Reporting FY"], dropna=False):
        fy_end = grp[grp["Month"] == fy_end_month]
        if not fy_end.empty:
            parts.append(fy_end.sort_values(["Calendar Year", "Month"], kind="stable").iloc[-1])
            continue
        ordered = grp.assign(
            FY_MONTH_ORDER=lambda d: ((d["Month"] - fiscal_start_month) % 12),
        ).sort_values(["FY_MONTH_ORDER", "Month", "Calendar Year"], kind="stable")
        parts.append(ordered.iloc[-1])

    if not parts:
        return mov.iloc[0:0]
    return pd.DataFrame(parts)


def select_bs_ytd_rows(
    bs_long: pd.DataFrame,
    index_cols: List[str],
    *,
    fiscal_start_month: int,
    ytd_reporting_fy: int,
    ltm_year: int,
    ltm_month: int,
) -> pd.DataFrame:
    """BS balance at the as-of month within the open reporting FY."""
    mov = bs_long[
        (bs_long["Month"] > 0) & (bs_long["Reporting FY"] == ytd_reporting_fy)
    ].copy()
    if mov.empty:
        return mov.iloc[0:0]

    parts: List[pd.Series] = []
    for _, grp in mov.groupby(index_cols, dropna=False):
        exact = grp[(grp["Calendar Year"] == ltm_year) & (grp["Month"] == ltm_month)]
        if not exact.empty:
            parts.append(exact.sort_values(["Calendar Year", "Month"], kind="stable").iloc[-1])
            continue
        eligible = grp[
            (grp["Calendar Year"] < ltm_year)
            | ((grp["Calendar Year"] == ltm_year) & (grp["Month"] <= ltm_month))
        ]
        if eligible.empty:
            continue
        ordered = eligible.assign(
            FY_MONTH_ORDER=lambda d: ((d["Month"] - fiscal_start_month) % 12),
        ).sort_values(["Calendar Year", "FY_MONTH_ORDER", "Month"], kind="stable")
        parts.append(ordered.iloc[-1])

    if not parts:
        return mov.iloc[0:0]
    return pd.DataFrame(parts)


def _pl_ytd_mask(
    df: pd.DataFrame,
    ytd_reporting_fy: int,
    ltm_year: int,
    ltm_month: int,
) -> pd.Series:
    return (
        (df["Reporting FY"] == ytd_reporting_fy)
        & (df["Month"] > 0)
        & (
            (df["Calendar Year"] < ltm_year)
            | ((df["Calendar Year"] == ltm_year) & (df["Month"] <= ltm_month))
        )
    )


def build_net_income_bs_rows(
    df: pd.DataFrame,
    period_order: List[str],
    fy_values: List[int],
    fy_rename: Dict[int, str],
    period_cols: List[str],
    *,
    fiscal_start_month: int = 1,
    fy_end_month: int = 12,
    ytd_col: Optional[str] = None,
    ytd_reporting_fy: Optional[int] = None,
    ytd_ltm_year: Optional[int] = None,
    ytd_ltm_month: Optional[int] = None,
) -> pd.DataFrame:
    """
    One BS row per entity: cumulative sum of all PL monthly movements (Soll+/Haben-).
    Period columns are running saldos; FY columns mirror the FY-end period saldo.
    """
    pl = df[df["Account Type"] == "PL"].copy()
    if "Entity" in df.columns:
        entities = sorted(pl["Entity"].dropna().unique().tolist())
        if not entities:
            entities = sorted(df["Entity"].dropna().unique().tolist())
    else:
        entities = [""]

    if not entities:
        return pd.DataFrame()

    pl_monthly = (
        pl[pl["Month"] > 0]
        .groupby(["Entity", "Period"], as_index=False)["Balance"]
        .sum()
    )
    fy_end_periods = _fy_end_periods_by_fy(
        df,
        fy_values,
        fiscal_start_month=fiscal_start_month,
        fy_end_month=fy_end_month,
    )

    rows: List[dict] = []
    for entity in entities:
        ent_pl = pl_monthly[pl_monthly["Entity"] == entity] if not pl_monthly.empty else pl_monthly
        mov_by_period = (
            ent_pl.set_index("Period")["Balance"].to_dict()
            if not ent_pl.empty
            else {}
        )

        row: dict = {
            "Entity": entity,
            "Account": NET_INCOME_ACCOUNT,
            "Account description": NET_INCOME_DESCRIPTION,
            "L1 - BS/PL": "BS",
            "L2": "Equity",
            "L3": "Net retained profits",
            "L4": "Net income",
            "L5": "",
            "L6": "Reported",
            "NA": "Equity",
            "Comments": None,
            "Q&A": None,
            "Answer of Target": None,
        }

        cumulative = 0.0
        for period in period_order:
            if str(period).startswith("EB-"):
                row[period] = 0.0
                continue
            movement = mov_by_period.get(period, 0.0)
            movement = float(movement) if pd.notna(movement) else 0.0
            cumulative += movement
            row[period] = cumulative

        for fy in fy_values:
            fy_col = fy_rename.get(fy, f"FY{str(fy)[-2:]}A")
            end_period = fy_end_periods.get(fy)
            if end_period is not None and end_period in row:
                row[fy_col] = row[end_period]
            else:
                fy_mask = (pl["Entity"] == entity) & (pl["Reporting FY"] == fy) & (pl["Month"] > 0)
                row[fy_col] = float(pl.loc[fy_mask, "Balance"].sum())

        if ytd_col and ytd_reporting_fy and ytd_ltm_year and ytd_ltm_month:
            ltm_period = make_period_str(ytd_ltm_year, ytd_ltm_month)
            if ltm_period in row:
                row[ytd_col] = row[ltm_period]
            else:
                mask = (pl["Entity"] == entity) & _pl_ytd_mask(
                    pl, ytd_reporting_fy, ytd_ltm_year, ytd_ltm_month
                )
                row[ytd_col] = float(pl.loc[mask, "Balance"].sum())

        rows.append(row)

    out = pd.DataFrame(rows)
    if out.empty:
        return out
    return out.reindex(columns=META_COLS_BS + period_cols + period_order)


def validate_df_long(df_long: pd.DataFrame, strict_mapping: bool = False) -> List[str]:
    warnings: List[str] = []
    if df_long.empty:
        warnings.append("No records extracted from trial balances.")
        return warnings
    if strict_mapping and "L1 - BS/PL" in df_long.columns:
        unmapped = df_long[df_long["L1 - BS/PL"].isna()]["Account"].nunique()
        if unmapped:
            warnings.append(f"{unmapped} accounts without mapping (strict_mapping=True).")
    return warnings


def build_final_output(
    df_long: pd.DataFrame,
    fiscal_start_month: int = 1,
    *,
    value_type: str = "balances",
    fy_end_month: int = 12,
    ytd_reporting_fy: Optional[int] = None,
    ytd_ltm_year: Optional[int] = None,
    ytd_ltm_month: Optional[int] = None,
) -> Tuple[pd.DataFrame, pd.DataFrame]:
    df = finalize_periods(df_long, fiscal_start_month)
    for col, default in (("L5", ""), ("L6", "Reported"), ("NA", None)):
        if col not in df.columns:
            df[col] = default
    value_type = str(value_type or "balances").strip()
    if value_type == "balances":
        df = pl_movements_from_balances(df, fiscal_start_month=fiscal_start_month)

    # Period order should be chronological within each FY (EB first FY, then months FY-by-FY),
    # not grouped by month across years.
    period_order: list[str] = []
    fy_values = sorted(df["Reporting FY"].dropna().unique().astype(int).tolist())
    first_fy = min(fy_values) if fy_values else None
    if first_fy is not None:
        eb_label = f"EB-{first_fy}"
        if (df["Period"] == eb_label).any():
            period_order.append(eb_label)
    if fy_values:
        work = df[df["Month"] > 0].copy()
        if not work.empty:
            work["FY_MONTH_ORDER"] = ((work["Month"] - fiscal_start_month) % 12).astype(int)
            for fy in fy_values:
                grp = work[work["Reporting FY"] == fy]
                if grp.empty:
                    continue
                ordered = (
                    grp.dropna(subset=["Period"])
                    .drop_duplicates(subset=["Month", "Period"])
                    .sort_values(["FY_MONTH_ORDER", "Month", "Calendar Year"], kind="stable")
                )
                period_order.extend([p for p in ordered["Period"].tolist() if p not in period_order])

    fy_month_counts = (
        df[df["Month"] > 0]
        .groupby("Reporting FY")["Month"]
        .nunique()
        .to_dict()
    )
    latest_fy = max(fy_values) if fy_values else None
    fy_rename: Dict[int, str] = {}
    fy_cols: List[str] = []
    for fy in fy_values:
        new_name = f"FY{str(fy)[-2:]}A"
        fy_rename[fy] = new_name
        fy_cols.append(new_name)

    ytd_col: Optional[str] = None
    if ytd_reporting_fy and ytd_ltm_year and ytd_ltm_month:
        ytd_col = f"YTD{str(int(ytd_reporting_fy))[-2:]}A"
    period_cols = list(fy_cols)
    if ytd_col:
        period_cols.append(ytd_col)

    index_cols_pl = INDEX_COLS_PL
    index_cols_bs = INDEX_COLS_BS

    pl_long = df[df["Account Type"] == "PL"].copy()

    # Presentation sign for Master_PL:
    # Raw accounting sign is Soll-positive / Haben-negative.
    # For reporting, revenues should be positive and expenses negative.
    pl_long["Balance"] = -pd.to_numeric(pl_long["Balance"], errors="coerce")

    pl_pivot = (
        pl_long.pivot_table(
            index=index_cols_pl,
            columns="Period",
            values="Balance",
            aggfunc="sum",
        )
        .reindex(columns=period_order)
        .reset_index()
    )
    pl_pivot.columns.name = None

    pl_fy_pivot = (
        pl_long.pivot_table(
            index=index_cols_pl,
            columns="Reporting FY",
            values="Balance",
            aggfunc="sum",
        )
        .reindex(columns=fy_values)
        .reset_index()
    )
    if fy_values:
        pl_fy_pivot = pl_fy_pivot.rename(columns=fy_rename)

    if ytd_col and ytd_reporting_fy and ytd_ltm_year and ytd_ltm_month:
        pl_ytd = (
            pl_long.loc[
                _pl_ytd_mask(pl_long, int(ytd_reporting_fy), int(ytd_ltm_year), int(ytd_ltm_month))
            ]
            .groupby(index_cols_pl, as_index=False)["Balance"]
            .sum()
            .rename(columns={"Balance": ytd_col})
        )
        master_pl = pl_pivot.merge(pl_fy_pivot, on=index_cols_pl, how="outer")
        master_pl = master_pl.merge(pl_ytd, on=index_cols_pl, how="left")
    else:
        master_pl = pl_pivot.merge(pl_fy_pivot, on=index_cols_pl, how="outer")

    bs_long = df[df["Account Type"] == "BS"].copy()
    if not bs_long.empty and fy_values:
        bs_fy_end = select_bs_fy_end_rows(
            bs_long,
            index_cols_bs,
            fiscal_start_month=fiscal_start_month,
            fy_end_month=fy_end_month,
        )
        bs_fy_pivot = (
            bs_fy_end.pivot_table(
                index=index_cols_bs,
                columns="Reporting FY",
                values="Balance",
                aggfunc="sum",
            )
            .reindex(columns=fy_values)
            .reset_index()
        )
        bs_fy_pivot = bs_fy_pivot.rename(columns=fy_rename)
        bs_period_pivot = (
            bs_long.pivot_table(
                index=index_cols_bs,
                columns="Period",
                values="Balance",
                aggfunc="sum",
            )
            .reindex(columns=period_order)
            .reset_index()
        )
        master_bs = bs_period_pivot.merge(bs_fy_pivot, on=index_cols_bs, how="outer")
        if ytd_col and ytd_reporting_fy and ytd_ltm_year and ytd_ltm_month:
            bs_ytd = select_bs_ytd_rows(
                bs_long,
                index_cols_bs,
                fiscal_start_month=fiscal_start_month,
                ytd_reporting_fy=int(ytd_reporting_fy),
                ltm_year=int(ytd_ltm_year),
                ltm_month=int(ytd_ltm_month),
            )
            if not bs_ytd.empty:
                bs_ytd_pivot = (
                    bs_ytd.groupby(index_cols_bs, as_index=False)["Balance"]
                    .sum()
                    .rename(columns={"Balance": ytd_col})
                )
                master_bs = master_bs.merge(bs_ytd_pivot, on=index_cols_bs, how="left")
    else:
        master_bs = pd.DataFrame(columns=META_COLS_BS + period_cols + period_order)

    net_income = build_net_income_bs_rows(
        df,
        period_order,
        fy_values,
        fy_rename,
        period_cols,
        fiscal_start_month=fiscal_start_month,
        fy_end_month=fy_end_month,
        ytd_col=ytd_col,
        ytd_reporting_fy=ytd_reporting_fy,
        ytd_ltm_year=ytd_ltm_year,
        ytd_ltm_month=ytd_ltm_month,
    )
    if not net_income.empty:
        master_bs = pd.concat([master_bs, net_income], ignore_index=True)

    for col in META_COLS_PL:
        if col not in master_pl.columns:
            master_pl[col] = None
    for col in META_COLS_BS:
        if col not in master_bs.columns:
            master_bs[col] = None

    desired_cols_pl = META_COLS_PL + period_cols + period_order
    desired_cols_bs = META_COLS_BS + period_cols + period_order
    master_pl = master_pl.reindex(columns=desired_cols_pl)
    master_bs = master_bs.reindex(columns=desired_cols_bs)

    for master in (master_pl, master_bs):
        master["L5"] = ""
        master["L6"] = "Reported"
        master["__sort"] = pd.to_numeric(master["Account"], errors="coerce")
        master.sort_values(["Entity", "__sort", "Account"], inplace=True)
        master.drop(columns=["__sort"], inplace=True)

    if "L1 - BS/PL" in master_pl.columns and master_pl["L1 - BS/PL"].notna().any():
        master_pl = master_pl[master_pl["L1 - BS/PL"] == "PL"].copy()
    if "L1 - BS/PL" in master_bs.columns and master_bs["L1 - BS/PL"].notna().any():
        master_bs = master_bs[master_bs["L1 - BS/PL"] == "BS"].copy()

    return master_bs, master_pl


# ==============================================================
# Excel formatting (GST theme)
# ==============================================================
FMT_KEUR = '#,##0;(#,##0);"-"'

CHECK_FILL = PatternFill("solid", fgColor="FFFEF3C7")
RED_FONT = Font(name=THEME.font_name, size=THEME.font_size, bold=True, color="FF9C0006")


def paint_canvas_white(ws, max_col: int, max_row: int):
    for r in range(1, max_row + 1):
        for c in range(1, max_col + 1):
            cell = ws.cell(row=r, column=c)
            cell.fill = THEME.fill_white
            cell.font = THEME.font_base


def set_row_heights(ws, max_row: int, height: float = 12):
    for r in range(1, max_row + 1):
        ws.row_dimensions[r].height = height


def format_header(ws, table_cols: int, n_text_cols: int):
    for c in range(1, table_cols + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = THEME.font_header
        cell.fill = THEME.fill_header
        cell.border = THEME.border_header_bottom
        cell.alignment = Alignment(
            horizontal="left" if c <= n_text_cols else "right",
            vertical="center",
        )


def auto_size_account_columns(ws, n_text_cols: int = 2):
    for col_idx in range(1, n_text_cols + 1):
        max_len = max(
            len(str(ws.cell(row=r, column=col_idx).value))
            for r in range(1, ws.max_row + 1)
            if ws.cell(row=r, column=col_idx).value is not None
        )
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 2, 60)


def format_data_cells(ws, n_rows: int, n_cols: int, n_text_cols: int):
    for r in range(2, n_rows + 2):
        for c in range(1, n_text_cols + 1):
            ws.cell(row=r, column=c).font = THEME.font_base
            ws.cell(row=r, column=c).alignment = Alignment(horizontal="left")
        for c in range(n_text_cols + 1, n_cols + 1):
            cell = ws.cell(row=r, column=c)
            cell.font = THEME.font_base
            cell.number_format = FMT_KEUR
            cell.alignment = Alignment(horizontal="right")


def add_total_pl_row(ws, df_shape: Tuple[int, int], label: str, n_text_cols: int) -> int:
    n_rows, n_cols = df_shape
    first_data_row = 2
    last_data_row = n_rows + 1
    total_row = last_data_row + 2
    for c in range(1, n_cols + 1):
        cell = ws.cell(row=total_row, column=c)
        cell.fill = CHECK_FILL
        cell.font = THEME.font_bold
        cell.alignment = Alignment(horizontal="left" if c <= n_text_cols else "right")
    ws.cell(row=total_row, column=1).value = label
    for c in range(n_text_cols + 1, n_cols + 1):
        col = get_column_letter(c)
        ws.cell(
            row=total_row,
            column=c,
            value=f"=SUBTOTAL(9,{col}{first_data_row}:{col}{last_data_row})",
        ).number_format = FMT_KEUR
    return total_row


def _write_master_sheet(
    wb,
    sheet_name: str,
    master: pd.DataFrame,
    config: dict,
    *,
    add_total: bool = False,
) -> None:
    if sheet_name in wb.sheetnames:
        del wb[sheet_name]
    ws = wb.create_sheet(sheet_name)
    n_text_cols = TEXT_OUTPUT_COLS_BS if sheet_name == SHEET_BS else TEXT_OUTPUT_COLS_PL
    for row in dataframe_to_rows(master, index=False, header=True):
        ws.append(row)
    n_rows, n_cols = master.shape
    canvas_bottom = n_rows + 3 + int(config.get("canvas_extra_rows_below_check", CANVAS_EXTRA_ROWS_BELOW_CHECK))
    canvas_right = n_cols + int(config.get("canvas_extra_cols", CANVAS_EXTRA_COLS))
    paint_canvas_white(ws, canvas_right, canvas_bottom)
    set_row_heights(ws, canvas_bottom, height=12)
    format_header(ws, n_cols, n_text_cols)
    format_data_cells(ws, n_rows, n_cols, n_text_cols)
    auto_size_account_columns(ws, n_text_cols)
    add_check_row(
        ws,
        master.shape,
        float(config.get("check_tolerance", CHECK_TOLERANCE)),
        str(config.get("check_row_label") or CHECK_ROW_LABEL),
        n_text_cols + 1,
    )
    if add_total:
        add_total_pl_row(ws, master.shape, "Total PL", n_text_cols)


def write_output(
    master_bs: pd.DataFrame,
    master_pl: pd.DataFrame,
    output_path: str,
    config: dict,
) -> None:
    output_file = output_path
    if os.path.exists(output_file):
        wb = load_workbook(output_file)
        for sn in (SHEET_BS, SHEET_PL):
            if sn in wb.sheetnames:
                del wb[sn]
    else:
        wb = Workbook()
        if "Sheet" in wb.sheetnames:
            del wb["Sheet"]
    _write_master_sheet(wb, SHEET_BS, master_bs, config)
    _write_master_sheet(wb, SHEET_PL, master_pl, config, add_total=True)
    wb.save(output_file)


def add_check_row(ws, df_shape, tolerance: float, label: str, first_amount_col: int):
    n_rows, n_cols = df_shape
    first_data_row = 2
    last_data_row = n_rows + 1
    check_row = last_data_row + 2

    for c in range(1, n_cols + 1):
        cell = ws.cell(row=check_row, column=c)
        cell.fill = CHECK_FILL
        cell.font = THEME.font_bold
        cell.alignment = Alignment(horizontal="left" if c < first_amount_col else "right")

    ws.cell(row=check_row, column=1).value = label

    for c in range(first_amount_col, n_cols + 1):
        col = get_column_letter(c)
        ws.cell(
            row=check_row,
            column=c,
            value=f'=IF(ABS(SUM({col}{first_data_row}:{col}{last_data_row}))<1E-7,"-",SUM({col}{first_data_row}:{col}{last_data_row}))'
        ).number_format = FMT_KEUR

    start_letter = get_column_letter(first_amount_col)
    rule = FormulaRule(
        formula=[f"ABS({start_letter}{check_row})>{tolerance}"],
        font=RED_FONT,
    )
    ws.conditional_formatting.add(
        f"{start_letter}{check_row}:{get_column_letter(n_cols)}{check_row}",
        rule,
    )
    return check_row


def main(config: dict) -> None:
    if not _has_column_mapping(config):
        raise SystemExit("column_mapping.default is required for SuSa processing.")

    layout = str(config.get("layout_format") or "monthly_sheets").strip()
    cells = config.get("entity_year_files") or []
    if not cells:
        raise SystemExit("Config must contain non-empty entity_year_files.")

    fy_end_month = int(config.get("fy_end_month") or 12)
    fiscal_start_month = int(config.get("fiscal_start_month") or ((fy_end_month % 12) + 1))

    all_long: List[pd.DataFrame] = []
    value_type = str(config.get("value_type") or "balances").strip()
    # Only include Eröffnungsbilanz for the first FY in the configured timeframe.
    fy_cells = [
        c
        for c in cells
        if not str(c.get("fy_label") or "").strip().upper().startswith("YTD")
    ]
    try:
        first_fy = min(int(c.get("year") or 0) for c in fy_cells if int(c.get("year") or 0) > 0)
    except ValueError:
        first_fy = None

    for cell in cells:
        paths = cell.get("paths") or []
        year = int(cell.get("year") or 0)
        if year <= 0:
            raise SystemExit(f"Invalid year in cell: {cell!r}")
        entity_name = str(cell.get("entity_name") or "")
        source_file = paths[0] if paths else ""
        is_ytd_cell = str(cell.get("fy_label") or "").strip().upper().startswith("YTD")

        include_opening = bool(
            not is_ytd_cell and first_fy is not None and year == first_fy
        )
        if layout == "monthly_workbooks":
            df = process_monthly_workbooks_cell(
                paths, year, config, entity_name, include_opening_balance=include_opening
            )
        elif layout == "monthly_sheets":
            if len(paths) != 1:
                raise SystemExit(f"monthly_sheets expects 1 workbook per cell, got {len(paths)}")
            df = process_monthly_sheets_workbook(
                paths[0], year, config, entity_name, include_opening_balance=include_opening
            )
        elif layout == "single_sheet":
            if len(paths) != 1:
                raise SystemExit(f"single_sheet expects 1 workbook per cell, got {len(paths)}")
            df = process_single_sheet_workbook(
                paths[0], year, config, entity_name, include_opening_balance=include_opening
            )
        else:
            raise SystemExit(f"Unknown layout_format: {layout!r}")

        all_long.append(
            enrich_extracted_block(
                df,
                entity=entity_name,
                source_fy=year,
                source_file=source_file,
            )
        )

    combined = pd.concat(all_long, ignore_index=True)

    df_mapping = load_account_type_mapping(config)
    combined = _attach_mapping(combined, df_mapping)
    if value_type == "movements" and (combined["Month"] == 0).any():
        has_typed_mapping = (
            not df_mapping.empty
            and "Account Type" in combined.columns
            and combined["Account Type"].isin(["BS", "PL"]).any()
        )
        combined = apply_bilanz_eb_plus_movements(combined, bs_only=has_typed_mapping)
    for w in validate_df_long(combined, strict_mapping=bool(config.get("strict_mapping"))):
        print(f"[WARN] {w}")
    if bool(config.get("strict_mapping")) and "L1 - BS/PL" in combined.columns:
        if combined["L1 - BS/PL"].isna().any():
            raise SystemExit("strict_mapping: unmapped accounts remain.")

    ytd_reporting_fy = config.get("ytd_reporting_fy_end_year")
    ytd_ltm_year = config.get("ytd_ltm_year")
    ytd_ltm_month = config.get("ytd_ltm_month")

    master_bs, master_pl = build_final_output(
        combined,
        fiscal_start_month=fiscal_start_month,
        value_type=value_type,
        fy_end_month=fy_end_month,
        ytd_reporting_fy=int(ytd_reporting_fy) if ytd_reporting_fy else None,
        ytd_ltm_year=int(ytd_ltm_year) if ytd_ltm_year else None,
        ytd_ltm_month=int(ytd_ltm_month) if ytd_ltm_month else None,
    )

    output_path = config.get("output_path")
    if not output_path:
        out_dir = str(config.get("output_file_path") or ".")
        case_id = str(config.get("case_id") or "output")
        project_name = str(config.get("title") or config.get("project_name") or "Project")
        try:
            from databook_paths import resolve_existing_master_path

            resolved = resolve_existing_master_path(
                out_dir,
                session_id=case_id,
                project_name=project_name,
            )
            if resolved is not None:
                output_path = str(resolved)
        except OSError:
            pass
        if not output_path:
            output_path = str(Path(out_dir) / f"{case_id}_SuSa_Master.xlsx")

    write_output(master_bs, master_pl, str(output_path), config)
    print("Finished:", output_path)


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Usage: python SuSabyYear.py <config.json>", file=sys.stderr)
        sys.exit(2)
    cfg_path = Path(sys.argv[1])
    with cfg_path.open(encoding="utf-8") as fh:
        cfg: dict[str, Any] = json.load(fh)
    main(cfg)
