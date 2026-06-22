import json
import os
import re
import sys
from typing import Any

import pandas as pd
from funktionssammlung import (
    apply_filters,
    build_output_file_path,
    ensure_output_writable,
    fdd_as_of_end,
    fdd_as_of_is_fy_end,
    fdd_current_fy_end_year,
)

DROP_INDEX_PRINT_LIMIT = 200

CONFIG = {                     #Alles was potenziell durch Eingabe veränderbar sein soll, manche Punkte folgen zwangsläufig a...

    "title": "Project Draft",
    "table": "ARR Bridge",
    "company": "Draft AG",

    # Fiscal year end
    "fy_end_month": 12,
    "fy_end_day": 31,
    "fy_end_year": 2024,

    # Core columns-- potenziell änderbar, alles davon wird aber zwangsläufig in irgendeiner Form gebraucht
    "value_col": "Contract Value After Discount",   #End Contract Value?
    "customer_col": "End Customer ID",              #Customer ID?
    "product_col": "Product Category",              #Product Name, Line?
    "start_col": "Contract Start Date",             #End...?
    "end_col": "Contract End Date",                 #End
    "invoice_col": "Invoice Date",

    # Filter
    "filters": {
        "enabled": True,
        "rules": [#{"col": "Service Type", "op": "in", "values": {"Hardware", "License"}},
            #{"col": "Entity", "op": "not_in", "values": {"Internal"}},
            #{"col": "Revenue", "op": ">=", "value": 0},
            #{"col": "Invoice Date", "op": "between", "start": "2024-01-01", "end": "2024-12-31"},
            #{"col": "Customer Name", "op": "regex", "pattern": r"^SAP|^Oracle"},
            #{"col": "Quantity", "op": "notna"}
        ],
    },

    #FX
    "apply_fx": True,
    "fx_col": "Functional FX Rate",                 #FX Rate, Constant FX Rate?

    # Universal grouping
    # Hierarchie-Reihenfolge: erst Parent (z.B. Region), dann Kind (z.B. Country), oder eben auch nur eine Spalte
    "group_cols": ["Entity"],

    # Sorting rules (
    "sort": {
        "top_level_by": "ARR2",                     # top_level: Parent (bei zwei group_cols) oder die group_col (bei nur einer), "ARR..."
        "top_level_desc": True,
        "leaf_within_parent": "ARR2",               # "ARR2" oder "name", nur relevant wenn es zwei group_cols gibt
        "leaf_desc": True
    },

    # Optional Top-Bucket, nur anwendbar bei 1 group_col
    "top_bucket": {
        "enabled": False,
        "thresholds": (0.2, 0.5, 0.8),
        "labels": ("Top {a}", "Top {a_plus1}-{b}", "Top {b_plus1}-{c}", "Other"),
        "based_on": "current"  # auch "previous" möglich: dann Top Bucket nach älterem ARR
    },

    "total_label": "Total",
}


def normalize_config(cfg: dict) -> dict:                #CONFIG Eingaben prüfen

    out = dict(cfg)

    if out.get("fy_end_year") in (None, ""):
        fy_end_m = int(out.get("fy_end_month", 12) or 12)
        fy_end_d = int(out.get("fy_end_day", 31) or 31)
        cy = int(out.get("current_year") or out.get("as_of_year") or 0)
        cm = int(out.get("current_month") or out.get("as_of_month") or 0)
        if cy and cm:
            as_of_end = fdd_as_of_end(cy, cm)
            current_fy_end_year = fdd_current_fy_end_year(as_of_end, fy_end_m, fy_end_d)
            as_of_is_fy_end = fdd_as_of_is_fy_end(as_of_end, fy_end_m, fy_end_d)
            out["fy_end_year"] = current_fy_end_year if as_of_is_fy_end else current_fy_end_year - 1

    # 1. Pflichtfelder prüfen
    required = [
        "fy_end_year",
        "fy_end_month",
        "fy_end_day",
        "value_col",
        "customer_col",
        "product_col",
        "start_col",
        "end_col",
        "invoice_col",
        "group_cols"
    ]

    for k in required:
        if k not in out:
            raise ValueError(f"CONFIG['{k}'] fehlt.")

    # 2. group_cols normalisieren
    cols = out.get("group_cols")

    if isinstance(cols, str):
        cols = [cols]

    if not isinstance(cols, (list, tuple)) or len(cols) == 0:
        raise ValueError("CONFIG['group_cols'] muss eine nicht-leere Liste sein.")

    out["group_cols"] = list(cols)
    out["group_labels"] = list(cols)

    # 3. FX Validierung
    if out.get("apply_fx", False):
        if not out.get("fx_col"):
            raise ValueError("apply_fx=True aber fx_col fehlt.")

    out.setdefault("total_label", "Total")
    out.setdefault("table", "ARR Bridge")
    out.setdefault("first_fy", int(out.get("fy_end_year", 2024)) - 2)

    return out


def resolve_fy_years(cfg: dict) -> tuple[int, int]:
    """Return (first_fy, latest_closed_fy_end_year) from bot date settings."""
    cfg = normalize_config(cfg)
    first_fy = int(cfg["first_fy"])
    fy_end_m = int(cfg.get("fy_end_month", 12))
    fy_end_d = int(cfg.get("fy_end_day", 31))
    cy = int(cfg.get("current_year") or cfg.get("as_of_year") or cfg.get("fy_end_year"))
    cm = int(cfg.get("current_month") or cfg.get("as_of_month") or 12)
    as_of_end = fdd_as_of_end(cy, cm)
    current_fy_end_year = fdd_current_fy_end_year(as_of_end, fy_end_m, fy_end_d)
    as_of_is_fy_end = fdd_as_of_is_fy_end(as_of_end, fy_end_m, fy_end_d)
    latest_closed = current_fy_end_year if as_of_is_fy_end else current_fy_end_year - 1
    if first_fy > latest_closed:
        raise ValueError(
            f"CONFIG['first_fy']={first_fy} liegt nach dem letzten abgeschlossenen FY ({latest_closed})."
        )
    return first_fy, latest_closed


def bridge_year_pairs(cfg: dict) -> list[tuple[int, int]]:
    first_fy, latest_closed = resolve_fy_years(cfg)
    return [(y, y + 1) for y in range(first_fy, latest_closed)]


def days_inclusive(start: pd.Series, end: pd.Series) -> pd.Series:    #Funktion die Vertragslaufzeit in Tagen berechnet
    days = (end - start).dt.days + 1  # +1 weil Mathe
    return days.clip(lower=1)


def report_drops(before_idx: pd.Index, after_idx: pd.Index, label: str):    #Funktion die nicht verwendbare Zeilen mit Index printt
    dropped = before_idx.difference(after_idx)
    n = len(dropped)
    if n == 0:
        return
    head = list(dropped[:DROP_INDEX_PRINT_LIMIT])
    more = "" if n <= DROP_INDEX_PRINT_LIMIT else f" ... (+{n - DROP_INDEX_PRINT_LIMIT} weitere)"
    print(f"[DROP] {label}: {n} Zeilen entfernt. Indizes (erste {min(n, DROP_INDEX_PRINT_LIMIT)}): {head}{more}")


def preprocess_contract_lines(df: pd.DataFrame, cfg: dict) -> pd.DataFrame:    #Funktion um relevante Daten vorzuverarbeiten
    cfg = normalize_config(cfg)

    d = df.copy()
    d.columns = d.columns.str.strip()

    #filter
    d = apply_filters(d, cfg)

    # Daten
    d[cfg["start_col"]] = pd.to_datetime(d[cfg["start_col"]], errors="coerce", dayfirst=True)
    d[cfg["end_col"]] = pd.to_datetime(d[cfg["end_col"]], errors="coerce", dayfirst=True)
    d[cfg["invoice_col"]] = pd.to_datetime(d[cfg["invoice_col"]], errors="coerce", dayfirst=True)

    #Core Columns
    d["cust_id"] = d[cfg["customer_col"]].astype(str).str.strip()
    d["product_key"] = d[cfg["product_col"]].astype(str).str.strip()

    d["contract_value"] = pd.to_numeric(d[cfg["value_col"]], errors="coerce")

    #FX Rate anwenden
    if cfg.get("apply_fx", False):
        fx_col = cfg.get("fx_col")
        if not fx_col or fx_col not in d.columns:
            raise ValueError("CONFIG['apply_fx']=True, aber fx_col fehlt oder existiert nicht im Input.")
        d["fx_rate"] = pd.to_numeric(d[fx_col], errors="coerce")
        # ungültige FX-Rates (0 oder NaN) -> 1.0
        d["fx_rate"] = d["fx_rate"].where(
            d["fx_rate"].notna() & (d["fx_rate"] != 0), 1.0)

    # Jetzt alle multiplizieren
    if cfg.get("apply_fx", False):
        d["contract_value"] = d["contract_value"] * d["fx_rate"]
    else:
        d["contract_value"] = d["contract_value"].fillna(0.0)

    #group_0, group_1 als universelle Keys
    group_cols = cfg["group_cols"]
    for i, col in enumerate(group_cols):
        if col not in d.columns:
            raise ValueError(f"group_col '{col}' existiert nicht im DataFrame.")
        d[f"group_{i}"] = d[col].astype(str).str.strip()

    #Report Dropped zeilen inklusive Begründung (columnangabe)
    core_fields = [
        cfg["start_col"],
        cfg["end_col"],
        "contract_value",
        "cust_id",
        "product_key",
        cfg["invoice_col"]
    ]

    for col in core_fields:
        before = d.index
        d = d[d[col].notna()]
        report_drops(before, d.index, f"Drop NA ({col})")

    # Drop leere group keys
    for i in range(len(cfg["group_cols"])):
        before = d.index
        d = d[d[f"group_{i}"].notna() & (d[f"group_{i}"] != "")]
        report_drops(before, d.index, f"Drop empty group_{i} ({cfg['group_cols'][i]})")

    # Vertragslaufzeit in Tagen
    d["contract_days"] = days_inclusive(d[cfg["start_col"]], d[cfg["end_col"]])

    return d


def snapshot_arr(d: pd.DataFrame, as_of: pd.Timestamp, cfg: dict) -> pd.DataFrame:    #Zieht Daten aus dem Monatssnapsho...
    start_col, end_col = cfg["start_col"], cfg["end_col"]
    inv_col = cfg["invoice_col"]

    # Monat des Stichtags
    month_start = as_of.replace(day=1)  #1. Dezember
    month_end = as_of

    # 1) Nur bereits invoiced bis Stichtag
    dd = d[d[inv_col].notna() & (d[inv_col] <= as_of)].copy()

    # 2) Overlap days zwischen Vertragsperiode und Monat
    overlap_start = dd[start_col].where(dd[start_col] > month_start, month_start)
    overlap_end = dd[end_col].where(dd[end_col] < month_end, month_end)

    overlap_days = (overlap_end - overlap_start).dt.days + 1
    overlap_days = overlap_days.clip(lower=0)              #Keine negativen vertragstage

    # 3) Tagesrate
    daily_rate = dd["contract_value"] / dd["contract_days"]

    # 4) MRR
    dd["mrr_month"] = daily_rate * overlap_days

    # 5) ARR = MRR * 12
    dd["arr"] = dd["mrr_month"] * 12.0

    #dynamische Gruppierung
    group_keys = [c for c in dd.columns if c.startswith("group_")]
    group_keys = sorted(group_keys, key=lambda x: int(x.split("_")[1]))  # group_0, group_1, ...

    agg_keys = group_keys + ["cust_id", "product_key"]

    # Aggregation auf group_keys-customer-product
    snap = (
        dd.loc[dd["arr"].notna() & (dd["arr"] != 0), agg_keys + ["arr",]]
        .groupby(agg_keys, as_index=False)
        .agg(arr=("arr", "sum"),)
    )
    return snap


def _prepare_for_line_arr(df: pd.DataFrame, cfg: dict) -> pd.DataFrame:
    """Row-level prep for helper columns (keeps source row alignment)."""
    cfg = normalize_config(cfg)
    d = df.copy()
    d.columns = d.columns.str.strip()
    d = apply_filters(d, cfg)
    d[cfg["start_col"]] = pd.to_datetime(d[cfg["start_col"]], errors="coerce", dayfirst=True)
    d[cfg["end_col"]] = pd.to_datetime(d[cfg["end_col"]], errors="coerce", dayfirst=True)
    d[cfg["invoice_col"]] = pd.to_datetime(d[cfg["invoice_col"]], errors="coerce", dayfirst=True)
    d["contract_value"] = pd.to_numeric(d[cfg["value_col"]], errors="coerce").fillna(0.0)
    if cfg.get("apply_fx", False):
        fx_col = cfg.get("fx_col")
        if fx_col and fx_col in d.columns:
            fx = pd.to_numeric(d[fx_col], errors="coerce")
            fx = fx.where(fx.notna() & (fx != 0), 1.0)
            d["contract_value"] = d["contract_value"] * fx
    d["contract_days"] = days_inclusive(d[cfg["start_col"]], d[cfg["end_col"]])
    return d


def contract_line_arr(d: pd.DataFrame, as_of: pd.Timestamp, cfg: dict) -> pd.Series:
    """Per contract-line ARR at as_of (same rules as snapshot_arr, before groupby)."""
    start_col, end_col = cfg["start_col"], cfg["end_col"]
    inv_col = cfg["invoice_col"]
    month_start = as_of.replace(day=1)
    month_end = as_of

    dd = d[d[inv_col].notna() & (d[inv_col] <= as_of)].copy()
    overlap_start = dd[start_col].where(dd[start_col] > month_start, month_start)
    overlap_end = dd[end_col].where(dd[end_col] < month_end, month_end)
    overlap_days = (overlap_end - overlap_start).dt.days + 1
    overlap_days = overlap_days.clip(lower=0)
    daily_rate = dd["contract_value"] / dd["contract_days"]
    mrr = daily_rate * overlap_days
    arr = mrr * 12.0
    out = pd.Series(0.0, index=d.index)
    out.loc[dd.index] = arr.fillna(0.0)
    return out


def compute_bridge(fy1: pd.DataFrame, fy2: pd.DataFrame) -> pd.DataFrame:
    #Group-keys weiterhin dynamisch halten
    group_keys = [c for c in fy1.columns if c.startswith("group_")]
    group_keys = sorted(set(group_keys) | set([c for c in fy2.columns if c.startswith("group_")]),
                        key=lambda x: int(x.split("_")[1]))

    # 1) Group-level ARR1/ARR2
    arr1_g = fy1.groupby(group_keys, as_index=False)["arr"].sum().rename(columns={"arr": "ARR1"})
    arr2_g = fy2.groupby(group_keys, as_index=False)["arr"].sum().rename(columns={"arr": "ARR2"})

    # 2) Customer nach Group (New/Lost/Retained)
    c1 = fy1.groupby(group_keys + ["cust_id"], as_index=False)["arr"].sum().rename(columns={"arr": "arr1_cust"})    #1. Jahr
    c2 = fy2.groupby(group_keys + ["cust_id"], as_index=False)["arr"].sum().rename(columns={"arr": "arr2_cust"})    #2. Jahr

    cust = c1.merge(c2, on=group_keys + ["cust_id"], how="outer")                                                   #Merge
    cust["arr1_cust"] = cust["arr1_cust"].fillna(0.0)
    cust["arr2_cust"] = cust["arr2_cust"].fillna(0.0)

    cust["is_new"] = (cust["arr1_cust"] == 0) & (cust["arr2_cust"] != 0)                                           #New: Customer...
    cust["is_lost"] = (cust["arr1_cust"] != 0) & (cust["arr2_cust"] == 0)                                          #Lost: C hat...
    cust["is_retained"] = (cust["arr1_cust"] != 0) & (cust["arr2_cust"] != 0)                                      #Retained: C...

    new_g = cust.loc[cust["is_new"]].groupby(group_keys, as_index=False)["arr2_cust"].sum().rename(columns={"arr2_cust": "New"})
    lost_g = cust.loc[cust["is_lost"]].groupby(group_keys, as_index=False)["arr1_cust"].sum().rename(columns={"arr1_cust": "Lost"})

    # 3) Product-level merge für retained Kunden (Cross/Up/Down)
    p = fy1.merge(
        fy2,
        on=group_keys + ["cust_id", "product_key"],
        how="outer",
        suffixes=("1", "2"),
    )

    # arr1/arr2 sicher ziehen
    for c in ["arr1", "arr2"]:
        p[c] = pd.to_numeric(p[c], errors="coerce").fillna(0.0)

    # retained flag hinzufügen
    retained = cust.loc[cust["is_retained"], group_keys + ["cust_id"]].copy()
    retained["retained"] = True
    p = p.merge(retained, on=group_keys + ["cust_id"], how="left")
    p["retained"] = p["retained"].fillna(False)

    # Cross-sell: retained & arr1=0
    cross = p[(p["retained"]) & (p["arr1"] == 0)]                                                                   #Cross-sell: Kun...
    cross_g = cross.groupby(group_keys, as_index=False)["arr2"].sum().rename(columns={"arr2": "CrossSell"})         #Output

    # Up/Downsell: retained & arr1!=0
    both = p[(p["retained"]) & (p["arr1"] != 0)].copy()                                                            #Für Up-/Downsell: Kunde hatte Umsätze...
    both["delta"] = both["arr2"] - both["arr1"]
    both["upsell_delta"] = both["delta"].clip(lower=0.0)
    both["downsell_delta"] = (-both["delta"]).clip(lower=0.0)

    upsell_g = both.groupby(group_keys, as_index=False)["upsell_delta"].sum().rename(columns={"upsell_delta": "Upsell"})
    downsell_g = both.groupby(group_keys, as_index=False)["downsell_delta"].sum().rename(columns={"downsell_delta": "Downsell"})

    # 4) Alles mergen
    out = arr1_g.merge(arr2_g, on=group_keys, how="outer")

    for t in [new_g, lost_g, cross_g, upsell_g, downsell_g]:
        out = out.merge(t, on=group_keys, how="left")

    for c in ["ARR1", "ARR2", "New", "Lost", "CrossSell", "Upsell", "Downsell"]:
        if c in out.columns:
            out[c] = out[c].fillna(0.0)

    # 5) NRR
    out["NRR"] = (out["ARR1"] - out["Lost"] + out["Upsell"] + out["CrossSell"] - out["Downsell"])

    # 6) Naming/Order (noch ohne Hierarchie-Layout, dafür nächste Funktion)
    out = out.rename(columns={"CrossSell": "Cross-sell"})
    out = out[group_keys + [
        "ARR1",
        "Upsell",
        "Downsell",
        "Cross-sell",
        "Lost",
        "NRR",
        "New",
        "ARR2",
    ]]

    return out


def finalize_bridge_layout(bridge: pd.DataFrame, cfg: dict) -> pd.DataFrame:    #Tabellen anordnen, sortieren, Groupier...
    cfg = normalize_config(cfg)

    #group keys finden
    group_keys = [c for c in bridge.columns if c.startswith("group_")]
    group_keys = sorted(group_keys, key=lambda x: int(x.split("_")[1]))
    n_levels = len(group_keys)

    #Anzeigenamen für group-Spalten bestimmen
    # Standard: aus CONFIG["group_cols"] übernehmen (1:1)
    display_names = cfg["group_cols"][:]  # z.B. ["Entity"] oder ["Customer Region","Customer Geography"]

    #Spalten umbenennen
    rename_map = {group_keys[i]: display_names[i] for i in range(n_levels)}

    out = bridge.copy()
    out = out.rename(columns=rename_map)

    # FY-Spalten dynamisch erkennen
    fy_cols = sorted([c for c in out.columns if c.startswith("FY")])

    # Sicherstellen: älteres FY links, aktuelles rechts
    if len(fy_cols) >= 2:
        fy_left = fy_cols[0]    # z.B. FY23A
        fy_right = fy_cols[-1]  # z.B. FY24A
    else:
        fy_left = fy_cols[0] if fy_cols else None
        fy_right = None

    other_metrics = ["Upsell", "Downsell", "Cross-sell", "Lost", "NRR", "New"]

    metric_cols = []

    if fy_left:
        metric_cols.append(fy_left)

    metric_cols += [c for c in other_metrics if c in out.columns]

    if fy_right:
        metric_cols.append(fy_right)

    # row_type hinzufügen (leaf/parent/total) für spätere Formatierung
    out["row_type"] = "leaf"

    #ab hier Fallunterscheidung für Sortierung
    # Hierarchie: Parent-Totals einziehen falls 2 groupkeys
    if n_levels >= 2:
        parent_col = display_names[0]
        leaf_col = display_names[1]

        # Parent totals (Summe je parent_col, für jede metric_col)
        parent_totals = (
            out.groupby(parent_col, as_index=False)[metric_cols]
            .sum()
        )
        parent_totals["row_type"] = "parent"

        #Parent in leaf-Spalte übernehmen
        parent_totals[leaf_col] = parent_totals[parent_col]

        # Falls es mehr als 2 Ebenen gibt (sollte es aber nicht!!!): alle tieferen Ebenen leer lassen
        for extra in display_names[2:]:
            parent_totals[extra] = ""

        # Leaf hat evtl. weitere Ebenen -> sicherstellen, dass parent_totals alle Spalten hat
        parent_totals = parent_totals[display_names + metric_cols + ["row_type"]]
        out = out[display_names + metric_cols + ["row_type"]]

        combined = pd.concat([out, parent_totals], ignore_index=True)

        #Hier alles ein DataFrame aber unsortiert
        # Sortierung
        s = cfg.get("sort", {}) or {}
        top_level_by = s.get("top_level_by", "ARR2")          # Standard: ARR2
        top_level_desc = bool(s.get("top_level_desc", True))
        leaf_within_parent = s.get("leaf_within_parent", "ARR2")
        leaf_desc = bool(s.get("leaf_desc", True))

        # 1) Parent sort order nach ARR2 (oder Name)
        parent_rank = parent_totals[[parent_col] + metric_cols].copy()
        if top_level_by == "name":
            parent_rank["_parent_sort"] = parent_rank[parent_col].astype(str)
            parent_rank = parent_rank.sort_values("_parent_sort", ascending=not top_level_desc)
        else:
            # default ARR2
            parent_rank["_parent_sort"] = pd.to_numeric(parent_rank[fy_right], errors="coerce").fillna(0.0)
            parent_rank = parent_rank.sort_values("_parent_sort", ascending=not top_level_desc)

        parent_order = parent_rank[parent_col].tolist()
        order_map = {p: i for i, p in enumerate(parent_order)}
        combined["_parent_order"] = combined[parent_col].map(order_map).fillna(10**9)

        # 2) Leaf sort innerhalb Parent
        if leaf_within_parent == "name":
            combined["_leaf_sort"] = combined[leaf_col].astype(str)
            leaf_asc = True  # alphabetisch
        else:
            combined["_leaf_sort"] = pd.to_numeric(combined[fy_right], errors="coerce").fillna(0.0)
            leaf_asc = not leaf_desc  # bei desc -> asc=False

        # Parent-Zeilen sollen innerhalb Parent ganz unten stehen
        type_order = {"leaf": 0, "parent": 1}
        combined["_type_order"] = combined["row_type"].map(type_order).fillna(99)

        combined = combined.sort_values(
            by=["_parent_order", "_type_order", "_leaf_sort", leaf_col],
            ascending=[True, True, leaf_asc, True]
        ).reset_index(drop=True)

        out = combined.drop(columns=["_parent_order", "_type_order", "_leaf_sort"], errors="ignore")

    else:
        #Kein Parent: einfach sortieren
        s = cfg.get("sort", {}) or {}
        col0 = display_names[0]

        if s.get("top_level_by", "ARR2") == "name":
            out = out.sort_values(col0, ascending=True).reset_index(drop=True)
        else:
            out["_sort"] = pd.to_numeric(out[fy_right], errors="coerce").fillna(0.0)
            out = out.sort_values("_sort", ascending=not bool(s.get("top_level_desc", True))).drop(columns="_sort").reset_index(drop=True)

        out = out[display_names + metric_cols + ["row_type"]]

    #Total-Zeile
    total = {c: "" for c in out.columns}
    total[display_names[0]] = cfg.get("total_label", "Total")
    total["row_type"] = "total"
    for c in metric_cols:
        total[c] = pd.to_numeric(out.loc[out["row_type"].isin(["leaf", "parent"]), c], errors="coerce").fillna(0.0).sum()

    out = pd.concat([out, pd.DataFrame([total])], ignore_index=True)

    # Wenn wir ein Parent-Layout haben: Parent-Spalte im Excel nicht anzeigen
    # (Parent-Name steht ja in der Leaf-Spalte bei row_type parent/bucket)
    if n_levels >= 2 and cfg.get("hide_parent_col", True):
        parent_col = display_names[0]
        out = out.drop(columns=[parent_col], errors="ignore")

    if "row_type" in out.columns:
        total_mask = out["row_type"] == "total"
        if total_mask.any():
            group_display_cols = [c for c in out.columns if c not in metric_cols + ["row_type"]]
            if group_display_cols:
                out.loc[total_mask, group_display_cols[0]] = cfg.get("total_label", "Total")

    return out


def apply_top_buckets(df: pd.DataFrame, cfg: dict) -> pd.DataFrame:
    cfg = normalize_config(cfg)

    if len(cfg["group_cols"]) != 1:
        return df

    tb = cfg.get("top_bucket", {}) or {}
    if not tb.get("enabled", False):
        return df

    out = df.copy()
    if "row_type" not in out.columns:
        out["row_type"] = "leaf"

    fy_cols = sorted([c for c in out.columns if isinstance(c, str) and c.startswith("FY")])
    other_metrics = ["Upsell", "Downsell", "Cross-sell", "Lost", "NRR", "New"]
    metric_cols = fy_cols + [c for c in other_metrics if c in out.columns]

    if not fy_cols:
        raise ValueError("Top Buckets: Keine FY-Spalten gefunden (erwarte Spalten wie 'FY24A').")

    based_on = tb.get("based_on", "current")
    if based_on == "current":
        metric = fy_cols[-1]
    elif based_on == "previous":
        metric = fy_cols[0]
    else:
        metric = based_on

    if metric not in out.columns:
        raise KeyError(f"Top Buckets: based_on='{based_on}' ergibt metric='{metric}', aber diese Spalte existiert nicht.")

    bucket_mode = str(tb.get("bucket_mode", "threshold")).strip().lower()
    if bucket_mode in {"number", "numbers"}:
        bucket_mode = "number"
    else:
        bucket_mode = "threshold"

    thresholds = tuple(tb.get("thresholds", ()) or (0.2, 0.5, 0.8))
    numbers = tuple(int(x) for x in (tb.get("numbers", ()) or ()))
    create_other_bucket = bool(tb.get("create_other_bucket", True))
    other_bucket_label = str(tb.get("other_bucket_label", "Other")).strip() or "Other"

    total_row = out[out["row_type"] == "total"].copy()
    out = out[out["row_type"] == "leaf"].copy()

    metric_positions = [i for i, c in enumerate(out.columns) if c == metric]
    metric_pos = metric_positions[-1] if metric_positions else 0
    out = out.assign(
        _sort_val=pd.to_numeric(out.iloc[:, metric_pos], errors="coerce").fillna(0.0)
    )
    out = out.sort_values("_sort_val", ascending=False).reset_index(drop=True)
    out = out.drop(columns=["_sort_val"])

    metric_vals = pd.to_numeric(out.iloc[:, metric_pos], errors="coerce").fillna(0.0)
    total_value = float(metric_vals.sum())
    if total_value == 0:
        return df

    out["_share"] = metric_vals / total_value
    out["_cum_share"] = out["_share"].cumsum()

    group_col = cfg["group_cols"][0]

    def label_for_range(start: int, end: int) -> str:
        a, b = start + 1, end
        if a == b:
            return f"Top {a}"
        if a == 1:
            return f"Top {b}"
        return f"Top {a}-{b}"

    use_formulas = bool(cfg.get("formula_mode", True))
    numeric_positions = _churn_numeric_col_positions(out)

    def make_bucket_row(label: str, subset: pd.DataFrame) -> pd.DataFrame:
        row_data: list = []
        for i, c in enumerate(out.columns):
            if c == group_col:
                row_data.append(label)
            elif c == "row_type":
                row_data.append("bucket")
            elif i in numeric_positions:
                if use_formulas:
                    row_data.append("")
                else:
                    row_data.append(
                        pd.to_numeric(subset.iloc[:, i], errors="coerce").fillna(0.0).sum()
                    )
            else:
                row_data.append("")
        return pd.DataFrame([row_data], columns=out.columns)

    bucket_ranges: list[tuple[int, int, str]] = []
    if bucket_mode == "threshold":
        prev_end = 0
        for th in thresholds:
            idx = out.index[out["_cum_share"] >= float(th)]
            end = int(idx[0]) + 1 if len(idx) > 0 else len(out)
            end = max(end, prev_end)
            end = min(end, len(out))
            if end > prev_end:
                bucket_ranges.append((prev_end, end, label_for_range(prev_end, end)))
                prev_end = end
        if create_other_bucket and prev_end < len(out):
            bucket_ranges.append((prev_end, len(out), other_bucket_label))
    else:
        start = 0
        for n in numbers:
            end = min(start + int(n), len(out))
            if end > start:
                bucket_ranges.append((start, end, label_for_range(start, end)))
            start = end
            if start >= len(out):
                break
        if create_other_bucket and start < len(out):
            bucket_ranges.append((start, len(out), other_bucket_label))

    if not bucket_ranges:
        return df

    pieces: list[pd.DataFrame] = []
    for start, end, label in bucket_ranges:
        seg = out.iloc[start:end].copy()
        pieces.append(seg)
        pieces.append(make_bucket_row(label, seg))

    out = pd.concat(pieces, ignore_index=True)
    out = out.drop(columns=["_share", "_cum_share"], errors="ignore")

    if not total_row.empty:
        out = pd.concat([out, total_row], ignore_index=True)

    return out


def apply_kEUR_formatting(results: dict) -> dict:    #kEUR Format anwenden
    def to_kEUR(series):
        return (
            pd.to_numeric(series, errors="coerce")
            .div(1000)
            .round(0)
        )

    formatted = {}

    for label, tbl in results.items():
        tbl = tbl.copy()

        # dynamisch FY-Spalten erkennen
        fy_cols = [c for c in tbl.columns if c.startswith("FY")]

        metric_cols = fy_cols + [
            "Upsell",
            "Downsell",
            "Cross-sell",
            "Lost",
            "New",
            "NRR"
        ]

        for c in metric_cols:
            if c in tbl.columns:
                tbl[c] = to_kEUR(tbl[c])

        formatted[label] = tbl

    return formatted


CHURN_BRIDGE_METRICS = ["Upsell", "Downsell", "Cross-sell", "Lost", "NRR", "New"]
CHURN_META_COLS = frozenset({"row_type", "level", "ui_level", "key", "parent_key", "label"})

_SCRIPTS_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "scripts")
if _SCRIPTS_DIR not in sys.path:
    sys.path.insert(0, _SCRIPTS_DIR)
from gst_excel_theme import THEME  # noqa: E402

from openpyxl.styles import Alignment, Border, Font, Side  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402
from openpyxl.utils.dataframe import dataframe_to_rows  # noqa: E402


def _churn_group_display_cols(tbl: pd.DataFrame) -> list[str]:
    skip = set(CHURN_BRIDGE_METRICS) | set(CHURN_META_COLS)
    return [c for c in tbl.columns if c not in skip and not str(c).startswith("FY")]


def split_key_to_n_parts(key: str, n: int) -> list[str | None]:
    raw = [p.strip() for p in str(key or "").split(" | ")]
    return [
        raw[i] if i < len(raw) and raw[i] != "" else None
        for i in range(n)
    ]


def _churn_hierarchy_key(parts: list[str], level: int) -> str:
    return " | ".join(str(p) for p in parts[:level])


def _churn_hierarchy_parent_key(parts: list[str], level: int) -> str:
    if level <= 1:
        return ""
    return " | ".join(str(p) for p in parts[: level - 1])


def _churn_export_metric_columns(columns: list, numeric_positions: list[int]) -> list:
    return [columns[i] for i in numeric_positions]


def build_churn_hierarchy_rows(leaves: pd.DataFrame, cfg: dict) -> pd.DataFrame:
    """Build GST-style hierarchy rows (all levels) from deepest-level leaf bridge data."""
    display_names = list(cfg["group_cols"])
    n_levels = len(display_names)
    if n_levels < 2:
        raise ValueError("build_churn_hierarchy_rows requires at least two group_cols.")

    columns = list(leaves.columns)
    numeric_positions = _churn_numeric_col_positions(leaves)
    metric_headers = _churn_export_metric_columns(columns, numeric_positions)
    meta_headers = ["level", "ui_level", "key", "parent_key", "row_type"]
    export_headers = ["label"] + metric_headers + meta_headers

    def metric_values_from_group(grp: pd.DataFrame) -> list[float]:
        return [
            float(pd.to_numeric(grp.iloc[:, i], errors="coerce").fillna(0.0).sum())
            for i in numeric_positions
        ]

    rows: list[list] = []

    for level in range(1, n_levels + 1):
        gcols = display_names[:level]
        if level == n_levels:
            for _, src in leaves.iterrows():
                parts = [str(src[c]) for c in display_names]
                metrics = [float(pd.to_numeric(src.iloc[i], errors="coerce") or 0.0) for i in numeric_positions]
                rows.append(
                    [
                        parts[-1],
                        *metrics,
                        level,
                        level,
                        _churn_hierarchy_key(parts, level),
                        _churn_hierarchy_parent_key(parts, level),
                        "leaf",
                    ]
                )
        else:
            for keys, grp in leaves.groupby(gcols, dropna=False):
                key_tuple = keys if isinstance(keys, tuple) else (keys,)
                parts = [str(k) for k in key_tuple]
                metrics = metric_values_from_group(grp)
                rows.append(
                    [
                        parts[-1],
                        *metrics,
                        level,
                        level,
                        _churn_hierarchy_key(parts, level),
                        _churn_hierarchy_parent_key(parts, level),
                        "hierarchy",
                    ]
                )

    return pd.DataFrame(rows, columns=export_headers)


def build_churn_display_order(df: pd.DataFrame, cfg: dict) -> pd.DataFrame:
    """Post-order walk: children first, then parent row (GST build_display_order)."""
    out = df.copy()
    s = cfg.get("sort", {}) or {}
    sort_by_name = s.get("top_level_by", "ARR2") == "name"

    fy_cols = sorted(c for c in out.columns if isinstance(c, str) and c.startswith("FY"))
    sort_col = fy_cols[-1] if fy_cols else None
    if sort_col and sort_col in out.columns:
        pos = [i for i, c in enumerate(out.columns) if c == sort_col][-1]
        out["_sort_value"] = pd.to_numeric(out.iloc[:, pos], errors="coerce").fillna(0.0)
    else:
        out["_sort_value"] = 0.0

    def sort_block(block: pd.DataFrame) -> pd.DataFrame:
        if block.empty:
            return block
        tmp = block.copy()
        if sort_by_name:
            return tmp.sort_values(["label"], ascending=True, kind="mergesort")
        return tmp.sort_values(["_sort_value", "label"], ascending=[False, True], kind="mergesort")

    order_rows: list[tuple[str, int]] = []
    visited: set[str] = set()
    pos = 0

    def walk(parent_key: str) -> None:
        nonlocal pos
        block = out[out["parent_key"].astype(str) == str(parent_key)].copy()
        block = block[~block["key"].astype(str).isin(visited)]
        block = sort_block(block)
        for _, row in block.iterrows():
            key = str(row["key"])
            walk(key)
            if key not in visited:
                pos += 1
                order_rows.append((key, pos))
                visited.add(key)

    walk("")

    leftovers = out[~out["key"].astype(str).isin(visited)].copy()
    leftovers = sort_block(leftovers)
    for _, row in leftovers.iterrows():
        pos += 1
        order_rows.append((str(row["key"]), pos))

    order_df = pd.DataFrame(order_rows, columns=["key", "_display_order"])
    out = out.merge(order_df, on="key", how="left")
    out = out.sort_values("_display_order", kind="mergesort").reset_index(drop=True)
    return out.drop(columns=["_sort_value", "_display_order"], errors="ignore")


def _churn_is_numeric_col(name) -> bool:
    return (isinstance(name, str) and name.startswith("FY")) or name in CHURN_BRIDGE_METRICS


def _churn_numeric_col_positions(df: pd.DataFrame) -> list[int]:
    return [i for i, c in enumerate(df.columns) if _churn_is_numeric_col(c)]


def _rightmost_fy_col(df: pd.DataFrame) -> str | None:
    fy_cols = sorted(c for c in df.columns if isinstance(c, str) and c.startswith("FY"))
    return fy_cols[-1] if fy_cols else None


def _rightmost_fy_col_position(df: pd.DataFrame) -> int | None:
    fy_name = _rightmost_fy_col(df)
    if fy_name is None:
        return None
    positions = [i for i, c in enumerate(df.columns) if c == fy_name]
    return positions[-1] if positions else None


def _rename_bridge_groups(bridge: pd.DataFrame, cfg: dict) -> pd.DataFrame:
    group_keys = sorted(
        [c for c in bridge.columns if c.startswith("group_")],
        key=lambda x: int(x.split("_")[1]),
    )
    display_names = cfg["group_cols"][:]
    rename_map = {group_keys[i]: display_names[i] for i in range(len(group_keys))}
    return bridge.rename(columns=rename_map)


def finalize_merged_bridge_layout(bridge: pd.DataFrame, cfg: dict) -> pd.DataFrame:
    """Sort, GST hierarchy (multi-level), and total row on the horizontally merged bridge."""
    cfg = normalize_config(cfg)
    display_names = list(cfg["group_cols"])
    n_levels = len(display_names)
    out = bridge.copy()

    fy_right_pos = _rightmost_fy_col_position(out)
    numeric_positions = _churn_numeric_col_positions(out)
    columns = list(out.columns)
    s = cfg.get("sort", {}) or {}

    if n_levels >= 2:
        leaves = out.copy()
        if s.get("top_level_by", "ARR2") == "name":
            leaves = leaves.sort_values(display_names[-1], ascending=True).reset_index(drop=True)
        elif fy_right_pos is not None:
            leaves = leaves.assign(
                _sort=pd.to_numeric(leaves.iloc[:, fy_right_pos], errors="coerce").fillna(0.0)
            )
            leaves = leaves.sort_values(
                "_sort", ascending=not bool(s.get("top_level_desc", True))
            ).drop(columns="_sort").reset_index(drop=True)

        hier = build_churn_hierarchy_rows(leaves, cfg)
        hier = build_churn_display_order(hier, cfg)

        deepest = hier[hier["row_type"] == "leaf"]
        n_metrics = len(_churn_export_metric_columns(columns, numeric_positions))
        total_data = [
            float(pd.to_numeric(deepest.iloc[:, i], errors="coerce").fillna(0.0).sum())
            for i in range(1, 1 + n_metrics)
        ]
        total_row = (
            [cfg.get("total_label", "Total")]
            + total_data
            + [0, 0, "__TOTAL__", "", "total"]
        )
        hier = pd.concat([hier, pd.DataFrame([total_row], columns=hier.columns)], ignore_index=True)
        return hier

    out["row_type"] = "leaf"
    col0 = display_names[0]
    if s.get("top_level_by", "ARR2") == "name":
        out = out.sort_values(col0, ascending=True).reset_index(drop=True)
    elif fy_right_pos is not None:
        out = out.assign(
            _sort=pd.to_numeric(out.iloc[:, fy_right_pos], errors="coerce").fillna(0.0)
        )
        out = out.sort_values(
            "_sort", ascending=not bool(s.get("top_level_desc", True))
        ).drop(columns="_sort").reset_index(drop=True)

    label_col = col0
    leaf_mask = out["row_type"] == "leaf"
    total_data: list = []
    for i, c in enumerate(out.columns):
        if c == "row_type":
            total_data.append("total")
        elif c == label_col:
            total_data.append(cfg.get("total_label", "Total"))
        elif i in numeric_positions:
            total_data.append(
                pd.to_numeric(out.iloc[leaf_mask.values, i], errors="coerce").fillna(0.0).sum()
            )
        else:
            total_data.append("")
    out = pd.concat([out, pd.DataFrame([total_data], columns=out.columns)], ignore_index=True)
    out.loc[out["row_type"] == "total", label_col] = cfg.get("total_label", "Total")
    return out


def _churn_index_key(values, n_levels: int):
    if n_levels == 1:
        return values[0]
    return tuple(values)


def build_merged_horizontal_bridge(df: pd.DataFrame, cfg: dict) -> pd.DataFrame:
    """One wide table: FY22A | bridge metrics | FY23A | metrics | FY24A | …"""
    tables = list(apply_kEUR_formatting(build_fy_bridge_tables(df, cfg)).values())
    if not tables:
        return pd.DataFrame()

    group_cols = _churn_group_display_cols(tables[0])
    idx_cols = group_cols
    n_idx = len(idx_cols)

    merged = tables[0].set_index(idx_cols)
    row_order = [
        _churn_index_key(tuple(x), n_idx)
        for x in tables[0][idx_cols].to_numpy()
    ]

    for tbl in tables[1:]:
        piece = tbl.set_index(idx_cols)
        fy_cols = sorted(c for c in piece.columns if str(c).startswith("FY"))
        if fy_cols:
            piece = piece.drop(columns=[fy_cols[0]], errors="ignore")
        merged = pd.concat([merged, piece], axis=1)
        for t in tbl[idx_cols].itertuples(index=False, name=None):
            key = _churn_index_key(t if isinstance(t, tuple) else (t,), n_idx)
            if key not in row_order:
                row_order.append(key)

    merged = merged.reindex(row_order)
    for col_idx in range(merged.shape[1]):
        c = merged.columns[col_idx]
        if str(c).startswith("FY") or c in CHURN_BRIDGE_METRICS:
            merged.iloc[:, col_idx] = pd.to_numeric(merged.iloc[:, col_idx], errors="coerce").fillna(0.0)

    merged = merged.reset_index()
    merged = finalize_merged_bridge_layout(merged, cfg)
    merged = apply_top_buckets(merged, cfg)

    if "row_type" in merged.columns:
        non_total = merged[merged["row_type"] != "total"]
        total = merged[merged["row_type"] == "total"]
        merged = pd.concat([non_total, total], ignore_index=True)

    return merged


def strip_churn_fy_values(export_df: pd.DataFrame) -> pd.DataFrame:
    out = export_df.copy()
    for c in out.columns:
        if isinstance(c, str) and c.startswith("FY"):
            out[c] = ""
    return out


def get_churn_excel_layout(cfg: dict, *, formula_mode: bool = True) -> dict[str, Any]:
    n_groups = len(cfg.get("group_cols") or [])
    col_offset = max(int(cfg.get("excel_formatting", {}).get("col_offset", 3)), n_groups, 3)
    header_row = 4
    return {
        "formula_mode": formula_mode,
        "col_offset": col_offset,
        "title_row": 1,
        "subtitle_row": 2,
        "company_row": 3,
        "header_row": header_row,
        "data_start_row": header_row + 1,
        "insert_rows": header_row - 1,
        "label_col": col_offset + 1,
        "key_cols": list(range(1, min(len(cfg.get("group_cols") or []), col_offset) + 1)),
        "helper_right_col": col_offset,
    }


def _churn_header_map(ws, header_row: int) -> dict[str, int]:
    m: dict[str, int] = {}
    for c in range(1, (ws.max_column or 0) + 1):
        v = ws.cell(row=header_row, column=c).value
        if v is not None and str(v).strip():
            m[str(v).strip()] = c
    return m


def get_churn_excel_layout_from_ws(ws, cfg: dict, *, formula_mode: bool = True) -> dict[str, Any]:
    layout = get_churn_excel_layout(cfg, formula_mode=formula_mode)
    header_row = layout["header_row"]
    header_map = _churn_header_map(ws, header_row)
    row_type_col = header_map.get("row_type", ws.max_column)
    fy_col_indices = [(h, header_map[h]) for h in header_map if h.startswith("FY")]
    layout.update({
        "row_type_col": row_type_col,
        "table_bottom_row": ws.max_row,
        "header_map": header_map,
        "fy_col_indices": fy_col_indices,
    })
    return layout


def write_churn_export_to_workbook(
    wb,
    export_df: pd.DataFrame,
    sheet_name: str,
    cfg: dict,
    *,
    formula_mode: bool = False,
) -> None:
    if sheet_name in wb.sheetnames:
        del wb[sheet_name]
    ws = wb.create_sheet(sheet_name)
    write_df = strip_churn_fy_values(export_df) if formula_mode else export_df
    for row in dataframe_to_rows(write_df, index=False, header=True):
        ws.append(row)


def _churn_col_header_map(ws, header_row: int) -> dict[int, str]:
    m: dict[int, str] = {}
    for c in range(1, (ws.max_column or 0) + 1):
        v = ws.cell(row=header_row, column=c).value
        m[c] = str(v).strip() if v is not None else ""
    return m


def _is_churn_money_header(header: str) -> bool:
    if not header:
        return False
    if header.startswith("FY") or header.startswith("YTD") or header.startswith("LTM"):
        return True
    return header in CHURN_BRIDGE_METRICS


def _churn_table_bottom_from_row_type(ws, row_type_col: int, data_start: int) -> int:
    last = data_start - 1
    for r in range(data_start, (ws.max_row or 0) + 1):
        rt = str(ws.cell(row=r, column=row_type_col).value or "").strip().lower()
        if rt in {"leaf", "parent", "hierarchy", "bucket", "total"}:
            last = r
    return last


def _trim_sheet_below_total(ws, row_type_col: int, data_start: int) -> int:
    bottom = _churn_table_bottom_from_row_type(ws, row_type_col, data_start)
    if bottom < (ws.max_row or 0):
        ws.delete_rows(bottom + 1, ws.max_row - bottom)
    return bottom


def apply_churn_vertical_outlines(ws, layout: dict) -> None:
    """Row groups: GST ui_level hierarchy or bucket blocks."""
    data_start = layout["data_start_row"]
    table_bottom = layout["table_bottom_row"]
    row_type_col = layout["row_type_col"]
    ui_level_col = layout.get("ui_level_col")
    max_level = layout.get("max_hierarchy_level")

    ws.sheet_properties.outlinePr.summaryBelow = True

    has_buckets = False
    for r in range(data_start, table_bottom + 1):
        rt = str(ws.cell(row=r, column=row_type_col).value or "").strip().lower()
        if rt == "bucket":
            has_buckets = True
            break

    for r in range(data_start, table_bottom + 1):
        rt = str(ws.cell(row=r, column=row_type_col).value or "").strip().lower()
        rd = ws.row_dimensions[r]

        if rt in {"total", "bucket"}:
            rd.outlineLevel = 0
            rd.hidden = False
            rd.collapsed = False
            continue

        if ui_level_col and not has_buckets:
            try:
                ui_level = int(ws.cell(row=r, column=ui_level_col).value)
            except (TypeError, ValueError):
                ui_level = max_level or 1
            rd.outlineLevel = max(0, min(7, ui_level - 1))
            rd.hidden = False
            rd.collapsed = False
            continue

        if rt == "hierarchy":
            rd.outlineLevel = 1
            rd.hidden = False
            rd.collapsed = True
        elif rt == "parent":
            rd.outlineLevel = 1
            rd.hidden = False
            rd.collapsed = True
        elif rt == "leaf":
            rd.outlineLevel = 2 if has_buckets else 1
            rd.hidden = False
            rd.collapsed = False


def write_churn_formula_key_columns(ws, cfg: dict, layout: dict) -> None:
    red_font = Font(name=THEME.font_name, size=THEME.font_size, color="FFFF5149")
    red_bold_font = Font(name=THEME.font_name, size=THEME.font_size, color="FFFF5149", bold=True)
    header_row = layout["header_row"]
    data_start = layout["data_start_row"]
    table_bottom = layout["table_bottom_row"]
    row_type_col = layout["row_type_col"]
    key_col = layout.get("key_col")
    level_col = layout.get("level_col")
    group_cols = list(cfg.get("group_cols") or [])
    n_keys = len(group_cols)
    col_offset = layout["helper_right_col"]

    for c in range(1, col_offset + 1):
        ws.cell(row=header_row, column=c).value = group_cols[c - 1] if c <= n_keys else None
        if c <= n_keys:
            ws.cell(row=header_row, column=c).font = red_bold_font
        ws.column_dimensions[get_column_letter(c)].hidden = True
        ws.column_dimensions[get_column_letter(c)].outlineLevel = 1

    max_level = layout.get("max_hierarchy_level") or n_keys

    for r in range(data_start, table_bottom + 1):
        rt = str(ws.cell(row=r, column=row_type_col).value or "").strip().lower()
        for c in range(1, col_offset + 1):
            ws.cell(row=r, column=c).font = red_font
            ws.cell(row=r, column=c).alignment = Alignment(horizontal="left", vertical="center")
            ws.cell(row=r, column=c).value = None

        if rt in {"total", "hierarchy", "bucket"} and n_keys >= 2:
            continue
        if rt == "parent" and n_keys >= 2:
            continue

        if n_keys == 1:
            ws.cell(row=r, column=1).value = ws.cell(row=r, column=layout["label_col"]).value
        elif key_col and level_col:
            try:
                row_level = int(ws.cell(row=r, column=level_col).value)
            except (TypeError, ValueError):
                row_level = 0
            if row_level == max_level and rt == "leaf":
                raw_key = ws.cell(row=r, column=key_col).value
                parts = split_key_to_n_parts(str(raw_key or ""), n_keys)
                for i, val in enumerate(parts, start=1):
                    ws.cell(row=r, column=i).value = val
        elif n_keys >= 2:
            ws.cell(row=r, column=1).value = ws.cell(row=r, column=layout["label_col"]).value


def format_churn_excel(
    wb,
    target_sheet_name: str,
    export_df: pd.DataFrame,
    cfg: dict,
    *,
    formula_mode: bool = False,
) -> dict[str, Any]:
    if target_sheet_name not in wb.sheetnames:
        raise ValueError(f"Sheet '{target_sheet_name}' nicht im Workbook gefunden.")
    ws = wb[target_sheet_name]
    layout = get_churn_excel_layout(cfg, formula_mode=formula_mode)
    col_offset = layout["col_offset"]
    header_row = layout["header_row"]
    data_start = layout["data_start_row"]

    ws.insert_rows(1, amount=layout["insert_rows"])
    ws.insert_cols(1, amount=col_offset)

    table_left = layout["label_col"]
    header_map = _churn_header_map(ws, header_row)
    row_type_col = header_map.get("row_type", ws.max_column)
    table_bottom = _trim_sheet_below_total(ws, row_type_col, data_start)

    fmt_money = '#,##0;(#,##0);"-"'
    fill_tech = THEME.fill_tech
    white_fill = THEME.fill_white
    header_fill = THEME.fill_header
    period_fill = THEME.fill_period
    subtotal_fill = THEME.fill_subtotal
    border_subtotal = THEME.border_subtotal_top
    border_subtotal_tb = Border(
        top=THEME.border_subtotal_top.top,
        bottom=THEME.border_subtotal_top.top,
    )
    fill_extent = 501

    for r in range(1, fill_extent + 1):
        for c in range(1, fill_extent + 1):
            ws.cell(row=r, column=c).font = THEME.font_base

    for r in range(1, fill_extent + 1):
        for c in range(col_offset + 1, fill_extent + 1):
            ws.cell(row=r, column=c).fill = white_fill
    for r in range(1, fill_extent + 1):
        for c in range(1, col_offset + 1):
            ws.cell(row=r, column=c).fill = fill_tech

    ws.sheet_properties.outlinePr.summaryRight = True
    for c in range(1, col_offset + 1):
        letter = get_column_letter(c)
        ws.column_dimensions[letter].outlineLevel = 1
        ws.column_dimensions[letter].hidden = True
    ws.column_dimensions[get_column_letter(col_offset + 1)].collapsed = True

    title_font = Font(name=THEME.font_name, size=THEME.font_size_title, color=THEME.text_brand_title)
    subtitle_font = Font(name=THEME.font_name, size=THEME.font_size_subtitle, color=THEME.text_brand_title)
    company_font = Font(name=THEME.font_name, size=THEME.font_size_company, color=THEME.text_brand_title)
    ws.cell(row=1, column=col_offset + 1).value = cfg.get("title", "")
    ws.cell(row=1, column=col_offset + 1).font = title_font
    ws.cell(row=2, column=col_offset + 1).value = cfg.get("table", "")
    ws.cell(row=2, column=col_offset + 1).font = subtitle_font
    ws.cell(row=3, column=col_offset + 1).value = cfg.get("company", "")
    ws.cell(row=3, column=col_offset + 1).font = company_font

    col_headers = _churn_col_header_map(ws, header_row)
    fy_col_indices = [(h, c) for c, h in sorted(col_headers.items()) if h.startswith("FY")]
    fy_col_set = {c for _, c in fy_col_indices}
    bridge_metric_cols = [c for c, h in col_headers.items() if h in CHURN_BRIDGE_METRICS]
    money_cols = [c for c, h in col_headers.items() if _is_churn_money_header(h)]
    table_right = max(
        (c for c in col_headers if c > col_offset and c != row_type_col),
        default=table_left,
    )

    meta_cols = {}
    for name in ("level", "ui_level", "key", "parent_key"):
        if name in header_map:
            meta_cols[name] = header_map[name]
            ws.column_dimensions[get_column_letter(meta_cols[name])].hidden = True
            ws.cell(row=header_row, column=meta_cols[name]).value = ""

    max_hierarchy_level = None
    if "level" in meta_cols and len(cfg.get("group_cols") or []) >= 2:
        try:
            max_hierarchy_level = int(pd.to_numeric(export_df["level"], errors="coerce").max())
        except (TypeError, ValueError, KeyError):
            max_hierarchy_level = len(cfg.get("group_cols") or [])

    ws.column_dimensions[get_column_letter(row_type_col)].hidden = True
    ws.cell(row=header_row, column=row_type_col).value = ""
    ws.cell(row=header_row, column=table_left).value = "kEUR"

    ws.row_dimensions[header_row].height = 24
    for c in range(table_left, table_right + 1):
        cell = ws.cell(row=header_row, column=c)
        cell.font = THEME.font_header
        cell.fill = period_fill if c in fy_col_set else header_fill
        cell.alignment = Alignment(
            horizontal="left" if c == table_left else "right",
            vertical="bottom",
            wrap_text=True,
        )
        cell.border = THEME.border_header_bottom

    group_cols_count = len(cfg.get("group_cols") or [])
    indent_col = table_left
    n_levels = group_cols_count

    for r in range(data_start, table_bottom + 1):
        ws.row_dimensions[r].height = 12
        rt = str(ws.cell(row=r, column=row_type_col).value or "").strip().lower()

        row_level = None
        if meta_cols.get("level"):
            try:
                row_level = int(ws.cell(row=r, column=meta_cols["level"]).value)
            except (TypeError, ValueError):
                row_level = None

        for c in money_cols:
            if c == row_type_col:
                continue
            ws.cell(row=r, column=c).number_format = fmt_money

        label_cell = ws.cell(row=r, column=indent_col)
        val = label_cell.value
        if isinstance(val, str) and val.strip():
            label_cell.value = val.strip().lower().capitalize() if len(val.strip()) > 1 else val.strip().upper()

        indent = 0
        if n_levels >= 3 and row_level is not None:
            if row_level == n_levels:
                indent = 0
            elif row_level == 1:
                indent = 0
            else:
                indent = 1
        elif n_levels == 2 and row_level == 2:
            indent = 1

        if rt == "total":
            for c in range(table_left, table_right + 1):
                cell = ws.cell(row=r, column=c)
                cell.font = THEME.font_bold
                if c in fy_col_set:
                    cell.fill = subtotal_fill
                elif c in money_cols:
                    cell.fill = white_fill
                cell.border = border_subtotal_tb
        elif rt == "bucket":
            for c in range(table_left, table_right + 1):
                cell = ws.cell(row=r, column=c)
                cell.font = THEME.font_bold
                cell.fill = subtotal_fill if c in fy_col_set else white_fill
                cell.border = border_subtotal
        elif rt in {"hierarchy", "parent"} and row_level == 1:
            for c in range(table_left, table_right + 1):
                cell = ws.cell(row=r, column=c)
                cell.font = THEME.font_bold
                if c in fy_col_set:
                    cell.fill = subtotal_fill
                elif c in money_cols:
                    cell.fill = white_fill
                cell.border = border_subtotal
        elif rt in {"hierarchy", "parent"}:
            for c in fy_col_set:
                if table_left <= c <= table_right:
                    ws.cell(row=r, column=c).fill = period_fill
            if indent > 0:
                label_cell.alignment = Alignment(horizontal="left", vertical="center", indent=indent)
        elif rt == "leaf":
            if indent > 0:
                label_cell.alignment = Alignment(horizontal="left", vertical="center", indent=indent)
            else:
                label_cell.alignment = Alignment(horizontal="left", vertical="center")
            for c in fy_col_set:
                if table_left <= c <= table_right:
                    ws.cell(row=r, column=c).fill = period_fill

    apply_churn_vertical_outlines(ws, {
        "data_start_row": data_start,
        "table_bottom_row": table_bottom,
        "row_type_col": row_type_col,
        "ui_level_col": meta_cols.get("ui_level"),
        "max_hierarchy_level": max_hierarchy_level,
    })

    layout.update({
        "row_type_col": row_type_col,
        "table_bottom_row": table_bottom,
        "table_left_col": table_left,
        "table_right_col": table_right,
        "label_col": table_left,
        "level_col": meta_cols.get("level"),
        "key_col": meta_cols.get("key"),
        "ui_level_col": meta_cols.get("ui_level"),
        "parent_key_col": meta_cols.get("parent_key"),
        "max_hierarchy_level": max_hierarchy_level,
        "header_map": _churn_header_map(ws, header_row),
        "col_headers": col_headers,
        "fy_col_indices": fy_col_indices,
        "bridge_metric_col_indices": bridge_metric_cols,
        "money_col_indices": money_cols,
    })

    if formula_mode:
        write_churn_formula_key_columns(ws, cfg, layout)

    return layout


def fy_end_date(year: int, cfg: dict) -> pd.Timestamp:    #Holt Enddaten aus CONFIG
    return pd.Timestamp(
        year=year,
        month=cfg["fy_end_month"],
        day=cfg["fy_end_day"]
    )


def build_fy_bridge_tables(df: pd.DataFrame, cfg: dict) -> dict:
    cfg = normalize_config(cfg)
    d = preprocess_contract_lines(df, cfg)
    _, latest_closed = resolve_fy_years(cfg)
    cfg = dict(cfg)
    cfg["fy_end_year"] = latest_closed

    results: dict[str, pd.DataFrame] = {}

    for y1, y2 in bridge_year_pairs(cfg):
        as_of_1 = fy_end_date(y1, cfg)
        as_of_2 = fy_end_date(y2, cfg)

        snap1 = snapshot_arr(d, as_of_1, cfg)
        snap2 = snapshot_arr(d, as_of_2, cfg)
        bridge = compute_bridge(snap1, snap2)

        bridge = bridge.rename(columns={
            "ARR1": f"FY{str(y1)[-2:]}A",
            "ARR2": f"FY{str(y2)[-2:]}A",
        })

        bridge = _rename_bridge_groups(bridge, cfg)

        label = f"{cfg.get('company', 'Company')} | ARR Bridge FY{str(y1)[-2:]}A to FY{str(y2)[-2:]}A"
        results[label] = bridge

    return results


def build_last_two_fy_tables(df: pd.DataFrame, cfg: dict) -> dict:
    """Backward-compatible alias."""
    return build_fy_bridge_tables(df, cfg)


def build_output_path(output_dir: str, cfg: dict, prefix: str = "churn") -> str:
    group_cols = cfg["group_cols"]
    group_part = "_".join(group_cols)
    group_part = re.sub(r"[^\w_]", "", group_part.replace(" ", "_").lower())
    first_fy, latest = resolve_fy_years(cfg)
    filename = f"{prefix}_{group_part}_fy{str(first_fy)[-2:]}_fy{str(latest)[-2:]}.xlsx"
    return os.path.join(output_dir, filename)


GAP_ROWS = 3
TITLE_ROWS = 1


def write_churn_tables_to_workbook(wb, tables: dict, *, sheet_name: str = "Churn") -> None:
    from openpyxl.utils.dataframe import dataframe_to_rows

    if sheet_name in wb.sheetnames:
        del wb[sheet_name]
    ws = wb.create_sheet(sheet_name)
    start_row = 0
    for label, tbl in tables.items():
        export_tbl = tbl.drop(columns=["row_type"], errors="ignore")
        ws.cell(row=start_row + 1, column=1, value=label)
        for r_idx, row in enumerate(dataframe_to_rows(export_tbl, index=False, header=True), start=start_row + 2):
            for c_idx, value in enumerate(row, start=1):
                ws.cell(row=r_idx, column=c_idx, value=value)
        start_row += TITLE_ROWS + (len(export_tbl) + 1) + GAP_ROWS


def write_churn_tables_to_excel(output_path: str, tables: dict, *, sheet_name: str = "Churn") -> None:
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        start_row = 0
        for label, tbl in tables.items():
            export_tbl = tbl.drop(columns=["row_type"], errors="ignore")
            pd.DataFrame({export_tbl.columns[0]: [label]}).to_excel(
                writer,
                sheet_name=sheet_name,
                index=False,
                header=False,
                startrow=start_row,
            )
            export_tbl.to_excel(
                writer,
                sheet_name=sheet_name,
                index=False,
                startrow=start_row + 1,
            )
            start_row += TITLE_ROWS + (len(export_tbl) + 1) + GAP_ROWS


from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter


def format_churn_workbook(output_path_churn_universal: str, tables: dict, cfg: dict,
                          GAP_ROWS: int = 3, TITLE_ROWS: int = 1,
                          ROW_OFFSET: int = 3, COL_OFFSET: int = 3):

    #Tabellenpositionen rekonstruieren
    table_items = list(tables.items())
    starts = []
    sr = 0
    for label, tbl in table_items:
        starts.append((label, sr, tbl.shape[0], tbl.shape[1]))  # start_row (0-basiert), nrows, ncols
        sr += TITLE_ROWS + (len(tbl) + 1) + GAP_ROWS  # +1 wegen Headerzeile

    # Workbook laden + Offsets einfügen
    wb = load_workbook(output_path_churn_universal)
    ws = wb.active

    if ws.cell(row=1, column=1 + COL_OFFSET).value in (cfg.get("title", ""), cfg.get("table", "")):
        # schon formatiert -> keine Offsets erneut einfügen
        pass
    else:
        ws.insert_rows(1, amount=ROW_OFFSET)
        ws.insert_cols(1, amount=COL_OFFSET)

    #  Styles
    fmt_money = "#,##0;(#,##0)"
    fmt_pct   = "0.0%"

    base_font = Font(name="GT Walsheim LC Light", size=8)
    bold_font = Font(name="GT Walsheim LC Light", bold=True, size=8)

    header_fill = PatternFill(fill_type="solid", fgColor="FFF2F2F2")
    white_fill  = PatternFill(fill_type="solid", fgColor="FFFFFFFF")
    fill_grey   = PatternFill(fill_type="solid", fgColor="FFF3F1EF")
    arr_fill    = PatternFill(fill_type="solid", fgColor="FFE4DFD7")

    header_alignment = Alignment(wrap_text=True, vertical="bottom")
    thin = Side(style="thin")
    border_top = Border(top=thin)
    border_top_bottom = Border(top=thin, bottom=thin)

    # Basis-Font überall
    for r in range(1, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            ws.cell(row=r, column=c).font = base_font

    # Weiß überall rechts (ab Spalte D nach Offset) -
    for r in range(1, 501):
        for c in range(1, 501):
            if c >= 1 + COL_OFFSET:   # ab "D"
                ws.cell(row=r, column=c).fill = white_fill

    # Linke Spalten A-C grau + einklappen
    for r in range(1, 501):
        for c in range(1, 1 + COL_OFFSET):
            ws.cell(row=r, column=c).fill = fill_grey

    ws.sheet_properties.outlinePr.summaryRight = True
    for c in range(1, 1 + COL_OFFSET):
        letter = get_column_letter(c)
        ws.column_dimensions[letter].outlineLevel = 1
        ws.column_dimensions[letter].hidden = True
    ws.column_dimensions[get_column_letter(1 + COL_OFFSET)].collapsed = True  # D

    # Titel etc.
    titel_cell = ws.cell(row=1, column=1 + COL_OFFSET)
    titel_cell.value = cfg.get("title", "")
    titel_cell.font = Font(name="GT Walsheim LC Light", size=24, color="FF4F2D7F")

    table_cell = ws.cell(row=2, column=1 + COL_OFFSET)
    table_cell.value = cfg.get("table", "")
    table_cell.font = Font(name="GT Walsheim LC Light", size=12, color="FF4F2D7F")

    # For: alles was pro Tabelle formatiert wird
    for label, start_row, nrows, ncols in starts:
        label_row = (start_row + 1) + ROW_OFFSET
        HEADER_ROW = (start_row + 1 + TITLE_ROWS) + ROW_OFFSET
        DATA_START_ROW = HEADER_ROW + 1

        TABLE_LEFT_COL = 1 + COL_OFFSET
        TABLE_RIGHT_COL = COL_OFFSET + ncols
        TABLE_BOTTOM_ROW = DATA_START_ROW + nrows - 1
        TOTAL_ROW = TABLE_BOTTOM_ROW

        # row_type ist letzte Spalte in der Tabelle
        ROWTYPE_COL = TABLE_RIGHT_COL

        # row_type Spalte verstecken + Header leeren
        ws.column_dimensions[get_column_letter(ROWTYPE_COL)].hidden = True
        ws.cell(row=HEADER_ROW, column=ROWTYPE_COL).value = ""

        # Label formatieren
        ws.cell(row=label_row, column=TABLE_LEFT_COL).font = Font(
            name="GT Walsheim LC Light", size=9, color="FF4F2D7F", bold=True
        )

        # Datenbereich weiß + Zeilenhöhe
        for r in range(DATA_START_ROW, TABLE_BOTTOM_ROW + 1):
            ws.row_dimensions[r].height = 12
            for c in range(1, TABLE_RIGHT_COL + 1):
                ws.cell(row=r, column=c).fill = white_fill

        # Header fett + grau + alignment
        ws.row_dimensions[HEADER_ROW].height = 24
        for c in range(1, TABLE_RIGHT_COL + 1):
            cell = ws.cell(row=HEADER_ROW, column=c)
            cell.font = bold_font
            cell.fill = header_fill
            cell.alignment = header_alignment
        ws.cell(row=HEADER_ROW, column=TABLE_LEFT_COL).value = "kEUR"

        # Header: Alignment links erste Spalte, sonst rechts
        for c in range(1, TABLE_RIGHT_COL + 1):
            cell = ws.cell(row=HEADER_ROW, column=c)
            if c == TABLE_LEFT_COL:
                cell.alignment = Alignment(horizontal="left", vertical="bottom", wrap_text=True)
            else:
                cell.alignment = Alignment(horizontal="right", vertical="bottom", wrap_text=True)

        # Data: vertikal center, horizontal beibehalten
        for r in range(DATA_START_ROW, TABLE_BOTTOM_ROW + 1):
            for c in range(1, TABLE_RIGHT_COL + 1):
                cell = ws.cell(row=r, column=c)
                existing = cell.alignment or Alignment()
                cell.alignment = Alignment(
                    horizontal=existing.horizontal,
                    vertical="center",
                    wrap_text=existing.wrap_text,
                    indent=existing.indent
                )

        #Alles links klein schreiben
        for r in range(HEADER_ROW, TABLE_BOTTOM_ROW + 1):
            cell = ws.cell(row=r, column=TABLE_LEFT_COL)
            val = cell.value
            if isinstance(val, str) and val.strip():
                v = val.strip().lower()
                cell.value = v[0].upper() + v[1:]

        # Metrik-Spalten finden (dynamisch)
        # Header-Zeile
        headers = {}

        tbl = tables[label]

        # Metrikspalten erkennen: FY + bekannte Bridge-Metriken + row_type
        known_metrics = {"Upsell", "Downsell", "Cross-sell", "Lost", "NRR", "New"}
        metric_like = [c for c in tbl.columns if (isinstance(c, str) and (c.startswith("FY") or c in known_metrics))]
        # row_type ist Hilfsspalte
        has_row_type = ("row_type" in tbl.columns)

        # group_cols_count: alles vor den Metriken
        # ersten Index einer Metric-Spalte als Trenner
        first_metric_idx = min([tbl.columns.get_loc(c) for c in metric_like], default=len(tbl.columns)) - (1 if has_row_type else 0)

        # Welche Excel-Spalte wird eingerückt?
        # - bei >=2 Group-Spalten: Leaf ist 2. Group-Spalte => Index 1
        # - bei 1 Group-Spalte: Leaf ist die einzige Group-Spalte => Index 0
        group_cols_count = len(cfg["group_cols"])
        leaf_group_idx = 1 if group_cols_count >= 2 else 0
        INDENT_COL = TABLE_LEFT_COL + leaf_group_idx

        for c in range(TABLE_LEFT_COL, TABLE_RIGHT_COL + 1):
            headers[c] = ws.cell(row=HEADER_ROW, column=c).value

        # group columns: erste Spalte als "name"-Spalte
        name_col = TABLE_LEFT_COL
        # ----------

        # FY-Spalten erkennen + ARR-Highlight: erste FY links, letzte FY rechts
        fy_cols = [c for c, h in headers.items() if isinstance(h, str) and h.startswith("FY")]
        fy_cols_sorted = sorted(fy_cols)  # sortiert nach Spaltenposition
        fy_left_col = fy_cols_sorted[0] if fy_cols_sorted else None
        fy_right_col = fy_cols_sorted[-1] if len(fy_cols_sorted) >= 2 else (fy_cols_sorted[0] if fy_cols_sorted else None)

        # Money cols: alle Spalten außer name_col und row_type_col
        money_cols = [c for c in range(TABLE_LEFT_COL + 1, TABLE_RIGHT_COL + 1) if c != ROWTYPE_COL]

        # Zahlenformat Geld
        for c in money_cols:
            for r in range(DATA_START_ROW, TABLE_BOTTOM_ROW + 1):
                cell = ws.cell(row=r, column=c)
                if isinstance(cell.value, (int, float)):
                    cell.number_format = fmt_money

        # ARR-Spalten fill (FY links + FY rechts)
        for r in range(DATA_START_ROW, TABLE_BOTTOM_ROW + 1):
            if fy_left_col:
                ws.cell(row=r, column=fy_left_col).fill = arr_fill
            if fy_right_col:
                ws.cell(row=r, column=fy_right_col).fill = arr_fill

        #Outline/Portfolio-Logik nach row_type
        # Reset outline in Tabellenbereich
        for r in range(DATA_START_ROW, TABLE_BOTTOM_ROW + 1):
            rd = ws.row_dimensions[r]
            rd.outlineLevel = 0
            rd.hidden = False
            rd.collapsed = False

        # Prüfen, ob parent/bucket vorkommt
        row_types = set()
        for r in range(DATA_START_ROW, TABLE_BOTTOM_ROW + 1):
            rt = ws.cell(row=r, column=ROWTYPE_COL).value
            if rt is not None:
                row_types.add(str(rt).strip())

        has_parent_like = ("parent" in row_types) or ("bucket" in row_types)

        ws.sheet_properties.outlinePr.summaryBelow = True  # Totals unter Gruppe

        for r in range(DATA_START_ROW, TABLE_BOTTOM_ROW + 1):
            rt = ws.cell(row=r, column=ROWTYPE_COL).value
            rt = str(rt).strip() if rt is not None else ""

            rd = ws.row_dimensions[r]

            if rt == "total":
                rd.outlineLevel = 0
                rd.hidden = False
                rd.collapsed = False
                # Total fett + Linie oben & unten
                for c in range(TABLE_LEFT_COL, TABLE_RIGHT_COL + 1):
                    cell = ws.cell(row=r, column=c)
                    cell.font = bold_font
                    cell.border = border_top_bottom

            elif rt in ("parent", "bucket"):
                # Parent/Bucket: Level 1
                rd.outlineLevel = 1
                rd.hidden = False
                rd.collapsed = True

                # Fett über ganze Breite
                for c in range(TABLE_LEFT_COL, TABLE_RIGHT_COL + 1):
                    ws.cell(row=r, column=c).font = bold_font

                # Summenstrich ÜBER parent/bucket
                for c in range(TABLE_LEFT_COL, TABLE_RIGHT_COL + 1):
                    ws.cell(row=r, column=c).border = border_top

            elif rt == "leaf":
                # Leaf: Level 2 wenn Portfolio, sonst Level 1 (Entity-only)
                rd.outlineLevel = 2 if has_parent_like else 1
                rd.hidden = False
                rd.collapsed = True

                # Einrücken nur wenn Portfolio vorhanden (wie Country in deinem Beispiel)
                if has_parent_like:
                    cell = ws.cell(row=r, column=INDENT_COL)
                    existing = cell.alignment or Alignment()
                    cell.alignment = Alignment(
                        horizontal="left",
                        vertical=existing.vertical if existing.vertical else "center",
                        wrap_text=existing.wrap_text,
                        indent=1
                    )

        # Total collapsed, damit Excel das +/- korrekt anzeigt
        ws.row_dimensions[TOTAL_ROW].collapsed = True

    # Auto-fit Spaltenbreite
    col_max_length = {}
    for label, start_row, nrows, ncols in starts:
        HEADER_ROW = (start_row + 1 + TITLE_ROWS) + ROW_OFFSET
        DATA_START_ROW = HEADER_ROW + 1
        TABLE_LEFT_COL = 1 + COL_OFFSET
        TABLE_RIGHT_COL = COL_OFFSET + ncols
        TABLE_BOTTOM_ROW = DATA_START_ROW + nrows - 1

        for c in range(TABLE_LEFT_COL, TABLE_RIGHT_COL + 1):
            for r in range(HEADER_ROW, TABLE_BOTTOM_ROW + 1):
                cell = ws.cell(row=r, column=c)
                if cell.value is None:
                    continue
                if isinstance(cell.value, str):
                    length = len(cell.value)
                elif isinstance(cell.value, (int, float)):
                    if cell.number_format == fmt_pct:
                        length = 6
                    else:
                        formatted = f"{cell.value:,.0f}"
                        length = len(formatted)
                else:
                    length = len(str(cell.value))
                col_max_length[c] = max(col_max_length.get(c, 0), length)

    for c, max_len in col_max_length.items():
        letter = get_column_letter(c)
        adjusted_width = min(max_len + 1, 25)
        ws.column_dimensions[letter].width = adjusted_width

    wb.save(output_path_churn_universal)


def run_churn_pipeline(df: pd.DataFrame, cfg: dict) -> tuple[pd.DataFrame, dict]:
    """Build kEUR-formatted horizontal bridge table for export."""
    cfg = normalize_config(cfg)
    merged = build_merged_horizontal_bridge(df, cfg)
    return merged, cfg


def main(cfg: dict | None = None, df: pd.DataFrame | None = None) -> str:
    if cfg is None:
        if len(sys.argv) < 2:
            raise ValueError("Bitte den Pfad zur churn config.json als Argument übergeben.")
        with open(sys.argv[1], "r", encoding="utf-8-sig") as f:
            cfg = json.load(f)
    cfg = normalize_config(cfg)
    if df is None:
        sheet = cfg.get("sheet_name") or "Data Template"
        df = pd.read_excel(cfg["file_path"], sheet_name=sheet, engine="openpyxl")

    export_df, cfg = run_churn_pipeline(df, cfg)
    if cfg.get("case_id"):
        output_path = build_output_file_path(cfg)
    else:
        output_dir = str(cfg.get("output_file_path") or ".").strip()
        output_path = build_output_path(output_dir, cfg)
    os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)
    ensure_output_writable(output_path)

    from openpyxl import Workbook

    report_sheet = str(cfg.get("base_sheet_name") or "Churn")
    wb = Workbook()
    if wb.sheetnames[0] == "Sheet":
        del wb["Sheet"]
    write_churn_export_to_workbook(wb, export_df, report_sheet, cfg, formula_mode=False)
    format_churn_excel(wb, report_sheet, export_df, cfg, formula_mode=False)
    wb.save(output_path)
    wb.close()
    return output_path


if __name__ == "__main__":
    path = main()
    print(f"Churn output written: {path}")
