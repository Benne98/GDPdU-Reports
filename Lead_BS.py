import re
import sys
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.styles import Border, Side

PROJECT_ROOT = Path(__file__).resolve().parent
SCRIPTS_DIR = PROJECT_ROOT / "scripts"
BACKEND_DIR = PROJECT_ROOT / "backend"
if str(SCRIPTS_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPTS_DIR))
if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))

from gst_excel_theme import THEME, apply_zero_row_conditional_formatting  # noqa: E402
from databook_periods import ordered_reporting_columns_from_df, split_fy_and_ytd  # noqa: E402
from report_row_layout import build_lead_bs_row_structure, l3_order_from_mapping  # noqa: E402

# ================================================
# DESKTOP DEFAULTS
# ================================================
DESKTOP_DIR = PROJECT_ROOT / "Desktop"

INPUT_FILE = str(DESKTOP_DIR / "BS_Reconciliation_output.xlsx")
MAPPING_FILE_BS = str(DESKTOP_DIR / "BS_recon_Mapping.xlsx")
PL_LEAD_FILE = DESKTOP_DIR / "PL_Reconciliation_output.xlsx"
PL_LEAD_FILENAME = PL_LEAD_FILE.name

SHEET_MASTER = "Master_BS"
SHEET_OUT = "Lead_BS"
LEAD_IS_SHEET = "Lead_IS"
BS_BUCKET_SHEET = "BS_Bucket"

PROJECT_NAME = "Desktop Test"
GROUP_NAME = "Group"
UNIT_LABEL = "kEUR"

SOURCE_COL_CANDIDATES = ("L5", "L6")
REPORTED_FILTER_VALUE = "reported"
L4_SORT_BASIS = "latest_fy"

L5_ORDER = ["FA", "TWC", "OWC", "ND", "Other", "Equity"]
L5_TOTAL_LABEL = {
    "FA": "Fixed assets",
    "TWC": "Trade working capital",
    "OWC": "Other working capital",
    "ND": "Net debt",
    "Other": "Other",
    "Equity": "Equity",
}

NET_ASSETS_LABEL = "Net assets"
NET_ASSETS_COMPONENT_TOTALS = [
    "Fixed assets",
    "Trade working capital",
    "Other working capital",
    "Net debt",
    "Other",
]

MAP_START_COL = 5
MAP_END_COL = 9
POS_COL = 10

PROJECT_TITLE_ROW = 1
SUBTITLE_ROW = 2
TITLE_ROW = 6
HEADER_ROW = 8
HEADER_ROW_7 = 7
DATA_START_ROW = 9

ROW_HEIGHT = 12
POS_COL_WIDTH = 32
MAP_COL_WIDTH = 12
VALUE_COL_WIDTH = 7.86

# ================================================
# STYLES (gst_excel_theme / BS recon work)
# ================================================
FONT_BASE = THEME.font_base
FONT_BASE_BOLD = THEME.font_bold
FONT_HEADER = THEME.font_header
FONT_MAPPING = Font(name=THEME.font_name, size=THEME.font_size, color=THEME.text_header)
FONT_TITLE = Font(
    name=THEME.font_name,
    size=THEME.font_size,
    color=THEME.text_brand_title,
    bold=True,
)
FONT_PROJECT = Font(
    name=THEME.font_name,
    size=THEME.font_size_title,
    color=THEME.text_brand_title,
    bold=False,
)
FONT_SUBTITLE = Font(
    name=THEME.font_name,
    size=THEME.font_size_subtitle,
    color=THEME.text_brand_title,
    bold=False,
)
FONT_CHECK_RED = Font(name=THEME.font_name, size=THEME.font_size, color=THEME.delta_negative)

FILL_GREY = THEME.fill_tech
FILL_WHITE = THEME.fill_white
FILL_HEADER = THEME.fill_header
FILL_SUBTOTAL = THEME.fill_subtotal
FILL_SOURCE = PatternFill("solid", fgColor="FFFF99")
FILL_KPI = THEME.fill_subtotal

FONT_KPI = Font(name=THEME.font_name, size=THEME.font_size, color=THEME.text_kpi, italic=False, bold=False)
FONT_KPI_ITALIC = Font(name=THEME.font_name, size=THEME.font_size, color=THEME.text_kpi, italic=True, bold=False)

ALIGN_LEFT = Alignment(horizontal="left", vertical="center")
ALIGN_CENTER = Alignment(horizontal="center", vertical="center")
ALIGN_RIGHT = Alignment(horizontal="right", vertical="center")

THIN_SIDE = Side(style="thin", color=THEME.border_color)
TOP_BORDER = Border(top=THIN_SIDE)
BOTTOM_BORDER = Border(bottom=THIN_SIDE)
TOP_BOTTOM_BORDER = Border(top=THIN_SIDE, bottom=THIN_SIDE)
BORDER_SUBTOTAL_TOP = THEME.border_subtotal_top

NUM_FMT_INT = "#,##0;(#,##0);-"
NUM_FMT_KPI = "0.0;(0.0);-"

# ================================================
# HELPERS
# ================================================
def col_letter(idx: int) -> str:
    return get_column_letter(idx)


def detect_year_columns(df_: pd.DataFrame) -> list[str]:
    fy, _ = split_fy_and_ytd(ordered_reporting_columns_from_df(df_))
    if not fy:
        raise ValueError("Keine FY-Spalten im Master_BS gefunden.")
    return fy


def detect_reporting_period_columns(df_: pd.DataFrame) -> list[str]:
    periods = ordered_reporting_columns_from_df(df_)
    if not periods:
        raise ValueError("Keine FY/YTD-Spalten im Master_BS gefunden.")
    return periods


def detect_source_col(df_):
    candidates = ["Quelle", "Source", "Reported/Adjusted", "Type", "Scenario"]
    sc = next((c for c in candidates if c in df_.columns), None)
    if sc is not None:
        return sc
    for c in df_.columns:
        try:
            s = df_[c].astype(str).str.strip().str.lower()
            if (s == "reported").any():
                return c
        except Exception:
            pass
    return None


def resolve_source_section_column(df_: pd.DataFrame) -> str | None:
    markers = {REPORTED_FILTER_VALUE, "adjusted", "consolidation"}

    def marker_count(col: str) -> int:
        if col not in df_.columns:
            return 0
        series = df_[col].astype(str).str.strip().str.lower()
        return int(series.isin(markers).sum())

    best_col = None
    best_count = 0
    for col in SOURCE_COL_CANDIDATES:
        count = marker_count(col)
        if count > best_count:
            best_col = col
            best_count = count
    if best_col and best_count > 0:
        return best_col
    legacy = detect_source_col(df_)
    if legacy:
        return legacy
    for col in SOURCE_COL_CANDIDATES:
        if col in df_.columns:
            return col
    return None


def detect_na_bucket_col(df_):
    """Net-assets bucket column: prefer explicit NA, else legacy L5 view columns."""
    for cand in [
        "NA",
        "L5 - GT NA View",
        "L5 - GT NA view",
        "L5 - GT NA",
        "L5",
        "L5 - Management defined Net Asset View",
    ]:
        if cand in df_.columns:
            return cand
    for c in df_.columns:
        if str(c).strip() == "NA":
            return c
        if str(c).strip().startswith("L5"):
            return c
    raise ValueError("Keine NA-/L5-Bucket-Spalte gefunden (z.B. 'NA' oder 'L5').")


detect_l5_col = detect_na_bucket_col


def normalize_l5_bucket(x: str) -> str:
    if x is None or (isinstance(x, float) and pd.isna(x)):
        return "Other"
    s = str(x).strip()
    if s == "":
        return "Other"
    if s in L5_ORDER:
        return s
    for b in L5_ORDER:
        if s.startswith(b):
            return b
    return "Other"


FY_PERIOD_PATTERN = re.compile(r"^FY(\d{2}|\d{4})A$", re.IGNORECASE)


def is_fy_period_label(period: str) -> bool:
    return bool(FY_PERIOD_PATTERN.match(period.strip()))


def compute_lead_is_layout(n_years: int):
    pos1 = 10
    y1 = list(range(pos1 + 1, pos1 + 1 + n_years))
    cagr1 = y1[-1] + 1
    spacer = cagr1 + 1
    pos2 = spacer + 1
    y2 = list(range(pos2 + 1, pos2 + 1 + n_years))
    return pos2, y2


def detect_lead_is_proforma_year_cols(ws_is, years: list[str]) -> list[int]:
    _, y2_cols = compute_lead_is_layout(len(years))
    found = {}
    for cc in y2_cols:
        v = ws_is.cell(HEADER_ROW, cc).value
        if isinstance(v, str) and v.strip() in years:
            found[v.strip()] = cc
    missing = [y for y in years if y not in found]
    if missing:
        raise RuntimeError(f"Lead_IS pro-forma block missing year headers: {missing}")
    return [found[y] for y in years]


def lead_is_external_ref(col_idx: int, row_num: int) -> str:
    col_l = col_letter(col_idx)
    return f"='[{PL_LEAD_FILENAME}]{LEAD_IS_SHEET}'!{col_l}{row_num}"


def find_row_in_col_any(ws_, col_idx, needles):
    needles = [n.strip() for n in needles]
    for rr in range(1, ws_.max_row + 1):
        v = ws_.cell(rr, col_idx).value
        if isinstance(v, str) and v.strip() in needles:
            return rr
    return None


def find_row_by_label_contains(ws_, col_idx, needle):
    needle = needle.lower()
    for rr in range(1, ws_.max_row + 1):
        v = ws_.cell(rr, col_idx).value
        if isinstance(v, str) and needle in v.lower():
            return rr
    return None


def detect_period_bucket_cols(ws_):
    period_to_cols = {}
    for cc in range(1, ws_.max_column + 1):
        v = ws_.cell(HEADER_ROW_7, cc).value
        if not isinstance(v, str):
            continue
        period = v.strip()
        if not is_fy_period_label(period):
            continue
        bucket_map = {}
        steps = 0
        c2 = cc
        while c2 <= ws_.max_column and steps < 60 and len(bucket_map) < len(L5_ORDER):
            hv = ws_.cell(HEADER_ROW, c2).value
            if isinstance(hv, str):
                hvs = hv.strip()
                if hvs in L5_ORDER:
                    bucket_map[hvs] = c2
            c2 += 1
            steps += 1
        if all(b in bucket_map for b in L5_ORDER):
            period_to_cols[period] = bucket_map
    if not period_to_cols:
        raise RuntimeError("Konnte in BS_Bucket keine FY-Periodenblöcke in Zeile 7 finden.")
    return period_to_cols


# ================================================
# MAIN
# ================================================
print("Lead_BS_work — Desktop defaults:")
print(f"  input:    {INPUT_FILE}")
print(f"  mapping:  {MAPPING_FILE_BS}")
print(f"  PL lead:  {PL_LEAD_FILE}")

if not PL_LEAD_FILE.exists():
    raise RuntimeError(f"PL Lead-Datei nicht gefunden: {PL_LEAD_FILE} (zuerst Lead_IS_work ausführen).")

wb_pl = load_workbook(PL_LEAD_FILE, read_only=True, data_only=False)
if LEAD_IS_SHEET not in wb_pl.sheetnames:
    wb_pl.close()
    raise RuntimeError(f"Sheet '{LEAD_IS_SHEET}' fehlt in {PL_LEAD_FILE}.")
ws_is = wb_pl[LEAD_IS_SHEET]
IS_LABEL_COL = 10
is_row_cogs = find_row_by_label_contains(ws_is, IS_LABEL_COL, "cost of goods sold")
is_row_ns = find_row_by_label_contains(ws_is, IS_LABEL_COL, "net sales")
if is_row_cogs is None or is_row_ns is None:
    wb_pl.close()
    raise RuntimeError("Konnte 'Cost of goods sold' oder 'Net sales' in Lead_IS (Spalte J) nicht finden.")

wb = load_workbook(INPUT_FILE)
if SHEET_MASTER not in wb.sheetnames:
    raise RuntimeError(f"Sheet '{SHEET_MASTER}' nicht gefunden in {INPUT_FILE}")

if SHEET_OUT in wb.sheetnames:
    del wb[SHEET_OUT]
ws = wb.create_sheet(SHEET_OUT)

df_bs = pd.read_excel(INPUT_FILE, sheet_name=SHEET_MASTER, engine="openpyxl")
PERIODS = detect_reporting_period_columns(df_bs)
FY_COLS, YTD_COLS = split_fy_and_ytd(PERIODS)
Y_COLS = list(range(POS_COL + 1, POS_COL + 1 + len(PERIODS)))
IS_YEAR_COLS = detect_lead_is_proforma_year_cols(ws_is, PERIODS)
wb_pl.close()

source_col = resolve_source_section_column(df_bs)
l5_col = detect_l5_col(df_bs)

required = {"Entity", "L3", "L4", l5_col} | set(PERIODS)
missing = required - set(df_bs.columns)
if missing:
    raise ValueError(f"Master_BS fehlt Spalten: {missing}")

map_df = pd.read_excel(MAPPING_FILE_BS, sheet_name=0, engine="openpyxl")
if "L3" not in map_df.columns:
    raise ValueError("BS-Mapping muss Spalte 'L3' enthalten.")

map_df = map_df.copy()
map_df["L3"] = map_df["L3"].astype(str).str.strip()
mapping_l3_order = l3_order_from_mapping(map_df)

MASTER_START = 2
MASTER_END = len(df_bs) + 1


def master_range(colname):
    L = get_column_letter(df_bs.columns.get_loc(colname) + 1)
    return f"{SHEET_MASTER}!${L}${MASTER_START}:${L}${MASTER_END}"


l3_rng = master_range("L3")
l4_rng = master_range("L4")
l5_rng = master_range(l5_col)
src_rng = master_range(source_col) if source_col is not None else None
year_rng = {y: master_range(y) for y in PERIODS}

row_structure = build_lead_bs_row_structure(
    df_bs,
    mapping_l3_order,
    {"l4_sort_basis": L4_SORT_BASIS},
    source_col=source_col or l5_col,
    bucket_col=l5_col,
    bucket_order=L5_ORDER,
    normalize_bucket_fn=normalize_l5_bucket,
    l5_total_labels=L5_TOTAL_LABEL,
)

existing_total_labels = [r["label"] for r in row_structure if r["type"] == "total_l5"]
net_assets_components_existing = [x for x in NET_ASSETS_COMPONENT_TOTALS if x in existing_total_labels]
missing_na = [x for x in NET_ASSETS_COMPONENT_TOTALS if x not in existing_total_labels]
if missing_na:
    print(f"Warnung: Net-assets-Komponenten fehlen in L5-Totals: {missing_na}")

insert_idx = None
if net_assets_components_existing:
    last_pos = -1
    for i, r in enumerate(row_structure):
        if r["type"] == "total_l5" and r["label"] in net_assets_components_existing:
            last_pos = i
    insert_idx = last_pos + 1 if last_pos >= 0 else None

net_assets_row = {
    "type": "total_extra",
    "label": NET_ASSETS_LABEL,
    "components_total_labels": net_assets_components_existing,
}
row_structure.insert(insert_idx if insert_idx is not None else len(row_structure), net_assets_row)

for i, r in enumerate(row_structure):
    r["_idx"] = i
    r["_excel_row"] = DATA_START_ROW + i

LAST_TABLE_ROW = DATA_START_ROW + len(row_structure) - 1

bucket_total_row = {}
net_assets_rownum = None
equity_rownum = None
for r in row_structure:
    if r["type"] == "total_l5":
        bucket_total_row[r.get("bucket")] = r["_excel_row"]
        if r.get("label") == "Equity":
            equity_rownum = r["_excel_row"]
    if r["type"] == "total_extra" and r.get("label") == NET_ASSETS_LABEL:
        net_assets_rownum = r["_excel_row"]

pt = ws.cell(PROJECT_TITLE_ROW, POS_COL, f"Project {PROJECT_NAME}")
pt.font = FONT_PROJECT
pt.alignment = ALIGN_LEFT
ws.row_dimensions[PROJECT_TITLE_ROW].height = 36

st = ws.cell(SUBTITLE_ROW, POS_COL, f"{GROUP_NAME} Lead balance sheet")
st.font = FONT_SUBTITLE
st.alignment = ALIGN_LEFT

tt = ws.cell(TITLE_ROW, POS_COL, f"{GROUP_NAME} | Lead balance sheet {FY_COLS[0]} - {FY_COLS[-1]}")
tt.font = FONT_TITLE
tt.alignment = ALIGN_LEFT

for hdr_col, hdr_txt in enumerate(
    ["Reported", "", "L5", "L3", "L4"],
    start=MAP_START_COL,
):
    ws.cell(HEADER_ROW, hdr_col, hdr_txt).font = FONT_HEADER

for c in range(MAP_START_COL, MAP_END_COL + 1):
    ws.cell(HEADER_ROW, c).alignment = ALIGN_LEFT
    ws.cell(HEADER_ROW, c).fill = FILL_HEADER
    ws.cell(HEADER_ROW_7, c).fill = FILL_HEADER

for cc in [POS_COL] + Y_COLS:
    ws.cell(HEADER_ROW_7, cc).fill = FILL_HEADER

ws.cell(HEADER_ROW, POS_COL, UNIT_LABEL).font = FONT_HEADER
ws.cell(HEADER_ROW, POS_COL).alignment = ALIGN_LEFT
ws.cell(HEADER_ROW, POS_COL).fill = FILL_HEADER

for idx, y in enumerate(PERIODS):
    cc = Y_COLS[idx]
    h = ws.cell(HEADER_ROW, cc, y)
    h.font = FONT_HEADER
    h.alignment = ALIGN_RIGHT
    h.fill = FILL_HEADER

for cc in range(MAP_START_COL, Y_COLS[-1] + 1):
    ws.cell(HEADER_ROW, cc).border = BOTTOM_BORDER

for c in range(MAP_START_COL, MAP_END_COL + 1):
    ws.column_dimensions[col_letter(c)].width = MAP_COL_WIDTH
ws.column_dimensions[col_letter(POS_COL)].width = POS_COL_WIDTH
for c in Y_COLS:
    ws.column_dimensions[col_letter(c)].width = VALUE_COL_WIDTH

map_rep_col = col_letter(MAP_START_COL + 0)
map_l5_col = col_letter(MAP_START_COL + 2)
map_l3_col = col_letter(MAP_START_COL + 3)
map_l4_col = col_letter(MAP_START_COL + 4)


def subtotal_span_for_l3(struct, idx, bucket, l3_label):
    k = idx - 1
    while (
        k >= 0
        and struct[k]["type"] == "detail"
        and struct[k].get("L5") == bucket
        and struct[k].get("L3") == l3_label
    ):
        k -= 1
    start = DATA_START_ROW + (k + 1)
    end = DATA_START_ROW + (idx - 1)
    return start, end


def bucket_sum_refs(struct, bucket, colL):
    refs = []
    for rr in struct:
        if rr.get("L5") != bucket:
            continue
        if rr["type"] == "subtotal_l3":
            refs.append(f"{colL}{rr['_excel_row']}")
        elif (
            rr["type"] == "detail_single"
            and str(rr["L3"]).strip() == str(rr["L4"]).strip()
        ):
            refs.append(f"{colL}{rr['_excel_row']}")
    return refs


for i, r in enumerate(row_structure):
    excel_row = r["_excel_row"]

    ws.cell(excel_row, MAP_START_COL + 0, "Reported").font = FONT_MAPPING
    ws.cell(excel_row, MAP_START_COL + 1, "").font = FONT_MAPPING
    ws.cell(excel_row, MAP_START_COL + 2, r.get("L5", "")).font = FONT_MAPPING
    ws.cell(excel_row, MAP_START_COL + 3, r.get("L3", "")).font = FONT_MAPPING
    ws.cell(excel_row, MAP_START_COL + 4, r.get("L4", "")).font = FONT_MAPPING
    for c in range(MAP_START_COL, MAP_END_COL + 1):
        ws.cell(excel_row, c).alignment = ALIGN_LEFT

    pc = ws.cell(excel_row, POS_COL, r["label"])
    if r["type"] == "detail" and str(r["L3"]).strip() != str(r["L4"]).strip():
        pc.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    else:
        pc.alignment = ALIGN_LEFT

    is_bold = r["type"] in {"total_l5", "total_extra"}
    pc.font = FONT_BASE_BOLD if is_bold else FONT_BASE

    if r["type"] in {"detail", "detail_single"}:
        ws.row_dimensions[excel_row].outlineLevel = 1
        ws.row_dimensions[excel_row].hidden = False
    else:
        ws.row_dimensions[excel_row].outlineLevel = 0
        ws.row_dimensions[excel_row].hidden = False

    for y_idx, year in enumerate(PERIODS):
        cell = ws.cell(excel_row, Y_COLS[y_idx])
        cell.alignment = ALIGN_RIGHT
        cell.number_format = NUM_FMT_INT
        cell.font = FONT_BASE_BOLD if is_bold else FONT_BASE

        rep_crit = f"${map_rep_col}${excel_row}"
        l5_crit = f"${map_l5_col}${excel_row}"
        l3_crit = f"${map_l3_col}${excel_row}"
        l4_crit = f"${map_l4_col}${excel_row}"

        if r["type"] in {"detail", "detail_single"}:
            sum_rng = year_rng[year]
            if src_rng is not None:
                cell.value = (
                    f"=SUMIFS({sum_rng},"
                    f"{src_rng},{rep_crit},"
                    f"{l5_rng},{l5_crit},"
                    f"{l3_rng},{l3_crit},"
                    f"{l4_rng},{l4_crit}"
                    f")/1000"
                )
            else:
                cell.value = (
                    f"=SUMIFS({sum_rng},"
                    f"{l5_rng},{l5_crit},"
                    f"{l3_rng},{l3_crit},"
                    f"{l4_rng},{l4_crit}"
                    f")/1000"
                )
        elif r["type"] == "subtotal_l3":
            start, end = subtotal_span_for_l3(row_structure, i, r["L5"], r["label"])
            colL = col_letter(Y_COLS[y_idx])
            cell.value = f"=SUM({colL}{start}:{colL}{end})" if end >= start else 0
        elif r["type"] == "total_l5":
            colL = col_letter(Y_COLS[y_idx])
            refs = bucket_sum_refs(row_structure, r["bucket"], colL)
            cell.value = f"=SUM({','.join(refs)})" if refs else 0
        elif r["type"] == "total_extra":
            colL = col_letter(Y_COLS[y_idx])
            refs = []
            for lab in r.get("components_total_labels", []):
                for rr2 in row_structure:
                    if rr2["type"] == "total_l5" and rr2["label"] == lab:
                        refs.append(f"{colL}{rr2['_excel_row']}")
                        break
            cell.value = f"=SUM({','.join(refs)})" if refs else 0

    if r["type"] in {"total_l5", "total_extra"}:
        is_equity_total = r["type"] == "total_l5" and r.get("label") == "Equity"
        is_net_assets = r["type"] == "total_extra" and r.get("label") == NET_ASSETS_LABEL
        border_style = TOP_BOTTOM_BORDER if (is_equity_total or is_net_assets) else BORDER_SUBTOTAL_TOP
        for cc in range(MAP_START_COL, Y_COLS[-1] + 1):
            ws.cell(excel_row, cc).fill = FILL_SUBTOTAL
        ws.cell(excel_row, POS_COL).border = border_style
        for cc in Y_COLS:
            ws.cell(excel_row, cc).border = border_style
    else:
        for cc in range(MAP_START_COL, Y_COLS[-1] + 1):
            ws.cell(excel_row, cc).fill = FILL_WHITE

row_inventory = find_row_by_label_contains(ws, POS_COL, "Inventories")
row_recv = find_row_by_label_contains(ws, POS_COL, "trade receivables")
row_pay = find_row_by_label_contains(ws, POS_COL, "trade payables")
row_twc = bucket_total_row.get("TWC")
row_owc = bucket_total_row.get("OWC")

KPI_START_ROW = LAST_TABLE_ROW + 1
kpi_labels = [
    "TWC in % of net sales",
    "NWC in % of net sales",
    "DIO",
    "DSO",
    "DPO",
]
KPI_ROWS = set(range(KPI_START_ROW, KPI_START_ROW + len(kpi_labels)))

for i_kpi, label in enumerate(kpi_labels):
    rr = KPI_START_ROW + i_kpi
    is_italic_kpi = label in ("TWC in % of net sales", "NWC in % of net sales")
    ws.cell(rr, POS_COL).fill = FILL_KPI
    for cc in Y_COLS:
        ws.cell(rr, cc).fill = FILL_KPI
    lbl_cell = ws.cell(rr, POS_COL, label)
    lbl_cell.alignment = ALIGN_LEFT
    lbl_cell.font = FONT_KPI_ITALIC if is_italic_kpi else FONT_KPI
    for y_idx in range(len(PERIODS)):
        v = ws.cell(rr, Y_COLS[y_idx])
        v.number_format = NUM_FMT_KPI
        v.alignment = ALIGN_RIGHT
        v.font = FONT_KPI_ITALIC if is_italic_kpi else FONT_KPI

GAP1_ROW_1 = KPI_START_ROW + len(kpi_labels)
GAP1_ROW_2 = GAP1_ROW_1 + 1
for rr in (GAP1_ROW_1, GAP1_ROW_2):
    ws.cell(rr, POS_COL).fill = FILL_WHITE
    for cc in Y_COLS:
        ws.cell(rr, cc).fill = FILL_WHITE

COGS_HELPER_ROW = GAP1_ROW_2 + 1
NET_SALES_HELPER_ROW = COGS_HELPER_ROW + 1

for rr in (COGS_HELPER_ROW, NET_SALES_HELPER_ROW):
    ws.row_dimensions[rr].outlineLevel = 2
    ws.row_dimensions[rr].hidden = True

for rr, label in [(COGS_HELPER_ROW, "Cost of goods sold"), (NET_SALES_HELPER_ROW, "Net sales")]:
    ws.cell(rr, POS_COL, label).alignment = ALIGN_LEFT
    ws.cell(rr, POS_COL).font = FONT_KPI
    ws.cell(rr, POS_COL).fill = FILL_WHITE
    for cc in Y_COLS:
        ws.cell(rr, cc).fill = FILL_WHITE
        ws.cell(rr, cc).font = FONT_KPI
        ws.cell(rr, cc).alignment = ALIGN_RIGHT
        ws.cell(rr, cc).number_format = NUM_FMT_INT

for y_idx in range(len(PERIODS)):
    ws.cell(COGS_HELPER_ROW, Y_COLS[y_idx]).value = lead_is_external_ref(
        IS_YEAR_COLS[y_idx], is_row_cogs,
    )
    ws.cell(NET_SALES_HELPER_ROW, Y_COLS[y_idx]).value = lead_is_external_ref(
        IS_YEAR_COLS[y_idx], is_row_ns,
    )

for y_idx in range(len(PERIODS)):
    colL = col_letter(Y_COLS[y_idx])
    ns_cell = f"{colL}{NET_SALES_HELPER_ROW}"
    cogs_cell = f"{colL}{COGS_HELPER_ROW}"
    twc_cell = f"{colL}{row_twc}" if row_twc else None
    owc_cell = f"{colL}{row_owc}" if row_owc else None
    inv_cell = f"{colL}{row_inventory}" if row_inventory else None
    recv_cell = f"{colL}{row_recv}" if row_recv else None
    pay_cell = f"{colL}{row_pay}" if row_pay else None

    r_twc_pct = KPI_START_ROW + 0
    r_nwc_pct = KPI_START_ROW + 1
    r_dio = KPI_START_ROW + 2
    r_dso = KPI_START_ROW + 3
    r_dpo = KPI_START_ROW + 4

    ws.cell(r_twc_pct, Y_COLS[y_idx]).value = (
        f'=IFERROR({twc_cell}/{ns_cell}*100,"n/a")' if twc_cell else "n/a"
    )
    ws.cell(r_nwc_pct, Y_COLS[y_idx]).value = (
        f'=IFERROR(({twc_cell}+{owc_cell})/{ns_cell}*100,"n/a")' if twc_cell and owc_cell else "n/a"
    )
    ws.cell(r_dio, Y_COLS[y_idx]).value = (
        f'=IFERROR(-{inv_cell}/{cogs_cell}*365,"n/a")' if inv_cell else "n/a"
    )
    ws.cell(r_dso, Y_COLS[y_idx]).value = (
        f'=IFERROR({recv_cell}/{ns_cell}*365,"n/a")' if recv_cell else "n/a"
    )
    ws.cell(r_dpo, Y_COLS[y_idx]).value = (
        f'=IFERROR({pay_cell}/{cogs_cell}*365,"n/a")' if pay_cell else "n/a"
    )

KPI_END_ROW = NET_SALES_HELPER_ROW
CHECK_ROW = KPI_END_ROW + 2

check_rows_to_collapse = []
source_rows_yellow = set()

ws.cell(CHECK_ROW, POS_COL, "Check").font = FONT_CHECK_RED
ws.cell(CHECK_ROW, POS_COL).alignment = ALIGN_LEFT

if net_assets_rownum is not None and equity_rownum is not None:
    for y_idx in range(len(PERIODS)):
        colL = col_letter(Y_COLS[y_idx])
        na_cell = f"{colL}{net_assets_rownum}"
        eq_cell = f"{colL}{equity_rownum}"
        c = ws.cell(CHECK_ROW, Y_COLS[y_idx])
        c.value = f"=ROUND(SUM({na_cell},{eq_cell}),0)"
        c.font = FONT_CHECK_RED
        c.alignment = ALIGN_RIGHT
        c.number_format = NUM_FMT_INT

check_rows_to_collapse.append(CHECK_ROW)
current_row = CHECK_ROW + 1

if BS_BUCKET_SHEET not in wb.sheetnames:
    print(f"Hinweis: Sheet '{BS_BUCKET_SHEET}' fehlt — Bucket-Checks werden übersprungen.")
else:
    ws_bs = wb[BS_BUCKET_SHEET]
    period_bucket_cols = detect_period_bucket_cols(ws_bs)
    bs_total_assets_row = find_row_in_col_any(ws_bs, POS_COL, ["Total assets"])
    bs_total_el_row = find_row_in_col_any(
        ws_bs, POS_COL, ["Total equity & liabilities", "Total equity & liabilities "],
    )
    if bs_total_assets_row is None or bs_total_el_row is None:
        print("Warnung: BS_Bucket Total-Zeilen nicht gefunden — Bucket-Checks übersprungen.")
    else:
        for bucket in L5_ORDER:
            if bucket not in bucket_total_row:
                continue
            src_row = current_row
            delta_row = current_row + 1
            ws.cell(src_row, POS_COL, f"Source: {bucket}").font = FONT_BASE
            ws.cell(src_row, POS_COL).alignment = ALIGN_LEFT
            ws.cell(delta_row, POS_COL, "Check").font = FONT_CHECK_RED
            ws.cell(delta_row, POS_COL).alignment = ALIGN_LEFT
            ws.cell(src_row, POS_COL).fill = FILL_SOURCE
            for cc in Y_COLS:
                ws.cell(src_row, cc).fill = FILL_SOURCE
            source_rows_yellow.add(src_row)

            for y_idx, year in enumerate(PERIODS):
                period = year
                if period not in period_bucket_cols:
                    avail = sorted(period_bucket_cols.keys())
                    period = avail[y_idx] if y_idx < len(avail) else avail[0]
                bucket_colL = col_letter(period_bucket_cols[period][bucket])
                src_cell = ws.cell(src_row, Y_COLS[y_idx])
                src_cell.value = (
                    f"='{BS_BUCKET_SHEET}'!{bucket_colL}{bs_total_assets_row}"
                    f"+'{BS_BUCKET_SHEET}'!{bucket_colL}{bs_total_el_row}"
                )
                src_cell.number_format = NUM_FMT_INT
                src_cell.alignment = ALIGN_RIGHT
                src_cell.font = FONT_BASE
                lead_cell = f"{col_letter(Y_COLS[y_idx])}{bucket_total_row[bucket]}"
                d = ws.cell(delta_row, Y_COLS[y_idx])
                d.value = f"={lead_cell}-{src_cell.coordinate}"
                d.number_format = NUM_FMT_INT
                d.alignment = ALIGN_RIGHT
                d.font = FONT_CHECK_RED

            check_rows_to_collapse.extend([src_row, delta_row])
            current_row += 2

for rr in check_rows_to_collapse:
    ws.row_dimensions[rr].outlineLevel = 2
    ws.row_dimensions[rr].hidden = True

ws.sheet_view.showOutlineSymbols = True

FILL_END_ROW = current_row + 25
FILL_END_COL = Y_COLS[-1] + 10

for rr in range(1, FILL_END_ROW + 1):
    if rr != PROJECT_TITLE_ROW:
        ws.row_dimensions[rr].height = ROW_HEIGHT

for rr in range(1, FILL_END_ROW + 1):
    for cc in range(1, POS_COL):
        ws.cell(rr, cc).fill = FILL_GREY
    for cc in range(POS_COL, FILL_END_COL + 1):
        if rr in source_rows_yellow and cc in ([POS_COL] + Y_COLS):
            continue
        if rr in (HEADER_ROW_7, HEADER_ROW) and cc in ([POS_COL] + Y_COLS):
            continue
        if rr in KPI_ROWS and cc in ([POS_COL] + Y_COLS):
            ws.cell(rr, cc).fill = FILL_KPI
            continue
        ws.cell(rr, cc).fill = FILL_WHITE

for cc in range(MAP_START_COL, MAP_END_COL + 1):
    ws.cell(HEADER_ROW, cc).fill = FILL_HEADER

for cc in [POS_COL] + Y_COLS:
    ws.cell(HEADER_ROW_7, cc).fill = FILL_HEADER
    ws.cell(HEADER_ROW, cc).fill = FILL_HEADER

for cc in range(1, MAP_END_COL + 1):
    col_l = col_letter(cc)
    ws.column_dimensions[col_l].outlineLevel = 1
    ws.column_dimensions[col_l].hidden = True

apply_zero_row_conditional_formatting(
    ws,
    first_row=DATA_START_ROW,
    last_row=LAST_TABLE_ROW,
    year_col_indices=Y_COLS,
    style_start_col=POS_COL,
    style_end_col=Y_COLS[-1],
)

wb.save(INPUT_FILE)
print(f"Fertig. Reiter '{SHEET_OUT}' gespeichert in:\n{INPUT_FILE}")
print(f"NA-Bucket-Spalte: {l5_col}")
print(f"Verwendete Perioden: {PERIODS}")
if source_col is None:
    print("Hinweis: Keine Reported/Adjusted-Spalte erkannt.")
else:
    print(f"Source-Spalte: '{source_col}'")
