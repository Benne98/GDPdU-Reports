"""Tests for etl/bs_pl_master.py — synthetic data only."""
from __future__ import annotations

import pandas as pd
import pytest

from etl.bs_pl_master import (
    BS_PL_REPLACE_MODES,
    ENTITY_PREFIX_COL,
    build_mapping_frames,
    bs_pl_master_profile,
    detect_level_columns,
    filter_mapping_append_only,
    is_bs_pl_master_workbook,
    preview_stats,
    read_bs_pl_master,
    resolve_entity_prefixes,
)
from etl.mapping_account import apply_account_mapping


def _raw_bs_pl() -> pd.DataFrame:
    return pd.DataFrame(
        {
            "Entity": ["Atlas", "Atlas"],
            "Account": ["16100", "81910"],
            "Account description": ["Asset acct", "PL acct"],
            "L1": ["BS", "PL"],
            "L2": ["Fixed assets", "Income"],
            "L3": ["Tangible", "Net sales"],
            "L4": ["PPE", "Reduction"],
        }
    )


def test_detect_level_columns_l4_only():
    cols = detect_level_columns(["Entity", "L1", "L2", "L3", "L4"])
    assert cols == ["L1", "L2", "L3", "L4"]


def test_detect_level_columns_with_optional_l5_l6():
    cols = detect_level_columns(["L1", "L2", "L3", "L4", "L5", "L6"])
    assert cols == ["L1", "L2", "L3", "L4", "L5", "L6"]


def test_is_bs_pl_master_workbook():
    assert is_bs_pl_master_workbook(["Master_BS", "Master_PL", "Other"])
    assert not is_bs_pl_master_workbook(["Master_BS"])
    assert not is_bs_pl_master_workbook(["PL", "BS"])


def test_resolve_entity_prefixes():
    raw = _raw_bs_pl()
    lookup = {"Atlas": "01", "Meridian": "02"}
    resolved, unknown = resolve_entity_prefixes(raw, lookup)
    assert unknown == []
    assert list(resolved[ENTITY_PREFIX_COL]) == ["01", "01"]


def test_build_account_number_group_via_profile():
    raw = _raw_bs_pl()
    lookup = {"Atlas": "01"}
    resolved, _ = resolve_entity_prefixes(raw, lookup)
    prof = bs_pl_master_profile(list(raw.columns))
    out = apply_account_mapping(resolved, prof, fiscal_year=2024)
    assert out.loc[0, "account_number_group"] == "01016100"
    assert out.loc[1, "account_number_group"] == "01081910"
    assert out.loc[0, "level_0"] == "BS"
    assert out.loc[1, "level_3"] == "Reduction"
    assert pd.isna(out.loc[1, "level_4"])


def test_optional_l5_l6_mapping():
    raw = _raw_bs_pl().copy()
    raw["L5"] = ["Detail", pd.NA]
    raw["L6"] = ["Sub", pd.NA]
    lookup = {"Atlas": "01"}
    resolved, _ = resolve_entity_prefixes(raw, lookup)
    prof = bs_pl_master_profile(list(raw.columns))
    out = apply_account_mapping(resolved, prof, fiscal_year=2024)
    assert out.loc[0, "level_4"] == "Detail"
    assert out.loc[0, "l4_sub"] == "Sub"


class _FakeSession:
    def execute(self, *_args, **_kwargs):
        class _R:
            def fetchall(self):
                return [
                    ("Atlas", "Atlas GmbH", "01"),
                    ("Meridian", "Meridian GmbH", "02"),
                ]

        return _R()


def test_build_mapping_frames_replicates_fiscal_years():
    raw = _raw_bs_pl()
    session = _FakeSession()
    frames, meta = build_mapping_frames(raw, session, [2023, 2024])
    assert len(frames) == 4
    assert set(frames["fiscal_year"].astype(int)) == {2023, 2024}
    assert meta["skipped_rows"] == 0
    assert meta["level_columns"] == ["L1", "L2", "L3", "L4"]


def test_bs_pl_master_profile_maps_l4_to_level_3():
    prof = bs_pl_master_profile()
    assert prof.columns["level_3"] == "L4"
    assert "level_4" not in prof.columns


def test_read_legacy_template_structure_entity_present():
    """Smoke-read a legacy desktop template that carries an Entity column.

    This test is intentionally SKIPPED in CI (file not on disk).  It serves as
    backward-compatibility insurance: legacy two-sheet files WITH an Entity column
    must still parse correctly (the Entity column is first in the output frame).
    """
    from pathlib import Path

    path = Path(r"C:\Users\bened\OneDrive\Desktop\BS_PL_Master Test input.xlsx")
    if not path.exists():
        pytest.skip("Legacy desktop template not present — backward-compat check skipped")
    df = read_bs_pl_master(path)
    # Legacy format: Entity column leads the frame.
    assert list(df.columns) == [
        "Entity",
        "Account",
        "Account description",
        "L1",
        "L2",
        "L3",
        "L4",
    ]
    assert len(df) > 0


# =========================================================================== #
# Synthetic xlsx helpers for Entity-less template tests (CoA wizard rework)
# =========================================================================== #

def _make_bs_xlsx(tmp_path) -> "Path":
    """Synthetic BS-only workbook WITHOUT Entity column."""
    from openpyxl import Workbook
    from pathlib import Path as _P
    wb = Workbook()
    ws = wb.active
    ws.title = "Master_BS"
    ws.append(["Account", "Account description", "L1", "L2", "L3", "L4"])
    ws.append(["16100", "Trade receivables", "BS", "Assets", "Current", "Trade AR"])
    ws.append(["18000", "Bank", "BS", "Assets", "Current", "Cash"])
    p = _P(tmp_path) / "bs_only.xlsx"
    wb.save(str(p))
    return p


def _make_pl_xlsx(tmp_path) -> "Path":
    """Synthetic PL-only workbook WITHOUT Entity column, using 'L1 - BS/PL' header."""
    from openpyxl import Workbook
    from pathlib import Path as _P
    wb = Workbook()
    ws = wb.active
    ws.title = "Master_PL"
    ws.append(["Account", "Account description", "L1 - BS/PL", "L2", "L3", "L4"])
    ws.append(["81910", "Net sales", "PL", "Revenue", "Sales", "Net sales"])
    p = _P(tmp_path) / "pl_only.xlsx"
    wb.save(str(p))
    return p


def _make_two_sheet_no_entity_xlsx(tmp_path) -> "Path":
    """Two-sheet workbook WITHOUT Entity column."""
    from openpyxl import Workbook
    from pathlib import Path as _P
    wb = Workbook()
    ws_bs = wb.active
    ws_bs.title = "Master_BS"
    ws_bs.append(["Account", "Account description", "L1", "L2", "L3", "L4"])
    ws_bs.append(["16100", "Trade receivables", "BS", "Assets", "Current", "Trade AR"])
    ws_pl = wb.create_sheet("Master_PL")
    ws_pl.append(["Account", "Account description", "L1 - BS/PL", "L2", "L3", "L4"])
    ws_pl.append(["81910", "Net sales", "PL", "Revenue", "Sales", "Net sales"])
    p = _P(tmp_path) / "two_sheet_no_entity.xlsx"
    wb.save(str(p))
    return p


# =========================================================================== #
# Entity-less template tests
# =========================================================================== #

def test_read_bs_pl_master_no_entity_column(tmp_path):
    """read_bs_pl_master on an Entity-less two-sheet workbook omits Entity col."""
    p = _make_two_sheet_no_entity_xlsx(tmp_path)
    df = read_bs_pl_master(p)
    assert "Account" in df.columns
    assert "Account description" in df.columns
    assert "L1" in df.columns
    assert "Entity" not in df.columns
    assert len(df) == 2          # 1 BS row + 1 PL row
    assert set(df["L1"]) == {"BS", "PL"}


def test_read_bs_pl_master_single_sheet_bs(tmp_path):
    """statement='bs' reads only Master_BS; PL rows absent; no 'L1 - BS/PL' col."""
    p = _make_bs_xlsx(tmp_path)
    df = read_bs_pl_master(p, statement="bs")
    assert len(df) == 2
    assert set(df["L1"]) == {"BS"}
    assert "L1 - BS/PL" not in df.columns


def test_read_bs_pl_master_single_sheet_pl_renames_l1(tmp_path):
    """statement='pl' reads Master_PL and renames 'L1 - BS/PL' → 'L1'."""
    p = _make_pl_xlsx(tmp_path)
    df = read_bs_pl_master(p, statement="pl")
    assert len(df) == 1
    assert "L1" in df.columns
    assert "L1 - BS/PL" not in df.columns
    assert df["L1"].iloc[0] == "PL"


def test_build_mapping_frames_with_entity_prefix(tmp_path):
    """entity_prefix='02' injected externally → account_number_group starts '02'.

    Key formula: entity_prefix(2) + zfill(account_number, 6) = 8-char key.
    E.g. prefix='02', account='16100' → '02' + '016100' = '02016100'.
    """
    raw = pd.DataFrame({
        "Account": ["16100", "81910"],
        "Account description": ["Trade AR", "Net sales"],
        "L1": ["BS", "PL"],
        "L2": ["Assets", "Revenue"],
        "L3": ["Current", "Sales"],
        "L4": ["Trade AR", "Net sales"],
    })
    # Session is None: entity_prefix path skips dim_legal_entity lookup entirely.
    frames, meta = build_mapping_frames(raw, None, [2024], entity_prefix="02")
    assert len(frames) == 2
    assert frames.loc[0, "account_number_group"] == "02016100"
    assert frames.loc[1, "account_number_group"] == "02081910"
    assert all(str(v).startswith("02") for v in frames["account_number_group"])
    assert meta["skipped_rows"] == 0
    assert meta["skipped_entities"] == []


def test_build_mapping_frames_entity_prefix_multi_fy():
    """entity_prefix path replicates rows across all requested fiscal years."""
    raw = pd.DataFrame({
        "Account": ["16100"],
        "Account description": ["Trade AR"],
        "L1": ["BS"],
        "L2": ["Assets"],
        "L3": ["Current"],
        "L4": ["Trade AR"],
    })
    frames, meta = build_mapping_frames(raw, None, [2023, 2024], entity_prefix="01")
    assert len(frames) == 2
    assert set(frames["fiscal_year"].astype(int)) == {2023, 2024}
    assert all(str(v).startswith("01") for v in frames["account_number_group"])


def test_key_equivalence_entity_prefix_vs_entity_column():
    """External entity_prefix='01' produces the same account_number_group as the
    legacy Entity-column path when both resolve to prefix '01' for the same account.

    Worked example: account='16100', prefix='01'
      → account_number_group = '01' + zfill('16100', 6) = '01' + '016100' = '01016100'
    """
    account = "16100"
    expected_key = "01016100"   # '01' + '016100' (zfill 6)

    # New path: no Entity column, prefix supplied externally.
    raw_no_entity = pd.DataFrame({
        "Account": [account],
        "Account description": ["Trade AR"],
        "L1": ["BS"],
        "L2": ["Assets"],
        "L3": ["Current"],
        "L4": ["Trade AR"],
    })
    frames_new, _ = build_mapping_frames(
        raw_no_entity, None, [2024], entity_prefix="01"
    )

    # Legacy path: Entity column resolved via dim_legal_entity.
    raw_with_entity = pd.DataFrame({
        "Entity": ["Atlas"],
        "Account": [account],
        "Account description": ["Trade AR"],
        "L1": ["BS"],
        "L2": ["Assets"],
        "L3": ["Current"],
        "L4": ["Trade AR"],
    })

    class _LegacySession:
        """Returns Atlas → prefix '01' from dim_legal_entity."""
        def execute(self, *_args, **_kwargs):
            class _R:
                def fetchall(self):
                    return [("Atlas", "Atlas GmbH", "01")]
            return _R()

    frames_legacy, _ = build_mapping_frames(
        raw_with_entity, _LegacySession(), [2024]
    )

    assert frames_new.loc[0, "account_number_group"] == expected_key
    assert frames_legacy.loc[0, "account_number_group"] == expected_key


def test_bs_pl_replace_modes():
    assert BS_PL_REPLACE_MODES == frozenset({"replace", "append"})


def test_filter_mapping_append_only_skips_existing():
    raw = _raw_bs_pl()
    lookup = {"Atlas": "01"}
    resolved, _ = resolve_entity_prefixes(raw, lookup)
    prof = bs_pl_master_profile(list(raw.columns))
    mapping_df = apply_account_mapping(resolved, prof, fiscal_year=2024)

    class _Session:
        def execute(self, stmt, params=None):
            sql = str(stmt)

            class _R:
                def __init__(self, rows):
                    self._rows = rows

                def fetchall(self):
                    return self._rows

            if "account_number_group = ANY" in sql and "entity_prefix" not in sql:
                return _R([("01016100", 2024)])
            return _R([])

    filtered = filter_mapping_append_only(mapping_df, _Session(), [2024])
    assert len(filtered) == 1
    assert filtered.iloc[0]["account_number_group"] == "01081910"


def test_preview_stats_append_reports_skip_count():
    raw = _raw_bs_pl()
    lookup = {"Atlas": "01"}
    resolved, _ = resolve_entity_prefixes(raw, lookup)
    prof = bs_pl_master_profile(list(raw.columns))
    mapping_df = apply_account_mapping(resolved, prof, fiscal_year=2024)

    class _Session:
        def execute(self, stmt, params=None):
            sql = str(stmt)

            class _R:
                def __init__(self, rows):
                    self._rows = rows

                def fetchall(self):
                    return self._rows

            if "dim_legal_entity" in sql:
                return _R([("Atlas", "Atlas GmbH", "01")])
            if "entity_prefix = ANY" in sql and "ORDER BY" not in sql:
                return _R([("01016100", 2024), ("01081910", 2024)])
            if "account_number_group = ANY" in sql:
                return _R([("01016100", 2024)])
            return _R([])

    stats = preview_stats(
        mapping_df,
        _Session(),
        [2024],
        1,
        1,
        replace_mode="append",
    )
    assert stats["would_insert"] == 1
    assert stats["would_skip"] == 1
    assert stats["would_delete"] == 0


def test_preview_stats_replace_reports_delete_count():
    raw = _raw_bs_pl()
    lookup = {"Atlas": "01"}
    resolved, _ = resolve_entity_prefixes(raw, lookup)
    prof = bs_pl_master_profile(list(raw.columns))
    mapping_df = apply_account_mapping(resolved, prof, fiscal_year=2024)

    class _Session:
        def execute(self, stmt, params=None):
            sql = str(stmt)

            class _R:
                def __init__(self, rows):
                    self._rows = rows

                def fetchall(self):
                    return self._rows

            if "dim_legal_entity" in sql:
                return _R([("Atlas", "Atlas GmbH", "01")])
            if "entity_prefix = ANY" in sql and "ORDER BY" not in sql:
                return _R(
                    [
                        ("01016100", 2024),
                        ("01081910", 2024),
                        ("01999999", 2024),
                    ]
                )
            if "fact_gl_line" in sql:
                return _R([])
            if "gl_account_id, account_name" in sql:
                return _R(
                    [
                        ("01999999", 2024, "01999999", "Orphan", "BS", "Assets", "Cash"),
                    ]
                )
            if "account_number_group = ANY" in sql:
                return _R([("01016100", 2024), ("01081910", 2024)])
            return _R([])

    stats = preview_stats(
        mapping_df,
        _Session(),
        [2024],
        1,
        1,
        replace_mode="replace",
    )
    assert stats["would_delete"] == 1
    assert len(stats["delete_details"]) == 1
    assert stats["delete_details"][0]["account_number_group"] == "01999999"
