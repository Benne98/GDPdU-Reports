"""Shared databook master workbook and mapping file paths."""
from __future__ import annotations

import re
import os
import tempfile
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill
from openpyxl.utils import get_column_letter

PROJECT_ROOT = Path(__file__).resolve().parent
DESKTOP_DIR = PROJECT_ROOT / "Desktop"
SORT_FILES_DIR = DESKTOP_DIR / "Sort_files"
MAPPINGFILES_DIR = DESKTOP_DIR / "Mappingfiles"
BACKEND_DIR = PROJECT_ROOT / "backend"
TEMPLATES_FDD_DIR = BACKEND_DIR / "templates" / "fdd"

MASTER_WORKBOOK = DESKTOP_DIR / "BS_PL_Master.xlsx"
MASTER_WORKBOOK_STR = str(MASTER_WORKBOOK)

CF_NA_L3_ORDER_FILE = str(TEMPLATES_FDD_DIR / "CF_NA_L3_order.xlsx")


def _first_existing_path(*candidates: Path) -> str:
    for path in candidates:
        if path.is_file():
            return str(path)
    return str(candidates[-1])


PL_RECON_MAPPING_FILE = _first_existing_path(
    TEMPLATES_FDD_DIR / "PL_recon_Mapping.xlsx",
    SORT_FILES_DIR / "PL_recon_Mapping.xlsx",
    DESKTOP_DIR / "PL_recon_Mapping.xlsx",
)
BS_RECON_MAPPING_FILE = _first_existing_path(
    TEMPLATES_FDD_DIR / "BS_recon_Mapping.xlsx",
    SORT_FILES_DIR / "BS_recon_Mapping.xlsx",
    DESKTOP_DIR / "BS_recon_Mapping.xlsx",
)
# Cashflow L4 order under Financial result — PL recon mapping has L3+L4 (not SuSa Kontenmapping).
PL_KONTOMAPPING_FILE = PL_RECON_MAPPING_FILE

MASTER_SHEET_ORDER = ("Master_BS", "Master_PL")

GROUP_BS_RECON_SHEET = "BS_Reconciliation"
GROUP_PL_RECON_SHEET = "PL_Reconciliation"
ENTITY_RECON_BS_SUFFIX = "_BS_Reconciliation"
ENTITY_RECON_PL_SUFFIX = "_PL_Reconciliation"
_EXCEL_SHEET_INVALID_RE = re.compile(r"[\[\]\:\*\?/\\]")

DATABOOK_SHEET_ORDER = (
    GROUP_BS_RECON_SHEET,
    GROUP_PL_RECON_SHEET,
    "BS_Bucket",
    "Lead_BS",
    "Lead_IS",
    "Working_Capital",
    "Cashflow",
)

_FTE_SHEET_RE = re.compile(r"^FTE[_\s]", re.I)
_FY_IN_NAME_RE = re.compile(r"(?:FY|YTD)?(19|20)\d{2}")
_LEGACY_FTE_OUTPUT_RE = re.compile(r"^Personnel\s+\d{4}$", re.I)

OPOS_REPORT_SHEETS = frozenset(
    {
        "Trade debtors aging detail",
        "Trade debtors aging",
        "Trade creditors aging detail",
        "Trade creditors aging",
    }
)

FA_ROLLF_OUTPUT_SHEET = "FA roll forward"

# Fast Track tab groups — left→right; shades of theme navy/blue (distinguishable, not rainbow).
SHEET_GROUP_SPECS: tuple[dict[str, str], ...] = (
    {"key": "executive", "label": "Executive Summary", "color": "1E3A5F"},
    {"key": "earnings", "label": "Earnings", "color": "2E5A8A"},
    {"key": "financial", "label": "Assets", "color": "3D7AB5"},
    {"key": "liquidity", "label": "Liquidity", "color": "4A90C8"},
    {"key": "appendix", "label": "Appendix", "color": "5B6B82"},
    {"key": "source", "label": "Source", "color": "6B7C93"},
)

SECTION_SHEET_PREFIX = "==> "
_SECTION_FILL_ROWS = 80
_SECTION_FILL_COLS = 40
_GROUP_RECON_BLOCK_TITLE_ROW = 7
_GROUP_RECON_SKIP_TITLES = frozenset(
    {
        "aggregated",
        "consolidation",
        "ic eliminations",
        "difference",
        "financial statements",
        "reported",
        "adjustments",
    }
)

_REVENUE_TABLE_BASES = frozenset(
    {
        "general sales table",
        "hierarchy_report",
        "hierarchy report",
        "pvm",
        "top",
        "churn",
    }
)
_REVENUE_GRAPH_BASES = frozenset(
    {
        "bubble",
        "bubble scatter",
        "bubble scatter plot",
        "margin analyses",
        "gp bridge",
        "np bridge",
        "arr bridge",
    }
)
_REVENUE_EXEC_BASES = frozenset(
    {
        "vertical bars",
        "vertical_bars",
        "horizontal bars",
        "horizontal_bars",
        "hierarchy",
        "breakdown",
        "net sales breakdown",
        "gross sales breakdown",
        "net sales hierarchy",
        "gross sales hierarchy",
    }
)


def open_batch_workbook(source_path: str | Path | None = None):
    """Load an existing workbook unchanged, or create a clean batch workbook."""
    if source_path:
        path = Path(source_path).expanduser().resolve()
        if not path.is_file():
            raise FileNotFoundError(f"Master workbook not found: {path}")
        return load_workbook(path)
    wb = Workbook()
    wb.active.title = "Fast Track"
    return wb


def atomic_save_workbook(wb, output_path: str | Path) -> Path:
    """Persist a workbook once via a same-directory temporary file and rename."""
    target = Path(output_path).expanduser().resolve()
    target.parent.mkdir(parents=True, exist_ok=True)
    fd, temp_name = tempfile.mkstemp(
        prefix=f".{target.stem}.",
        suffix=target.suffix or ".xlsx",
        dir=str(target.parent),
    )
    os.close(fd)
    temp = Path(temp_name)
    try:
        wb.save(temp)
        os.replace(temp, target)
    except Exception:
        temp.unlink(missing_ok=True)
        raise
    return target


def _header_columns(ws) -> dict[str, int]:
    return {
        str(ws.cell(1, col).value or "").strip(): col
        for col in range(1, (ws.max_column or 1) + 1)
    }


def derive_entity_order_by_latest_revenue(wb) -> list[str]:
    """Order Master_PL entities by latest FY net-sales/revenue value, descending."""
    if "Master_PL" not in wb.sheetnames:
        return []
    ws = wb["Master_PL"]
    headers = _header_columns(ws)
    entity_col = headers.get("Entity")
    if not entity_col:
        return []
    fy_columns: list[tuple[int, int]] = []
    for header, col in headers.items():
        match = re.fullmatch(r"FY\s*(\d{2,4})A?", header, re.I)
        if match:
            year = int(match.group(1))
            fy_columns.append((year + (2000 if year < 100 else 0), col))
    if not fy_columns:
        return []
    latest_col = max(fy_columns)[1]
    hierarchy_cols = [
        col
        for header, col in headers.items()
        if re.fullmatch(r"L[1-6](?:\s*-\s*BS/PL)?", header, re.I)
    ]
    values: dict[str, float] = {}
    first_seen: list[str] = []
    for row in range(2, (ws.max_row or 1) + 1):
        entity = str(ws.cell(row, entity_col).value or "").strip()
        if not entity or entity.lower() in {"entity", "consolidation"}:
            continue
        if entity not in values:
            values[entity] = 0.0
            first_seen.append(entity)
        labels = [str(ws.cell(row, col).value or "").strip().lower() for col in hierarchy_cols]
        if not any(label in {"net sales", "revenue", "sales"} for label in labels):
            continue
        raw = ws.cell(row, latest_col).value
        if isinstance(raw, str) and raw.startswith("="):
            continue
        try:
            values[entity] += float(raw or 0)
        except (TypeError, ValueError):
            continue
    ranked = sorted(first_seen, key=lambda entity: (-values[entity], first_seen.index(entity)))
    return ranked


def sanitize_entity_recon_sheet_name(entity: str, kind: str) -> str:
    """Build Excel-safe per-entity recon tab name ({Entity}_BS|PL_Reconciliation)."""
    suffix = ENTITY_RECON_BS_SUFFIX if str(kind).strip().lower() == "bs" else ENTITY_RECON_PL_SUFFIX
    stem = _EXCEL_SHEET_INVALID_RE.sub("", str(entity or "").strip()) or "Entity"
    max_stem = max(31 - len(suffix), 1)
    return f"{stem[:max_stem]}{suffix}"


def is_entity_recon_sheet(name: str) -> bool:
    n = str(name or "")
    if n in {GROUP_BS_RECON_SHEET, GROUP_PL_RECON_SHEET}:
        return False
    return n.endswith(ENTITY_RECON_BS_SUFFIX) or n.endswith(ENTITY_RECON_PL_SUFFIX)


def entity_recon_sheet_names(entities: list[str]) -> list[str]:
    """Per entity: BS then PL (entity_order sequence)."""
    out: list[str] = []
    for entity in entities:
        out.append(sanitize_entity_recon_sheet_name(entity, "bs"))
        out.append(sanitize_entity_recon_sheet_name(entity, "pl"))
    return out


def remove_stale_entity_recon_sheets(wb, entities: list[str]) -> None:
    """Drop per-entity recon tabs that are no longer in entity_order."""
    keep = set(entity_recon_sheet_names(entities))
    for sn in list(wb.sheetnames):
        if is_entity_recon_sheet(sn) and sn not in keep:
            del wb[sn]


def collect_entity_recon_sheets_ordered(
    existing: list[str],
    *,
    entity_order: list[str] | None = None,
) -> list[str]:
    """BS before PL per entity; prefer Group-Recon block order when provided."""
    entity_first_idx: dict[str, int] = {}
    sheets_by_entity: dict[str, dict[str, str]] = {}
    for idx, name in enumerate(existing):
        if not is_entity_recon_sheet(name):
            continue
        if name.endswith(ENTITY_RECON_BS_SUFFIX):
            entity = name[: -len(ENTITY_RECON_BS_SUFFIX)]
            kind = "bs"
        else:
            entity = name[: -len(ENTITY_RECON_PL_SUFFIX)]
            kind = "pl"
        entity_first_idx.setdefault(entity, idx)
        sheets_by_entity.setdefault(entity, {})[kind] = name

    if entity_order:
        ordered_entities: list[str] = []
        seen: set[str] = set()
        for entity in entity_order:
            if entity in sheets_by_entity and entity not in seen:
                ordered_entities.append(entity)
                seen.add(entity)
        for entity in sorted(entity_first_idx.keys(), key=lambda e: entity_first_idx[e]):
            if entity not in seen:
                ordered_entities.append(entity)
    else:
        ordered_entities = sorted(entity_first_idx.keys(), key=lambda e: entity_first_idx[e])

    out: list[str] = []
    for entity in ordered_entities:
        pair = sheets_by_entity.get(entity, {})
        if "bs" in pair:
            out.append(pair["bs"])
        if "pl" in pair:
            out.append(pair["pl"])
    return out


def section_sheet_name(group_label: str) -> str:
    raw = f"{SECTION_SHEET_PREFIX}{str(group_label).strip()}"
    return raw[:31]


def is_section_sheet(name: str) -> bool:
    return str(name or "").startswith(SECTION_SHEET_PREFIX)


def entity_order_from_group_recon(wb) -> list[str]:
    """Entity sequence as shown in Group PL/BS Reconciliation block titles."""
    existing_entities = {
        (
            name[: -len(ENTITY_RECON_BS_SUFFIX)]
            if name.endswith(ENTITY_RECON_BS_SUFFIX)
            else name[: -len(ENTITY_RECON_PL_SUFFIX)]
        )
        for name in wb.sheetnames
        if is_entity_recon_sheet(name)
    }
    if not existing_entities:
        return []

    for sheet in (GROUP_PL_RECON_SHEET, GROUP_BS_RECON_SHEET):
        if sheet not in wb.sheetnames:
            continue
        ws = wb[sheet]
        titles: list[str] = []
        seen_titles: set[str] = set()
        for cc in range(1, (ws.max_column or 0) + 1):
            raw = ws.cell(_GROUP_RECON_BLOCK_TITLE_ROW, cc).value
            if not isinstance(raw, str):
                continue
            title = raw.strip()
            if not title or title in seen_titles:
                continue
            seen_titles.add(title)
            titles.append(title)

        ordered: list[str] = []
        for title in titles:
            if title.lower() in _GROUP_RECON_SKIP_TITLES:
                continue
            if title in existing_entities:
                ordered.append(title)
        if ordered:
            return ordered
    return []


def _source_group_rank(name: str) -> int:
    n = str(name)
    if is_opos_ar_ap_source_sheet(n):
        return 0
    if is_fa_rollf_source_sheet(n):
        return 1
    if is_fte_source_sheet(n):
        return 2
    return 3


def _source_sort_key(sheet_name: str) -> tuple:
    name = str(sheet_name)
    return (_source_group_rank(name), name.lower())


def is_opos_report_sheet(name: str) -> bool:
    return str(name) in OPOS_REPORT_SHEETS


def is_opos_ar_ap_source_sheet(name: str) -> bool:
    n = str(name)
    return n.startswith("__SOURCE__AR_") or n.startswith("__SOURCE__AP_")


def is_fa_rollf_output_sheet(name: str) -> bool:
    lower = str(name).strip().lower()
    return lower == FA_ROLLF_OUTPUT_SHEET.lower() or "fixed assets roll" in lower


def is_fa_rollf_source_sheet(name: str) -> bool:
    n = str(name)
    return n.startswith("__SOURCE__FA_")


def is_fte_output_sheet(name: str) -> bool:
    n = str(name).strip().lower()
    return n == "fte development" or bool(_LEGACY_FTE_OUTPUT_RE.match(str(name)))


def is_fte_source_sheet(name: str) -> bool:
    n = str(name)
    if is_fte_output_sheet(n):
        return False
    if re.match(r"^FTE_", n):
        return True
    if n.startswith("__SOURCE__FTE_"):
        return True
    # Legacy names from before the FTE_ prefix (cleanup on re-run).
    if not n.startswith("__SOURCE__"):
        return False
    if is_fa_rollf_source_sheet(n) or is_opos_ar_ap_source_sheet(n) or is_sales_source_sheet(n):
        return False
    rest = n[len("__SOURCE__") :]
    return bool(re.match(r"^(FY|YTD)\d{2}A$", rest))


def clear_opos_workbook_sheets(wb) -> None:
    """Remove only OPOS aging report tabs and AR/AP source sheets (re-run cleanup)."""
    for sn in list(wb.sheetnames):
        name = str(sn)
        if name in OPOS_REPORT_SHEETS:
            del wb[name]
            continue
        if is_opos_ar_ap_source_sheet(name):
            del wb[name]


def clear_fa_rollf_workbook_sheets(wb) -> None:
    """Remove only Fixed Assets rollforward output and FA period source sheets."""
    for sn in list(wb.sheetnames):
        name = str(sn)
        if is_fa_rollf_output_sheet(name) or is_fa_rollf_source_sheet(name):
            del wb[name]


def clear_fte_workbook_sheets(wb) -> None:
    """Remove only FTE Development output and FTE period source sheets."""
    for sn in list(wb.sheetnames):
        name = str(sn)
        if is_fte_output_sheet(name) or is_fte_source_sheet(name):
            del wb[name]


def _strand_sort_key(sheet_name: str) -> tuple:
    name = str(sheet_name)
    lower = name.lower()
    if lower == FA_ROLLF_OUTPUT_SHEET.lower() or "fixed assets roll" in lower or "fa roll forward" in lower:
        return (0, name)
    if _FTE_SHEET_RE.match(name) or "fte development" in lower:
        years = [int(y) for y in _FY_IN_NAME_RE.findall(name)]
        return (1, min(years) if years else 9999, name)
    if "debtor" in lower:
        return (2, 0 if "summary" in lower or not lower.endswith("detail") else 1, name)
    if "creditor" in lower:
        return (3, 0 if "summary" in lower or not lower.endswith("detail") else 1, name)
    return (4, name)


def is_workbook_source_sheet(name: str) -> bool:
    """True for hidden data sources (__SOURCE__*, legacy FTE_*), not strand outputs."""
    name = str(name)
    if name.startswith("__SOURCE__"):
        return True
    if name.lower() == "fte development":
        return False
    return bool(re.match(r"^FTE_", name))


def is_sales_source_sheet(name: str) -> bool:
    n = str(name)
    return n in {"__SOURCE__", "__SOURCE__Sales"}


def _base_sheet_token(name: str) -> str:
    """Strip trailing _N suffix used by get_next_sheet_name_from_wb."""
    token = str(name).strip()
    match = re.fullmatch(r"(.+)_(\d+)$", token)
    if match:
        return match.group(1).strip().lower()
    return token.lower()


def is_revenue_table_sheet(name: str) -> bool:
    return _base_sheet_token(name) in _REVENUE_TABLE_BASES


def is_revenue_graph_sheet(name: str) -> bool:
    return _base_sheet_token(name) in _REVENUE_GRAPH_BASES


def is_vertical_bars_sheet(name: str) -> bool:
    token = _base_sheet_token(name)
    if token in _REVENUE_EXEC_BASES:
        return True
    return token.endswith("sales breakdown") or token.endswith("sales hierarchy")


# Backward-compatible alias
is_horizontal_bars_sheet = is_vertical_bars_sheet


def is_earnings_source_sheet(name: str) -> bool:
    return is_sales_source_sheet(name) or is_fte_source_sheet(name)


def is_financial_source_sheet(name: str) -> bool:
    return is_fa_rollf_source_sheet(name) or is_opos_ar_ap_source_sheet(name)


def _apply_tab_color(ws, color: str) -> None:
    rgb = str(color or "").strip().lstrip("#").upper()
    if len(rgb) == 8 and rgb.startswith("FF"):
        rgb = rgb[2:]
    if len(rgb) == 6:
        ws.sheet_properties.tabColor = rgb


def style_group_section_sheet(ws, *, group_label: str, color: str) -> None:
    """Cover divider: full blue canvas, white Finssentials title + group subtitle."""
    fill = PatternFill("solid", fgColor=str(color).lstrip("#").upper()[-6:])
    title_font = Font(name="Calibri", size=24, bold=True, color="FFFFFF")
    subtitle_font = Font(name="Calibri", size=12, color="FFFFFF")
    left = Alignment(horizontal="left", vertical="center")

    for row in range(1, _SECTION_FILL_ROWS + 1):
        ws.row_dimensions[row].height = 18
        for col in range(1, _SECTION_FILL_COLS + 1):
            cell = ws.cell(row, col)
            cell.value = None
            cell.fill = fill
            cell.border = Border()

    for col in range(1, _SECTION_FILL_COLS + 1):
        letter = get_column_letter(col)
        ws.column_dimensions[letter].width = 14 if col > 1 else 4

    ws.row_dimensions[2].height = 36
    ws.row_dimensions[3].height = 22
    title = ws.cell(2, 2, "Finssentials")
    title.font = title_font
    title.alignment = left
    title.fill = fill
    subtitle = ws.cell(3, 2, str(group_label).strip())
    subtitle.font = subtitle_font
    subtitle.alignment = left
    subtitle.fill = fill

    ws.sheet_view.showGridLines = False
    _apply_tab_color(ws, color)


def ensure_group_section_sheets(wb) -> dict[str, str]:
    """Create/refresh ==> Group divider sheets. Returns group_key → sheet name."""
    names: dict[str, str] = {}
    for spec in SHEET_GROUP_SPECS:
        sheet_name = section_sheet_name(spec["label"])
        names[spec["key"]] = sheet_name
        if sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
        else:
            ws = wb.create_sheet(sheet_name)
        style_group_section_sheet(ws, group_label=spec["label"], color=spec["color"])
    return names


def _earnings_sort_key(name: str) -> tuple:
    lower = name.lower()
    if lower == "lead_is":
        return (0, name.lower())
    if is_revenue_table_sheet(name):
        return (1, name.lower())
    if is_revenue_graph_sheet(name):
        return (2, name.lower())
    if is_fte_output_sheet(name):
        years = [int(y) for y in _FY_IN_NAME_RE.findall(name)]
        return (3, min(years) if years else 9999, name.lower())
    if is_earnings_source_sheet(name):
        # Sales source before personnel / FTE sources
        if is_sales_source_sheet(name):
            return (4, 0, name.lower())
        return (4, 1, name.lower())
    return (5, name.lower())


def _financial_sort_key(name: str) -> tuple:
    lower = name.lower()
    if lower == "lead_bs":
        return (0, name.lower())
    if lower in {"bs_bucket", "bucket_bs"}:
        return (1, name.lower())
    if lower in {"working_capital", "working capital"}:
        return (2, name.lower())
    if is_fa_rollf_output_sheet(name):
        return (3, name.lower())
    if is_opos_report_sheet(name):
        return (4, *_strand_sort_key(name))
    if is_financial_source_sheet(name):
        return (5, *_source_sort_key(name))
    return (6, name.lower())


def _classify_sheet_group(name: str) -> str | None:
    """Map a content sheet to a group key (section sheets handled separately)."""
    if is_section_sheet(name) or name == "Fast Track":
        return None
    lower = name.strip().lower()

    if name in MASTER_SHEET_ORDER or is_workbook_source_sheet(name):
        return "source"

    if is_vertical_bars_sheet(name):
        return "executive"

    if lower == "lead_is" or is_revenue_table_sheet(name) or is_revenue_graph_sheet(name):
        return "earnings"
    if is_fte_output_sheet(name):
        return "earnings"

    if lower in {"lead_bs", "bs_bucket", "bucket_bs", "working_capital", "working capital"}:
        return "financial"
    if is_fa_rollf_output_sheet(name) or is_opos_report_sheet(name):
        return "financial"

    if lower == "cashflow":
        return "liquidity"

    if (
        name in {GROUP_BS_RECON_SHEET, GROUP_PL_RECON_SHEET}
        or is_entity_recon_sheet(name)
    ):
        return "appendix"

    # Unknown leftovers go to Appendix.
    return "appendix"


def _move_sheets_to_order(wb, ordered: list[str]) -> None:
    for target_idx, name in enumerate(ordered):
        if name not in wb.sheetnames:
            continue
        while wb.sheetnames.index(name) > target_idx:
            wb.move_sheet(wb[name], offset=-1)


def apply_group_tab_colors(wb) -> None:
    color_by_section = {
        section_sheet_name(spec["label"]): spec["color"] for spec in SHEET_GROUP_SPECS
    }
    color_by_group = {spec["key"]: spec["color"] for spec in SHEET_GROUP_SPECS}
    for name in wb.sheetnames:
        if name in color_by_section:
            _apply_tab_color(wb[name], color_by_section[name])
            continue
        group = _classify_sheet_group(name)
        if group:
            _apply_tab_color(wb[name], color_by_group[group])


def remove_placeholder_fast_track_sheet(wb) -> None:
    """Drop the empty bootstrap sheet created when no master workbook is loaded."""
    if "Fast Track" not in wb.sheetnames:
        return
    ws = wb["Fast Track"]
    if ws.max_row == 1 and ws.max_column == 1 and ws.cell(1, 1).value in (None, ""):
        del wb["Fast Track"]


def reorder_workbook_sheets(wb) -> None:
    """Group Fast Track tabs: section dividers + blue tab colors + fixed left→right order."""
    remove_placeholder_fast_track_sheet(wb)
    section_names = ensure_group_section_sheets(wb)

    existing = list(wb.sheetnames)
    entity_order = entity_order_from_group_recon(wb)
    buckets: dict[str, list[str]] = {spec["key"]: [] for spec in SHEET_GROUP_SPECS}

    for name in existing:
        if is_section_sheet(name):
            continue
        group = _classify_sheet_group(name)
        if group:
            buckets[group].append(name)

    buckets["executive"].sort(key=lambda n: (0 if is_vertical_bars_sheet(n) else 1, n.lower()))
    buckets["earnings"].sort(key=_earnings_sort_key)
    buckets["financial"].sort(key=_financial_sort_key)
    buckets["liquidity"].sort(key=lambda n: (0 if n.lower() == "cashflow" else 1, n.lower()))

    source_masters = [n for n in MASTER_SHEET_ORDER if n in buckets["source"]]
    source_rest = sorted(
        [n for n in buckets["source"] if n not in source_masters],
        key=_source_sort_key,
    )
    buckets["source"] = source_masters + source_rest

    appendix_fixed = [
        n
        for n in (GROUP_BS_RECON_SHEET, GROUP_PL_RECON_SHEET)
        if n in buckets["appendix"]
    ]
    appendix_entities = collect_entity_recon_sheets_ordered(
        buckets["appendix"],
        entity_order=entity_order,
    )
    appendix_used = set(appendix_fixed) | set(appendix_entities)
    appendix_other_outputs = sorted(
        [
            n
            for n in buckets["appendix"]
            if n not in appendix_used and not is_workbook_source_sheet(n)
        ],
        key=str.lower,
    )
    buckets["appendix"] = (
        appendix_fixed
        + appendix_entities
        + appendix_other_outputs
    )

    # Ensure sources are visible under ==> Source
    for name in buckets["source"]:
        if name in wb.sheetnames:
            wb[name].sheet_state = "visible"

    ordered: list[str] = []
    for spec in SHEET_GROUP_SPECS:
        key = spec["key"]
        section = section_names[key]
        ordered.append(section)
        ordered.extend(buckets[key])

    # Preserve any unexpected sheets after Source (should be rare).
    for name in existing:
        if name not in ordered and not is_section_sheet(name):
            ordered.append(name)

    _move_sheets_to_order(wb, ordered)
    apply_group_tab_colors(wb)


def reorder_workbook_file(path: str | Path) -> None:
    """Load workbook at path, reorder sheets, save."""
    from openpyxl import load_workbook

    p = Path(path)
    if not p.is_file():
        return
    wb = load_workbook(p)
    reorder_workbook_sheets(wb)
    wb.save(p)
    wb.close()
