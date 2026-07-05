"""FTE Development — verformelte FTE/Payroll-Tabelle aus Personaltabellen (EURk)."""
from __future__ import annotations

import json
import os
import re
import sys
from collections import defaultdict
from typing import Any

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.dataframe import dataframe_to_rows

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
if BASE_DIR not in sys.path:
    sys.path.insert(0, BASE_DIR)

from funktionssammlung import (  # noqa: E402
    _get_sheet_header_map,
    build_output_file_path,
    build_source_sheet_name,
    build_sumifs_formula_body,
    ensure_output_writable,
    ensure_source_sheet_in_output,
    get_next_sheet_name_from_wb,
    refresh_source_sheet_from_df,
    source_range_ref,
    write_source_df_to_ws,
)

HELPER_FTE = "_fte_avg"
HELPER_PAYROLL = "_payroll"
HELPER_FY = "_fte_fy"


def normalize_config(cfg: dict[str, Any]) -> dict[str, Any]:
    out = dict(cfg)
    dims = list(out.get("dimensions") or [])
    if not dims and out.get("department_map"):
        dims = [
            {"source_col": "Bereich", "output_label": "Department"},
        ]
    if len(dims) < 1 or len(dims) > 3:
        raise ValueError("dimensions must contain 1 to 3 entries.")
    out["dimensions"] = dims
    metrics = list(out.get("preset_metrics") or ["fte", "payroll"])
    if not metrics:
        raise ValueError("At least one preset metric is required.")
    out["preset_metrics"] = metrics
    out["group_cols"] = [d["source_col"] for d in dims]
    out["group_labels"] = [d.get("output_label") or d["source_col"] for d in dims]
    out["formula_mode"] = bool(out.get("formula_mode", True))
    out["first_fy"] = int(out.get("first_fy") or 2022)
    out["last_fy"] = int(out.get("last_fy") or out["first_fy"])
    return out


def _parse_fy_year(label: str) -> int:
    m = re.search(r"(19|20)\d{2}", str(label))
    if not m:
        raise ValueError(f"Cannot parse FY from label: {label!r}")
    return int(m.group(0))


def _fy_col_header(year: int) -> str:
    return f"FY{str(year)[-2:]}A"


def _load_personaltable(path: str) -> pd.DataFrame:
    xl = pd.ExcelFile(path)
    return pd.read_excel(path, sheet_name=xl.sheet_names[0])


def merge_uploads(cfg: dict[str, Any]) -> pd.DataFrame:
    resolved = list(cfg.get("resolved_entity_year_files") or [])
    frames: list[pd.DataFrame] = []
    upload_mode = str(cfg.get("upload_mode") or "per_fy_grid")
    mapping = dict(cfg.get("fte_mapping") or {})
    year_col = mapping.get("year_col")

    for cell in resolved:
        paths = cell.get("paths") or []
        if not paths:
            continue
        df = _load_personaltable(paths[0])
        if upload_mode == "per_fy_grid":
            try:
                fy = _parse_fy_year(str(cell.get("fy_label") or ""))
            except ValueError:
                fy = int(cell.get("year") or 0)
            df = df.copy()
            df[HELPER_FY] = fy
        elif year_col and year_col in df.columns:
            df = df.copy()
            df[HELPER_FY] = pd.to_numeric(df[year_col], errors="coerce").fillna(0).astype(int)
        else:
            df = df.copy()
            df[HELPER_FY] = cfg["first_fy"]
        frames.append(df)
    if not frames:
        raise ValueError("No personaltable data loaded from uploads.")
    return pd.concat(frames, ignore_index=True)


def _col_letter(header_map: dict[str, int], col_name: str) -> str | None:
    from funktionssammlung import _normalize_header_name

    key = _normalize_header_name(col_name)
    idx = header_map.get(key)
    if not idx:
        for k, v in header_map.items():
            if key in k or k in key:
                idx = v
                break
    return get_column_letter(idx) if idx else None


def _cell_ref(sheet: str, col_letter: str, row: int) -> str:
    return f"'{sheet}'!{col_letter}{row}"


def compute_row_fte(row: pd.Series, cfg: dict[str, Any]) -> float:
    mapping = dict(cfg.get("fte_mapping") or {})
    tenure = str(cfg.get("fte_tenure_mode") or mapping.get("tenure_mode") or "months_col")
    pct_col = mapping.get("employment_pct_col") or "Beschäftigungsgrad"
    try:
        pct = float(row.get(pct_col) or 100)
    except (TypeError, ValueError):
        pct = 100.0
    if tenure == "entry_exit_dates":
        # simplified: use months col if present else 12
        months_col = mapping.get("months_col") or "Summe"
        try:
            months = float(row.get(months_col) or 12)
        except (TypeError, ValueError):
            months = 12.0
    else:
        months_col = mapping.get("months_col") or "Summe"
        try:
            months = float(row.get(months_col) or 0)
        except (TypeError, ValueError):
            months = 0.0
        if months <= 0:
            months = sum(int(row.get(m, 0) or 0) for m in [
                "Januar", "Februar", "März", "April", "Mai", "Juni",
                "Juli", "August", "September", "Oktober", "November", "Dezember",
            ])
    return months * pct / 100.0 / 12.0


def compute_row_payroll(row: pd.Series, cfg: dict[str, Any]) -> float:
    mapping = dict(cfg.get("payroll_mapping") or {})
    mode = str(cfg.get("fte_payroll_mode") or mapping.get("payroll_mode") or "sum_components")
    if mode == "total_col":
        col = mapping.get("total_col") or "Gesamtsumme"
        return float(row.get(col) or 0)
    if mode == "monthly_col":
        col = mapping.get("monthly_col") or "gehalt mon."
        months_col = dict(cfg.get("fte_mapping") or {}).get("months_col") or "Summe"
        try:
            months = float(row.get(months_col) or 12)
        except (TypeError, ValueError):
            months = 12.0
        return float(row.get(col) or 0) * months
    total = 0.0
    for col in mapping.get("component_cols") or ["Grundgehalt"]:
        total += float(row.get(col) or 0)
    social = mapping.get("social_col")
    if social:
        total += float(row.get(social) or 0)
    return total


def _write_helper_formulas(ws, source_name: str, cfg: dict[str, Any], n_rows: int) -> dict[str, int]:
    """Append _fte_avg and _payroll helper columns with Excel formulas."""
    header_map = _get_sheet_header_map(ws)
    base_max = ws.max_column or 1
    fte_col = base_max + 1
    pay_col = base_max + 2
    ws.cell(row=1, column=fte_col).value = HELPER_FTE
    ws.cell(row=1, column=pay_col).value = HELPER_PAYROLL

    mapping = dict(cfg.get("fte_mapping") or {})
    pay_map = dict(cfg.get("payroll_mapping") or {})
    pct_l = _col_letter(header_map, mapping.get("employment_pct_col") or "Beschäftigungsgrad")
    months_l = _col_letter(header_map, mapping.get("months_col") or "Summe")
    tenure = str(cfg.get("fte_tenure_mode") or "months_col")

    for r in range(2, n_rows + 2):
        if tenure == "months_col" and pct_l and months_l:
            fte_formula = (
                f"=IF({months_l}{r}=0,0,{months_l}{r}*{pct_l}{r}/100/12)"
            )
        elif pct_l and months_l:
            fte_formula = f"=IF({months_l}{r}=0,0,{months_l}{r}*{pct_l}{r}/100/12)"
        else:
            fte_formula = "=0"
        ws.cell(row=r, column=fte_col).value = fte_formula

        mode = str(cfg.get("fte_payroll_mode") or "sum_components")
        if mode == "total_col":
            tc = _col_letter(header_map, pay_map.get("total_col") or "Gesamtsumme")
            pay_formula = f"={tc}{r}" if tc else "=0"
        elif mode == "monthly_col":
            mc = _col_letter(header_map, pay_map.get("monthly_col") or "gehalt mon.")
            pay_formula = f"={mc}{r}*{months_l}{r}" if mc and months_l else "=0"
        else:
            parts = []
            for col in pay_map.get("component_cols") or ["Grundgehalt"]:
                cl = _col_letter(header_map, col)
                if cl:
                    parts.append(f"{cl}{r}")
            sc = pay_map.get("social_col")
            if sc:
                scl = _col_letter(header_map, sc)
                if scl:
                    parts.append(f"{scl}{r}")
            pay_formula = "=" + "+".join(parts) if parts else "=0"
        ws.cell(row=r, column=pay_col).value = pay_formula

    return {HELPER_FTE: fte_col, HELPER_PAYROLL: pay_col}


def _distinct_hierarchy(df: pd.DataFrame, group_cols: list[str]) -> list[tuple]:
    if not group_cols:
        return [()]
    sub = df[group_cols].fillna("").astype(str)
    return sorted({tuple(row) for row in sub.drop_duplicates().itertuples(index=False, name=None)})


def _metric_label(metric: str) -> str:
    return {
        "fte": "Average FTEs #",
        "payroll": "Payroll accounting",
        "avg_cost_per_fte": "Average cost per FTE",
    }.get(metric, metric)


def format_fte_report_excel(ws, cfg: dict[str, Any], fy_years: list[int], n_metric_rows: int) -> None:
    bold = Font(bold=True)
    eurk_fill = PatternFill(fill_type="solid", fgColor="F8FAFC")
    ws.cell(row=1, column=1).value = "EURk"
    ws.cell(row=1, column=1).font = bold
    for ci, y in enumerate(fy_years, start=2):
        c = ws.cell(row=2, column=ci, value=_fy_col_header(y))
        c.font = bold
        c.fill = eurk_fill
    for r in range(3, 3 + n_metric_rows):
        ws.cell(row=r, column=1).alignment = Alignment(horizontal="left")
        for ci in range(2, 2 + len(fy_years)):
            cell = ws.cell(row=r, column=ci)
            cell.alignment = Alignment(horizontal="right")
            label = str(ws.cell(row=r, column=1).value or "")
            if "FTE" in label and "#" in label:
                cell.number_format = "#,##0"
            else:
                cell.number_format = "#,##0;(#,##0)"


def _resolve_output_path(cfg: dict[str, Any]) -> str:
    path = build_output_file_path(cfg)
    custom = str(cfg.get("output_filename") or "").strip()
    if custom:
        return os.path.join(os.path.dirname(path), custom)
    return path


def build_fte_workbook(cfg: dict[str, Any]) -> str:
    cfg = normalize_config(cfg)
    df = merge_uploads(cfg)
    df[HELPER_FTE] = df.apply(lambda r: compute_row_fte(r, cfg), axis=1)
    df[HELPER_PAYROLL] = df.apply(lambda r: compute_row_payroll(r, cfg), axis=1)

    output_path = _resolve_output_path(cfg)
    ensure_output_writable(output_path)

    first_path = (cfg.get("resolved_entity_year_files") or [{}])[0]
    paths = first_path.get("paths") if isinstance(first_path, dict) else None
    cfg = dict(cfg)
    cfg["file_path"] = paths[0] if paths else ""
    cfg["sheet_name"] = cfg.get("sheet_name") or "Personaltabelle"

    source_sheet_name = ensure_source_sheet_in_output(cfg, output_path)
    wb = load_workbook(output_path)
    if source_sheet_name in wb.sheetnames:
        ws_src = wb[source_sheet_name]
        refresh_source_sheet_from_df(ws_src, df.drop(columns=[HELPER_FTE, HELPER_PAYROLL], errors="ignore"), cfg)
    else:
        ws_src = wb.active
        ws_src.title = source_sheet_name
        write_source_df_to_ws(ws_src, df.drop(columns=[HELPER_FTE, HELPER_PAYROLL], errors="ignore"), cfg)

    n_rows = len(df)
    helpers = _write_helper_formulas(ws_src, source_sheet_name, cfg, n_rows)
    header_map = _get_sheet_header_map(ws_src)

    report_name = get_next_sheet_name_from_wb(wb, cfg.get("base_sheet_name") or "FTE Development")
    if report_name in wb.sheetnames:
        del wb[report_name]
    ws = wb.create_sheet(report_name)

    fy_years = list(range(int(cfg["first_fy"]), int(cfg["last_fy"]) + 1))
    group_cols = list(cfg["group_cols"])
    hierarchy = _distinct_hierarchy(df, group_cols)

    row_ptr = 3
    metrics = list(cfg["preset_metrics"])
    source_header_map = {k: v for k, v in header_map.items()}

    fte_helper_col = get_column_letter(helpers[HELPER_FTE])
    pay_helper_col = get_column_letter(helpers[HELPER_PAYROLL])
    fte_rng = f"'{source_sheet_name}'!{fte_helper_col}2:{fte_helper_col}{n_rows + 1}"
    pay_rng = f"'{source_sheet_name}'!{pay_helper_col}2:{pay_helper_col}{n_rows + 1}"

    fy_col_name = HELPER_FY
    if fy_col_name not in df.columns:
        fy_col_name = mapping.get("year_col") if (mapping := dict(cfg.get("fte_mapping") or {})) else None

    for metric in metrics:
        ws.cell(row=row_ptr, column=1, value=_metric_label(metric)).font = Font(bold=True)
        row_ptr += 1
        for keys in hierarchy:
            label_parts = list(keys) if keys else ["Total"]
            ws.cell(row=row_ptr, column=1, value=" / ".join(label_parts))
            for ci, fy in enumerate(fy_years, start=2):
                base_pairs: list[tuple[str, str]] = []
                for gc, val in zip(group_cols, keys if keys else []):
                    gr = source_range_ref(source_sheet_name, source_header_map, gc)
                    base_pairs.append((gr, f'"{val}"'))
                if fy_col_name and fy_col_name in df.columns:
                    fyr = source_range_ref(source_sheet_name, source_header_map, fy_col_name)
                    base_pairs.append((fyr, str(fy)))
                sum_rng = fte_rng if metric == "fte" else pay_rng
                divide = metric != "fte"
                body = build_sumifs_formula_body(sum_rng, base_pairs, [[]], divide_by_1000=divide)
                if metric == "fte":
                    formula = f"=ROUND({body.lstrip('=')},0)"
                elif metric == "avg_cost_per_fte":
                    fte_body = build_sumifs_formula_body(fte_rng, base_pairs, [[]], divide_by_1000=False)
                    pay_body = build_sumifs_formula_body(pay_rng, base_pairs, [[]], divide_by_1000=True)
                    formula = f'=IF({fte_body.lstrip("=")}=0,"",-{pay_body.lstrip("=")}/{fte_body.lstrip("=")}*1000)'
                else:
                    formula = f"=-ABS({body.lstrip('=')})"
                ws.cell(row=row_ptr, column=ci).value = formula
            row_ptr += 1
        row_ptr += 1

    # PEX rows from GL
    pex = dict(cfg.get("pex_values") or {})
    if pex:
        ws.cell(row=row_ptr, column=1, value="Personnel expenses (GL)").font = Font(bold=True)
        row_ptr += 1
        pex_entities = sorted({k.rsplit("_", 1)[0] for k in pex if k.startswith("entity")})
        for ent_key in pex_entities or ["entity1"]:
            ei = int(re.search(r"\d+", ent_key).group()) - 1 if re.search(r"\d+", ent_key) else 0
            ws.cell(row=row_ptr, column=1, value=str(ent_key))
            for ci, fy in enumerate(fy_years, start=2):
                key = f"{ent_key}_FY{fy}"
                alt = f"entity{ei + 1}_FY{fy}"
                val = pex.get(key, pex.get(alt, pex.get(f"{ent_key}_FY{str(fy)[-2:]}", "")))
                if val not in (None, ""):
                    ws.cell(row=row_ptr, column=ci, value=-abs(float(val)))
            row_ptr += 1

    format_fte_report_excel(ws, cfg, fy_years, row_ptr - 3)
    wb.save(output_path)
    wb.close()
    return output_path


def main() -> None:
    config_path = sys.argv[1] if len(sys.argv) > 1 else ""
    if not config_path or not os.path.isfile(config_path):
        print(json.dumps({"success": False, "message": "Config path required"}))
        sys.exit(1)
    with open(config_path, encoding="utf-8") as f:
        cfg = json.load(f)
    try:
        out = build_fte_workbook(cfg)
        print(json.dumps({"success": True, "output_file": out, "output_path": out}))
    except Exception as exc:
        print(json.dumps({"success": False, "message": str(exc)}))
        sys.exit(1)


if __name__ == "__main__":
    main()
