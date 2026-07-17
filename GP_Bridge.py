import os
import re
import textwrap
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from typing import Dict, Any, List, Tuple, Optional, Set

import numpy as np
import pandas as pd
import matplotlib.pyplot as plt

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image as XLImage

import matplotlib.font_manager as fm
from matplotlib.colors import to_rgba


def _xl_image_from_path(path: str) -> XLImage:
    """Embed PNG bytes so the temp file can be deleted before workbook.save()."""
    return XLImage(BytesIO(Path(path).read_bytes()))


try:
    from funktionssammlung import apply_filters, compute_amount, recognized_value
except ImportError:  # pragma: no cover
    apply_filters = None  # type: ignore
    compute_amount = None  # type: ignore
    recognized_value = None  # type: ignore

try:
    from revenue_reconciliation import sales_basis_labels
except ImportError:  # pragma: no cover
    def sales_basis_labels(sales_basis=None):
        basis = "gross" if str(sales_basis or "").strip().lower() == "gross" else "net"
        metric = "Gross sales" if basis == "gross" else "Net sales"
        return type(
            "L",
            (),
            {
                "sales_basis": basis,
                "metric_label": metric,
                "bridge_sheet_title": "GP Bridge" if basis == "gross" else "NP Bridge",
                "bridge_table_title": "Gross Profit Bridge" if basis == "gross" else "Net Profit Bridge",
            },
        )()


# ==========================
# CONFIG
# ==========================
CONFIG: Dict[str, Any] = {
    "title": "Project Draft",
    "table": "Gross Profit Bridge",
    "company": "Draft AG",
    "subtitle_suffix": "by invoiced amounts",
    "period_label_prefix": "XX",

    "current_year": 2024,
    "current_month": 12,
    "fiscal_year_end_month": 7,
    "fiscal_year_end_day": 31,

    "period_mode": "FY",        # "FY" | "YTD" | "LTM"
    "calc_mode": "invoice",     # "invoice" | "accrual"
    "invoice_mapping_mode": "date",  # "date" | "year"
    "profit_mode": "cost",      # "cost" | "profit"
    "sales_basis": "net",

    "invoice_col": "Invoice Date",
    "start_col": "Contract Start Date",
    "end_col": "Contract End Date",

    # GP = Revenue - Cost  (or profit_col when profit_mode=profit)
    "revenue_col": "Contract Value After Discount",
    "cost_col": "Third-Party Contract Value",
    "profit_col": "",

    "filter": {"enabled": False, "col": "", "values": set()},
    "filters": {"enabled": False, "rules": []},

    "apply_fx": False,
    "fx_col": "Functional FX Rate",

    "bridge": {
        "dim_col": "Product Line",
        "top_n": 3,
        "other_label": "Other",
        "missing_token": "__MISSING__",

        # Combine Input-Other
        "combine_input_other": True,
        "input_other_tokens": {"Other"},
        "other_bucket_label": "Other (rest)",

        # Sortierung
        "sort_mode": "value",  # "abs" | "value" | "pos_then_neg_abs"
    },

    # --------------------------
    # Plot / Export controls
    # --------------------------
    "plot": {
        # Höhe fix, Breite dynamisch
        "fixed_height_in": 1.5,
        "bars_per_inch": 1.925,         # EINZIGER Breiten-Parameter (je kleiner, desto breiter)
        "min_width_in": 0,              # optionaler Sicherheitsboden (damit kurze Charts nicht winzig werden)
        "max_width_in": 100,            # optionaler Deckel

        # Rendering / Saving
        "render_dpi": 180,              # Figure DPI im Speicher
        "save_dpi": 300,                # PNG DPI (Pixel-Output*0.65
        # Balken / Layout
        "bar_width": 0.85,              # <- Balkenbreite (0..1 typisch)
        "x_wrap_width": 8,
        "x_max_lines": 3,
        "xtick_rotation": 45,

        "show_values": True,
        "annotate_min_abs_keUR": 0.0,   # Mindestgröße Zahlenlabel

        # Fonts
        "font": {
            "base": 8,          # rcParams["font.size"]
            "title": 11,        # Titel eh deaktiviert
            "x_ticks": 5.5,
            "value": 8,
            "value_total": 8,   # Periodenbalken = gleiche Schrift wie Steps
        },

        # Dünne Linien/Striche
        "style": {
            "bar_edge_lw": 0.6,         # Balken-Rand (weiß/black)
            "connector_lw": 0.8,        # Verbindungslinien
            "connector_alpha": 0.35,
            "other_edge_lw": 0.3,       # Rand vom "Other"-Balken
        },

        # falls du später wieder Titel willst:
        "show_title": False,
    },

    "output_dir": "",
    "output_prefix": "gp_bridge",

    # Excel embedding (wie Bubble Plot: feste Anzeigegröße)
    "excel": {
        "enabled": True,
        "col_offset": 3,
        "sheet_name": "BRIDGE",
        "image_anchor_cell": "D6",
        "set_column_widths": True,

        # wie beim Bubble Plot:
        "display_dpi": 180,     # Excel-Displaygröße unabhängig vom PNG
        "image_scale": 0.5,     # <--- hier skalieren
    },
}

# ==========================
# FONT / GLOBAL STYLE
# ==========================
from gst_excel_theme import THEME, apply_matplotlib_theme, matplotlib_font_properties

apply_matplotlib_theme(
    size=CONFIG.get("plot", {}).get("font", {}).get("base", 8),
    weight="light",
)
FONT_PROP = matplotlib_font_properties()

BRAND_BLUE = "#1E3A5F"
BRAND_BLUE_LIGHT = "#C5D8EB"


def _deep_merge_dict(base: Dict[str, Any], override: Dict[str, Any]) -> Dict[str, Any]:
    """Recursively merge override onto a copy of base (dicts only)."""
    out: Dict[str, Any] = dict(base or {})
    for k, v in (override or {}).items():
        if isinstance(v, dict) and isinstance(out.get(k), dict):
            out[k] = _deep_merge_dict(out[k], v)
        else:
            out[k] = v
    return out


def apply_config_defaults(cfg: Dict[str, Any]) -> Dict[str, Any]:
    """Ensure Fast Track partial cfgs still get CONFIG plot/excel defaults."""
    return _deep_merge_dict(CONFIG, cfg or {})


def load_bridge_colors_from_theme(_cfg: Optional[Dict[str, Any]] = None) -> Dict[str, str]:
    """Finssentials chart palette (shared theme; no external theme JSON)."""
    return {
        "total": BRAND_BLUE,
        "pos": "#0AB6DF",
        "neg": "#F88281",
        "other": BRAND_BLUE_LIGHT,
        "neutral": "#666666",
    }


def bridge_sheet_title(cfg: Dict[str, Any]) -> str:
    """Sheet name: sales_basis (Fast Track Net/Gross) wins; else profit_mode."""
    if "sales_basis" in (cfg or {}):
        return sales_basis_labels(cfg.get("sales_basis")).bridge_sheet_title
    mode = str(cfg.get("profit_mode") or "cost").strip().lower()
    return "NP Bridge" if mode == "profit" else "GP Bridge"


def bridge_table_title(cfg: Dict[str, Any]) -> str:
    if "sales_basis" in (cfg or {}):
        return sales_basis_labels(cfg.get("sales_basis")).bridge_table_title
    mode = str(cfg.get("profit_mode") or "cost").strip().lower()
    return "Net Profit Bridge" if mode == "profit" else "Gross Profit Bridge"


def calc_mode_subtitle_suffix(calc_mode: Any) -> str:
    mode = str(calc_mode or "invoice").strip().lower()
    if mode == "accrual":
        return "by accrued amounts"
    return "by invoiced amounts"



# --------------------------
# Validation
# --------------------------
def normalize_config(cfg: Dict[str, Any]) -> Dict[str, Any]:
    out = dict(cfg)

    required = [
        "current_year", "current_month", "calc_mode", "invoice_col",
        "period_mode", "revenue_col", "bridge",
    ]
    for k in required:
        if k not in out or out[k] in (None, ""):
            raise ValueError(f"CONFIG['{k}'] muss gesetzt sein.")

    if out["calc_mode"] not in {"invoice", "accrual"}:
        raise ValueError("CONFIG['calc_mode'] muss 'invoice' oder 'accrual' sein.")
    if out["period_mode"] not in {"FY", "YTD", "LTM"}:
        raise ValueError("CONFIG['period_mode'] muss 'FY', 'YTD' oder 'LTM' sein.")

    m = int(out["current_month"])
    if not (1 <= m <= 12):
        raise ValueError("CONFIG['current_month'] muss 1..12 sein.")

    out["fiscal_year_end_month"] = int(out.get("fiscal_year_end_month", out.get("fy_end_month", 12)))
    out["fiscal_year_end_day"] = int(out.get("fiscal_year_end_day", out.get("fy_end_day", 31)))

    out["profit_mode"] = str(out.get("profit_mode") or "cost").strip().lower()
    if out["profit_mode"] not in {"cost", "profit"}:
        raise ValueError("CONFIG['profit_mode'] muss 'cost' oder 'profit' sein.")
    if out["profit_mode"] == "cost":
        if not out.get("cost_col"):
            raise ValueError("CONFIG['cost_col'] muss gesetzt sein, wenn profit_mode='cost'.")
    else:
        if not out.get("profit_col"):
            raise ValueError("CONFIG['profit_col'] muss gesetzt sein, wenn profit_mode='profit'.")

    out["invoice_mapping_mode"] = str(out.get("invoice_mapping_mode") or "date").strip().lower()
    if out["invoice_mapping_mode"] not in {"year", "date"}:
        raise ValueError("CONFIG['invoice_mapping_mode'] muss 'year' oder 'date' sein.")
    if out["invoice_mapping_mode"] == "year" and out["period_mode"] != "FY":
        raise ValueError("invoice_mapping_mode='year' unterstützt nur period_mode='FY'.")

    labels = sales_basis_labels(out.get("sales_basis"))
    out["sales_basis"] = labels.sales_basis
    sheet_title = bridge_sheet_title(out)
    table_title = bridge_table_title(out)
    if not str(out.get("table") or "").strip() or str(out.get("table")).strip() in {
        "Gross Profit Bridge",
        "Net Profit Bridge",
        "BRIDGE",
    }:
        out["table"] = table_title
    mode_suffix = calc_mode_subtitle_suffix(out.get("calc_mode"))
    prev_suffix = str(out.get("subtitle_suffix") or "").strip().lower()
    # Always keep the recognition phrase in sync with calc_mode (ignore stale defaults).
    if prev_suffix in {
        "",
        "by invoiced amounts",
        "by accrued amounts",
        "net sales bridge",
        "gross sales bridge",
    } or prev_suffix.endswith(" bridge"):
        out["subtitle_suffix"] = mode_suffix
    excel_cfg = dict(out.get("excel") or {})
    prev = str(excel_cfg.get("sheet_name") or out.get("base_sheet_name") or "").strip().lower()
    # Keep Fast Track Net/Gross sheet names; only fill blanks / generic BRIDGE
    if prev in {"", "bridge"}:
        excel_cfg["sheet_name"] = sheet_title
    out["excel"] = excel_cfg
    base_prev = str(out.get("base_sheet_name") or "").strip().lower()
    if base_prev in {"", "bridge"}:
        out["base_sheet_name"] = str(excel_cfg.get("sheet_name") or sheet_title)

    # Legacy singular filter
    f = out.get("filter") or {}
    if f.get("enabled", False):
        if not f.get("col"):
            raise ValueError("filter.enabled=True, aber filter.col fehlt.")
        vals = f.get("values")
        if vals is None or len(vals) == 0:
            raise ValueError("filter.enabled=True, aber filter.values ist leer.")
        f["values"] = set(vals)
    out["filter"] = f

    # FX
    if out.get("apply_fx", False) and not out.get("fx_col"):
        raise ValueError("apply_fx=True, aber fx_col fehlt.")

    # Bridge
    b = out.get("bridge") or {}
    if not b.get("dim_col"):
        raise ValueError("bridge.dim_col muss gesetzt sein.")
    b["top_n"] = int(b.get("top_n", 6))
    if b["top_n"] < 1:
        raise ValueError("bridge.top_n muss >= 1 sein.")
    b.setdefault("other_label", "Other")
    b.setdefault("missing_token", "__MISSING__")

    b.setdefault("combine_input_other", True)
    toks = b.get("input_other_tokens", {b["other_label"]})
    if toks is None:
        toks = {b["other_label"]}
    b["input_other_tokens"] = set(map(str, toks))

    b.setdefault("other_bucket_label", "Other")
    b.setdefault("sort_mode", "pos_then_neg_abs")
    if b["sort_mode"] not in {"abs", "value", "pos_then_neg_abs"}:
        raise ValueError("bridge.sort_mode muss 'abs', 'value' oder 'pos_then_neg_abs' sein.")

    out["bridge"] = b

    out.setdefault("fmt_money", "#,##0;(#,##0)")
    return out


# --------------------------
# Period logic (3 Perioden => 2 Bridges)
# --------------------------
def safe_month_day_ts(year: int, month: int, day: int) -> pd.Timestamp:
    first = pd.Timestamp(year, month, 1)
    last = first + pd.offsets.MonthEnd(0)
    d = min(int(day), int(last.day))
    return pd.Timestamp(year, month, d)


def fiscal_year_bounds(fy_end_year: int, fy_end_m: int, fy_end_d: int) -> Tuple[pd.Timestamp, pd.Timestamp]:
    end = safe_month_day_ts(fy_end_year, fy_end_m, fy_end_d)
    start = (end - pd.DateOffset(years=1)) + pd.Timedelta(days=1)
    return start, end


def compute_current_fy_end_year(as_of_end: pd.Timestamp, fy_end_m: int, fy_end_d: int) -> int:
    CY = as_of_end.year
    fye_this_year = safe_month_day_ts(CY, fy_end_m, fy_end_d)
    return CY if as_of_end <= fye_this_year else CY + 1


def _fy_label(y: int) -> str:
    return f"FY{str(y)[-2:]}A"


def get_three_periods(cfg: Dict[str, Any]) -> Tuple[List[Tuple[pd.Timestamp, pd.Timestamp, str]], str]:
    """
    Three periods left→right (oldest→newest), aligned with Fast Track Bars:
    - At FYE: last 3 completed FYs
    - Else: previous FY, last completed FY, YTD of the open FY
    YTD/LTM modes keep their dedicated windows when period_mode is set explicitly.
    """
    CY = int(cfg["current_year"])
    m = int(cfg["current_month"])
    fy_end_m = int(cfg["fiscal_year_end_month"])
    fy_end_d = int(cfg["fiscal_year_end_day"])

    as_of_end = pd.Timestamp(CY, m, 1) + pd.offsets.MonthEnd(0)
    cur_fy_end_year = compute_current_fy_end_year(as_of_end, fy_end_m, fy_end_d)
    fye_this_calendar_year = safe_month_day_ts(as_of_end.year, fy_end_m, fy_end_d)
    at_fiscal_year_end = int(as_of_end.month) == int(fy_end_m)
    year_mapping = str(cfg.get("invoice_mapping_mode") or "date").strip().lower() == "year"

    mode = cfg["period_mode"]

    if mode == "FY":
        last_completed = as_of_end.year if as_of_end >= fye_this_calendar_year else as_of_end.year - 1
        if at_fiscal_year_end or year_mapping:
            years = [last_completed - 2, last_completed - 1, last_completed]
            periods = [
                (*fiscal_year_bounds(y, fy_end_m, fy_end_d), _fy_label(y))
                for y in years
            ]
        else:
            older_fy = last_completed - 1
            prev_fy = last_completed
            ytd_start, _ = fiscal_year_bounds(cur_fy_end_year, fy_end_m, fy_end_d)
            periods = [
                (*fiscal_year_bounds(older_fy, fy_end_m, fy_end_d), _fy_label(older_fy)),
                (*fiscal_year_bounds(prev_fy, fy_end_m, fy_end_d), _fy_label(prev_fy)),
                (ytd_start, as_of_end, f"YTD{str(cur_fy_end_year)[-2:]}"),
            ]
        subtitle = f"{periods[0][2]} \u2192 {periods[2][2]}"
        return periods, subtitle

    if mode == "YTD":
        s2, _ = fiscal_year_bounds(cur_fy_end_year, fy_end_m, fy_end_d)
        e2 = as_of_end
        delta_days = int((e2 - s2).days)

        prev1 = cur_fy_end_year - 1
        s1, _ = fiscal_year_bounds(prev1, fy_end_m, fy_end_d)
        e1 = s1 + pd.Timedelta(days=delta_days)

        prev2 = cur_fy_end_year - 2
        s0, _ = fiscal_year_bounds(prev2, fy_end_m, fy_end_d)
        e0 = s0 + pd.Timedelta(days=delta_days)

        l2 = f"YTD{str(cur_fy_end_year)[-2:]}"
        l1 = f"YTD{str(prev1)[-2:]}"
        l0 = f"YTD{str(prev2)[-2:]}"

        periods = [(s0, e0, l0), (s1, e1, l1), (s2, e2, l2)]
        subtitle = f"{l0} \u2192 {l2}"
        return periods, subtitle

    if mode == "LTM":
        e2 = as_of_end
        s2 = e2 - pd.DateOffset(years=1) + pd.Timedelta(days=1)

        e1 = e2 - pd.DateOffset(years=1)
        s1 = e1 - pd.DateOffset(years=1) + pd.Timedelta(days=1)

        e0 = e2 - pd.DateOffset(years=2)
        s0 = e0 - pd.DateOffset(years=1) + pd.Timedelta(days=1)

        l2 = f"LTM {e2.month}-{str(e2.year)[-2:]}"
        l1 = f"LTM {e1.month}-{str(e1.year)[-2:]}"
        l0 = f"LTM {e0.month}-{str(e0.year)[-2:]}"

        periods = [(s0, e0, l0), (s1, e1, l1), (s2, e2, l2)]
        subtitle = f"{l0} \u2192 {l2}"
        return periods, subtitle

    raise ValueError("Unbekannter period_mode.")


# --------------------------
# Preprocess + Calc
# --------------------------
def preprocess_input(df: pd.DataFrame, cfg: Dict[str, Any]) -> pd.DataFrame:
    d = df.copy()
    d.columns = d.columns.str.strip()

    # Fast Track multi-rule filters + legacy singular filter
    if apply_filters is not None:
        d = apply_filters(d, cfg)
    f = cfg.get("filter", {}) or {}
    if f.get("enabled", False):
        col = f["col"]
        if col not in d.columns:
            raise ValueError(f"Filter-Spalte '{col}' fehlt im Input.")
        d = d[d[col].isin(f["values"])].copy()

    mapping = str(cfg.get("invoice_mapping_mode") or "date").strip().lower()
    date_cols: List[str] = []
    if mapping != "year":
        date_cols.append(cfg["invoice_col"])
    if cfg["calc_mode"] == "accrual":
        for c in [cfg.get("start_col"), cfg.get("end_col")]:
            if not c or c not in d.columns:
                raise ValueError(f"calc_mode='accrual' braucht Spalte '{c}' im Input.")
            date_cols.append(str(c))

    for c in date_cols:
        if c not in d.columns:
            raise ValueError(f"Spalte '{c}' fehlt im Input.")
        d[c] = pd.to_datetime(d[c], errors="coerce", dayfirst=True)

    rcol = cfg["revenue_col"]
    if rcol not in d.columns:
        raise ValueError(f"revenue_col '{rcol}' fehlt im Input.")
    d["revenue"] = pd.to_numeric(d[rcol], errors="coerce").fillna(0.0)

    profit_mode = str(cfg.get("profit_mode") or "cost").strip().lower()
    if profit_mode == "profit":
        pcol = cfg["profit_col"]
        if pcol not in d.columns:
            raise ValueError(f"profit_col '{pcol}' fehlt im Input.")
        d["gp"] = pd.to_numeric(d[pcol], errors="coerce").fillna(0.0)
        d["cost"] = d["revenue"] - d["gp"]
    else:
        ccol = cfg["cost_col"]
        if ccol not in d.columns:
            raise ValueError(f"cost_col '{ccol}' fehlt im Input.")
        d["cost"] = pd.to_numeric(d[ccol], errors="coerce").fillna(0.0)
        d["gp"] = d["revenue"] - d["cost"]

    if cfg.get("apply_fx", False):
        fx_col = cfg["fx_col"]
        if fx_col not in d.columns:
            raise ValueError(f"apply_fx=True, aber fx_col '{fx_col}' fehlt im Input.")
        d["fx_rate"] = pd.to_numeric(d[fx_col], errors="coerce")
        d["fx_rate"] = d["fx_rate"].where(d["fx_rate"].notna() & (d["fx_rate"] != 0), 1.0)
        d["revenue"] = d["revenue"] * d["fx_rate"]
        d["cost"] = d["cost"] * d["fx_rate"]
        d["gp"] = d["gp"] * d["fx_rate"]

    # Match GST: keep rows with invoice date; contract dates only for accrual
    if mapping != "year":
        d = d.dropna(subset=[cfg["invoice_col"]]).copy()
    if cfg["calc_mode"] == "accrual":
        d = d.dropna(subset=[cfg["start_col"], cfg["end_col"]]).copy()
        d = d[d[cfg["end_col"]] >= d[cfg["start_col"]]].copy()
    return d


def compute_period_amount(d: pd.DataFrame, start: pd.Timestamp, end: pd.Timestamp, cfg: Dict[str, Any], src_col: str) -> pd.Series:
    """Shared invoice/accrual amounts via funktionssammlung (same as GST)."""
    if compute_amount is not None:
        return compute_amount(d, start, end, cfg, src_col)
    if cfg["calc_mode"] == "invoice":
        invoice_dates = pd.to_datetime(d[cfg["invoice_col"]], errors="coerce", dayfirst=True)
        mask = (invoice_dates >= start) & (invoice_dates <= end)
        amt = pd.to_numeric(d[src_col], errors="coerce").fillna(0.0)
        return amt.where(mask, 0.0)
    # Accrual fallback = recognized_value (no invoice gating)
    if recognized_value is not None:
        tmp = d.copy()
        tmp["value"] = pd.to_numeric(tmp[src_col], errors="coerce").fillna(0.0)
        return recognized_value(tmp, start, end, cfg)
    S = pd.to_datetime(d[cfg["start_col"]], errors="coerce")
    E = pd.to_datetime(d[cfg["end_col"]], errors="coerce")
    aktiv_von = S.where(S > start, start)
    aktiv_bis = E.where(E < end, end)
    tage = ((aktiv_bis - aktiv_von) / np.timedelta64(1, "D"))
    tage = (np.floor(tage) + 1).clip(min=0)
    contract_days = ((E - S) / np.timedelta64(1, "D"))
    contract_days = (np.floor(contract_days) + 1).clip(min=1)
    value = pd.to_numeric(d[src_col], errors="coerce").fillna(0.0)
    out = value * (tage / contract_days)
    return out.where((S <= end) & (E >= start), 0.0)



# --------------------------
# Bridge build
# --------------------------
def _clean_dim_series(s: pd.Series, missing_token: str) -> pd.Series:
    x = s.astype(str).str.strip()
    x = x.replace({"": missing_token, "nan": missing_token, "None": missing_token, "NaT": missing_token})
    x = x.where(~s.isna(), missing_token)
    return x


@dataclass
class BridgeItem:
    label: str
    value_keUR: float
    is_total: bool = False
    is_other: bool = False


def build_two_bridge_items(d: pd.DataFrame, cfg: Dict[str, Any]) -> Tuple[List[BridgeItem], str]:
    bcfg = cfg["bridge"]
    dim_col = bcfg["dim_col"]
    top_n = int(bcfg["top_n"])
    other_label = str(bcfg["other_label"])
    missing_token = str(bcfg["missing_token"])

    combine_input_other = bool(bcfg.get("combine_input_other", True))
    input_other_tokens: Set[str] = set(map(str, bcfg.get("input_other_tokens", {other_label})))

    other_bucket_label = str(bcfg.get("other_bucket_label", "Other"))
    sort_mode = str(bcfg.get("sort_mode", "pos_then_neg_abs"))

    if dim_col not in d.columns:
        raise ValueError(f"bridge.dim_col '{dim_col}' fehlt im Input.")

    periods, subtitle = get_three_periods(cfg)
    (s0, e0, l0), (s1, e1, l1), (s2, e2, l2) = periods

    tmp = d.copy()
    tmp["dim"] = _clean_dim_series(tmp[dim_col], missing_token)

    profit_mode = str(cfg.get("profit_mode") or "cost").strip().lower()
    for idx, (s, e, _) in enumerate(periods):
        if profit_mode == "profit":
            tmp[f"gp{idx}"] = compute_period_amount(tmp, s, e, cfg, "gp")
        else:
            rev = compute_period_amount(tmp, s, e, cfg, "revenue")
            cost = compute_period_amount(tmp, s, e, cfg, "cost")
            tmp[f"gp{idx}"] = rev - cost

    gp0_total = float(tmp["gp0"].sum())
    gp1_total = float(tmp["gp1"].sum())
    gp2_total = float(tmp["gp2"].sum())

    def delta_by_dim(a: int, b: int) -> pd.DataFrame:
        ga = tmp.groupby("dim", as_index=False)[f"gp{a}"].sum().rename(columns={f"gp{a}": "gpa"})
        gb = tmp.groupby("dim", as_index=False)[f"gp{b}"].sum().rename(columns={f"gp{b}": "gpb"})
        m = pd.merge(gb, ga, on="dim", how="outer").fillna(0.0)
        m["delta"] = m["gpb"] - m["gpa"]
        return m[["dim", "delta"]]

    def _sort_steps_df(m: pd.DataFrame) -> pd.DataFrame:
        if m.empty:
            return m

        if sort_mode == "abs":
            return m.assign(abs_delta=m["delta"].abs()).sort_values("abs_delta", ascending=False).drop(columns=["abs_delta"])
        if sort_mode == "value":
            return m.sort_values("delta", ascending=False)

        mm = m.copy()
        mm["is_neg"] = mm["delta"] < 0
        mm["abs_delta"] = mm["delta"].abs()
        mm = mm.sort_values(["is_neg", "abs_delta"], ascending=[True, False]).drop(columns=["is_neg", "abs_delta"])
        return mm

    def bridge_steps(a: int, b: int, label_a: str, label_b: str, total_a: float, total_b: float) -> List[BridgeItem]:
        m = delta_by_dim(a, b)

        miss_mask = (m["dim"].astype(str) == missing_token)
        missing_delta = float(m.loc[miss_mask, "delta"].sum())
        m = m.loc[~miss_mask].copy()

        input_other_delta = 0.0
        if combine_input_other:
            other_mask = m["dim"].astype(str).isin(input_other_tokens)
            input_other_delta = float(m.loc[other_mask, "delta"].sum())
            m = m.loc[~other_mask].copy()

        m_rank = _sort_steps_df(m)
        top = m_rank.head(top_n).copy()
        rest = m_rank.iloc[top_n:].copy()

        other_delta = float(rest["delta"].sum() + missing_delta + input_other_delta)
        bucket_label = other_bucket_label

        items: List[BridgeItem] = [BridgeItem(label=label_a, value_keUR=total_a / 1000.0, is_total=True)]
        for _, r in top.iterrows():
            items.append(BridgeItem(label=str(r["dim"]), value_keUR=float(r["delta"]) / 1000.0))

        if abs(other_delta) > 1e-9:
            items.append(BridgeItem(label=bucket_label, value_keUR=other_delta / 1000.0, is_other=True))

        items.append(BridgeItem(label=label_b, value_keUR=total_b / 1000.0, is_total=True))
        return items

    items_01 = bridge_steps(0, 1, l0, l1, gp0_total, gp1_total)
    items_12 = bridge_steps(1, 2, l1, l2, gp1_total, gp2_total)

    merged = items_01 + items_12[1:]
    return merged, subtitle

def fmt_keUR(x: float) -> str:
    v = int(round(x))
    if v < 0:
        return f"({abs(v):,})"
    return f"{v:,}"


def wrap_xtick(s: str, width: int, max_lines: int) -> str:
    s = str(s)
    lines = textwrap.wrap(s, width=width)
    if len(lines) > max_lines:
        lines = lines[:max_lines]
        lines[-1] = lines[-1].rstrip(".") + "..."
    return "\n".join(lines)


def compute_dynamic_figsize(cfg: Dict[str, Any], n_bars: int) -> Tuple[float, float]:
    pcfg = cfg.get("plot", {}) or {}
    h = float(pcfg.get("fixed_height_in", 3.4))

    bars_per_inch = float(pcfg.get("bars_per_inch", 2.2))  # <- einziger "Tuning"-Regler
    if bars_per_inch <= 0:
        raise ValueError("plot.bars_per_inch muss > 0 sein.")

    # Breite so, dass die Balken "gerade gut" draufpassen:
    w = n_bars / bars_per_inch

    # optionale Sicherheitsgrenzen
    min_w = pcfg.get("min_width_in", None)
    max_w = pcfg.get("max_width_in", None)
    if min_w is not None:
        w = max(float(min_w), w)
    if max_w is not None:
        w = min(float(max_w), w)

    return w, h


def _waterfall_geometry(items: List[BridgeItem]) -> Tuple[List[float], List[float]]:
    running = 0.0
    bottoms: List[float] = []
    heights: List[float] = []
    for it in items:
        if it.is_total:
            bottoms.append(0.0)
            heights.append(float(it.value_keUR))
            running = float(it.value_keUR)
        else:
            delta = float(it.value_keUR)
            btm = running if delta >= 0 else running + delta
            bottoms.append(btm)
            heights.append(abs(delta))
            running += delta
    return bottoms, heights


def detect_axis_break(
    bottoms: List[float],
    heights: List[float],
    items: List[BridgeItem],
    plot_cfg: Dict[str, Any],
) -> Optional[Dict[str, float]]:
    """
    Detect outlier peak and return compression params.

    Honest band follows the second-largest total ("body"). Values above that
    are compressed into a short upper band so the staircase stays readable.
    """
    if not bool(plot_cfg.get("axis_break", True)):
        return None
    if not bottoms:
        return None

    tops = [float(b + h) for b, h in zip(bottoms, heights)]
    y_ceil = max(tops) if tops else 0.0
    y_floor = float(min(0.0, min(bottoms)))
    if y_ceil - y_floor <= 1e-9:
        return None

    total_abs = sorted(abs(float(it.value_keUR)) for it in items if it.is_total)
    if len(total_abs) < 2:
        return None

    body = float(total_abs[-2])
    include_pad = float(plot_cfg.get("axis_break_include_pad", 1.02))
    for b, h, it in zip(bottoms, heights, items):
        if it.is_total:
            continue
        top = float(b + h)
        if top <= body * include_pad:
            body = max(body, top)

    ratio = float(plot_cfg.get("axis_break_ratio", 1.60))
    if body <= 1e-9 or y_ceil / body < ratio:
        return None

    scale_low = body * float(plot_cfg.get("axis_break_headroom", 1.0))
    if scale_low >= y_ceil * 0.90:
        return None

    low_frac = float(plot_cfg.get("axis_break_low_frac", 0.70))
    low_frac = min(0.88, max(0.55, low_frac))

    return {
        "scale_low": float(scale_low),
        "y_ceil": float(y_ceil),
        "y_floor": float(y_floor),
        "low_frac": low_frac,
        "axis_max": float(scale_low),
    }


def _make_compress_mapper(brk: Dict[str, float]):
    """
    Continuous map data-y → display-y (no gap — keeps the staircase connected).
    [y_floor, scale_low] is 1:1; above that values are compressed into the remaining height.
    """
    scale_low = float(brk["scale_low"])
    y_ceil = float(brk["y_ceil"])
    y_floor = float(brk["y_floor"])
    low_frac = float(brk["low_frac"])
    upper_frac = max(0.08, 1.0 - low_frac)

    unit = max(scale_low - y_floor, 1e-9)
    upper_h = unit * (upper_frac / low_frac)
    display_top = scale_low + upper_h

    def map_y(y: float, band: str = "auto") -> float:
        y = float(y)
        if y <= scale_low + 1e-12:
            return y
        span = max(y_ceil - scale_low, 1e-12)
        t = max(0.0, min(1.0, (y - scale_low) / span))
        return scale_low + t * upper_h

    return map_y, display_top


def _draw_axis_break_slashes(
    ax,
    xc: float,
    bar_width: float,
    y_mid: float,
    color: str = "#1E3A5F",
    angle_deg: float = 20.0,
    overhang_frac: float = 0.12,
) -> None:
    """
    Two short parallel diagonals (~20°) with a slight overhang;
    white only between them and only inside the bar width — no bar edge seams.
    """
    from matplotlib.patches import Polygon

    overhang = overhang_frac * bar_width
    bar_left = xc - bar_width / 2
    bar_right = xc + bar_width / 2
    left = bar_left - overhang
    right = bar_right + overhang
    width_data = right - left

    fig = ax.figure
    bbox = ax.get_position()
    fig_w_in, fig_h_in = fig.get_size_inches()
    ax_w_in = max(fig_w_in * bbox.width, 1e-9)
    ax_h_in = max(fig_h_in * bbox.height, 1e-9)
    x0, x1 = ax.get_xlim()
    y0lim, y1lim = ax.get_ylim()
    data_per_in_x = abs(x1 - x0) / ax_w_in
    data_per_in_y = abs(y1lim - y0lim) / ax_h_in

    width_in = width_data / max(data_per_in_x, 1e-12)
    rise_in = width_in * np.tan(np.radians(angle_deg))
    rise_data = rise_in * data_per_in_y

    # Thin white gap between the two parallels (~1.1 mm)
    sep_data = 0.042 * data_per_in_y

    def y_on_slash(x: float, y_at_center_offset: float) -> float:
        # linear interp from left to right
        t = (x - left) / max(width_data, 1e-12)
        return (y_mid - 0.5 * rise_data + y_at_center_offset) + t * rise_data

    # White wipe only inside the bar (not in the overhang), so no stray horizontal edges
    poly = Polygon(
        [
            (bar_left, y_on_slash(bar_left, 0.0)),
            (bar_right, y_on_slash(bar_right, 0.0)),
            (bar_right, y_on_slash(bar_right, sep_data)),
            (bar_left, y_on_slash(bar_left, sep_data)),
        ],
        closed=True,
        facecolor="white",
        edgecolor="none",
        zorder=5,
        clip_on=False,
    )
    ax.add_patch(poly)
    line_kw = dict(color=color, linewidth=1.15, solid_capstyle="butt", zorder=6, clip_on=False)
    ax.plot(
        [left, right],
        [y_on_slash(left, 0.0), y_on_slash(right, 0.0)],
        **line_kw,
    )
    ax.plot(
        [left, right],
        [y_on_slash(left, sep_data), y_on_slash(right, sep_data)],
        **line_kw,
    )


# --------------------------
# Plot
# --------------------------

# --------------------------
# Plot
# --------------------------
def plot_gp_bridge(items: List[BridgeItem], cfg: Dict[str, Any], subtitle: str) -> plt.Figure:
    plot_cfg = cfg.get("plot", {}) or {}

    n = len(items)
    fig_w, fig_h = compute_dynamic_figsize(cfg, n_bars=n)
    dpi = int(plot_cfg.get("render_dpi", 180))

    show_values = bool(plot_cfg.get("show_values", True))
    min_abs = float(plot_cfg.get("annotate_min_abs_keUR", 0.0))
    wrap_w = int(plot_cfg.get("x_wrap_width", 12))
    max_lines = int(plot_cfg.get("x_max_lines", 2))

    bar_width = float(plot_cfg.get("bar_width", 0.68))

    fcfg = plot_cfg.get("font", {}) or {}
    xt_fs = float(fcfg.get("x_ticks", 5.5))
    val_fs = float(fcfg.get("value", 6))
    val_total_fs = float(fcfg.get("value_total", val_fs))
    title_fs = int(fcfg.get("title", 11))
    xtick_rotation = float(plot_cfg.get("xtick_rotation", 45 if n >= 4 else 0))

    style_cfg = plot_cfg.get("style", {}) or {}
    bar_edge_lw = float(style_cfg.get("bar_edge_lw", 0.6))
    connector_lw = float(style_cfg.get("connector_lw", 0.8))
    connector_alpha = float(style_cfg.get("connector_alpha", 0.35))
    other_edge_lw = float(style_cfg.get("other_edge_lw", 0.9))

    theme_colors = load_bridge_colors_from_theme(cfg)
    total_color = theme_colors["total"]
    pos_color = theme_colors["pos"]
    neg_color = theme_colors["neg"]

    fig, ax = plt.subplots(figsize=(fig_w, fig_h), dpi=dpi)

    # More bottom room when labels are rotated (like vertical bars / ARR)
    bottom_margin = 0.40 if xtick_rotation else 0.28
    fig.subplots_adjust(left=0.03, right=0.995, top=0.92, bottom=bottom_margin)

    labels_raw = [it.label for it in items]
    labels = [wrap_xtick(lab, width=wrap_w, max_lines=max_lines) for lab in labels_raw]

    x = np.arange(n)

    running = 0.0
    bottoms: List[float] = []
    heights: List[float] = []
    facecolors: List[Any] = []
    is_other_flags: List[bool] = []

    # Waterfall
    for it in items:
        if it.is_total:
            bottoms.append(0.0)
            heights.append(it.value_keUR)
            facecolors.append(total_color)
            is_other_flags.append(False)
            running = it.value_keUR
        else:
            delta = it.value_keUR
            btm = running if delta >= 0 else running + delta
            bottoms.append(btm)
            heights.append(abs(delta))

            if it.is_other:
                facecolors.append((0, 0, 0, 0))
                is_other_flags.append(True)
            else:
                facecolors.append(pos_color if delta >= 0 else neg_color)
                is_other_flags.append(False)

            running += delta

    bars = ax.bar(
        x,
        heights,
        bar_width,
        bottom=bottoms,
        color=facecolors,
        edgecolor="white",
        linewidth=bar_edge_lw
    )

    # Apply "Other" hatch style
    for rect, is_other in zip(bars, is_other_flags):
        if is_other:
            rect.set_facecolor((0, 0, 0, 0))
            rect.set_edgecolor("black")
            rect.set_linewidth(other_edge_lw)
            rect.set_hatch("////")

    # Connector lines
    running = 0.0
    for i, it in enumerate(items[:-1]):
        if it.is_total:
            running = it.value_keUR
        else:
            running += it.value_keUR

        ax.plot(
            [x[i] + bar_width / 2, x[i + 1] - bar_width / 2],
            [running, running],
            linewidth=connector_lw,
            color="black",
            alpha=connector_alpha
        )

    # Werte
    if show_values:
        ax.relim()
        ax.autoscale_view()
        yr = (ax.get_ylim()[1] - ax.get_ylim()[0] + 1e-9)
        dy = 0.02 * yr

        for i, it in enumerate(items):
            val = it.value_keUR
            if abs(val) < min_abs:
                continue

            if it.is_total:
                y_center = it.value_keUR / 2.0
                ax.text(
                    x[i], y_center,
                    fmt_keUR(val),
                    ha="center", va="center",
                    fontsize=val_total_fs, color="white",
                    fontweight="light",
                    zorder=10,
                    bbox=dict(
                        facecolor=total_color,
                        edgecolor="none",
                        alpha=1.0,
                        boxstyle="round,pad=0.20",
                    ),
                    clip_on=False,
                )
            else:
                delta = it.value_keUR
                if delta >= 0:
                    y = bottoms[i] + heights[i] + dy
                    va = "bottom"
                else:
                    y = bottoms[i] - dy
                    va = "top"

                if it.is_other:
                    ax.text(
                        x[i], y,
                        fmt_keUR(delta),
                        ha="center", va=va,
                        fontsize=val_fs, color="black",
                        fontweight="light",
                        bbox=dict(
                            facecolor="white",
                            edgecolor="none",
                            alpha=0.75,
                            boxstyle="round,pad=0.18"
                        )
                    )
                else:
                    ax.text(
                        x[i], y,
                        fmt_keUR(delta),
                        ha="center", va=va,
                        fontsize=val_fs, color="black",
                        fontweight="light"
                    )

    ax.set_xticks(x)
    ax.set_xticklabels(
        labels,
        fontproperties=FONT_PROP,
        rotation=xtick_rotation,
        ha="right" if xtick_rotation else "center",
        rotation_mode="anchor" if xtick_rotation else "default",
    )

    for tick in ax.get_xticklabels():
        tick.set_fontsize(xt_fs)
        tick.set_fontweight("light")
        tick.set_color(BRAND_BLUE)
        if xtick_rotation:
            tick.set_rotation(xtick_rotation)
            tick.set_ha("right")
            tick.set_rotation_mode("anchor")
    ax.tick_params(axis="x", pad=2, length=3.0, width=0.6, direction="out", color=BRAND_BLUE)

    ax.set_yticks([])

    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    ax.spines["left"].set_visible(False)
    ax.grid(False)

    if bool(plot_cfg.get("show_title", False)):
        company = str(cfg.get("company", "")).strip()
        title = str(cfg.get("table", "")).strip()
        prefix = str(cfg.get("period_label_prefix", "")).strip()
        parts = [p for p in [company, prefix, title, subtitle] if p]
        ax.set_title(" | ".join(parts), fontsize=title_fs, fontproperties=FONT_PROP, fontweight="light")

    return fig


# --------------------------
# Output naming
# --------------------------
def clean_filename(s: str) -> str:
    s = str(s).strip().replace(" ", "_")
    return re.sub(r"[^\w\-_.]", "", s)


def build_output_paths(cfg: Dict[str, Any], suffix: str) -> Tuple[str, str]:
    out_dir = cfg.get("output_dir", ".")
    os.makedirs(out_dir, exist_ok=True)
    prefix = clean_filename(cfg.get("output_prefix", "gp_bridge"))
    mode = cfg.get("calc_mode", "invoice")
    pm = cfg.get("period_mode", "FY")
    y = int(cfg["current_year"])
    m = int(cfg["current_month"])
    base = f"{prefix}_{clean_filename(suffix)}_{pm}_cy{str(y)[-2:]}_{m:02d}_{mode}"
    png_path = os.path.join(out_dir, base + ".png")
    xlsx_path = os.path.join(out_dir, base + ".xlsx")
    return png_path, xlsx_path


# --------------------------
# Excel export (wie Bubble Plot: feste Anzeigegröße + scale)
# --------------------------

def _write_gp_bridge_sheet(
    wb,
    cfg: Dict[str, Any],
    png_path: str,
    subtitle: str,
    fig_size_in: Tuple[float, float],
    sheet_name: str,
) -> str:
    excel_cfg = cfg.get("excel", {}) or {}
    col_offset = int(excel_cfg.get("col_offset", 3))
    anchor_cell = excel_cfg.get("image_anchor_cell", "D6")
    set_widths = bool(excel_cfg.get("set_column_widths", True))

    if sheet_name in wb.sheetnames:
        del wb[sheet_name]
    ws = wb.create_sheet(sheet_name)

    brand = BRAND_BLUE.replace("#", "FF")
    title_font = Font(name=THEME.font_name, size=24, color=brand)
    sub_font = Font(name=THEME.font_name, size=12, color=brand)
    info_font = Font(name=THEME.font_name, size=9, color=brand, bold=True)

    ws["D1"].value = cfg.get("title", "")
    ws["D1"].font = title_font

    suffix = str(cfg.get("subtitle_suffix", "")).strip()
    suffix_txt = f" ({suffix})" if suffix else ""
    table_name = f"{cfg.get('table', '')}{suffix_txt}"

    ws["D2"].value = f"{table_name} | {subtitle}"
    ws["D2"].font = sub_font

    ws["D4"].value = f"{cfg.get('company', '')} | {table_name}"
    ws["D4"].font = info_font

    fill_grey = PatternFill(fill_type="solid", fgColor="FFF3F1EF")
    for r in range(1, 200):
        for c in range(1, 1 + col_offset):
            ws.cell(row=r, column=c).fill = fill_grey

    fill_white = PatternFill(fill_type="solid", fgColor="FFFFFFFF")
    for r in range(1, 200):
        for c in range(1 + col_offset, 30):
            ws.cell(row=r, column=c).fill = fill_white

    if os.path.exists(png_path):
        img = _xl_image_from_path(png_path)

        fig_w, fig_h = fig_size_in
        display_dpi = int(excel_cfg.get("display_dpi", 180))
        scale = float(excel_cfg.get("image_scale", 1.0))

        img.width = int(fig_w * display_dpi * scale)
        img.height = int(fig_h * display_dpi * scale)

        ws.add_image(img, anchor_cell)
    else:
        ws[anchor_cell].value = f"Bild nicht gefunden: {png_path}"

    if set_widths:
        for c in range(1, 30):
            ws.column_dimensions[get_column_letter(c)].width = 14
        ws.column_dimensions["A"].width = 3
        ws.column_dimensions["B"].width = 3
        ws.column_dimensions["C"].width = 3
        ws.column_dimensions["D"].width = 20

    ws.sheet_properties.outlinePr.summaryRight = True
    for col_letter in ["A", "B", "C"]:
        ws.column_dimensions[col_letter].outlineLevel = 1
        ws.column_dimensions[col_letter].hidden = True
    ws.column_dimensions["D"].collapsed = True
    return sheet_name


def export_excel_with_image(
    cfg: Dict[str, Any],
    png_path: str,
    xlsx_path: str,
    subtitle: str,
    fig_size_in: Tuple[float, float],
    wb=None,
    *,
    save: bool = True,
    sheet_name: Optional[str] = None,
) -> Optional[str]:
    excel_cfg = cfg.get("excel", {}) or {}
    if not excel_cfg.get("enabled", True):
        return None

    target = sheet_name or excel_cfg.get("sheet_name") or bridge_sheet_title(cfg)
    owns_wb = wb is None
    if owns_wb:
        wb = Workbook()
        default = wb.active
        if default is not None and default.title == "Sheet":
            wb.remove(default)

    written = _write_gp_bridge_sheet(wb, cfg, png_path, subtitle, fig_size_in, str(target))
    if save and xlsx_path:
        wb.save(xlsx_path)
    return written


def _render_gp_bridge_assets(df: pd.DataFrame, cfg: Dict[str, Any]):
    cfg = apply_config_defaults(cfg)
    cfg = normalize_config(cfg)
    d = preprocess_input(df, cfg)

    items, subtitle = build_two_bridge_items(d, cfg)

    fig_w, fig_h = compute_dynamic_figsize(cfg, n_bars=len(items))
    fig = plot_gp_bridge(items, cfg, subtitle)

    png_path, xlsx_path = build_output_paths(cfg, subtitle)

    save_dpi = int(cfg.get("plot", {}).get("save_dpi", 300))
    fig.savefig(png_path, dpi=save_dpi, bbox_inches=None)
    plt.close(fig)
    return cfg, subtitle, png_path, xlsx_path, (fig_w, fig_h)


# --------------------------
# MAIN runner
# --------------------------
def run_gp_bridge(df: pd.DataFrame, cfg: Dict[str, Any]) -> Tuple[Optional[str], Optional[str]]:
    cfg, subtitle, png_path, xlsx_path, fig_size = _render_gp_bridge_assets(df, cfg)
    excel_path_out = None
    try:
        if cfg.get("excel", {}).get("enabled", True):
            export_excel_with_image(cfg, png_path, xlsx_path, subtitle, fig_size_in=fig_size)
            excel_path_out = xlsx_path
    finally:
        try:
            if os.path.exists(png_path):
                os.remove(png_path)
        except Exception as e:
            print("Warnung: Konnte temporäres PNG nicht löschen:", e)
    return None, excel_path_out


def run_gp_bridge_in_workbook(cfg: dict, wb) -> str:
    """Fast Track / session mode: write GP/NP Bridge into an open workbook."""
    file_path = str(cfg.get("file_path") or "").strip()
    sheet_name = str(cfg.get("sheet_name") or "").strip()
    if not file_path:
        raise ValueError("GP_Bridge requires file_path")
    df = pd.read_excel(file_path, sheet_name=sheet_name or 0, engine="openpyxl")

    out_dir = str(cfg.get("output_file_path") or cfg.get("output_dir") or ".")
    cfg = {**cfg, "output_dir": out_dir, "output_prefix": cfg.get("output_prefix") or "gp_bridge"}
    excel_cfg = dict(cfg.get("excel") or {})
    excel_cfg["enabled"] = True
    cfg["excel"] = excel_cfg

    if "fy_end_month" in cfg and "fiscal_year_end_month" not in cfg:
        cfg["fiscal_year_end_month"] = cfg["fy_end_month"]
    if "fy_end_day" in cfg and "fiscal_year_end_day" not in cfg:
        cfg["fiscal_year_end_day"] = cfg["fy_end_day"]

    cfg, subtitle, png_path, _xlsx, fig_size = _render_gp_bridge_assets(df, cfg)
    sheet_title = str(
        (cfg.get("excel") or {}).get("sheet_name")
        or cfg.get("base_sheet_name")
        or bridge_sheet_title(cfg)
    )
    try:
        from funktionssammlung import get_next_sheet_name_from_wb

        target = get_next_sheet_name_from_wb(wb, sheet_title)
    except Exception:
        target = sheet_title
        if target in wb.sheetnames:
            target = f"{target}_1"

    try:
        written = export_excel_with_image(
            cfg,
            png_path,
            xlsx_path="",
            subtitle=subtitle,
            fig_size_in=fig_size,
            wb=wb,
            save=False,
            sheet_name=target,
        )
        return str(written or target)
    finally:
        try:
            if os.path.exists(png_path):
                os.remove(png_path)
        except Exception:
            pass


if __name__ == "__main__":
    file_path = ""
    df = pd.read_excel(file_path, sheet_name="Data Template", engine="openpyxl")

    _, xlsx_path = run_gp_bridge(df, CONFIG)
    print("Excel gespeichert unter:", xlsx_path)
