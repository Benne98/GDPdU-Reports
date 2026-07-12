"""Fixed assets rollforward — annual Anlagengitter extracts to grouped bridge table."""
from __future__ import annotations

import json
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

BASE_DIR = Path(__file__).resolve().parent
SCRIPTS_DIR = BASE_DIR / "scripts"
if str(BASE_DIR) not in sys.path:
    sys.path.insert(0, str(BASE_DIR))
if str(SCRIPTS_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPTS_DIR))

from funktionssammlung import (  # noqa: E402
    _get_sheet_header_map,
    apply_filters,
    build_chunked_row_sum_formula,
    build_formula_filter_branches,
    build_output_file_path,
    build_sumifs_formula_body,
    ensure_output_writable,
    open_session_workbook,
    replace_workbook_sheet,
    resolve_column_names_to_df,
    source_range_ref,
    write_source_df_to_ws,
)
from gst_excel_theme import THEME  # noqa: E402
from databook_excel_layout import (  # noqa: E402
    AGGREGATED_BLOCK_TITLE,
    BS_RECON_SHEET,
    FONT_MAPPING,
    bs_recon_aggregated_ref,
    check_row_groups_after_table,
    collapse_check_portfolio,
    find_bs_recon_row_contains,
    find_recon_block_start_col,
)
from databook_workbook import FA_ROLLF_OUTPUT_SHEET  # noqa: E402

FA_TABLE_TITLE = "FA roll forward"

FREE_COL = 1
PROJECT_TITLE_ROW = 1
SUBTITLE_ROW = 2
TABLE_TITLE_ROW = 6
HEADER_ROW_7 = 7
HEADER_ROW = 8
DATA_START_ROW = 9

MAP_COL_WIDTH = 12  # databook helper block (Lead_BS, cashflow, Working_capital, …)
VALUE_COL_WIDTH = 8.0
LABEL_COL_WIDTH = 32.0
ROW_HEIGHT = 12
NUM_FMT = "#,##0;(#,##0);-"
UNIT_LABEL = "kEUR"
INDENT = "    "
_ZERO_TOL = 1e-6
_MONTH_ABBR = {
    1: "Jan",
    2: "Feb",
    3: "Mar",
    4: "Apr",
    5: "May",
    6: "Jun",
    7: "Jul",
    8: "Aug",
    9: "Sep",
    10: "Oct",
    11: "Nov",
    12: "Dec",
}

FONT_BASE = THEME.font_base
FONT_BOLD = THEME.font_bold
FONT_HEADER = THEME.font_header
FONT_CHECK_RED = Font(name=THEME.font_name, size=THEME.font_size, color=THEME.delta_negative)
FONT_PROJECT = Font(name=THEME.font_name, size=THEME.font_size_title, color=THEME.text_brand_title)
FONT_SUBTITLE = Font(name=THEME.font_name, size=THEME.font_size_subtitle, color=THEME.text_brand_title)
FILL_GREY = THEME.fill_tech
FILL_HEADER = THEME.fill_header
FILL_SUBTOTAL = THEME.fill_subtotal
FILL_WHITE = THEME.fill_white
FILL_PERIOD = THEME.fill_period
FILL_YELLOW = PatternFill("solid", fgColor="FFFF00")
ALIGN_LEFT = Alignment(horizontal="left", vertical="center")
ALIGN_RIGHT = Alignment(horizontal="right", vertical="center")
ALIGN_CENTER = Alignment(horizontal="center", vertical="center")
BORDER_SUBTOTAL_TOP = THEME.border_subtotal_top
BORDER_HEADER_BOTTOM = Border(bottom=Side(style="medium", color=THEME.border_strong))

MOVEMENT_KEYS = ("additions", "disposals", "depreciation")
MOVEMENT_LABELS = {
    "additions": "Add.",
    "disposals": "Disp.",
    "depreciation": "D&A",
}
@dataclass(frozen=True)
class FaLayout:
    """Grey helper block A..(POS-1) with spacer cols between group criteria; POS = table start."""

    helper_start_col: int
    helper_end_col: int
    group_crit_cols: tuple[int, ...]
    pos_col: int
    first_data_col: int


def fa_column_layout(num_group_cols: int) -> FaLayout:
    """A..POS-1 grey (placeholders + 1–2 criteria cols + internal spacers), like databook."""
    n = max(0, min(2, int(num_group_cols)))
    col = 1
    # leading placeholder (A)
    group_cols: list[int] = []
    col += 1
    for i in range(n):
        if i > 0:
            col += 1  # spacer between criteria columns
        group_cols.append(col)
        col += 1
    col += 1  # trailing spacer before POS
    helper_end = col - 1
    pos_col = col
    return FaLayout(
        helper_start_col=1,
        helper_end_col=helper_end,
        group_crit_cols=tuple(group_cols),
        pos_col=pos_col,
        first_data_col=pos_col + 1,
    )


# Default export for tests / one-level grouping.
_DEFAULT_LAYOUT = fa_column_layout(1)
POS_COL = _DEFAULT_LAYOUT.pos_col
FIRST_DATA_COL = _DEFAULT_LAYOUT.first_data_col
GROUP_CRIT_COL_START = _DEFAULT_LAYOUT.group_crit_cols[0] if _DEFAULT_LAYOUT.group_crit_cols else 2

_FY_GRID_RE = re.compile(r"^FY(19|20)\d{2}$", re.IGNORECASE)
_YTD_GRID_RE = re.compile(r"^YTD(19|20)\d{2}$", re.IGNORECASE)


@dataclass
class PeriodSpec:
    label: str
    file_path: str
    sheet_name: str = ""


@dataclass
class RowSpec:
    row_type: str  # leaf | parent | total
    label: str
    level: int
    group_key: tuple[str, ...] | None = None


@dataclass
class BridgeColumn:
    kind: str  # period | movement
    col_idx: int
    header: str
    movement_key: str | None = None
    source_period_idx: int | None = None
    segment_idx: int | None = None
    is_first_period: bool = False
    hidden: bool = False


def col_letter(idx: int) -> str:
    return get_column_letter(idx)


def _fy_end_month_from_cfg(cfg: dict) -> int:
    try:
        m = int(cfg.get("fy_end_month") or 12)
    except (TypeError, ValueError):
        m = 12
    return m if 1 <= m <= 12 else 12


def grid_label_to_display(label: str, *, fy_end_month: int = 12) -> str:
    """Stichtag label: month before FY start (e.g. FY2024 → Dec23A)."""
    s = str(label or "").strip()
    month = _MONTH_ABBR.get(fy_end_month, "Dec")
    m_fy = _FY_GRID_RE.match(s)
    m_ytd = _YTD_GRID_RE.match(s)
    if m_fy or m_ytd:
        raw_yr = s[2:] if m_fy else s[3:]
        try:
            yr = int(raw_yr)
        except ValueError:
            return s
        if yr < 100:
            yr += 2000
        stichtag_year = yr - 1
        return f"{month}{str(stichtag_year)[-2:]}A"
    return s


def source_sheet_name(period_label: str, *, fy_end_month: int = 12) -> str:
    display = grid_label_to_display(period_label, fy_end_month=fy_end_month)
    return f"__SOURCE__FA_{display}"[:31]


def _sentence_case_title(text: str) -> str:
    parts = str(text or "").strip().split()
    if not parts:
        return ""
    return parts[0].capitalize() + (
        (" " + " ".join(p.lower() for p in parts[1:])) if len(parts) > 1 else ""
    )


def _parse_numeric(series: pd.Series) -> pd.Series:
    return pd.to_numeric(series, errors="coerce").fillna(0.0)


def normalize_config(cfg: dict) -> dict:
    out = dict(cfg)
    cols = dict(out.get("columns") or {})
    for key in ("opening", "additions", "disposals"):
        if not str(cols.get(key) or "").strip():
            raise ValueError(f"columns.{key} is required")
    dep_cols = cols.get("depreciation_cols")
    if not isinstance(dep_cols, list) or not dep_cols:
        dep = cols.get("depreciation")
        if isinstance(dep, list):
            dep_cols = [str(x).strip() for x in dep if str(x).strip()]
        elif dep:
            dep_cols = [str(dep).strip()]
        else:
            dep_cols = []
    dep_cols = [str(x).strip() for x in dep_cols if str(x).strip()]
    if not dep_cols:
        raise ValueError("columns.depreciation_cols is required")
    cols["depreciation_cols"] = dep_cols
    cols.pop("depreciation", None)
    out["columns"] = cols
    out["group_cols"] = [str(c).strip() for c in (out.get("group_cols") or []) if str(c).strip()][:2]
    periods = out.get("periods")
    if not isinstance(periods, list) or not periods:
        raise ValueError("periods[] is required")
    norm_periods: list[dict] = []
    for p in periods:
        if not isinstance(p, dict):
            continue
        label = str(p.get("label") or "").strip()
        fp = str(p.get("file_path") or "").strip()
        if not label or not fp:
            continue
        norm_periods.append(
            {
                "label": label,
                "file_path": fp,
                "sheet_name": str(p.get("sheet_name") or "").strip(),
            }
        )
    if not norm_periods:
        raise ValueError("periods[] has no valid entries")
    out["periods"] = norm_periods
    out["header_row"] = int(out.get("header_row") or 0)
    out["unit_label"] = str(out.get("unit_label") or UNIT_LABEL)
    scale = out.get("amount_scale")
    out["amount_scale"] = float(scale) if scale not in (None, "", 0) else 1000.0
    out["sheet_name"] = str(out.get("sheet_name") or FA_ROLLF_OUTPUT_SHEET)
    out["total_label"] = str(out.get("total_label") or "Fixed assets")
    out["formula_mode"] = bool(out.get("formula_mode", True))
    bal = str(out.get("balance_check_col") or "Lfd Buchwert").strip()
    out["balance_check_col"] = bal
    raw_filters = out.get("filters", {}) or {}
    if isinstance(raw_filters, list):
        filters = {"enabled": len(raw_filters) > 0, "rules": raw_filters}
    else:
        filters = dict(raw_filters)
    filters.setdefault("enabled", False)
    filters.setdefault("rules", [])
    out["filters"] = filters
    return out


def _resolve_sheet_name(file_path: str, sheet_name: str) -> str:
    if sheet_name:
        return sheet_name
    xl = pd.ExcelFile(file_path)
    return xl.sheet_names[0] if xl.sheet_names else "Sheet1"


def load_period_dataframe(spec: PeriodSpec, header_row: int) -> pd.DataFrame:
    sheet = _resolve_sheet_name(spec.file_path, spec.sheet_name)
    return pd.read_excel(spec.file_path, sheet_name=sheet, header=header_row)



def _amount_column_names(cfg: dict) -> list[str]:
    cols = cfg["columns"]
    return [cols["opening"], cols["additions"], cols["disposals"], *cols["depreciation_cols"]]


def _depreciation_cols_from_cfg(cfg: dict) -> list[str]:
    return list(cfg["columns"]["depreciation_cols"])


def aggregate_period(
    df: pd.DataFrame,
    cfg: dict,
) -> tuple[dict[tuple[str, ...], dict[str, float]], pd.DataFrame]:
    group_cols = list(cfg.get("group_cols") or [])
    amount_names = _amount_column_names(cfg)
    all_needed = list(group_cols) + amount_names
    resolved = resolve_column_names_to_df(df, all_needed)

    g_resolved = resolved[: len(group_cols)]
    a_resolved = resolved[len(group_cols) :]

    work = apply_filters(df.copy(), cfg)
    dep_cols = _depreciation_cols_from_cfg(cfg)
    amount_resolved = list(a_resolved[:3]) + dep_cols
    for res_name in amount_resolved:
        if res_name in work.columns:
            work[res_name] = _parse_numeric(work[res_name])

    opening_col = amount_resolved[0]
    work = work[work[opening_col].notna() | work[amount_resolved[1]].notna()]

    key_map: dict[str, str | list[str]] = {
        "opening": amount_resolved[0],
        "additions": amount_resolved[1],
        "disposals": amount_resolved[2],
        "depreciation": dep_cols,
    }

    if group_cols:
        grouped = work.groupby(g_resolved, dropna=False)
        agg: dict[tuple[str, ...], dict[str, float]] = {}
        for key, grp in grouped:
            if not isinstance(key, tuple):
                key = (key,)
            key = tuple("" if pd.isna(v) else str(v).strip() for v in key)
            row: dict[str, float] = {}
            for k, col in key_map.items():
                if k == "depreciation" and isinstance(col, list):
                    row[k] = sum(float(grp[c].sum()) for c in col)
                else:
                    row[k] = float(grp[col].sum())
            row["closing"] = compute_closing(row)
            agg[key] = row
    else:
        row = {}
        for k, col in key_map.items():
            if k == "depreciation" and isinstance(col, list):
                row[k] = sum(float(work[c].sum()) for c in col)
            else:
                row[k] = float(work[col].sum())
        row["closing"] = compute_closing(row)
        agg[()] = row

    return agg, work


def compute_closing(row: dict[str, float]) -> float:
    total = float(row.get("opening") or 0)
    for k in MOVEMENT_KEYS:
        total += float(row.get(k) or 0)
    return total


def visible_movements_for_agg(agg: dict[tuple[str, ...], dict[str, float]]) -> list[str]:
    visible: list[str] = []
    for mk in MOVEMENT_KEYS:
        total = sum(float(row.get(mk) or 0) for row in agg.values())
        if abs(total) > _ZERO_TOL:
            visible.append(mk)
    return visible


def build_row_specs(
    group_cols: list[str],
    keys: list[tuple[str, ...]],
    sort_values: dict[tuple[str, ...], float] | None = None,
) -> list[RowSpec]:
    if not group_cols:
        return []

    def _sort_key(k: tuple[str, ...]) -> tuple:
        val = -(sort_values.get(k, 0.0) if sort_values else 0.0)
        return (val, (k[0] if k else "").lower())

    if len(group_cols) == 1:
        return [
            RowSpec(row_type="leaf", label=k[0] if k else "", level=1, group_key=k)
            for k in sorted(keys, key=_sort_key)
        ]
    by_l1: dict[str, list[tuple[str, ...]]] = defaultdict(list)
    for k in keys:
        by_l1[k[0] if k else ""].append(k)
    rows: list[RowSpec] = []
    for l1 in sorted(by_l1.keys(), key=str.lower):
        children = sorted(by_l1[l1], key=_sort_key)
        for k in children:
            child_label = k[1] if len(k) > 1 else ""
            rows.append(RowSpec(row_type="leaf", label=child_label, level=2, group_key=k))
        rows.append(RowSpec(row_type="parent", label=l1, level=1, group_key=None))
    return rows


def build_bridge_columns(
    period_labels: list[str],
    source_aggs: list[dict[tuple[str, ...], dict[str, float]]],
    *,
    fy_end_month: int = 12,
    first_data_col: int = FIRST_DATA_COL,
) -> tuple[list[BridgeColumn], list[int]]:
    """Dec22A | Add | Disp | D&A | Dec23A | … — each segment from its own source period."""
    columns: list[BridgeColumn] = []
    col = first_data_col
    n = len(period_labels)
    if n == 0:
        return columns, []

    columns.append(
        BridgeColumn(
            kind="period",
            col_idx=col,
            header=grid_label_to_display(period_labels[0], fy_end_month=fy_end_month),
            source_period_idx=0,
            is_first_period=True,
        )
    )
    col += 1

    for seg in range(n - 1):
        mov_agg = source_aggs[seg]
        visible = visible_movements_for_agg(mov_agg)
        for mk in MOVEMENT_KEYS:
            if mk not in visible:
                continue
            columns.append(
                BridgeColumn(
                    kind="movement",
                    col_idx=col,
                    header=MOVEMENT_LABELS[mk],
                    movement_key=mk,
                    source_period_idx=seg,
                    segment_idx=seg,
                )
            )
            col += 1

        period_idx = seg + 1
        columns.append(
            BridgeColumn(
                kind="period",
                col_idx=col,
                header=grid_label_to_display(period_labels[period_idx], fy_end_month=fy_end_month),
                source_period_idx=period_idx,
                segment_idx=seg,
            )
        )
        col += 1

    return columns, []


def _divide_by_scale(cfg: dict) -> bool:
    return float(cfg.get("amount_scale") or 1000) == 1000.0


def _group_sumifs_pairs(
    source_sheet: str,
    header_map: dict[str, int],
    group_cols: list[str],
    row_idx: int,
    *,
    group_crit_cols: tuple[int, ...],
) -> list[tuple[str, str]]:
    pairs: list[tuple[str, str]] = []
    for i, gc in enumerate(group_cols):
        if i >= len(group_crit_cols):
            break
        col_idx = group_crit_cols[i]
        rng = source_range_ref(source_sheet, header_map, gc)
        crit = f"${col_letter(col_idx)}{row_idx}"
        pairs.append((rng, crit))
    return pairs


def _formula_filter_branches(cfg: dict, source_sheet: str, header_map: dict[str, int]) -> list[list[tuple[str, str]]]:
    try:
        return build_formula_filter_branches(cfg, source_sheet, header_map)
    except ValueError:
        return [[]]


def _fa_sumifs_formula(
    source_sheet: str,
    header_map: dict[str, int],
    amount_header: str,
    row_idx: int,
    group_cols: list[str],
    cfg: dict,
    *,
    group_crit_cols: tuple[int, ...],
) -> str:
    amount_rng = source_range_ref(source_sheet, header_map, amount_header)
    base_pairs = _group_sumifs_pairs(
        source_sheet, header_map, group_cols, row_idx, group_crit_cols=group_crit_cols
    )
    filter_branches = _formula_filter_branches(cfg, source_sheet, header_map)
    body = build_sumifs_formula_body(
        amount_rng,
        base_pairs,
        filter_branches,
        divide_by_1000=_divide_by_scale(cfg),
    )
    return f"=IFERROR({body},\"\")"


def _fa_depreciation_formula(
    source_sheet: str,
    header_map: dict[str, int],
    row_idx: int,
    group_cols: list[str],
    cfg: dict,
    *,
    group_crit_cols: tuple[int, ...],
) -> str:
    headers = _depreciation_cols_from_cfg(cfg)
    if len(headers) == 1:
        return _fa_sumifs_formula(
            source_sheet,
            header_map,
            headers[0],
            row_idx,
            group_cols,
            cfg,
            group_crit_cols=group_crit_cols,
        )
    parts: list[str] = []
    base_pairs = _group_sumifs_pairs(
        source_sheet, header_map, group_cols, row_idx, group_crit_cols=group_crit_cols
    )
    filter_branches = _formula_filter_branches(cfg, source_sheet, header_map)
    for hdr in headers:
        amount_rng = source_range_ref(source_sheet, header_map, hdr)
        parts.append(
            build_sumifs_formula_body(
                amount_rng,
                base_pairs,
                filter_branches,
                divide_by_1000=_divide_by_scale(cfg),
            )
        )
    return f"=IFERROR({'+'.join(parts)},\"\")"


def _parent_sum_formula(row_idx: int, child_rows: list[int], col_idx: int) -> str:
    if not child_rows:
        return "=0"
    refs = [f"{col_letter(col_idx)}{r}" for r in child_rows]
    if len(refs) == 1:
        return f"={refs[0]}"
    inner = "+".join(f"IFERROR({r},0)" for r in refs)
    return f"=IFERROR({inner},\"\")"


def _bridge_delta_formula(total_row: int, next_period_col: int, prev_period_col: int, movement_cols: list[int]) -> str:
    nxt = f"{col_letter(next_period_col)}{total_row}"
    parts = [f"IFERROR({col_letter(prev_period_col)}{total_row},0)"]
    for c in movement_cols:
        parts.append(f"IFERROR({col_letter(c)}{total_row},0)")
    return f"={nxt}-({'+'.join(parts)})"


def _build_bridge_segments(columns: list[BridgeColumn]) -> list[tuple[int, list[int], int]]:
    """(prev_period_col, movement_cols, next_period_col) per segment."""
    segments: list[tuple[int, list[int], int]] = []
    prev_period: int | None = None
    movement_cols: list[int] = []
    for bc in columns:
        if bc.kind == "period":
            if prev_period is not None and movement_cols:
                segments.append((prev_period, movement_cols, bc.col_idx))
            prev_period = bc.col_idx
            movement_cols = []
        elif bc.kind == "movement":
            movement_cols.append(bc.col_idx)
    return segments


def _source_total_formula(
    source_sheet: str,
    header_map: dict[str, int],
    balance_header: str,
    cfg: dict,
) -> str:
    try:
        rng = source_range_ref(source_sheet, header_map, balance_header)
    except ValueError:
        rng = source_range_ref(source_sheet, header_map, cfg["columns"]["opening"])
    if _divide_by_scale(cfg):
        return f"=IFERROR(SUM({rng})/1000,\"\")"
    scale = float(cfg["amount_scale"])
    return f"=IFERROR(SUM({rng})/{scale},\"\")"


def _apply_fa_row_border_band(
    ws,
    row: int,
    layout: FaLayout,
    border: Border,
    *,
    last_used_col: int,
) -> None:
    """Subtotal/header borders on filled helper cols + table; skip A and spacer columns."""
    for cc in layout.group_crit_cols:
        ws.cell(row, cc).border = border
    for cc in range(layout.pos_col, last_used_col + 1):
        ws.cell(row, cc).border = border


def _apply_sheet_layout(
    ws,
    layout: FaLayout,
    *,
    last_used_col: int,
    last_fill_row: int,
) -> None:
    """Grey full helper block A..spacer; white from POS onward."""
    fill_end_col = last_used_col + 40
    for rr in range(1, last_fill_row + 1):
        ws.row_dimensions[rr].height = ROW_HEIGHT if rr != PROJECT_TITLE_ROW else 36
        for cc in range(layout.helper_start_col, layout.pos_col):
            ws.cell(rr, cc).fill = FILL_GREY
        for cc in range(layout.pos_col, fill_end_col + 1):
            ws.cell(rr, cc).fill = FILL_WHITE

    for cc in range(layout.helper_start_col, layout.pos_col):
        cl = col_letter(cc)
        ws.column_dimensions[cl].outlineLevel = 2
        ws.column_dimensions[cl].hidden = True
    ws.column_dimensions[col_letter(layout.helper_end_col)].collapsed = True

    pos_letter = col_letter(layout.pos_col)
    ws.column_dimensions[pos_letter].outlineLevel = 0
    ws.column_dimensions[pos_letter].hidden = False


def _paint_header_band(
    ws,
    layout: FaLayout,
    columns: list[BridgeColumn],
    *,
    unit_label: str,
    group_col_labels: list[str] | None = None,
) -> None:
    labels = group_col_labels or []
    for i, crit_col in enumerate(layout.group_crit_cols):
        label = labels[i] if i < len(labels) else ""
        ws.cell(HEADER_ROW_7, crit_col).fill = FILL_HEADER
        ws.cell(HEADER_ROW_7, crit_col).font = FONT_HEADER
        hc = ws.cell(HEADER_ROW, crit_col, label)
        hc.font = FONT_HEADER
        hc.alignment = ALIGN_RIGHT
        hc.border = BORDER_HEADER_BOTTOM
        hc.fill = FILL_HEADER

    for bc in columns:
        ws.cell(HEADER_ROW_7, bc.col_idx).fill = (
            FILL_PERIOD if bc.kind == "period" else FILL_HEADER
        )
        ws.cell(HEADER_ROW_7, bc.col_idx).font = FONT_HEADER
        ws.cell(HEADER_ROW_7, bc.col_idx).alignment = ALIGN_RIGHT

        hc = ws.cell(HEADER_ROW, bc.col_idx, bc.header)
        hc.font = FONT_HEADER
        hc.alignment = ALIGN_RIGHT
        hc.border = BORDER_HEADER_BOTTOM
        hc.fill = FILL_PERIOD if bc.kind == "period" else FILL_HEADER

    pos_top = ws.cell(HEADER_ROW_7, layout.pos_col)
    pos_top.fill = FILL_HEADER
    pos_top.font = FONT_HEADER
    pos_top.alignment = ALIGN_LEFT

    pos_hdr = ws.cell(HEADER_ROW, layout.pos_col, unit_label)
    pos_hdr.fill = FILL_HEADER
    pos_hdr.border = BORDER_HEADER_BOTTOM
    pos_hdr.font = FONT_HEADER
    pos_hdr.alignment = ALIGN_LEFT


def _table_band_cols(layout: FaLayout, columns: list[BridgeColumn], last_used_col: int) -> list[int]:
    cols = {layout.pos_col, *{bc.col_idx for bc in columns if not bc.hidden}}
    return sorted(c for c in cols if layout.pos_col <= c <= last_used_col)


def _paint_check_source_yellow(
    ws,
    layout: FaLayout,
    source_row: int,
    table_cols: list[int],
) -> None:
    for col_idx in table_cols:
        ws.cell(source_row, col_idx).fill = FILL_YELLOW


def _style_value_cell(cell, *, bold: bool = False, fill=None, red: bool = False) -> None:
    cell.number_format = NUM_FMT
    cell.alignment = ALIGN_RIGHT
    if red:
        cell.font = FONT_CHECK_RED
    else:
        cell.font = FONT_BOLD if bold else FONT_BASE
    if fill is not None:
        cell.fill = fill


def write_workbook(
    cfg: dict,
    source_aggs: list[dict],
    source_dfs: list[pd.DataFrame],
    source_header_maps: list[dict[str, int]],
) -> str:
    cfg = normalize_config(cfg)
    out_path = build_output_file_path(cfg)
    ensure_output_writable(out_path)

    group_cols = cfg["group_cols"]
    layout = fa_column_layout(len(group_cols))
    period_labels = [p["label"] for p in cfg["periods"]]
    period_display_labels = [
        grid_label_to_display(lbl, fy_end_month=_fy_end_month_from_cfg(cfg)) for lbl in period_labels
    ]
    formula_mode = bool(cfg.get("formula_mode", True))

    all_keys: set[tuple[str, ...]] = set()
    for agg in source_aggs:
        all_keys.update(agg.keys())
    keys = list(all_keys)
    last_agg = source_aggs[-1] if source_aggs else {}
    sort_values = {k: float(last_agg.get(k, {}).get("opening") or 0) for k in keys}

    row_specs = build_row_specs(group_cols, keys, sort_values)
    fy_end_month = _fy_end_month_from_cfg(cfg)
    columns, _spacer_cols = build_bridge_columns(
        period_labels,
        source_aggs,
        fy_end_month=fy_end_month,
        first_data_col=layout.first_data_col,
    )
    period_col_indices = [c.col_idx for c in columns if c.kind == "period"]
    last_used_col = max((c.col_idx for c in columns), default=layout.first_data_col)
    opening_header = cfg["columns"]["opening"]
    group_crit_cols = layout.group_crit_cols

    detail_rows: list[int] = []
    parent_rows: list[int] = []
    excel_rows: list[tuple[RowSpec, int]] = []
    r = DATA_START_ROW
    for spec in row_specs:
        excel_rows.append((spec, r))
        if spec.row_type == "leaf":
            detail_rows.append(r)
        elif spec.row_type == "parent":
            parent_rows.append(r)
        r += 1
    total_row = r
    last_table_row = total_row

    wb = open_session_workbook(cfg)
    ws = replace_workbook_sheet(wb, cfg["sheet_name"][:31])

    project = _sentence_case_title(str(cfg.get("title") or cfg.get("project_name") or "Project"))
    company = str(cfg.get("company") or cfg.get("group_name") or "").strip()
    sheet_label = str(cfg.get("sheet_name") or "Fixed assets rollforward").strip()

    ws.cell(PROJECT_TITLE_ROW, layout.pos_col, project).font = FONT_PROJECT
    ws.cell(SUBTITLE_ROW, layout.pos_col, company).font = FONT_SUBTITLE

    period_range = ""
    if period_display_labels:
        period_range = f" {period_display_labels[0]} - {period_display_labels[-1]}"
    table_heading = (
        f"{company} | {FA_TABLE_TITLE}{period_range}" if company else f"{FA_TABLE_TITLE}{period_range}"
    )
    table_title = ws.cell(TABLE_TITLE_ROW, layout.pos_col, table_heading)
    table_title.font = FONT_HEADER
    table_title.alignment = ALIGN_LEFT

    _paint_header_band(
        ws,
        layout,
        columns,
        unit_label=cfg["unit_label"],
        group_col_labels=group_cols,
    )

    for bc in columns:
        if bc.hidden:
            ws.column_dimensions[col_letter(bc.col_idx)].hidden = True
            ws.column_dimensions[col_letter(bc.col_idx)].width = 0

    for spec, er in excel_rows:
        display = spec.label
        if spec.level == 2:
            display = f"{INDENT}{display}"
        ws.cell(er, layout.pos_col, display).alignment = ALIGN_LEFT
        ws.cell(er, layout.pos_col).font = FONT_BASE
        if spec.group_key is not None and layout.group_crit_cols:
            for i, val in enumerate(spec.group_key):
                if i < len(group_cols) and i < len(layout.group_crit_cols):
                    cell = ws.cell(er, layout.group_crit_cols[i], val)
                    cell.font = FONT_MAPPING

        is_parent = spec.row_type == "parent"
        if is_parent:
            ws.row_dimensions[er].outlineLevel = 0
            _apply_fa_row_border_band(
                ws, er, layout, BORDER_SUBTOTAL_TOP, last_used_col=last_used_col
            )
        elif spec.row_type == "leaf" and len(group_cols) >= 2:
            ws.row_dimensions[er].outlineLevel = 1

    total_spec_label = cfg["total_label"]
    ws.cell(total_row, layout.pos_col, total_spec_label).alignment = ALIGN_LEFT
    ws.cell(total_row, layout.pos_col).font = FONT_BOLD
    _apply_fa_row_border_band(
        ws, total_row, layout, BORDER_SUBTOTAL_TOP, last_used_col=last_used_col
    )

    leaf_rows_by_parent: dict[str, list[int]] = defaultdict(list)
    if len(group_cols) >= 2:
        for spec, er in excel_rows:
            if spec.row_type == "leaf" and spec.group_key:
                leaf_rows_by_parent[spec.group_key[0]].append(er)

    for spec, er in excel_rows:
        is_parent = spec.row_type == "parent"
        row_fill = FILL_SUBTOTAL if is_parent else FILL_WHITE

        for bc in columns:
            cell = ws.cell(er, bc.col_idx)
            period_fill = FILL_PERIOD if bc.kind == "period" and not is_parent else row_fill
            cell.fill = period_fill

            if not formula_mode:
                continue

            pi = bc.source_period_idx or 0
            src = source_sheet_name(period_labels[pi], fy_end_month=fy_end_month)
            hmap = source_header_maps[pi]

            if is_parent:
                child_rows = leaf_rows_by_parent.get(spec.label, [])
                cell.value = _parent_sum_formula(er, child_rows, bc.col_idx)
            elif bc.kind == "period":
                cell.value = _fa_sumifs_formula(
                    src, hmap, opening_header, er, group_cols, cfg, group_crit_cols=group_crit_cols
                )
            elif bc.kind == "movement":
                mk = bc.movement_key or "additions"
                if mk == "depreciation":
                    cell.value = _fa_depreciation_formula(
                        src, hmap, er, group_cols, cfg, group_crit_cols=group_crit_cols
                    )
                else:
                    cell.value = _fa_sumifs_formula(
                        src, hmap, cfg["columns"][mk], er, group_cols, cfg, group_crit_cols=group_crit_cols
                    )

            _style_value_cell(cell, bold=is_parent, fill=period_fill if bc.kind == "period" and not is_parent else row_fill)

    for bc in columns:
        if bc.kind == "period":
            leaf_only = [r for spec, r in excel_rows if spec.row_type == "leaf"]
            cell = ws.cell(total_row, bc.col_idx)
            if formula_mode and leaf_only:
                cell.value = build_chunked_row_sum_formula(leaf_only, bc.col_idx)
            cell.font = FONT_BOLD
            cell.number_format = NUM_FMT
            cell.alignment = ALIGN_RIGHT
            cell.fill = FILL_SUBTOTAL
            cell.border = BORDER_SUBTOTAL_TOP
        elif bc.kind == "movement":
            leaf_only = [r for spec, r in excel_rows if spec.row_type == "leaf"]
            cell = ws.cell(total_row, bc.col_idx)
            if formula_mode and leaf_only:
                cell.value = build_chunked_row_sum_formula(leaf_only, bc.col_idx)
            cell.font = FONT_BOLD
            cell.number_format = NUM_FMT
            cell.alignment = ALIGN_RIGHT
            cell.fill = FILL_SUBTOTAL
            cell.border = BORDER_SUBTOTAL_TOP

    for i, df in enumerate(source_dfs):
        sheet_title = source_sheet_name(period_labels[i], fy_end_month=fy_end_month)
        if sheet_title in wb.sheetnames:
            del wb[sheet_title]
        src_ws = wb.create_sheet(sheet_title)
        stub_cfg = {
            "file_path": cfg["periods"][i]["file_path"],
            "sheet_name": _resolve_sheet_name(
                cfg["periods"][i]["file_path"], cfg["periods"][i].get("sheet_name", "")
            ),
        }
        write_source_df_to_ws(src_ws, df, stub_cfg)

    (
        check_source_row,
        check_delta_row,
        check_bsrec_row,
        check_bsrec_delta_row,
        bridge_check_row,
    ) = check_row_groups_after_table(last_table_row, [2, 2, 1])
    bridge_segments = _build_bridge_segments(columns)

    for label, rr in (
        ("Source - Fixed asset files", check_source_row),
        ("Check", check_delta_row),
        ("Source - BS Reconciliation", check_bsrec_row),
        ("Check", check_bsrec_delta_row),
        ("Bridge check", bridge_check_row),
    ):
        c = ws.cell(rr, layout.pos_col, label)
        c.alignment = ALIGN_LEFT
        c.font = FONT_BASE

    for bc in columns:
        if bc.kind != "period":
            continue
        pi = bc.source_period_idx or 0
        src = source_sheet_name(period_labels[pi], fy_end_month=fy_end_month)
        hmap = source_header_maps[pi]
        col = bc.col_idx

        src_cell = ws.cell(check_source_row, col)
        if formula_mode:
            src_cell.value = _source_total_formula(src, hmap, opening_header, cfg)
        _style_value_cell(src_cell)

        total_cell = ws.cell(total_row, col)
        delta_cell = ws.cell(check_delta_row, col)
        delta_cell.value = f"={total_cell.coordinate}-{src_cell.coordinate}"
        _style_value_cell(delta_cell, red=True)

        bs_cell = ws.cell(check_bsrec_row, col)
        if formula_mode and BS_RECON_SHEET in wb.sheetnames:
            ws_rec = wb[BS_RECON_SHEET]
            agg_start = find_recon_block_start_col(ws_rec, AGGREGATED_BLOCK_TITLE)
            fa_row = find_bs_recon_row_contains(ws_rec, str(cfg.get("total_label") or "Fixed assets"))
            if agg_start is not None and fa_row is not None:
                rec_col = agg_start + (bc.source_period_idx or 0)
                bs_cell.value = bs_recon_aggregated_ref(fa_row, rec_col)
        _style_value_cell(bs_cell)

        bs_delta = ws.cell(check_bsrec_delta_row, col)
        bs_delta.value = f"={total_cell.coordinate}-{bs_cell.coordinate}"
        _style_value_cell(bs_delta, red=True)

    for prev_col, mov_cols, next_col in bridge_segments:
        delta_cell = ws.cell(bridge_check_row, next_col)
        delta_cell.value = _bridge_delta_formula(total_row, next_col, prev_col, mov_cols)
        _style_value_cell(delta_cell, red=True)

    ws.column_dimensions[col_letter(layout.pos_col)].width = LABEL_COL_WIDTH
    for cc in range(layout.helper_start_col, layout.pos_col):
        ws.column_dimensions[col_letter(cc)].width = MAP_COL_WIDTH
    for bc in columns:
        if not bc.hidden:
            ws.column_dimensions[col_letter(bc.col_idx)].width = VALUE_COL_WIDTH

    collapse_check_portfolio(ws, [check_source_row, bridge_check_row])

    ws.sheet_view.showOutlineSymbols = True
    ws.sheet_properties.outlinePr.summaryBelow = True
    ws.sheet_properties.outlinePr.applyStyles = True

    _apply_sheet_layout(
        ws,
        layout,
        last_used_col=last_used_col,
        last_fill_row=bridge_check_row + 100,
    )

    _paint_header_band(
        ws,
        layout,
        columns,
        unit_label=cfg["unit_label"],
        group_col_labels=group_cols,
    )
    for spec, er in excel_rows:
        if spec.row_type == "parent":
            for bc in columns:
                ws.cell(er, bc.col_idx).fill = FILL_SUBTOTAL
        else:
            for bc in columns:
                if bc.kind == "period":
                    ws.cell(er, bc.col_idx).fill = FILL_PERIOD
    for bc in columns:
        ws.cell(total_row, bc.col_idx).fill = FILL_SUBTOTAL
        ws.cell(total_row, bc.col_idx).border = BORDER_SUBTOTAL_TOP
    table_cols = _table_band_cols(layout, columns, last_used_col)
    _paint_check_source_yellow(ws, layout, check_source_row, table_cols)
    _paint_check_source_yellow(ws, layout, check_bsrec_row, table_cols)

    from databook_workbook import reorder_workbook_sheets

    reorder_workbook_sheets(wb)
    wb.save(out_path)
    return out_path


def run_fixed_assets_rollf(cfg: dict) -> str:
    cfg = normalize_config(cfg)
    header_row = int(cfg["header_row"])
    period_specs = [
        PeriodSpec(p["label"], p["file_path"], p.get("sheet_name", "")) for p in cfg["periods"]
    ]

    source_aggs: list[dict[tuple[str, ...], dict[str, float]]] = []
    source_dfs: list[pd.DataFrame] = []
    source_header_maps: list[dict[str, int]] = []

    for spec in period_specs:
        df = load_period_dataframe(spec, header_row)
        agg, work = aggregate_period(df, cfg)
        source_aggs.append(agg)
        source_dfs.append(work)
        tmp_wb = Workbook()
        tmp_ws = tmp_wb.active
        stub = {
            "file_path": spec.file_path,
            "sheet_name": _resolve_sheet_name(spec.file_path, spec.sheet_name),
        }
        write_source_df_to_ws(tmp_ws, work, stub)
        source_header_maps.append(_get_sheet_header_map(tmp_ws))

    return write_workbook(cfg, source_aggs, source_dfs, source_header_maps)


def main() -> None:
    if len(sys.argv) < 2:
        raise SystemExit("Usage: python fixed_assets_rollf.py <config.json>")
    cfg_path = Path(sys.argv[1])
    cfg = json.loads(cfg_path.read_text(encoding="utf-8-sig"))
    out = run_fixed_assets_rollf(cfg)
    print(json.dumps({"success": True, "output_file_path": out}))


if __name__ == "__main__":
    main()
