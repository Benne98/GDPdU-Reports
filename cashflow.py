"""Cashflow statement — EBITDA bridge, Δ working capital from Master_BS, FCF + KPIs."""
from __future__ import annotations

import json
import sys
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import Alignment, Border, Font, Side

PROJECT_ROOT = Path(__file__).resolve().parent
SCRIPTS_DIR = PROJECT_ROOT / "scripts"
BACKEND_DIR = PROJECT_ROOT / "backend"
if str(SCRIPTS_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPTS_DIR))
if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))

from gst_excel_theme import THEME  # noqa: E402
from databook_excel_layout import (  # noqa: E402
    LAYOUT_NA,
    LAYOUT_PL,
    apply_databook_row_border_band,
    find_recon_aggregated_year_cols,
    find_recon_block_year_cols,
    hide_helper_column_group,
    paint_grey_white_canvas,
    paint_header_band,
    prune_zero_value_rows,
    write_kpi_section_title_row,
)
from databook_periods import (  # noqa: E402
    cashflow_delta_spec,
    cashflow_display_periods,
    display_reporting_columns_from_df,
    ordered_month_columns_from_df,
    ordered_reporting_columns_from_df,
    split_fy_and_ytd,
)
from report_row_layout import (  # noqa: E402
    CF_BUCKET_ORDER,
    CF_BUCKET_TOTAL_LABELS,
    build_cf_na_bucket_row_structure,
    compute_l4_sort_metric,
    l3_order_from_mapping,
    l4_order_from_pl_l3_mapping,
    load_na_l3_orders,
    period_columns_for_sort,
    sort_l4_labels,
)

from databook_workbook import (  # noqa: E402
    BS_RECON_MAPPING_FILE,
    CF_NA_L3_ORDER_FILE,
    MASTER_WORKBOOK_STR,
    PL_KONTOMAPPING_FILE,
)

DESKTOP_DIR = PROJECT_ROOT / "Desktop"
DEFAULT_INPUT = MASTER_WORKBOOK_STR
DEFAULT_MAPPING_BS = BS_RECON_MAPPING_FILE
DEFAULT_CF_NA_L3_ORDER = CF_NA_L3_ORDER_FILE
DEFAULT_PL_MAPPING = PL_KONTOMAPPING_FILE

FINANCIAL_RESULT_L3 = "Financial result"

SHEET_MASTER_BS = "Master_BS"
SHEET_MASTER_PL = "Master_PL"
SHEET_PL_RECON = "PL_Reconciliation"
SHEET_LEAD_IS = "Lead_IS"
SHEET_OUT = "Cashflow"

SOURCE_COL_CANDIDATES = ("L5", "L6")
REPORTED_FILTER_VALUE = "reported"
L4_SORT_BASIS = "latest_fy"

_COL_LAYOUT = LAYOUT_NA
MAP_START_COL = _COL_LAYOUT.map_start_col
MAP_END_COL = _COL_LAYOUT.map_end_col
TECH_SPACER_COL = _COL_LAYOUT.spacer_col
POS_COL = _COL_LAYOUT.pos_col

PROJECT_TITLE_ROW = 1
SUBTITLE_ROW = 2
TITLE_ROW = 6
HEADER_ROW = 8
HEADER_ROW_7 = 7
DATA_START_ROW = 9
RECON_BLOCK_TITLE_ROW = 7

ROW_HEIGHT = 12
POS_COL_WIDTH = 32
MAP_COL_WIDTH = 12
VALUE_COL_WIDTH = 7.86

FONT_BASE = THEME.font_base
FONT_BASE_BOLD = THEME.font_bold
FONT_HEADER = THEME.font_header
FONT_MAPPING = Font(name=THEME.font_name, size=THEME.font_size, color=THEME.text_header)
FONT_TITLE = Font(
    name=THEME.font_name, size=THEME.font_size, color=THEME.text_brand_title, bold=True
)
FONT_PROJECT = Font(
    name=THEME.font_name, size=THEME.font_size_title, color=THEME.text_brand_title
)
FONT_SUBTITLE = Font(
    name=THEME.font_name, size=THEME.font_size_subtitle, color=THEME.text_brand_title
)
FONT_KPI = Font(name=THEME.font_name, size=THEME.font_size, color=THEME.text_kpi)

FILL_GREY = THEME.fill_tech
FILL_WHITE = THEME.fill_white
FILL_HEADER = THEME.fill_header
FILL_SUBTOTAL = THEME.fill_subtotal
FILL_KPI = THEME.fill_subtotal

ALIGN_LEFT = Alignment(horizontal="left", vertical="center")
ALIGN_RIGHT = Alignment(horizontal="right", vertical="center")

THIN_SIDE = Side(style="thin", color=THEME.border_color)
TOP_BOTTOM_BORDER = Border(top=THIN_SIDE, bottom=THIN_SIDE)
BORDER_SUBTOTAL_TOP = THEME.border_subtotal_top

NUM_FMT_INT = "#,##0;(#,##0);-"
NUM_FMT_KPI = "0.0;(0.0);-"

EBITDA_REPORTED_LABEL = "EBITDA (reported)"
EBITDA_ADJUSTED_LABEL = "EBITDA"
EBITDA_ADJ_LABEL = "EBITDA adjustments"

FIXED_ROWS_BEFORE_WC = [
    {"type": "ebitda_reported", "label": EBITDA_REPORTED_LABEL},
]

FIXED_ROWS_AFTER_EBITDA_BRIDGE = [
    {"type": "ebitda_adj", "label": EBITDA_ADJ_LABEL},
    {"type": "ebitda_adjusted", "label": EBITDA_ADJUSTED_LABEL},
]

PL_ADJ_SOURCE_VALUE = "Adjusted"

FIXED_ROWS_AFTER_WC = [
    {"type": "operating_cf", "label": "Cash flow from operating activities"},
    {"type": "capex", "label": "Capex", "NA": "FA"},
    {"type": "investing_cf", "label": "Cash flow from investing activities"},
    {"type": "fcf_before_tax", "label": "Free cash flow before tax"},
    {"type": "tax_income", "label": "Taxes on income", "pl_kind": "tax_income"},
    {"type": "tax_other", "label": "Other taxes", "pl_kind": "tax_other"},
    {"type": "free_cash_flow", "label": "Free cash flow"},
]

CF_SECTION_HEADER_TYPES = frozenset({"operating_cf", "investing_cf", "financial_result"})
CF_SUBTOTAL_TB_TYPES = frozenset({"free_cash_flow", "net_cash_flow"})

KPI_SPECS = [
    ("OCF as % of EBITDA", "ocf_pct"),
    ("FCF (bef. taxes) as % of EBITDA", "fcf_pct"),
]

PL_LABEL_FALLBACKS = {
    "tax_income": "taxes on income",
    "tax_other": "other taxes",
}


def col_letter(idx: int) -> str:
    return get_column_letter(idx)


def load_argv_config() -> dict:
    if len(sys.argv) < 2:
        return {}
    path = Path(sys.argv[1])
    if not path.is_file():
        return {}
    return json.loads(path.read_text(encoding="utf-8"))


def path_value(cfg: dict, key: str, default: str) -> str:
    paths = cfg.get("paths") or {}
    val = paths.get(key)
    return str(val) if val else default


def resolve_config() -> dict:
    raw = load_argv_config()
    paths = raw.get("paths") or {}
    return {
        "input_file": path_value(raw, "input_file", path_value(raw, "master_file", DEFAULT_INPUT)),
        "mapping_file": paths.get("mapping_file") or DEFAULT_MAPPING_BS,
        "cf_na_l3_order_file": paths.get("cf_na_l3_order_file") or DEFAULT_CF_NA_L3_ORDER,
        "pl_mapping_file": paths.get("pl_mapping_file") or DEFAULT_PL_MAPPING,
        "project_name": str(raw.get("project_name") or "Desktop Test").strip(),
        "company_name": str(raw.get("company_name") or "Group").strip(),
        "fy_end_month": int(raw.get("fy_end_month") or 12),
    }


def detect_source_col(df_: pd.DataFrame) -> str | None:
    candidates = ["Quelle", "Source", "Reported/Adjusted", "Type", "Scenario"]
    sc = next((c for c in candidates if c in df_.columns), None)
    if sc is not None:
        return sc
    for c in df_.columns:
        try:
            s = df_[c].astype(str).str.strip().str.lower()
            if (s == "reported").any():
                return c
        except Exception:
            pass
    return None


def resolve_source_section_column(df_: pd.DataFrame) -> str | None:
    markers = {REPORTED_FILTER_VALUE, "adjusted", "consolidation"}

    def marker_count(col: str) -> int:
        if col not in df_.columns:
            return 0
        series = df_[col].astype(str).str.strip().str.lower()
        return int(series.isin(markers).sum())

    best_col = None
    best_count = 0
    for col in SOURCE_COL_CANDIDATES:
        count = marker_count(col)
        if count > best_count:
            best_col = col
            best_count = count
    if best_col and best_count > 0:
        return best_col
    legacy = detect_source_col(df_)
    if legacy:
        return legacy
    for col in SOURCE_COL_CANDIDATES:
        if col in df_.columns:
            return col
    return None


def detect_na_bucket_col(df_: pd.DataFrame) -> str:
    for cand in ["NA", "L5 - GT NA View", "L5 - GT NA view", "L5 - GT NA", "L5"]:
        if cand in df_.columns:
            return cand
    for c in df_.columns:
        if str(c).strip() == "NA" or str(c).strip().startswith("L5"):
            return c
    raise ValueError("Keine NA-/L5-Bucket-Spalte gefunden.")


def normalize_cf_na_bucket(x) -> str:
    if x is None or (isinstance(x, float) and pd.isna(x)):
        return "Other"
    s = str(x).strip()
    if s in ("TWC", "OWC", "Other", "ND", "FA", "Equity"):
        return s
    if s.startswith("TWC"):
        return "TWC"
    if s.startswith("OWC"):
        return "OWC"
    if s.startswith("ND"):
        return "ND"
    if s.startswith("FA"):
        return "FA"
    return "Other"


def find_row_by_label(ws_, col_idx: int, label: str) -> int | None:
    target = label.strip().lower()
    for rr in range(1, ws_.max_row + 1):
        v = ws_.cell(rr, col_idx).value
        if isinstance(v, str) and v.strip().lower() == target:
            return rr
    return None


def find_row_by_label_contains(ws_, col_idx: int, needle: str) -> int | None:
    needle_l = needle.lower()
    for rr in range(1, ws_.max_row + 1):
        v = ws_.cell(rr, col_idx).value
        if isinstance(v, str) and needle_l in v.lower():
            return rr
    return None


def lead_is_reported_year_cols(ws_is, years: list[str]) -> dict[str, int]:
    rep_s, rep_e, _, _ = detect_is_blocks(ws_is)
    found: dict[str, int] = {}
    for cc in range(rep_s, rep_e + 1):
        v = ws_is.cell(HEADER_ROW, cc).value
        if isinstance(v, str) and v.strip() in years:
            found[v.strip()] = cc
    return found


def detect_is_blocks(ws_is):
    title_row = TITLE_ROW
    reported_start = None
    proforma_start = None
    for cc in range(1, ws_is.max_column + 1):
        v = ws_is.cell(title_row, cc).value
        if isinstance(v, str):
            vv = v.lower()
            if reported_start is None and "reported income statement" in vv:
                reported_start = cc
            if proforma_start is None and "pro forma income statement" in vv:
                proforma_start = cc
    if proforma_start is None:
        raise RuntimeError("Pro-forma income statement (Zeile 6) nicht in Lead_IS gefunden.")
    if reported_start is None:
        reported_start = 1
    if reported_start < proforma_start:
        reported_end = proforma_start - 1
        proforma_end = ws_is.max_column
    else:
        proforma_end = reported_start - 1
        reported_end = ws_is.max_column
    return reported_start, reported_end, proforma_start, proforma_end


def lead_is_proforma_year_cols(ws_is, years: list[str]) -> dict[str, int]:
    _, _, pf_s, pf_e = detect_is_blocks(ws_is)
    found: dict[str, int] = {}
    for cc in range(pf_s, pf_e + 1):
        v = ws_is.cell(HEADER_ROW, cc).value
        if isinstance(v, str) and v.strip() in years:
            found[v.strip()] = cc
    return found


def resolve_pl_l3_label(wb, kind: str) -> str:
    fallback_needle = PL_LABEL_FALLBACKS[kind]
    if SHEET_PL_RECON in wb.sheetnames:
        ws_pl = wb[SHEET_PL_RECON]
        for col in range(1, min(ws_pl.max_column or 0, 30) + 1):
            hdr = ws_pl.cell(HEADER_ROW, col).value
            if isinstance(hdr, str) and "aggregated" in hdr.lower():
                row = find_row_by_label_contains(ws_pl, col, fallback_needle)
                if row is not None:
                    val = ws_pl.cell(row, col).value
                    if isinstance(val, str) and val.strip():
                        return val.strip()
        row = find_row_by_label_contains(ws_pl, POS_COL, fallback_needle)
        if row is not None:
            val = ws_pl.cell(row, POS_COL).value
            if isinstance(val, str) and val.strip():
                return val.strip()
    if kind == "tax_income":
        return "Taxes on income"
    return "Other taxes"


def subtotal_span_for_l3(struct, idx: int, bucket: str, l3_label: str) -> tuple[int, int]:
    k = idx - 1
    while (
        k >= 0
        and struct[k]["type"] in {"detail", "detail_single"}
        and struct[k].get("NA") == bucket
        and struct[k].get("L3") == l3_label
    ):
        k -= 1
    start = DATA_START_ROW + (k + 1)
    end = DATA_START_ROW + (idx - 1)
    return start, end


def bucket_l3_level_refs(struct, bucket: str, col_l: str) -> list[str]:
    """Sum L3 subtotals / single-line L3 rows in a bucket — exclude L4 detail lines."""
    refs: list[str] = []
    for rr in struct:
        if rr.get("bucket") != bucket and rr.get("NA") != bucket:
            continue
        if rr["type"] in {"subtotal_l3", "detail_single"}:
            refs.append(f"{col_l}{rr['_excel_row']}")
    return refs


def fr_detail_row_refs(struct, col_l: str) -> list[str]:
    return [f"{col_l}{rr['_excel_row']}" for rr in struct if rr["type"] == "fr_detail"]


def build_pl_adjustment_detail_rows(df_pl: pd.DataFrame) -> list[dict]:
    """EBITDA adjustment lines from Master_PL (L6 = Adjusted), sorted by latest FY amount."""
    source_col = resolve_source_section_column(df_pl)
    if not source_col or "L3" not in df_pl.columns or "L4" not in df_pl.columns:
        return []
    src = df_pl[source_col].astype(str).str.strip().str.lower()
    mask = src.eq(PL_ADJ_SOURCE_VALUE.lower())
    sub = df_pl.loc[mask]
    if sub.empty:
        return []

    period_cols = period_columns_for_sort(df_pl, L4_SORT_BASIS)
    pairs: list[tuple[str, str]] = []
    seen: set[tuple[str, str]] = set()
    for _, row in sub.iterrows():
        l3 = str(row["L3"]).strip()
        l4 = str(row["L4"]).strip()
        key = (l3, l4)
        if key in seen:
            continue
        seen.add(key)
        pairs.append(key)

    def pair_metric(l3: str, l4: str) -> float:
        return compute_l4_sort_metric(
            df_pl, l3, l4, period_cols, source_col, reported_value=PL_ADJ_SOURCE_VALUE
        )

    pairs.sort(key=lambda p: (-abs(pair_metric(p[0], p[1])), p[1].lower()))

    rows: list[dict] = []
    for l3, l4 in pairs:
        label = l4 if l4 else l3
        row_type = "pl_adj_detail_single" if l3 == l4 else "pl_adj_detail"
        rows.append(
            {
                "type": row_type,
                "label": label,
                "L3": l3,
                "L4": l4,
            }
        )
    return rows


def pl_adj_detail_row_refs(struct, col_l: str) -> list[str]:
    return [
        f"{col_l}{rr['_excel_row']}"
        for rr in struct
        if rr["type"] in {"pl_adj_detail", "pl_adj_detail_single"}
    ]


def build_financial_result_rows(pl_map_df: pd.DataFrame, df_pl: pd.DataFrame) -> list[dict]:
    l4_order = l4_order_from_pl_l3_mapping(pl_map_df, FINANCIAL_RESULT_L3)
    source_col = resolve_source_section_column(df_pl)
    if source_col:
        rep = df_pl[source_col].astype(str).str.strip().str.lower().eq(REPORTED_FILTER_VALUE)
        master_l4s = set(
            df_pl.loc[
                rep & df_pl["L3"].astype(str).str.strip().str.lower().eq(FINANCIAL_RESULT_L3.lower()),
                "L4",
            ]
            .astype(str)
            .str.strip()
        )
        l4_order = [l4 for l4 in l4_order if l4 in master_l4s]
    rows: list[dict] = []
    for l4 in l4_order:
        rows.append(
            {
                "type": "fr_detail",
                "label": l4,
                "L3": FINANCIAL_RESULT_L3,
                "L4": l4,
            }
        )
    rows.append({"type": "financial_result", "label": "Cash flow from financing activities"})
    rows.append({"type": "net_cash_flow", "label": "Net cash flow"})
    return rows


def display_label(r: dict) -> str:
    rt = r["type"]
    if rt in {"detail", "detail_single"}:
        return f"Δ {r['label']}"
    if rt == "subtotal_l3":
        return f"Δ {r['label']}"
    return r["label"]


def _sumifs_expr(
    data_rng: str,
    *,
    src_rng: str | None,
    rep_crit: str,
    extra_criteria: list[tuple[str, str]],
) -> str:
    parts = [f"SUMIFS({data_rng}"]
    if src_rng:
        parts.append(f"{src_rng},{rep_crit}")
    for rng, crit in extra_criteria:
        parts.append(f"{rng},{crit}")
    return ",".join(parts) + ")"


def delta_sumifs(
    curr_rng: str,
    prior_rng: str,
    *,
    src_rng: str | None,
    rep_crit: str,
    extra_criteria: list[tuple[str, str]],
) -> str:
    curr = _sumifs_expr(curr_rng, src_rng=src_rng, rep_crit=rep_crit, extra_criteria=extra_criteria)
    prior = _sumifs_expr(prior_rng, src_rng=src_rng, rep_crit=rep_crit, extra_criteria=extra_criteria)
    return f"=({curr}-{prior})/1000"


def delta_sumifs_for_period(
    spec,
    period_rng: dict[str, str],
    *,
    src_rng: str | None,
    rep_crit: str,
    extra_criteria: list[tuple[str, str]],
) -> str:
    if spec.curr_col and spec.prior_col and spec.curr_col in period_rng and spec.prior_col in period_rng:
        return delta_sumifs(
            period_rng[spec.curr_col],
            period_rng[spec.prior_col],
            src_rng=src_rng,
            rep_crit=rep_crit,
            extra_criteria=extra_criteria,
        )
    if spec.kind == "fy":
        return "=0"
    curr_terms = [
        _sumifs_expr(period_rng[mc], src_rng=src_rng, rep_crit=rep_crit, extra_criteria=extra_criteria)
        for mc in spec.curr_months
        if mc in period_rng
    ]
    prior_terms = [
        _sumifs_expr(period_rng[mc], src_rng=src_rng, rep_crit=rep_crit, extra_criteria=extra_criteria)
        for mc in spec.prior_months
        if mc in period_rng
    ]
    if not curr_terms or not prior_terms:
        return "=0"
    curr_sum = "+".join(curr_terms)
    prior_sum = "+".join(prior_terms)
    return f"=({curr_sum}-{prior_sum})/1000"


def main() -> None:
    cfg = resolve_config()
    input_file = cfg["input_file"]
    mapping_file = cfg["mapping_file"]
    cf_na_l3_order_file = cfg["cf_na_l3_order_file"]
    pl_mapping_file = cfg["pl_mapping_file"]
    project_name = cfg["project_name"]
    group_name = cfg["company_name"]
    fy_end_month = cfg["fy_end_month"]

    print("cashflow —")
    print(f"  input:   {input_file}")
    print(f"  mapping: {mapping_file}")
    print(f"  na_l3:   {cf_na_l3_order_file}")
    print(f"  pl_map:  {pl_mapping_file}")

    wb = load_workbook(input_file)
    for req in (SHEET_MASTER_BS, SHEET_MASTER_PL, SHEET_PL_RECON, SHEET_LEAD_IS):
        if req not in wb.sheetnames:
            raise RuntimeError(
                f"Sheet '{req}' fehlt in {input_file}. Pipeline: recon_pl → lead_is → … → cashflow."
            )

    if SHEET_OUT in wb.sheetnames:
        del wb[SHEET_OUT]
    ws = wb.create_sheet(SHEET_OUT)

    df_bs = pd.read_excel(input_file, sheet_name=SHEET_MASTER_BS, engine="openpyxl")
    master_periods = ordered_reporting_columns_from_df(df_bs)
    display_periods = display_reporting_columns_from_df(df_bs)
    fy_cols, ytd_cols = split_fy_and_ytd(display_periods)
    if len(fy_cols) < 2 and not ytd_cols:
        raise ValueError("Mindestens zwei FY/YTD-Spalten im Master_BS für Cashflow-Deltas erforderlich.")

    DISPLAY_PERIODS, _prior_fy_unused = cashflow_display_periods(fy_cols)
    if ytd_cols:
        DISPLAY_PERIODS.append(ytd_cols[-1])
    Y_COLS = list(range(POS_COL + 1, POS_COL + 1 + len(DISPLAY_PERIODS)))

    source_col = resolve_source_section_column(df_bs)
    na_col = detect_na_bucket_col(df_bs)

    map_df = pd.read_excel(mapping_file, sheet_name=0, engine="openpyxl")
    if "L3" not in map_df.columns:
        raise ValueError("BS-Mapping muss Spalte 'L3' enthalten.")
    mapping_l3_order = l3_order_from_mapping(map_df)

    na_l3_map_df = pd.read_excel(cf_na_l3_order_file, sheet_name=0, engine="openpyxl")
    na_l3_order = load_na_l3_orders(na_l3_map_df, CF_BUCKET_ORDER)

    master_bs_ref = SHEET_MASTER_BS
    master_start = 2
    master_end = len(df_bs) + 1

    def master_bs_range(colname: str) -> str:
        loc = df_bs.columns.get_loc(colname) + 1
        letter = get_column_letter(loc)
        return f"{master_bs_ref}!${letter}${master_start}:${letter}${master_end}"

    na_rng = master_bs_range(na_col)
    l3_rng = master_bs_range("L3")
    l4_rng = master_bs_range("L4")
    src_rng = master_bs_range(source_col) if source_col else None

    month_cols = ordered_month_columns_from_df(df_bs)
    master_ytd_cols = [p for p in master_periods if str(p).upper().startswith("YTD")]
    period_rng: dict[str, str] = {}
    for p in set(DISPLAY_PERIODS) | set(master_periods) | set(month_cols) | set(master_ytd_cols):
        if p in df_bs.columns:
            period_rng[p] = master_bs_range(p)

    display_fy, prior_fy = cashflow_display_periods(fy_cols)
    delta_specs = {}
    for period in DISPLAY_PERIODS:
        prior = prior_fy[display_fy.index(period)] if period in display_fy else None
        delta_specs[period] = cashflow_delta_spec(
            period,
            prior,
            list(df_bs.columns),
            fy_end_month=fy_end_month,
        )

    pl_l3_rng = None
    pl_l4_rng = None
    pl_src_rng = None
    pl_period_rng: dict[str, str] = {}
    pl_tax_income = resolve_pl_l3_label(wb, "tax_income")
    pl_tax_other = resolve_pl_l3_label(wb, "tax_other")

    df_pl = pd.read_excel(input_file, sheet_name=SHEET_MASTER_PL, engine="openpyxl")
    pl_source = resolve_source_section_column(df_pl)
    if pl_source and "L3" in df_pl.columns:
        pl_master_ref = SHEET_MASTER_PL
        pl_end = len(df_pl) + 1

        def master_pl_range(colname: str) -> str:
            loc = df_pl.columns.get_loc(colname) + 1
            letter = get_column_letter(loc)
            return f"{pl_master_ref}!${letter}${master_start}:${letter}${pl_end}"

        pl_l3_rng = master_pl_range("L3")
        pl_l4_rng = master_pl_range("L4")
        pl_src_rng = master_pl_range(pl_source)
        for p in DISPLAY_PERIODS:
            if p in df_pl.columns:
                pl_period_rng[p] = master_pl_range(p)

    wc_rows = build_cf_na_bucket_row_structure(
        df_bs,
        mapping_l3_order,
        {"l4_sort_basis": L4_SORT_BASIS},
        source_col=source_col or na_col,
        bucket_col=na_col,
        bucket_order=CF_BUCKET_ORDER,
        bucket_total_labels=CF_BUCKET_TOTAL_LABELS,
        normalize_bucket_fn=normalize_cf_na_bucket,
        na_l3_order=na_l3_order,
    )
    wc_rows = prune_zero_value_rows(
        wc_rows,
        df_bs,
        month_cols,
        source_col=source_col or SOURCE_COL_CANDIDATES[0],
    )

    row_structure: list[dict] = []
    row_structure.extend(FIXED_ROWS_BEFORE_WC)
    row_structure.extend(build_pl_adjustment_detail_rows(df_pl))
    row_structure.extend(FIXED_ROWS_AFTER_EBITDA_BRIDGE)
    row_structure.extend(wc_rows)
    row_structure.extend(FIXED_ROWS_AFTER_WC)

    pl_map_df = pd.read_excel(pl_mapping_file, sheet_name=0, engine="openpyxl")
    row_structure.extend(build_financial_result_rows(pl_map_df, df_pl))

    for i, r in enumerate(row_structure):
        r["_idx"] = i
        r["_excel_row"] = DATA_START_ROW + i

    LAST_TABLE_ROW = DATA_START_ROW + len(row_structure) - 1

    ws_recon = wb[SHEET_PL_RECON]
    ws_is = wb[SHEET_LEAD_IS]
    recon_year_cols = find_recon_aggregated_year_cols(ws_recon, DISPLAY_PERIODS)
    is_pf_cols = lead_is_proforma_year_cols(ws_is, DISPLAY_PERIODS)
    is_rep_cols = lead_is_reported_year_cols(ws_is, DISPLAY_PERIODS)

    if not is_pf_cols:
        raise RuntimeError("Pro-forma Lead_IS enthält keine passenden Perioden-Spalten.")

    row_ebitda_reported = None
    for label_col in (LAYOUT_PL.pos_col,):
        row_ebitda_reported = find_row_by_label(ws_recon, label_col, "EBITDA")
        if row_ebitda_reported:
            break
    if row_ebitda_reported is None:
        row_ebitda_reported = find_row_by_label_contains(ws_recon, LAYOUT_PL.pos_col, "ebitda")
    if row_ebitda_reported is None:
        raise RuntimeError("EBITDA-Zeile in PL_Reconciliation nicht gefunden.")

    _, _, pf_label_col, _ = detect_is_blocks(ws_is)
    row_ebitda_adj = find_row_by_label(ws_is, pf_label_col, EBITDA_ADJUSTED_LABEL)
    if row_ebitda_adj is None:
        row_ebitda_adj = find_row_by_label_contains(ws_is, pf_label_col, "ebitda")
    if row_ebitda_adj is None:
        raise RuntimeError(f"EBITDA-Zeile in Lead_IS Pro forma nicht gefunden.")

    row_ebitda_rep_is = row_ebitda_adj
    rep_s, rep_e, _, _ = detect_is_blocks(ws_is)
    for label_col in (LAYOUT_PL.pos_col, rep_s):
        if label_col > rep_e:
            continue
        found = find_row_by_label(ws_is, label_col, EBITDA_ADJUSTED_LABEL) or find_row_by_label(
            ws_is, label_col, "EBITDA"
        )
        if found:
            row_ebitda_rep_is = found
            break

    row_by_type: dict[str, int] = {}
    bucket_total_row: dict[str, int] = {}
    for r in row_structure:
        row_by_type[r["type"]] = r["_excel_row"]
        if r["type"] == "total_na":
            bucket_total_row[r.get("bucket", "")] = r["_excel_row"]

    # Titles / headers
    pt = ws.cell(PROJECT_TITLE_ROW, POS_COL, f"Project {project_name}")
    pt.font = FONT_PROJECT
    pt.alignment = ALIGN_LEFT
    ws.row_dimensions[PROJECT_TITLE_ROW].height = 36

    st = ws.cell(SUBTITLE_ROW, POS_COL, f"{group_name} Cash flow")
    st.font = FONT_SUBTITLE
    st.alignment = ALIGN_LEFT

    tt = ws.cell(
        TITLE_ROW,
        POS_COL,
        f"{group_name} | Cash flow {DISPLAY_PERIODS[0]} - {DISPLAY_PERIODS[-1]}",
    )
    tt.font = FONT_TITLE
    tt.alignment = ALIGN_LEFT

    for hdr_col, hdr_txt in enumerate(["Reported", "", "NA", "L3", "L4"], start=MAP_START_COL):
        ws.cell(HEADER_ROW, hdr_col, hdr_txt).font = FONT_HEADER

    for c in range(MAP_START_COL, MAP_END_COL + 1):
        ws.cell(HEADER_ROW, c).alignment = ALIGN_LEFT
        ws.cell(HEADER_ROW, c).fill = FILL_HEADER
        ws.cell(HEADER_ROW_7, c).fill = FILL_HEADER

    for cc in [POS_COL] + Y_COLS:
        ws.cell(HEADER_ROW_7, cc).fill = FILL_HEADER

    ws.cell(HEADER_ROW, POS_COL, "kEUR").font = FONT_HEADER
    ws.cell(HEADER_ROW, POS_COL).alignment = ALIGN_LEFT
    ws.cell(HEADER_ROW, POS_COL).fill = FILL_HEADER

    for idx, period in enumerate(DISPLAY_PERIODS):
        cc = Y_COLS[idx]
        h = ws.cell(HEADER_ROW, cc, period)
        h.font = FONT_HEADER
        h.alignment = ALIGN_RIGHT
        h.fill = FILL_HEADER

    for c in range(MAP_START_COL, MAP_END_COL + 1):
        ws.column_dimensions[col_letter(c)].width = MAP_COL_WIDTH
    ws.column_dimensions[col_letter(POS_COL)].width = POS_COL_WIDTH
    for c in Y_COLS:
        ws.column_dimensions[col_letter(c)].width = VALUE_COL_WIDTH

    map_rep_col = col_letter(MAP_START_COL)
    map_na_col = col_letter(MAP_START_COL + 2)
    map_l3_col = col_letter(MAP_START_COL + 3)
    map_l4_col = col_letter(MAP_START_COL + 4)

    ws.sheet_properties.outlinePr.summaryBelow = True
    ws.sheet_properties.outlinePr.summaryRight = True

    highlight_types = CF_SECTION_HEADER_TYPES | CF_SUBTOTAL_TB_TYPES
    bold_types = highlight_types | {
        "ebitda_reported",
        "ebitda_adjusted",
        "fcf_before_tax",
        "total_na",
        "financial_result",
    }
    table_right_col = Y_COLS[-1]

    for i, r in enumerate(row_structure):
        excel_row = r["_excel_row"]
        ws.row_dimensions[excel_row].height = ROW_HEIGHT
        rt = r["type"]

        na_val = r.get("NA", "")
        l3_val = r.get("L3", "")
        l4_val = r.get("L4", "")
        if rt == "tax_income":
            l3_val = pl_tax_income
            l4_val = pl_tax_income
        elif rt == "tax_other":
            l3_val = pl_tax_other
            l4_val = pl_tax_other
        elif rt == "fr_detail":
            l3_val = r.get("L3", FINANCIAL_RESULT_L3)
            l4_val = r.get("L4", r["label"])
        elif rt in {"pl_adj_detail", "pl_adj_detail_single"}:
            l3_val = r.get("L3", "")
            l4_val = r.get("L4", r["label"])
        elif rt == "capex":
            na_val = "FA"

        if rt not in {"ebitda_reported", "ebitda_adjusted", "ebitda_adj"}:
            map_src = PL_ADJ_SOURCE_VALUE if rt in {"pl_adj_detail", "pl_adj_detail_single"} else "Reported"
            ws.cell(excel_row, MAP_START_COL, map_src).font = FONT_MAPPING
            ws.cell(excel_row, MAP_START_COL + 1, "").font = FONT_MAPPING
            ws.cell(excel_row, MAP_START_COL + 2, na_val).font = FONT_MAPPING
            ws.cell(excel_row, MAP_START_COL + 3, l3_val).font = FONT_MAPPING
            ws.cell(excel_row, MAP_START_COL + 4, l4_val).font = FONT_MAPPING
            for c in range(MAP_START_COL, MAP_END_COL + 1):
                ws.cell(excel_row, c).alignment = ALIGN_LEFT

        pc = ws.cell(
            excel_row,
            POS_COL,
            display_label(r)
            if rt in {"detail", "detail_single", "subtotal_l3"}
            else r["label"],
        )
        pc.alignment = ALIGN_LEFT
        is_bold = rt in bold_types
        is_highlight = rt in highlight_types
        pc.font = FONT_BASE_BOLD if is_bold else FONT_BASE

        if rt == "detail" or rt == "pl_adj_detail":
            pc.alignment = Alignment(horizontal="left", vertical="center", indent=1)
            ws.row_dimensions[excel_row].outlineLevel = 1
            ws.row_dimensions[excel_row].collapsed = True
        elif rt in {"detail_single", "subtotal_l3", "pl_adj_detail_single"}:
            ws.row_dimensions[excel_row].outlineLevel = 0
        else:
            ws.row_dimensions[excel_row].outlineLevel = 0

        rep_crit = f"${map_rep_col}${excel_row}"
        na_crit = f"${map_na_col}${excel_row}"
        l3_crit = f"${map_l3_col}${excel_row}"
        l4_crit = f"${map_l4_col}${excel_row}"

        for p_idx, period in enumerate(DISPLAY_PERIODS):
            spec = delta_specs[period]
            cell = ws.cell(excel_row, Y_COLS[p_idx])
            cell.alignment = ALIGN_RIGHT
            cell.number_format = NUM_FMT_INT
            cell.font = FONT_BASE_BOLD if is_bold else FONT_BASE
            col_l = col_letter(Y_COLS[p_idx])

            if rt == "ebitda_reported":
                if period in recon_year_cols:
                    rc = recon_year_cols[period]
                    cell.value = f"='{SHEET_PL_RECON}'!{col_letter(rc)}{row_ebitda_reported}"
                elif period in is_rep_cols:
                    ic = is_rep_cols[period]
                    cell.value = f"='{SHEET_LEAD_IS}'!{col_letter(ic)}{row_ebitda_rep_is}"
                else:
                    cell.value = "n/a"
            elif rt == "ebitda_adjusted":
                cell.value = (
                    f"={col_l}{row_by_type['ebitda_reported']}"
                    f"+{col_l}{row_by_type['ebitda_adj']}"
                )
            elif rt == "ebitda_adj":
                refs = pl_adj_detail_row_refs(row_structure, col_l)
                cell.value = f"=SUM({','.join(refs)})" if refs else 0
            elif rt in {"pl_adj_detail", "pl_adj_detail_single"}:
                if pl_period_rng.get(period) and pl_l3_rng and pl_l4_rng and pl_src_rng:
                    cell.value = (
                        f"=SUMIFS({pl_period_rng[period]},{pl_src_rng},{rep_crit},"
                        f"{pl_l3_rng},{l3_crit},{pl_l4_rng},{l4_crit})/1000"
                    )
                else:
                    cell.value = 0
            elif rt in {"detail", "detail_single"}:
                cell.value = delta_sumifs_for_period(
                    spec,
                    period_rng,
                    src_rng=src_rng,
                    rep_crit=rep_crit,
                    extra_criteria=[
                        (na_rng, na_crit),
                        (l3_rng, l3_crit),
                        (l4_rng, l4_crit),
                    ],
                )
            elif rt == "subtotal_l3":
                start, end = subtotal_span_for_l3(row_structure, i, r["NA"], r["label"])
                cell.value = f"=SUM({col_l}{start}:{col_l}{end})" if end >= start else 0
            elif rt == "total_na":
                refs = bucket_l3_level_refs(row_structure, r["bucket"], col_l)
                cell.value = f"=SUM({','.join(refs)})" if refs else 0
            elif rt == "capex":
                cell.value = delta_sumifs_for_period(
                    spec,
                    period_rng,
                    src_rng=src_rng,
                    rep_crit=rep_crit,
                    extra_criteria=[(na_rng, na_crit)],
                )
            elif rt == "tax_income" or rt == "tax_other":
                if pl_period_rng.get(period) and pl_l3_rng and pl_src_rng:
                    cell.value = (
                        f"=SUMIFS({pl_period_rng[period]},{pl_src_rng},{rep_crit},"
                        f"{pl_l3_rng},{l3_crit})/1000"
                    )
                else:
                    cell.value = 0
            elif rt == "operating_cf":
                parts = [f"{col_l}{row_by_type['ebitda_adjusted']}"]
                for b in CF_BUCKET_ORDER:
                    br = bucket_total_row.get(b)
                    if br:
                        parts.append(f"{col_l}{br}")
                cell.value = f"=SUM({','.join(parts)})"
            elif rt == "investing_cf":
                cell.value = f"={col_l}{row_by_type['capex']}"
            elif rt == "fcf_before_tax":
                cell.value = (
                    f"={col_l}{row_by_type['operating_cf']}"
                    f"+{col_l}{row_by_type['investing_cf']}"
                )
            elif rt == "free_cash_flow":
                cell.value = (
                    f"=SUM({col_l}{row_by_type['fcf_before_tax']},"
                    f"{col_l}{row_by_type['tax_income']},"
                    f"{col_l}{row_by_type['tax_other']})"
                )
            elif rt == "fr_detail":
                if pl_period_rng.get(period) and pl_l3_rng and pl_l4_rng and pl_src_rng:
                    cell.value = (
                        f"=SUMIFS({pl_period_rng[period]},{pl_src_rng},{rep_crit},"
                        f"{pl_l3_rng},{l3_crit},{pl_l4_rng},{l4_crit})/1000"
                    )
                else:
                    cell.value = 0
            elif rt == "financial_result":
                refs = fr_detail_row_refs(row_structure, col_l)
                cell.value = f"=SUM({','.join(refs)})" if refs else 0
            elif rt == "net_cash_flow":
                cell.value = (
                    f"={col_l}{row_by_type['free_cash_flow']}"
                    f"+{col_l}{row_by_type['financial_result']}"
                )

        if rt in CF_SUBTOTAL_TB_TYPES:
            ws.cell(excel_row, POS_COL).fill = FILL_SUBTOTAL
            for cc in Y_COLS:
                ws.cell(excel_row, cc).fill = FILL_SUBTOTAL
            border = TOP_BOTTOM_BORDER
            pc.border = border
            apply_databook_row_border_band(
                ws, excel_row, _COL_LAYOUT, border, last_used_col=table_right_col
            )
        elif rt in CF_SECTION_HEADER_TYPES:
            ws.cell(excel_row, POS_COL).fill = FILL_SUBTOTAL
            for cc in Y_COLS:
                ws.cell(excel_row, cc).fill = FILL_SUBTOTAL
            border = BORDER_SUBTOTAL_TOP
            pc.border = border
            apply_databook_row_border_band(
                ws, excel_row, _COL_LAYOUT, border, last_used_col=table_right_col
            )
        else:
            for cc in Y_COLS:
                ws.cell(excel_row, cc).fill = FILL_WHITE

    # KPI block
    KPI_TITLE_ROW = LAST_TABLE_ROW + 1
    write_kpi_section_title_row(
        ws,
        KPI_TITLE_ROW,
        title_cols=Y_COLS,
        label_col=POS_COL,
        title_text="KPIs",
        fill_end_col=table_right_col,
    )
    KPI_START = LAST_TABLE_ROW + 2
    kpi_rows: dict[str, int] = {}
    for i_kpi, (label, key) in enumerate(KPI_SPECS):
        rr = KPI_START + i_kpi
        kpi_rows[key] = rr
        ws.row_dimensions[rr].height = ROW_HEIGHT
        ws.cell(rr, POS_COL, label).font = FONT_KPI
        ws.cell(rr, POS_COL, label).alignment = ALIGN_LEFT
        for p_idx, period in enumerate(DISPLAY_PERIODS):
            cc = Y_COLS[p_idx]
            col_l = col_letter(cc)
            cell = ws.cell(rr, cc)
            cell.number_format = NUM_FMT_KPI
            cell.alignment = ALIGN_RIGHT
            cell.font = FONT_KPI
            ebitda = f"{col_l}{row_by_type['ebitda_adjusted']}"
            if key == "ocf_pct":
                cell.value = f'=IFERROR({col_l}{row_by_type["operating_cf"]}/{ebitda}*100,"n/a")'
            elif key == "fcf_pct":
                cell.value = (
                    f'=IFERROR({col_l}{row_by_type["fcf_before_tax"]}/{ebitda}*100,"n/a")'
                )

    current_row = KPI_START + len(KPI_SPECS) - 1
    KPI_ROW_SET = set(range(KPI_START, KPI_START + len(KPI_SPECS)))
    KPI_TITLE_ROW_SET = {KPI_TITLE_ROW}

    ws.sheet_view.showOutlineSymbols = True

    LAST_VALUE_COL = Y_COLS[-1]
    FILL_END_ROW = max(LAST_TABLE_ROW, current_row) + 100
    FILL_END_COL = LAST_VALUE_COL + 40
    layout = LAYOUT_NA

    for rr in range(1, FILL_END_ROW + 1):
        if rr != PROJECT_TITLE_ROW:
            ws.row_dimensions[rr].height = ROW_HEIGHT

    highlight_row_nums = {r["_excel_row"] for r in row_structure if r["type"] in highlight_types}

    paint_grey_white_canvas(ws, layout, last_row=FILL_END_ROW, last_col=LAST_VALUE_COL)

    for rr in range(1, FILL_END_ROW + 1):
        for cc in range(POS_COL, FILL_END_COL + 1):
            ws.cell(rr, cc).border = Border()
            if rr in KPI_ROW_SET and cc in ([POS_COL] + Y_COLS):
                continue
            if rr in (HEADER_ROW_7, HEADER_ROW) and cc in ([POS_COL] + Y_COLS):
                continue
            if rr in highlight_row_nums and cc <= LAST_VALUE_COL:
                continue
            if ws.cell(rr, cc).fill.fill_type is None:
                ws.cell(rr, cc).fill = FILL_WHITE

    paint_header_band(
        ws,
        layout,
        header_rows=[HEADER_ROW_7, HEADER_ROW],
        period_cols=Y_COLS,
    )

    for rr in KPI_ROW_SET | KPI_TITLE_ROW_SET:
        ws.cell(rr, POS_COL).fill = FILL_KPI
        for cc in Y_COLS:
            ws.cell(rr, cc).fill = FILL_KPI

    for r in row_structure:
        if r["type"] not in highlight_types:
            continue
        excel_row = r["_excel_row"]
        ws.cell(excel_row, POS_COL).fill = FILL_SUBTOTAL
        for c in Y_COLS:
            ws.cell(excel_row, c).fill = FILL_SUBTOTAL

    hide_helper_column_group(ws, layout)

    wb.calculation.fullCalcOnLoad = True
    wb.save(input_file)
    print(f"Saved sheet '{SHEET_OUT}' to {input_file}")
    print(f"Perioden: {DISPLAY_PERIODS}")


if __name__ == "__main__":
    main()
