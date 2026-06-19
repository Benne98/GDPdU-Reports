import os
import sys
import json
import re
from copy import copy

import pandas as pd
import numpy as np

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
PARENT_DIR = os.path.dirname(BASE_DIR)

if PARENT_DIR not in sys.path:
    sys.path.insert(0, PARENT_DIR)

from funktionssammlung import (  # noqa: E402
    apply_filters,
    safe_month_day_ts,
    fiscal_year_bounds,
    fdd_as_of_end,
    fdd_current_fy_end_year,
    fdd_as_of_is_fy_end,
    fdd_ytd_bounds,
    fdd_ltm_bounds,
    recognized_value,
    fy_label,
    ytd_label,
    ltm_label,
    to_kEUR,
    ensure_output_writable,
    build_output_file_path,
    ensure_source_sheet_in_output,
    refresh_source_sheet_from_df,
    write_export_df_to_sheet,
    get_next_sheet_name_from_wb,
    build_sumifs_formula_body,
    build_chunked_row_sum_formula,
    source_range_ref,
    source_column_range_ref,
    add_accrual_col_in_ws,
    _get_sheet_header_map as _get_source_header_map,
)

from openpyxl import load_workbook  # noqa: E402
from openpyxl.formatting.rule import CellIsRule  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

_SCRIPTS_DIR = os.path.join(BASE_DIR, "scripts")
if _SCRIPTS_DIR not in sys.path:
    sys.path.insert(0, _SCRIPTS_DIR)
from gst_excel_theme import THEME  # noqa: E402


DROP_INDEX_PRINT_LIMIT = 200


def _set_cell_formula(ws, row: int, col: int, formula: str) -> None:
    """Write Excel formula explicitly (openpyxl + Excel repair-safe)."""
    cell = ws.cell(row=row, column=col)
    if formula and str(formula).startswith("="):
        cell.value = formula
        cell.data_type = "f"
    else:
        cell.value = formula


def _report_data_bottom_row(ws, col_row_type: int, data_start_row: int) -> int:
    last = data_start_row
    for r in range(data_start_row, ws.max_row + 1):
        rt = ws.cell(row=r, column=col_row_type).value
        if rt is not None and str(rt).strip() != "":
            last = r
    return last


def _count_sheet_formulas(ws, max_row: int = 2000, max_col: int = 40) -> int:
    n = 0
    for r in range(1, min(ws.max_row or 1, max_row) + 1):
        for c in range(1, min(ws.max_column or 1, max_col) + 1):
            v = ws.cell(row=r, column=c).value
            if isinstance(v, str) and v.startswith("="):
                n += 1
    return n


def _warn_formula_integrity(ws, label: str) -> None:
    """Lightweight post-formula sanity check (logs only)."""
    n = _count_sheet_formulas(ws)
    if n == 0:
        print(f"[WARN] {label}: no formulas found on report sheet.")
        return
    for r in range(1, min(ws.max_row or 1, 2000) + 1):
        for c in range(1, min(ws.max_column or 1, 40) + 1):
            v = ws.cell(row=r, column=c).value
            if isinstance(v, str) and v.startswith("=") and "#REF!" in v.upper():
                print(f"[WARN] {label}: broken formula (#REF!) at {get_column_letter(c)}{r}")
                return


def get_top_excel_layout(cfg: dict) -> dict:
    """Layout wie GST/PVM: A–C eingeklappt, TOP-Schlüssel in Spalte B, sichtbare Tabelle ab Spalte D."""
    extra = 2 if cfg.get("formula_mode", True) else 0
    col_offset = int(cfg.get("excel_formatting", {}).get("col_offset", 3))
    company_row = 3 + extra
    header_row = 4 + extra
    return {
        "extra_period_rows": extra,
        "col_offset": col_offset,
        "title_row": 1,
        "subtitle_row": 2,
        "period_start_row": 3 if extra == 2 else None,
        "period_end_row": 4 if extra == 2 else None,
        "company_row": company_row,
        "header_row": header_row,
        "data_start_row": header_row + 1,
        "insert_rows": header_row - 1,
        "key_col": 2,
        "label_col": col_offset + 1,
        "helper_right_col": col_offset,
    }


def strip_formula_value_columns(df: pd.DataFrame, meta: dict) -> pd.DataFrame:
    """Leave period and delta columns empty for Excel formulas."""
    out = df.copy()
    pct_col = meta["pct_col"]
    cum_col = meta["cum_col"]
    skip = {"kEUR", "Rank", "ABC", pct_col, cum_col, "row_type"}
    for c in out.columns:
        if c in skip:
            continue
        if isinstance(c, str) and (
            c.startswith("FY") or c.startswith("YTD") or c.startswith("LTM") or c.startswith("Δ ")
        ):
            out[c] = ""
    return out


def _normalize_bucket_mode(mode: str) -> str:
    """Map wizard values (custom_threshold, custom_numbers, numbers) to script values."""
    m = str(mode or "threshold").strip().lower()
    if m in {"number", "numbers", "custom_numbers"}:
        return "number"
    return "threshold"


def _normalize_filters_cfg(raw) -> dict:
    if isinstance(raw, list):
        return {"enabled": len(raw) > 0, "rules": raw}
    if isinstance(raw, dict):
        out = dict(raw)
        out.setdefault("enabled", False)
        out.setdefault("rules", [])
        return out
    return {"enabled": False, "rules": []}


def report_drops(before_idx: pd.Index, after_idx: pd.Index, label: str):
    dropped = before_idx.difference(after_idx)
    n = len(dropped)
    if n == 0:
        return
    head = list(dropped[:DROP_INDEX_PRINT_LIMIT])
    more = "" if n <= DROP_INDEX_PRINT_LIMIT else f" ... (+{n - DROP_INDEX_PRINT_LIMIT} weitere)"
    print(f"[DROP] {label}: {n} Zeilen entfernt. Indexes (erste {min(n, DROP_INDEX_PRINT_LIMIT)}): {head}{more}")


def _parse_ltm_month(ltm: str) -> tuple[int, int] | None:
    """Parse 'YYYY-MM' from Date Settings into (year, month)."""
    parts = str(ltm or "").strip().split("-")
    try:
        y, m = int(parts[0]), int(parts[1])
        if 1 <= m <= 12:
            return y, m
    except (ValueError, IndexError):
        pass
    return None


def normalize_config(cfg: dict) -> dict:
    out = dict(cfg)
    ltm = str(out.get("ltm_month") or "").strip()
    if ltm:
        parsed = _parse_ltm_month(ltm)
        if parsed:
            out["current_year"], out["current_month"] = parsed
    required = [
        "current_year",
        "current_month",
        "calc_mode",
        "value_col",
        "top_col",
        "invoice_col",
        "file_path",
        "output_file_path",
        "sheet_name",
        "case_id",
    ]
    for k in required:
        if k not in out or out[k] in (None, ""):
            raise ValueError(f"CONFIG['{k}'] muss gesetzt sein.")

    out["current_year"] = int(out["current_year"])
    out["current_month"] = int(out["current_month"])
    if not 1 <= out["current_month"] <= 12:
        raise ValueError("CONFIG['current_month'] muss 1..12 sein.")

    out["calc_mode"] = str(out["calc_mode"]).strip().lower()
    if out["calc_mode"] not in {"invoice", "accrual"}:
        raise ValueError("CONFIG['calc_mode'] muss 'invoice' oder 'accrual' sein.")

    if out["calc_mode"] == "accrual":
        for k in ("start_col", "end_col"):
            if not str(out.get(k) or "").strip():
                raise ValueError(f"CONFIG['{k}'] muss bei calc_mode='accrual' gesetzt sein.")

    out.setdefault("invoice_mapping_mode", "date" if out["calc_mode"] == "invoice" else "accrual")

    out.setdefault("subtitle_suffix", "")
    out.setdefault("company", "")
    out.setdefault("title", "")
    out.setdefault("table", "")
    out.setdefault("base_sheet_name", "TOP")
    out.setdefault("run_id", "")
    out.setdefault("apply_fx", False)
    out.setdefault("fx_col", None)
    out.setdefault("formula_mode", True)
    cy = int(out["current_year"])
    out.setdefault("first_fy", cy - 3)
    out["first_fy"] = int(out["first_fy"])
    if out["first_fy"] < 1900:
        raise ValueError("CONFIG['first_fy'] muss ein plausibles Geschäftsjahr sein.")

    # FY end defaults
    out["fiscal_year_end_month"] = int(out.get("fiscal_year_end_month", 12))
    out["fiscal_year_end_day"] = int(out.get("fiscal_year_end_day", 31))
    if not 1 <= out["fiscal_year_end_month"] <= 12:
        raise ValueError("CONFIG['fiscal_year_end_month'] muss 1..12 sein.")
    if not 1 <= out["fiscal_year_end_day"] <= 31:
        raise ValueError("CONFIG['fiscal_year_end_day'] muss 1..31 sein.")

    # ABC thresholds
    abc = out.get("abc_thresholds", (0.2, 0.4))
    if abc is None or len(abc) != 2:
        raise ValueError("CONFIG['abc_thresholds'] muss Tuple Länge 2 sein, z.B. (0.2, 0.4).")
    a_th, b_th = float(abc[0]), float(abc[1])
    if not (0 < a_th < b_th < 1):
        raise ValueError("ABC thresholds müssen gelten: 0 < A < B < 1 (z.B. 0.2, 0.4).")
    out["abc_thresholds"] = (a_th, b_th)

    # top bucket config
    tb = dict(out.get("top_bucket", {}) or {})
    tb.setdefault("enabled", True)
    tb.setdefault("bucket_mode", "threshold")  # threshold | number
    tb.setdefault("thresholds", (0.2, 0.5, 0.8))
    tb.setdefault("numbers", ())
    tb.setdefault("create_other_bucket", True)
    tb.setdefault("other_bucket_label", "Other")
    tb["bucket_mode"] = _normalize_bucket_mode(tb.get("bucket_mode"))
    out["top_bucket"] = tb

    if tb["bucket_mode"] not in {"threshold", "number"}:
        raise ValueError("CONFIG['top_bucket']['bucket_mode'] muss 'threshold' oder 'number' sein.")

    out["filters"] = _normalize_filters_cfg(out.get("filters"))

    if bool(out.get("apply_fx", False)):
        if not out.get("fx_col"):
            raise ValueError("apply_fx=True, aber CONFIG['fx_col'] fehlt.")

    return out


def preprocess_top_input(df: pd.DataFrame, cfg: dict) -> pd.DataFrame:
    cfg = normalize_config(cfg)
    d = df.copy()
    d.columns = d.columns.astype(str).str.replace("\u00a0", " ", regex=False).str.strip()
    d = apply_filters(d, cfg)

    date_cols = [cfg["invoice_col"]]
    if cfg["calc_mode"] == "accrual":
        for c in [cfg.get("start_col"), cfg.get("end_col")]:
            if not c:
                continue
            if c not in d.columns:
                raise ValueError(f"calc_mode=accrual braucht Spalte '{c}' im Input.")
            date_cols.append(c)

    for c in date_cols:
        if c not in d.columns:
            raise ValueError(f"Spalte '{c}' fehlt im Input.")
        d[c] = pd.to_datetime(d[c], errors="coerce", dayfirst=True)

    if cfg["value_col"] not in d.columns:
        raise ValueError(f"value_col '{cfg['value_col']}' fehlt im Input.")
    d["value"] = pd.to_numeric(d[cfg["value_col"]], errors="coerce")

    # FX
    if cfg.get("apply_fx", False):
        fx_col = cfg.get("fx_col")
        if not fx_col or fx_col not in d.columns:
            raise ValueError("apply_fx=True, aber fx_col fehlt im Input.")
        d["fx_rate"] = pd.to_numeric(d[fx_col], errors="coerce")
        valid_fx = d["fx_rate"].notna() & (d["fx_rate"] != 0)
        d.loc[valid_fx, "value"] = d.loc[valid_fx, "value"] * d.loc[valid_fx, "fx_rate"]

    # top key
    if cfg["top_col"] not in d.columns:
        raise ValueError(f"top_col '{cfg['top_col']}' fehlt im Input.")
    d["top_key"] = d[cfg["top_col"]].astype(str).str.strip()

    # drop NA
    before = d.index
    bad = d[d[cfg["invoice_col"]].isna() | d["value"].isna() | d["top_key"].isna()]
    if len(bad) > 0:
        print("[DROP] Invoice Date NaT:", bad[cfg["invoice_col"]].isna().sum())
        print("[DROP] Value NaN:", bad["value"].isna().sum())
        print("[DROP] top_key missing:", bad["top_key"].isna().sum())

    d = d.dropna(subset=[cfg["invoice_col"], "value", "top_key"])
    report_drops(before, d.index, "Drop NA invoice/value/top_key")

    if cfg["calc_mode"] == "accrual":
        before = d.index
        d = d.dropna(subset=[cfg.get("start_col"), cfg.get("end_col")])
        report_drops(before, d.index, "Drop NA contract start/end")

        before = d.index
        d = d[d[cfg.get("end_col")] >= d[cfg.get("start_col")]].copy()
        report_drops(before, d.index, "Drop end < start")

    before = d.index
    d = d[d["top_key"].notna() & (d["top_key"] != "")].copy()
    report_drops(before, d.index, "Drop empty top_key")

    return d


def build_period_definitions(cfg: dict) -> dict:
    cfg = normalize_config(cfg)

    CY = int(cfg["current_year"])
    m = int(cfg["current_month"])
    fy_end_m = int(cfg["fiscal_year_end_month"])
    fy_end_d = int(cfg["fiscal_year_end_day"])

    as_of_end = fdd_as_of_end(CY, m)
    current_fy_end_year = fdd_current_fy_end_year(as_of_end, fy_end_m, fy_end_d)
    as_of_is_fy_end = fdd_as_of_is_fy_end(as_of_end, fy_end_m, fy_end_d)

    first_fy = int(cfg["first_fy"])
    latest_closed_fy_end_year = current_fy_end_year if as_of_is_fy_end else current_fy_end_year - 1
    if first_fy > latest_closed_fy_end_year:
        raise ValueError(
            f"CONFIG['first_fy']={first_fy} liegt nach dem letzten abgeschlossenen FY "
            f"({latest_closed_fy_end_year})."
        )

    periods: dict[str, tuple[pd.Timestamp, pd.Timestamp]] = {}
    FY_end_years = list(range(first_fy, latest_closed_fy_end_year + 1))
    for y in FY_end_years:
        s, e = fiscal_year_bounds(y, fy_end_m, fy_end_d)
        periods[fy_label(y)] = (s, e)

    # YTD / LTM bounds via the shared canonical helpers (consistent with GST/PVM).
    end_cy = as_of_end
    end_py = fdd_as_of_end(end_cy.year - 1, end_cy.month)

    start_ytd_cy, _ = fdd_ytd_bounds(end_cy, fy_end_m, fy_end_d)
    start_ytd_py, _ = fdd_ytd_bounds(end_py, fy_end_m, fy_end_d)
    start_ltm_cy, _ = fdd_ltm_bounds(end_cy)
    start_ltm_py, _ = fdd_ltm_bounds(end_py)

    FYs = [fy_label(y) for y in FY_end_years]

    if as_of_is_fy_end:
        extra: list[str] = []
        fy_delta = f"Δ {fy_label(current_fy_end_year)} - {fy_label(current_fy_end_year - 1)}"
        ytd_delta = None
    else:
        periods.update(
            {
                ytd_label(end_py.year): (start_ytd_py, end_py),
                ytd_label(end_cy.year): (start_ytd_cy, end_cy),
                ltm_label(end_py.year): (start_ltm_py, end_py),
                ltm_label(end_cy.year): (start_ltm_cy, end_cy),
            }
        )
        extra = [
            ytd_label(end_py.year),
            ytd_label(end_cy.year),
            ltm_label(end_py.year),
            ltm_label(end_cy.year),
        ]
        fy_delta = f"Δ {fy_label(current_fy_end_year - 1)} - {fy_label(current_fy_end_year - 2)}"
        ytd_delta = f"Δ {ytd_label(end_cy.year)} - {ytd_label(end_py.year)}"

    return {
        "FYs": FYs,
        "extra": extra,
        "period_map": periods,
        "fy_delta_label": fy_delta,
        "ytd_delta_label": ytd_delta,
        "as_of_end": as_of_end,
        "as_of_is_fy_end": as_of_is_fy_end,
        "current_fy_end_year": current_fy_end_year,
        "latest_closed_fy_end_year": latest_closed_fy_end_year,
        "first_fy": first_fy,
    }


def aggregate_period(d: pd.DataFrame, start: pd.Timestamp, end: pd.Timestamp, cfg: dict) -> pd.DataFrame:
    if cfg["calc_mode"] == "invoice":
        mask = (d[cfg["invoice_col"]] >= start) & (d[cfg["invoice_col"]] <= end)
        tmp = d.loc[mask, ["top_key", "value"]].copy()
        out = tmp.groupby("top_key", as_index=False)["value"].sum().rename(columns={"value": "amount"})
        return out.fillna(0.0)

    rv = recognized_value(d, start, end, cfg)
    tmp = pd.DataFrame({"top_key": d["top_key"], "amount": rv})
    out = tmp.groupby("top_key", as_index=False)["amount"].sum()
    return out.fillna(0.0)


def build_wide_table(d: pd.DataFrame, period_defs: dict, cfg: dict) -> pd.DataFrame:
    period_map = period_defs["period_map"]
    period_cols = period_defs["FYs"] + period_defs["extra"]

    wide = None
    for col_name in period_cols:
        start, end = period_map[col_name]
        agg = aggregate_period(d, start, end, cfg).rename(columns={"amount": col_name})
        if wide is None:
            wide = agg
        else:
            wide = wide.merge(agg, on="top_key", how="outer")

    assert wide is not None
    wide = wide.fillna(0.0)
    wide = wide.rename(columns={"top_key": "kEUR"})
    return wide


def drop_rows_all_zero_periods(wide: pd.DataFrame) -> pd.DataFrame:
    out = wide.copy()
    period_cols = [c for c in out.columns if isinstance(c, str) and (c.startswith("FY") or c.startswith("YTD") or c.startswith("LTM"))]
    if not period_cols:
        return out
    vals = out[period_cols].apply(pd.to_numeric, errors="coerce").fillna(0.0)
    keep_mask = vals.abs().sum(axis=1) != 0
    return out.loc[keep_mask].reset_index(drop=True)


def drop_rows_all_zero_relevant_values(df: pd.DataFrame) -> pd.DataFrame:
    out = df.copy()
    relevant_cols = []
    for c in out.columns:
        if isinstance(c, str) and (c.startswith("FY") or c.startswith("YTD") or c.startswith("LTM") or c.startswith("Δ ")):
            relevant_cols.append(c)
    if not relevant_cols:
        return out
    vals = out[relevant_cols].apply(pd.to_numeric, errors="coerce").fillna(0.0)
    keep_mask = vals.abs().sum(axis=1) != 0
    return out.loc[keep_mask].reset_index(drop=True)


def add_metrics(wide: pd.DataFrame, period_defs: dict, cfg: dict):
    out = wide.copy()

    rank_col = period_defs["FYs"][-1]
    if cfg.get("rank_based_on") and str(cfg["rank_based_on"]).upper() == "FY_LAST":
        rank_col = str(cfg.get("rank_based_on_col", rank_col))

    if rank_col not in out.columns:
        fy_cols = sorted([c for c in out.columns if isinstance(c, str) and c.startswith("FY")])
        if not fy_cols:
            raise ValueError("Keine FY-Spalten gefunden – kann Rank nicht berechnen.")
        rank_col = fy_cols[-1]

    out[rank_col] = pd.to_numeric(out[rank_col], errors="coerce").fillna(0.0)
    out = out.sort_values([rank_col, "kEUR"], ascending=[False, True]).reset_index(drop=True)

    out["Rank"] = out[rank_col].rank(method="dense", ascending=False).astype("Int64")
    total_val = float(out[rank_col].sum())
    out["% of"] = out[rank_col] / total_val if total_val != 0 else 0.0
    out["cum. % of"] = out["% of"].cumsum()

    a_th, b_th = cfg.get("abc_thresholds", (0.2, 0.4))
    out["ABC"] = np.select(
        [out["cum. % of"] <= a_th, out["cum. % of"] <= b_th],
        ["A", "B"],
        default="C",
    )

    MIN_B_PCT = 0.05
    if len(out) > 0 and not (out["ABC"] == "A").any():
        out.loc[out.index[0], "ABC"] = "A"
    if len(out) > 0 and not (out["ABC"] == "B").any():
        cand_idx_list = out.index[out["ABC"] == "C"].tolist()
        if cand_idx_list:
            cand_idx = cand_idx_list[0]
            cand_pct = pd.to_numeric(out.loc[cand_idx, "% of"], errors="coerce")
            cand_pct = 0.0 if pd.isna(cand_pct) else float(cand_pct)
            if cand_pct >= MIN_B_PCT:
                out.loc[cand_idx, "ABC"] = "B"

    meta = {
        "rank_col": rank_col,
        "pct_col": "% of",
        "cum_col": "cum. % of",
        "fy_delta_col": period_defs.get("fy_delta_label"),
        "ytd_delta_col": period_defs.get("ytd_delta_label"),
    }

    def compute_delta(label: str | None):
        if not label:
            return
        parts = label.replace("Δ ", "").split(" - ")
        if len(parts) != 2:
            return
        a, b = parts[0], parts[1]
        if a in out.columns and b in out.columns:
            out[label] = pd.to_numeric(out[a], errors="coerce").fillna(0.0) - pd.to_numeric(out[b], errors="coerce").fillna(0.0)

    compute_delta(meta["fy_delta_col"])
    compute_delta(meta["ytd_delta_col"])

    return out, meta


def apply_top_buckets_topreport(df: pd.DataFrame, cfg: dict, meta: dict) -> pd.DataFrame:
    tb = cfg.get("top_bucket", {}) or {}
    if not tb.get("enabled", False):
        out = df.copy()
        out["row_type"] = "leaf"
        return out

    out = df.copy()
    out["row_type"] = "leaf"

    pct_col = meta["pct_col"]
    cum_col = meta["cum_col"]

    out = out.sort_values(["Rank", "kEUR"], ascending=[True, True]).reset_index(drop=True)
    out[cum_col] = pd.to_numeric(out[cum_col], errors="coerce").fillna(0.0)

    bucket_mode = tb.get("bucket_mode", "threshold")
    thresholds = tuple(tb.get("thresholds", ()) or ())
    numbers = tuple(tb.get("numbers", ()) or ())
    create_other_bucket = bool(tb.get("create_other_bucket", True))
    other_bucket_label = str(tb.get("other_bucket_label", "Other")).strip() or "Other"

    non_sum_cols = {"kEUR", "Rank", "ABC", pct_col, cum_col, "row_type", "__top_key"}
    sum_cols = [c for c in out.columns if c not in non_sum_cols]

    use_formulas = bool(cfg.get("formula_mode", True))

    def make_bucket_row(label: str, start: int, end: int) -> dict:
        block = out.iloc[start:end]
        row = {c: "" for c in out.columns}
        row["kEUR"] = label
        row["row_type"] = "bucket"
        if not use_formulas:
            for c in sum_cols:
                row[c] = pd.to_numeric(block[c], errors="coerce").fillna(0.0).sum()
        return row

    def label_for_range(start: int, end: int) -> str:
        a = start + 1
        b = end
        if a == b:
            return f"Top {a}"
        if a == 1:
            return f"Top {b}"
        return f"Top {a}-{b}"

    bucket_ranges: list[tuple[int, int, str]] = []
    if bucket_mode == "threshold":
        prev_end = 0
        for th in thresholds:
            idx = int(out.index[out[cum_col] >= float(th)][0]) if (out[cum_col] >= float(th)).any() else len(out) - 1
            end = max(int(idx) + 1, prev_end)
            end = min(end, len(out))
            if end > prev_end:
                bucket_ranges.append((prev_end, end, label_for_range(prev_end, end)))
                prev_end = end
        if create_other_bucket and prev_end < len(out):
            bucket_ranges.append((prev_end, len(out), other_bucket_label))
    elif bucket_mode == "number":
        start = 0
        for n in numbers:
            end = min(start + int(n), len(out))
            if end > start:
                bucket_ranges.append((start, end, label_for_range(start, end)))
            start = end
            if start >= len(out):
                break
        if create_other_bucket and start < len(out):
            bucket_ranges.append((start, len(out), other_bucket_label))
    else:
        raise ValueError("top_bucket['bucket_mode'] muss 'threshold' oder 'number' sein.")

    # Leaf-Zeilen zuerst, Bucket-Summenzeile darunter (wie Excel-Gruppierung summaryBelow)
    if bucket_ranges:
        pieces: list[pd.DataFrame] = []
        for start, end, label in bucket_ranges:
            pieces.append(out.iloc[start:end].copy())
            pieces.append(pd.DataFrame([make_bucket_row(label, start, end)]))
        out = pd.concat(pieces, ignore_index=True)

    out.loc[out["row_type"] == "bucket", cum_col] = ""
    out.loc[out["row_type"] == "bucket", pct_col] = ""
    return out


def append_total_row(df: pd.DataFrame, cfg: dict, meta: dict) -> pd.DataFrame:
    out = df.copy()
    total_label = str(cfg.get("total_label", "Total")).strip() or "Total"
    pct_col = meta["pct_col"]
    cum_col = meta["cum_col"]
    base = out[out.get("row_type", "leaf").isin(["leaf"])].copy()
    use_formulas = bool(cfg.get("formula_mode", True))

    total = {c: "" for c in out.columns}
    total["kEUR"] = total_label
    total["row_type"] = "total"
    if not use_formulas:
        for c in out.columns:
            if c in {"kEUR", "Rank", "ABC", pct_col, cum_col, "row_type"}:
                continue
            if isinstance(c, str) and (
                c.startswith("FY") or c.startswith("YTD") or c.startswith("LTM") or c.startswith("Δ ")
            ):
                total[c] = pd.to_numeric(base[c], errors="coerce").fillna(0.0).sum()

    return pd.concat([out, pd.DataFrame([total])], ignore_index=True)


def arrange_columns(df: pd.DataFrame, period_defs: dict, meta: dict, cfg: dict) -> pd.DataFrame:
    out = df.copy()
    FYs = period_defs["FYs"]
    extra = period_defs["extra"]
    pct_col = meta["pct_col"]
    cum_col = meta["cum_col"]
    fy_delta = meta.get("fy_delta_col")
    ytd_delta = meta.get("ytd_delta_col")

    cols = ["kEUR"]
    cols += [c for c in FYs if c in out.columns]
    if fy_delta and fy_delta in out.columns:
        cols += [fy_delta]
    cols += ["Rank", "ABC", pct_col, cum_col]

    ytds = [c for c in extra if isinstance(c, str) and c.startswith("YTD") and c in out.columns]
    ltms = [c for c in extra if isinstance(c, str) and c.startswith("LTM") and c in out.columns]
    cols += ytds + ltms
    if ytd_delta and ytd_delta in out.columns:
        cols += [ytd_delta]

    if "row_type" in out.columns:
        cols += ["row_type"]

    cols = [c for c in cols if c in out.columns]
    return out[cols]


def _top_red_fonts():
    """PVM/GST helper styling (red dates + keys)."""
    red = "FFFF5149"
    red_font = Font(name=THEME.font_name, size=THEME.font_size, color=red)
    red_bold_font = Font(name=THEME.font_name, size=THEME.font_size, color=red, bold=True)
    return red_font, red_bold_font


def write_top_period_helper_rows(
    ws,
    period_defs: dict,
    *,
    period_col_indices: list[int],
    label_by_col: dict[int, str],
    period_start_row: int,
    period_end_row: int,
):
    red_font, _ = _top_red_fonts()
    period_map = period_defs["period_map"]

    for c in period_col_indices:
        label = str(label_by_col.get(c) or "").strip()
        if label not in period_map:
            continue
        s, e = period_map[label]
        c_start = ws.cell(row=period_start_row, column=c)
        c_end = ws.cell(row=period_end_row, column=c)
        c_start.value = pd.Timestamp(s).to_pydatetime()
        c_end.value = pd.Timestamp(e).to_pydatetime()
        c_start.number_format = "DD.MM.YYYY"
        c_end.number_format = "DD.MM.YYYY"
        c_start.font = red_font
        c_end.font = red_font
        c_start.alignment = Alignment(horizontal="right", vertical="center")
        c_end.alignment = Alignment(horizontal="right", vertical="center")


def _top_period_columns(header_to_col: dict[str, int], period_defs: dict) -> tuple[list[int], dict[int, str]]:
    """FY/YTD/LTM columns and col → period label (for hidden bounds + SUMIFS)."""
    period_map = period_defs.get("period_map") or {}
    indices: list[int] = []
    label_by_col: dict[int, str] = {}
    for label, col_idx in header_to_col.items():
        if label not in period_map:
            continue
        if not (
            label.startswith("FY")
            or label.startswith("YTD")
            or label.startswith("LTM")
        ):
            continue
        c = int(col_idx)
        indices.append(c)
        label_by_col[c] = label
    return sorted(indices), label_by_col


def _top_period_columns_from_col_headers(
    col_to_header: dict[int, str],
    period_defs: dict,
) -> tuple[list[int], dict[int, str]]:
    header_to_col = {
        str(hdr).strip(): c
        for c, hdr in col_to_header.items()
        if hdr is not None and str(hdr).strip()
    }
    return _top_period_columns(header_to_col, period_defs)


def write_top_formula_key_column(
    ws,
    cfg: dict,
    layout: dict,
    *,
    label_col: int,
    key_col: int,
    row_type_col: int,
    table_bottom_row: int,
):
    red_font, red_bold_font = _top_red_fonts()
    header_row = layout["header_row"]
    data_start_row = layout["data_start_row"]
    helper_right = layout["helper_right_col"]

    ws.cell(row=header_row, column=key_col).value = str(cfg.get("top_col", "")).strip()
    for c in range(1, helper_right + 1):
        ws.cell(row=header_row, column=c).font = red_bold_font
        ws.cell(row=header_row, column=c).alignment = Alignment(
            horizontal="left", vertical="bottom", wrap_text=True
        )

    for r in range(data_start_row, table_bottom_row + 1):
        rt = ws.cell(row=r, column=row_type_col).value
        rt_s = str(rt or "").strip().lower()
        visible = ws.cell(row=r, column=label_col).value

        for c in range(1, helper_right + 1):
            ws.cell(row=r, column=c).font = red_font
            ws.cell(row=r, column=c).alignment = Alignment(horizontal="left", vertical="center")
            if c != key_col:
                ws.cell(row=r, column=c).value = None

        if rt_s in {"bucket", "total"}:
            ws.cell(row=r, column=key_col).value = None
        elif rt_s == "leaf":
            ws.cell(row=r, column=key_col).value = visible


def format_top_report_excel(
    wb,
    target_sheet_name: str,
    project_name: str,
    table_name: str,
    company_name: str,
    cfg: dict,
    period_defs: dict,
    *,
    source_sheet_name: str | None = None,
    source_df: pd.DataFrame | None = None,
):
    layout = get_top_excel_layout(cfg)
    col_offset = layout["col_offset"]
    HEADER_ROW = layout["header_row"]
    DATA_START_ROW = layout["data_start_row"]
    COMPANY_ROW = layout["company_row"]
    PERIOD_START_ROW = layout["period_start_row"]
    PERIOD_END_ROW = layout["period_end_row"]

    if target_sheet_name not in wb.sheetnames:
        raise ValueError(f"Sheet '{target_sheet_name}' nicht im Workbook gefunden.")
    ws = wb[target_sheet_name]

    ws.insert_rows(1, amount=layout["insert_rows"])
    ws.insert_cols(1, amount=col_offset)

    TABLE_LEFT_COL = layout["label_col"]
    TABLE_RIGHT_COL = ws.max_column
    TABLE_TOP_ROW = HEADER_ROW
    TABLE_BOTTOM_ROW = ws.max_row

    fmt_money = '#,##0;(#,##0);"-"'
    fmt_pct = '0.0%'
    WIDTH_FACTOR = 1.105

    def set_width(col_idx: int, width: float):
        ws.column_dimensions[get_column_letter(col_idx)].width = float(width) * WIDTH_FACTOR

    WIDTHS = {
        "kEUR": 12.5,
        "period": 6.43,
        "rank_abc": 4.5,
        "delta": 6.75,
        "percent": 5.2,
        "row_type": 2,
        "top_key": 10,
    }

    base_font = THEME.font_base
    bold_font = THEME.font_bold
    header_fill = THEME.fill_header
    white_fill = THEME.fill_white
    fill_tech = THEME.fill_tech
    fill_subtotal = THEME.fill_subtotal

    font_A = Font(name=THEME.font_name, size=THEME.font_size, color="FF5E2ABA")
    font_B = Font(name=THEME.font_name, size=THEME.font_size, color="FF8E5BBE")
    font_C = Font(name=THEME.font_name, size=THEME.font_size, color="FFBFA6D9")

    POS_COLOR = THEME.delta_positive
    NEG_COLOR = THEME.delta_negative
    GREY_TEXT = THEME.delta_zero

    dashed = Side(style="dashed", color=THEME.border_color)
    border_subtotal_top = THEME.border_subtotal_top
    border_subtotal_tb = Border(
        top=THEME.border_subtotal_top.top,
        bottom=THEME.border_subtotal_top.top,
    )
    header_alignment = Alignment(wrap_text=True, vertical="bottom")
    header_border = THEME.border_header_bottom

    header_map = {}
    for c in range(1, ws.max_column + 1):
        header_map[c] = ws.cell(row=HEADER_ROW, column=c).value

    def h(col: int):
        v = header_map.get(col)
        return v.strip() if isinstance(v, str) else v

    col_kEUR = None
    col_abc = None
    col_row_type = None
    money_cols = []
    percent_cols = []
    delta_cols = []

    for c in range(1, TABLE_RIGHT_COL + 1):
        hv = h(c)
        if hv == "kEUR":
            col_kEUR = c
        elif hv == "ABC":
            col_abc = c
        elif hv == "row_type":
            col_row_type = c

        if isinstance(hv, str) and (hv.startswith("%") or hv.startswith("cum")):
            percent_cols.append(c)

        if isinstance(hv, str) and (hv.startswith("FY") or hv.startswith("YTD") or hv.startswith("LTM") or hv.startswith("Δ ")):
            money_cols.append(c)

        if isinstance(hv, str) and hv.startswith("Δ "):
            delta_cols.append(c)

    if col_kEUR is None:
        raise ValueError("Header 'kEUR' nicht gefunden – kann Layout nicht anwenden.")
    if col_row_type is None:
        raise ValueError("Spalte 'row_type' fehlt. Bitte in arrange_columns() behalten und exportieren.")

    key_col = layout["key_col"]
    label_col = layout["label_col"]
    if col_kEUR is None:
        col_kEUR = label_col

    # A–C eingeklappt (wie GST/PVM); row_type am Ende versteckt
    ws.sheet_properties.outlinePr.summaryRight = True
    for c in range(1, col_offset + 1):
        letter = get_column_letter(c)
        ws.column_dimensions[letter].hidden = True
        ws.column_dimensions[letter].outlineLevel = 1
    ws.column_dimensions[get_column_letter(col_offset + 1)].collapsed = True
    ws.column_dimensions[get_column_letter(col_row_type)].hidden = True

    # Basic fill/background
    for r in range(1, 501):
        for c in range(1, 501):
            inside_table = (TABLE_TOP_ROW <= r <= TABLE_BOTTOM_ROW and TABLE_LEFT_COL <= c <= TABLE_RIGHT_COL)
            ws.cell(row=r, column=c).fill = white_fill if inside_table else white_fill

    for r in range(1, 501):
        for c in range(1, 1 + col_offset):
            ws.cell(row=r, column=c).fill = fill_tech

    ws.row_dimensions[HEADER_ROW].height = 24
    for r in range(DATA_START_ROW, ws.max_row + 1):
        ws.row_dimensions[r].height = 12

    for c in range(TABLE_LEFT_COL, TABLE_RIGHT_COL + 1):
        hdr = ws.cell(row=HEADER_ROW, column=c)
        hdr.fill = header_fill
        hdr.font = THEME.font_header
        hdr.alignment = header_alignment
        hdr.border = header_border

    subtitle_suffix = str(cfg.get("subtitle_suffix", "")).strip()
    company_line = str(company_name or "").strip()
    if subtitle_suffix:
        company_line = f"{company_line} | {subtitle_suffix}" if company_line else subtitle_suffix

    blue_text = THEME.text_brand_title
    title_cell = ws.cell(row=layout["title_row"], column=label_col)
    title_cell.value = project_name
    title_cell.font = Font(name=THEME.font_name, size=THEME.font_size_title, color=blue_text)

    table_cell = ws.cell(row=layout["subtitle_row"], column=label_col)
    display_table_name = str(table_name or "").strip()
    # If table name is "TOP (…)", show only the content (no TOP / no parentheses)
    m = re.match(r"(?is)\s*TOP\s*\((.+)\)\s*$", display_table_name)
    if m:
        display_table_name = m.group(1).strip()
    table_cell.value = display_table_name
    table_cell.font = Font(name=THEME.font_name, size=THEME.font_size_subtitle, color=blue_text)

    company_cell = ws.cell(row=COMPANY_ROW, column=label_col)
    company_cell.value = company_line
    company_cell.font = Font(
        name=THEME.font_name, size=THEME.font_size_company, color=blue_text, bold=True
    )

    # column widths
    set_width(col_kEUR, WIDTHS["kEUR"])
    col_rank = None
    for c in range(TABLE_LEFT_COL, TABLE_RIGHT_COL + 1):
        if h(c) == "Rank":
            col_rank = c
            break
    if col_rank:
        set_width(col_rank, WIDTHS["rank_abc"])
    if col_abc:
        set_width(col_abc, WIDTHS["rank_abc"])
    for c in money_cols:
        set_width(c, WIDTHS["period"])
    for c in delta_cols:
        set_width(c, WIDTHS["delta"])
    for c in percent_cols:
        set_width(c, WIDTHS["percent"])

    # Base font (ohne Perioden-Hilfszeilen — die bleiben rot)
    for r in range(HEADER_ROW, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            ws.cell(row=r, column=c).font = base_font

    # Header alignment
    for c in range(1, ws.max_column + 1):
        cell = ws.cell(row=HEADER_ROW, column=c)
        if c == col_kEUR:
            cell.alignment = Alignment(horizontal="left", vertical="bottom", wrap_text=True)
        elif c == col_abc:
            cell.alignment = Alignment(horizontal="center", vertical="bottom", wrap_text=True)
        else:
            cell.alignment = Alignment(horizontal="right", vertical="bottom", wrap_text=True)

    # Data alignment: kEUR linksbündig, Rest rechts
    for r in range(DATA_START_ROW, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            horiz = "left" if c == col_kEUR else (ws.cell(row=r, column=c).alignment.horizontal or "right")
            ws.cell(row=r, column=c).alignment = Alignment(
                horizontal=horiz,
                vertical="center",
                wrap_text=ws.cell(row=r, column=c).alignment.wrap_text,
            )
        ws.cell(row=r, column=key_col).alignment = Alignment(
            horizontal="left", vertical="center", wrap_text=True
        )

    # Number formats (apply even for formulas; Excel evaluates later)
    for c in money_cols:
        for r in range(DATA_START_ROW, ws.max_row + 1):
            ws.cell(row=r, column=c).number_format = fmt_money

    for c in percent_cols:
        for r in range(DATA_START_ROW, ws.max_row + 1):
            ws.cell(row=r, column=c).number_format = fmt_pct

    # ABC formatting
    if col_abc:
        abc_center = Alignment(horizontal="center", vertical="center")
        for r in range(DATA_START_ROW, ws.max_row + 1):
            cell = ws.cell(row=r, column=col_abc)
            if isinstance(cell.value, str):
                cell.alignment = abc_center
                if cell.value == "A":
                    cell.font = font_A
                elif cell.value == "B":
                    cell.font = font_B
                elif cell.value == "C":
                    cell.font = font_C

    # Delta coloring (use conditional formatting; works with formulas)
    pos_font = Font(color=POS_COLOR)
    neg_font = Font(color=NEG_COLOR)
    zero_font = Font(color=GREY_TEXT)
    for c in delta_cols:
        col_letter = get_column_letter(c)
        rng = f"{col_letter}{DATA_START_ROW}:{col_letter}{TABLE_BOTTOM_ROW}"
        ws.conditional_formatting.add(rng, CellIsRule(operator="greaterThan", formula=["0"], font=pos_font))
        ws.conditional_formatting.add(rng, CellIsRule(operator="lessThan", formula=["0"], font=neg_font))
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=["0"], font=zero_font))

    # dashed borders around delta cols
    for r in range(HEADER_ROW, TABLE_BOTTOM_ROW + 1):
        for c in delta_cols:
            cell = ws.cell(row=r, column=c)
            b = copy(cell.border)
            b.left = dashed
            b.right = dashed
            cell.border = b

    # Lines above bucket/total and bottom at total
    total_row_idx = None
    bucket_rows: list[int] = []
    for r in range(DATA_START_ROW, ws.max_row + 1):
        rt = ws.cell(row=r, column=col_row_type).value
        rt_s = str(rt).strip().lower() if isinstance(rt, str) else ""
        if rt_s == "bucket":
            bucket_rows.append(r)
            for c in range(TABLE_LEFT_COL, TABLE_RIGHT_COL + 1):
                cell = ws.cell(row=r, column=c)
                cell.fill = fill_subtotal
                cell.font = bold_font
                cell.border = border_subtotal_top
        elif rt_s == "total":
            total_row_idx = r
            for c in range(TABLE_LEFT_COL, TABLE_RIGHT_COL + 1):
                cell = ws.cell(row=r, column=c)
                cell.fill = fill_subtotal
                cell.font = bold_font
                cell.border = border_subtotal_tb

    # Perioden-Hilfszeilen + Schlüsselspalte B (nach Basis-Font, damit Rot erhalten bleibt)
    if PERIOD_START_ROW and PERIOD_END_ROW and cfg.get("formula_mode", True):
        period_col_indices, label_by_col = _top_period_columns_from_col_headers(
            header_map, period_defs,
        )
        write_top_period_helper_rows(
            ws,
            period_defs,
            period_col_indices=period_col_indices,
            label_by_col=label_by_col,
            period_start_row=PERIOD_START_ROW,
            period_end_row=PERIOD_END_ROW,
        )
        for r in (PERIOD_START_ROW, PERIOD_END_ROW):
            ws.row_dimensions[r].hidden = True
            ws.row_dimensions[r].outlineLevel = 1
            ws.row_dimensions[r].collapsed = False
            ws.row_dimensions[r].height = 12
        ws.row_dimensions[PERIOD_END_ROW].collapsed = True

    if cfg.get("formula_mode", True):
        write_top_formula_key_column(
            ws,
            cfg,
            layout,
            label_col=label_col,
            key_col=key_col,
            row_type_col=col_row_type,
            table_bottom_row=TABLE_BOTTOM_ROW,
        )

    # Outline: Detailzeilen oben, Bucket-Summe darunter (summaryBelow)
    ws.sheet_properties.outlinePr.summaryBelow = True
    if total_row_idx is not None:
        for r in range(DATA_START_ROW, total_row_idx + 1):
            rt = ws.cell(row=r, column=col_row_type).value
            rt_s = str(rt).strip().lower() if isinstance(rt, str) else ""
            rd = ws.row_dimensions[r]
            if rt_s == "bucket":
                rd.outlineLevel = 0
                rd.hidden = False
                rd.collapsed = False
            elif rt_s == "total":
                rd.outlineLevel = 0
                rd.hidden = False
                rd.collapsed = False
            elif rt_s == "leaf":
                rd.outlineLevel = 1
                rd.hidden = False
                rd.collapsed = False

    if cfg.get("formula_mode", True) and source_df is not None and source_sheet_name:
        n_before = _count_sheet_formulas(ws)
        apply_top_report_formulas(
            wb=wb,
            report_sheet_name=target_sheet_name,
            source_sheet_name=source_sheet_name,
            cfg=cfg,
            period_defs=period_defs,
            source_df=source_df,
        )
        n_after = _count_sheet_formulas(ws)
        print(f"[TOP] Formeln im Sheet: {n_after} (vorher {n_before})")
        hdr = _get_report_header_map(ws, layout["header_row"])
        if "row_type" in hdr:
            col_rt = hdr["row_type"]
            bottom = _report_data_bottom_row(ws, col_rt, layout["data_start_row"])
            has_leaf = any(
                str(ws.cell(r, col_rt).value or "").strip().lower() == "leaf"
                for r in range(layout["data_start_row"], bottom + 1)
            )
            if has_leaf and n_after == 0:
                raise RuntimeError(
                    "TOP formula_mode: Es wurden keine Excel-Formeln geschrieben, "
                    "obwohl Detailzeilen (leaf) vorhanden sind. "
                    "Bitte top_col, value_col und invoice_col prüfen."
                )


def _get_report_header_map(ws, header_row: int) -> dict[str, int]:
    out: dict[str, int] = {}
    for c in range(1, ws.max_column + 1):
        v = ws.cell(row=header_row, column=c).value
        if v is None:
            continue
        out[str(v).strip()] = c
    return out


def _col_ref(col_idx: int, row_idx: int, *, abs_col: bool = False, abs_row: bool = False) -> str:
    c = get_column_letter(col_idx)
    return f"{'$' if abs_col else ''}{c}{'$' if abs_row else ''}{row_idx}"


def _excel_abs_ref(col_idx: int, row_idx: int) -> str:
    return f"${get_column_letter(col_idx)}${row_idx}"


def apply_top_report_formulas(
    *,
    wb,
    report_sheet_name: str,
    source_sheet_name: str,
    cfg: dict,
    period_defs: dict,
    source_df: pd.DataFrame,
):
    """
    Verformelt-Modus (analog zu GST/PVM):
    - invoice: SUMIFS auf value_col mit Datums-Grenzen aus Periodenzeilen
    - accrual: Hilfsspalten via add_accrual_col_in_ws (wiederverwendet falls vorhanden), dann SUMIFS
    - Deltas: Differenz zweier Perioden-Spalten im Report
  """
    if not cfg.get("formula_mode", True):
        return
    if bool(cfg.get("apply_fx", False)):
        print("[WARN] formula_mode aktuell nicht mit apply_fx kombiniert – überspringe Formeln.")
        return

    calc_mode = str(cfg.get("calc_mode", "invoice")).strip().lower()

    if report_sheet_name not in wb.sheetnames:
        raise ValueError(f"Report-Sheet '{report_sheet_name}' nicht gefunden.")
    if source_sheet_name not in wb.sheetnames:
        raise ValueError(f"Source-Sheet '{source_sheet_name}' nicht gefunden.")

    ws = wb[report_sheet_name]
    ws_source = wb[source_sheet_name]

    refresh_source_sheet_from_df(ws_source, source_df, cfg)

    layout = get_top_excel_layout(cfg)
    HEADER_ROW = layout["header_row"]
    DATA_START_ROW = layout["data_start_row"]
    PERIOD_START_ROW = layout["period_start_row"]
    PERIOD_END_ROW = layout["period_end_row"]

    header_map = _get_report_header_map(ws, HEADER_ROW)
    if "row_type" not in header_map:
        raise ValueError("Erwarte Spalte 'row_type' im Report (Header-Zeile).")

    col_key = layout["key_col"]
    label_col = layout["label_col"]
    col_row_type = header_map["row_type"]

    table_bottom = _report_data_bottom_row(ws, col_row_type, DATA_START_ROW)

    write_top_formula_key_column(
        ws,
        cfg,
        layout,
        label_col=label_col,
        key_col=col_key,
        row_type_col=col_row_type,
        table_bottom_row=table_bottom,
    )

    src_header = _get_source_header_map(ws_source)
    required_src = [cfg["top_col"], cfg["invoice_col"], cfg["value_col"]]
    if calc_mode == "accrual":
        required_src += [cfg["start_col"], cfg["end_col"]]
    for required in required_src:
        source_range_ref(source_sheet_name, src_header, str(required).strip())

    top_rng = source_range_ref(source_sheet_name, src_header, cfg["top_col"])

    active_period_labels = list(period_defs["FYs"]) + list(period_defs.get("extra") or [])
    period_map: dict[str, tuple[pd.Timestamp, pd.Timestamp]] = {
        k: period_defs["period_map"][k] for k in active_period_labels if k in period_defs["period_map"]
    }

    if calc_mode == "invoice" and (not PERIOD_START_ROW or not PERIOD_END_ROW):
        raise ValueError("Period helper rows fehlen – format_top_report_excel zuerst ausführen.")

    period_col_indices, label_by_col = _top_period_columns(header_map, period_defs)
    if calc_mode == "invoice" and period_col_indices:
        write_top_period_helper_rows(
            ws,
            period_defs,
            period_col_indices=period_col_indices,
            label_by_col=label_by_col,
            period_start_row=PERIOD_START_ROW,
            period_end_row=PERIOD_END_ROW,
        )

    # Pre-materialize accrual helper columns once per period (not per leaf cell).
    accrual_sum_ranges: dict[str, str] = {}
    if calc_mode == "accrual":
        helper_col_by_label: dict[str, str] = {}
        helper_col_idx_by_label: dict[str, int] = {}
        for period_label, (p_start, p_end) in period_map.items():
            _helper_col, helper_idx = add_accrual_col_in_ws(
                ws_source=ws_source,
                source_df=source_df,
                value_col=cfg["value_col"],
                invoice_col=cfg["invoice_col"],
                start_col=cfg["start_col"],
                end_col=cfg["end_col"],
                period_label=period_label,
                period_start=p_start,
                period_end=p_end,
            )
            helper_col_idx_by_label[period_label] = helper_idx
        for period_label, helper_idx in helper_col_idx_by_label.items():
            accrual_sum_ranges[period_label] = source_column_range_ref(
                source_sheet_name, helper_idx
            )

    def period_sumifs_formula(row_idx: int, period_col_idx: int, period_label: str) -> str:
        key_ref = f"${get_column_letter(col_key)}{row_idx}"
        base_pairs = [(top_rng, key_ref)]

        if calc_mode == "invoice":
            inv_rng = source_range_ref(source_sheet_name, src_header, cfg["invoice_col"])
            val_rng = source_range_ref(source_sheet_name, src_header, cfg["value_col"])
            start_ref = _excel_abs_ref(period_col_idx, PERIOD_START_ROW)
            end_ref = _excel_abs_ref(period_col_idx, PERIOD_END_ROW)
            base_pairs += [
                (inv_rng, f'">="&{start_ref}'),
                (inv_rng, f'"<="&{end_ref}'),
            ]
            sum_rng = val_rng
        else:
            if period_label in accrual_sum_ranges:
                sum_rng = accrual_sum_ranges[period_label]
            else:
                period_start, period_end = period_map[period_label]
                _helper_col, helper_idx = add_accrual_col_in_ws(
                    ws_source=ws_source,
                    source_df=source_df,
                    value_col=cfg["value_col"],
                    invoice_col=cfg["invoice_col"],
                    start_col=cfg["start_col"],
                    end_col=cfg["end_col"],
                    period_label=period_label,
                    period_start=period_start,
                    period_end=period_end,
                )
                sum_rng = source_column_range_ref(source_sheet_name, helper_idx)

        body = build_sumifs_formula_body(sum_rng, base_pairs, [[]], divide_by_1000=True)
        return f'=IFERROR({body},"")'

    # row ranges (buckets/totals)
    leaf_rows: list[int] = []
    bucket_rows: list[int] = []
    total_row: int | None = None
    for r in range(DATA_START_ROW, table_bottom + 1):
        rt = ws.cell(row=r, column=col_row_type).value
        rt_s = str(rt).strip().lower() if isinstance(rt, str) else ""
        if rt_s == "leaf":
            leaf_rows.append(r)
        elif rt_s == "bucket":
            bucket_rows.append(r)
        elif rt_s == "total":
            total_row = r

    if not leaf_rows:
        print("[WARN] TOP formula_mode: keine leaf-Zeilen gefunden – überspringe Periodenformeln.")
        return

    def sum_formula(rows: list[int], col_idx: int) -> str:
        return build_chunked_row_sum_formula(rows, col_idx)

    # period columns: formulas for leaf/bucket/total
    for header in active_period_labels:
        if header not in header_map:
            continue
        col_idx = header_map[header]
        if header in period_map:
            # leaf
            for r in leaf_rows:
                _set_cell_formula(ws, r, col_idx, period_sumifs_formula(r, col_idx, header))
            for i, br in enumerate(bucket_rows):
                start_r = (bucket_rows[i - 1] + 1) if i > 0 else DATA_START_ROW
                rows_in_bucket = [rr for rr in leaf_rows if start_r <= rr < br]
                _set_cell_formula(ws, br, col_idx, sum_formula(rows_in_bucket, col_idx))
            if total_row is not None:
                _set_cell_formula(ws, total_row, col_idx, sum_formula(leaf_rows, col_idx))

    # Delta columns: '=IFERROR(curr-prev,"")'
    for header, col_idx in header_map.items():
        if isinstance(header, str) and header.startswith("Δ "):
            parts = header.replace("Δ ", "", 1).split(" - ")
            if len(parts) != 2:
                continue
            a, b = parts[0].strip(), parts[1].strip()
            if a not in header_map or b not in header_map:
                continue
            a_col = header_map[a]
            b_col = header_map[b]

            def delta_formula(row_idx: int) -> str:
                a_ref = f"{get_column_letter(a_col)}{row_idx}"
                b_ref = f"{get_column_letter(b_col)}{row_idx}"
                return f'=IFERROR({a_ref}-{b_ref},"")'

            for r in leaf_rows:
                _set_cell_formula(ws, r, col_idx, delta_formula(r))
            for i, br in enumerate(bucket_rows):
                start_r = (bucket_rows[i - 1] + 1) if i > 0 else DATA_START_ROW
                rows_in_bucket = [rr for rr in leaf_rows if start_r <= rr < br]
                _set_cell_formula(ws, br, col_idx, sum_formula(rows_in_bucket, col_idx))
            if total_row is not None:
                _set_cell_formula(ws, total_row, col_idx, sum_formula(leaf_rows, col_idx))

    _warn_formula_integrity(ws, "TOP report")
    wb.calculation.forceFullCalc = True
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.calcMode = "auto"


def main():
    if len(sys.argv) < 2:
        raise ValueError("Bitte den Pfad zur topconfig.json als Argument übergeben.")

    config_path = sys.argv[1]
    if not os.path.exists(config_path):
        raise FileNotFoundError(f"Config-Datei nicht gefunden: {config_path}")

    with open(config_path, "r", encoding="utf-8-sig") as f:
        config = json.load(f)

    cfg = normalize_config(config)

    df = pd.read_excel(cfg["file_path"], sheet_name=cfg["sheet_name"], engine="openpyxl")
    d = preprocess_top_input(df, cfg)

    period_defs = build_period_definitions(cfg)
    wide = build_wide_table(d, period_defs, cfg)
    wide = drop_rows_all_zero_periods(wide)
    with_metrics, meta = add_metrics(wide, period_defs, cfg)
    with_metrics = drop_rows_all_zero_relevant_values(with_metrics)
    bucketed = apply_top_buckets_topreport(with_metrics, cfg, meta)
    final = append_total_row(bucketed, cfg, meta)
    final = arrange_columns(final, period_defs, meta, cfg)
    if cfg.get("formula_mode", True):
        final = strip_formula_value_columns(final, meta)
    else:
        for c in final.columns:
            if c in {"kEUR", "Rank", "ABC", "row_type", meta["pct_col"], meta["cum_col"]}:
                continue
            if isinstance(c, str) and (
                c.startswith("FY") or c.startswith("YTD") or c.startswith("LTM") or c.startswith("Δ ")
            ):
                final[c] = to_kEUR(final[c])

    output_path = build_output_file_path(cfg)
    os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)
    ensure_output_writable(output_path)
    source_sheet_name = ensure_source_sheet_in_output(cfg, output_path)

    wb = load_workbook(output_path)
    target_sheet_name = get_next_sheet_name_from_wb(wb, cfg.get("base_sheet_name", "TOP"))
    write_export_df_to_sheet(wb, final, target_sheet_name)

    print(f"Excel erfolgreich geschrieben: {output_path}")
    print(f"Neues Sheet: {target_sheet_name}")

    suffix = str(cfg.get("subtitle_suffix", "")).strip()
    table_name_for_sheet = suffix if suffix else str(cfg.get("table", "TOP") or "TOP").strip()

    format_top_report_excel(
        wb=wb,
        target_sheet_name=target_sheet_name,
        project_name=str(cfg.get("title", "")),
        table_name=table_name_for_sheet,
        company_name=str(cfg.get("company", "")),
        cfg=cfg,
        period_defs=period_defs,
        source_sheet_name=source_sheet_name,
        source_df=df,
    )

    print("Formatierung angewendet.")

    wb.save(output_path)
    wb.close()


if __name__ == "__main__":
    main()

