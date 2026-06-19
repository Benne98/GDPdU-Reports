"""
Extract financial statement data from PDF annual reports via PDF→Markdown→Claude.
JSON config as first CLI argument (FDD backend).
"""
from __future__ import annotations

import json
import os
import re
import shutil
import sys
from pathlib import Path
from typing import Any, Dict, List, Optional

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Border, Font, PatternFill, Side

PROJECT_ROOT = Path(__file__).resolve().parent
SCRIPTS_DIR = PROJECT_ROOT / "scripts"
if str(SCRIPTS_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPTS_DIR))

try:
    from dotenv import load_dotenv

    load_dotenv(PROJECT_ROOT / "backend" / ".env")
except ImportError:
    pass

from pdf_to_md import pdf_to_markdown

CONFIG: Dict[str, Any] = {}
DEFAULT_ANTHROPIC_MODEL = "claude-sonnet-4-20250514"


def _anthropic_error_message(exc: Exception) -> str:
    body = str(exc)
    lowered = body.lower()
    if "organization has been disabled" in lowered:
        return (
            "Anthropic API: Die Organisation zu diesem API-Key ist deaktiviert. "
            "Bitte in der Anthropic Console (console.anthropic.com) prüfen und "
            "einen neuen API-Key einer aktiven Organisation in backend/.env eintragen."
        )
    if "invalid x-api-key" in lowered or "authentication" in lowered:
        return (
            "Anthropic API: Ungültiger API-Key. Bitte ANTHROPIC_API_KEY in backend/.env prüfen."
        )
    if "credit balance" in lowered or "billing" in lowered:
        return (
            "Anthropic API: Kein Guthaben / Billing-Problem. Bitte Zahlungsdaten in der Anthropic Console prüfen."
        )
    return f"Anthropic API error: {body}"


def load_config(path: Path) -> Dict[str, Any]:
    raw = json.loads(path.read_text(encoding="utf-8-sig"))
    anthropic = dict(raw.get("anthropic") or {})
    anthropic.setdefault("api_key", os.environ.get("ANTHROPIC_API_KEY", ""))
    anthropic.setdefault(
        "model",
        os.environ.get("ANTHROPIC_MODEL") or DEFAULT_ANTHROPIC_MODEL,
    )
    raw["anthropic"] = anthropic
    raw.setdefault("fy_month_bs", "Dec")
    return raw


def normalize_header(h: str) -> str:
    if not h:
        return ""
    return h.strip().lower().replace("\n", " ").replace("  ", " ")


def sign_factor(val) -> int:
    return -1 if str(val).strip() == "(1)" else 1


def ensure_template(template_path: Path, num_years: int = 3) -> Path:
    if template_path.is_file():
        return template_path
    template_path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)
    headers = [
        "German Mapping", "English Mapping", "Sign", "Total", "Group", "Level",
    ] + [f"FY_{i}" for i in range(1, num_years + 1)]
    sample_rows = [
        ("Anlagevermögen", "Fixed assets", 1, "Yes", "Fixed assets", "L2"),
        ("Umlaufvermögen", "Current assets", 1, "Yes", "Current assets", "L2"),
        ("Eigenkapital", "Equity", 1, "Yes", "Equity", "L2"),
        ("Umsatzerlöse", "Net sales", 1, "No", "Net sales", "L4"),
        ("Materialaufwand", "Cost of goods sold", 1, "No", "Cost of goods sold", "L4"),
        ("Jahresüberschuss", "Net result", 1, "Yes", "All", "L4"),
    ]
    for sheet in ("BS", "PL"):
        ws = wb.create_sheet(sheet)
        for c, h in enumerate(headers, start=1):
            ws.cell(1, c, h)
        for r, row in enumerate(sample_rows, start=2):
            for c, val in enumerate(row, start=1):
                ws.cell(r, c, val)
    wb.save(template_path)
    return template_path


def extract_json_from_claude_message(message) -> Dict:
    text_parts = [
        b.text.strip()
        for b in message.content
        if hasattr(b, "text") and b.text and b.text.strip()
    ]
    if not text_parts:
        raise ValueError("Claude returned no text blocks.")
    text = "\n".join(text_parts).strip()
    if text.startswith("```"):
        text = text.strip("`")
        if text.lower().startswith("json"):
            text = text[4:].strip()
    return json.loads(text)


def infer_year_from_filename(name: str) -> Optional[int]:
    m = re.search(r"(20\d{2})", name)
    return int(m.group(1)) if m else None


def convert_pdf_to_markdown(pdf_path: Path, cfg: Dict[str, Any]) -> str:
    cache_dir = cfg.get("pdf_md_cache_dir")
    if not cache_dir:
        output_excel = Path(str(cfg.get("output_excel") or ""))
        cache_dir = output_excel.parent / ".pdf-md-cache"
    cache_dir = Path(str(cache_dir))
    cache_dir.mkdir(parents=True, exist_ok=True)

    markdown = pdf_to_markdown(
        pdf_path,
        cache_dir=cache_dir,
        show_progress=sys.stderr.isatty(),
    )

    md_output_dir = cfg.get("md_output_dir")
    if md_output_dir:
        md_dir = Path(str(md_output_dir))
        md_dir.mkdir(parents=True, exist_ok=True)
        md_path = md_dir / f"{pdf_path.stem}.md"
        md_path.write_text(markdown, encoding="utf-8")
        print(f"Saved markdown: {md_path}")

    return markdown


def extract_year_data(pdf_path: Path, cfg: Dict[str, Any]) -> Dict:
    from anthropic import Anthropic

    api_key = str(cfg.get("anthropic", {}).get("api_key") or "").strip()
    if not api_key:
        raise ValueError("ANTHROPIC_API_KEY is not set.")

    model = str(cfg.get("anthropic", {}).get("model") or DEFAULT_ANTHROPIC_MODEL)
    client = Anthropic(api_key=api_key)

    print(f"Converting {pdf_path.name} to markdown...")
    markdown = convert_pdf_to_markdown(pdf_path, cfg)

    hint_year = infer_year_from_filename(pdf_path.name)
    year_hint = f" Filename hint year: {hint_year}." if hint_year else ""

    system_prompt = f"""
Extract all balance sheet and P&L positions from this German annual financial statement.
The document content is provided as Markdown converted from the PDF.{year_hint}

For EACH position return:
- numeric value (as number, no thousand separators)
- source column: either "GESCHAEFTSJAHR" or "EURO"

Rules:
- Values in column "Geschäftsjahr" are TOTALS.
- Values in column "Euro" are SUBPOSITIONS.
- Totals belong to the last over-position without a Euro value.
- Do NOT calculate anything.
- Do NOT use Vorjahr (prior year).
- Detect the fiscal year from document title, balance sheet date, or filename.
- Use markdown tables and headings to locate BS and PL sections.

Return STRICT JSON:
{{
    "year": 2023,
    "BS": {{ "<german label>": {{"value": 123, "source": "EURO|GESCHAEFTSJAHR"}} }},
    "PL": {{ "<german label>": {{"value": 456, "source": "EURO|GESCHAEFTSJAHR"}} }}
}}
"""

    # Keep prompt within reasonable token limits while preserving tables
    max_chars = 180_000
    if len(markdown) > max_chars:
        print(f"Warning: markdown truncated from {len(markdown)} to {max_chars} chars")
        markdown = markdown[:max_chars]

    try:
        message = client.messages.create(
            model=model,
            system=system_prompt,
            max_tokens=4096,
            temperature=0,
            messages=[{
                "role": "user",
                "content": (
                    f"PDF file: {pdf_path.name}\n\n"
                    f"--- MARKDOWN START ---\n{markdown}\n--- MARKDOWN END ---\n\n"
                    "Return strictly as JSON."
                ),
            }],
        )
    except Exception as exc:
        raise ValueError(_anthropic_error_message(exc)) from exc

    data = extract_json_from_claude_message(message)
    if "year" not in data:
        if hint_year:
            data["year"] = hint_year
        else:
            raise ValueError(f"Could not detect year for {pdf_path.name}")
    data.setdefault("BS", {})
    data.setdefault("PL", {})
    return data


def write_final_excel(all_data: Dict[int, Dict], cfg: Dict[str, Any]) -> None:
    template_path = Path(str(cfg.get("template_path") or ""))
    output_excel = Path(str(cfg["output_excel"]))
    ensure_template(template_path, num_years=max(3, len(all_data)))
    shutil.copy(template_path, output_excel)
    wb = load_workbook(output_excel)
    years = sorted(all_data.keys())
    fy_month_bs = str(cfg.get("fy_month_bs") or "Dec")

    for sheet in ["BS", "PL"]:
        ws = wb[sheet]
        raw_headers = {
            normalize_header(ws.cell(1, c).value): c
            for c in range(1, ws.max_column + 1)
        }

        def get_col(*names):
            for n in names:
                key = normalize_header(n)
                if key in raw_headers:
                    return raw_headers[key]
            raise KeyError(f"Missing column: {names}")

        col_german = get_col("German Mapping", "Deutsch Mapping")
        col_english = get_col("English Mapping", "Englisch Mapping")
        col_total = get_col("Total")
        col_group = get_col("Group")
        col_sign = get_col("Sign")

        review_col = ws.max_column + 1
        ws.cell(1, review_col, "Review")
        header_cell = ws.cell(1, review_col)
        header_cell.font = Font(name="GT Walsheim LC Light", size=8, bold=True)
        header_cell.fill = PatternFill(fill_type="solid", fgColor="F2F2F2")
        header_cell.border = Border(bottom=Side(style="thin"))

        fy_cols = {}
        for i, y in enumerate(years, start=1):
            fy_col = get_col(f"FY_{i}")
            header_name = (
                f"{fy_month_bs}{y % 100:02d}A"
                if sheet == "BS"
                else f"FY{y % 100:02d}A"
            )
            ws.cell(1, fy_col, header_name)
            fy_cols[y] = fy_col

        calculated_flags = {r: [] for r in range(2, ws.max_row + 1)}

        for r in range(2, ws.max_row + 1):
            german = ws.cell(r, col_german).value
            if not german:
                continue
            for y in years:
                entry = all_data[y][sheet].get(str(german))
                if entry:
                    ws.cell(r, fy_cols[y], float(entry["value"]))

        for r in range(2, ws.max_row + 1):
            eng = ws.cell(r, col_english).value
            total_flag = ws.cell(r, col_total).value
            group = ws.cell(r, col_group).value
            if str(total_flag).strip().lower() != "yes":
                continue
            for y in years:
                existing = ws.cell(r, fy_cols[y]).value
                if isinstance(existing, (int, float)):
                    continue
                values = []
                for rr in range(2, ws.max_row + 1):
                    if rr == r:
                        continue
                    child_total = ws.cell(rr, col_total).value
                    child_group = ws.cell(rr, col_group).value
                    if sheet == "PL" and group == "All":
                        if str(child_total).strip().lower() == "yes":
                            continue
                        child_level = ws.cell(rr, get_col("Level")).value
                        if child_level != "L4":
                            continue
                        v = ws.cell(rr, fy_cols[y]).value
                        if isinstance(v, (int, float)):
                            factor = sign_factor(ws.cell(rr, col_sign).value)
                            values.append(v * factor)
                    else:
                        if child_group == eng:
                            v = ws.cell(rr, fy_cols[y]).value
                            if isinstance(v, (int, float)):
                                values.append(v)
                if values:
                    ws.cell(r, fy_cols[y], sum(values))
                    calculated_flags[r].append(f"FY{str(y)[-2:]}")

        for r, years_list in calculated_flags.items():
            if years_list:
                ws.cell(r, review_col, "Calculated for " + ", ".join(years_list))

        last_data_row = 1
        for r in range(2, ws.max_row + 1):
            if ws.cell(r, col_german).value:
                last_data_row = r
        review_font = Font(name="GT Walsheim LC Light", size=8)
        white_fill = PatternFill(fill_type="solid", fgColor="FFFFFF")
        for r in range(2, last_data_row + 1):
            c = ws.cell(r, review_col)
            c.font = review_font
            c.fill = white_fill

    wb.save(output_excel)


def run(cfg: Dict[str, Any]) -> None:
    global CONFIG
    CONFIG = cfg
    pdf_dir = Path(str(cfg.get("pdf_input_dir") or ""))
    if not pdf_dir.is_dir():
        raise SystemExit(f"pdf_input_dir not found: {pdf_dir}")

    pdfs = sorted(pdf_dir.glob("*.pdf"))
    if not pdfs:
        raise SystemExit(f"No PDF files in {pdf_dir}")

    all_data: Dict[int, Dict] = {}
    for pdf in pdfs:
        print(f"Processing {pdf.name}")
        result = extract_year_data(pdf, cfg)
        year = int(result["year"])
        if year in all_data:
            print(f"Warning: duplicate year {year} from {pdf.name}, overwriting.")
        all_data[year] = result

    write_final_excel(all_data, cfg)
    print(f"Wrote {cfg['output_excel']}")


def main() -> None:
    if len(sys.argv) < 2:
        raise SystemExit("Usage: python PDF_Extraction_recon.py <config.json>")
    cfg = load_config(Path(sys.argv[1]))
    try:
        run(cfg)
    except ValueError as exc:
        raise SystemExit(str(exc)) from None


if __name__ == "__main__":
    main()
