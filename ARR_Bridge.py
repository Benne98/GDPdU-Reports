import os
import re
import textwrap
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from typing import Dict, Any, List, Tuple, Optional, Set

import numpy as np
import pandas as pd
import matplotlib.pyplot as plt

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image as XLImage
from matplotlib.patches import Rectangle
from matplotlib.colors import to_rgba
import matplotlib.font_manager as fm


def _xl_image_from_path(path: str) -> XLImage:
    """Embed PNG bytes so temp files can be deleted before workbook.save()."""
    return XLImage(BytesIO(Path(path).read_bytes()))

try:
    from funktionssammlung import apply_filters, compute_amount, recognized_value
except ImportError:  # pragma: no cover
    apply_filters = None  # type: ignore
    compute_amount = None  # type: ignore
    recognized_value = None  # type: ignore

try:
    from revenue_reconciliation import sales_basis_labels
except ImportError:  # pragma: no cover
    def sales_basis_labels(sales_basis=None):
        basis = "gross" if str(sales_basis or "").strip().lower() == "gross" else "net"
        metric = "Gross sales" if basis == "gross" else "Net sales"
        return type("L", (), {"sales_basis": basis, "metric_label": metric})()


# ----------------------------
# CONFIG
# ----------------------------
CONFIG: Dict[str, Any] = {
    "title": "Project Draft",
    "table": "ARR Component Bridge",
    "company": "Draft AG",
    "subtitle_suffix": "",
    "period_label_prefix": "XX",

    # Betrachtungszeitpunkt: Monatsende für "ARR" und auch als "as_of" für Periodenlogik
    "current_year": 2024,
    "current_month": 12,

    # FY Ende (NUR für calc_mode invoice/accrual relevant)
    "fy_end_month": 7,
    "fy_end_day": 31,

    # Periodenmodus (NUR für calc_mode invoice/accrual relevant): "FY" | "YTD" | "LTM"
    "period_mode": "FY",

    # Kalk-Modus: "ARR" | "invoice" | "accrual"
    "calc_mode": "ARR",

    # Input Columns
    "value_col": "Contract Value After Discount",
    "customer_col": "End Customer ID",
    "product_col": "Product Category",
    "start_col": "Contract Start Date",
    "end_col": "Contract End Date",
    "invoice_col": "Invoice Date",

    # Fast Track multi-rule filters (funktionssammlung.apply_filters)
    "filters": {"enabled": False, "rules": []},

    # Optional legacy recurring filter (off by default for Fast Track)
    "recurring_filter": {
        "enabled": False,
        "col": "Service Type",
        "values": {"Maintenance", "Subscription", "Managed Service"},
    },

    "invoice_mapping_mode": "date",  # "date" | "year" (period column; invoice/accrual only)
    "sales_basis": "net",

    # FX
    "apply_fx": False,
    "fx_col": "Functional FX Rate",

    # Grouping (1 Group-Col)
    "group_col": "Entity",
    "group_value_selection": "Total",

    # Bridge Darstellung
    "bridge": {
        "order": ["Upsell", "Downsell", "Cross-sell", "Lost", "New"],
        "top_n_per_step": 3,
        "missing_token": "__MISSING__",
        "combine_input_other": True,
        "input_other_tokens": {"Other"},
        "other_bucket_label": "Other",
        "sort_mode": "value",  # "abs" | "value"
    },

    # Plot
    "plot": {
        # Höhe fix, Breite dynamisch
        "fixed_height_in": 1.5,
        "bars_per_inch": 1.925,  # wie GP/NP Bridge
        "min_width_in": 0,
        "max_width_in": 100,

        # Rendering / Saving
        "render_dpi": 180,
        "save_dpi": 300,

        # Balken / Layout
        "bar_width": 0.85,
        "x_wrap_width": 10,
        "x_max_lines": 2,
        "xtick_rotation": 45,
        "show_values": True,
        "annotate_min_abs_keUR": 0.0,

        # Fonts (Zahlen + Achsen kleiner; Kasten-Labels unverändert via step_boxes)
        "font": {
            "base": 7,
            "title": 10,
            "x_ticks": 5.0,
            "value": 5.0,
        },

        # Dünne Linien/Striche
        "style": {
            "bar_edge_lw": 0.6,       # Balken-Rand (weiß)
            "connector_lw": 0.8,
            "connector_alpha": 0.35,
            "other_edge_lw": 0.3,
        },

        # Optional: wenn du irgendwann Titel brauchst
        "show_title": False,

        # Step-Boxes (bleibt inhaltlich wie vorher, nur weiterhin config-driven)
        "step_boxes": {
            "enabled": True,
            "pad_x_frac": 0.06,
            "pad_y_frac": 0.17,
            "linestyle": (0, (3, 2)),
            "linewidth": 0.9,
            "alpha": 0.45,
            "show_labels": True,
            "label_fontsize": 7,
            "label_pad_y_frac": 0.015,
            "label_ha": "left",
        },
    },

    # Output
    "output_dir": ".",
    "output_prefix": "arr_component_bridge",

    # Excel (JETZT wie Code 1: feste Anzeigegröße)
    "excel": {
        "enabled": True,
        "col_offset": 3,
        "sheet_name": "BRIDGE",
        "image_anchor_cell": "D6",
        "set_column_widths": True,

        "display_dpi": 180,
        "image_scale": 0.5,
    },

    "keep_png": False,
}

# ----------------------------
# FONT setup (shared theme — no local TTF paths)
# ----------------------------
from gst_excel_theme import THEME, apply_matplotlib_theme, matplotlib_font_properties

apply_matplotlib_theme(
    size=CONFIG.get("plot", {}).get("font", {}).get("base", 8),
    weight="light",
)
FONT_PROP = matplotlib_font_properties()


def _deep_merge_dict(base: Dict[str, Any], override: Dict[str, Any]) -> Dict[str, Any]:
    """Recursively merge override onto a copy of base (dicts only)."""
    out: Dict[str, Any] = dict(base or {})
    for k, v in (override or {}).items():
        if isinstance(v, dict) and isinstance(out.get(k), dict):
            out[k] = _deep_merge_dict(out[k], v)
        else:
            out[k] = v
    return out


def apply_config_defaults(cfg: Dict[str, Any]) -> Dict[str, Any]:
    """Ensure Fast Track partial cfgs still get CONFIG plot/excel defaults."""
    return _deep_merge_dict(CONFIG, cfg or {})


# ============================================================
BRAND_BLUE = "#1E3A5F"
BRAND_BLUE_LIGHT = "#C5D8EB"


def bridge_colors(_cfg: Optional[Dict[str, Any]] = None) -> Dict[str, str]:
    """Finssentials chart palette (no external theme JSON)."""
    return {
        "total": BRAND_BLUE,
        "pos": "#0AB6DF",
        "neg": "#F88281",
        "other": BRAND_BLUE_LIGHT,
        "neutral": "#666666",
    }


# ============================================================
# Period logic, NUR für invoice/accrual genutzt (1zu1)
# ============================================================
def safe_month_day_ts(year: int, month: int, day: int) -> pd.Timestamp:
    first = pd.Timestamp(year, month, 1)
    last  = first + pd.offsets.MonthEnd(0)
    d     = min(int(day), int(last.day))
    return pd.Timestamp(year, month, d)


def fiscal_year_bounds(fy_end_year: int, fy_end_m: int, fy_end_d: int) -> Tuple[pd.Timestamp, pd.Timestamp]:
    end   = safe_month_day_ts(fy_end_year, fy_end_m, fy_end_d)
    start = (end - pd.DateOffset(years=1)) + pd.Timedelta(days=1)
    return start, end


def compute_current_fy_end_year(as_of_end: pd.Timestamp, fy_end_m: int, fy_end_d: int) -> int:
    CY            = as_of_end.year
    fye_this_year = safe_month_day_ts(CY, fy_end_m, fy_end_d)
    return CY if as_of_end <= fye_this_year else CY + 1


def _fy_label(y: int) -> str:
    return f"FY{str(y)[-2:]}A"


def _fy_end_md(cfg: Dict[str, Any]) -> Tuple[int, int]:
    fy_end_m = int(cfg.get("fiscal_year_end_month", cfg.get("fy_end_month", 12)))
    fy_end_d = int(cfg.get("fiscal_year_end_day", cfg.get("fy_end_day", 31)))
    return fy_end_m, fy_end_d


def get_three_periods(cfg: Dict[str, Any]) -> Tuple[List[Tuple[pd.Timestamp, pd.Timestamp, str]], str]:
    """Same FY/YTD/LTM windows as GP_Bridge (chart: two steps; subtitle: first→last)."""
    CY = int(cfg["current_year"])
    m = int(cfg["current_month"])
    fy_end_m, fy_end_d = _fy_end_md(cfg)

    as_of_end = pd.Timestamp(CY, m, 1) + pd.offsets.MonthEnd(0)
    cur_fy_end_year = compute_current_fy_end_year(as_of_end, fy_end_m, fy_end_d)
    fye_this_calendar_year = safe_month_day_ts(as_of_end.year, fy_end_m, fy_end_d)
    at_fiscal_year_end = int(as_of_end.month) == int(fy_end_m)
    year_mapping = str(cfg.get("invoice_mapping_mode") or "date").strip().lower() == "year"

    mode = str(cfg.get("period_mode", "FY")).upper()

    if mode == "FY":
        last_completed = as_of_end.year if as_of_end >= fye_this_calendar_year else as_of_end.year - 1
        if at_fiscal_year_end or year_mapping:
            years = [last_completed - 2, last_completed - 1, last_completed]
            periods = [
                (*fiscal_year_bounds(y, fy_end_m, fy_end_d), _fy_label(y))
                for y in years
            ]
        else:
            older_fy = last_completed - 1
            prev_fy = last_completed
            ytd_start, _ = fiscal_year_bounds(cur_fy_end_year, fy_end_m, fy_end_d)
            periods = [
                (*fiscal_year_bounds(older_fy, fy_end_m, fy_end_d), _fy_label(older_fy)),
                (*fiscal_year_bounds(prev_fy, fy_end_m, fy_end_d), _fy_label(prev_fy)),
                (ytd_start, as_of_end, f"YTD{str(cur_fy_end_year)[-2:]}"),
            ]
        subtitle = f"{periods[0][2]} → {periods[2][2]}"
        return periods, subtitle

    if mode == "YTD":
        s2, _ = fiscal_year_bounds(cur_fy_end_year, fy_end_m, fy_end_d)
        e2 = as_of_end
        delta_days = int((e2 - s2).days)

        prev1 = cur_fy_end_year - 1
        s1, _ = fiscal_year_bounds(prev1, fy_end_m, fy_end_d)
        e1 = s1 + pd.Timedelta(days=delta_days)

        prev2 = cur_fy_end_year - 2
        s0, _ = fiscal_year_bounds(prev2, fy_end_m, fy_end_d)
        e0 = s0 + pd.Timedelta(days=delta_days)

        l2 = f"YTD{str(cur_fy_end_year)[-2:]}"
        l1 = f"YTD{str(prev1)[-2:]}"
        l0 = f"YTD{str(prev2)[-2:]}"

        periods = [(s0, e0, l0), (s1, e1, l1), (s2, e2, l2)]
        subtitle = f"{l0} → {l2}"
        return periods, subtitle

    if mode == "LTM":
        e2 = as_of_end
        s2 = e2 - pd.DateOffset(years=1) + pd.Timedelta(days=1)

        e1 = e2 - pd.DateOffset(years=1)
        s1 = e1 - pd.DateOffset(years=1) + pd.Timedelta(days=1)

        e0 = e2 - pd.DateOffset(years=2)
        s0 = e0 - pd.DateOffset(years=1) + pd.Timedelta(days=1)

        l2 = f"LTM {e2.month}-{str(e2.year)[-2:]}"
        l1 = f"LTM {e1.month}-{str(e1.year)[-2:]}"
        l0 = f"LTM {e0.month}-{str(e0.year)[-2:]}"

        periods = [(s0, e0, l0), (s1, e1, l1), (s2, e2, l2)]
        subtitle = f"{l0} → {l2}"
        return periods, subtitle

    raise ValueError("Unbekannter period_mode. Erlaubt: FY|YTD|LTM")


# ============================================================
# Gemeinsames Preprocessing (Filter/FX/Keys/Dates)
# ============================================================
def days_inclusive(start, end) -> pd.Series:
    """Inclusive day count; works for Series and DatetimeIndex (no .dt on TimedeltaIndex)."""
    start = pd.to_datetime(start)
    end = pd.to_datetime(end)
    days = (end - start) / np.timedelta64(1, "D")
    days = np.floor(days) + 1
    if isinstance(days, pd.Series):
        out = days
    else:
        idx = getattr(start, "index", None) or getattr(end, "index", None)
        out = pd.Series(days, index=idx)
    return out.clip(lower=1)


def preprocess_base(df: pd.DataFrame, cfg: Dict[str, Any]) -> pd.DataFrame:
    d = df.copy()
    d.columns = d.columns.str.strip()

    if apply_filters is not None:
        d = apply_filters(d, cfg)

    # optional recurring filter
    rf = cfg.get("recurring_filter") or {}
    if rf.get("enabled", False):
        col = rf.get("col")
        vals = rf.get("values")
        if col not in d.columns:
            raise ValueError(f"Recurring-Filter-Spalte fehlt: '{col}'")
        d = d[d[col].isin(set(vals))].copy()

    mapping = str(cfg.get("invoice_mapping_mode") or "date").strip().lower()
    cm = str(cfg.get("calc_mode") or "ARR").strip().lower()
    if mapping == "year" and cm == "arr":
        raise ValueError("invoice_mapping_mode='year' ist mit calc_mode='ARR' nicht unterstützt.")

    needs_contract = cm in {"arr", "accrual"}

    inv = cfg["invoice_col"]
    if inv not in d.columns:
        raise ValueError(f"Spalte fehlt im Input: '{inv}'")
    if mapping != "year":
        d[inv] = pd.to_datetime(d[inv], errors="coerce", dayfirst=True)

    if needs_contract:
        for c in [cfg["start_col"], cfg["end_col"]]:
            if c not in d.columns:
                raise ValueError(f"Spalte fehlt im Input: '{c}'")
            d[c] = pd.to_datetime(d[c], errors="coerce", dayfirst=True)

    # keys
    d["cust_id"] = d[cfg["customer_col"]].astype(str).str.strip()
    d["product_key"] = d[cfg["product_col"]].astype(str).str.strip()

    # value (align with GST: missing amounts → 0)
    d["base_value"] = pd.to_numeric(d[cfg["value_col"]], errors="coerce").fillna(0.0)

    # FX ( invalid * 1)
    if cfg.get("apply_fx", False):
        fx_col = cfg.get("fx_col")
        if not fx_col or fx_col not in d.columns:
            raise ValueError("apply_fx=True, aber fx_col fehlt oder existiert nicht im Input.")
        d["fx_rate"] = pd.to_numeric(d[fx_col], errors="coerce")
        d["fx_rate"] = d["fx_rate"].where(d["fx_rate"].notna() & (d["fx_rate"] != 0), 1.0)
        d.loc[:, "base_value"] = d["base_value"] * d["fx_rate"]

    # group
    gcol = cfg.get("group_col")
    if not gcol or gcol not in d.columns:
        raise ValueError(f"group_col fehlt oder existiert nicht: '{gcol}'")

    missing_token = str((cfg.get("bridge", {}) or {}).get("missing_token", "__MISSING__")).strip()
    x = d[gcol]
    x = x.astype(str).str.strip()
    x = x.replace({"": missing_token, "nan": missing_token, "None": missing_token, "NaT": missing_token})
    x = x.where(~x.isna(), missing_token)
    d["group_0"] = x

    # core drop NAs — contract dates only when ARR/accrual (same as GST/TOP/PVM)
    core_fields = [inv, "cust_id", "product_key", "group_0"]
    if needs_contract:
        core_fields = [inv, cfg["start_col"], cfg["end_col"], "cust_id", "product_key", "group_0"]
    for col in core_fields:
        d = d[d[col].notna()].copy()
    d = d[d["group_0"] != ""].copy()

    if needs_contract:
        d["contract_days"] = days_inclusive(d[cfg["start_col"]], d[cfg["end_col"]])
        d = d[d[cfg["end_col"]] >= d[cfg["start_col"]]].copy()

    return d


# ============================================================
# calc_mode = "ARR"
# ============================================================
def snapshot_arr(d: pd.DataFrame, as_of: pd.Timestamp, cfg: Dict[str, Any]) -> pd.DataFrame:
    start_col, end_col = cfg["start_col"], cfg["end_col"]
    inv_col = cfg["invoice_col"]

    month_start = as_of.replace(day=1)
    month_end   = as_of

    dd = d[d[inv_col].notna() & (d[inv_col] <= as_of)].copy()

    overlap_start = dd[start_col].where(dd[start_col] > month_start, month_start)
    overlap_end   = dd[end_col].where(dd[end_col] < month_end, month_end)

    overlap_days = _days_inclusive_generic(overlap_start, overlap_end)
    overlap_days = overlap_days.clip(lower=0)

    daily_rate = dd["base_value"] / dd["contract_days"]

    dd["mrr_month"] = daily_rate * overlap_days
    dd["metric"]    = dd["mrr_month"] * 12.0  # "arr"

    snap = (
        dd.loc[dd["metric"].notna() & (dd["metric"] != 0), ["group_0", "cust_id", "product_key", "metric"]]
        .groupby(["group_0", "cust_id", "product_key"], as_index=False)
        .agg(metric=("metric", "sum"))
    )

    return snap


# ============================================================
# calc_mode invoice/accrual -> Periodisierte Amounts
# ============================================================
def _days_inclusive_generic(a, b) -> pd.Series:
    a    = pd.to_datetime(a)
    b    = pd.to_datetime(b)
    days = (b - a) / np.timedelta64(1, "D")
    days = np.floor(days) + 1
    if isinstance(days, pd.Series):
        out = days
    else:
        idx = getattr(a, "index", None) or getattr(b, "index", None)
        out = pd.Series(days, index=idx)
    return out.clip(lower=0)


def recognized_amount(d: pd.DataFrame, period_start: pd.Timestamp, period_end: pd.Timestamp, cfg: Dict[str, Any]) -> pd.Series:
    """Same accrual as GST/TOP/PVM (`funktionssammlung.recognized_value`): contract ∩ period."""
    if recognized_value is not None:
        tmp = d.copy()
        tmp["value"] = pd.to_numeric(tmp["base_value"], errors="coerce").fillna(0.0)
        return recognized_value(tmp, period_start, period_end, cfg)

    S = pd.to_datetime(d[cfg["start_col"]], errors="coerce")
    E = pd.to_datetime(d[cfg["end_col"]], errors="coerce")
    aktiv_von = S.where(S > period_start, period_start)
    aktiv_bis = E.where(E < period_end, period_end)
    tage_realisiert = _days_inclusive_generic(aktiv_von, aktiv_bis)
    contract_days = _days_inclusive_generic(S, E).clip(lower=1)
    value = pd.to_numeric(d["base_value"], errors="coerce").fillna(0.0)
    out = value * (tage_realisiert / contract_days)
    return out.where((S <= period_end) & (E >= period_start), 0.0)


def compute_period_amount(d: pd.DataFrame, start: pd.Timestamp, end: pd.Timestamp, cfg: Dict[str, Any]) -> pd.Series:
    mode = str(cfg.get("calc_mode", "invoice")).lower()
    if mode == "invoice":
        if compute_amount is not None:
            return compute_amount(d, start, end, cfg, "base_value")
        mapping = str(cfg.get("invoice_mapping_mode") or "date").strip().lower()
        amt = pd.to_numeric(d["base_value"], errors="coerce").fillna(0.0)
        if mapping == "year":
            try:
                from funktionssammlung import parse_invoice_fy_year

                inv_y = parse_invoice_fy_year(d[cfg["invoice_col"]])
            except Exception:
                inv_y = pd.to_numeric(d[cfg["invoice_col"]], errors="coerce")
            return amt.where(inv_y == int(end.year), 0.0)
        invoice_dates = pd.to_datetime(d[cfg["invoice_col"]], errors="coerce", dayfirst=True)
        mask = (invoice_dates >= start) & (invoice_dates <= end)
        return amt.where(mask, 0.0)
    if compute_amount is not None:
        return compute_amount(d, start, end, cfg, "base_value")
    return recognized_amount(d, start, end, cfg)


def snapshot_period_amount(d: pd.DataFrame, period_start: pd.Timestamp, period_end: pd.Timestamp, cfg: Dict[str, Any]) -> pd.DataFrame:
    dd = d.copy()
    dd["metric"] = compute_period_amount(dd, period_start, period_end, cfg)

    snap = (
        dd.loc[dd["metric"].notna() & (dd["metric"] != 0), ["group_0", "cust_id", "product_key", "metric"]]
        .groupby(["group_0", "cust_id", "product_key"], as_index=False)
        .agg(metric=("metric", "sum"))
    )

    return snap


# ============================================================
# Bridge-Logik
# ============================================================
def compute_bridge_components(snap1: pd.DataFrame, snap2: pd.DataFrame) -> Dict[str, Any]:
    group_keys = ["group_0"]

    m1_g = snap1.groupby(group_keys, as_index=False)["metric"].sum().rename(columns={"metric": "M1"})
    m2_g = snap2.groupby(group_keys, as_index=False)["metric"].sum().rename(columns={"metric": "M2"})

    c1 = snap1.groupby(group_keys + ["cust_id"], as_index=False)["metric"].sum().rename(columns={"metric": "m1_cust"})
    c2 = snap2.groupby(group_keys + ["cust_id"], as_index=False)["metric"].sum().rename(columns={"metric": "m2_cust"})

    cust = c1.merge(c2, on=group_keys + ["cust_id"], how="outer")
    cust["m1_cust"] = cust["m1_cust"].fillna(0.0)
    cust["m2_cust"] = cust["m2_cust"].fillna(0.0)

    cust["is_new"]      = (cust["m1_cust"] == 0) & (cust["m2_cust"] != 0)
    cust["is_lost"]     = (cust["m1_cust"] != 0) & (cust["m2_cust"] == 0)
    cust["is_retained"] = (cust["m1_cust"] != 0) & (cust["m2_cust"] != 0)

    p = snap1.merge(
        snap2,
        on=group_keys + ["cust_id", "product_key"],
        how="outer",
        suffixes=("1", "2"),
    )

    p["m1"] = pd.to_numeric(p["metric1"], errors="coerce").fillna(0.0)
    p["m2"] = pd.to_numeric(p["metric2"], errors="coerce").fillna(0.0)

    retained = cust.loc[cust["is_retained"], group_keys + ["cust_id"]].copy()
    retained["retained"] = True
    p = p.merge(retained, on=group_keys + ["cust_id"], how="left")
    p["retained"] = p["retained"].fillna(False)

    cross        = p[(p["retained"]) & (p["m1"] == 0)]
    cross_detail = cross.groupby(group_keys + ["product_key"], as_index=False)["m2"].sum().rename(columns={"m2": "value"})

    both = p[(p["retained"]) & (p["m1"] != 0)].copy()
    both["delta"]          = both["m2"] - both["m1"]
    both["upsell_delta"]   = both["delta"].clip(lower=0.0)
    both["downsell_delta"] = (-both["delta"]).clip(lower=0.0)

    upsell_detail   = both.groupby(group_keys + ["product_key"], as_index=False)["upsell_delta"].sum().rename(columns={"upsell_delta": "value"})
    downsell_detail = both.groupby(group_keys + ["product_key"], as_index=False)["downsell_delta"].sum().rename(columns={"downsell_delta": "value"})

    new_customers  = cust.loc[cust["is_new"],  group_keys + ["cust_id"]].copy()
    lost_customers = cust.loc[cust["is_lost"], group_keys + ["cust_id"]].copy()

    p_new = p.merge(new_customers.assign(is_new=True), on=group_keys + ["cust_id"], how="left")
    p_new["is_new"] = p_new["is_new"].fillna(False)
    new_detail = (
        p_new.loc[p_new["is_new"]]
        .groupby(group_keys + ["product_key"], as_index=False)["m2"]
        .sum()
        .rename(columns={"m2": "value"})
    )

    p_lost = p.merge(lost_customers.assign(is_lost=True), on=group_keys + ["cust_id"], how="left")
    p_lost["is_lost"] = p_lost["is_lost"].fillna(False)
    lost_detail = (
        p_lost.loc[p_lost["is_lost"]]
        .groupby(group_keys + ["product_key"], as_index=False)["m1"]
        .sum()
        .rename(columns={"m1": "value"})
    )

    def sum_detail(df_detail: pd.DataFrame) -> pd.DataFrame:
        return df_detail.groupby(group_keys, as_index=False)["value"].sum()

    cross_g    = sum_detail(cross_detail).rename(columns={"value": "Cross-sell"})
    upsell_g   = sum_detail(upsell_detail).rename(columns={"value": "Upsell"})
    downsell_g = sum_detail(downsell_detail).rename(columns={"value": "Downsell"})
    new_g      = sum_detail(new_detail).rename(columns={"value": "New"})
    lost_g     = sum_detail(lost_detail).rename(columns={"value": "Lost"})

    out = m1_g.merge(m2_g, on=group_keys, how="outer")
    for t in [upsell_g, downsell_g, cross_g, lost_g, new_g]:
        out = out.merge(t, on=group_keys, how="left")

    for c in ["M1", "M2", "Upsell", "Downsell", "Cross-sell", "Lost", "New"]:
        if c in out.columns:
            out[c] = out[c].fillna(0.0)

    out["NRR"] = (out["M1"] - out["Lost"] + out["Upsell"] + out["Cross-sell"] - out["Downsell"])

    return {
        "totals_by_group": out,
        "detail": {
            "Upsell":     upsell_detail,
            "Downsell":   downsell_detail,
            "Cross-sell": cross_detail,
            "Lost":       lost_detail,
            "New":        new_detail,
        }
    }


# ============================================================
# Waterfall Darstellung
# ============================================================
@dataclass
class BridgeItem:
    label: str
    value_keUR: float
    is_total: bool = False
    is_other: bool = False
    step: str = ""
    dim_label: str = ""


def fmt_keUR(x: float) -> str:
    v = int(round(x))
    if v < 0:
        return f"({abs(v):,})"
    return f"{v:,}"


def wrap_xtick(s: str, width: int, max_lines: int) -> str:
    s = str(s)
    lines = textwrap.wrap(s, width=width)
    if len(lines) > max_lines:
        lines = lines[:max_lines]
        lines[-1] = lines[-1].rstrip(".") + "..."
    return "\n".join(lines)


def _clean_dim_series(s: pd.Series, missing_token: str) -> pd.Series:
    x = s.astype(str).str.strip()
    x = x.replace({"": missing_token, "nan": missing_token, "None": missing_token, "NaT": missing_token})
    x = x.where(~s.isna(), missing_token)
    return x


def _sort_steps_df(m: pd.DataFrame, sort_mode: str) -> pd.DataFrame:
    if m.empty:
        return m
    if sort_mode == "abs":
        return m.assign(abs_delta=m["value"].abs()).sort_values("abs_delta", ascending=False).drop(columns=["abs_delta"])
    if sort_mode == "value":
        return m.sort_values("value", ascending=False)
    mm = m.copy()
    mm["is_neg"]    = mm["value"] < 0
    mm["abs_delta"] = mm["value"].abs()
    mm = mm.sort_values(["is_neg", "abs_delta"], ascending=[True, False]).drop(columns=["is_neg", "abs_delta"])
    return mm


def build_items_step_blocks_by_entity(
    totals_by_group: pd.DataFrame,
    cfg: Dict[str, Any],
    label_left: str,
    label_right: str,
    total_left: float,
    total_right: float,
    entity_sort_by: str = "M2",
) -> List[BridgeItem]:
    order = cfg["bridge"]["order"]

    ent = totals_by_group.copy()
    ent["group_0"] = ent["group_0"].astype(str)
    ent = ent[ent["group_0"].str.lower() != "total"].copy()

    if entity_sort_by == "name":
        ent = ent.sort_values("group_0", ascending=True)
    else:
        ent["_sort"] = pd.to_numeric(ent.get(entity_sort_by, 0.0), errors="coerce").fillna(0.0)
        ent = ent.sort_values("_sort", ascending=False).drop(columns=["_sort"], errors="ignore")

    entities = ent["group_0"].tolist()

    items: List[BridgeItem] = [
        BridgeItem(label=label_left, dim_label=label_left, value_keUR=total_left / 1000.0, is_total=True, step="Total")
    ]

    top_n                = int(cfg["bridge"].get("top_n_per_step", 1))
    combine_input_other  = bool(cfg["bridge"].get("combine_input_other", True))
    input_other_tokens   = set(cfg["bridge"].get("input_other_tokens") or set())
    missing_token        = str(cfg["bridge"].get("missing_token", "__MISSING__"))
    other_bucket_label   = str(cfg["bridge"].get("other_bucket_label", "Other (Rest)"))

    sort_mode = str(cfg["bridge"].get("sort_mode", "value")).lower()  # "abs"|"value"

    for step in order:
        # Step-Werte je Entity ziehen
        rows = []
        for g in entities:
            row = ent.loc[ent["group_0"] == g]
            if row.empty:
                continue
            v = float(row.iloc[0].get(step, 0.0))
            if abs(v) < 1e-12:
                continue
            rows.append((g, v))

        if not rows:
            continue

        # Input-"Other" ggf. direkt in Other-Bucket schieben
        forced_other_sum = 0.0
        if combine_input_other and input_other_tokens:
            kept = []
            for g, v in rows:
                if str(g).strip() in input_other_tokens:
                    forced_other_sum += float(v)
                else:
                    kept.append((g, float(v)))
            rows = kept
        else:
            # konsistent casten
            rows = [(str(g).strip(), float(v)) for g, v in rows]

        # Falls nur "Other" im Input war (und sonst nichts), trotzdem weiter:
        if not rows and abs(forced_other_sum) < 1e-12:
            continue

        # Sortierung
        if sort_mode == "abs":
            rows.sort(key=lambda t: abs(t[1]), reverse=True)
        else:
            rows.sort(key=lambda t: t[1], reverse=True)

        # Top N
        top_rows  = rows[:top_n]
        rest_rows = rows[top_n:]

        # Sign-Logik wie bei dir
        sign = -1.0 if step in {"Downsell", "Lost"} else 1.0

        # Top hinzufügen
        for g, v in top_rows:
            items.append(
                BridgeItem(
                    label=f"{step} | {g}",
                    dim_label=g,
                    step=step,
                    value_keUR=sign * (v / 1000.0),
                    is_total=False,
                    is_other=False
                )
            )

        # Rest → Other
        rest_sum = sum(v for _, v in rest_rows) + forced_other_sum
        if abs(rest_sum) > 1e-12:
            items.append(
                BridgeItem(
                    label=f"{step} | {other_bucket_label}",
                    dim_label=other_bucket_label,
                    step=step,
                    value_keUR=sign * (rest_sum / 1000.0),
                    is_total=False,
                    is_other=True
                )
            )

    items.append(
        BridgeItem(label=label_right, dim_label=label_right, value_keUR=total_right / 1000.0, is_total=True, step="Total")
    )
    return items


# ----------------------------
# DYNAMIC FIGSIZE
# ----------------------------
def compute_dynamic_figsize(cfg: Dict[str, Any], n_bars: int) -> Tuple[float, float]:
    pcfg = cfg.get("plot", {}) or {}
    h    = float(pcfg.get("fixed_height_in", 3.4))

    bars_per_inch = float(pcfg.get("bars_per_inch", 2.2))
    if bars_per_inch <= 0:
        raise ValueError("plot.bars_per_inch muss > 0 sein.")

    w = n_bars / bars_per_inch

    min_w = pcfg.get("min_width_in", None)
    max_w = pcfg.get("max_width_in", None)
    if min_w is not None:
        w = max(float(min_w), w)
    if max_w is not None:
        w = min(float(max_w), w)

    return w, h


def draw_step_segment_boxes(
    ax,
    items,
    x,
    width,
    bottoms,
    heights,
    cfg,
    value_label_ys: Optional[Dict[str, Tuple[float, float]]] = None,
):
    sb_cfg = cfg.get("plot", {}).get("step_boxes", {})
    if not sb_cfg.get("enabled", True):
        return

    pad_x_frac      = float(sb_cfg.get("pad_x_frac", 0.25))
    pad_y_frac      = float(sb_cfg.get("pad_y_frac", 0.06))
    linestyle       = sb_cfg.get("linestyle", (0, (3, 2)))
    linewidth       = float(sb_cfg.get("linewidth", 0.9))
    alpha           = float(sb_cfg.get("alpha", 0.45))

    show_labels      = bool(sb_cfg.get("show_labels", True))
    label_fs         = int(sb_cfg.get("label_fontsize", 7))
    label_pad_y_frac = float(sb_cfg.get("label_pad_y_frac", 0.015))
    label_ha         = str(sb_cfg.get("label_ha", "left"))

    idx_by_step = {}
    for i, it in enumerate(items):
        if getattr(it, "is_total", False):
            continue
        step = (getattr(it, "step", "") or "").strip()
        if not step:
            continue
        idx_by_step.setdefault(step, []).append(i)

    y0_lim, y1_lim = ax.get_ylim()
    y_span      = (y1_lim - y0_lim + 1e-9)
    # Cap pad so oversized pad_y_frac cannot invent huge empty axis space
    pad_y       = min(pad_y_frac * y_span, 0.12 * y_span)
    pad_x       = pad_x_frac * width
    label_pad_y = label_pad_y_frac * y_span
    value_label_ys = value_label_ys or {}

    min_needed = y0_lim
    max_needed = y1_lim

    for step, idxs in idx_by_step.items():
        if not idxs:
            continue

        i0, i1 = min(idxs), max(idxs)

        left  = (x[i0] - width/2) - pad_x
        right = (x[i1] + width/2) + pad_x
        w     = right - left

        y_low  = min(bottoms[i] for i in idxs)
        y_high = max((bottoms[i] + heights[i]) for i in idxs)
        # Expand box to enclose value labels (avoids dashed edge cutting through numbers)
        if step in value_label_ys:
            ly0, ly1 = value_label_ys[step]
            y_low = min(y_low, ly0)
            y_high = max(y_high, ly1)
        if np.isclose(y_low, y_high):
            continue

        y_rect = y_low - pad_y
        h      = (y_high - y_low) + 2 * pad_y
        y_top  = y_rect + h

        rect = Rectangle(
            (left, y_rect),
            w, h,
            fill=False,
            linestyle=linestyle,
            linewidth=linewidth,
            edgecolor="black",
            alpha=alpha,
            zorder=5,
            clip_on=False,
        )
        ax.add_patch(rect)

        min_needed = min(min_needed, y_rect)
        max_needed = max(max_needed, y_top)

        if show_labels:
            if label_ha == "center":
                tx = left + w / 2
            elif label_ha == "right":
                tx = left + w
            else:
                tx = left
            ty = y_top + label_pad_y
            max_needed = max(max_needed, ty)

            ax.text(
                tx, ty, step,
                ha=label_ha, va="bottom",
                fontsize=label_fs, color="black",
                fontweight="light",
                zorder=6,
                bbox=dict(facecolor="white", edgecolor="none", alpha=0.75, pad=1.2),
                clip_on=False,
            )

    if (max_needed > y1_lim) or (min_needed < y0_lim):
        ax.set_ylim(min_needed, max_needed)


def plot_waterfall(items: List[BridgeItem], cfg: Dict[str, Any], subtitle: str) -> plt.Figure:
    plot_cfg = cfg.get("plot", {}) or {}

    n              = len(items)
    fig_w, fig_h   = compute_dynamic_figsize(cfg, n_bars=n)
    dpi            = int(plot_cfg.get("render_dpi", plot_cfg.get("dpi", 180)))

    show_values = bool(plot_cfg.get("show_values", True))
    min_abs     = float(plot_cfg.get("annotate_min_abs_keUR", 0.0))
    wrap_w      = int(plot_cfg.get("x_wrap_width", 12))
    max_lines   = int(plot_cfg.get("x_max_lines", 2))

    bar_width = float(plot_cfg.get("bar_width", 0.68))

    fcfg     = plot_cfg.get("font", {}) or {}
    xt_fs    = float(fcfg.get("x_ticks", 5.5))
    val_fs   = float(fcfg.get("value", 5.5))
    title_fs = float(fcfg.get("title", 10))
    xtick_rotation = float(plot_cfg.get("xtick_rotation", 45 if n >= 4 else 0))

    style_cfg       = plot_cfg.get("style", {}) or {}
    bar_edge_lw     = float(style_cfg.get("bar_edge_lw", 0.6))
    connector_lw    = float(style_cfg.get("connector_lw", 0.8))
    connector_alpha = float(style_cfg.get("connector_alpha", 0.35))
    other_edge_lw   = float(style_cfg.get("other_edge_lw", 0.9))

    colors = bridge_colors(cfg)
    total_color, pos_color, neg_color = colors["total"], colors["pos"], colors["neg"]
    other_color = colors.get("other", BRAND_BLUE_LIGHT)

    fig, ax = plt.subplots(figsize=(fig_w, fig_h), dpi=dpi)

    # More bottom room when labels are rotated
    bottom_margin = 0.40 if xtick_rotation else 0.28
    fig.subplots_adjust(left=0.03, right=0.995, top=0.92, bottom=bottom_margin)

    labels_raw = [(it.dim_label or it.label) for it in items]
    labels     = [wrap_xtick(lab, width=wrap_w, max_lines=max_lines) for lab in labels_raw]

    x = np.arange(len(items))

    running = 0.0
    bottoms, heights, facecolors, is_other_flags = [], [], [], []

    for it in items:
        if it.is_total:
            btm = 0.0 if it.value_keUR >= 0 else it.value_keUR
            bottoms.append(btm)
            heights.append(abs(it.value_keUR))
            facecolors.append(total_color)
            is_other_flags.append(False)
            running = it.value_keUR
        else:
            delta = it.value_keUR
            btm   = running if delta >= 0 else running + delta
            bottoms.append(btm)
            heights.append(abs(delta))

            if it.is_other:
                facecolors.append(other_color)
                is_other_flags.append(True)
            else:
                facecolors.append(pos_color if delta >= 0 else neg_color)
                is_other_flags.append(False)

            running += delta

    bars = ax.bar(
        x, heights, bar_width,
        bottom=bottoms,
        color=facecolors,
        edgecolor="white",
        linewidth=bar_edge_lw
    )

    for rect, is_other in zip(bars, is_other_flags):
        if is_other:
            rect.set_facecolor(to_rgba(other_color, 0.45))
            rect.set_edgecolor(total_color)
            rect.set_linewidth(other_edge_lw)
            rect.set_hatch("////")

    running = 0.0
    for i, it in enumerate(items[:-1]):
        if it.is_total:
            running = it.value_keUR
        else:
            running += it.value_keUR

        ax.plot(
            [x[i] + bar_width/2, x[i+1] - bar_width/2],
            [running, running],
            linewidth=connector_lw,
            color=BRAND_BLUE,
            alpha=connector_alpha
        )

    ax.relim()
    ax.autoscale_view()

    # Place value labels first; collect step extents so boxes enclose bars + numbers
    value_label_ys: Dict[str, Tuple[float, float]] = {}
    if show_values:
        yr = (ax.get_ylim()[1] - ax.get_ylim()[0] + 1e-9)
        dy = 0.02 * yr

        for i, it in enumerate(items):
            val = it.value_keUR
            if abs(val) < min_abs:
                continue

            if it.is_total:
                y0       = bottoms[i]
                h        = heights[i]
                y_center = y0 + h/2.0
                ax.text(
                    x[i], y_center, fmt_keUR(val),
                    ha="center", va="center",
                    fontsize=val_fs, color="white",
                    fontweight="light",
                    zorder=10,
                    bbox=dict(
                        facecolor=total_color,
                        edgecolor="none",
                        alpha=1.0,
                        boxstyle="round,pad=0.20",
                    ),
                    clip_on=False,
                )
            else:
                delta = it.value_keUR
                if delta >= 0:
                    y  = bottoms[i] + heights[i] + dy
                    va = "bottom"
                else:
                    y  = bottoms[i] - dy
                    va = "top"

                step = (getattr(it, "step", "") or "").strip()
                if step:
                    lo, hi = value_label_ys.get(step, (y, y))
                    value_label_ys[step] = (min(lo, y), max(hi, y))

                if it.is_other:
                    ax.text(
                        x[i], y, fmt_keUR(delta),
                        ha="center", va=va,
                        fontsize=val_fs, color=BRAND_BLUE,
                        fontweight="light",
                        zorder=10,
                        bbox=dict(facecolor="white", edgecolor="none", alpha=0.75, boxstyle="round,pad=0.18")
                    )
                else:
                    ax.text(
                        x[i], y, fmt_keUR(delta),
                        ha="center", va=va,
                        fontsize=val_fs, color=BRAND_BLUE,
                        fontweight="light",
                        zorder=10,
                    )

    draw_step_segment_boxes(
        ax, items, x, bar_width, bottoms, heights, cfg, value_label_ys=value_label_ys
    )

    ax.set_xticks(x)
    ax.set_xticklabels(
        labels,
        fontproperties=FONT_PROP,
        rotation=xtick_rotation,
        ha="right" if xtick_rotation else "center",
        rotation_mode="anchor" if xtick_rotation else "default",
    )
    for tick in ax.get_xticklabels():
        tick.set_fontsize(xt_fs)
        tick.set_fontweight("light")
        tick.set_color(BRAND_BLUE)
        if xtick_rotation:
            tick.set_rotation(xtick_rotation)
            tick.set_ha("right")
            tick.set_rotation_mode("anchor")
    # Small mid-bar axis ticks for orientation
    ax.tick_params(axis="x", pad=2, length=3.5, width=0.7, direction="out", color=BRAND_BLUE, bottom=True)

    ax.set_yticks([])
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    ax.spines["left"].set_visible(False)
    ax.spines["bottom"].set_visible(True)
    ax.spines["bottom"].set_color(BRAND_BLUE)
    ax.spines["bottom"].set_linewidth(0.6)
    ax.grid(False)

    if bool(plot_cfg.get("show_title", False)):
        company = str(cfg.get("company", "")).strip()
        title   = str(cfg.get("table", "")).strip()
        prefix  = str(cfg.get("period_label_prefix", "")).strip()
        parts   = [p for p in [company, prefix, title, subtitle] if p]
        ax.set_title(" | ".join(parts), fontsize=title_fs, fontweight="light", fontproperties=FONT_PROP)

    return fig


# ============================================================
# Output + Excel
# ============================================================
def clean_filename(s: str) -> str:
    s = str(s).strip().replace(" ", "_")
    return re.sub(r"[^\w\-_.]", "", s)


def build_output_paths(cfg: Dict[str, Any], suffix: str) -> Tuple[str, str]:
    out_dir = cfg.get("output_dir", ".")
    os.makedirs(out_dir, exist_ok=True)

    prefix       = clean_filename(cfg.get("output_prefix", "arr_component_bridge"))
    group_string = clean_filename(cfg.get("group_col"))
    cm           = str(cfg.get("calc_mode", "ARR")).lower()

    y = int(cfg["current_year"])
    m = int(cfg["current_month"])

    pm   = str(cfg.get("period_mode", "NA")) if cm in {"invoice", "accrual"} else "NA"
    gsel = clean_filename(str(cfg.get("group_value_selection", "Total")))

    base      = f"{prefix}_{group_string}_{clean_filename(suffix)}_{pm}_cy{str(y)[-2:]}m{m:02d}_{cm}_{gsel}"
    png_path  = os.path.join(out_dir, base + ".png")
    xlsx_path = os.path.join(out_dir, base + ".xlsx")
    return png_path, xlsx_path


def _write_arr_bridge_sheet(
    wb,
    cfg: Dict[str, Any],
    png_path: str,
    subtitle: str,
    fig_size_in: Tuple[float, float],
    sheet_name: str,
) -> str:
    from funktionssammlung import replace_workbook_sheet

    excel_cfg = cfg.get("excel", {}) or {}
    col_offset = int(excel_cfg.get("col_offset", 3))
    anchor_cell = excel_cfg.get("image_anchor_cell", "D6")
    set_widths = bool(excel_cfg.get("set_column_widths", True))

    title = str(sheet_name or excel_cfg.get("sheet_name") or "ARR Bridge")[:31]
    ws = replace_workbook_sheet(wb, title)

    brand = BRAND_BLUE.replace("#", "FF")
    title_font = Font(name=THEME.font_name, size=24, color=brand)
    sub_font = Font(name=THEME.font_name, size=12, color=brand)
    info_font = Font(name=THEME.font_name, size=9, color=brand, bold=True)

    ws["D1"].value = cfg.get("title", "")
    ws["D1"].font = title_font

    suffix = str(cfg.get("subtitle_suffix", "")).strip()
    suffix_txt = f" ({suffix})" if suffix else ""
    table_name = f"{cfg.get('table','')}{suffix_txt}"

    group_col = cfg.get("group_col", "")
    group_val = cfg.get("group_value_selection", "Total")

    cm = str(cfg.get("calc_mode", "ARR"))
    pm = str(cfg.get("period_mode", "NA")) if str(cfg.get("calc_mode", "")).lower() in {"invoice", "accrual"} else "NA"

    ws["D2"].value = f"{table_name} | {subtitle} | {cm} | period={pm} | {group_col}: {group_val}"
    ws["D2"].font = sub_font

    ws["D4"].value = f"{cfg.get('company','')} | {table_name}"
    ws["D4"].font = info_font

    fill_grey = PatternFill(fill_type="solid", fgColor="FFF3F1EF")
    for r in range(1, 200):
        for c in range(1, 1 + col_offset):
            ws.cell(row=r, column=c).fill = fill_grey

    fill_white = PatternFill(fill_type="solid", fgColor="FFFFFFFF")
    for r in range(1, 200):
        for c in range(1 + col_offset, 30):
            ws.cell(row=r, column=c).fill = fill_white

    if os.path.exists(png_path):
        img = _xl_image_from_path(png_path)
        fig_w, fig_h = fig_size_in
        display_dpi = int(excel_cfg.get("display_dpi", 180))
        scale = float(excel_cfg.get("image_scale", 0.5))
        img.width = int(fig_w * display_dpi * scale)
        img.height = int(fig_h * display_dpi * scale)
        ws.add_image(img, anchor_cell)
    else:
        ws[anchor_cell].value = f"Bild nicht gefunden: {png_path}"

    if set_widths:
        for c in range(1, 30):
            ws.column_dimensions[get_column_letter(c)].width = 14
        ws.column_dimensions["A"].width = 3
        ws.column_dimensions["B"].width = 3
        ws.column_dimensions["C"].width = 3
        ws.column_dimensions["D"].width = 20

    ws.sheet_properties.outlinePr.summaryRight = True
    for col_letter in ["A", "B", "C"]:
        ws.column_dimensions[col_letter].outlineLevel = 1
        ws.column_dimensions[col_letter].hidden = True
    ws.column_dimensions["D"].collapsed = True
    return ws.title


def export_excel_with_image(
    cfg: Dict[str, Any],
    png_path: str,
    xlsx_path: str,
    subtitle: str,
    fig_size_in: Tuple[float, float],
    wb=None,
    *,
    save: bool = True,
    sheet_name: Optional[str] = None,
) -> Optional[str]:
    excel_cfg = cfg.get("excel", {}) or {}
    if not excel_cfg.get("enabled", True):
        return None

    target = sheet_name or excel_cfg.get("sheet_name") or cfg.get("base_sheet_name") or "ARR Bridge"
    owns_wb = wb is None
    if owns_wb:
        wb = Workbook()
        default = wb.active
        if default is not None and default.title == "Sheet":
            wb.remove(default)

    written = _write_arr_bridge_sheet(wb, cfg, png_path, subtitle, fig_size_in, str(target))
    if save and xlsx_path:
        wb.save(xlsx_path)
    return written


# ============================================================
# Endpoint-Definition (3 Zeitpunkte => 2 Bridges)
# ============================================================
def get_three_endpoints_for_calc(cfg: Dict[str, Any]) -> Tuple[List[Tuple[Any, str]], str]:
    cm = str(cfg.get("calc_mode", "ARR")).lower()

    if cm == "arr":
        CY = int(cfg["current_year"])
        m  = int(cfg["current_month"])
        e2 = pd.Timestamp(CY, m, 1) + pd.offsets.MonthEnd(0)
        e1 = e2 - pd.DateOffset(years=1)
        e0 = e2 - pd.DateOffset(years=2)

        l2 = f"FY{str(e2.year)[-2:]}A"
        l1 = f"FY{str(e1.year)[-2:]}A"
        l0 = f"FY{str(e0.year)[-2:]}A"
        subtitle = f"{l0} → {l2}"
        return [(e0, l0), (e1, l1), (e2, l2)], subtitle

    periods, subtitle = get_three_periods(cfg)
    return periods, subtitle


# ============================================================
# MAIN Runner
# ============================================================
def _render_arr_bridge_assets(df: pd.DataFrame, cfg: Dict[str, Any]):
    cfg = apply_config_defaults(cfg)
    d = preprocess_base(df, cfg)

    gsel = str(cfg.get("group_value_selection", "Total"))
    cm = str(cfg.get("calc_mode", "ARR")).lower()

    endpoints, subtitle = get_three_endpoints_for_calc(cfg)

    if cm == "arr":
        (e0, l0), (e1, l1), (e2, l2) = endpoints  # type: ignore
        s1 = snapshot_arr(d, e1, cfg)
        s2 = snapshot_arr(d, e2, cfg)
        label_left, label_right = l1, l2
    else:
        (s0s, s0e, l0), (s1s, s1e, l1), (s2s, s2e, l2) = endpoints  # type: ignore
        s1 = snapshot_period_amount(d, s1s, s1e, cfg)
        s2 = snapshot_period_amount(d, s2s, s2e, cfg)
        label_left, label_right = l1, l2

    comp = compute_bridge_components(s1, s2)
    totals = comp["totals_by_group"]
    if gsel.lower() == "total":
        total_left = float(pd.to_numeric(totals["M1"], errors="coerce").fillna(0.0).sum())
        total_right = float(pd.to_numeric(totals["M2"], errors="coerce").fillna(0.0).sum())
    else:
        raise ValueError("group_value_selection muss aktuell 'Total' sein.")

    items = build_items_step_blocks_by_entity(
        totals_by_group=totals,
        cfg=cfg,
        label_left=label_left,
        label_right=label_right,
        total_left=total_left,
        total_right=total_right,
        entity_sort_by="M2",
    )

    fig_w, fig_h = compute_dynamic_figsize(cfg, n_bars=len(items))
    fig = plot_waterfall(items, cfg, subtitle)
    png_path, xlsx_path = build_output_paths(cfg, subtitle)
    save_dpi = int(cfg.get("plot", {}).get("save_dpi", 300))
    fig.savefig(png_path, dpi=save_dpi, bbox_inches=None)
    plt.close(fig)
    return cfg, subtitle, png_path, xlsx_path, (fig_w, fig_h)


def run_component_bridge(df: pd.DataFrame, cfg: Dict[str, Any]) -> Tuple[Optional[str], Optional[str]]:
    cfg, subtitle, png_path, xlsx_path, fig_size = _render_arr_bridge_assets(df, cfg)
    excel_out = None
    try:
        if cfg.get("excel", {}).get("enabled", True):
            export_excel_with_image(cfg, png_path, xlsx_path, subtitle, fig_size_in=fig_size)
            excel_out = xlsx_path
    finally:
        if not cfg.get("keep_png", False):
            try:
                if os.path.exists(png_path):
                    os.remove(png_path)
            except Exception as e:
                print("Warnung: Konnte temporäres PNG nicht löschen:", e)
    return None, excel_out


def run_arr_bridge_in_workbook(cfg: dict, wb) -> str:
    """Fast Track / session mode: write ARR Bridge into an open workbook."""
    file_path = str(cfg.get("file_path") or "").strip()
    sheet_name = str(cfg.get("sheet_name") or "").strip()
    if not file_path:
        raise ValueError("ARR_Bridge requires file_path")
    df = pd.read_excel(file_path, sheet_name=sheet_name or 0, engine="openpyxl")

    out_dir = str(cfg.get("output_file_path") or cfg.get("output_dir") or ".")
    cfg = {**cfg, "output_dir": out_dir, "output_prefix": cfg.get("output_prefix") or "arr_component_bridge"}
    excel_cfg = dict(cfg.get("excel") or {})
    excel_cfg.setdefault("sheet_name", cfg.get("base_sheet_name") or "ARR Bridge")
    excel_cfg["enabled"] = True
    cfg["excel"] = excel_cfg

    # Align FY keys from Fast Track common config
    if "fy_end_month" in cfg and "fiscal_year_end_month" not in cfg:
        cfg["fiscal_year_end_month"] = cfg["fy_end_month"]
    if "fy_end_day" in cfg and "fiscal_year_end_day" not in cfg:
        cfg["fiscal_year_end_day"] = cfg["fy_end_day"]

    cfg, subtitle, png_path, _xlsx, fig_size = _render_arr_bridge_assets(df, cfg)
    sheet_title = str(
        (cfg.get("excel") or {}).get("sheet_name")
        or cfg.get("base_sheet_name")
        or "ARR Bridge"
    )
    try:
        from funktionssammlung import get_next_sheet_name_from_wb

        target = get_next_sheet_name_from_wb(wb, sheet_title)
    except Exception:
        target = sheet_title
        if target in wb.sheetnames:
            target = f"{target}_1"

    try:
        written = export_excel_with_image(
            cfg,
            png_path,
            xlsx_path="",
            subtitle=subtitle,
            fig_size_in=fig_size,
            wb=wb,
            save=False,
            sheet_name=target,
        )
        return str(written or target)
    finally:
        try:
            if os.path.exists(png_path):
                os.remove(png_path)
        except Exception:
            pass


# ============================================================
# Beispiel MAIN
# ============================================================
if __name__ == "__main__":
    _desktop = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Desktop")
    _src = os.path.join(_desktop, "Sales data", "SaaS-Sales.xlsx")
    if not os.path.isfile(_src):
        raise SystemExit(f"Demo file not found: {_src}")

    _cfg = dict(CONFIG)
    _cfg.update(
        {
            "calc_mode": "accrual",
            "period_mode": "FY",
            "current_year": 2025,
            "current_month": 7,
            "fy_end_month": 12,
            "fy_end_day": 31,
            "value_col": "Contract Value After Discount",
            "customer_col": "End Customer ID",
            "product_col": "Product Name",
            "start_col": "Contract Start Date",
            "end_col": "Contract End Date",
            "invoice_col": "Invoice Date",
            "group_col": "Entity",
            "group_value_selection": "Total",
            "apply_fx": False,
            "recurring_filter": {
                "enabled": True,
                "col": "Product Revenue Model",
                "values": {"Recurring"},
            },
            "table": "Revenue Component Bridge",
            "subtitle_suffix": "by accrued amounts",
            "output_dir": _desktop if os.path.isdir(_desktop) else ".",
            "output_prefix": "arr_component_bridge",
            "excel": {
                **(CONFIG.get("excel") or {}),
                "enabled": True,
                "sheet_name": "ARR Bridge",
            },
        }
    )
    _df = pd.read_excel(_src, sheet_name="Database", engine="openpyxl")
    _, xlsx_path = run_component_bridge(_df, _cfg)
    print("Excel gespeichert unter:", xlsx_path)
