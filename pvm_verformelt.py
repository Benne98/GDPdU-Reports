import os
import sys
import json
import re
from typing import List, Optional

import numpy as np
import pandas as pd

BASE_DIR   = os.path.dirname(os.path.abspath(__file__))
PARENT_DIR = os.path.dirname(BASE_DIR)

if BASE_DIR not in sys.path:
    sys.path.insert(0, BASE_DIR)
if PARENT_DIR not in sys.path:
    sys.path.insert(0, PARENT_DIR)

from funktionssammlung import (
    apply_filters,
    ensure_output_writable,
    build_output_file_path,
    to_kEUR,
    fy_label,
    ytd_label,
    ltm_label,
    fiscal_year_bounds,
    fdd_as_of_end,
    fdd_current_fy_end_year,
    fdd_as_of_is_fy_end,
    fdd_ytd_bounds,
    fdd_ltm_bounds,
    ensure_source_sheet_in_output,
    refresh_source_sheet_from_df,
    source_range_ref,
    build_formula_filter_branches,
    build_sumifs_formula_body,
    build_chunked_row_sum_formula,
    write_export_df_to_sheet,
    get_next_sheet_name_from_wb,
    ensure_text_filter_helper_columns_from_cfg_in_ws,
    add_accrual_col_in_ws,
    _get_sheet_header_map,
    _normalize_header_name,
    parse_invoice_fy_year,
)

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.dataframe import dataframe_to_rows

_SCRIPTS_DIR = os.path.join(BASE_DIR, "scripts")
if _SCRIPTS_DIR not in sys.path:
    sys.path.insert(0, _SCRIPTS_DIR)
from gst_excel_theme import THEME  # noqa: E402

# -----------------------
# Config laden
# -----------------------
if len(sys.argv) < 2:
    raise ValueError("Bitte den Pfad zur topconfig.json als Argument übergeben.")

CONFIG_PATH = sys.argv[1]

if not os.path.exists(CONFIG_PATH):
    raise FileNotFoundError(f"Config-Datei nicht gefunden: {CONFIG_PATH}")

with open(CONFIG_PATH, "r", encoding="utf-8-sig") as f:
    CONFIG = json.load(f)


# -----------------------
# Helpers / Validation
# -----------------------
def normalize_config(cfg: dict) -> dict:
    out = dict(cfg)

    required = [
        "file_path",
        "sheet_name",
        "output_file_path",
        "case_id",
        "group_col",
        "quantity_col",
        "revenue_col",
        "invoice_col",
        "as_of_year",
        "as_of_month",
        "period_mode",
        "fy_end_month",
        "fy_end_day",
        "pvm_method",
        "first_fy",
    ]

    for k in required:
        if k not in out or out[k] in (None, ""):
            raise ValueError(f"CONFIG['{k}'] muss gesetzt sein.")

    out.setdefault("run_id",              "")
    out.setdefault("base_sheet_name",     "pvm")
    out.setdefault("apply_fx",            False)
    out.setdefault("sort_by_bridge",      "revenue")
    if out.get("total_label_revenue") in (None, ""):
        out["total_label_revenue"] = "Gross sales"
    if out.get("total_label_cost") in (None, ""):
        out["total_label_cost"] = "Cost of materials"
    if out.get("total_label_gp") in (None, ""):
        out["total_label_gp"] = "Gross profit"
    out.setdefault("profit_mode",         "cost")
    out.setdefault("formula_mode",        True)
    out.setdefault("invoice_mapping_mode", "date")

    raw_filters = out.get("filters", {}) or {}
    if isinstance(raw_filters, list):
        filters = {"enabled": len(raw_filters) > 0, "rules": raw_filters}
    else:
        filters = dict(raw_filters)
    filters.setdefault("enabled", False)
    filters.setdefault("rules",   [])
    out["filters"] = filters

    out["profit_mode"] = str(out["profit_mode"]).strip().lower()
    if out["profit_mode"] not in {"cost", "profit"}:
        raise ValueError("CONFIG['profit_mode'] muss 'cost' oder 'profit' sein.")

    if out["profit_mode"] == "cost":
        if "cost_col" not in out or out["cost_col"] in (None, ""):
            raise ValueError("CONFIG['cost_col'] muss gesetzt sein, wenn profit_mode='cost'.")
    else:
        if "profit_col" not in out or out["profit_col"] in (None, ""):
            raise ValueError("CONFIG['profit_col'] muss gesetzt sein, wenn profit_mode='profit'.")

    out["invoice_mapping_mode"] = str(out["invoice_mapping_mode"]).strip().lower()
    if out["invoice_mapping_mode"] not in {"year", "date"}:
        raise ValueError("CONFIG['invoice_mapping_mode'] muss 'year' oder 'date' sein.")

    if out["invoice_mapping_mode"] == "year":
        mode = str(out["period_mode"]).strip().upper()
        if mode != "FY":
            raise ValueError("invoice_mapping_mode='year' unterstützt nur period_mode='FY'.")

    if out["pvm_method"] not in {"three_components", "classical", "chicago"}:
        raise ValueError("CONFIG['pvm_method'] muss 'three_components', 'classical' oder 'chicago' sein.")

    mode = str(out["period_mode"]).strip().upper()
    if mode not in {"FY", "YTD", "LTM"}:
        raise ValueError("CONFIG['period_mode'] muss 'FY', 'YTD' oder 'LTM' sein.")
    out["period_mode"] = mode

    out["as_of_year"]  = int(out["as_of_year"])
    out["as_of_month"] = int(out["as_of_month"])
    out["first_fy"]    = int(out["first_fy"])

    if not 1 <= out["as_of_month"] <= 12:
        raise ValueError("CONFIG['as_of_month'] muss 1..12 sein.")

    out["fy_end_month"] = int(out["fy_end_month"])
    out["fy_end_day"]   = int(out["fy_end_day"])

    if not 1 <= out["fy_end_month"] <= 12:
        raise ValueError("CONFIG['fy_end_month'] muss 1..12 sein.")
    if not 1 <= out["fy_end_day"] <= 31:
        raise ValueError("CONFIG['fy_end_day'] muss 1..31 sein.")
    if out["first_fy"] < 1900:
        raise ValueError("CONFIG['first_fy'] muss ein plausibles Geschäftsjahr sein.")

    out["apply_fx"] = bool(out["apply_fx"])
    if out["apply_fx"]:
        if "fx_col" not in out or out["fx_col"] in (None, ""):
            raise ValueError("CONFIG['fx_col'] muss gesetzt sein, wenn apply_fx=True.")

    if out["sort_by_bridge"] not in {"revenue", "cost", "gross_profit"}:
        raise ValueError("CONFIG['sort_by_bridge'] muss 'revenue', 'cost' oder 'gross_profit' sein.")

    rep = dict(out.get("reported_numbers", {}) or {})

    sales_vals  = rep.get("sales",  [])
    profit_vals = rep.get("profit", [])

    if not isinstance(sales_vals, (list, tuple)):
        sales_vals  = []
    if not isinstance(profit_vals, (list, tuple)):
        profit_vals = []

    out["reported_numbers"] = {
        "sales":  list(sales_vals),
        "profit": list(profit_vals),
    }

    return out


def fy_end_date(year: int, cfg: dict) -> pd.Timestamp:
    first = pd.Timestamp(year=year, month=int(cfg["fy_end_month"]), day=1)
    last  = first + pd.offsets.MonthEnd(0)
    d     = min(int(cfg["fy_end_day"]), int(last.day))
    return pd.Timestamp(year=year, month=int(cfg["fy_end_month"]), day=d)


def build_periods(cfg: dict):
    cfg = normalize_config(cfg)

    mode     = str(cfg["period_mode"]).strip().upper()
    first_fy = int(cfg["first_fy"])

    fy_end_m = int(cfg["fy_end_month"])
    fy_end_d = int(cfg["fy_end_day"])
    as_of_month = int(cfg["as_of_month"])
    as_of_end = fdd_as_of_end(int(cfg["as_of_year"]), as_of_month)

    periods = {}

    if mode == "FY":
        current_fy_end_year = fdd_current_fy_end_year(as_of_end, fy_end_m, fy_end_d)
        as_of_is_fy_end = fdd_as_of_is_fy_end(as_of_end, fy_end_m, fy_end_d)
        latest_closed = current_fy_end_year if as_of_is_fy_end else (current_fy_end_year - 1)

        if first_fy > latest_closed:
            raise ValueError(
                f"CONFIG['first_fy']={first_fy} liegt nach dem letzten abgeschlossenen FY ({latest_closed})."
            )

        labels = []
        for year in range(first_fy, latest_closed + 1):
            start, end   = fiscal_year_bounds(year, fy_end_m, fy_end_d)
            lbl          = fy_label(year)
            periods[lbl] = (start, end)
            labels.append(lbl)

        return periods, labels

    if mode == "YTD":
        latest_ytd_year = as_of_end.year
        if first_fy > latest_ytd_year:
            raise ValueError(
                f"CONFIG['first_fy']={first_fy} liegt nach dem letzten YTD-Jahr ({latest_ytd_year})."
            )

        labels = []
        for year in range(first_fy, latest_ytd_year + 1):
            end       = fdd_as_of_end(year, as_of_month)
            start, _  = fdd_ytd_bounds(end, fy_end_m, fy_end_d)

            lbl          = ytd_label(year)
            periods[lbl] = (start, end)
            labels.append(lbl)

        return periods, labels

    if mode == "LTM":
        latest_ltm_year = as_of_end.year
        if first_fy > latest_ltm_year:
            raise ValueError(
                f"CONFIG['first_fy']={first_fy} liegt nach dem letzten LTM-Jahr ({latest_ltm_year})."
            )

        labels = []
        for year in range(first_fy, latest_ltm_year + 1):
            end        = fdd_as_of_end(year, as_of_month)
            start, _   = fdd_ltm_bounds(end)

            lbl          = ltm_label(year)
            periods[lbl] = (start, end)
            labels.append(lbl)

        return periods, labels

    raise ValueError("CONFIG['period_mode'] muss 'FY', 'YTD' oder 'LTM' sein.")


# -----------------------
# Core Data Processing
# -----------------------
def preprocess_pvm_input(df: pd.DataFrame, cfg: dict) -> pd.DataFrame:
    cfg = normalize_config(cfg)

    d         = df.copy()
    d.columns = d.columns.astype(str).str.replace("\u00a0", " ", regex=False).str.strip()

    d = apply_filters(d, cfg)

    required_cols = [
        cfg["group_col"],
        cfg["quantity_col"],
        cfg["revenue_col"],
    ]

    required_cols.append(cfg["invoice_col"])

    if cfg["profit_mode"] == "cost":
        required_cols.append(cfg["cost_col"])
    else:
        required_cols.append(cfg["profit_col"])

    for c in required_cols:
        if c not in d.columns:
            raise ValueError(f"Spalte '{c}' fehlt im Input.")

    if cfg["invoice_mapping_mode"] == "date":
        d[cfg["invoice_col"]] = pd.to_datetime(
            d[cfg["invoice_col"]],
            errors="coerce",
            dayfirst=True
        )
    else:
        d["_invoice_fy_year"] = parse_invoice_fy_year(d[cfg["invoice_col"]])

    d[cfg["quantity_col"]] = pd.to_numeric(d[cfg["quantity_col"]], errors="coerce")
    d[cfg["revenue_col"]]  = pd.to_numeric(d[cfg["revenue_col"]],  errors="coerce")

    if cfg["profit_mode"] == "cost":
        d[cfg["cost_col"]]   = pd.to_numeric(d[cfg["cost_col"]],   errors="coerce").fillna(0.0)
    else:
        d[cfg["profit_col"]] = pd.to_numeric(d[cfg["profit_col"]], errors="coerce").fillna(0.0)

    d["volume"]  = d[cfg["quantity_col"]]
    d["revenue"] = d[cfg["revenue_col"]]

    if cfg["profit_mode"] == "cost":
        d["cost"]         = d[cfg["cost_col"]]
        d["gross_profit"] = d["revenue"] - d["cost"]
    else:
        d["gross_profit"] = d[cfg["profit_col"]]
        d["cost"]         = d["revenue"] - d["gross_profit"]

    if cfg.get("apply_fx", False):
        fx_col = cfg.get("fx_col")
        if not fx_col or fx_col not in d.columns:
            raise ValueError("apply_fx=True, aber fx_col fehlt oder existiert nicht im Input.")

        d["fx_rate"]  = pd.to_numeric(d[fx_col], errors="coerce")
        valid_fx      = d["fx_rate"].notna() & (d["fx_rate"] != 0)

        d.loc[valid_fx, "revenue"] = d.loc[valid_fx, "revenue"] * d.loc[valid_fx, "fx_rate"]

        if cfg["profit_mode"] == "cost":
            d.loc[valid_fx, "cost"]         = d.loc[valid_fx, "cost"]         * d.loc[valid_fx, "fx_rate"]
            d["gross_profit"]               = d["revenue"] - d["cost"]
        else:
            d.loc[valid_fx, "gross_profit"] = d.loc[valid_fx, "gross_profit"] * d.loc[valid_fx, "fx_rate"]
            d["cost"]                       = d["revenue"] - d["gross_profit"]

    if cfg["invoice_mapping_mode"] == "date":
        bad = d[
            d[cfg["invoice_col"]].isna()
            | d[cfg["group_col"]].isna()
            | d["volume"].isna()
            | d["revenue"].isna()
        ]
    else:
        bad = d[
            d["_invoice_fy_year"].isna()
            | d[cfg["group_col"]].isna()
            | d["volume"].isna()
            | d["revenue"].isna()
        ]

    if cfg["invoice_mapping_mode"] == "date":
        d = d.dropna(subset=[cfg["invoice_col"], cfg["group_col"], "volume", "revenue"])
    else:
        d = d.dropna(subset=["_invoice_fy_year", cfg["group_col"], "volume", "revenue"])

    d["group_key"] = d[cfg["group_col"]].astype(str).str.strip()

    before = d.index
    d = d[d["group_key"].notna() & (d["group_key"] != "")]

    return d


def aggregate_year(d: pd.DataFrame, start: pd.Timestamp, end: pd.Timestamp, cfg: dict) -> pd.DataFrame:
    if cfg["invoice_mapping_mode"] == "date":
        mask = (d[cfg["invoice_col"]] >= start) & (d[cfg["invoice_col"]] <= end)
    else:
        target_fy_year = int(end.year)
        mask           = pd.to_numeric(d["_invoice_fy_year"], errors="coerce") == target_fy_year

    tmp = d.loc[mask].copy()

    out = tmp.groupby("group_key", as_index=False).agg(
        volume=("volume", "sum"),
        revenue=("revenue", "sum"),
        cost=("cost", "sum"),
        gross_profit=("gross_profit", "sum"),
    )

    return out.fillna(0)


def compute_pvm_effects(df1: pd.DataFrame, df2: pd.DataFrame, method: str, value_col: str) -> pd.DataFrame:
    merged = df1.merge(df2, on="group_key", how="outer", suffixes=("1", "2")).fillna(0)

    V1 = merged["volume1"]
    V2 = merged["volume2"]

    X1 = merged[f"{value_col}1"]
    X2 = merged[f"{value_col}2"]

    P1 = X1 / V1.replace(0, np.nan)
    P2 = X2 / V2.replace(0, np.nan)

    P1 = P1.fillna(0)
    P2 = P2.fillna(0)

    merged["FY1"] = X1
    merged["FY2"] = X2

    if method == "three_components":
        merged["Volume"] = (V2 - V1) * P1
        merged["Price"]  = (P2 - P1) * V1
        merged["Mix"]    = (P2 - P1) * (V2 - V1)

    elif method == "classical":
        merged["Volume"] = (V2 - V1) * P1
        merged["Price"]  = (P2 - P1) * V2

    elif method == "chicago":
        min_vol = np.minimum(V1, V2)
        P_ref   = np.where(V2 >= V1, P2, P1)

        merged["Volume"] = (V2 - V1) * P_ref
        merged["Price"]  = (P2 - P1) * min_vol

    else:
        raise ValueError(f"Unbekannte Methode: {method}")

    if method != "three_components" and "Mix" in merged.columns:
        merged = merged.drop(columns=["Mix"])

    # Nur Bridge-Spalten behalten — sonst bleiben volume1/2, revenue1/2 … stehen und
    # der nächste merge() in build_metric_bridge_table erzeugt doppelte Spalten (_x/_y).
    keep = ["group_key", "FY1", "FY2", "Price", "Volume"]
    if method == "three_components" and "Mix" in merged.columns:
        keep.append("Mix")
    keep = [c for c in keep if c in merged.columns]
    return merged[keep].copy()


def raw_period_col(label: str) -> str:
    return f"PERIOD_{label}"


def raw_price_col(newer_label: str) -> str:
    return f"PRICE_{newer_label}"


def raw_volume_col(newer_label: str) -> str:
    return f"VOLUME_{newer_label}"


def raw_mix_col(newer_label: str) -> str:
    return f"MIX_{newer_label}"


def display_header_from_raw(raw: str) -> str:
    if raw == "group_key":
        return "kEUR"
    if raw.startswith("PERIOD_"):
        return raw.split("_", 1)[1]
    if raw.startswith("PRICE_"):
        return "Price"
    if raw.startswith("VOLUME_"):
        return "Volume"
    if raw.startswith("MIX_"):
        return "Mix"
    return str(raw)


def build_prev_label_map(labels: list[str]) -> dict[str, str]:
    out = {}
    for i in range(1, len(labels)):
        out[labels[i]] = labels[i - 1]
    return out


def coerce_reported_number(value):
    return pd.to_numeric(pd.Series([value]), errors="coerce").iloc[0]


def build_left_helper_layout(labels: list[str]) -> dict:
    period_helpers = {}
    current_col    = 2  # A bleibt frei

    for lbl in labels:
        period_helpers[lbl] = {
            "qty_col":   current_col,
            "price_col": current_col + 1,
        }
        current_col += 2

    key_col          = current_col
    spacer_col       = current_col + 1
    helper_right_col = spacer_col
    table_left_col   = helper_right_col + 1

    return {
        "LEFT_FREE_COL":    1,
        "PERIOD_HELPERS":   period_helpers,
        "KEY_COL":          key_col,
        "RIGHT_SPACER_COL": spacer_col,
        "HELPER_RIGHT_COL": helper_right_col,
        "TABLE_LEFT_COL":   table_left_col,
        "COL_OFFSET":       helper_right_col,
    }


def build_reported_value_map(cfg: dict, labels: list[str], metric_key: str) -> dict:
    vals = list((cfg.get("reported_numbers", {}) or {}).get(metric_key, []) or [])
    out  = {}

    for i, lbl in enumerate(labels):
        out[raw_period_col(lbl)] = coerce_reported_number(vals[i]) if i < len(vals) else np.nan

    return out


def append_recon_and_reported_rows(
    df: pd.DataFrame,
    cfg: dict,
    labels: list[str],
    metric_key: str,
    section_label: str,
) -> pd.DataFrame:
    out = df.copy()

    recon_row              = {c: "" for c in out.columns}
    recon_row["group_key"] = "Recon. difference"

    reported_row              = {c: "" for c in out.columns}
    reported_row["group_key"] = f"{section_label} (reported)"

    rep_map = build_reported_value_map(cfg, labels, metric_key)
    for col_name, val in rep_map.items():
        if col_name in reported_row:
            reported_row[col_name] = val

    return pd.concat(
        [out, pd.DataFrame([recon_row, reported_row])],
        ignore_index=True
    )


def get_reported_row_labels(cfg: dict) -> set[str]:
    return {
        f"{cfg['total_label_revenue']} (reported)",
        f"{cfg['total_label_gp']} (reported)",
    }


def build_bridge_export_columns(cfg: dict, labels: list[str]) -> list[str]:
    cols = ["group_key", raw_period_col(labels[0])]

    for newer_label in labels[1:]:
        cols.append(raw_price_col(newer_label))
        cols.append(raw_volume_col(newer_label))
        if cfg["pvm_method"] == "three_components":
            cols.append(raw_mix_col(newer_label))
        cols.append(raw_period_col(newer_label))

    return cols


def rename_bridge_columns(df_bridge: pd.DataFrame, old_label: str, new_label: str, method: str) -> pd.DataFrame:
    out = df_bridge.copy()

    rename_map = {
        "FY1":    raw_period_col(old_label),
        "FY2":    raw_period_col(new_label),
        "Price":  raw_price_col(new_label),
        "Volume": raw_volume_col(new_label),
    }

    if method == "three_components" and "Mix" in out.columns:
        rename_map["Mix"] = raw_mix_col(new_label)

    out = out.rename(columns=rename_map)

    if method != "three_components" and "Mix" in out.columns:
        out = out.drop(columns=["Mix"])

    return out


def build_metric_bridge_table(
    agg_by_label: dict[str, pd.DataFrame],
    labels: list[str],
    method: str,
    value_col: str,
    section_label: str,
) -> pd.DataFrame:
    if len(labels) < 2:
        raise ValueError("Für eine Bridge-Auswertung werden mindestens zwei Perioden benötigt.")

    first_old = labels[0]
    first_new = labels[1]

    out = compute_pvm_effects(
        agg_by_label[first_old][["group_key", "volume", value_col]],
        agg_by_label[first_new][["group_key", "volume", value_col]],
        method,
        value_col=value_col,
    )
    out = rename_bridge_columns(out, first_old, first_new, method)

    for i in range(2, len(labels)):
        old_label = labels[i - 1]
        new_label = labels[i]

        nxt = compute_pvm_effects(
            agg_by_label[old_label][["group_key", "volume", value_col]],
            agg_by_label[new_label][["group_key", "volume", value_col]],
            method,
            value_col=value_col,
        )
        nxt = rename_bridge_columns(nxt, old_label, new_label, method)

        out = out.merge(
            nxt.drop(columns=[raw_period_col(old_label)], errors="ignore"),
            on="group_key",
            how="outer",
        )

    return out.fillna(0)


def apply_bridge_order(x: pd.DataFrame, order: list[str]) -> pd.DataFrame:
    order_map     = {k: i for i, k in enumerate(order)}
    out           = x.copy()
    out["_order"] = out["group_key"].map(order_map)
    out           = out.sort_values("_order").drop(columns="_order")
    return out


def append_section_and_total(df: pd.DataFrame, section_label: str) -> pd.DataFrame:
    out = df.copy()

    # Use missing values (not "") so concat does not coerce section headers to 0.0;
    # zeros make build_pvm_row_sets treat the header row as a total → 0 formulas written.
    section_df = pd.DataFrame([{"group_key": section_label}])

    total              = {c: "" for c in out.columns}
    total["group_key"] = section_label

    for c in out.columns:
        if c == "group_key":
            continue
        total[c] = pd.to_numeric(out[c], errors="coerce").fillna(0).sum()

    total_df = pd.DataFrame([total])

    return pd.concat([section_df, out, total_df], ignore_index=True).reindex(columns=out.columns)


def build_two_bridges(df: pd.DataFrame, cfg: dict) -> pd.DataFrame:
    periods, labels = build_periods(cfg)
    d               = preprocess_pvm_input(df, cfg)
    method          = cfg["pvm_method"]

    if len(labels) < 2:
        raise ValueError("Es werden mindestens zwei Perioden benötigt, um Bridges zu berechnen.")

    agg_by_label = {
        lbl: aggregate_year(d, *periods[lbl], cfg)
        for lbl in labels
    }

    rev = build_metric_bridge_table(
        agg_by_label=agg_by_label,
        labels=labels,
        method=method,
        value_col="revenue",
        section_label=cfg["total_label_revenue"],
    )

    cost_tbl = build_metric_bridge_table(
        agg_by_label=agg_by_label,
        labels=labels,
        method=method,
        value_col="cost",
        section_label=cfg["total_label_cost"],
    )

    gp = build_metric_bridge_table(
        agg_by_label=agg_by_label,
        labels=labels,
        method=method,
        value_col="gross_profit",
        section_label=cfg["total_label_gp"],
    )

    latest_period_raw = raw_period_col(labels[-1])

    sort_source = cfg.get("sort_by_bridge", "revenue")
    if sort_source == "revenue":
        sort_df = rev.copy()
    elif sort_source == "cost":
        sort_df = cost_tbl.copy()
    else:
        sort_df = gp.copy()

    sort_df[latest_period_raw] = pd.to_numeric(sort_df[latest_period_raw], errors="coerce").fillna(0)

    order = (
        sort_df.sort_values(by=latest_period_raw, ascending=False)["group_key"]
        .astype(str)
        .tolist()
    )

    rev      = apply_bridge_order(rev,      order)
    cost_tbl = apply_bridge_order(cost_tbl, order)
    gp       = apply_bridge_order(gp,       order)

    rev = append_section_and_total(rev, cfg["total_label_revenue"])
    rev = append_recon_and_reported_rows(
        rev,
        cfg=cfg,
        labels=labels,
        metric_key="sales",
        section_label=cfg["total_label_revenue"],
    )

    cost_tbl = append_section_and_total(cost_tbl, cfg["total_label_cost"])

    gp = append_section_and_total(gp, cfg["total_label_gp"])
    gp = append_recon_and_reported_rows(
        gp,
        cfg=cfg,
        labels=labels,
        metric_key="profit",
        section_label=cfg["total_label_gp"],
    )

    final = pd.concat([rev, cost_tbl, gp], ignore_index=True)
    return final


def finalize_column_order(df: pd.DataFrame, cfg: dict) -> pd.DataFrame:
    out        = df.copy()
    _, labels  = build_periods(cfg)

    cols = [c for c in build_bridge_export_columns(cfg, labels) if c in out.columns]
    out  = out[cols]
    return out


def apply_kEUR_formatting(df: pd.DataFrame, cfg: dict) -> pd.DataFrame:
    """Scale value columns to kEUR; preserve section headers as non-numeric."""
    out = df.copy()
    bridge_labels = {
        str(cfg["total_label_revenue"]).strip().casefold(),
        str(cfg["total_label_cost"]).strip().casefold(),
        str(cfg["total_label_gp"]).strip().casefold(),
    }
    bridge_occurrence: dict[str, int] = {}

    for idx in out.index:
        gk = str(out.at[idx, "group_key"] if "group_key" in out.columns else "").strip()
        gk_cf = gk.casefold()

        if gk == "Recon. difference":
            for col in out.columns:
                if col != "group_key":
                    out.at[idx, col] = np.nan
            continue

        if gk_cf in bridge_labels:
            bridge_occurrence[gk_cf] = bridge_occurrence.get(gk_cf, 0) + 1
            if bridge_occurrence[gk_cf] == 1:
                for col in out.columns:
                    if col != "group_key":
                        out.at[idx, col] = np.nan
                continue

        for col in out.columns:
            if col == "group_key":
                continue
            out.at[idx, col] = to_kEUR(pd.Series([out.at[idx, col]])).iloc[0]

    return out


def excel_abs_ref(col_idx: int, row_idx: int) -> str:
    return f"${get_column_letter(col_idx)}${row_idx}"


def _append_source_helper_col(ws_source, header: str, values: list, number_format: str | None = None) -> int:
    target_col = ws_source.max_column + 1
    ws_source.cell(row=1, column=target_col).value = header

    for i, v in enumerate(values, start=2):
        cell = ws_source.cell(row=i, column=target_col)
        cell.value = v
        if number_format:
            cell.number_format = number_format

    return target_col


def ensure_pvm_fx_helper_columns(ws_source, source_df: pd.DataFrame, cfg: dict) -> dict[str, str]:
    if not cfg.get("apply_fx", False):
        return {}

    d         = source_df.copy()
    d.columns = d.columns.astype(str).str.replace("\u00a0", " ", regex=False).str.strip()

    fx_col = cfg["fx_col"]
    if fx_col not in d.columns:
        raise ValueError(f"FX-Spalte '{fx_col}' fehlt im Source-Sheet.")

    fx           = pd.to_numeric(d[fx_col], errors="coerce")
    fx_effective = fx.where(fx.notna() & (fx != 0), 1.0)

    out        = {}
    header_map = _get_sheet_header_map(ws_source)

    def ensure_one(key: str, source_col: str):
        helper_header = f"_PVM_FX_{key.upper()}_{source_col}"
        if _normalize_header_name(helper_header) not in header_map:
            vals = (
                pd.to_numeric(d[source_col], errors="coerce").fillna(0.0) * fx_effective
            ).tolist()
            _append_source_helper_col(ws_source, helper_header, vals)
        out[key] = helper_header

    ensure_one("revenue", cfg["revenue_col"])

    if cfg["profit_mode"] == "cost":
        ensure_one("cost",   cfg["cost_col"])
    else:
        ensure_one("profit", cfg["profit_col"])

    return out


def _write_period_bounds_to_columns(
    ws,
    cfg: dict,
    ctx: dict,
    col_indices: List[int],
    period_start,
    period_end,
    red_font: Font,
) -> None:
    """Write hidden period metadata (year or date range) into given columns — GST-style."""
    for col_idx in col_indices:
        if cfg["invoice_mapping_mode"] == "year":
            if ctx["PERIOD_YEAR_ROW"] is None:
                continue
            cell = ws.cell(row=ctx["PERIOD_YEAR_ROW"], column=col_idx)
            cell.value = int(period_end.year)
            cell.font = red_font
            cell.alignment = Alignment(horizontal="right", vertical="center")
            cell.number_format = "0"
        else:
            if ctx["PERIOD_START_ROW"] is None or ctx["PERIOD_END_ROW"] is None:
                continue
            start_cell = ws.cell(row=ctx["PERIOD_START_ROW"], column=col_idx)
            end_cell = ws.cell(row=ctx["PERIOD_END_ROW"], column=col_idx)
            start_cell.value = pd.Timestamp(period_start).to_pydatetime()
            end_cell.value = pd.Timestamp(period_end).to_pydatetime()
            start_cell.font = red_font
            end_cell.font = red_font
            start_cell.alignment = Alignment(horizontal="right", vertical="center")
            end_cell.alignment = Alignment(horizontal="right", vertical="center")
            start_cell.number_format = "DD.MM.YYYY"
            end_cell.number_format = "DD.MM.YYYY"


def build_period_helper_rows(ws, cfg: dict, export_columns, ctx, export_col_map: Optional[dict] = None):
    """Period bounds in helper qty/price cols and in visible FY columns (like GST)."""
    red_font = Font(name=THEME.font_name, size=THEME.font_size, color="FFFF5149")
    if export_col_map is None:
        export_col_map = {
            raw: ctx["TABLE_LEFT_COL"] + i
            for i, raw in enumerate(export_columns)
        }

    for period_label in ctx["PERIOD_LABELS"]:
        period_start, period_end = ctx["periods"][period_label]
        helper_info = ctx["PERIOD_HELPERS"][period_label]
        helper_cols = [helper_info["qty_col"], helper_info["price_col"]]

        visible_col = export_col_map.get(raw_period_col(period_label))
        cols_to_write = list(helper_cols)
        if visible_col is not None and visible_col not in cols_to_write:
            cols_to_write.append(visible_col)

        _write_period_bounds_to_columns(
            ws, cfg, ctx, cols_to_write, period_start, period_end, red_font
        )


def get_pvm_layout_context(ws, cfg: dict) -> dict:
    labels        = build_periods(cfg)[1]
    helper_layout = build_left_helper_layout(labels)

    formula_mode = bool(cfg.get("formula_mode", True))
    extra_period_rows = 0
    if formula_mode:
        extra_period_rows = 2 if cfg["invoice_mapping_mode"] == "date" else 1

    title_row    = 1
    subtitle_row = 2

    if formula_mode and cfg["invoice_mapping_mode"] == "date":
        period_year_row  = None
        period_start_row = 3
        period_end_row   = 4
    elif formula_mode:
        period_year_row  = 3
        period_start_row = 3
        period_end_row   = None
    else:
        period_year_row  = None
        period_start_row = None
        period_end_row   = None

    company_row    = 2 + extra_period_rows + 1
    header_row     = company_row + 1
    data_start_row = header_row + 1

    return {
        "COL_OFFSET":        helper_layout["COL_OFFSET"],
        "INSERT_ROW_COUNT":  header_row - 1,
        "EXTRA_PERIOD_ROWS": extra_period_rows,
        "PERIOD_YEAR_ROW":   period_year_row,
        "PERIOD_START_ROW":  period_start_row,
        "PERIOD_END_ROW":    period_end_row,
        "TITLE_ROW":         title_row,
        "SUBTITLE_ROW":      subtitle_row,
        "COMPANY_ROW":       company_row,
        "HEADER_ROW":        header_row,
        "DATA_START_ROW":    data_start_row,
        "PERIOD_LABELS":     labels,
        "PERIOD_HELPERS":    helper_layout["PERIOD_HELPERS"],
        "LEFT_FREE_COL":     helper_layout["LEFT_FREE_COL"],
        "KEY_COL":           helper_layout["KEY_COL"],
        "RIGHT_SPACER_COL":  helper_layout["RIGHT_SPACER_COL"],
        "HELPER_RIGHT_COL":  helper_layout["HELPER_RIGHT_COL"],
        "TABLE_LEFT_COL":    helper_layout["TABLE_LEFT_COL"],
        "TABLE_RIGHT_COL":   ws.max_column,
        "TABLE_BOTTOM_ROW":  ws.max_row,
        "LABEL_COL":         helper_layout["TABLE_LEFT_COL"],
    }


def period_bounds_col_idx(ctx: dict, period_label: str, *, for_price: bool = False) -> int:
    """Column with period start/end (or year) for SUMIFS bounds — qty or price helper col."""
    helpers = ctx["PERIOD_HELPERS"][period_label]
    key = "price_col" if for_price else "qty_col"
    return int(helpers[key])


def row_has_table_numeric(ws, row_idx: int, ctx: dict) -> bool:
    for c in range(ctx["LABEL_COL"] + 1, ctx["TABLE_RIGHT_COL"] + 1):
        if is_real_numeric_excel_value(ws.cell(row=row_idx, column=c).value):
            return True
    return False


def is_real_numeric_excel_value(val) -> bool:
    if isinstance(val, (int, float, np.integer, np.floating)) and pd.notna(val):
        return True
    if isinstance(val, str):
        s = val.strip().replace("\u00a0", " ").replace(" ", "")
        if not s or s == "-":
            return False
        if s.startswith("="):
            return False
        try:
            float(s.replace(",", ".").replace("'", ""))
            return True
        except ValueError:
            return False
    return False


def build_pvm_bridge_blocks(ws, ctx: dict, cfg: dict) -> List[dict]:
    """One block per bridge: section → details → total → recon → reported (optional)."""

    def _norm(s: str) -> str:
        return (s or "").strip().casefold()

    label_to_metric = {
        _norm(cfg["total_label_revenue"]): "revenue",
        _norm(cfg["total_label_cost"]):    "cost",
        _norm(cfg["total_label_gp"]):      "gross_profit",
    }
    reported_labels = get_reported_row_labels(cfg)
    label_col       = ctx["LABEL_COL"]
    data_start      = ctx["DATA_START_ROW"]
    bottom          = ctx["TABLE_BOTTOM_ROW"]

    blocks: List[dict] = []
    r = data_start

    while r <= bottom:
        raw = ws.cell(row=r, column=label_col).value
        if not isinstance(raw, str):
            r += 1
            continue

        nk = _norm(raw)
        if nk not in label_to_metric or row_has_table_numeric(ws, r, ctx):
            r += 1
            continue

        metric     = label_to_metric[nk]
        sec_row    = r
        detail_rows: List[int] = []
        r += 1

        total_row = None
        while r <= bottom:
            lab = ws.cell(row=r, column=label_col).value
            if isinstance(lab, str):
                nk2 = _norm(lab)
                if nk2 in label_to_metric and not row_has_table_numeric(ws, r, ctx):
                    break
                if nk2 == nk and row_has_table_numeric(ws, r, ctx):
                    total_row = r
                    r += 1
                    break
            detail_rows.append(r)
            r += 1

        if total_row is None:
            continue

        recon_row = None
        if r <= bottom:
            lab = ws.cell(row=r, column=label_col).value
            if isinstance(lab, str) and lab.strip() == "Recon. difference":
                recon_row = r
                r += 1

        reported_row = None
        if r <= bottom:
            lab = ws.cell(row=r, column=label_col).value
            if isinstance(lab, str) and lab.strip() in reported_labels:
                reported_row = r
                r += 1

        blocks.append(
            {
                "metric":        metric,
                "section_row":   sec_row,
                "detail_rows":   detail_rows,
                "total_row":     total_row,
                "recon_row":     recon_row,
                "reported_row":  reported_row,
            }
        )

    return blocks


def build_pvm_row_sets(ws, ctx: dict, cfg: dict):
    sections: dict = {}
    totals: dict   = {}
    for block in build_pvm_bridge_blocks(ws, ctx, cfg):
        sections[block["metric"]] = {
            "section_row": block["section_row"],
            "detail_rows": block["detail_rows"],
        }
        totals[block["metric"]] = block["total_row"]
    return sections, totals


def write_pvm_helper_area(ws, cfg: dict, ctx: dict):
    red_font      = Font(name=THEME.font_name, size=THEME.font_size, color="FFFF5149")
    red_bold_font = Font(name=THEME.font_name, size=THEME.font_size, color="FFFF5149", bold=True)

    header_row       = ctx["HEADER_ROW"]
    data_start_row   = ctx["DATA_START_ROW"]
    table_bottom_row = ctx["TABLE_BOTTOM_ROW"]
    label_col        = ctx["LABEL_COL"]
    key_col          = ctx["KEY_COL"]
    helper_right_col = ctx["HELPER_RIGHT_COL"]

    sections, totals = build_pvm_row_sets(ws, ctx, cfg)

    non_detail_rows = set()
    for info in sections.values():
        non_detail_rows.add(info["section_row"])
    for tot_row in totals.values():
        non_detail_rows.add(tot_row)

    special_non_detail_labels = {"Recon. difference"} | get_reported_row_labels(cfg)

    ws.cell(row=header_row, column=ctx["LEFT_FREE_COL"]).value = ""

    for lbl in ctx["PERIOD_LABELS"]:
        ws.cell(
            row=header_row,
            column=ctx["PERIOD_HELPERS"][lbl]["qty_col"]
        ).value = f"Quantity {lbl}"

        ws.cell(
            row=header_row,
            column=ctx["PERIOD_HELPERS"][lbl]["price_col"]
        ).value = f"Price {lbl}"

    ws.cell(row=header_row, column=key_col).value              = cfg["group_col"]
    ws.cell(row=header_row, column=ctx["RIGHT_SPACER_COL"]).value = ""

    for c in range(1, helper_right_col + 1):
        ws.cell(row=header_row, column=c).font      = red_bold_font
        ws.cell(row=header_row, column=c).alignment = Alignment(
            horizontal="left",
            vertical="bottom",
            wrap_text=True,
        )

    for r in range(data_start_row, table_bottom_row + 1):
        visible_label = ws.cell(row=r, column=label_col).value

        for c in range(1, helper_right_col + 1):
            ws.cell(row=r, column=c).font      = red_font
            ws.cell(row=r, column=c).alignment = Alignment(horizontal="left", vertical="center")
            if c != key_col:
                ws.cell(row=r, column=c).value = None

        if r in non_detail_rows:
            ws.cell(row=r, column=key_col).value = None
            continue

        if isinstance(visible_label, str) and visible_label.strip() in special_non_detail_labels:
            ws.cell(row=r, column=key_col).value = None
            continue

        ws.cell(row=r, column=key_col).value = visible_label


def build_report_sum_formula(row_numbers: list[int], col_idx: int) -> str:
    return build_chunked_row_sum_formula(row_numbers, col_idx)


def build_metric_body_for_period(
    *,
    metric: str,
    period_label: str,
    row_idx: int,
    period_col_idx: int,
    cfg: dict,
    source_sheet_name: str,
    source_header_map: dict,
    filter_branches: list[list[tuple[str, str]]],
    fx_helper_cols: dict[str, str],
    ctx: dict,
) -> str:
    key_ref = f"${get_column_letter(ctx['KEY_COL'])}{row_idx}"

    group_rng  = source_range_ref(source_sheet_name, source_header_map, cfg["group_col"])
    base_pairs = [(group_rng, key_ref)]

    # period_col_idx must be the column that contains period year / start / end (helper or visible FY col).
    bounds_col_idx = period_col_idx

    if cfg["invoice_mapping_mode"] == "date":
        invoice_rng = source_range_ref(source_sheet_name, source_header_map, cfg["invoice_col"])
        start_ref   = excel_abs_ref(bounds_col_idx, ctx["PERIOD_START_ROW"])
        end_ref     = excel_abs_ref(bounds_col_idx, ctx["PERIOD_END_ROW"])

        base_pairs += [
            (invoice_rng, f'">="&{start_ref}'),
            (invoice_rng, f'"<="&{end_ref}'),
        ]
    else:
        year_rng = source_range_ref(source_sheet_name, source_header_map, cfg["invoice_col"])
        year_ref = excel_abs_ref(bounds_col_idx, ctx["PERIOD_YEAR_ROW"])

        base_pairs += [
            (year_rng, year_ref),
        ]

    def build_body_for_source_col(source_col_name: str, divide_by_1000: bool = True) -> str:
        sum_rng = source_range_ref(source_sheet_name, source_header_map, source_col_name)
        return build_sumifs_formula_body(
            sum_rng,
            base_pairs,
            filter_branches,
            divide_by_1000=divide_by_1000,
        )

    revenue_source = fx_helper_cols.get("revenue", cfg["revenue_col"])
    cost_source    = fx_helper_cols.get("cost",    cfg.get("cost_col"))
    profit_source  = fx_helper_cols.get("profit",  cfg.get("profit_col"))

    if metric == "quantity":
        return build_body_for_source_col(cfg["quantity_col"], divide_by_1000=False)

    if metric == "revenue":
        return build_body_for_source_col(revenue_source, divide_by_1000=True)

    if metric == "cost":
        if cfg["profit_mode"] == "cost":
            return build_body_for_source_col(cost_source, divide_by_1000=True)

        rev_body    = build_body_for_source_col(revenue_source, divide_by_1000=True)
        profit_body = build_body_for_source_col(profit_source,  divide_by_1000=True)
        return f"(({rev_body})-({profit_body}))"

    if metric == "gross_profit":
        if cfg["profit_mode"] == "profit":
            return build_body_for_source_col(profit_source, divide_by_1000=True)

        rev_body  = build_body_for_source_col(revenue_source, divide_by_1000=True)
        cost_body = build_body_for_source_col(cost_source,    divide_by_1000=True)
        return f"(({rev_body})-({cost_body}))"

    raise ValueError(f"Unbekannte Kennzahl: {metric}")


def build_pvm_effect_formula(
    *,
    effect_kind: str,
    method: str,
    row_idx: int,
    old_amount_col_idx: int,
    new_amount_col_idx: int,
    old_qty_col_idx: int,
    new_qty_col_idx: int,
    old_price_col_idx: int,
    new_price_col_idx: int,
) -> str:
    old_amount_ref = f"{get_column_letter(old_amount_col_idx)}{row_idx}"
    new_amount_ref = f"{get_column_letter(new_amount_col_idx)}{row_idx}"

    oq_ref = f"{get_column_letter(old_qty_col_idx)}{row_idx}"
    nq_ref = f"{get_column_letter(new_qty_col_idx)}{row_idx}"
    p1_ref = f"{get_column_letter(old_price_col_idx)}{row_idx}"
    p2_ref = f"{get_column_letter(new_price_col_idx)}{row_idx}"

    if effect_kind == "volume":
        if method == "three_components":
            expr = f"(({nq_ref})-({oq_ref}))*({p1_ref})"
        elif method == "classical":
            expr = f"(({nq_ref})-({oq_ref}))*({p1_ref})"
        elif method == "chicago":
            expr = f"(({nq_ref})-({oq_ref}))*IF(({nq_ref})>=({oq_ref}),({p2_ref}),({p1_ref}))"
        else:
            raise ValueError(f"Unbekannte Methode: {method}")

    elif effect_kind == "price":
        if method == "three_components":
            expr = f"(({p2_ref})-({p1_ref}))*({oq_ref})"
        elif method == "classical":
            expr = f"(({p2_ref})-({p1_ref}))*({nq_ref})"
        elif method == "chicago":
            expr = f"(({p2_ref})-({p1_ref}))*MIN(({oq_ref}),({nq_ref}))"
        else:
            raise ValueError(f"Unbekannte Methode: {method}")

    elif effect_kind == "mix":
        expr = f"(({p2_ref})-({p1_ref}))*(({nq_ref})-({oq_ref}))"

    else:
        raise ValueError(f"Unbekannter Effect-Typ: {effect_kind}")

    return f'=IF(OR(COUNT({old_amount_ref})=0,COUNT({new_amount_ref})=0),"n/a",{expr})'


def write_pvm_helper_formulas(
    ws,
    cfg: dict,
    ctx: dict,
    sections: dict,
    labels: list[str],
    source_sheet_name: str,
    source_header_map: dict,
    filter_branches: list[list[tuple[str, str]]],
    fx_helper_cols: dict[str, str],
):
    qty_fmt   = '#,##0.00;(#,##0.00);"-"'
    price_fmt = '#,##0.0000;(#,##0.0000);"-"'

    for metric, info in sections.items():
        detail_rows = info["detail_rows"]
        if not detail_rows:
            continue

        for r in detail_rows:
            for period_label in labels:
                qty_col   = ctx["PERIOD_HELPERS"][period_label]["qty_col"]
                price_col = ctx["PERIOD_HELPERS"][period_label]["price_col"]

                qty_body = build_metric_body_for_period(
                    metric="quantity",
                    period_label=period_label,
                    row_idx=r,
                    period_col_idx=qty_col,
                    cfg=cfg,
                    source_sheet_name=source_sheet_name,
                    source_header_map=source_header_map,
                    filter_branches=filter_branches,
                    fx_helper_cols=fx_helper_cols,
                    ctx=ctx,
                )

                amount_body = build_metric_body_for_period(
                    metric=metric,
                    period_label=period_label,
                    row_idx=r,
                    period_col_idx=price_col,
                    cfg=cfg,
                    source_sheet_name=source_sheet_name,
                    source_header_map=source_header_map,
                    filter_branches=filter_branches,
                    fx_helper_cols=fx_helper_cols,
                    ctx=ctx,
                )

                qty_ref = f"{get_column_letter(qty_col)}{r}"

                qty_cell               = ws.cell(row=r, column=qty_col)
                qty_cell.value         = f'=IFERROR({qty_body},0)'
                qty_cell.number_format = qty_fmt

                price_cell               = ws.cell(row=r, column=price_col)
                price_cell.value         = f'=IFERROR(IF({qty_ref}=0,0,({amount_body})/{qty_ref}),0)'
                price_cell.number_format = price_fmt


def apply_recon_difference_formulas(
    ws,
    cfg: dict,
    ctx: dict,
    export_col_map: dict,
    labels: list[str],
):
    label_col        = ctx["LABEL_COL"]
    data_start_row   = ctx["DATA_START_ROW"]
    table_bottom_row = ctx["TABLE_BOTTOM_ROW"]

    label_to_rows = {}
    for r in range(data_start_row, table_bottom_row + 1):
        val = ws.cell(row=r, column=label_col).value
        if isinstance(val, str):
            label_to_rows.setdefault(val.strip(), []).append(r)

    targets = [
        (cfg["total_label_revenue"], f"{cfg['total_label_revenue']} (reported)"),
        (cfg["total_label_gp"],      f"{cfg['total_label_gp']} (reported)"),
    ]

    for total_label, reported_label in targets:
        reported_rows = label_to_rows.get(reported_label, [])
        if not reported_rows:
            continue

        reported_row = reported_rows[0]
        recon_row    = reported_row - 1
        total_row    = reported_row - 2

        if recon_row < data_start_row or total_row < data_start_row:
            continue

        for lbl in labels:
            col_idx = export_col_map.get(raw_period_col(lbl))
            if not col_idx:
                continue

            col_letter = get_column_letter(col_idx)
            ws.cell(row=recon_row, column=col_idx).value        = f"={col_letter}{reported_row}-{col_letter}{total_row}"
            ws.cell(row=recon_row, column=col_idx).number_format = '#,##0;(#,##0);"-"'


def apply_pvm_formulas(
    wb,
    report_sheet_name: str,
    source_sheet_name: str,
    cfg: dict,
    source_df: pd.DataFrame,
    export_columns: list[str],
):
    if not cfg.get("formula_mode", True):
        return

    ws        = wb[report_sheet_name]
    ws_source = wb[source_sheet_name]

    refresh_source_sheet_from_df(ws_source, source_df, cfg)

    text_helper_cols_by_index = ensure_text_filter_helper_columns_from_cfg_in_ws(
        ws_source=ws_source,
        source_df=source_df,
        cfg=cfg,
    )

    fx_helper_cols = ensure_pvm_fx_helper_columns(
        ws_source=ws_source,
        source_df=source_df,
        cfg=cfg,
    )

    source_header_map = _get_sheet_header_map(ws_source)

    filter_branches = build_formula_filter_branches(
        cfg=cfg,
        source_sheet_name=source_sheet_name,
        source_header_map=source_header_map,
        text_helper_cols_by_index=text_helper_cols_by_index,
    )

    periods, labels = build_periods(cfg)
    prev_label_map  = build_prev_label_map(labels)

    ctx                     = get_pvm_layout_context(ws, cfg)
    ctx["periods"]          = periods
    ctx["TABLE_RIGHT_COL"]  = ctx["TABLE_LEFT_COL"] + len(export_columns) - 1
    ctx["TABLE_BOTTOM_ROW"] = ws.max_row

    export_col_map = {
        raw: ctx["TABLE_LEFT_COL"] + i
        for i, raw in enumerate(export_columns)
    }

    build_period_helper_rows(ws, cfg, export_columns, ctx, export_col_map)

    write_pvm_helper_area(ws, cfg, ctx)

    sections, totals = build_pvm_row_sets(ws, ctx, cfg)
    if not sections:
        blocks = build_pvm_bridge_blocks(ws, ctx, cfg)
        for block in blocks:
            sections[block["metric"]] = {
                "section_row": block["section_row"],
                "detail_rows": block["detail_rows"],
            }
            totals[block["metric"]] = block["total_row"]

    formulas_written = 0

    # 1) Helper-Spalten
    write_pvm_helper_formulas(
        ws=ws,
        cfg=cfg,
        ctx=ctx,
        sections=sections,
        labels=labels,
        source_sheet_name=source_sheet_name,
        source_header_map=source_header_map,
        filter_branches=filter_branches,
        fx_helper_cols=fx_helper_cols,
    )

    # 2) Sichtbare Perioden + PVM-Effekte
    for metric, info in sections.items():
        detail_rows = info["detail_rows"]
        total_row   = totals.get(metric)

        for r in detail_rows:
            for raw_col, col_idx in export_col_map.items():
                cell = ws.cell(row=r, column=col_idx)

                if raw_col == "group_key":
                    continue

                if raw_col.startswith("PERIOD_"):
                    period_label = raw_col.split("_", 1)[1]
                    body = build_metric_body_for_period(
                        metric=metric,
                        period_label=period_label,
                        row_idx=r,
                        period_col_idx=col_idx,
                        cfg=cfg,
                        source_sheet_name=source_sheet_name,
                        source_header_map=source_header_map,
                        filter_branches=filter_branches,
                        fx_helper_cols=fx_helper_cols,
                        ctx=ctx,
                    )
                    cell.value         = f'=IFERROR({body},"n/a")'
                    cell.number_format = '#,##0;(#,##0);"-"'
                    formulas_written  += 1
                    continue

                if raw_col.startswith("PRICE_") or raw_col.startswith("VOLUME_") or raw_col.startswith("MIX_"):
                    newer_label = raw_col.split("_", 1)[1]
                    older_label = prev_label_map.get(newer_label)
                    if not older_label:
                        continue

                    old_amount_col_idx = export_col_map[raw_period_col(older_label)]
                    new_amount_col_idx = export_col_map[raw_period_col(newer_label)]

                    old_qty_col_idx   = ctx["PERIOD_HELPERS"][older_label]["qty_col"]
                    new_qty_col_idx   = ctx["PERIOD_HELPERS"][newer_label]["qty_col"]
                    old_price_col_idx = ctx["PERIOD_HELPERS"][older_label]["price_col"]
                    new_price_col_idx = ctx["PERIOD_HELPERS"][newer_label]["price_col"]

                    if raw_col.startswith("PRICE_"):
                        effect_kind = "price"
                    elif raw_col.startswith("VOLUME_"):
                        effect_kind = "volume"
                    else:
                        effect_kind = "mix"

                    cell.value = build_pvm_effect_formula(
                        effect_kind=effect_kind,
                        method=cfg["pvm_method"],
                        row_idx=r,
                        old_amount_col_idx=old_amount_col_idx,
                        new_amount_col_idx=new_amount_col_idx,
                        old_qty_col_idx=old_qty_col_idx,
                        new_qty_col_idx=new_qty_col_idx,
                        old_price_col_idx=old_price_col_idx,
                        new_price_col_idx=new_price_col_idx,
                    )
                    cell.number_format = '#,##0;(#,##0);"-"'
                    formulas_written  += 1
                    continue

        if total_row is not None:
            for raw_col, col_idx in export_col_map.items():
                if raw_col == "group_key":
                    continue
                ws.cell(row=total_row, column=col_idx).value        = build_report_sum_formula(detail_rows, col_idx)
                ws.cell(row=total_row, column=col_idx).number_format = '#,##0;(#,##0);"-"'
                formulas_written += 1

    apply_recon_difference_formulas(
        ws=ws,
        cfg=cfg,
        ctx=ctx,
        export_col_map=export_col_map,
        labels=labels,
    )

    if formulas_written == 0:
        raise ValueError(
            "Es wurden 0 Formeln geschrieben. "
            "Wahrscheinlich wurden keine Section-/Detail-Zeilen erkannt oder die Exportspalten nicht korrekt gemappt."
        )

    wb.calculation.forceFullCalc  = True
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.calcMode       = "auto"


def apply_pvm_vertical_outlines(ws, ctx: dict, cfg: dict) -> None:
    """Per-bridge row groups: section + details collapsible; total/recon/reported always visible."""
    for block in build_pvm_bridge_blocks(ws, ctx, cfg):
        sec_row   = block["section_row"]
        details   = block["detail_rows"]
        tot_row   = block["total_row"]
        recon_row = block.get("recon_row")
        rep_row   = block.get("reported_row")
        collapse  = block["metric"] == "cost"

        ws.row_dimensions[sec_row].outlineLevel = 1
        ws.row_dimensions[sec_row].hidden       = collapse
        ws.row_dimensions[sec_row].collapsed  = collapse

        for r in details:
            ws.row_dimensions[r].outlineLevel = 1
            ws.row_dimensions[r].hidden       = collapse

        ws.row_dimensions[tot_row].outlineLevel = 0
        ws.row_dimensions[tot_row].hidden       = False
        ws.row_dimensions[tot_row].collapsed    = False

        for r in (recon_row, rep_row):
            if r is None:
                continue
            ws.row_dimensions[r].outlineLevel = 0
            ws.row_dimensions[r].hidden       = False
            ws.row_dimensions[r].collapsed    = False


# -----------------------
# Excel Formatting
# -----------------------
def format_pvm_excel(
    wb,
    target_sheet_name: str,
    cfg: dict,
    export_columns: list[str],
):
    if target_sheet_name not in wb.sheetnames:
        raise ValueError(f"Sheet '{target_sheet_name}' nicht in Workbook gefunden.")

    ws = wb[target_sheet_name]
    periods, labels = build_periods(cfg)

    ctx            = get_pvm_layout_context(ws, cfg)
    ctx["periods"] = periods

    ws.insert_rows(1, amount=ctx["INSERT_ROW_COUNT"])
    ws.insert_cols(1, amount=ctx["COL_OFFSET"])

    ctx            = get_pvm_layout_context(ws, cfg)
    ctx["periods"] = periods

    TABLE_LEFT_COL   = ctx["TABLE_LEFT_COL"]
    TABLE_RIGHT_COL  = TABLE_LEFT_COL + len(export_columns) - 1
    TABLE_BOTTOM_ROW = ws.max_row
    HELPER_RIGHT_COL = ctx["HELPER_RIGHT_COL"]

    PAINT_RIGHT_COL  = TABLE_RIGHT_COL + 50
    PAINT_BOTTOM_ROW = TABLE_BOTTOM_ROW + 50

    ctx["TABLE_RIGHT_COL"]  = TABLE_RIGHT_COL
    ctx["TABLE_BOTTOM_ROW"] = TABLE_BOTTOM_ROW

    fmt_money = '#,##0;(#,##0);"-"'

    base_font   = THEME.font_base
    bold_font   = THEME.font_bold
    header_fill = THEME.fill_header
    white_fill  = THEME.fill_white
    fill_grey   = THEME.fill_tech
    period_fill = THEME.fill_subtotal

    border_top        = THEME.border_subtotal_top
    border_top_bottom = Border(
        top=THEME.border_subtotal_top.top,
        bottom=THEME.border_subtotal_top.top,
    )

    section_font = Font(
        name=THEME.font_name,
        size=THEME.font_size,
        bold=True,
        color=THEME.text_brand_title,
    )
    tech_font = Font(
        name=THEME.font_name,
        size=THEME.font_size,
        color="FFFF5149",
    )
    tech_bold_font = Font(
        name=THEME.font_name,
        size=THEME.font_size,
        bold=True,
        color="FFFF5149",
    )

    for r in range(1, PAINT_BOTTOM_ROW + 1):
        for c in range(1, PAINT_RIGHT_COL + 1):
            ws.cell(row=r, column=c).font = base_font

    for r in range(1, PAINT_BOTTOM_ROW + 1):
        for c in range(TABLE_LEFT_COL, PAINT_RIGHT_COL + 1):
            ws.cell(row=r, column=c).fill = white_fill

    for r in range(1, PAINT_BOTTOM_ROW + 1):
        for c in range(1, HELPER_RIGHT_COL + 1):
            ws.cell(row=r, column=c).fill = fill_grey

    ws.sheet_properties.outlinePr.summaryRight = True
    ws.sheet_properties.outlinePr.summaryBelow = True

    for c in range(1, HELPER_RIGHT_COL + 1):
        letter = get_column_letter(c)
        ws.column_dimensions[letter].outlineLevel = 1
        ws.column_dimensions[letter].hidden       = True

    ws.column_dimensions[get_column_letter(TABLE_LEFT_COL)].collapsed = True

    if cfg.get("formula_mode", True):
        helper_rows = []
        if cfg["invoice_mapping_mode"] == "year":
            helper_rows = [ctx["PERIOD_YEAR_ROW"]]
        else:
            helper_rows = [ctx["PERIOD_START_ROW"], ctx["PERIOD_END_ROW"]]

        for r in helper_rows:
            if r is None:
                continue
            ws.row_dimensions[r].hidden       = True
            ws.row_dimensions[r].outlineLevel = 1
            ws.row_dimensions[r].collapsed    = False

        if helper_rows:
            ws.row_dimensions[helper_rows[-1]].collapsed = False

    method_label = cfg["pvm_method"].replace("_", " ")

    ws.cell(row=ctx["TITLE_ROW"], column=TABLE_LEFT_COL).value = cfg.get("title", "")
    ws.cell(row=ctx["TITLE_ROW"], column=TABLE_LEFT_COL).font = Font(
        name=THEME.font_name,
        size=THEME.font_size_title,
        color=THEME.text_brand_title,
    )

    subtitle_text = f"{cfg.get('table', '').strip()} ({method_label})".strip()
    if subtitle_text.startswith("("):
        subtitle_text = method_label
    ws.cell(row=ctx["SUBTITLE_ROW"], column=TABLE_LEFT_COL).value = subtitle_text
    ws.cell(row=ctx["SUBTITLE_ROW"], column=TABLE_LEFT_COL).font = Font(
        name=THEME.font_name,
        size=THEME.font_size_subtitle,
        color=THEME.text_brand_title,
    )

    company_line = str(cfg.get("company", "") or "").strip()
    ws.cell(row=ctx["COMPANY_ROW"], column=TABLE_LEFT_COL).value = company_line
    ws.cell(row=ctx["COMPANY_ROW"], column=TABLE_LEFT_COL).font = Font(
        name=THEME.font_name,
        size=THEME.font_size_company,
        color=THEME.text_brand_title,
        bold=True,
    )

    ws.row_dimensions[ctx["HEADER_ROW"]].height = 24

    for i, raw in enumerate(export_columns):
        c           = TABLE_LEFT_COL + i
        cell        = ws.cell(row=ctx["HEADER_ROW"], column=c)
        cell.value  = display_header_from_raw(raw)
        cell.font   = THEME.font_header
        cell.fill   = header_fill
        cell.border = THEME.border_header_bottom
        cell.alignment = Alignment(
            horizontal="left" if c == TABLE_LEFT_COL else "right",
            vertical="bottom",
            wrap_text=True
        )

    if cfg.get("formula_mode", True):
        export_col_map_fmt = {
            raw: ctx["TABLE_LEFT_COL"] + i
            for i, raw in enumerate(export_columns)
        }
        build_period_helper_rows(ws, cfg, export_columns, ctx, export_col_map_fmt)

    for r in range(ctx["DATA_START_ROW"], TABLE_BOTTOM_ROW + 1):
        ws.row_dimensions[r].height = 12

    for r in range(ctx["DATA_START_ROW"], TABLE_BOTTOM_ROW + 1):
        for c in range(TABLE_LEFT_COL, TABLE_RIGHT_COL + 1):
            cell = ws.cell(row=r, column=c)

            if isinstance(cell.value, (int, float, np.integer, np.floating)):
                cell.number_format = fmt_money

            cell.alignment = Alignment(
                horizontal="left" if c == TABLE_LEFT_COL else "right",
                vertical="center",
            )

    fy_cols = []
    for i, raw in enumerate(export_columns):
        c       = TABLE_LEFT_COL + i
        visible = display_header_from_raw(raw)
        if isinstance(visible, str) and (visible.startswith("FY") or visible.startswith("YTD") or visible.startswith("LTM")):
            fy_cols.append(c)

    for r in range(ctx["DATA_START_ROW"], TABLE_BOTTOM_ROW + 1):
        for c in fy_cols:
            ws.cell(row=r, column=c).fill = period_fill

    reported_labels = get_reported_row_labels(cfg)
    section_rows: set[int] = set()
    total_rows: set[int]   = set()
    reported_rows: set[int] = set()

    for block in build_pvm_bridge_blocks(ws, ctx, cfg):
        section_rows.add(block["section_row"])
        total_rows.add(block["total_row"])
        if block.get("reported_row"):
            reported_rows.add(block["reported_row"])

    for r in range(ctx["DATA_START_ROW"], TABLE_BOTTOM_ROW + 1):
        rd              = ws.row_dimensions[r]
        rd.outlineLevel = 0
        rd.hidden       = False
        rd.collapsed    = False

    apply_pvm_vertical_outlines(ws, ctx, cfg)

    for r in section_rows:
        for c in range(TABLE_LEFT_COL, TABLE_RIGHT_COL + 1):
            cell      = ws.cell(row=r, column=c)
            cell.font = section_font
            cell.fill = period_fill if c in fy_cols else white_fill

    for r in total_rows:
        for c in range(TABLE_LEFT_COL, TABLE_RIGHT_COL + 1):
            cell      = ws.cell(row=r, column=c)
            cell.font = bold_font
            cell.fill = period_fill if c in fy_cols else white_fill
            cell.border = border_top

    for r in reported_rows:
        for c in range(TABLE_LEFT_COL, TABLE_RIGHT_COL + 1):
            cell      = ws.cell(row=r, column=c)
            cell.font = bold_font
            cell.fill = period_fill if c in fy_cols else white_fill
            cell.border = border_top_bottom

    col_max_length = {}
    for c in range(TABLE_LEFT_COL, TABLE_RIGHT_COL + 1):
        for r in range(ctx["HEADER_ROW"], TABLE_BOTTOM_ROW + 1):
            cell  = ws.cell(row=r, column=c)
            value = cell.value

            if value is None:
                continue

            if isinstance(value, str):
                length = len(value)
            elif isinstance(value, (int, float, np.integer, np.floating)):
                if cell.number_format == fmt_money:
                    if float(value) == 0:
                        formatted = "-"
                    else:
                        formatted = f"{float(value):,.0f}"
                    length = len(formatted)
                else:
                    length = len(str(value))
            else:
                length = len(str(value))

            col_max_length[c] = max(col_max_length.get(c, 0), length)

    for c, max_len in col_max_length.items():
        letter         = get_column_letter(c)
        adjusted_width = min(max_len + 1, 90)
        ws.column_dimensions[letter].width = adjusted_width

# -----------------------
# MAIN
# -----------------------
def main():
    cfg = normalize_config(CONFIG)

    df = pd.read_excel(
        cfg["file_path"],
        sheet_name=cfg["sheet_name"],
        engine="openpyxl"
    )

    final = build_two_bridges(df, cfg)
    final = finalize_column_order(final, cfg)
    final = apply_kEUR_formatting(final, cfg)

    output_path = build_output_file_path(cfg)
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    ensure_output_writable(output_path)

    source_sheet_name = ensure_source_sheet_in_output(cfg, output_path)

    wb = load_workbook(output_path)

    target_sheet_name = get_next_sheet_name_from_wb(
        wb,
        cfg.get("base_sheet_name", "pvm")
    )

    write_export_df_to_sheet(
        wb=wb,
        export_df=final,
        target_sheet_name=target_sheet_name,
    )

    format_pvm_excel(
        wb=wb,
        target_sheet_name=target_sheet_name,
        cfg=cfg,
        export_columns=final.columns.tolist(),
    )

    if cfg.get("formula_mode", True):
        apply_pvm_formulas(
            wb=wb,
            report_sheet_name=target_sheet_name,
            source_sheet_name=source_sheet_name,
            cfg=cfg,
            source_df=df.copy(),
            export_columns=final.columns.tolist(),
        )

    wb.save(output_path)
    wb.close()

    print(f"PVM Excel erfolgreich erstellt: {output_path}")
    print(f"Neues Sheet: {target_sheet_name}")


if __name__ == "__main__":
    main()
