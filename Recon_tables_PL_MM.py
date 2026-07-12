import os
import sys
import json
import re
from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.styles import Border, Side
from openpyxl.formatting.rule import CellIsRule

PROJECT_ROOT = Path(__file__).resolve().parent
SCRIPTS_DIR = PROJECT_ROOT / "scripts"
if str(SCRIPTS_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPTS_DIR))

from databook_workbook import (  # noqa: E402
    BS_RECON_MAPPING_FILE,
    CF_NA_L3_ORDER_FILE,
    MASTER_WORKBOOK_STR,
    PL_RECON_MAPPING_FILE,
    remove_stale_entity_recon_sheets,
    sanitize_entity_recon_sheet_name,
)

from gst_excel_theme import THEME, apply_recon_portfolio_layout  # noqa: E402
from databook_excel_layout import (  # noqa: E402
    LAYOUT_PL,
    apply_databook_row_border_band,
    assign_recon_block_columns,
    build_entity_recon_blocks,
    check_row_groups_after_table,
    collapse_check_portfolio,
    paint_grey_white_canvas,
    prune_zero_value_rows,
    refresh_master_sheet,
    write_fs_check_section,
)
from report_row_layout import (  # noqa: E402
    build_pl_row_structure,
    l3_order_from_mapping,
    normalize_l4_sort_basis,
)
# =============================================
# HARD-CODED DEFAULTS / STABLE SETTINGS
# =============================================

DESKTOP_DIR = PROJECT_ROOT / "Desktop"
DEFAULT_SOURCE_FILE = MASTER_WORKBOOK_STR
DEFAULT_SOURCE_SHEET = "Master_PL"
DEFAULT_TARGET_FILE = MASTER_WORKBOOK_STR
DEFAULT_MAPPING_FILE = PL_RECON_MAPPING_FILE
DEFAULT_MAPPING_SHEET_INDEX = 0
DEFAULT_MAPPING_ENGINE = "openpyxl"

DEFAULT_SOURCE_ENGINE = "openpyxl"
DEFAULT_SOURCE_HEADER_ROW = 0

# Master workbooks may store Reported/Adjusted in L5 or L6 depending on pipeline version.
SOURCE_COL_CANDIDATES = ("L5", "L6")
SOURCE_COL_NAME = "L5"
REPORTED_FILTER_VALUE = "reported"

MAPPING_REQUIRED_COLUMNS = {"L3"}

CHECK_TITLE_LABEL   = "Source - Financial statements"
CHECK_DIFF_LABEL    = "Difference to trial balances"
NET_RESULT_TECH_KEY = "Net result"

DEFAULT_TECHNICAL_PL_KEY_MAP = {
    "Other operating expenses":        "Other operating expenses",
    "Wages & salaries":    "Wages & salaries",
    "Reduction in earnings": "Reduction in earnings",
    "Other operating income":          "Other operating income",
    "Personnel expenses":    "Personnel expenses",
    "Social security":       "Social security",
    "Taxes on income":                   "Taxes on income",
}

DEFAULT_TOP_BORDER_ONLY = [
    "Net sales",
    "Gross profit",
    "EBIT",
    "EBT",
]

DEFAULT_TOP_AND_BOTTOM = [
    "EBITDA",
    "Net result",
]

DEFAULT_PL_TOTALS_CONFIG = [
    {
        "label":        "Total output",
        "components":   ["Net sales", "Δ Finished goods & WIP", "Own work capitalised"],
        "insert_after": "Own work capitalised",
    },
    {
        "label":        "Gross profit",
        "components":   ["Total output", "Cost of goods sold"],
        "insert_after": "Cost of goods sold",
    },
    {
        "label":        "Net operating expenses",
        "components":   ["Other operating income", "Other operating expenses"],
        "insert_after": "Other operating expenses",
    },
    {
        "label":        "EBITDA",
        "components":   ["Gross profit", "Personnel expenses", "Net operating expenses"],
        "insert_after": "Net operating expenses",
    },
    {
        "label":        "EBIT",
        "components":   ["EBITDA", "D&A"],
        "insert_after": "D&A",
    },
    {
        "label":        "EBT",
        "components":   ["EBIT", "Financial result"],
        "insert_after": "Financial result",
    },
    {
        "label":        "Net result",
        "components":   ["EBT", "Taxes on income", "Other taxes"],
        "insert_after": "Other taxes",
    },
]

DEFAULT_TITLES = {
    "aggregated_title":           "Aggregated",
    "consolidation_title":        "Consolidation",
    "difference_title":           "Difference",
    "financial_statements_title": "Financial statements",
    "ic_display_name":            "IC eliminations",
}

DEFAULT_SORT_METRIC_LABELS = {
    "revenue": "Net sales",
    "assets":  "Total assets",
}

IC_MASTER_ENTITY_DEFAULT = "Consolidation"

BLANK_TOKENS = {"", "nan", "none", "null"}

# Layout — A free | B..E helpers | F spacer | G POS
_COL_LAYOUT = LAYOUT_PL
MAP_START_COL = _COL_LAYOUT.map_start_col
MAP_END_COL = _COL_LAYOUT.map_end_col
TECH_SPACER_COL = _COL_LAYOUT.spacer_col
POS_COL = _COL_LAYOUT.pos_col
FIRST_ENTITY_COL = _COL_LAYOUT.first_value_col

PROJECT_TITLE_ROW = 1
SUBTITLE_ROW      = 2

TITLE_ROW = 6
TITLE_COL = POS_COL

ENTITY_CODE_ROW = 3
HEADER_ROW      = 8
BLOCK_TITLE_ROW = HEADER_ROW - 1
DATA_START_ROW  = HEADER_ROW + 1

SPACER_WIDTH             = 1.14
VALUE_COL_WIDTH          = 7.86
ROW_HEIGHT               = 12
PROJECT_TITLE_ROW_HEIGHT = 36
REPEAT_POS_WIDTH         = 32

FILL_PADDING_ROWS = 100
FILL_PADDING_COLS = 40

# Styles — aligned with gst_excel_theme / SuSa master output
FONT_BASE = THEME.font_base
FONT_BASE_BOLD = THEME.font_bold
FONT_HEADER = THEME.font_header
FONT_MAPPING = Font(name=THEME.font_name, size=THEME.font_size, color=THEME.text_header)
FONT_TITLE = Font(
    name=THEME.font_name,
    size=THEME.font_size,
    bold=True,
    color=THEME.text_brand_title,
)
FONT_PROJECT_TITLE = Font(
    name=THEME.font_name,
    size=THEME.font_size_title,
    color=THEME.text_brand_title,
    bold=False,
)
FONT_SUBTITLE = Font(
    name=THEME.font_name,
    size=THEME.font_size_subtitle,
    color=THEME.text_brand_title,
    bold=False,
)

FILL_TECH = THEME.fill_tech
FILL_WHITE = THEME.fill_white
FILL_HEADER = THEME.fill_header
FILL_SUBTOTAL = THEME.fill_subtotal
CHECK_FILL = PatternFill("solid", fgColor="FFFEF3C7")

ALIGN_LEFT = Alignment(horizontal="left", vertical="center")
ALIGN_CENTER = Alignment(horizontal="center", vertical="center")
ALIGN_RIGHT = Alignment(horizontal="right", vertical="center")

BORDER_HEADER_BOTTOM = THEME.border_header_bottom
BORDER_SUBTOTAL_TOP = THEME.border_subtotal_top
BORDER_SUBTOTAL_TOP_BOTTOM = Border(
    top=Side(style="medium", color=THEME.border_strong),
    bottom=Side(style="medium", color=THEME.border_strong),
)

NUM_FMT_INT = '#,##0;(#,##0);"-"'
DIFF_FONT = Font(name=THEME.font_name, size=THEME.font_size, color=THEME.delta_negative)

YEARS = []  # set dynamically in main()

# =============================================
# CONFIG
# =============================================

def normalize_technical_pl_key(x):
    if x is None or pd.isna(x):
        return x
    s = str(x).strip()
    if s in DEFAULT_TECHNICAL_PL_KEY_MAP:
        return DEFAULT_TECHNICAL_PL_KEY_MAP[s]
    lower_map = {k.lower(): v for k, v in DEFAULT_TECHNICAL_PL_KEY_MAP.items()}
    return lower_map.get(s.lower(), s)


def resolve_source_section_column(df: pd.DataFrame) -> str:
    """
    Pick the column that carries Reported/Adjusted markers in the current master layout.
    Newer masters use L6 when L5 is reserved or empty.
    """
    markers = {REPORTED_FILTER_VALUE, "adjusted", "consolidation"}

    def marker_count(col: str) -> int:
        if col not in df.columns:
            return 0
        series = df[col].astype(str).str.strip().str.lower()
        return int(series.isin(markers).sum())

    best_col = None
    best_count = 0
    for col in SOURCE_COL_CANDIDATES:
        count = marker_count(col)
        if count > best_count:
            best_col = col
            best_count = count

    if best_col and best_count > 0:
        return best_col

    for col in SOURCE_COL_CANDIDATES:
        if col in df.columns:
            return col

    return SOURCE_COL_CANDIDATES[0]


def normalize_config(cfg: dict) -> dict:
    out = dict(cfg or {})

    required = ["project_name", "company_name", "sort_by", "paths"]
    for k in required:
        if k not in out or out[k] in (None, "", []):
            raise ValueError(f"CONFIG['{k}'] muss gesetzt sein.")

    out["project_name"] = str(out["project_name"]).strip()
    out["company_name"] = str(out["company_name"]).strip()

    out["sort_by"] = str(out["sort_by"]).strip().lower()
    if out["sort_by"] not in {"assets", "revenue", "custom"}:
        raise ValueError("CONFIG['sort_by'] muss 'assets', 'revenue' oder 'custom' sein.")

    out["l4_sort_basis"] = normalize_l4_sort_basis(out, default="latest_fy")

    entity_order = list(out.get("entity_order") or [])
    out["entity_order"] = [str(e).strip() for e in entity_order if str(e).strip()]
    if out["sort_by"] == "custom" and not out["entity_order"]:
        raise ValueError("CONFIG['entity_order'] ist bei sort_by='custom' erforderlich.")

    paths = dict(out.get("paths", {}) or {})
    for k in ["source_file", "source_sheet", "target_file"]:
        if k not in paths or paths[k] in (None, ""):
            raise ValueError(f"CONFIG['paths']['{k}'] muss gesetzt sein.")

    paths.setdefault("source_engine", DEFAULT_SOURCE_ENGINE)
    paths.setdefault("mapping_file", DEFAULT_MAPPING_FILE)
    paths.setdefault("append_to_master", False)
    paths.setdefault("master_file", paths.get("target_file"))
    paths.setdefault("report_sheet",       "PL_Reconciliation")
    paths.setdefault("audit_master_sheet", "Master_PL")
    out["paths"] = paths

    display = dict(out.get("display", {}) or {})
    display.setdefault("unit_label", "kEUR")

    titles = dict(display.get("titles", {}) or {})
    titles.setdefault("aggregated_title",           DEFAULT_TITLES["aggregated_title"])
    titles.setdefault("consolidation_title",         DEFAULT_TITLES["consolidation_title"])
    titles.setdefault("difference_title",            DEFAULT_TITLES["difference_title"])
    titles.setdefault("financial_statements_title",  DEFAULT_TITLES["financial_statements_title"])
    titles.setdefault("ic_display_name",             DEFAULT_TITLES["ic_display_name"])
    display["titles"]  = titles
    out["display"]     = display

    sort_metric_labels = dict(out.get("sort_metric_labels", {}) or {})
    sort_metric_labels.setdefault("revenue", DEFAULT_SORT_METRIC_LABELS["revenue"])
    sort_metric_labels.setdefault("assets",  DEFAULT_SORT_METRIC_LABELS["assets"])
    sort_metric_labels = {
        str(k).strip().lower(): normalize_technical_pl_key(v)
        for k, v in sort_metric_labels.items()
    }
    out["sort_metric_labels"] = sort_metric_labels

    pl_cfg = dict(out.get("pl_config", {}) or {})

    raw_display_map = (
        pl_cfg.get("display_label_map")
        or pl_cfg.get("pl_position_rename_map")
        or {}
    )
    display_label_map = {}
    for k, v in dict(raw_display_map).items():
        tech_key = normalize_technical_pl_key(k)
        if tech_key is None or pd.isna(tech_key):
            continue
        display_label_map[str(tech_key).strip()] = str(v).strip()

    pl_cfg["display_label_map"] = display_label_map

    top_border_only = list(pl_cfg.get("top_border_only", DEFAULT_TOP_BORDER_ONLY) or [])
    top_and_bottom  = list(pl_cfg.get("top_and_bottom",  DEFAULT_TOP_AND_BOTTOM)  or [])
    totals_config   = list(pl_cfg.get("totals_config",   DEFAULT_PL_TOTALS_CONFIG) or [])

    pl_cfg["top_border_only"] = [
        normalize_technical_pl_key(x) for x in top_border_only
        if x not in (None, "")
    ]
    pl_cfg["top_and_bottom"] = [
        normalize_technical_pl_key(x) for x in top_and_bottom
        if x not in (None, "")
    ]

    normalized_totals = []
    for t in totals_config:
        td           = dict(t or {})
        label        = normalize_technical_pl_key(td.get("label"))
        insert_after = normalize_technical_pl_key(td.get("insert_after"))
        components   = [normalize_technical_pl_key(x) for x in list(td.get("components", []) or [])]

        if label in (None, "") or insert_after in (None, ""):
            continue

        normalized_totals.append(
            {
                "label":        label,
                "components":   [c for c in components if c not in (None, "")],
                "insert_after": insert_after,
            }
        )

    pl_cfg["totals_config"] = normalized_totals
    out["pl_config"] = pl_cfg

    out.setdefault("show_fs_check", True)

    fs_cfg           = dict(out.get("fs_check_values", {}) or {})
    fs_entities      = list(fs_cfg.get("entities",      []) or [])
    fs_consolidation = list(fs_cfg.get("consolidation", []) or [])
    out["fs_check_values"] = {
        "entities":      fs_entities,
        "consolidation": fs_consolidation,
    }

    return out


def build_desktop_default_config() -> dict:
    source_file = DEFAULT_SOURCE_FILE
    df = pd.read_excel(source_file, sheet_name=DEFAULT_SOURCE_SHEET, engine="openpyxl")
    entities = []
    for v in pd.unique(df["Entity"]):
        if is_blank_entity(v):
            continue
        name = str(v).strip()
        if name.lower() == IC_MASTER_ENTITY_DEFAULT.lower():
            continue
        entities.append(name)

    return {
        "project_name": "Desktop Test",
        "company_name": "Group",
        "sort_by": "custom",
        "entity_order": entities,
        "l4_sort_basis": "latest_fy",
        "show_fs_check": False,
        "fs_check_values": {"entities": [], "consolidation": []},
        "display": {"unit_label": "kEUR"},
        "pl_config": {},
        "paths": {
            "source_file": source_file,
            "source_sheet": DEFAULT_SOURCE_SHEET,
            "source_engine": DEFAULT_SOURCE_ENGINE,
            "target_file": DEFAULT_TARGET_FILE,
            "mapping_file": DEFAULT_MAPPING_FILE,
            "report_sheet": "PL_Reconciliation",
            "audit_master_sheet": "Master_PL",
            "master_file": DEFAULT_TARGET_FILE,
            "append_to_master": False,
        },
    }


def load_config_from_argv() -> dict:
    if len(sys.argv) < 2:
        cfg = build_desktop_default_config()
        print("Keine Config übergeben — verwende Desktop-Defaults:")
        print(f"  source:  {cfg['paths']['source_file']}")
        print(f"  mapping: {cfg['paths']['mapping_file']}")
        print(f"  target:  {cfg['paths']['target_file']}")
        return normalize_config(cfg)

    config_path = sys.argv[1]

    if not os.path.exists(config_path):
        raise FileNotFoundError(f"Config-Datei nicht gefunden: {config_path}")

    with open(config_path, "r", encoding="utf-8-sig") as f:
        cfg = json.load(f)

    return normalize_config(cfg)


# =============================================
# HELPERS
# =============================================

def is_blank_entity(x):
    if x is None or pd.isna(x):
        return True
    s = str(x).strip()
    return (s == "") or (s.lower() in BLANK_TOKENS)


def is_blank_label(x):
    if x is None or pd.isna(x):
        return True
    s = str(x).strip()
    return (s == "") or (s.lower() in BLANK_TOKENS)


def col_letter(idx: int) -> str:
    return get_column_letter(idx)


def technical_to_display_label(tech_label: str, cfg: dict) -> str:
    if tech_label is None or pd.isna(tech_label):
        return tech_label
    s = str(tech_label).strip()
    return cfg["pl_config"]["display_label_map"].get(s, s)


def detect_year_columns(df_: pd.DataFrame) -> list[str]:
    pattern = re.compile(r"^FY(\d{2}|\d{4})A$", re.IGNORECASE)

    years = []
    for c in df_.columns:
        s = str(c).strip()
        if pattern.match(s):
            years.append(s)

    if not years:
        raise ValueError("Es konnten keine FY-Spalten erkannt werden (erwartet z. B. FY23A, FY24A, FY25A).")

    def year_sort_key(x: str):
        m   = pattern.match(str(x).strip())
        val = m.group(1)
        return int(val)

    years = sorted(dict.fromkeys(years), key=year_sort_key)
    return years


def build_fs_check_value_map(fs_cfg: dict, individual_entities: list[str], years: list[str]) -> dict:
    result = {}

    entity_lists         = list(fs_cfg.get("entities",      []) or [])
    consolidation_values = list(fs_cfg.get("consolidation", []) or [])

    def coerce_value(v):
        if v in (None, ""):
            return 0
        try:
            return float(v)
        except Exception:
            return 0

    for idx, entity in enumerate(individual_entities):
        raw_values = entity_lists[idx] if idx < len(entity_lists) else []
        raw_values = list(raw_values or [])

        year_map = {}
        for y_idx, year in enumerate(years):
            value          = raw_values[y_idx] if y_idx < len(raw_values) else 0
            year_map[year] = coerce_value(value)

        result[entity] = year_map

    cons_map = {}
    for y_idx, year in enumerate(years):
        value          = consolidation_values[y_idx] if y_idx < len(consolidation_values) else 0
        cons_map[year] = coerce_value(value)

    result["Consolidation"] = cons_map
    return result


def build_fs_check_value_map_for_entity(
    fs_cfg: dict,
    entity: str,
    individual_entities: list[str],
    years: list[str],
) -> dict:
    full = build_fs_check_value_map(fs_cfg, individual_entities, years)
    return {entity: dict(full.get(entity, {}))}


def load_mapping_df(cfg: dict) -> pd.DataFrame:
    mapping_file = str(cfg.get("paths", {}).get("mapping_file") or DEFAULT_MAPPING_FILE)
    map_df = pd.read_excel(
        mapping_file,
        sheet_name=DEFAULT_MAPPING_SHEET_INDEX,
        engine=DEFAULT_MAPPING_ENGINE,
    )

    if not MAPPING_REQUIRED_COLUMNS.issubset(map_df.columns):
        raise ValueError("Mapping-Datei muss Spalte 'L3' enthalten.")

    map_df = map_df.copy()
    map_df["L3"] = map_df["L3"].astype(str).apply(normalize_technical_pl_key)

    return map_df


def load_source_df(cfg: dict) -> pd.DataFrame:
    global SOURCE_COL_NAME

    source_engine = str(cfg["paths"].get("source_engine") or DEFAULT_SOURCE_ENGINE)
    df = pd.read_excel(
        cfg["paths"]["source_file"],
        sheet_name=cfg["paths"]["source_sheet"],
        engine=source_engine,
        header=DEFAULT_SOURCE_HEADER_ROW,
    )

    SOURCE_COL_NAME = resolve_source_section_column(df)

    years = detect_year_columns(df)

    required = {"Entity", "L3", "L4", SOURCE_COL_NAME} | set(years)
    missing = required - set(df.columns)
    if missing:
        raise ValueError(f"Source-Sheet fehlt Spalten: {missing}")

    df_norm = df.copy()
    for col in ("L3", "L4"):
        df_norm[col] = df_norm[col].apply(normalize_technical_pl_key)

    return df_norm


def insert_after_anchor(struct: list[dict], anchor_label: str, new_row: dict) -> bool:
    for i, row in enumerate(struct):
        if row["tech_label"] == anchor_label:
            struct.insert(i + 1, new_row)
            return True
    return False


def apply_totals_to_row_structure(row_structure: list[dict], cfg: dict) -> list[dict]:
    for t in cfg["pl_config"]["totals_config"]:
        tech_label = t["label"]
        new_row = {
            "type":          "total",
            "tech_label":    tech_label,
            "display_label": technical_to_display_label(tech_label, cfg),
            "components":    t["components"],
            "L3":            tech_label,
            "L4":            "",
        }

        anchor = t["insert_after"]
        inserted = insert_after_anchor(row_structure, anchor, new_row)
        if not inserted:
            for candidate in reversed(t["components"]):
                if candidate != anchor and insert_after_anchor(row_structure, candidate, new_row):
                    inserted = True
                    break
        if not inserted:
            raise RuntimeError(
                f"Anchor '{anchor}' für Total '{t['label']}' nicht gefunden."
            )

    return row_structure


def prune_row_structure(struct: list[dict]) -> list[dict]:
    labels = set([r["tech_label"] for r in struct])
    pruned = []

    for r in struct:
        if r["type"] == "detail":
            pruned.append(r)
            continue

        if r["type"] == "subtotal":
            pruned.append(r)
            continue

        if r["type"] == "total":
            comps = r.get("components", [])
            valid = [c for c in comps if c in labels]
            if len(valid) == 0:
                continue
            r2 = dict(r)
            r2["components"] = valid
            pruned.append(r2)
            continue

        pruned.append(r)

    return pruned


def prune_nan_zero_blocks(struct: list[dict], df_master: pd.DataFrame, years: list[str], source_col: str) -> list[dict]:
    rep_mask = (
        df_master[source_col]
        .astype(str)
        .str.strip()
        .str.lower()
        .eq(REPORTED_FILTER_VALUE)
    )

    to_remove = set()

    for i, row in enumerate(struct):
        if row.get("type") != "subtotal":
            continue

        l3_label = row.get("tech_label")
        if is_blank_label(l3_label):
            continue

        j           = i - 1
        details_idx = []
        while j >= 0 and struct[j].get("type") == "detail" and struct[j].get("L3") == l3_label:
            details_idx.append(j)
            j -= 1
        details_idx = list(reversed(details_idx))

        if len(details_idx) != 1:
            continue

        detail_row = struct[details_idx[0]]
        l4_label   = detail_row.get("tech_label")

        if not is_blank_label(l4_label):
            continue

        l3_mask       = df_master["L3"].astype(str).str.strip().eq(str(l3_label).strip())
        l4_series     = df_master["L4"]
        l4_blank_mask = l4_series.isna() | l4_series.astype(str).str.strip().str.lower().isin(BLANK_TOKENS)

        mask = rep_mask & l3_mask & l4_blank_mask

        if mask.any():
            sums     = df_master.loc[mask, years].sum(numeric_only=True)
            all_zero = bool((sums.abs() < 1e-9).all())
        else:
            all_zero = True

        if all_zero:
            to_remove.add(details_idx[0])
            to_remove.add(i)

    if not to_remove:
        return struct

    return [r for k, r in enumerate(struct) if k not in to_remove]


def build_entity_sort_values(
    df: pd.DataFrame,
    entities: list[str],
    years: list[str],
    sort_metric_label: str,
) -> dict:
    rep_mask = (
        df[SOURCE_COL_NAME]
        .astype(str)
        .str.strip()
        .str.lower()
        .eq(REPORTED_FILTER_VALUE)
    )

    out = {}

    for entity in entities:
        ent_df   = df[rep_mask & df["Entity"].astype(str).str.strip().eq(str(entity).strip())].copy()

        l4_match = ent_df["L4"].astype(str).str.strip().eq(str(sort_metric_label).strip())
        l3_match = ent_df["L3"].astype(str).str.strip().eq(str(sort_metric_label).strip())

        if l4_match.any():
            metric_df = ent_df[l4_match].copy()
        else:
            metric_df = ent_df[l3_match].copy()

        if metric_df.empty:
            out[entity] = {
                "latest_year_value": float("-inf"),
                "total_value":       float("-inf"),
            }
            continue

        vals              = metric_df[years].apply(pd.to_numeric, errors="coerce").fillna(0)
        summed            = vals.sum(axis=0)
        latest_year_value = float(summed.iloc[-1]) if len(summed) > 0 else float("-inf")
        total_value       = float(summed.sum())

        out[entity] = {
            "latest_year_value": latest_year_value,
            "total_value":       total_value,
        }

    return out


def build_sorted_entities(df: pd.DataFrame, cfg: dict, years: list[str]) -> tuple[list[str], bool]:
    raw_entities = []
    for v in pd.unique(df["Entity"]):
        if not is_blank_entity(v):
            raw_entities.append(str(v).strip())

    ic_master_entity = IC_MASTER_ENTITY_DEFAULT

    has_ic = ic_master_entity in raw_entities
    if has_ic:
        idx_ic              = raw_entities.index(ic_master_entity)
        individual_entities = raw_entities[:idx_ic]
    else:
        individual_entities = raw_entities[:]

    if cfg["sort_by"] == "custom":
        order = [e for e in cfg.get("entity_order", []) if e in individual_entities]
        remaining = [e for e in individual_entities if e not in order]
        individual_entities = order + remaining
    else:
        sort_metric_label = cfg["sort_metric_labels"][cfg["sort_by"]]
        sort_values       = build_entity_sort_values(df, individual_entities, years, sort_metric_label)

        if any(v["latest_year_value"] == float("-inf") for v in sort_values.values()):
            missing_entities = [
                e for e, v in sort_values.items()
                if v["latest_year_value"] == float("-inf")
            ]
            raise ValueError(
                f"Für sort_by='{cfg['sort_by']}' konnte das technische Sort-Label "
                f"'{sort_metric_label}' nicht für alle Entities gefunden werden. "
                f"Betroffen: {missing_entities}"
            )

        original_index      = {e: i for i, e in enumerate(individual_entities)}
        individual_entities = sorted(
            individual_entities,
            key=lambda e: (
                -sort_values[e]["latest_year_value"],
                -sort_values[e]["total_value"],
                original_index[e],
            )
        )

    return individual_entities, has_ic


def build_blocks(cfg: dict, individual_entities: list[str], has_ic: bool) -> list[dict]:
    titles = cfg["display"]["titles"]

    blocks = []

    for e in individual_entities:
        blocks.append(
            {
                "kind":     "entity",
                "key":      e,
                "code":     e,
                "title":    e,
                "has_plpos": False,
            }
        )

    if len(individual_entities) > 0:
        blocks.append(
            {
                "kind":     "aggregated",
                "key":      "Aggregated",
                "code":     "Aggregated",
                "title":    titles["aggregated_title"],
                "has_plpos": False,
            }
        )

    if has_ic:
        blocks.append(
            {
                "kind":     "ic",
                "key":      titles["ic_display_name"],
                "code":     IC_MASTER_ENTITY_DEFAULT,
                "title":    titles["ic_display_name"],
                "has_plpos": False,
            }
        )

    if len(individual_entities) > 0 and has_ic:
        blocks.append(
            {
                "kind":     "consolidation",
                "key":      "Consolidation",
                "code":     "Consolidation",
                "title":    titles["consolidation_title"],
                "has_plpos": False,
            }
        )

    blocks.append(
        {
            "kind":     "difference",
            "key":      "Difference",
            "code":     "Difference",
            "title":    titles["difference_title"],
            "has_plpos": True,
        }
    )

    blocks.append(
        {
            "kind":     "fs",
            "key":      "Financial statements",
            "code":     "Financial statements",
            "title":    titles["financial_statements_title"],
            "has_plpos": True,
        }
    )

    current_col = FIRST_ENTITY_COL
    for b in blocks:
        b["startcol"]        = current_col
        extra                = 1 if b.get("has_plpos") else 0
        b["poscol"]          = current_col if extra == 1 else None
        b["year_startcol"]   = current_col + extra
        b["year_endcol"]     = b["year_startcol"] + len(YEARS) - 1
        b["spacer_col"]      = b["year_endcol"] + 1
        current_col          = b["spacer_col"] + 1

    return blocks


def build_entity_blocks(cfg: dict, entity: str) -> list[dict]:
    return build_entity_recon_blocks(
        entity,
        cfg["display"]["titles"],
        first_col=FIRST_ENTITY_COL,
        n_years=len(YEARS),
    )


def entity_block_lookup(blocks: list[dict], key: str) -> dict:
    return next(bb for bb in blocks if bb["key"] == key)


def ic_sign(block: dict) -> int:
    return -1 if block["kind"] == "ic" else 1


# =============================================
# SHEET WRITER
# =============================================

def write_pl_reconciliation_sheet(
    wb,
    sheet_name: str,
    *,
    cfg: dict,
    df: pd.DataFrame,
    row_structure: list,
    blocks: list[dict],
    individual_entities: list[str],
    fs_check_value_map: dict,
    is_entity_sheet: bool = False,
) -> None:
    if sheet_name in wb.sheetnames:
        del wb[sheet_name]
    ws = wb.create_sheet(sheet_name)

    # PROJECT & REPORT TITLE
    # ---------------------------------------------------
    pt           = ws.cell(PROJECT_TITLE_ROW, POS_COL, f"Project {cfg['project_name']}")
    pt.font      = FONT_PROJECT_TITLE
    pt.alignment = ALIGN_LEFT
    ws.row_dimensions[PROJECT_TITLE_ROW].height = PROJECT_TITLE_ROW_HEIGHT

    title_name = (
        individual_entities[0]
        if is_entity_sheet and individual_entities
        else cfg["company_name"]
    )

    st = ws.cell(
        SUBTITLE_ROW,
        POS_COL,
        f"{title_name} Reconciliation - Income statement"
    )
    st.font      = FONT_SUBTITLE
    st.alignment = ALIGN_LEFT

    title_text = (
        f"{title_name} | Reconciliation (trial balances) "
        f"{YEARS[0]} - {YEARS[-1]}"
    )
    tc           = ws.cell(TITLE_ROW, TITLE_COL, title_text)
    tc.font      = FONT_TITLE
    tc.alignment = ALIGN_LEFT

    # ---------------------------------------------------
    # MASTER RANGES
    # ---------------------------------------------------
    MASTER_START = 2
    MASTER_END   = len(df) + 1

    def master_range(colname):
        if colname not in df.columns:
            raise ValueError(f"Spalte '{colname}' wurde im Master nicht gefunden.")
        ltr = get_column_letter(df.columns.get_loc(colname) + 1)
        return f"{cfg['paths']['audit_master_sheet']}!${ltr}${MASTER_START}:${ltr}${MASTER_END}"

    src_rng = master_range(SOURCE_COL_NAME)
    ent_rng = master_range("Entity")
    l3_rng  = master_range("L3")
    l4_rng  = master_range("L4")

    # ---------------------------------------------------
    # BLOCK PLAN
    # ---------------------------------------------------
    LAST_USED_COL = blocks[-1]["spacer_col"] if blocks else FIRST_ENTITY_COL
    LAST_VALUE_COL = max(b["year_endcol"] for b in blocks) if blocks else POS_COL

    # ---------------------------------------------------
    # HEADERS
    # ---------------------------------------------------
    ws.cell(HEADER_ROW, MAP_START_COL + 0, "Reported").font = FONT_HEADER
    ws.cell(HEADER_ROW, MAP_START_COL + 1, "").font         = FONT_HEADER
    ws.cell(HEADER_ROW, MAP_START_COL + 2, "L3").font       = FONT_HEADER
    ws.cell(HEADER_ROW, MAP_START_COL + 3, "L4").font       = FONT_HEADER

    for c in range(MAP_START_COL, MAP_END_COL + 1):
        ws.cell(HEADER_ROW, c).alignment = ALIGN_LEFT

    ph           = ws.cell(HEADER_ROW, POS_COL, cfg["display"]["unit_label"])
    ph.font      = FONT_HEADER
    ph.alignment = ALIGN_LEFT

    for b in blocks:
        if b.get("has_plpos"):
            ws.merge_cells(
                start_row=BLOCK_TITLE_ROW,
                start_column=b["year_startcol"],
                end_row=BLOCK_TITLE_ROW,
                end_column=b["year_endcol"],
            )
            t           = ws.cell(BLOCK_TITLE_ROW, b["year_startcol"], b["title"])
            t.font      = FONT_BASE_BOLD
            t.alignment = ALIGN_CENTER
            if b.get("kind") not in {"difference", "fs"}:
                ws.cell(BLOCK_TITLE_ROW, b["poscol"]).value = None
                ws.cell(BLOCK_TITLE_ROW, b["poscol"]).fill = FILL_WHITE
        else:
            ws.merge_cells(
                start_row=BLOCK_TITLE_ROW,
                start_column=b["startcol"],
                end_row=BLOCK_TITLE_ROW,
                end_column=b["year_endcol"],
            )
            t           = ws.cell(BLOCK_TITLE_ROW, b["startcol"], b["title"])
            t.font      = FONT_BASE_BOLD
            t.alignment = ALIGN_CENTER

        if b.get("has_plpos"):
            ws.merge_cells(
                start_row=ENTITY_CODE_ROW,
                start_column=b["year_startcol"],
                end_row=ENTITY_CODE_ROW,
                end_column=b["year_endcol"],
            )
            ec = ws.cell(ENTITY_CODE_ROW, b["year_startcol"], b["code"])
        else:
            ws.merge_cells(
                start_row=ENTITY_CODE_ROW,
                start_column=b["startcol"],
                end_row=ENTITY_CODE_ROW,
                end_column=b["year_endcol"],
            )
            ec = ws.cell(ENTITY_CODE_ROW, b["startcol"], b["code"])

        ec.font      = FONT_BASE
        ec.alignment = ALIGN_CENTER

        if b.get("has_plpos"):
            hpos           = ws.cell(HEADER_ROW, b["poscol"], cfg["display"]["unit_label"])
            hpos.font      = FONT_HEADER
            hpos.alignment = ALIGN_LEFT

        for y_idx, year in enumerate(YEARS):
            h           = ws.cell(HEADER_ROW, b["year_startcol"] + y_idx, year)
            h.font      = FONT_HEADER
            h.alignment = ALIGN_RIGHT

        ws.column_dimensions[col_letter(b["spacer_col"])].width = SPACER_WIDTH

    # ---------------------------------------------------
    # WRITE ROWS
    # ---------------------------------------------------
    row_index = {}

    for i, r in enumerate(row_structure):
        r["_idx"]          = i
        excel_row          = DATA_START_ROW + i
        r["_excel_row"]    = excel_row
        row_index[r["tech_label"]] = excel_row

        ws.cell(excel_row, MAP_START_COL + 0, "Reported").font = FONT_MAPPING
        ws.cell(excel_row, MAP_START_COL + 1, "").font         = FONT_MAPPING
        ws.cell(excel_row, MAP_START_COL + 2, r["L3"]).font    = FONT_MAPPING
        ws.cell(excel_row, MAP_START_COL + 3, r["L4"]).font    = FONT_MAPPING

        for c in range(MAP_START_COL, MAP_END_COL + 1):
            ws.cell(excel_row, c).alignment = ALIGN_LEFT

        is_real_subtotal = (
            r["type"] == "subtotal"
            and any(
                rr["type"] == "detail" and rr["L3"] == r["tech_label"]
                for rr in row_structure
            )
        )

        pc           = ws.cell(excel_row, POS_COL, r["display_label"])
        pc.alignment = ALIGN_LEFT
        pc.font      = FONT_BASE_BOLD if (r["type"] == "total" or is_real_subtotal) else FONT_BASE

        for b in blocks:
            if not b.get("has_plpos"):
                continue
            pcell           = ws.cell(excel_row, b["poscol"])
            pcell.value     = f"={col_letter(POS_COL)}{excel_row}"
            pcell.alignment = ALIGN_LEFT
            pcell.font      = FONT_BASE_BOLD if (r["type"] == "total" or is_real_subtotal) else FONT_BASE

        if r["type"] == "detail":
            ws.row_dimensions[excel_row].outlineLevel = 1
            ws.row_dimensions[excel_row].hidden       = False
        else:
            ws.row_dimensions[excel_row].outlineLevel = 0
            ws.row_dimensions[excel_row].hidden       = False

    LAST_TABLE_ROW = DATA_START_ROW + len(row_structure) - 1

    # ---------------------------------------------------
    # VALUES
    # ---------------------------------------------------
    map_reported_ref_col = col_letter(MAP_START_COL + 0)
    map_l3_ref_col       = col_letter(MAP_START_COL + 2)
    map_l4_ref_col       = col_letter(MAP_START_COL + 3)

    for b in blocks:
        if b["kind"] in {"entity", "ic"}:
            entity_ref = f"${col_letter(b['startcol'])}${ENTITY_CODE_ROW}"
        else:
            entity_ref = None

        sign = ic_sign(b)

        for r in row_structure:
            excel_row = r["_excel_row"]

            for y_idx, year in enumerate(YEARS):
                cell               = ws.cell(excel_row, b["year_startcol"] + y_idx)
                cell.alignment     = ALIGN_RIGHT
                cell.number_format = NUM_FMT_INT

                is_real_subtotal = (
                    r["type"] == "subtotal"
                    and any(
                        rr["type"] == "detail" and rr["L3"] == r["tech_label"]
                        for rr in row_structure
                    )
                )

                cell.font = FONT_BASE_BOLD if (r["type"] == "total" or is_real_subtotal) else FONT_BASE

                if b["kind"] in {"entity", "ic"}:
                    if r["type"] == "detail":
                        sum_rng  = master_range(year)
                        l3_crit  = f"${map_l3_ref_col}${excel_row}"
                        l4_crit  = f"${map_l4_ref_col}${excel_row}"
                        rep_crit = f"${map_reported_ref_col}${excel_row}"

                        cell.value = (
                            f"={sign}*SUMIFS({sum_rng},"
                            f"{ent_rng},{entity_ref},"
                            f"{src_rng},{rep_crit},"
                            f"{l3_rng},{l3_crit},"
                            f"{l4_rng},{l4_crit}"
                            f")/1000"
                        )

                    elif r["type"] == "subtotal":
                        idx = r["_idx"]
                        k   = idx - 1
                        while k >= 0 and row_structure[k]["type"] == "detail" and row_structure[k]["L3"] == r["tech_label"]:
                            k -= 1
                        detail_start_idx = k + 1
                        detail_end_idx   = idx - 1

                        col_l = col_letter(b["year_startcol"] + y_idx)
                        if detail_end_idx >= detail_start_idx:
                            start_row  = DATA_START_ROW + detail_start_idx
                            end_row    = DATA_START_ROW + detail_end_idx
                            cell.value = f"=SUM({col_l}{start_row}:{col_l}{end_row})"
                        else:
                            sum_rng  = master_range(year)
                            l3_crit  = f"${map_l3_ref_col}${excel_row}"
                            l4_crit  = f"${map_l4_ref_col}${excel_row}"
                            rep_crit = f"${map_reported_ref_col}${excel_row}"

                            cell.value = (
                                f"={sign}*SUMIFS({sum_rng},"
                                f"{ent_rng},{entity_ref},"
                                f"{src_rng},{rep_crit},"
                                f"{l3_rng},{l3_crit},"
                                f"{l4_rng},{l4_crit}"
                                f")/1000"
                            )

                    else:  # total
                        refs = []
                        for comp in r.get("components", []):
                            if comp in row_index:
                                refs.append(f"{col_letter(b['year_startcol'] + y_idx)}{row_index[comp]}")
                        cell.value = f"=SUM({','.join(refs)})" if refs else 0

                elif b["kind"] == "aggregated":
                    refs = []
                    for e in individual_entities:
                        bb = entity_block_lookup(blocks, e)
                        refs.append(f"{col_letter(bb['year_startcol'] + y_idx)}{excel_row}")
                    cell.value = f"=SUM({','.join(refs)})" if refs else 0

                elif b["kind"] == "consolidation":
                    agg_block = entity_block_lookup(blocks, "Aggregated")
                    ic_block  = entity_block_lookup(blocks, cfg["display"]["titles"]["ic_display_name"])
                    agg_ref   = f"{col_letter(agg_block['year_startcol'] + y_idx)}{excel_row}"
                    ic_ref    = f"{col_letter(ic_block['year_startcol'] + y_idx)}{excel_row}"
                    cell.value = f"={agg_ref}+{ic_ref}"

                elif b["kind"] == "fs":
                    cell.value = None

                elif b["kind"] == "difference":
                    fs_block = entity_block_lookup(blocks, "Financial statements")
                    fs_ref = f"{col_letter(fs_block['year_startcol'] + y_idx)}{excel_row}"
                    if any(bb["kind"] == "consolidation" for bb in blocks):
                        con_block = entity_block_lookup(blocks, "Consolidation")
                        con_ref = f"{col_letter(con_block['year_startcol'] + y_idx)}{excel_row}"
                        cell.value = f"={fs_ref}-{con_ref}"
                    elif is_entity_sheet:
                        ent_block = next(bb for bb in blocks if bb["kind"] == "entity")
                        ent_ref = f"{col_letter(ent_block['year_startcol'] + y_idx)}{excel_row}"
                        cell.value = f"={fs_ref}-{ent_ref}"
                    else:
                        cell.value = None

    LAST_CONTENT_ROW = LAST_TABLE_ROW

    # ---------------------------------------------------
    # COLUMN WIDTHS
    # ---------------------------------------------------
    ws.column_dimensions[col_letter(POS_COL)].width = 32

    for b in blocks:
        if b.get("has_plpos"):
            ws.column_dimensions[col_letter(b["poscol"])].width = REPEAT_POS_WIDTH
        for y_idx in range(len(YEARS)):
            ws.column_dimensions[col_letter(b["year_startcol"] + y_idx)].width = VALUE_COL_WIDTH
        ws.column_dimensions[col_letter(b["spacer_col"])].width = SPACER_WIDTH

    # ---------------------------------------------------
    # FILLS
    # ---------------------------------------------------
    FILL_END_ROW = LAST_CONTENT_ROW + FILL_PADDING_ROWS
    FILL_END_COL = LAST_USED_COL + FILL_PADDING_COLS
    sheet_layout = LAYOUT_PL

    for rr in range(1, FILL_END_ROW + 1):
        if rr == PROJECT_TITLE_ROW:
            continue
        ws.row_dimensions[rr].height = ROW_HEIGHT

    paint_grey_white_canvas(
        ws,
        sheet_layout,
        last_row=FILL_END_ROW,
        last_col=LAST_USED_COL,
    )

    for rr in range(DATA_START_ROW, LAST_TABLE_ROW + 1):
        for cc in range(MAP_START_COL, MAP_END_COL + 1):
            ws.cell(rr, cc).font = FONT_MAPPING

    for cc in range(MAP_START_COL, MAP_END_COL + 1):
        ws.cell(HEADER_ROW, cc).font = FONT_HEADER
    ws.cell(HEADER_ROW, POS_COL).font = FONT_HEADER

    spacer_cols = {b["spacer_col"] for b in blocks}

    for r in row_structure:
        if r["type"] not in ("subtotal", "total"):
            continue
        excel_row = r["_excel_row"]
        is_real_subtotal = (
            r["type"] == "subtotal"
            and any(
                rr["type"] == "detail" and rr["L3"] == r["tech_label"]
                for rr in row_structure
            )
        )
        if r["type"] == "subtotal" and not is_real_subtotal:
            continue
        for cc in range(POS_COL, LAST_VALUE_COL + 1):
            cell = ws.cell(excel_row, cc)
            cell.fill = FILL_SUBTOTAL

    # ---------------------------------------------------
    # BORDERS
    # ---------------------------------------------------
    top_border_only = set(cfg["pl_config"]["top_border_only"])
    top_and_bottom  = set(cfg["pl_config"]["top_and_bottom"])

    for r in row_structure:
        tech_label = r["tech_label"]
        if tech_label not in top_border_only and tech_label not in top_and_bottom:
            continue

        excel_row     = row_index[tech_label]
        border_to_set = (
            BORDER_SUBTOTAL_TOP_BOTTOM
            if tech_label in top_and_bottom
            else BORDER_SUBTOTAL_TOP
        )

        apply_databook_row_border_band(
            ws,
            excel_row,
            LAYOUT_PL,
            border_to_set,
            last_used_col=LAST_VALUE_COL,
            skip_cols=spacer_cols,
        )

    apply_recon_portfolio_layout(
        ws,
        blocks=blocks,
        pos_col=POS_COL,
        header_row=HEADER_ROW,
        block_title_row=BLOCK_TITLE_ROW,
        entity_code_row=ENTITY_CODE_ROW,
        last_used_col=LAST_USED_COL,
        spacer_cols=spacer_cols,
        tech_layout=sheet_layout,
        visible_block_kinds=frozenset({"entity", "difference", "fs"})
        if is_entity_sheet
        else frozenset({"consolidation", "difference", "fs"}),
    )

    ws.row_dimensions[HEADER_ROW].height = ROW_HEIGHT
    ws.row_dimensions[BLOCK_TITLE_ROW].height = ROW_HEIGHT

    for sc in spacer_cols:
        for rr in range(1, LAST_TABLE_ROW + 120):
            ws.cell(rr, sc).border = Border()
            ws.cell(rr, sc).fill = FILL_WHITE

    if cfg.get("show_fs_check", True):
        _check_rows = check_row_groups_after_table(LAST_TABLE_ROW, [2])
        CHECK_TITLE_ROW, CHECK_DIFF_ROW = _check_rows
        net_label = normalize_technical_pl_key(NET_RESULT_TECH_KEY)
        NET_RESULT_ROW = row_index.get(net_label, LAST_TABLE_ROW)
        write_fs_check_section(
            ws,
            blocks=blocks,
            years=YEARS,
            source_row=CHECK_TITLE_ROW,
            delta_row=CHECK_DIFF_ROW,
            anchor_row=NET_RESULT_ROW,
            pos_col=POS_COL,
            block_kinds=frozenset({"entity"}) if is_entity_sheet else frozenset({"entity", "consolidation"}),
            yellow_block_kinds=frozenset({"entity"})
            if is_entity_sheet
            else frozenset({"entity", "aggregated", "ic", "consolidation"}),
            spacer_cols=spacer_cols,
            source_values=fs_check_value_map if any(fs_check_value_map.values()) else None,
            source_label=CHECK_TITLE_LABEL,
            delta_label=CHECK_DIFF_LABEL,
            num_fmt=NUM_FMT_INT,
        )
        collapse_check_portfolio(ws, [CHECK_TITLE_ROW, CHECK_DIFF_ROW])


# =============================================
# MAIN
# =============================================

def main():
    global YEARS

    cfg = load_config_from_argv()

    df = load_source_df(cfg)
    YEARS = detect_year_columns(df)
    print(f"Verwende Source-Spalte für Reported/Adjusted: {SOURCE_COL_NAME}")

    map_df        = load_mapping_df(cfg)
    l3_order      = l3_order_from_mapping(map_df)
    row_structure = build_pl_row_structure(
        df,
        l3_order,
        cfg,
        source_col=SOURCE_COL_NAME,
        display_label_fn=lambda lbl: technical_to_display_label(lbl, cfg),
    )
    row_structure = apply_totals_to_row_structure(row_structure, cfg)
    row_structure = prune_nan_zero_blocks(row_structure, df, YEARS, source_col=SOURCE_COL_NAME)
    row_structure = prune_zero_value_rows(
        row_structure,
        df,
        YEARS,
        source_col=SOURCE_COL_NAME,
    )
    row_structure = prune_row_structure(row_structure)

    individual_entities, has_ic = build_sorted_entities(df, cfg, YEARS)
    group_blocks = build_blocks(cfg, individual_entities, has_ic)
    fs_check_value_map = build_fs_check_value_map(cfg["fs_check_values"], individual_entities, YEARS)

    target_file = cfg["paths"]["target_file"]
    target_path = Path(target_file).resolve()
    report_sheet = cfg["paths"]["report_sheet"]
    audit_sheet = cfg["paths"]["audit_master_sheet"]

    if target_path.is_file():
        wb = load_workbook(target_file)
    else:
        wb = Workbook()
        if wb.sheetnames:
            wb.remove(wb.active)

    remove_stale_entity_recon_sheets(wb, individual_entities)

    if audit_sheet not in wb.sheetnames:
        ws_master = wb.create_sheet(audit_sheet)
        for c_idx, colname in enumerate(df.columns, start=1):
            ws_master.cell(1, c_idx, colname)
        for r_idx, row in enumerate(df.itertuples(index=False, name=None), start=2):
            for c_idx, val in enumerate(row, start=1):
                ws_master.cell(r_idx, c_idx, val)

    refresh_master_sheet(wb, df, audit_sheet)

    write_pl_reconciliation_sheet(
        wb,
        report_sheet,
        cfg=cfg,
        df=df,
        row_structure=row_structure,
        blocks=group_blocks,
        individual_entities=individual_entities,
        fs_check_value_map=fs_check_value_map,
        is_entity_sheet=False,
    )

    for entity in individual_entities:
        entity_sheet = sanitize_entity_recon_sheet_name(entity, "pl")
        entity_blocks = build_entity_blocks(cfg, entity)
        entity_fs_map = build_fs_check_value_map_for_entity(
            cfg["fs_check_values"], entity, individual_entities, YEARS,
        )
        write_pl_reconciliation_sheet(
            wb,
            entity_sheet,
            cfg=cfg,
            df=df,
            row_structure=row_structure,
            blocks=entity_blocks,
            individual_entities=[entity],
            fs_check_value_map=entity_fs_map,
            is_entity_sheet=True,
        )

    # ---------------------------------------------------
    # SAVE
    # ---------------------------------------------------
    wb.save(target_file)
    print(f"Fertig. Datei gespeichert unter: {target_file}")

    if cfg["paths"].get("append_to_master"):
        from copy import copy

        master_file = str(cfg["paths"].get("master_file") or target_file)
        if Path(master_file).resolve() != target_path:
            master_wb = load_workbook(master_file)
            if report_sheet in master_wb.sheetnames:
                del master_wb[report_sheet]
            recon_wb = load_workbook(target_file)
            src = recon_wb[report_sheet]
            dest = master_wb.create_sheet(report_sheet)
            for row in src.iter_rows():
                for cell in row:
                    dest[cell.coordinate].value = cell.value
                    if cell.has_style:
                        dest[cell.coordinate].font = copy(cell.font)
                        dest[cell.coordinate].fill = copy(cell.fill)
                        dest[cell.coordinate].border = copy(cell.border)
                        dest[cell.coordinate].alignment = copy(cell.alignment)
                        dest[cell.coordinate].number_format = cell.number_format
            master_wb.save(master_file)
            print(f"Sheet '{report_sheet}' appended to {master_file}")

    print(f"Sortierung erfolgte über: {cfg['sort_by']}")
    print(f"Verwendete Jahre: {YEARS}")


if __name__ == "__main__":
    main()
