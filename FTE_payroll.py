"""FTE & payroll development — personnel tables to grouped period report (FA-style layout)."""
from __future__ import annotations

import json
import re
import sys
from collections import defaultdict
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

BASE_DIR = Path(__file__).resolve().parent
SCRIPTS_DIR = BASE_DIR / "scripts"
BACKEND_DIR = BASE_DIR / "backend"
if str(BASE_DIR) not in sys.path:
    sys.path.insert(0, str(BASE_DIR))
if str(SCRIPTS_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPTS_DIR))
if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))

from databook_periods import master_ytd_label, ytd_reporting_fy_end_year  # noqa: E402

from fixed_assets_rollf import (  # noqa: E402
    RowSpec,
    _FY_GRID_RE,
    _YTD_GRID_RE,
    _resolve_sheet_name,
    build_row_specs,
)
from report_row_layout import (  # noqa: E402
    compute_l4_sort_metric,
    discover_l4_under_l3,
    period_columns_for_sort,
    sort_l4_labels,
)
from funktionssammlung import (  # noqa: E402
    _get_sheet_header_map,
    _normalize_header_name,
    apply_filters,
    build_output_file_path,
    build_sumifs_formula_body,
    ensure_output_writable,
    open_session_workbook,
    replace_workbook_sheet,
    resolve_column_names_to_df,
    source_range_ref,
    write_source_df_to_ws,
)
from gst_excel_theme import THEME  # noqa: E402
from databook_excel_layout import FONT_MAPPING  # noqa: E402

PROJECT_TITLE_ROW = 1
SUBTITLE_ROW = 2
TABLE_TITLE_ROW = 6
HEADER_ROW_7 = 7
HEADER_ROW = 8
DATA_START_ROW = 9

MAP_COL_WIDTH = 12
VALUE_COL_WIDTH = 8.0
LABEL_COL_WIDTH = 32.0
ROW_HEIGHT = 12
NUM_FMT_INT = "#,##0;(#,##0);-"
NUM_FMT_K = "#,##0;(#,##0);-"
UNIT_LABEL = "kEUR"
INDENT = "    "
_FTE_HELPER_COL = "__FTE_WT__"
_PL_L2 = "Expense"
_PL_L3 = "Personnel expenses"
_PL_EXCLUDE_L4 = "Wages & salaries"
_PL_REPORTED_VALUE = "Reported"
_PL_SOURCE_COL = "L6"
_ZERO_TOL = 1e-6

FONT_BASE = THEME.font_base
FONT_BOLD = THEME.font_bold
FONT_HEADER = THEME.font_header
FONT_PROJECT = Font(name=THEME.font_name, size=THEME.font_size_title, color=THEME.text_brand_title)
FONT_SUBTITLE = Font(name=THEME.font_name, size=THEME.font_size_subtitle, color=THEME.text_brand_title)
FILL_GREY = THEME.fill_tech
FILL_HEADER = THEME.fill_header
FILL_SUBTOTAL = THEME.fill_subtotal
FILL_WHITE = THEME.fill_white
FILL_PERIOD = THEME.fill_period
ALIGN_LEFT = Alignment(horizontal="left", vertical="center")
ALIGN_RIGHT = Alignment(horizontal="right", vertical="center")
ALIGN_CENTER = Alignment(horizontal="center", vertical="center")
BORDER_SUBTOTAL_TOP = THEME.border_subtotal_top
_BORDER_SIDE = Side(style="medium", color=THEME.border_strong)
BORDER_SUBTOTAL_BOTH = Border(top=_BORDER_SIDE, bottom=_BORDER_SIDE)
BORDER_HEADER_BOTTOM = Border(bottom=_BORDER_SIDE)


@dataclass(frozen=True)
class FteLayout:
    helper_start_col: int
    helper_end_col: int
    group_crit_cols: tuple[int, ...]
    pl_reported_col: int
    pl_l2_col: int
    pl_l3_col: int
    pl_l4_col: int
    pos_col: int
    first_data_col: int


@dataclass
class PeriodColumn:
    col_idx: int
    header: str
    period_idx: int


@dataclass(frozen=True)
class PeriodSpec:
    label: str
    file_path: str
    sheet_name: str = ""


def fte_column_layout(num_group_cols: int) -> FteLayout:
    n = max(0, min(2, int(num_group_cols)))
    col = 1
    group_cols: list[int] = []
    col += 1
    for i in range(n):
        if i > 0:
            col += 1
        group_cols.append(col)
        col += 1
    pl_reported_col = col
    pl_l2_col = col + 1
    pl_l3_col = col + 2
    pl_l4_col = col + 3
    col += 4
    col += 1  # trailing spacer before POS (databook-style)
    helper_end = col - 1
    pos_col = col
    return FteLayout(
        helper_start_col=1,
        helper_end_col=helper_end,
        group_crit_cols=tuple(group_cols),
        pl_reported_col=pl_reported_col,
        pl_l2_col=pl_l2_col,
        pl_l3_col=pl_l3_col,
        pl_l4_col=pl_l4_col,
        pos_col=pos_col,
        first_data_col=pos_col + 1,
    )


def fte_mapping_column_indices(layout: FteLayout) -> list[int]:
    cols = list(layout.group_crit_cols)
    cols.extend(
        [
            layout.pl_reported_col,
            layout.pl_l2_col,
            layout.pl_l3_col,
            layout.pl_l4_col,
        ]
    )
    return sorted(set(cols))


def fte_spacer_column_indices(layout: FteLayout) -> list[int]:
    mapping = set(fte_mapping_column_indices(layout))
    return [
        cc
        for cc in range(layout.helper_start_col + 1, layout.pos_col)
        if cc not in mapping
    ]


def fte_period_header(label: str, *, cfg: dict | None = None) -> str:
    """Map upload grid label (FY2023 / YTD2025) to master period header (FY23A / YTD25A)."""
    s = str(label or "").strip()
    upper = s.upper()
    if re.match(r"^FY\d{2}A$", upper):
        return upper
    if re.match(r"^YTD\d{2}A$", upper):
        return upper
    m_fy = _FY_GRID_RE.match(s)
    if m_fy:
        fy_year = int(s[2:])
        if fy_year < 100:
            fy_year += 2000
        if cfg:
            try:
                fy_end_m = int(cfg.get("fy_end_month") or 12)
                fy_end_d = int(cfg.get("fy_end_day") or 31)
                ytd_fy = ytd_reporting_fy_end_year(cfg.get("ltm_month"), fy_end_m, fy_end_d)
                if ytd_fy is not None and int(ytd_fy) == fy_year:
                    return master_ytd_label(ytd_fy)
            except (TypeError, ValueError):
                pass
        return f"FY{s[-2:]}A"
    m_ytd = _YTD_GRID_RE.match(s)
    if m_ytd:
        return f"YTD{s[-2:]}A"
    return s


def resolve_period_header(period: dict, cfg: dict) -> str:
    explicit = str(period.get("header") or "").strip()
    if explicit:
        return explicit
    return fte_period_header(str(period.get("label") or ""), cfg=cfg)


def fte_source_sheet_name(period_label: str, *, header: str | None = None) -> str:
    hdr = str(header or "").strip() or fte_period_header(period_label)
    return f"__SOURCE__{hdr}"[:31]


def col_letter(idx: int) -> str:
    return get_column_letter(idx)


def normalize_employment_rate(value: float) -> float:
    v = float(value or 0)
    if abs(v) > 1.5:
        return v / 100.0
    return v


def fte_row_weight_decimal(employment_rate: float, months_sum: float) -> float:
    rate = normalize_employment_rate(employment_rate)
    months = max(0.0, float(months_sum or 0))
    return rate * (months / 12.0)


def fte_row_weight(employment_rate: float, months_sum: float) -> int:
    return int(round(fte_row_weight_decimal(employment_rate, months_sum)))


def normalize_config(cfg: dict) -> dict:
    out = dict(cfg)
    cols = dict(out.get("columns") or {})
    for key in ("employment", "months_sum"):
        if not str(cols.get(key) or "").strip():
            raise ValueError(f"columns.{key} is required")
    payroll_cols = cols.get("payroll_cols")
    if not isinstance(payroll_cols, list) or not payroll_cols:
        raise ValueError("columns.payroll_cols is required")
    cols["payroll_cols"] = [str(x).strip() for x in payroll_cols if str(x).strip()]
    out["columns"] = cols
    out["group_cols"] = [str(c).strip() for c in (out.get("group_cols") or []) if str(c).strip()][:2]
    periods = out.get("periods")
    if not isinstance(periods, list) or not periods:
        raise ValueError("periods[] is required")
    norm_periods: list[dict] = []
    for p in periods:
        if not isinstance(p, dict):
            continue
        label = str(p.get("label") or "").strip()
        fp = str(p.get("file_path") or "").strip()
        if not label or not fp:
            continue
        entry = {
            "label": label,
            "file_path": fp,
            "sheet_name": str(p.get("sheet_name") or "").strip(),
        }
        hdr = str(p.get("header") or "").strip()
        if hdr:
            entry["header"] = hdr
        norm_periods.append(entry)
    if not norm_periods:
        raise ValueError("periods[] has no valid entries")
    out["periods"] = norm_periods
    out["header_row"] = int(out.get("header_row") or 0)
    out["unit_label"] = str(out.get("unit_label") or UNIT_LABEL)
    scale = out.get("amount_scale")
    out["amount_scale"] = float(scale) if scale not in (None, "", 0) else 1000.0
    out["sheet_name"] = str(out.get("sheet_name") or "FTE development")
    out["formula_mode"] = bool(out.get("formula_mode", True))
    out["master_pl_path"] = str(out.get("master_pl_path") or "").strip()
    out["master_pl_sheet"] = str(out.get("master_pl_sheet") or "Master_PL").strip()
    raw_filters = out.get("filters", {}) or {}
    if isinstance(raw_filters, list):
        filters = {"enabled": len(raw_filters) > 0, "rules": raw_filters}
    else:
        filters = dict(raw_filters)
    filters.setdefault("enabled", False)
    filters.setdefault("rules", [])
    out["filters"] = filters
    out.setdefault("use_session_workbook", True)
    return out


def load_period_dataframe(spec: PeriodSpec, header_row: int) -> pd.DataFrame:
    sheet = _resolve_sheet_name(spec.file_path, spec.sheet_name)
    return pd.read_excel(spec.file_path, sheet_name=sheet, header=header_row)


def _divide_by_scale(cfg: dict) -> bool:
    return float(cfg.get("amount_scale") or 1000) == 1000.0


def _fte_wt_cell_formula(emp_cell: str, months_cell: str) -> str:
    return f"=IF(ABS({emp_cell})>1.5,{emp_cell}/100,{emp_cell})*{months_cell}/12"


def _apply_fte_wt_formulas(
    ws,
    employment_header: str,
    months_header: str,
) -> None:
    """Append __FTE_WT__ as row formulas on the source sheet (decimal, unrounded)."""
    hmap = _get_sheet_header_map(ws)
    emp_idx = hmap.get(_normalize_header_name(employment_header))
    months_idx = hmap.get(_normalize_header_name(months_header))
    if not emp_idx or not months_idx:
        raise ValueError(
            f"Source sheet missing employment/months columns: {employment_header!r}, {months_header!r}"
        )
    wt_norm = _normalize_header_name(_FTE_HELPER_COL)
    wt_idx = hmap.get(wt_norm)
    if not wt_idx:
        wt_idx = (ws.max_column or 0) + 1
        ws.cell(1, wt_idx, _FTE_HELPER_COL)
    emp_l = col_letter(emp_idx)
    mon_l = col_letter(months_idx)
    last_row = ws.max_row or 1
    for r in range(2, last_row + 1):
        cell = ws.cell(r, wt_idx)
        cell.value = _fte_wt_cell_formula(f"{emp_l}{r}", f"{mon_l}{r}")
        cell.number_format = "0.0000"


def aggregate_period_keys(
    df: pd.DataFrame,
    cfg: dict,
) -> tuple[set[tuple[str, ...]], pd.DataFrame, dict[str, str]]:
    group_cols = list(cfg.get("group_cols") or [])
    cols = cfg["columns"]
    needed = group_cols + [cols["employment"], cols["months_sum"], *cols["payroll_cols"]]
    resolved = resolve_column_names_to_df(df, needed)
    g_resolved = resolved[: len(group_cols)]
    emp_col = resolved[len(group_cols)]
    months_col = resolved[len(group_cols) + 1]
    payroll_resolved = resolved[len(group_cols) + 2 :]

    work = apply_filters(df.copy(), cfg)
    col_map = {
        "employment": emp_col,
        "months_sum": months_col,
        "payroll_cols": payroll_resolved,
    }
    for pc in payroll_resolved:
        if pc in work.columns:
            work[pc] = pd.to_numeric(work[pc], errors="coerce").fillna(0.0)

    keys: set[tuple[str, ...]] = set()
    if group_cols:
        for key, _grp in work.groupby(g_resolved, dropna=False):
            if not isinstance(key, tuple):
                key = (key,)
            keys.add(tuple("" if pd.isna(v) else str(v).strip() for v in key))
    else:
        keys.add(())
    return keys, work, col_map


def build_period_columns(
    periods: list[dict],
    cfg: dict,
    *,
    first_data_col: int,
) -> list[PeriodColumn]:
    cols: list[PeriodColumn] = []
    for i, period in enumerate(periods):
        cols.append(
            PeriodColumn(
                col_idx=first_data_col + i,
                header=resolve_period_header(period, cfg),
                period_idx=i,
            )
        )
    return cols


def _group_sumifs_pairs(
    source_sheet: str,
    header_map: dict[str, int],
    group_cols: list[str],
    row_idx: int,
    *,
    group_crit_cols: tuple[int, ...],
) -> list[tuple[str, str]]:
    pairs: list[tuple[str, str]] = []
    for i, gc in enumerate(group_cols):
        if i >= len(group_crit_cols):
            break
        col_idx = group_crit_cols[i]
        rng = source_range_ref(source_sheet, header_map, gc)
        crit = f"${col_letter(col_idx)}{row_idx}"
        pairs.append((rng, crit))
    return pairs


def _fte_sumifs_formula(
    source_sheet: str,
    header_map: dict[str, int],
    row_idx: int,
    group_cols: list[str],
    cfg: dict,
    *,
    group_crit_cols: tuple[int, ...],
) -> str:
    if _normalize_header_name(_FTE_HELPER_COL) not in header_map:
        raise ValueError(f"Source sheet missing {_FTE_HELPER_COL}")
    amount_rng = source_range_ref(source_sheet, header_map, _FTE_HELPER_COL)
    base_pairs = _group_sumifs_pairs(
        source_sheet, header_map, group_cols, row_idx, group_crit_cols=group_crit_cols
    )
    body = build_sumifs_formula_body(amount_rng, base_pairs, [[]], divide_by_1000=False)
    return f"=ROUND({body},0)"


def _payroll_sumifs_formula(
    source_sheet: str,
    header_map: dict[str, int],
    payroll_headers: list[str],
    row_idx: int,
    group_cols: list[str],
    cfg: dict,
    *,
    group_crit_cols: tuple[int, ...],
) -> str:
    parts: list[str] = []
    base_pairs = _group_sumifs_pairs(
        source_sheet, header_map, group_cols, row_idx, group_crit_cols=group_crit_cols
    )
    for ph in payroll_headers:
        try:
            rng = source_range_ref(source_sheet, header_map, ph)
        except ValueError:
            continue
        body = build_sumifs_formula_body(
            rng,
            base_pairs,
            [[]],
            divide_by_1000=_divide_by_scale(cfg),
        )
        parts.append(body)
    if not parts:
        return "=0"
    inner = "+".join(parts)
    return f"=-({inner})"


def _parent_sum_formula(child_rows: list[int], col_idx: int) -> str:
    if not child_rows:
        return "=0"
    refs = [f"{col_letter(col_idx)}{r}" for r in child_rows]
    if len(refs) == 1:
        return f"={refs[0]}"
    return f"=SUM({','.join(refs)})"


def _chunk_sum_formula(rows: list[int], col_idx: int) -> str:
    if not rows:
        return "=0"
    refs = [f"{col_letter(col_idx)}{r}" for r in rows]
    if len(refs) <= 12:
        return f"=SUM({','.join(refs)})"
    chunks = [refs[i : i + 12] for i in range(0, len(refs), 12)]
    parts = [f"SUM({','.join(c)})" for c in chunks]
    return f"=({' + '.join(parts)})"


def _load_master_pl_frame(master_path: str, sheet_name: str) -> pd.DataFrame | None:
    if not master_path or not Path(master_path).is_file():
        return None
    return pd.read_excel(master_path, sheet_name=sheet_name, header=0, engine="openpyxl")


def _pl_source_col(df: pd.DataFrame) -> str:
    if _PL_SOURCE_COL in df.columns:
        return _PL_SOURCE_COL
    if "L5" in df.columns:
        return "L5"
    raise ValueError("Master_PL missing L5/L6 source column")


def build_pl_personnel_rows(master_path: str, sheet_name: str, cfg: dict) -> list[dict[str, str]]:
    """Personnel expense L4 lines (excl. W&S), amount-sorted with Other last."""
    df = _load_master_pl_frame(master_path, sheet_name)
    if df is None:
        return [{"display_label": "Other", "L2": _PL_L2, "L3": _PL_L3, "L4": "Other"}]
    source_col = _pl_source_col(df)
    l4_labels = discover_l4_under_l3(df, _PL_L3, source_col, reported_value=_PL_REPORTED_VALUE)
    l4_labels = [
        x for x in l4_labels if str(x).strip().lower() != _PL_EXCLUDE_L4.lower()
    ]
    if not l4_labels:
        return [{"display_label": "Other", "L2": _PL_L2, "L3": _PL_L3, "L4": "Other"}]
    period_cols = period_columns_for_sort(df, cfg)
    metrics = {
        l4: compute_l4_sort_metric(
            df, _PL_L3, l4, period_cols, source_col, reported_value=_PL_REPORTED_VALUE
        )
        for l4 in l4_labels
    }
    sorted_l4s = sort_l4_labels(l4_labels, metrics)
    others = [x for x in sorted_l4s if str(x).strip().lower() == "other"]
    non_others = [x for x in sorted_l4s if str(x).strip().lower() != "other"]
    ordered = non_others + others
    return [
        {"display_label": l4, "L2": _PL_L2, "L3": _PL_L3, "L4": l4}
        for l4 in ordered
    ]


def load_pl_personnel_lines(master_path: str, sheet_name: str = "Master_PL") -> list[tuple[str, int]]:
    """Legacy helper — (L4 label, excel_row) for tests."""
    rows = build_pl_personnel_rows(master_path, sheet_name, {})
    if not master_path or not Path(master_path).is_file():
        return [(r["display_label"], 0) for r in rows]
    df = _load_master_pl_frame(master_path, sheet_name)
    if df is None:
        return []
    out: list[tuple[str, int]] = []
    for row in rows:
        l4 = row["L4"]
        match = df[
            (df["L3"].astype(str).str.strip() == _PL_L3)
            & (df["L4"].astype(str).str.strip() == str(l4).strip())
        ]
        if not match.empty:
            out.append((l4, int(match.index[0]) + 2))
    return out


def _master_pl_sheet_prefix(cfg: dict, wb) -> str:
    sheet = cfg.get("master_pl_sheet", "Master_PL")
    if sheet in wb.sheetnames:
        safe = sheet.replace("'", "''")
        return f"'{safe}'"
    mp = str(cfg.get("master_pl_path") or "").strip()
    if mp and Path(mp).is_file():
        fname = Path(mp).name.replace("'", "''")
        return f"'[{fname}]{sheet}'"
    safe = sheet.replace("'", "''")
    return f"'{safe}'"


def _master_pl_column_ranges(
    master_path: str,
    sheet_name: str,
    prefix: str,
) -> dict[str, str]:
    df = _load_master_pl_frame(master_path, sheet_name)
    if df is None:
        return {}
    start, end = 2, len(df) + 1
    ranges: dict[str, str] = {}
    for col_name in ("L2", "L3", "L4", _pl_source_col(df)):
        ltr = col_letter(int(df.columns.get_loc(col_name)) + 1)
        ranges[col_name] = f"{prefix}!${ltr}${start}:${ltr}${end}"
    return ranges


def _pl_sumifs_formula(
    master_path: str,
    sheet_name: str,
    period_header: str,
    row_idx: int,
    layout: FteLayout,
    *,
    sheet_prefix: str,
    col_ranges: dict[str, str],
) -> str:
    period_col = _master_period_col_letter(master_path, sheet_name, period_header)
    if not period_col or not col_ranges:
        return "=0"
    df = _load_master_pl_frame(master_path, sheet_name)
    end_row = (len(df) + 1) if df is not None else 2
    sum_rng = f"{sheet_prefix}!${period_col}$2:${period_col}${end_row}"
    src_rng = col_ranges.get("L6") or col_ranges.get("L5", "")
    l2_rng = col_ranges.get("L2", "")
    l3_rng = col_ranges.get("L3", "")
    l4_rng = col_ranges.get("L4", "")
    if not all([src_rng, l2_rng, l3_rng, l4_rng]):
        return "=0"
    rep_crit = f"${col_letter(layout.pl_reported_col)}${row_idx}"
    l2_crit = f"${col_letter(layout.pl_l2_col)}${row_idx}"
    l3_crit = f"${col_letter(layout.pl_l3_col)}${row_idx}"
    l4_crit = f"${col_letter(layout.pl_l4_col)}${row_idx}"
    return (
        f"=SUMIFS({sum_rng},"
        f"{src_rng},{rep_crit},"
        f"{l2_rng},{l2_crit},"
        f"{l3_rng},{l3_crit},"
        f"{l4_rng},{l4_crit})/1000"
    )


def _master_period_col_letter(master_path: str, sheet_name: str, period_header: str) -> str | None:
    if not master_path or not Path(master_path).is_file():
        return None
    from openpyxl import load_workbook

    wb = load_workbook(master_path, read_only=True, data_only=False)
    try:
        if sheet_name not in wb.sheetnames:
            return None
        ws = wb[sheet_name]
        for c in range(1, ws.max_column + 1):
            if str(ws.cell(1, c).value or "").strip() == period_header:
                return col_letter(c)
    finally:
        wb.close()
    return None


def _apply_sheet_layout(ws, layout: FteLayout, *, last_used_col: int, last_fill_row: int) -> None:
    fill_end_col = last_used_col + 40
    for rr in range(1, last_fill_row + 1):
        ws.row_dimensions[rr].height = ROW_HEIGHT if rr != PROJECT_TITLE_ROW else 36
        for cc in range(layout.helper_start_col, layout.pos_col):
            ws.cell(rr, cc).fill = FILL_GREY
        for cc in range(layout.pos_col, fill_end_col + 1):
            ws.cell(rr, cc).fill = FILL_WHITE
    for cc in range(layout.helper_start_col, layout.pos_col):
        cl = col_letter(cc)
        ws.column_dimensions[cl].outlineLevel = 2
        ws.column_dimensions[cl].hidden = True
    ws.column_dimensions[col_letter(layout.helper_end_col)].collapsed = True
    ws.column_dimensions[col_letter(layout.pos_col)].outlineLevel = 0
    ws.column_dimensions[col_letter(layout.pos_col)].hidden = False


def _paint_header_band(
    ws,
    layout: FteLayout,
    period_cols: list[PeriodColumn],
    *,
    unit_label: str,
    group_col_labels: list[str] | None = None,
) -> None:
    labels = group_col_labels or []
    for i, crit_col in enumerate(layout.group_crit_cols):
        label = labels[i] if i < len(labels) else ""
        ws.cell(HEADER_ROW_7, crit_col).fill = FILL_HEADER
        ws.cell(HEADER_ROW_7, crit_col).font = FONT_HEADER
        hc = ws.cell(HEADER_ROW, crit_col, label)
        hc.font = FONT_HEADER
        hc.alignment = ALIGN_LEFT
        hc.border = BORDER_HEADER_BOTTOM
        hc.fill = FILL_HEADER

    pl_headers = ("Reported", "L2", "L3", "L4")
    pl_cols = (
        layout.pl_reported_col,
        layout.pl_l2_col,
        layout.pl_l3_col,
        layout.pl_l4_col,
    )
    for hdr, col_idx in zip(pl_headers, pl_cols):
        ws.cell(HEADER_ROW_7, col_idx).fill = FILL_HEADER
        ws.cell(HEADER_ROW_7, col_idx).font = FONT_HEADER
        hc = ws.cell(HEADER_ROW, col_idx, hdr)
        hc.font = FONT_HEADER
        hc.alignment = ALIGN_LEFT
        hc.border = BORDER_HEADER_BOTTOM
        hc.fill = FILL_HEADER

    for pc in period_cols:
        ws.cell(HEADER_ROW_7, pc.col_idx).fill = FILL_HEADER
        ws.cell(HEADER_ROW_7, pc.col_idx).font = FONT_HEADER
        ws.cell(HEADER_ROW_7, pc.col_idx).alignment = ALIGN_RIGHT
        hc = ws.cell(HEADER_ROW, pc.col_idx, pc.header)
        hc.font = FONT_HEADER
        hc.alignment = ALIGN_RIGHT
        hc.border = BORDER_HEADER_BOTTOM
        hc.fill = FILL_HEADER
    pos_top = ws.cell(HEADER_ROW_7, layout.pos_col)
    pos_top.fill = FILL_HEADER
    pos_top.font = FONT_HEADER
    pos_top.alignment = ALIGN_LEFT
    pos_hdr = ws.cell(HEADER_ROW, layout.pos_col, unit_label)
    pos_hdr.fill = FILL_HEADER
    pos_hdr.border = BORDER_HEADER_BOTTOM
    pos_hdr.font = FONT_HEADER
    pos_hdr.alignment = ALIGN_LEFT


def _write_subtotal_row(
    ws,
    row: int,
    layout: FteLayout,
    label: str,
    period_cols: list[PeriodColumn],
    sum_rows: list[int],
    *,
    num_fmt: str = NUM_FMT_INT,
) -> None:
    c = ws.cell(row, layout.pos_col, label)
    c.font = FONT_BOLD
    c.alignment = ALIGN_LEFT
    c.border = BORDER_SUBTOTAL_BOTH
    for cc in fte_mapping_column_indices(layout):
        mc = ws.cell(row, cc)
        mc.border = BORDER_SUBTOTAL_BOTH
    for pc in period_cols:
        cell = ws.cell(row, pc.col_idx)
        cell.value = _chunk_sum_formula(sum_rows, pc.col_idx)
        cell.font = FONT_BOLD
        cell.number_format = num_fmt
        cell.alignment = ALIGN_RIGHT
        cell.fill = FILL_SUBTOTAL
        cell.border = BORDER_SUBTOTAL_BOTH


def _apply_hierarchy_outline(ws, row: int, spec: RowSpec, num_group_cols: int) -> None:
    if num_group_cols < 2:
        return
    if spec.row_type == "parent":
        ws.row_dimensions[row].outlineLevel = 0
    elif spec.row_type == "leaf":
        ws.row_dimensions[row].outlineLevel = 1


def _paint_total_personnel_row(
    ws,
    row: int,
    layout: FteLayout,
    *,
    last_used_col: int,
) -> None:
    for cc in fte_mapping_column_indices(layout):
        mc = ws.cell(row, cc)
        mc.border = BORDER_SUBTOTAL_BOTH
    for cc in range(layout.pos_col, last_used_col + 1):
        cell = ws.cell(row, cc)
        cell.fill = FILL_SUBTOTAL
        cell.font = FONT_BOLD
        if cc == layout.pos_col:
            cell.alignment = ALIGN_LEFT
            if not cell.value:
                cell.value = "Personnel expenses"
        elif cc >= layout.first_data_col:
            cell.alignment = ALIGN_RIGHT
        cell.border = BORDER_SUBTOTAL_BOTH


def _style_value_cell(cell, *, bold: bool = False, fill=None, num_fmt: str = NUM_FMT_K) -> None:
    cell.number_format = num_fmt
    cell.alignment = ALIGN_RIGHT
    cell.font = FONT_BOLD if bold else FONT_BASE
    if fill is not None:
        cell.fill = fill


def write_workbook(
    cfg: dict,
    source_works: list[pd.DataFrame],
    source_header_maps: list[dict[str, int]],
    col_maps: list[dict[str, Any]],
) -> str:
    cfg = normalize_config(cfg)
    out_path = build_output_file_path(cfg)
    ensure_output_writable(out_path)

    group_cols = cfg["group_cols"]
    layout = fte_column_layout(len(group_cols))
    period_entries = list(cfg["periods"])
    formula_mode = bool(cfg.get("formula_mode", True))

    all_keys: set[tuple[str, ...]] = set()
    sort_payroll: dict[tuple[str, ...], float] = {}
    for work, cm in zip(source_works, col_maps, strict=False):
        keys, _, _ = aggregate_period_keys(work, cfg)
        all_keys.update(keys)
    if source_works and col_maps:
        last_work = source_works[-1]
        cm = col_maps[-1]
        payroll_cols = cm["payroll_cols"]
        if group_cols:
            g_res = resolve_column_names_to_df(last_work, group_cols)[: len(group_cols)]
            for key, grp in last_work.groupby(g_res, dropna=False):
                if not isinstance(key, tuple):
                    key = (key,)
                key = tuple("" if pd.isna(v) else str(v).strip() for v in key)
                sort_payroll[key] = sum(
                    float(grp[pc].sum()) for pc in payroll_cols if pc in grp.columns
                )
    keys = list(all_keys)
    row_specs = build_row_specs(group_cols, keys, sort_payroll)

    period_cols = build_period_columns(
        period_entries,
        cfg,
        first_data_col=layout.first_data_col,
    )
    last_used_col = period_cols[-1].col_idx if period_cols else layout.first_data_col

    wb = open_session_workbook(cfg)
    ws = replace_workbook_sheet(wb, cfg["sheet_name"][:31])

    project = str(cfg.get("title") or cfg.get("project_name") or "Project").strip()
    company = str(cfg.get("company") or cfg.get("group_name") or "").strip()
    sheet_label = str(cfg.get("sheet_name") or "FTE development").strip()

    ws.cell(PROJECT_TITLE_ROW, layout.pos_col, project).font = FONT_PROJECT
    ws.cell(SUBTITLE_ROW, layout.pos_col, company).font = FONT_SUBTITLE
    period_hdrs = [pc.header for pc in period_cols]
    period_range = ""
    if period_hdrs:
        period_range = f" {period_hdrs[0]} - {period_hdrs[-1]}"
    table_label = "FTE Development"
    table_heading = (
        f"{company} | {table_label}{period_range}" if company else f"{table_label}{period_range}"
    )
    ws.cell(TABLE_TITLE_ROW, layout.pos_col, table_heading).font = FONT_HEADER

    _paint_header_band(
        ws,
        layout,
        period_cols,
        unit_label=cfg["unit_label"],
        group_col_labels=group_cols,
    )

    leaf_rows_by_parent: dict[str, list[int]] = defaultdict(list)
    excel_rows: list[tuple[RowSpec, int]] = []
    r = DATA_START_ROW

    for spec in row_specs:
        excel_rows.append((spec, r))
        display = spec.label
        if spec.level == 2:
            display = f"{INDENT}{display}"
        ws.cell(r, layout.pos_col, display).alignment = ALIGN_LEFT
        ws.cell(r, layout.pos_col).font = FONT_BASE
        if spec.group_key is not None and layout.group_crit_cols:
            for i, val in enumerate(spec.group_key):
                if i < len(group_cols) and i < len(layout.group_crit_cols):
                    cell = ws.cell(r, layout.group_crit_cols[i], val)
                    cell.font = FONT_MAPPING
        _apply_hierarchy_outline(ws, r, spec, len(group_cols))
        r += 1

    block1_detail_rows = [er for spec, er in excel_rows if spec.row_type == "leaf"]
    block1_subtotal_row = r
    _write_subtotal_row(
        ws,
        r,
        layout,
        "# Average FTEs",
        period_cols,
        block1_detail_rows,
        num_fmt=NUM_FMT_INT,
    )
    r += 1

    block2_detail_start = r
    block2_rows: list[int] = []
    for spec, _er in excel_rows:
        block2_rows.append(r)
        display = spec.label
        if spec.level == 2:
            display = f"{INDENT}{display}"
        ws.cell(r, layout.pos_col, display).alignment = ALIGN_LEFT
        ws.cell(r, layout.pos_col).font = FONT_BASE
        _apply_hierarchy_outline(ws, r, spec, len(group_cols))
        r += 1
    block2_subtotal_row = r
    ws.cell(r, layout.pos_col, "Average cost per FTE").font = FONT_BOLD
    ws.cell(r, layout.pos_col).alignment = ALIGN_LEFT
    ws.cell(r, layout.pos_col).border = BORDER_SUBTOTAL_BOTH
    for cc in fte_mapping_column_indices(layout):
        ws.cell(r, cc).border = BORDER_SUBTOTAL_BOTH
    for pc in period_cols:
        cell = ws.cell(r, pc.col_idx)
        cell.font = FONT_BOLD
        cell.number_format = NUM_FMT_K
        cell.alignment = ALIGN_RIGHT
        cell.fill = FILL_SUBTOTAL
        cell.border = BORDER_SUBTOTAL_BOTH
    r += 1

    block3_rows: list[int] = []
    for spec, _er in excel_rows:
        block3_rows.append(r)
        display = spec.label
        if spec.level == 2:
            display = f"{INDENT}{display}"
        ws.cell(r, layout.pos_col, display).alignment = ALIGN_LEFT
        ws.cell(r, layout.pos_col).font = FONT_BASE
        if spec.group_key is not None and layout.group_crit_cols:
            for i, val in enumerate(spec.group_key):
                if i < len(group_cols) and i < len(layout.group_crit_cols):
                    cell = ws.cell(r, layout.group_crit_cols[i], val)
                    cell.font = FONT_MAPPING
        _apply_hierarchy_outline(ws, r, spec, len(group_cols))
        r += 1
    block3_detail_rows = [
        block3_rows[j] for j, (spec, _) in enumerate(excel_rows) if spec.row_type == "leaf"
    ]
    block3_subtotal_row = r
    _write_subtotal_row(
        ws,
        r,
        layout,
        "Payroll accounting",
        period_cols,
        block3_detail_rows,
        num_fmt=NUM_FMT_K,
    )
    r += 1

    master_path = str(cfg.get("master_pl_path") or "").strip()
    master_sheet = cfg.get("master_pl_sheet", "Master_PL")
    pl_line_specs = build_pl_personnel_rows(master_path, master_sheet, cfg)
    pl_prefix = _master_pl_sheet_prefix(cfg, wb)
    pl_col_ranges = _master_pl_column_ranges(master_path, master_sheet, pl_prefix)

    pl_rows: list[int] = []
    for pl_spec in pl_line_specs:
        pl_rows.append(r)
        ws.cell(r, layout.pos_col, pl_spec["display_label"]).alignment = ALIGN_LEFT
        ws.cell(r, layout.pos_col).font = FONT_BASE
        for col_idx, val in (
            (layout.pl_reported_col, _PL_REPORTED_VALUE),
            (layout.pl_l2_col, pl_spec["L2"]),
            (layout.pl_l3_col, pl_spec["L3"]),
            (layout.pl_l4_col, pl_spec["L4"]),
        ):
            cell = ws.cell(r, col_idx, val)
            cell.font = FONT_MAPPING
        r += 1

    total_pe_row = r
    ws.cell(r, layout.pos_col, "Personnel expenses")

    leaf_rows_by_parent.clear()
    if len(group_cols) >= 2:
        for spec, er in excel_rows:
            if spec.row_type == "leaf" and spec.group_key:
                leaf_rows_by_parent[spec.group_key[0]].append(er)

    group_crit_cols = layout.group_crit_cols

    for pi, pc in enumerate(period_cols):
        period_entry = period_entries[pi]
        hdr = resolve_period_header(period_entry, cfg)
        src = fte_source_sheet_name(str(period_entry.get("label") or ""), header=hdr)
        hmap = source_header_maps[pi]
        payroll_headers = col_maps[pi]["payroll_cols"]

        for spec, er in excel_rows:
            is_parent = spec.row_type == "parent"
            fte_cell = ws.cell(er, pc.col_idx)
            if formula_mode:
                if is_parent:
                    child_rows = leaf_rows_by_parent.get(spec.label, [])
                    fte_cell.value = _parent_sum_formula(child_rows, pc.col_idx)
                elif spec.row_type == "leaf":
                    fte_cell.value = _fte_sumifs_formula(
                        src,
                        hmap,
                        er,
                        group_cols,
                        cfg,
                        group_crit_cols=group_crit_cols,
                    )
            _style_value_cell(fte_cell, bold=is_parent, fill=FILL_PERIOD if not is_parent else FILL_SUBTOTAL, num_fmt=NUM_FMT_INT)

        for idx, (spec, er) in enumerate(excel_rows):
            cost_r = block2_detail_start + idx
            pay_r = block3_rows[idx]
            fte_ref = f"{col_letter(pc.col_idx)}{er}"
            pay_ref = f"{col_letter(pc.col_idx)}{pay_r}"
            cost_cell = ws.cell(cost_r, pc.col_idx)
            cost_cell.value = f'=IF({fte_ref}=0,"",{pay_ref}/{fte_ref})'
            _style_value_cell(cost_cell, num_fmt=NUM_FMT_K)

        fte_tot = f"{col_letter(pc.col_idx)}{block1_subtotal_row}"
        pay_tot = f"{col_letter(pc.col_idx)}{block3_subtotal_row}"
        sub_cell = ws.cell(block2_subtotal_row, pc.col_idx)
        sub_cell.value = f'=IF({fte_tot}=0,"",{pay_tot}/{fte_tot})'
        _style_value_cell(sub_cell, bold=True, fill=FILL_SUBTOTAL, num_fmt=NUM_FMT_K)
        sub_cell.border = BORDER_SUBTOTAL_BOTH

        for idx, (spec, _er) in enumerate(excel_rows):
            pay_r = block3_rows[idx]
            pay_cell = ws.cell(pay_r, pc.col_idx)
            is_parent = spec.row_type == "parent"
            if formula_mode:
                if is_parent:
                    children = [
                        block3_rows[j]
                        for j, (s, _) in enumerate(excel_rows)
                        if s.row_type == "leaf" and s.group_key and s.group_key[0] == spec.label
                    ]
                    pay_cell.value = _parent_sum_formula(children, pc.col_idx)
                elif spec.row_type == "leaf":
                    pay_cell.value = _payroll_sumifs_formula(
                        src,
                        hmap,
                        payroll_headers,
                        pay_r,
                        group_cols,
                        cfg,
                        group_crit_cols=group_crit_cols,
                    )
            _style_value_cell(pay_cell, bold=is_parent, fill=FILL_SUBTOTAL if is_parent else FILL_WHITE, num_fmt=NUM_FMT_K)

        for i, _pl_spec in enumerate(pl_line_specs):
            pl_r = pl_rows[i]
            pl_cell = ws.cell(pl_r, pc.col_idx)
            if formula_mode:
                pl_cell.value = _pl_sumifs_formula(
                    master_path,
                    master_sheet,
                    pc.header,
                    pl_r,
                    layout,
                    sheet_prefix=pl_prefix,
                    col_ranges=pl_col_ranges,
                )
            else:
                pl_cell.value = 0
            _style_value_cell(pl_cell, num_fmt=NUM_FMT_K)

        pe_refs = [f"{col_letter(pc.col_idx)}{block3_subtotal_row}"]
        pe_refs.extend(f"{col_letter(pc.col_idx)}{pr}" for pr in pl_rows)
        total_cell = ws.cell(total_pe_row, pc.col_idx)
        total_cell.value = f"=SUM({','.join(pe_refs)})"
        _style_value_cell(total_cell, bold=True, fill=FILL_SUBTOTAL, num_fmt=NUM_FMT_K)
        total_cell.border = BORDER_SUBTOTAL_BOTH


    for name in list(wb.sheetnames):
        if name.startswith("__SOURCE__"):
            del wb[name]
    for i, df in enumerate(source_works):
        period_entry = period_entries[i]
        hdr = resolve_period_header(period_entry, cfg)
        sheet_title = fte_source_sheet_name(str(period_entry.get("label") or ""), header=hdr)
        if sheet_title in wb.sheetnames:
            del wb[sheet_title]
        src_ws = wb.create_sheet(sheet_title)
        stub_cfg = {
            "file_path": cfg["periods"][i]["file_path"],
            "sheet_name": _resolve_sheet_name(
                cfg["periods"][i]["file_path"], cfg["periods"][i].get("sheet_name", "")
            ),
        }
        write_source_df_to_ws(src_ws, df, stub_cfg)
        if formula_mode:
            _apply_fte_wt_formulas(
                src_ws,
                col_maps[i]["employment"],
                col_maps[i]["months_sum"],
            )
        source_header_maps[i] = _get_sheet_header_map(src_ws)

    ws.column_dimensions[col_letter(layout.pos_col)].width = LABEL_COL_WIDTH
    for cc in range(layout.helper_start_col, layout.pos_col):
        ws.column_dimensions[col_letter(cc)].width = MAP_COL_WIDTH
    for pc in period_cols:
        ws.column_dimensions[col_letter(pc.col_idx)].width = VALUE_COL_WIDTH

    last_row = total_pe_row
    _apply_sheet_layout(ws, layout, last_used_col=last_used_col, last_fill_row=last_row + 50)
    _paint_header_band(
        ws,
        layout,
        period_cols,
        unit_label=cfg["unit_label"],
        group_col_labels=group_cols,
    )
    _paint_total_personnel_row(ws, total_pe_row, layout, last_used_col=last_used_col)
    if len(group_cols) >= 2:
        ws.sheet_view.showOutlineSymbols = True
        ws.sheet_properties.outlinePr.summaryBelow = True
        ws.sheet_properties.outlinePr.applyStyles = True

    from databook_workbook import reorder_workbook_sheets

    reorder_workbook_sheets(wb)
    wb.save(out_path)
    return out_path


def run_fte_payroll(cfg: dict) -> str:
    cfg = normalize_config(cfg)
    header_row = int(cfg["header_row"])
    period_specs = [
        PeriodSpec(p["label"], p["file_path"], p.get("sheet_name", "")) for p in cfg["periods"]
    ]

    source_works: list[pd.DataFrame] = []
    source_header_maps: list[dict[str, int]] = []
    col_maps: list[dict[str, Any]] = []

    for spec in period_specs:
        df = load_period_dataframe(spec, header_row)
        _keys, work, col_map = aggregate_period_keys(df, cfg)
        source_works.append(work)
        col_maps.append(col_map)
        tmp_wb = Workbook()
        tmp_ws = tmp_wb.active
        stub = {
            "file_path": spec.file_path,
            "sheet_name": _resolve_sheet_name(spec.file_path, spec.sheet_name),
        }
        write_source_df_to_ws(tmp_ws, work, stub)
        if bool(cfg.get("formula_mode", True)):
            _apply_fte_wt_formulas(tmp_ws, col_map["employment"], col_map["months_sum"])
        source_header_maps.append(_get_sheet_header_map(tmp_ws))

    return write_workbook(cfg, source_works, source_header_maps, col_maps)


def main() -> None:
    if len(sys.argv) < 2:
        raise SystemExit("Usage: python FTE_payroll.py <config.json>")
    cfg_path = Path(sys.argv[1])
    cfg = json.loads(cfg_path.read_text(encoding="utf-8-sig"))
    out = run_fte_payroll(cfg)
    print(json.dumps({"success": True, "output_file_path": out}))


if __name__ == "__main__":
    main()
