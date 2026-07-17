import os
import re
import math
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from typing import Dict, Any, List, Tuple, Optional
from PIL import Image

import numpy as np
import pandas as pd
import matplotlib.pyplot as plt
from matplotlib.colors import LinearSegmentedColormap
import matplotlib.font_manager as fm
import textwrap
import matplotlib.patheffects as pe

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image as XLImage


def _xl_image_from_path(path: str) -> XLImage:
    """Embed PNG bytes so temp files can be deleted before workbook.save()."""
    return XLImage(BytesIO(Path(path).read_bytes()))

try:
    from revenue_reconciliation import sales_basis_labels
except ImportError:  # pragma: no cover - standalone fallback
    def sales_basis_labels(sales_basis=None):
        basis = "gross" if str(sales_basis or "").strip().lower() == "gross" else "net"
        metric = "Gross sales" if basis == "gross" else "Net sales"
        return type(
            "SalesBasisLabels",
            (),
            {
                "sales_basis": basis,
                "metric_label": metric,
                "total_label": f"Total {metric}",
                "reported_label": f"Reported {metric}",
            },
        )()


from gst_excel_theme import THEME, apply_matplotlib_theme, matplotlib_font_properties

apply_matplotlib_theme(size=6)
FONT_PROP = matplotlib_font_properties()
# Matches Excel header above the chart (theme font 12 / Finssentials navy).
HEADER_FONT_PROP = matplotlib_font_properties(size=12)

DROP_INDEX_PRINT_LIMIT = 200

BRAND_BLUE = "#1E3A5F"
BRAND_BLUE_LIGHT = "#C5D8EB"
_BLUE_CMAP = LinearSegmentedColormap.from_list(
    "finssentials_blue",
    [BRAND_BLUE_LIGHT, BRAND_BLUE],
)

# Universal Fast Track bar defaults (shared with batch builder).
DEFAULT_BUCKET_MODE = {
    "enabled": True,
    "top1_min_share": 0.05,
    "thresholds": [0.2, 0.4, 0.6, 0.8],
    "labels": {"top1": "Top 1", "bucket": "Top {a}-{b}", "other": "Rest"},
}
DEFAULT_TOP_N_MODE = {
    "enabled": False,
    "top1_min_share": 0.05,
    "thresholds": [0.2, 0.4, 0.6, 0.8],
    "labels": {"top1": "Top 1", "bucket": "Top {a}-{b}", "other": "Rest"},
}


def default_bars_config(
    *,
    entity_col: str = "",
    segment_col: str = "",
    customer_col: str = "",
    product_col: str = "",
    region_col: str = "",
) -> List[Dict[str, Any]]:
    """Build up to 5 bars; omit entries with empty dim columns (Segment/Region optional)."""
    # top_n is only a safety cap for non-bucket mode (share floor drives counts).
    specs = [
        ("Entity", entity_col, 25, False),
        ("Segment", segment_col, 25, False),
        ("Top customers", customer_col, 4, True),
        ("Top products", product_col, 4, True),
        ("Region", region_col, 25, False),
    ]
    bars: List[Dict[str, Any]] = []
    for title, dim_col, top_n, bucket_on in specs:
        dim = str(dim_col or "").strip()
        if not dim:
            continue
        bm = dict(DEFAULT_BUCKET_MODE if bucket_on else DEFAULT_TOP_N_MODE)
        bm["enabled"] = bool(bucket_on)
        bars.append({"title": title, "dim_col": dim, "top_n": top_n, "bucket_mode": bm})
    return bars


# Standalone demo CONFIG (not used by Fast Track imports).
CONFIG: Dict[str, Any] = {
    "title": "Project Draft",
    "table": "Net sales breakdown",
    "company": "Draft AG",
    "subtitle_suffix": "Revenue share by dimension",
    "period_label_prefix": "",
    "sales_basis": "net",
    "current_year": 2025,
    "current_month": 7,
    "fiscal_year_end_month": 7,
    "fiscal_year_end_day": 31,
    "period_mode": "FY",
    "calc_mode": "accrual",
    "value_col": "Contract Value After Discount",
    "invoice_col": "Invoice Date",
    "start_col": "Contract Start Date",
    "end_col": "Contract End Date",
    "filter": {"enabled": False, "col": "", "values": set()},
    "apply_fx": False,
    "fx_col": "Functional FX Rate",
    "bars": default_bars_config(
        entity_col="Entity",
        segment_col="Segment",
        customer_col="Customer Name",
        product_col="Product Name",
        region_col="End Customer Region",
    ),
    "plot": {
        "show_plot_title": False,
        "figsize": (5.0, 5.48),
        "dpi": 300,
        "show_values": False,
        "label_min_share": 0.05,
        "segment_fontsize": 7,
        "min_share_to_rest": 0.05,
        "x_step": 0.32,
        "xtick_rotation": 22,
    },
    "rest_label": "Misc.",
    "missing_token": "__MISSING__",
    "output_dir": ".",
    "output_prefix": "vertical_bars",
    "excel": {
        "enabled": True,
        "col_offset": 3,
        "sheet_name": "Net sales breakdown",
        "image_anchor_cell": "D6",
        "set_column_widths": True,
        "legend_anchor_cell": "D35",
        "legend_scale": 0.35,
    },
}


# Helpers / Validation
def normalize_config(cfg: Dict[str, Any]) -> Dict[str, Any]:
    out = dict(cfg)

    required = ["current_year", "current_month", "calc_mode", "value_col", "invoice_col", "period_mode", "bars"]
    for k in required:
        if k not in out or out[k] in (None, ""):
            raise ValueError(f"CONFIG['{k}'] muss gesetzt sein.")

    if out["calc_mode"] not in {"invoice", "accrual"}:
        raise ValueError("CONFIG['calc_mode'] muss 'invoice' oder 'accrual' sein.")
    if out["period_mode"] not in {"FY", "YTD", "LTM"}:
        raise ValueError("CONFIG['period_mode'] muss 'FY', 'YTD' oder 'LTM' sein.")

    m = int(out["current_month"])
    if m < 1 or m > 12:
        raise ValueError("CONFIG['current_month'] muss 1..12 sein.")

    fy_end_m = int(out.get("fiscal_year_end_month", out.get("fy_end_month", 12)))
    fy_end_d = int(out.get("fiscal_year_end_day", out.get("fy_end_day", 31)))
    if not (1 <= fy_end_m <= 12):
        raise ValueError("CONFIG['fiscal_year_end_month'] muss 1..12 sein.")
    if not (1 <= fy_end_d <= 31):
        raise ValueError("CONFIG['fiscal_year_end_day'] muss 1..31 sein.")
    out["fiscal_year_end_month"] = fy_end_m
    out["fiscal_year_end_day"] = fy_end_d

    labels = sales_basis_labels(out.get("sales_basis"))
    out["sales_basis"] = labels.sales_basis
    metric = labels.metric_label
    # Prefer explicit table; otherwise derive from sales basis.
    if not str(out.get("table") or "").strip() or str(out.get("table")).strip() in {
        "Gross sales breakdown",
        "Net sales breakdown",
    }:
        out["table"] = f"{metric} breakdown"
    if not str(out.get("subtitle_suffix") or "").strip():
        out["subtitle_suffix"] = "Revenue share by dimension"

    # Sheet name follows sales basis: "Net sales breakdown" / "Gross sales breakdown"
    sheet_title = f"{metric} breakdown"
    excel_cfg = dict(out.get("excel") or {})
    prev_sheet = str(excel_cfg.get("sheet_name") or out.get("base_sheet_name") or "").strip().lower()
    if prev_sheet in {
        "",
        "vertical bars",
        "horizontal bars",
        "breakdown",
        "net sales breakdown",
        "gross sales breakdown",
    }:
        excel_cfg["sheet_name"] = sheet_title
    out["excel"] = excel_cfg
    if not str(out.get("base_sheet_name") or "").strip() or str(out.get("base_sheet_name")).strip().lower() in {
        "vertical bars",
        "horizontal bars",
        "breakdown",
        "net sales breakdown",
        "gross sales breakdown",
    }:
        out["base_sheet_name"] = sheet_title

    # Filter (accept Fast Track "filters" passthrough as disabled default)
    f = out.get("filter") or {}
    if not f and out.get("filters"):
        f = {"enabled": False}
    if f.get("enabled", False):
        if not f.get("col"):
            raise ValueError("filter.enabled=True, aber filter.col fehlt.")
        vals = f.get("values")
        if vals is None or len(vals) == 0:
            raise ValueError("filter.enabled=True, aber filter.values ist leer.")
        f["values"] = set(vals)
    out["filter"] = f

    # FX
    if out.get("apply_fx", False):
        if not out.get("fx_col"):
            raise ValueError("apply_fx=True, aber fx_col fehlt.")

    # Bars — drop empty dim_col entries (optional Segment / Region)
    bars_in = out.get("bars") or []
    if not isinstance(bars_in, list):
        raise ValueError("CONFIG['bars'] muss eine Liste sein.")
    bars = [b for b in bars_in if str((b or {}).get("dim_col") or "").strip()]
    if len(bars) == 0:
        raise ValueError("CONFIG['bars'] muss mindestens eine Dimension mit dim_col enthalten.")
    if len(bars) > 5:
        raise ValueError("Maximal 5 Balken erlaubt (bars).")

    for i, b in enumerate(bars, start=1):
        if b.get("top_n") is None:
            raise ValueError(f"bars[{i}].top_n fehlt (pro Balken).")
        b["top_n"] = int(b["top_n"])
        if b["top_n"] < 1:
            raise ValueError(f"bars[{i}].top_n muss >= 1 sein.")
        bm = b.get("bucket_mode") or {"enabled": False}
        bm.setdefault("enabled", False)
        if bm["enabled"]:
            bm.setdefault("top1_min_share", 0.10)
            bm.setdefault("thresholds", [0.2, 0.4, 0.6, 0.8])
            bm.setdefault("labels", {"top1": "Top 1", "bucket": "Top {a}-{b}", "other": out.get("rest_label", "Rest")})
            # thresholds müssen streng monoton sein und zwischen 0 und 1
            th = [float(x) for x in bm["thresholds"]]
            if any((t <= 0 or t >= 1) for t in th):
                raise ValueError(f"bars[{i}].bucket_mode.thresholds müssen in (0,1) liegen.")
            if any(th[j] >= th[j + 1] for j in range(len(th) - 1)):
                raise ValueError(f"bars[{i}].bucket_mode.thresholds müssen strikt steigend sein.")
            bm["thresholds"] = th
        bm["top1_min_share"] = float(bm["top1_min_share"])
        b["bucket_mode"] = bm

    out["bars"] = bars
    return out


# Period logic
def safe_month_day_ts(year: int, month: int, day: int) -> pd.Timestamp:  # Helperfunktion, um gescheite Daten zu erhalten, korrigiert auch ungültige Tage
    first = pd.Timestamp(year, month, 1)
    last = first + pd.offsets.MonthEnd(0)
    d = min(int(day), int(last.day))
    return pd.Timestamp(year, month, d)


def fiscal_year_bounds(fy_end_year: int, fy_end_m: int, fy_end_d: int) -> Tuple[pd.Timestamp, pd.Timestamp]:  # Gibt Start und Ende eines FY zurück
    end = safe_month_day_ts(fy_end_year, fy_end_m, fy_end_d)
    start = (end - pd.DateOffset(years=1)) + pd.Timedelta(days=1)
    return start, end


def compute_current_fy_end_year(as_of_end: pd.Timestamp, fy_end_m: int, fy_end_d: int) -> int:  # Findet heraus in welchem Geschäftsjahr wir uns befinden
    CY = as_of_end.year
    fye_this_year = safe_month_day_ts(CY, fy_end_m, fy_end_d)
    return CY if as_of_end <= fye_this_year else CY + 1


def get_period(cfg: Dict[str, Any]) -> Tuple[pd.Timestamp, pd.Timestamp, str]:
    CY = int(cfg["current_year"])
    m = int(cfg["current_month"])

    fy_end_m = int(cfg["fiscal_year_end_month"])
    fy_end_d = int(cfg["fiscal_year_end_day"])

    as_of_end = pd.Timestamp(CY, m, 1) + pd.offsets.MonthEnd(0)  # Betrachtungszeitpunkt
    cur_fy_end_year = compute_current_fy_end_year(as_of_end, fy_end_m, fy_end_d)  # Aktuelles Geschäftsjahr

    def fy_label(y: int) -> str:
        return f"FY{str(y)[-2:]}A"

    mode = cfg["period_mode"]

    if mode == "FY":
        # letztes abgeschlossenes FY (FY-Ende <= as_of_end)
        fye_this_year = safe_month_day_ts(as_of_end.year, fy_end_m, fy_end_d)  # Stichtag im Jahr der as_of_end
        last_completed_fy_end_year = as_of_end.year if as_of_end >= fye_this_year else as_of_end.year - 1  # last_complete = aktuelles Jahr wenn FY schon vorbei
        s, e = fiscal_year_bounds(last_completed_fy_end_year, fy_end_m, fy_end_d)  # Start und ende des FY
        return s, e, fy_label(last_completed_fy_end_year)

    if mode == "YTD":
        # YTD: vom FY-Start bis as_of_end
        s, _ = fiscal_year_bounds(cur_fy_end_year, fy_end_m, fy_end_d)  # Start des aktuellen Geschäftsjahres
        return s, as_of_end, f"YTD{str(cur_fy_end_year)[-2:]}"  # Ende dann Betrachtungszeitpunkt

    if mode == "LTM":
        # letzte 12 Monate bis as_of_end
        start_ltm = as_of_end - pd.DateOffset(years=1) + pd.Timedelta(days=1)  # Start ist Betrachtungszeitpunkt - 1 Jahr + 1 Tag
        return start_ltm, as_of_end, f"LTM {m}-{str(as_of_end.year)[-2:]}"

    raise ValueError("Unbekannter period_mode.")


def get_comparison_periods(cfg: Dict[str, Any]) -> List[Tuple[pd.Timestamp, pd.Timestamp, str]]:
    """
    Three comparison windows for side-by-side charts, left→right oldest→newest:
    - If as-of month-end is the fiscal year-end (Stichtag): last 3 completed FYs
    - Else: previous FY, last completed FY, and YTD of the current FY
    """
    CY = int(cfg["current_year"])
    m = int(cfg["current_month"])
    fy_end_m = int(cfg["fiscal_year_end_month"])
    fy_end_d = int(cfg["fiscal_year_end_day"])

    as_of_end = pd.Timestamp(CY, m, 1) + pd.offsets.MonthEnd(0)
    fye_this_calendar_year = safe_month_day_ts(as_of_end.year, fy_end_m, fy_end_d)
    # True when the as-of month-end is exactly the FY closing month (no partial YTD).
    at_fiscal_year_end = int(as_of_end.month) == int(fy_end_m)

    def fy_label(y: int) -> str:
        return f"FY{str(y)[-2:]}A"

    if at_fiscal_year_end:
        last_completed = as_of_end.year if as_of_end >= fye_this_calendar_year else as_of_end.year - 1
        years = [last_completed - 2, last_completed - 1, last_completed]
        out: List[Tuple[pd.Timestamp, pd.Timestamp, str]] = []
        for y in years:
            s, e = fiscal_year_bounds(y, fy_end_m, fy_end_d)
            out.append((s, e, fy_label(y)))
        return out

    cur_fy_end_year = compute_current_fy_end_year(as_of_end, fy_end_m, fy_end_d)
    last_completed = as_of_end.year if as_of_end >= fye_this_calendar_year else as_of_end.year - 1
    # When not at FYE, last_completed is the closed FY just before the open YTD year.
    prev_fy = last_completed
    older_fy = last_completed - 1
    ytd_start, _ = fiscal_year_bounds(cur_fy_end_year, fy_end_m, fy_end_d)
    ytd_label = f"YTD{str(cur_fy_end_year)[-2:]}"

    s0, e0 = fiscal_year_bounds(older_fy, fy_end_m, fy_end_d)
    s1, e1 = fiscal_year_bounds(prev_fy, fy_end_m, fy_end_d)
    return [
        (s0, e0, fy_label(older_fy)),
        (s1, e1, fy_label(prev_fy)),
        (ytd_start, as_of_end, ytd_label),
    ]


# Preprocess + Revenue calc
def preprocess_input(df: pd.DataFrame, cfg: Dict[str, Any]) -> pd.DataFrame:
    d = df.copy()
    d.columns = d.columns.str.strip()

    # Fast Track multi-rule filters (preferred) + legacy singular filter.
    try:
        from funktionssammlung import apply_filters

        d = apply_filters(d, cfg)
    except Exception:
        pass
    f = cfg.get("filter", {}) or {}
    if f.get("enabled", False):
        col = f["col"]
        if col not in d.columns:
            raise ValueError(f"Filter-Spalte '{col}' fehlt im Input.")
        d = d[d[col].isin(f["values"])].copy()

    mapping = str(cfg.get("invoice_mapping_mode") or "date").strip().lower()
    # Date parse (period/year mode keeps invoice_col as FY labels — do not force datetime)
    date_cols: List[str] = []
    if mapping != "year":
        date_cols.append(cfg["invoice_col"])
    if cfg["calc_mode"] == "accrual":
        for c in [cfg["start_col"], cfg["end_col"]]:
            if c not in d.columns:
                raise ValueError(f"calc_mode='accrual' braucht Spalte '{c}' im Input.")
        date_cols += [cfg["start_col"], cfg["end_col"]]

    for c in date_cols:
        if c not in d.columns:
            raise ValueError(f"Spalte '{c}' fehlt im Input.")
        d[c] = pd.to_datetime(d[c], errors="coerce", dayfirst=True)

    # Numeric value
    if cfg["value_col"] not in d.columns:
        raise ValueError(f"value_col '{cfg['value_col']}' fehlt im Input.")
    d["value"] = pd.to_numeric(d[cfg["value_col"]], errors="coerce")

    # FX
    if cfg.get("apply_fx", False):
        fx_col = cfg["fx_col"]
        if fx_col not in d.columns:
            raise ValueError(f"apply_fx=True, aber fx_col '{fx_col}' fehlt im Input.")
        d["fx_rate"] = pd.to_numeric(d[fx_col], errors="coerce")
        valid_fx = d["fx_rate"].notna() & (d["fx_rate"] != 0)
        d.loc[valid_fx, "value"] = d.loc[valid_fx, "fx_rate"] * d.loc[valid_fx, "value"]

    # Drop rows where invoice/value missing (für Revenue brauchen wir value und invoice zumindest als Referenz)
    d = d.dropna(subset=[cfg["invoice_col"], "value"]).copy()

    # For accrual: require contract dates and valid range
    if cfg["calc_mode"] == "accrual":
        d = d.dropna(subset=[cfg["start_col"], cfg["end_col"]]).copy()
        d = d[d[cfg["end_col"]] >= d[cfg["start_col"]]].copy()

    return d


def days_inclusive(a, b) -> pd.Series:  # Datumsubtraktion --> Laufzeit Vertrag
    a = pd.to_datetime(a)
    b = pd.to_datetime(b)
    days = (b - a) / np.timedelta64(1, "D")
    days = np.floor(days) + 1
    if isinstance(days, pd.Series):
        out = days
    else:
        idx = getattr(a, "index", None) or getattr(b, "index", None)
        out = pd.Series(days, index=idx)
    return out.clip(lower=0)


def recognized_value(d: pd.DataFrame, period_start: pd.Timestamp, period_end: pd.Timestamp, cfg: Dict[str, Any]) -> pd.Series:  # accrual Logik
    S = d[cfg["start_col"]]
    E = d[cfg["end_col"]]
    I = d[cfg["invoice_col"]]

    aktiv_von = np.where(  ## aktiv von Vertragsstartdatum, falls Invoice nach Periodenstart oder Vertrag nach Periodenstart
        I >= period_start,
        S,
        np.where(S > period_start, S, period_start)
    )
    aktiv_von = pd.to_datetime(aktiv_von)

    aktiv_bis = pd.to_datetime(
        np.minimum(E.values.astype("datetime64[ns]"), np.datetime64(period_end))
    )

    tage_realisiert = days_inclusive(aktiv_von, aktiv_bis)
    contract_days = days_inclusive(S, E).clip(lower=1)

    value = pd.to_numeric(d["value"], errors="coerce").fillna(0.0)

    ko1 = (S > period_end) | (I > period_end)  ## KO wenn Vertragsstartdatum nach Periodenende oder Invoice nach Periodenende
    ko2 = (I < period_start) & (E < period_start)  ## KO wenn Invoice vor Periodenstart und Vertragsenddatum vor Periodenstart

    out = value * (tage_realisiert / contract_days)  ## Output=vertragswert* realisierte Tage/Vertragslaufzeit
    out = out.where(~(ko1 | ko2), 0.0)
    return out


def compute_revenue(d: pd.DataFrame, start: pd.Timestamp, end: pd.Timestamp, cfg: Dict[str, Any]) -> pd.Series:  # Wendet Invoice oder Accrual an
    try:
        from funktionssammlung import compute_amount

        # compute_amount expects the numeric column name already present.
        tmp = d.copy()
        tmp[cfg["value_col"]] = d["value"]
        return compute_amount(tmp, start, end, cfg, cfg["value_col"])
    except Exception:
        pass
    if cfg["calc_mode"] == "invoice":
        mapping = str(cfg.get("invoice_mapping_mode") or "date").strip().lower()
        if mapping == "year":
            years = pd.to_numeric(d[cfg["invoice_col"]], errors="coerce")
            # also accept FY24A-style labels
            if years.isna().any():
                from funktionssammlung import parse_invoice_fy_year

                years = parse_invoice_fy_year(d[cfg["invoice_col"]])
            mask = years == int(end.year)
            amt = pd.to_numeric(d["value"], errors="coerce").fillna(0.0)
            return amt.where(mask, 0.0)
        mask = (d[cfg["invoice_col"]] >= start) & (d[cfg["invoice_col"]] <= end)
        amt = pd.to_numeric(d["value"], errors="coerce").fillna(0.0)
        return amt.where(mask, 0.0)

    # accrual
    return recognized_value(d, start, end, cfg)


# Breakdown logic per bar
@dataclass
class Segment:  # Legt Eigenschaften der segmente fest, Segment ist ein Abschnitt eines Balkens
    label: str
    share: float  # Anteil am TOTAL (inkl. negative Anteile möglich)
    amount: float
    is_rest: bool = False
    draw: bool = True  # False => nicht im Balken zeichnen (z.B. negative Kategorie)


def _clean_dim_series(s: pd.Series, missing_token: str) -> pd.Series:  # Alle fehlenden Werte zu einer Missing Kategorie machen
    x = s.astype(str).str.strip()
    # normalize common missing strings
    x = x.replace({"": missing_token, "nan": missing_token, "None": missing_token, "NaT": missing_token})
    x = x.where(~s.isna(), missing_token)
    return x


def aggregate_by_dimension(d: pd.DataFrame, dim_col: str, missing_token: str) -> pd.DataFrame:  # Aggregiert die Umsatzspalte nach den Dimensionen
    if dim_col not in d.columns:
        raise ValueError(f"Dimension-Spalte '{dim_col}' fehlt im Input.")
    tmp = d[[dim_col, "amount"]].copy()
    tmp["dim"] = _clean_dim_series(tmp[dim_col], missing_token)
    out = tmp.groupby("dim", as_index=False)["amount"].sum()
    out = out.rename(columns={"dim": "label"})
    return out


def build_top_n_segments(
    df_dim: pd.DataFrame,
    top_n: int,
    rest_label: str,
    missing_token: str,
    min_share: float = 0.05,
) -> Tuple[List[Segment], List[Segment]]:
    """
    Non-bucket mode: keep categories (by size desc) while share >= min_share;
    everything smaller (plus missing) goes into Rest. ``top_n`` is only a safety cap.
    """
    df = df_dim.copy()

    missing_amt = float(df.loc[df["label"] == missing_token, "amount"].sum())
    df = df[df["label"] != missing_token].copy()
    df = df.sort_values("amount", ascending=False).reset_index(drop=True)

    total = float(df["amount"].sum() + missing_amt)
    if abs(total) < 1e-12:
        return [Segment(label=rest_label, share=1.0, amount=total, is_rest=True, draw=True)], []

    floor = float(min_share or 0.0)
    max_keep = max(1, int(top_n)) if top_n else 25
    segs: List[Segment] = []
    neg_list: List[Segment] = []
    rest_amt = float(missing_amt)
    kept = 0

    for _, r in df.iterrows():
        amt = float(r["amount"])
        sh = amt / total
        if amt < 0:
            rest_amt += amt
            neg_list.append(
                Segment(label=str(r["label"]), share=float(sh), amount=amt, is_rest=False, draw=False)
            )
            continue
        if kept < max_keep and sh >= floor - 1e-15:
            segs.append(
                Segment(label=str(r["label"]), share=float(sh), amount=amt, is_rest=False, draw=True)
            )
            kept += 1
        else:
            rest_amt += amt

    rest_share = rest_amt / total
    segs.append(
        Segment(
            label=rest_label,
            share=float(rest_share),
            amount=rest_amt,
            is_rest=True,
            draw=(rest_amt > 0),
        )
    )
    return segs, neg_list


def build_bucket_segments(
    df_dim: pd.DataFrame,
    bucket_cfg: Dict[str, Any],
    rest_label: str,
    missing_token: str
) -> Tuple[List[Segment], List[Segment]]:
    """
    Bucket Mode:
    - Buckets werden anhand positiver Rangfolge gebildet
    - Negative Einzelkategorien werden NICHT separat gelistet, sondern gehen netto in Rest ("Other")
    - Share immer relativ zu TOTAL netto
    Returns: (segments, negative_list)  # negative_list hier immer leer
    """
    df = df_dim.copy()

    missing_amt = float(df.loc[df["label"] == missing_token, "amount"].sum())
    df = df[df["label"] != missing_token].copy()

    # Total netto
    total = float(df["amount"].sum() + missing_amt)
    if abs(total) < 1e-12:
        return ([Segment(label=rest_label, share=1.0, amount=float(df["amount"].sum() + missing_amt), is_rest=True, draw=True)], [])

    # Für Bucket-Bildung sortieren wir nach amount DESC
    df = df.sort_values("amount", ascending=False).reset_index(drop=True)

    # Top1-Min-Share Logik: hier auf POSITIVE Anteil basieren (sonst weird bei negativen totals)
    df_pos = df[df["amount"] > 0].copy().reset_index(drop=True)
    if len(df_pos) == 0:
        # nichts Positives -> nichts zeichnen, aber mathematisch total existiert
        # => pragmatisch: alles als Rest (draw=False wenn Rest negativ)
        rest_amt = float(df["amount"].sum() + missing_amt)
        return ([Segment(label=rest_label, share=(rest_amt / total), amount=rest_amt, is_rest=True, draw=(rest_amt > 0))], [])

    total_pos_for_cut = float(df_pos["amount"].sum())
    top1_share_pos = float(df_pos.loc[0, "amount"] / total_pos_for_cut)

    thresholds = list(bucket_cfg.get("thresholds", [0.2, 0.4, 0.6, 0.8]))
    top1_min = float(bucket_cfg.get("top1_min_share", 0.10))
    labels = bucket_cfg.get("labels", {"top1": "Top 1", "bucket": "Top {a}-{b}", "other": rest_label})

    segments: List[Tuple[str, float]] = []
    start_idx = 0

    if top1_share_pos > top1_min:
        segments.append((str(df_pos.loc[0, "label"]), float(df_pos.loc[0, "amount"])))  # Top1 = echter Name (des Kunden bswp.)
        start_idx = 1

    amounts = df_pos["amount"].to_numpy()
    cum = np.cumsum(amounts) / total_pos_for_cut

    def cut_exclusive(th: float, cur: int) -> int:
        """
        'Bis zu Threshold' Logik:
        - Bucket enthält nur Elemente mit cum <= th
        - Das Element, das die Grenze überschreitet, startet den nächsten Bucket
        """
        # Indizes, die wir überhaupt noch betrachten (ab cur)
        mask = np.arange(len(cum)) >= cur
        valid = np.where(mask & (cum <= th))[0]

        if len(valid) == 0:
            # Noch kein Element passt in diesen Bucket (weil schon das erste > th)
            # -> damit es nicht leer bleibt, mind. 1 Element
            return min(cur + 1, len(df_pos))

        # letztes Element, das noch <= th ist; +1 weil slice-Ende exklusiv
        end = int(valid[-1]) + 1

        # Sicherheitsnetz: nie rückwärts / nie leer
        if end <= cur:
            end = min(cur + 1, len(df_pos))

        return end

    ranges: List[Tuple[int, int]] = []
    cur = start_idx
    for th in thresholds:
        nxt = cut_exclusive(th, cur)
        if nxt > cur:
            ranges.append((cur, nxt))
            cur = nxt

    if cur < len(df_pos):  # Falls noch Kunden übrig --> übrigeb werdeb in "Rest"-Bucket gepackt
        ranges.append((cur, len(df_pos)))

    # erst die echten Buckets (ohne Rest)
    for j, (a, b) in enumerate(ranges):
        amt = float(df_pos.iloc[a:b]["amount"].sum())
        if amt <= 0:
            continue
        is_last = (j == len(ranges) - 1)
        if is_last:
            # Rest-Bucket (erstmal nur positive Reste)
            segments.append((labels.get("other", rest_label), amt))
        else:
            rank_a = a + 1
            rank_b = b
            if b - a == 1:
                label = df_pos.iloc[a]["label"]
            else:
                label = labels.get("bucket", "Top {a}-{b}").format(a=rank_a, b=rank_b)
            segments.append((label, amt))

    # Jetzt: negative + missing netto in Rest-Bucket addieren
    rest_extra = float((df[df["amount"] <= 0]["amount"].sum()) + missing_amt)  # alle non-positive + missing
    # Finde Rest bucket
    rest_name = labels.get("other", rest_label)
    found = False
    for idx, (lab, amt) in enumerate(segments):
        if lab == rest_name:
            segments[idx] = (lab, amt + rest_extra)
            found = True
            break
    if not found:
        segments.append((rest_name, rest_extra))

    out: List[Segment] = []
    for lab, amt in segments:
        sh = float(amt / total)
        out.append(Segment(label=lab, share=sh, amount=float(amt), is_rest=(lab == rest_name), draw=(amt > 0)))

    # keine separate negative list im bucket-mode
    return out, []


def apply_min_share_to_rest(
    segs: List[Segment],
    threshold: float,
    rest_label: str
) -> List[Segment]:
    """
    Aggregiert alle Nicht-Rest-Segmente mit |share| <= threshold in den Rest.
    - funktioniert für Top-N und Bucket
    - negative shares werden ebenfalls (betragsmäßig) berücksichtigt
    """
    if threshold is None:
        return segs
    threshold = float(threshold)
    if threshold <= 0:
        return segs

    rest_idx = None
    for i, s in enumerate(segs):
        if s.is_rest or str(s.label) == str(rest_label):
            rest_idx = i
            break

    # wenn kein Rest existiert, lege einen an
    if rest_idx is None:
        segs = list(segs) + [Segment(label=rest_label, share=0.0, amount=0.0, is_rest=True, draw=True)]
        rest_idx = len(segs) - 1

    rest_seg = segs[rest_idx]

    keep: List[Segment] = []
    moved_amount = 0.0
    moved_share = 0.0

    for s in segs:
        if s is rest_seg:
            continue
        if s.is_rest:
            continue

        # alles <= threshold wandert in Rest
        if abs(float(s.share)) <= threshold:
            moved_amount += float(s.amount)
            moved_share += float(s.share)
        else:
            keep.append(s)

    # Rest aktualisieren
    new_rest_amount = float(rest_seg.amount) + moved_amount
    new_rest_share = float(rest_seg.share) + moved_share

    # Rest draw aktualisieren (nur zeichnen wenn netto positiv)
    new_rest = Segment(
        label=rest_label,
        share=new_rest_share,
        amount=new_rest_amount,
        is_rest=True,
        draw=(new_rest_amount > 0)
    )

    # Reihenfolge beibehalten: keep (in bestehender Reihenfolge), Rest am Ende
    return keep + [new_rest]


def prepare_bar_segments(d: pd.DataFrame, bar_cfg: Dict[str, Any], cfg: Dict[str, Any]) -> Tuple[str, List[Segment], List[Segment]]:
    title = bar_cfg.get("title") or str(bar_cfg.get("dim_col") or "Dimension")
    dim_col = str(bar_cfg.get("dim_col") or "").strip()
    if not dim_col:
        raise ValueError(f"Dimension '{title}' übersprungen (dim_col leer).")
    if dim_col not in d.columns:
        raise ValueError(f"Dimension '{title}' übersprungen (Spalte '{dim_col}' fehlt).")
    top_n = int(bar_cfg.get("top_n", 6))
    rest_label = cfg.get("rest_label", "Rest")
    missing_token = cfg.get("missing_token", "__MISSING__")

    df_dim = aggregate_by_dimension(d, dim_col, missing_token)

    bm = bar_cfg.get("bucket_mode") or {"enabled": False}

    min_share = float((cfg.get("plot", {}) or {}).get("min_share_to_rest", 0.05) or 0.0)

    if bm.get("enabled", False):
        segs, neg_list = build_bucket_segments(
            df_dim,
            bm,
            rest_label=rest_label,
            missing_token=missing_token
        )
        # Bucket mode: still fold tiny named buckets into Rest.
        segs = apply_min_share_to_rest(segs, threshold=min_share, rest_label=rest_label)
    else:
        segs, neg_list = build_top_n_segments(
            df_dim,
            top_n,
            rest_label=rest_label,
            missing_token=missing_token,
            min_share=min_share,
        )

    # Negative neu prüfen (falls kleine Negative jetzt im Rest sind)
    if neg_list:
        remaining_labels = {s.label for s in segs if not s.is_rest}
        neg_list = [n for n in neg_list if n.label in remaining_labels]

    # Negative-Prüfung
    if neg_list:
        details = ", ".join([f"{n.label} ({n.share * 100:.1f}%)" for n in neg_list])
        raise ValueError(
            f"Folgende Dimension wird aufgrund negativer Segmente nicht im Chart abgebildet: '{title}': {details}. "
            f"Um die Dimension abzubilden, muss das negative Segment als Teil von 'Rest', "
            f"zu einem insgesamt positivem Segment aggregiert werden --> weniger Segmente."
        )

    # Rest sicherstellen
    if not any(s.is_rest for s in segs):
        segs.append(Segment(label=rest_label, share=0.0, amount=0.0, is_rest=True))

    return title, segs, neg_list


# Plotting
# ----------------------------
def finssentials_blue_gradient(n: int) -> List:
    """Light (top/Top1) → darker (toward Rest). Rest itself stays hatched white."""
    if n <= 1:
        return [_BLUE_CMAP(0.55)]
    xs = np.linspace(0.15, 0.95, n)
    return [_BLUE_CMAP(float(x)) for x in xs]


def purple_gradient(n: int) -> List:
    """Backward-compatible alias."""
    return finssentials_blue_gradient(n)


def _wrap_label(s: str, width: int = 14, max_lines: int = 2) -> str:
    lines = textwrap.wrap(str(s), width=width)
    if len(lines) > max_lines:
        lines = lines[:max_lines]
        lines[-1] = lines[-1].rstrip(".") + "..."
    return "\n".join(lines)


def _draw_bars_on_ax(
    ax: Any,
    bars: List[Tuple[str, List[Segment], List[Segment]]],
    plot_cfg: Dict[str, Any],
    *,
    period_title: str = "",
) -> dict:
    """Draw one period's stacked bars onto ``ax``. Returns small-segment map."""
    label_min_share = float(plot_cfg.get("label_min_share", 0.05))
    x_step = float(plot_cfg.get("x_step", 0.3))
    x = np.arange(len(bars)) * x_step
    width = 0.15
    show_values = bool(plot_cfg.get("show_values", False))
    seg_fs = float(plot_cfg.get("segment_fontsize", 7))
    small_by_bar: dict[str, list[tuple[float, str]]] = {}

    if period_title:
        # Same look as Excel subtitle above the embedded image (Calibri 12, brand navy).
        ax.set_title(
            period_title,
            fontproperties=HEADER_FONT_PROP,
            fontsize=12,
            color=BRAND_BLUE,
            pad=8,
            loc="center",
        )

    for i, (bar_title, segs, _negs) in enumerate(bars):
        top_parts = [s for s in segs if not s.is_rest]
        rest_part = next((s for s in segs if s.is_rest), None)
        drawable_top = [s for s in top_parts if s.draw and s.share > 0]
        cols = finssentials_blue_gradient(len(drawable_top))
        color_map = {drawable_top[j].label: cols[j] for j in range(len(drawable_top))}
        # Lightest colour = first gradient stop = Top1 (drawn on top).
        lightest_label = drawable_top[0].label if drawable_top else None

        draw_list = []
        if rest_part and rest_part.draw and rest_part.share > 0:
            draw_list.append(rest_part)
        draw_list += list(reversed([s for s in drawable_top if s.share > 0]))

        bottom = 0.0
        for seg in draw_list:
            if seg.share <= 0:
                continue

            if seg.is_rest:
                ax.bar(
                    x[i],
                    seg.share,
                    width,
                    bottom=bottom,
                    facecolor="white",
                    edgecolor="black",
                    linewidth=0.8,
                    hatch="///",
                )
            else:
                ax.bar(
                    x[i],
                    seg.share,
                    width,
                    bottom=bottom,
                    color=color_map.get(seg.label),
                    edgecolor="white",
                    linewidth=0.6,
                )
                seg_color = color_map.get(seg.label)
                left_x = x[i] - width / 2
                right_x = x[i] + width / 2
                ax.vlines(
                    [left_x, right_x],
                    bottom,
                    bottom + seg.share,
                    colors=[seg_color, seg_color],
                    linewidth=0.9,
                    zorder=2,
                )

            if seg.share >= label_min_share:
                label_wrapped = _wrap_label(seg.label, width=14, max_lines=2)
                if show_values:
                    txt = f"{label_wrapped}\n{seg.share * 100:.1f}%\n{seg.amount:,.0f}"
                else:
                    txt = f"{label_wrapped}\n{seg.share * 100:.1f}%"

                text_kwargs: Dict[str, Any] = dict(ha="center", va="center", fontsize=seg_fs)
                if seg.is_rest:
                    text_kwargs.update(
                        dict(
                            color="black",
                            clip_on=True,
                            bbox=dict(
                                facecolor="white",
                                edgecolor="none",
                                alpha=0.80,
                                boxstyle="round,pad=0.25",
                            ),
                        )
                    )
                else:
                    seg_color = color_map.get(seg.label) or (0.3, 0.3, 0.3, 1.0)
                    # White on medium/dark blues; black only on the lightest (top) segment.
                    txt_color = "black" if seg.label == lightest_label else "white"
                    text_kwargs.update(
                        dict(
                            color=txt_color,
                            weight=400,
                            clip_on=False,
                            bbox=dict(
                                facecolor=seg_color,
                                edgecolor="none",
                                alpha=1.0,
                                boxstyle="round,pad=0.20",
                            ),
                        )
                    )
                ax.text(x[i], bottom + seg.share / 2, txt, zorder=10, **text_kwargs)
            else:
                if bar_title not in small_by_bar:
                    small_by_bar[bar_title] = []
                if show_values:
                    line = f"{seg.label} ({seg.share * 100:.1f}%) {seg.amount:,.0f}"
                else:
                    line = f"{seg.label} ({seg.share * 100:.1f}%)"
                small_by_bar[bar_title].append((seg.share, line))

            bottom += seg.share

    ax.set_xticks(x)
    labels = [b[0] for b in bars]
    rotation = float(plot_cfg.get("xtick_rotation", 22 if len(bars) >= 4 else 0))
    ax.set_xticklabels(
        labels,
        fontproperties=FONT_PROP,
        rotation=rotation,
        ha="right" if rotation else "center",
        rotation_mode="anchor" if rotation else "default",
    )
    ax.tick_params(axis="x", labelsize=8, pad=2)
    ax.set_ylim(0, 1.0)
    ax.set_yticks([])
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    ax.spines["left"].set_visible(False)
    ax.set_xlim(-0.2, max(x[-1] + 0.25, 0.5) if len(x) else 1.0)
    return small_by_bar


def plot_breakdown(
    bars: List[Tuple[str, List[Segment], List[Segment]]],
    cfg: Dict[str, Any],
    period_label: str,
) -> Tuple[plt.Figure, Any, List[str], Optional[plt.Text]]:
    """Draw a single period chart (one image per period)."""
    plot_cfg = cfg.get("plot", {}) or {}
    figsize = plot_cfg.get("figsize", (5.0, 5.48))
    dpi = int(plot_cfg.get("dpi", 180))

    fig, ax = plt.subplots(figsize=figsize, dpi=dpi)
    small_by_bar = _draw_bars_on_ax(ax, bars, plot_cfg, period_title=str(period_label or ""))

    legend_lines: List[str] = []
    for bar_title, items in small_by_bar.items():
        items_sorted = sorted(items, key=lambda x: x[0], reverse=True)
        for _share, line in items_sorted:
            legend_lines.append(f"{bar_title}: {line}")

    txt_obj = None
    if legend_lines:
        fig.tight_layout(rect=[0.0, 0.12, 1.0, 1.0])
        txt_obj = fig.text(
            0.02,
            0.02,
            "Small segments:\n" + "\n".join(legend_lines[:30]),
            ha="left",
            va="bottom",
            fontsize=6,
        )
    else:
        fig.tight_layout()

    fig.subplots_adjust(left=0.06, right=0.98)
    return fig, ax, legend_lines, txt_obj


# Output naming
def clean_filename(s: str) -> str:
    s = str(s).strip().replace(" ", "_")
    return re.sub(r"[^\w\-_.]", "", s)


def build_output_paths(cfg: Dict[str, Any], period_label: str) -> Tuple[str, str]:
    out_dir = cfg.get("output_dir", ".")
    os.makedirs(out_dir, exist_ok=True)
    prefix = clean_filename(cfg.get("output_prefix", "breakdown"))
    mode = cfg.get("calc_mode", "invoice")
    pm = cfg.get("period_mode", "FY")
    y = int(cfg["current_year"])
    m = int(cfg["current_month"])

    base = f"{prefix}_{period_label}_cy{str(y)[-2:]}_{m:02d}_{mode}"
    png_path = os.path.join(out_dir, base + ".png")
    xlsx_path = os.path.join(out_dir, base + ".xlsx")
    return png_path, xlsx_path


def _display_image_size(cfg: Dict[str, Any]) -> Tuple[int, int]:
    plot_cfg = cfg.get("plot", {}) or {}
    fig_w, fig_h = plot_cfg.get("figsize", (5.0, 5.48))
    display_dpi = 180
    scale = 0.5
    return int(float(fig_w) * display_dpi * scale), int(float(fig_h) * display_dpi * scale)


# Excel export with image + header offsets
def _write_vertical_bars_sheet(
    wb,
    cfg: Dict[str, Any],
    plot_png_paths: List[str],
    legend_png_paths: List[Optional[str]],
    period_label: str,
    sheet_name: str,
) -> str:
    excel_cfg = cfg.get("excel", {}) or {}
    col_offset = int(excel_cfg.get("col_offset", 3))
    set_widths = bool(excel_cfg.get("set_column_widths", True))
    legend_scale = float(excel_cfg.get("legend_scale", 0.35))
    legend_max_width_px = excel_cfg.get("legend_max_width_px", None)
    img_w, img_h = _display_image_size(cfg)
    # Pack charts side-by-side from image display width (minimal gap).
    approx_px_per_col = 7.0 * 14.0  # default Excel col width 14 ≈ 98px
    auto_step = max(1, int(math.ceil(float(img_w) / approx_px_per_col)))
    col_step = int(excel_cfg["chart_col_step"]) if excel_cfg.get("chart_col_step") not in (None, "") else auto_step
    start_col = 4  # D
    plot_row = 6
    legend_row = plot_row + max(1, int(math.ceil(float(img_h) / 20.0))) + 1

    if sheet_name in wb.sheetnames:
        del wb[sheet_name]
    ws = wb.create_sheet(sheet_name)

    brand = BRAND_BLUE.replace("#", "FF")
    title_font = Font(name=THEME.font_name, size=24, color=brand)
    sub_font = Font(name=THEME.font_name, size=12, color=brand)
    info_font = Font(name=THEME.font_name, size=9, color=brand, bold=True)

    ws["D1"].value = cfg.get("title", "")
    ws["D1"].font = title_font

    suffix = str(cfg.get("subtitle_suffix", "")).strip()
    suffix_txt = f" ({suffix})" if suffix else ""
    table_name = f"{cfg.get('table', '')}{suffix_txt}"

    ws["D2"].value = f"{table_name} | {period_label}"
    ws["D2"].font = sub_font

    ws["D4"].value = f"{cfg.get('company', '')} | {table_name}"
    ws["D4"].font = info_font

    fill_grey = PatternFill(fill_type="solid", fgColor="FFF3F1EF")
    for r in range(1, 200):
        for c in range(1, 1 + col_offset):
            ws.cell(row=r, column=c).fill = fill_grey

    fill_white = PatternFill(fill_type="solid", fgColor="FFFFFFFF")
    for r in range(1, 200):
        for c in range(1 + col_offset, 40):
            ws.cell(row=r, column=c).fill = fill_white
    paths = list(plot_png_paths or [])
    if not paths:
        ws.cell(row=plot_row, column=start_col).value = "(Plot-Bild nicht gefunden)"

    for i, plot_png_path in enumerate(paths):
        anchor = f"{get_column_letter(start_col + i * col_step)}{plot_row}"
        if plot_png_path and os.path.exists(plot_png_path):
            img = _xl_image_from_path(plot_png_path)
            img.width = img_w
            img.height = img_h
            ws.add_image(img, anchor)
        else:
            ws[anchor].value = f"(Plot-Bild nicht gefunden: {plot_png_path})"

        legend_png_path = legend_png_paths[i] if i < len(legend_png_paths) else None
        if legend_png_path and os.path.exists(legend_png_path):
            leg = _xl_image_from_path(legend_png_path)
            leg.width = int(leg.width * legend_scale)
            leg.height = int(leg.height * legend_scale)
            if legend_max_width_px is not None:
                legend_max_width_px_i = int(legend_max_width_px)
                if leg.width > legend_max_width_px_i:
                    ratio = legend_max_width_px_i / leg.width
                    leg.width = int(leg.width * ratio)
                    leg.height = int(leg.height * ratio)
            leg_anchor = f"{get_column_letter(start_col + i * col_step)}{legend_row}"
            ws.add_image(leg, leg_anchor)

    if set_widths:
        for c in range(1, 40):
            ws.column_dimensions[get_column_letter(c)].width = 14
        ws.column_dimensions["A"].width = 3
        ws.column_dimensions["B"].width = 3
        ws.column_dimensions["C"].width = 3
        ws.column_dimensions["D"].width = 20

    ws.sheet_properties.outlinePr.summaryRight = True
    for col_letter in ["A", "B", "C"]:
        ws.column_dimensions[col_letter].outlineLevel = 1
        ws.column_dimensions[col_letter].hidden = True
    ws.column_dimensions["D"].collapsed = True
    return sheet_name


def export_excel_with_images(
    cfg: Dict[str, Any],
    plot_png_path: Any,
    legend_png_path: Any,
    xlsx_path: str,
    period_label: str,
    wb=None,
    *,
    save: bool = True,
    sheet_name: Optional[str] = None,
) -> Optional[str]:
    excel_cfg = cfg.get("excel", {}) or {}
    if not excel_cfg.get("enabled", True):
        return None

    # Accept single path or list of paths (one image per period).
    plot_paths = list(plot_png_path) if isinstance(plot_png_path, (list, tuple)) else [plot_png_path]
    if legend_png_path is None:
        legend_paths: List[Optional[str]] = [None] * len(plot_paths)
    elif isinstance(legend_png_path, (list, tuple)):
        legend_paths = list(legend_png_path)
    else:
        legend_paths = [legend_png_path]

    target_name = sheet_name or excel_cfg.get("sheet_name") or cfg.get("base_sheet_name") or "Net sales breakdown"
    owns_wb = wb is None
    if owns_wb:
        wb = Workbook()
        default = wb.active
        if default is not None and default.title == "Sheet":
            wb.remove(default)

    written = _write_vertical_bars_sheet(
        wb, cfg, plot_paths, legend_paths, period_label, str(target_name)
    )
    if save and xlsx_path:
        wb.save(xlsx_path)
    return written


def split_by_text_bbox(
    fig: plt.Figure,
    buf_rgba: np.ndarray,
    txt_obj: Optional[plt.Text],
    out_plot_png: str,
    out_legend_png: Optional[str],
    dpi: int,
    pad_px: int = 12
) -> Tuple[str, Optional[str]]:

    H, W, _ = buf_rgba.shape

    if txt_obj is None or out_legend_png is None:
        Image.fromarray(buf_rgba).save(out_plot_png, dpi=(dpi, dpi))
        return out_plot_png, None

    renderer = fig.canvas.get_renderer()
    bbox = txt_obj.get_window_extent(renderer=renderer)

    x0 = max(0, int(bbox.x0) - pad_px)
    x1 = min(W, int(bbox.x1) + pad_px)

    y0_bl = max(0, int(bbox.y0) - pad_px)
    y1_bl = min(H, int(bbox.y1) + pad_px)

    y_top = max(0, H - y1_bl)
    y_bot = min(H, H - y0_bl)

    legend_rgba = buf_rgba[y_top:y_bot, x0:x1, :]
    Image.fromarray(legend_rgba).save(out_legend_png, dpi=(dpi, dpi))

    plot_cut = max(0, y_top - pad_px)
    plot_rgba = buf_rgba[:plot_cut, :, :]
    Image.fromarray(plot_rgba).save(out_plot_png, dpi=(dpi, dpi))

    return out_plot_png, out_legend_png


def _bars_for_period(
    d_base: pd.DataFrame, cfg: Dict[str, Any], start: pd.Timestamp, end: pd.Timestamp
) -> List[Tuple[str, List[Segment], List[Segment]]]:
    d = d_base.copy()
    d["amount"] = compute_revenue(d, start, end, cfg)
    bars_out: List[Tuple[str, List[Segment], List[Segment]]] = []
    for bar_cfg in cfg["bars"]:
        try:
            title, segs, negs = prepare_bar_segments(d, bar_cfg, cfg)
            bars_out.append((title, segs, negs))
        except ValueError as e:
            print(str(e))
            continue
    return bars_out


def _render_breakdown_assets(
    df: pd.DataFrame, cfg: Dict[str, Any]
) -> Tuple[Dict[str, Any], str, List[str], List[Optional[str]], str]:
    """Normalize config, build one PNG per period; returns paths lists."""
    cfg = normalize_config(cfg)
    d_base = preprocess_input(df, cfg)
    periods = get_comparison_periods(cfg)
    period_panels: List[Tuple[str, List[Tuple[str, List[Segment], List[Segment]]]]] = []
    for start, end, label in periods:
        bars_out = _bars_for_period(d_base, cfg, start, end)
        if bars_out:
            period_panels.append((label, bars_out))

    if not period_panels:
        raise ValueError("Keine Balken konnten berechnet werden (Dimensionen fehlen oder negativ).")

    period_label = " · ".join(p[0] for p in period_panels)
    file_period = "_".join(p[0] for p in period_panels)
    png_path, xlsx_path = build_output_paths(cfg, file_period)
    plot_cfg = cfg.get("plot", {}) or {}
    dpi = int(plot_cfg.get("dpi", 180))

    plot_pngs: List[str] = []
    legend_pngs: List[Optional[str]] = []
    for idx, (p_label, bars_out) in enumerate(period_panels):
        fig, _ax, _legend_items, txt_obj = plot_breakdown(bars_out, cfg, p_label)
        plot_png = png_path.replace(".png", f"_{idx + 1}_{clean_filename(p_label)}_plot.png")
        legend_png = png_path.replace(".png", f"_{idx + 1}_{clean_filename(p_label)}_legend.png")
        fig.canvas.draw()
        buf = np.asarray(fig.canvas.buffer_rgba())
        split_by_text_bbox(
            fig=fig,
            buf_rgba=buf,
            txt_obj=txt_obj,
            out_plot_png=plot_png,
            out_legend_png=legend_png,
            dpi=dpi,
            pad_px=12,
        )
        plt.close(fig)
        plot_pngs.append(plot_png)
        legend_pngs.append(legend_png if os.path.exists(legend_png) else None)

    return cfg, period_label, plot_pngs, legend_pngs, xlsx_path


def _cleanup_temp_pngs(plot_pngs: List[str], legend_pngs: List[Optional[str]]) -> None:
    for p in list(plot_pngs) + list(legend_pngs):
        if p and os.path.exists(p):
            try:
                os.remove(p)
            except PermissionError:
                pass


def run_breakdown_chart(df: pd.DataFrame, cfg: Dict[str, Any]):
    cfg, period_label, plot_pngs, legend_pngs, xlsx_path = _render_breakdown_assets(df, cfg)
    excel_path_out = None
    try:
        if cfg.get("excel", {}).get("enabled", True):
            export_excel_with_images(cfg, plot_pngs, legend_pngs, xlsx_path, period_label)
            excel_path_out = xlsx_path
    finally:
        _cleanup_temp_pngs(plot_pngs, legend_pngs)
    return None, excel_path_out


def run_vertical_bars_in_workbook(cfg: dict, wb) -> str:
    """Fast Track / session mode: write Vertical bars into an open workbook."""
    file_path = str(cfg.get("file_path") or "").strip()
    sheet_name = str(cfg.get("sheet_name") or "").strip()
    if not file_path:
        raise ValueError("vertical_bars requires file_path")
    df = pd.read_excel(file_path, sheet_name=sheet_name or 0, engine="openpyxl")

    out_dir = str(cfg.get("output_file_path") or cfg.get("output_dir") or ".")
    cfg = {**cfg, "output_dir": out_dir, "output_prefix": cfg.get("output_prefix") or "vertical_bars"}
    excel_cfg = dict(cfg.get("excel") or {})
    excel_cfg.setdefault("sheet_name", cfg.get("base_sheet_name") or "")
    excel_cfg["enabled"] = True
    cfg["excel"] = excel_cfg

    cfg, period_label, plot_pngs, legend_pngs, _xlsx = _render_breakdown_assets(df, cfg)
    sheet_title = str(
        (cfg.get("excel") or {}).get("sheet_name")
        or cfg.get("base_sheet_name")
        or f"{sales_basis_labels(cfg.get('sales_basis')).metric_label} breakdown"
    )
    try:
        from funktionssammlung import get_next_sheet_name_from_wb

        target = get_next_sheet_name_from_wb(wb, sheet_title)
    except Exception:
        target = sheet_title
        if target in wb.sheetnames:
            target = f"{target}_1"

    try:
        written = export_excel_with_images(
            cfg,
            plot_pngs,
            legend_pngs,
            xlsx_path="",
            period_label=period_label,
            wb=wb,
            save=False,
            sheet_name=target,
        )
        return str(written or target)
    finally:
        _cleanup_temp_pngs(plot_pngs, legend_pngs)


# Backward-compatible alias
run_horizontal_bars_in_workbook = run_vertical_bars_in_workbook


if __name__ == "__main__":
    _desktop = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Desktop")
    _src = os.path.join(_desktop, "Sales data", "SaaS-Sales.xlsx")
    if not os.path.isfile(_src):
        _src = os.path.join(
            os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
            "finssentials",
            "Desktop",
            "Sales data",
            "SaaS-Sales.xlsx",
        )
    _cfg = dict(CONFIG)
    # Mid-FY as-of → YTD + 2 completed FYs
    _cfg["current_year"] = 2025
    _cfg["current_month"] = 3
    _cfg["fiscal_year_end_month"] = 7
    _cfg["fiscal_year_end_day"] = 31
    _cfg["output_dir"] = _desktop if os.path.isdir(_desktop) else "."
    _df = pd.read_excel(_src, sheet_name="Database", engine="openpyxl")
    _png, _xlsx = run_breakdown_chart(_df, _cfg)
    print("Excel gespeichert unter:", _xlsx)
