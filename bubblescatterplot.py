"""
Bubble scatter plot (GP vs. GM) with optional hierarchy, Excel export, separate legend PNG.
Reconstructed from archived script photos; uses Funktionssammlung for filters, periods, amounts, colors.
"""

from __future__ import annotations

import ast
import copy
import json
import os
import sys
import textwrap
from io import BytesIO
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import matplotlib.font_manager as fm
import matplotlib.pyplot as plt
import matplotlib.ticker as ticker
import numpy as np
import pandas as pd
from matplotlib.colors import to_hex, to_rgba
from matplotlib.lines import Line2D
from matplotlib.patches import Rectangle
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from funktionssammlung import apply_filters, compute_amount, get_period, pick_colors_hardcoded
from revenue_reconciliation import (
    RECON_DIFFERENCE_LABEL,
    sales_basis_labels,
    write_reconciliation_cells,
)

# Font — shared Excel/chart theme (no local TTF paths)
from gst_excel_theme import THEME, apply_matplotlib_theme, matplotlib_font_properties

apply_matplotlib_theme(size=7)
FONT_PROP = matplotlib_font_properties()

# -----------------------------------------------------------------------------
# INPUT (Beispiel — Pfade anpassen)
# -----------------------------------------------------------------------------
# file_path = r"...\Database_Sonoro.xlsx"
# df = pd.read_excel(file_path, sheet_name="Database", engine="openpyxl")

# -----------------------------------------------------------------------------
# DEFAULT_CONFIG (overridden by JSON from FDD bot / CLI)
# -----------------------------------------------------------------------------
DEFAULT_CONFIG: Dict[str, Any] = {
    "title": "Project Draft",
    "table": "Margin analyses",
    "company": "Draft AG",
    "subtitle_suffix": "GP by segment",
    # Time input (as-of immer Monatsende)
    "current_year": 2025,
    "current_month": 7,
    "fiscal_year_end_month": 7,
    "fiscal_year_end_day": 31,
    # Period mode: "FY" | "YTD" | "LTM"
    "period_mode": "FY",
    # Calc mode: "invoice" | "accrual"
    "calc_mode": "invoice",
    # Date columns
    "invoice_col": "Invoice date",
    "start_col": "Contract Start Date",
    "end_col": "Contract End Date",
    # Revenue / COGS
    "revenue_col": "Revenue",
    "cogs_col": "Einstandspreis Gesamt",
    "bubblesize_col": "",
    "profit_mode": "cost",
    "profit_col": "",
    "filters": {
        "enabled": False,
        "rules": [],
    },
    "apply_fx": False,
    "fx_col": "Functional FX Rate",
    "group_cols": ["Core products IM", "Product group new"],
    "max_bubbles": 30,
    "rank_by": "revenue",
    "grouping": {
        "max_items_parent": 6,
        "max_items_child": 6,
        "max_children_per_parent": 6,
        "create_other_bucket_parent": True,
        "create_other_bucket_child": True,
        "combine_input_other": False,
        "input_other_tokens": {"Other"},
        "other_label": "Rest",
    },
    "plot": {
        "figsize": (1200 / 300, 600 / 300),
        "dpi": 300,
        "save_dpi": 300,
        "x_label": "GM (in%)",
        "y_label": "GP",
        "y_unit": "m",
        "y_unit_label": "in EURm",
        "x_as_percent": True,
        "x_tick_step": 5,
        "bubble_area_min": 60,
        "bubble_area_max": 900,
        "x_ref": None,
        "y_ref": None,
        "x_lim": None,
        "y_lim": None,
        "axis_label_size": 7,
        "tick_label_size": 7,
        "legend": {
            "enabled": True,
            "position": "bottom",
            "ncol": 6,
            "font_size": 4,
            "max_items": 36,
            "marker_size": 4.5,
            "sort_by": "revenue",
            "group_mode": "parent_cluster",
            "box_linestyle": "--",
            "box_linewidth": 0.5,
            "cluster_box_pad_px": 2,
            "cluster_title_gap_px": 2,
            "cluster_top_pad_px": 2,
            "cluster_gap_px": 8,
        },
        "show_labels": False,
        "label_wrap": 16,
        "label_max_lines": 2,
        "style": {
            "axis_linewidth": 0.45,
            "tick_width": 0.45,
            "tick_length": 4,
            "scatter_edge_width": 0.45,
            "ref_linewidth": 0.45,
        },
    },
    "output_file_path": "",
    "output_path": "",
    "output_dir": os.path.join(os.path.dirname(os.path.abspath(__file__)), "output", "bubble_scatter"),
    "output_prefix": "bubble_gp_gm",
    "file_path": "",
    "sheet_name": "Sheet1",
    "case_id": "",
    "excel": {
        "enabled": True,
        "row_offset": 4,
        "col_offset": 3,
        "sheet_name": "Margin analyses",
        "image_anchor_cell": "D6",
        "set_column_widths": True,
        "display_dpi": 180,
        "image_scale": 1.0,
        "legend_anchor_cell": "D28",
        "legend_scale": 0.5,
        "helper_table_row": 42,    },
}


def _deep_merge_defaults(base: Dict[str, Any], override: Dict[str, Any]) -> Dict[str, Any]:
    out = copy.deepcopy(base)
    for key, val in override.items():
        if isinstance(val, dict) and isinstance(out.get(key), dict):
            out[key] = _deep_merge_defaults(out[key], val)
        else:
            out[key] = val
    return out


def _parse_ltm_month(cfg: Dict[str, Any]) -> None:
    ltm = str(cfg.get("ltm_month") or "").strip()
    if not ltm:
        return
    parts = ltm.split("-")
    try:
        y, m = int(parts[0]), int(parts[1])
        if 1 <= m <= 12:
            cfg["current_year"] = y
            cfg["current_month"] = m
    except (ValueError, IndexError):
        pass


def normalize_config(cfg: Dict[str, Any]) -> Dict[str, Any]:
    out = _deep_merge_defaults(DEFAULT_CONFIG, cfg or {})
    _parse_ltm_month(out)

    for key in ("current_year", "current_month", "revenue_col", "invoice_col"):
        if key not in out or out[key] in (None, ""):
            raise ValueError(f"CONFIG['{key}'] muss gesetzt sein.")

    out["current_year"] = int(out["current_year"])
    out["current_month"] = int(out["current_month"])
    if not (1 <= out["current_month"] <= 12):
        raise ValueError("CONFIG['current_month'] muss 1..12 sein.")

    fy_end_m = int(out.get("fiscal_year_end_month", 12))
    fy_end_d = int(out.get("fiscal_year_end_day", 31))
    if not (1 <= fy_end_m <= 12):
        raise ValueError("CONFIG['fiscal_year_end_month'] muss 1..12 sein.")
    if not (1 <= fy_end_d <= 31):
        raise ValueError("CONFIG['fiscal_year_end_day'] muss 1..31 sein.")
    out["fiscal_year_end_month"] = fy_end_m
    out["fiscal_year_end_day"] = fy_end_d

    out["period_mode"] = str(out.get("period_mode", "FY")).strip().upper()
    if out["period_mode"] not in {"FY", "YTD", "LTM"}:
        raise ValueError("CONFIG['period_mode'] muss 'FY', 'YTD' oder 'LTM' sein.")

    out["calc_mode"] = str(out.get("calc_mode", "invoice")).strip().lower()
    if out["calc_mode"] not in {"invoice", "accrual"}:
        raise ValueError("CONFIG['calc_mode'] muss 'invoice' oder 'accrual' sein.")
    out["invoice_col"] = str(out["invoice_col"]).strip()
    out["invoice_mapping_mode"] = str(
        out.get("invoice_mapping_mode") or "date"
    ).strip().lower()
    if out["invoice_mapping_mode"] not in {"date", "year"}:
        raise ValueError("CONFIG['invoice_mapping_mode'] muss 'date' oder 'year' sein.")
    if out["calc_mode"] == "accrual":
        for key in ("start_col", "end_col"):
            if not str(out.get(key) or "").strip():
                raise ValueError(f"CONFIG['{key}'] muss bei calc_mode='accrual' gesetzt sein.")
            out[key] = str(out[key]).strip()
    else:
        out.setdefault("start_col", "")
        out.setdefault("end_col", "")

    out["profit_mode"] = str(out.get("profit_mode", "cost")).strip().lower()
    if out["profit_mode"] not in {"cost", "profit"}:
        raise ValueError("CONFIG['profit_mode'] muss 'cost' oder 'profit' sein.")
    out["revenue_col"] = str(out["revenue_col"]).strip()
    out["cogs_col"] = str(out.get("cogs_col") or "").strip()
    out["profit_col"] = str(out.get("profit_col") or "").strip()
    if out["profit_mode"] == "cost" and not out["cogs_col"]:
        raise ValueError("CONFIG['cogs_col'] muss bei profit_mode='cost' gesetzt sein.")
    if out["profit_mode"] == "profit" and not out["profit_col"]:
        raise ValueError("CONFIG['profit_col'] muss bei profit_mode='profit' gesetzt sein.")

    # bubblesize_col: only keep explicit user value; never inherit stale DEFAULT "Revenue"
    raw_cfg = cfg or {}
    user_bubblesize = str(raw_cfg.get("bubblesize_col") or "").strip()
    out["bubblesize_col"] = user_bubblesize if user_bubblesize else out["revenue_col"]

    out["apply_fx"] = bool(out.get("apply_fx", False))
    if out["apply_fx"]:
        if not out.get("fx_col"):
            raise ValueError("apply_fx=True, aber fx_col fehlt.")
        out["fx_col"] = str(out["fx_col"])

    flt = out.get("filters") or {}
    if not isinstance(flt, dict):
        flt = {"enabled": False, "rules": []}
    rules = flt.get("rules") or []
    flt["enabled"] = bool(flt.get("enabled", False)) and bool(rules)
    flt["rules"] = rules if isinstance(rules, list) else []
    out["filters"] = flt

    gc = out.get("group_cols", [])
    if gc is None:
        gc = []
    if not isinstance(gc, list):
        raise ValueError("CONFIG['group_cols'] muss eine Liste sein ([], [col], oder [parent, child]).")
    if len(gc) > 2:
        raise ValueError("Maximal 2 group_cols erlaubt (parent, child).")
    out["group_cols"] = [str(x) for x in gc]

    out["max_bubbles"] = int(out.get("max_bubbles", 30))
    out["rank_by"] = str(out.get("rank_by", "revenue")).lower()

    out.setdefault("title", "")
    out["company"] = str(
        out.get("company") or out.get("company_name") or ""
    ).strip()
    out.setdefault("table", "Margin analyses")
    labels = sales_basis_labels(out.get("sales_basis"))
    out["sales_basis"] = labels.sales_basis
    plot = dict(out.get("plot") or {})
    default_x = {"GM (in%)", "GM", "NM (in%)", "NM", ""}
    default_y = {"GP", "NP", ""}
    if str(plot.get("x_label") or "").strip() in default_x:
        plot["x_label"] = f"{labels.margin_abbrev} (in%)"
    if str(plot.get("y_label") or "").strip() in default_y:
        plot["y_label"] = labels.profit_abbrev
    out["plot"] = plot
    suffix = str(out.get("subtitle_suffix") or "").strip()
    if suffix in {"", "GP by segment", "NP by segment", "GM by segment", "NM by segment"}:
        out["subtitle_suffix"] = f"{labels.profit_abbrev} by segment"
    else:
        out["subtitle_suffix"] = suffix

    grp = out.get("grouping") or {}
    out["grouping"] = grp
    cap = lambda v, d=6: max(1, min(6, int(v if v is not None and v != "" else d)))
    grp["max_items_child"] = cap(
        grp.get("max_items_child", grp.get("max_children_per_parent", 6))
    )
    grp["max_items_parent"] = cap(
        grp.get("max_items_parent", grp.get("max_parents", grp["max_items_child"]))
    )
    grp["max_children_per_parent"] = grp["max_items_child"]
    grp["max_parents"] = grp["max_items_parent"]
    legacy_other = grp.get("create_other_bucket")
    grp["create_other_bucket_parent"] = bool(
        grp.get(
            "create_other_bucket_parent",
            legacy_other if legacy_other is not None else True,
        )
    )
    grp["create_other_bucket_child"] = bool(
        grp.get(
            "create_other_bucket_child",
            legacy_other if legacy_other is not None else True,
        )
    )
    grp["combine_input_other"] = bool(grp.get("combine_input_other", False))
    tok = grp.get("input_other_tokens", {"Other"})
    if isinstance(tok, str):
        tok = {tok}
    grp["input_other_tokens"] = set(tok) if not isinstance(tok, set) else tok
    grp["other_label"] = str(grp.get("other_label", "Other"))

    p = out.get("plot") or {}
    out["plot"] = p
    figsize = p.get("figsize", (6.0, 3.2))
    if isinstance(figsize, (list, tuple)) and len(figsize) >= 2:
        p["figsize"] = (float(figsize[0]), float(figsize[1]))
    else:
        p["figsize"] = (6.0, 3.2)

    for dk in ["dpi", "render_dpi", "save_dpi"]:
        if dk in p and p[dk] is not None:
            p[dk] = int(p[dk])

    p["axis_label_size"] = int(p.get("axis_label_size", 7))
    p["tick_label_size"] = int(p.get("tick_label_size", 7))

    p["y_unit"] = str(p.get("y_unit", "m")).lower()
    if p["y_unit"] not in {"k", "m"}:
        raise ValueError("plot.y_unit muss 'k' oder 'm' sein.")
    if not str(p.get("y_unit_label") or "").strip():
        p["y_unit_label"] = "in EURk" if p["y_unit"] == "k" else "in EURm"

    p["x_as_percent"] = bool(p.get("x_as_percent", True))
    if p.get("x_tick_step", None) is not None:
        p["x_tick_step"] = float(p["x_tick_step"])

    p["bubble_area_min"] = float(p.get("bubble_area_min", 60))
    p["bubble_area_max"] = float(p.get("bubble_area_max", 900))
    if p["bubble_area_min"] <= 0 or p["bubble_area_max"] <= 0:
        raise ValueError("plot.bubble_area_min/max müssen > 0 sein.")
    if p["bubble_area_max"] < p["bubble_area_min"]:
        raise ValueError("plot.bubble_area_max muss >= bubble_area_min sein.")

    style = p.get("style") or {}
    p["style"] = style
    style.setdefault("axis_linewidth", 0.45)
    style.setdefault("tick_width", 0.45)
    style.setdefault("tick_length", 4)
    style.setdefault("scatter_edge_width", 0.45)
    style.setdefault("ref_linewidth", 0.45)

    for k in ["axis_linewidth", "tick_width", "tick_length", "scatter_edge_width", "ref_linewidth"]:
        style[k] = float(style[k])

    leg = p.get("legend") or {}
    p["legend"] = leg
    leg["enabled"] = bool(leg.get("enabled", True))
    leg["position"] = str(leg.get("position", "bottom")).lower()
    if leg["position"] not in {"bottom", "right"}:
        raise ValueError("plot.legend.position muss 'bottom' oder 'right' sein.")
    leg["ncol"] = int(leg.get("ncol", 4))
    leg["font_size"] = int(leg.get("font_size", 4))
    leg["max_items"] = int(leg.get("max_items", 20))
    leg["marker_size"] = float(leg.get("marker_size", 4.5))

    e = out.get("excel") or {}
    out["excel"] = e
    e["enabled"] = bool(e.get("enabled", True))
    if e["enabled"]:
        e["row_offset"] = int(e.get("row_offset", 4))
        e["col_offset"] = int(e.get("col_offset", 3))
        e["sheet_name"] = str(e.get("sheet_name", "Margin analyses"))
        e["image_anchor_cell"] = str(e.get("image_anchor_cell", "D6"))
        e["set_column_widths"] = bool(e.get("set_column_widths", True))
        e["display_dpi"] = int(e.get("display_dpi", 180))
        e["image_scale"] = float(e.get("image_scale", 1.0))
        e["legend_anchor_cell"] = str(e.get("legend_anchor_cell", "D28"))
        e["legend_scale"] = float(e.get("legend_scale", 0.5))
        if e["legend_scale"] <= 0:
            raise ValueError("excel.legend_scale muss > 0 sein.")

    return out


def preprocess_input(df: pd.DataFrame, cfg: Dict[str, Any]) -> pd.DataFrame:
    d = df.copy()
    d.columns = d.columns.astype(str).str.replace("\u00a0", " ", regex=False).str.strip()
    if cfg["calc_mode"] == "invoice" and cfg["invoice_mapping_mode"] == "date":
        d[cfg["invoice_col"]] = pd.to_datetime(
            d[cfg["invoice_col"]], errors="coerce", dayfirst=True
        )
    if cfg["calc_mode"] == "accrual":
        d[cfg["start_col"]] = pd.to_datetime(
            d[cfg["start_col"]], errors="coerce", dayfirst=True
        )
        d[cfg["end_col"]] = pd.to_datetime(
            d[cfg["end_col"]], errors="coerce", dayfirst=True
        )
    if cfg.get("apply_fx"):
        fx_col = str(cfg["fx_col"])
        d["fx_rate"] = pd.to_numeric(d[fx_col], errors="coerce").fillna(1.0)
        valid_fx = d["fx_rate"].notna() & (d["fx_rate"] != 0)
        amt_cols = {cfg["revenue_col"], cfg["bubblesize_col"]}
        if str(cfg.get("profit_mode", "cost")).lower() == "profit":
            amt_cols.add(cfg["profit_col"])
        else:
            amt_cols.add(cfg["cogs_col"])
        for col in amt_cols:
            if col and col in d.columns:
                d.loc[valid_fx, col] = pd.to_numeric(d.loc[valid_fx, col], errors="coerce") * d.loc[valid_fx, "fx_rate"]

    d = d.dropna(subset=[cfg["invoice_col"]]).copy()
    if cfg["calc_mode"] == "accrual":
        d = d.dropna(subset=[cfg["start_col"], cfg["end_col"]]).copy()
        d = d[d[cfg["end_col"]] >= d[cfg["start_col"]]].copy()
    return d


def _wrap_label(s: str, width: int, max_lines: int) -> str:
    s = str(s)
    lines = textwrap.wrap(s, width=width)
    if len(lines) > max_lines:
        lines = lines[:max_lines]
        lines[-1] = lines[-1].rstrip(".") + "..."
    return "\n".join(lines)


def make_group_key(df: pd.DataFrame, group_cols: List[str]) -> pd.Series:
    if len(group_cols) == 0:
        return pd.Series(["All"] * len(df), index=df.index)
    if len(group_cols) == 1:
        return df[group_cols[0]].astype(str).fillna("Missing").str.strip()
    p = df[group_cols[0]].astype(str).fillna("Missing").str.strip()
    c = df[group_cols[1]].astype(str).fillna("Missing").str.strip()
    return p + " | " + c


def split_parent_child(key: str) -> Tuple[str, str]:
    key = str(key)
    if " | " in key:
        p, c = key.split(" | ", 1)
        return p.strip(), c.strip()
    return key.strip(), key.strip()


def _rollup_other_row(
    parent: str,
    child: str,
    label: str,
    group_key: str,
    rest: pd.DataFrame,
) -> dict:
    rest_rev = float(rest["revenue"].sum())
    rest_gp = float(rest["gp"].sum())
    rest_bub = float(rest["bubblesize"].sum())
    rest_gm = (rest_gp / rest_rev) if abs(rest_rev) > 1e-12 else np.nan
    row: Dict[str, Any] = {
        "__group__": group_key,
        "parent": parent,
        "child": child,
        "label": label,
        "revenue": rest_rev,
        "bubblesize": rest_bub,
        "gp": rest_gp,
        "gm": rest_gm,
    }
    if "cogs" in rest.columns:
        row["cogs"] = float(rest["cogs"].sum())
    return row


def _order_with_other_last(items: List[str], other_label: str) -> List[str]:
    """Keep revenue (or other) sort but force the aggregated Other parent/cluster to the end."""
    other = str(other_label).strip()
    ordered: List[str] = []
    seen: set[str] = set()
    for raw in items:
        s = str(raw)
        if s in seen:
            continue
        seen.add(s)
        ordered.append(s)
    without = [x for x in ordered if x != other]
    if other in seen:
        without.append(other)
    return without


def apply_parent_limit(g: pd.DataFrame, cfg: Dict[str, Any]) -> pd.DataFrame:
    """Single group_col: cap distinct items; overflow → Other."""
    grp_cfg = cfg.get("grouping", {})
    other_label = str(grp_cfg.get("other_label", "Other"))
    create_other_bucket = bool(grp_cfg.get("create_other_bucket_parent", True))
    max_items = int(grp_cfg.get("max_items_parent", grp_cfg.get("max_children_per_parent", 6)))
    max_items = max(1, min(6, max_items))
    g = g.sort_values("revenue", ascending=False)

    if not create_other_bucket:
        return g.head(max_items).reset_index(drop=True)

    top = g.head(max_items - 1)
    rest = g.iloc[max_items - 1 :]

    if rest.empty:
        return g.reset_index(drop=True)

    other_row = _rollup_other_row(other_label, other_label, other_label, other_label, rest)
    return pd.concat([top, pd.DataFrame([other_row])], ignore_index=True).reset_index(drop=True)


def apply_parent_group_limit(g: pd.DataFrame, cfg: Dict[str, Any]) -> pd.DataFrame:
    """Two group_cols: cap parent groups by revenue; overflow parents → one Other parent row."""
    grp_cfg = cfg.get("grouping", {})
    other_label = str(grp_cfg.get("other_label", "Other"))
    create_other_bucket = bool(grp_cfg.get("create_other_bucket_parent", True))
    max_parents = int(grp_cfg.get("max_items_parent", grp_cfg.get("max_parents", 6)))
    max_parents = max(1, min(6, max_parents))

    parent_totals = (
        g.groupby("parent", as_index=False)
        .agg(revenue=("revenue", "sum"))
        .sort_values("revenue", ascending=False)
    )
    parents_ordered = _order_with_other_last(
        parent_totals["parent"].astype(str).tolist(),
        other_label,
    )

    if not create_other_bucket:
        keep = set(parents_ordered[:max_parents])
        return g[g["parent"].astype(str).isin(keep)].reset_index(drop=True)

    keep_parents = parents_ordered[: max_parents - 1]
    keep_set = set(keep_parents)
    kept = g[g["parent"].astype(str).isin(keep_set)].copy()
    overflow = g[~g["parent"].astype(str).isin(keep_set)].copy()

    if overflow.empty:
        return g.reset_index(drop=True)

    other_row = _rollup_other_row(
        other_label,
        other_label,
        other_label,
        f"{other_label} | {other_label}",
        overflow,
    )
    return pd.concat([kept, pd.DataFrame([other_row])], ignore_index=True).reset_index(drop=True)


def apply_child_limit(g: pd.DataFrame, cfg: Dict[str, Any]) -> pd.DataFrame:
    grp_cfg = cfg.get("grouping", {})
    max_children = int(grp_cfg.get("max_items_child", grp_cfg.get("max_children_per_parent", 6)))
    max_children = max(1, min(6, max_children))
    create_other_bucket = bool(grp_cfg.get("create_other_bucket_child", True))
    combine_input_other = bool(grp_cfg.get("combine_input_other", False))
    raw_tokens = grp_cfg.get("input_other_tokens", {"Other"})
    if isinstance(raw_tokens, str):
        raw_tokens = {raw_tokens}
    tokens = {str(t).strip().lower() for t in raw_tokens}
    other_label = str(grp_cfg.get("other_label", "Other"))

    out_rows: List[pd.DataFrame] = []
    for parent, sub_all in g.groupby("parent", sort=False):
        sub_all = sub_all.sort_values("revenue", ascending=False)
        c_norm = sub_all["child"].astype(str).str.strip().str.lower()
        mask_io = c_norm.isin(tokens)
        input_other = sub_all.loc[mask_io].copy()
        sub = sub_all.loc[~mask_io].copy()

        if not create_other_bucket:
            sub_out = sub.head(max_children).copy()
            if not input_other.empty and not combine_input_other:
                sub_out = pd.concat([sub_out, input_other], ignore_index=True)
            out_rows.append(sub_out)
            continue

        top = sub.head(max_children - 1).copy()
        rest = sub.iloc[max_children - 1 :].copy()
        if combine_input_other and not input_other.empty:
            rest = pd.concat([rest, input_other], ignore_index=True)

        if not rest.empty:
            other_row = _rollup_other_row(
                parent,
                other_label,
                f"{other_label} ({parent})",
                f"{parent} | {other_label}",
                rest,
            )
            sub_out = pd.concat([top, pd.DataFrame([other_row])], ignore_index=True)
        else:
            sub_out = top
        if not input_other.empty and not combine_input_other:
            sub_out = pd.concat([sub_out, input_other], ignore_index=True)
        out_rows.append(sub_out)

    return pd.concat(out_rows, ignore_index=True)


def build_bubble_table(d: pd.DataFrame, cfg: Dict[str, Any]) -> pd.DataFrame:
    group_cols = cfg.get("group_cols") or []
    key = make_group_key(d, group_cols)
    dc = d.copy()
    dc["__group__"] = key

    profit_mode = str(cfg.get("profit_mode", "cost")).lower()
    if profit_mode == "profit":
        g = dc.groupby("__group__", as_index=False).agg(
            revenue=("revenue_amt", "sum"),
            gp=("profit_amt", "sum"),
            bubblesize=("bubble_amt", "sum"),
        )
    else:
        g = dc.groupby("__group__", as_index=False).agg(
            revenue=("revenue_amt", "sum"),
            cogs=("cogs_amt", "sum"),
            bubblesize=("bubble_amt", "sum"),
        )
        g["gp"] = g["revenue"] - g["cogs"]
    g["gm"] = np.where(g["revenue"].abs() > 1e-12, g["gp"] / g["revenue"], np.nan)

    if len(group_cols) == 2:
        g["parent"] = g["__group__"].map(lambda x: split_parent_child(str(x))[0])
        g["child"] = g["__group__"].map(lambda x: split_parent_child(str(x))[1])
        g["label"] = g["child"] + " (" + g["parent"] + ")"
        g = apply_parent_group_limit(g, cfg)
        g = apply_child_limit(g, cfg)
    else:
        g["parent"] = g["__group__"].astype(str)
        g["child"] = g["__group__"].astype(str)
        g["label"] = g["__group__"].astype(str)
        g = apply_parent_limit(g, cfg)

    rank_by = str(cfg.get("rank_by", "revenue")).lower()
    g = g.sort_values("gp" if rank_by == "gp" else "revenue", ascending=False).reset_index(drop=True)

    max_b = int(cfg.get("max_bubbles", 36))
    if max_b > 0 and len(g) > max_b:
        g = g.head(max_b).reset_index(drop=True)
    return g


def plot_bubble_chart(
    g: pd.DataFrame, cfg: Dict[str, Any], period_label: str
) -> Tuple[plt.Figure, plt.Axes, Optional[list]]:
    _ = period_label
    plot_cfg = cfg.get("plot", {}) or {}
    figsize = plot_cfg.get("figsize", (6.0, 3.2))
    dpi = int(plot_cfg.get("dpi", 180))
    fig, ax = plt.subplots(figsize=figsize, dpi=dpi)
    style_cfg = plot_cfg.get("style", {}) or {}
    axis_lw = float(style_cfg.get("axis_linewidth", 0.8))
    for spine in ax.spines.values():
        spine.set_linewidth(axis_lw)

    fig.subplots_adjust(left=0.1, right=0.98, top=0.95, bottom=0.2)

    x_as_percent = bool(plot_cfg.get("x_as_percent", True))
    x = g["gm"] * (100.0 if x_as_percent else 1.0)

    y_unit = str(plot_cfg.get("y_unit", "m")).lower()
    if y_unit not in {"k", "m"}:
        raise ValueError("plot.y_unit muss 'k' oder 'm' sein.")
    y_div = 1_000.0 if y_unit == "k" else 1_000_000.0
    y = g["gp"] / y_div

    s_raw = g["bubblesize"].abs().fillna(0.0).to_numpy()
    s_min = float(plot_cfg.get("bubble_area_min", 60))
    s_max = float(plot_cfg.get("bubble_area_max", 900))
    if np.nanmax(s_raw) <= 1e-12:
        s = np.full_like(s_raw, s_min)
    else:
        s = s_min + (s_raw - np.nanmin(s_raw)) * (s_max - s_min) / (np.nanmax(s_raw) - np.nanmin(s_raw) + 1e-12)

    group_cols = cfg.get("group_cols", [])
    group_keys = g["__group__"].tolist()

    if len(group_cols) == 2 and "parent" in g.columns and "child" in g.columns:
        rank_by = str(cfg.get("rank_by", "revenue")).lower()
        value_col_for_color_rank = "gp" if rank_by == "gp" else "revenue"
        other_label = str((cfg.get("grouping", {}) or {}).get("other_label", "Other"))
        parent_order = (
            g.groupby("parent", as_index=False)[value_col_for_color_rank]
            .sum()
            .sort_values(value_col_for_color_rank, ascending=False)["parent"]
            .astype(str)
            .tolist()
        )
        parent_order = _order_with_other_last(parent_order, other_label)
        ordered_keys: List[str] = []
        for p in parent_order:
            sub = g[g["parent"].astype(str) == str(p)].copy()
            sub = sub.sort_values(value_col_for_color_rank, ascending=False)
            ordered_keys.extend(sub["__group__"].astype(str).tolist())
        missing_keys = [k for k in group_keys if k not in ordered_keys]
        group_keys_for_colors = ordered_keys + missing_keys
    else:
        group_keys_for_colors = group_keys

    grp_cfg = cfg.get("grouping") or {}
    max_ch = int(grp_cfg.get("max_items_child", grp_cfg.get("max_children_per_parent", 6)))
    max_par = int(grp_cfg.get("max_items_parent", grp_cfg.get("max_parents", 6)))
    color_map = pick_colors_hardcoded(
        group_keys=group_keys_for_colors,
        group_cols=group_cols,
        max_bubbles=int(cfg.get("max_bubbles", 36)),
        max_parents=max(1, min(6, max_par)),
        max_children_per_parent=max(1, min(6, max_ch)),
    )
    colors = [color_map[k] for k in g["__group__"].tolist()]

    scatter_lw = float(style_cfg.get("scatter_edge_width", 0.8))
    ax.scatter(x, y, s=s, c=colors, alpha=1.0, edgecolors="white", linewidths=scatter_lw)

    leg_cfg = plot_cfg.get("legend", {}) or {}
    legend_entries: List[Dict[str, Any]] = []
    if bool(leg_cfg.get("enabled", True)):
        for pos, (_, row) in enumerate(g.iterrows()):
            k = str(row["__group__"])
            legend_entries.append(
                {
                    "key": k,
                    "parent": row.get("parent", k),
                    "child": row.get("child", k),
                    "label_full": row.get("label", k),
                    "x": float(x.iloc[pos]),
                    "size": float(row.get("bubblesize", 0.0)),
                    "revenue": float(row.get("revenue", 0.0)),
                    "color": color_map[k],
                }
            )

    x_label = str(plot_cfg.get("x_label", "GM"))
    y_unit_lab = str(plot_cfg.get("y_unit_label", "")).strip()
    y_lab_base = str(plot_cfg.get("y_label", "GP"))
    y_label = f"{y_lab_base} {y_unit_lab}".strip() if y_unit_lab else y_lab_base

    ax.set_xlabel(x_label, fontproperties=FONT_PROP)
    ax.set_ylabel(y_label, fontproperties=FONT_PROP)
    ax.xaxis.label.set_fontstyle("italic")
    ax.yaxis.label.set_fontstyle("italic")

    axis_fs = int(plot_cfg.get("axis_label_size", 6))
    tick_fs = int(plot_cfg.get("tick_label_size", 5))
    ax.xaxis.label.set_fontsize(axis_fs)
    ax.yaxis.label.set_fontsize(axis_fs)

    tick_width = float(style_cfg.get("tick_width", 0.8))
    tick_length = float(style_cfg.get("tick_length", 4))
    ax.tick_params(axis="both", labelsize=tick_fs, width=tick_width, length=tick_length)
    for label in ax.get_xticklabels():
        label.set_fontproperties(FONT_PROP)
    for label in ax.get_yticklabels():
        label.set_fontproperties(FONT_PROP)

    x_ref = plot_cfg.get("x_ref", None)
    y_ref = plot_cfg.get("y_ref", None)
    ref_lw = float(style_cfg.get("ref_linewidth", 0.8))
    if x_ref is not None:
        ax.axvline(float(x_ref), linestyle="--", linewidth=ref_lw, color="grey", alpha=0.8)
    if y_ref is not None:
        ax.axhline(float(y_ref), linestyle="--", linewidth=ref_lw, color="grey", alpha=0.8)

    x_lim = plot_cfg.get("x_lim", None)
    y_lim = plot_cfg.get("y_lim", None)
    if x_lim is not None:
        ax.set_xlim(x_lim[0], x_lim[1])
    if y_lim is not None:
        ax.set_ylim(y_lim[0], y_lim[1])

    x_tick_step = plot_cfg.get("x_tick_step", None)
    if x_tick_step is not None:
        ax.xaxis.set_major_locator(ticker.MultipleLocator(float(x_tick_step)))

    if bool(plot_cfg.get("show_labels", False)):
        wrap_w = int(plot_cfg.get("label_wrap", 16))
        max_lines = int(plot_cfg.get("label_max_lines", 2))
        for pos, (_, row) in enumerate(g.iterrows()):
            lab = _wrap_label(str(row["label"]), width=wrap_w, max_lines=max_lines)
            ax.text(
                float(x.iloc[pos]),
                float(y.iloc[pos]),
                lab,
                ha="center",
                va="center",
                fontsize=3,
                color="white",
            )

    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)

    fig.canvas.draw()

    if plot_cfg.get("x_lim", None) is None or plot_cfg.get("y_lim", None) is None:
        r_pts = float(np.sqrt(np.nanmax(s) / 2.0))
        r_px = r_pts * fig.dpi / 72.0
        inv = ax.transData.inverted()
        p0 = inv.transform((0, 0))
        px = inv.transform((r_px, 0))
        py = inv.transform((0, r_px))
        x_pad = abs(px[0] - p0[0]) * 1.20
        y_pad = abs(py[1] - p0[1]) * 1.20
        if plot_cfg.get("x_lim", None) is None:
            ax.set_xlim(float(np.nanmin(x) - x_pad), float(np.nanmax(x) + x_pad))
        if plot_cfg.get("y_lim", None) is None:
            ax.set_ylim(float(np.nanmin(y) - y_pad), float(np.nanmax(y) + y_pad))

    for t in ax.get_xticklabels():
        t.set_fontproperties(FONT_PROP)
        t.set_fontsize(tick_fs)
    for t in ax.get_yticklabels():
        t.set_fontproperties(FONT_PROP)
        t.set_fontsize(tick_fs)
    ax.xaxis.label.set_fontstyle("italic")
    ax.yaxis.label.set_fontstyle("italic")

    return fig, ax, legend_entries


def save_legend_png(
    legend_entries: list,
    out_png: str,
    dpi: int,
    ncol: int,
    font_size: int,
    marker_size: float,
    cfg: Dict[str, Any],
    pad: float = 0.10,
) -> Optional[str]:
    """
    Rendert Legend als eigenes PNG.
    - flat
    - parent_cluster: 1 Parent = 1 Spalte, Box eng um Einträge, Parent-Label oben drüber,
      Spalten sehr eng nebeneinander.
    """
    if not legend_entries:
        return None

    leg_cfg = (cfg.get("plot", {}) or {}).get("legend", {}) or {}
    sort_by = str(leg_cfg.get("sort_by", "x")).lower()
    group_mode = str(leg_cfg.get("group_mode", "flat")).lower()

    box_ls = str(leg_cfg.get("box_linestyle", "--"))
    box_lw = float(leg_cfg.get("box_linewidth", 0.9))

    box_pad_px = float(leg_cfg.get("cluster_box_pad_px", 2.0))
    title_gap_px = float(leg_cfg.get("cluster_title_gap_px", 6.0))
    top_pad_px = float(leg_cfg.get("cluster_top_pad_px", 2.0))
    # Gap between consecutive segment boxes (outer edge → outer edge), must stay > 0.
    cluster_gap_px = float(leg_cfg.get("cluster_gap_px", 3.0))
    if cluster_gap_px < 1.0:
        cluster_gap_px = 1.0

    # Explicit size so parent segment titles match child legend labels
    # (FONT_PROP alone can keep the theme default size).
    legend_fp = matplotlib_font_properties(size=int(font_size))

    def coerce_color(c: Any) -> str:
        if c is None:
            return "#999999"
        if isinstance(c, (tuple, list)):
            try:
                return to_hex(to_rgba(c), keep_alpha=False)
            except Exception:
                return "#999999"
        if isinstance(c, str):
            s = c.strip()
            if s.startswith("#"):
                return s
            if s.startswith("(") and s.endswith(")"):
                try:
                    t = ast.literal_eval(s)
                    if isinstance(t, (tuple, list)):
                        return to_hex(to_rgba(t), keep_alpha=False)
                except Exception:
                    pass
            try:
                return to_hex(to_rgba(s), keep_alpha=False)
            except Exception:
                return "#999999"
        try:
            return to_hex(to_rgba(c), keep_alpha=False)
        except Exception:
            return "#999999"

    def sort_key_row(row: Dict[str, Any]) -> float:
        if sort_by == "size":
            return float(row.get("size", 0.0) or 0.0)
        if sort_by == "revenue":
            return float(row.get("revenue", 0.0) or 0.0)
        return float(row.get("x", 0.0) or 0.0)

    if group_mode != "parent_cluster":
        items = sorted(legend_entries, key=sort_key_row, reverse=True)
        max_items = int(leg_cfg.get("max_items", 20))
        items = items[:max_items]
        handles: List[Line2D] = []
        labels: List[str] = []
        for it in items:
            handles.append(
                Line2D(
                    [0],
                    [0],
                    marker="o",
                    linestyle="",
                    markersize=float(marker_size),
                    markerfacecolor=coerce_color(it.get("color")),
                    markeredgecolor="white",
                    markeredgewidth=0.8,
                )
            )
            labels.append(str(it.get("label_full") or it.get("label") or it.get("child") or ""))

        n = max(1, len(items))
        nrows = max(1, (n + int(ncol) - 1) // int(ncol))
        fig = plt.figure(figsize=(max(3.5, ncol * 1.1), max(0.6, 0.35 * nrows + 0.3)), dpi=dpi)
        ax = fig.add_subplot(111)
        ax.axis("off")
        ax.legend(
            handles=handles,
            labels=labels,
            loc="center",
            ncol=int(ncol),
            frameon=False,
            fontsize=int(font_size),
            prop=legend_fp,
            handletextpad=0.35,
            columnspacing=0.8,
            borderaxespad=0.0,
        )
        fig.savefig(out_png, dpi=dpi, bbox_inches="tight", pad_inches=pad, transparent=True)
        plt.close(fig)
        return out_png

    df = pd.DataFrame(legend_entries)
    if "parent" not in df.columns or "child" not in df.columns:
        return save_legend_png(
            legend_entries=legend_entries,
            out_png=out_png,
            dpi=dpi,
            ncol=ncol,
            font_size=font_size,
            marker_size=marker_size,
            cfg={**cfg, "plot": {**cfg.get("plot", {}), "legend": {**leg_cfg, "group_mode": "flat"}}},
            pad=pad,
        )

    other_label = str((cfg.get("grouping", {}) or {}).get("other_label", "Rest"))
    if "revenue" in df.columns:
        parent_order = (
            df.groupby("parent", as_index=False)["revenue"]
            .sum()
            .sort_values("revenue", ascending=False)["parent"]
            .tolist()
        )
    else:
        parent_order = sorted(df["parent"].astype(str).unique().tolist())
    parent_order = _order_with_other_last([str(p) for p in parent_order], other_label)
    group_cols = cfg.get("group_cols") or []
    flat_single_column = len(group_cols) <= 1

    def _rows_to_handles_labels(sub: pd.DataFrame) -> Tuple[List[Line2D], List[str]]:
        handles_out: List[Line2D] = []
        labels_out: List[str] = []
        for _, r in sub.iterrows():
            handles_out.append(
                Line2D(
                    [0],
                    [0],
                    marker="o",
                    linestyle="",
                    markersize=float(marker_size),
                    markerfacecolor=coerce_color(r.get("color")),
                    markeredgecolor="white",
                    markeredgewidth=0.8,
                )
            )
            labels_out.append(str(r.get("child", r.get("label_full", ""))))
        return handles_out, labels_out

    def _sort_sub(sub: pd.DataFrame) -> pd.DataFrame:
        sub = sub.copy()
        sub["__is_other__"] = sub["child"].astype(str).str.strip().eq(other_label)
        if sort_by == "size":
            sub = sub.sort_values(["__is_other__", "size"], ascending=[True, False])
        elif sort_by == "x":
            sub = sub.sort_values(["__is_other__", "x"], ascending=[True, False])
        elif "revenue" in sub.columns:
            sub = sub.sort_values(["__is_other__", "revenue"], ascending=[True, False])
        else:
            sub = sub.sort_values(["__is_other__", "child"], ascending=[True, True])
        return sub.drop(columns=["__is_other__"])

    blocks: List[Tuple[Any, str, List[Line2D], List[str]]] = []
    if flat_single_column:
        sub = _sort_sub(df)
        handles, labels = _rows_to_handles_labels(sub)
        pcol = coerce_color(sub.iloc[0]["color"]) if len(sub) else "#999999"
        blocks.append(("", pcol, handles, labels))
    else:
        for p in parent_order:
            sub = _sort_sub(df[df["parent"] == p].copy())
            pcol = coerce_color(sub.iloc[0]["color"]) if len(sub) else "#999999"
            handles, labels = _rows_to_handles_labels(sub)
            blocks.append((p, pcol, handles, labels))

    # Use the figure's effective dpi for px↔inches (Retina often doubles fig.dpi).
    probe = plt.figure(dpi=dpi)
    eff_dpi = float(probe.dpi) or float(dpi)
    plt.close(probe)

    def _px_to_in(px: float) -> float:
        return float(px) / eff_dpi

    def _measure_block(handles: List[Line2D], labels: List[str]) -> Tuple[float, float]:
        tmp = plt.figure(figsize=(1.0, 1.0), dpi=dpi)
        ax_m = tmp.add_subplot(111)
        ax_m.axis("off")
        leg_m = ax_m.legend(
            handles=handles,
            labels=labels,
            loc="upper left",
            bbox_to_anchor=(0.0, 1.0),
            frameon=False,
            fontsize=int(font_size),
            prop=legend_fp,
            handletextpad=0.35,
            columnspacing=0.6,
            borderaxespad=0.0,
            labelspacing=0.25,
        )
        tmp.canvas.draw()
        bb_m = leg_m.get_window_extent(renderer=tmp.canvas.get_renderer())
        w_m, h_m = float(bb_m.width), float(bb_m.height)
        plt.close(tmp)
        return w_m, h_m

    # Measure each segment legend; store box size (= content + pad). The next box
    # starts immediately after the previous outer edge + cluster_gap_px.
    legend_sizes: List[Tuple[float, float]] = []
    box_widths: List[float] = []
    box_heights: List[float] = []
    draw_box: List[bool] = []
    show_title: List[bool] = []

    for _p, _pcol, handles, labels in blocks:
        w_raw, h_raw = _measure_block(handles, labels)
        legend_sizes.append((w_raw, h_raw))
        box_widths.append(w_raw + 2 * box_pad_px)
        box_heights.append(h_raw + 2 * box_pad_px)
        n_lbl = len(labels)
        if flat_single_column:
            draw_box.append(True)
            show_title.append(False)
        else:
            draw_box.append(n_lbl > 1)
            show_title.append(n_lbl > 1)

    title_line_px = font_size * (eff_dpi / 72.0) + title_gap_px
    title_h = max((title_line_px if st else 0.0) for st in show_title) if show_title else 0.0
    max_h = max(box_heights) if box_heights else 1.0
    n_blocks = len(blocks)
    side_margin_px = 1.0

    box_lefts: List[float] = []
    x_cursor = side_margin_px
    for i, w_box in enumerate(box_widths):
        box_lefts.append(x_cursor)
        x_cursor += w_box + (cluster_gap_px if i < n_blocks - 1 else 0.0)
    total_w_px = x_cursor + side_margin_px
    total_h_px = top_pad_px + title_h + max_h + 2.0

    fig = plt.figure(figsize=(_px_to_in(total_w_px), _px_to_in(total_h_px)), dpi=dpi)
    axes_list: List[Any] = []
    legends_data: List[Tuple[Any, Any, str]] = []
    y0_px = 1.0
    content_h_px = total_h_px - y0_px

    for i, (p, pcol, handles, labels) in enumerate(blocks):
        w_legend = legend_sizes[i][0]
        # Legend sits inside the box; axes left = box left + pad.
        left_px = box_lefts[i] + box_pad_px
        ax_c = fig.add_axes(
            [
                left_px / total_w_px,
                y0_px / total_h_px,
                max(w_legend, 1.0) / total_w_px,
                content_h_px / total_h_px,
            ]
        )
        ax_c.axis("off")
        leg = ax_c.legend(
            handles=handles,
            labels=labels,
            loc="upper left",
            bbox_to_anchor=(0.0, 1.0),
            frameon=False,
            fontsize=int(font_size),
            prop=legend_fp,
            handletextpad=0.35,
            columnspacing=0.6,
            borderaxespad=0.0,
            labelspacing=0.25,
        )
        axes_list.append(ax_c)
        legends_data.append((leg, p, pcol))

    fig.canvas.draw()
    renderer = fig.canvas.get_renderer()

    # Refine with true rendered sizes, then re-place tightly.
    measured_box_widths: List[float] = []
    measured_legend_sizes: List[Tuple[float, float]] = []
    for leg, _p, _pcol in legends_data:
        bb = leg.get_window_extent(renderer=renderer)
        measured_legend_sizes.append((float(bb.width), float(bb.height)))
        measured_box_widths.append(float(bb.width) + 2 * box_pad_px)

    box_lefts = []
    x_cursor = side_margin_px
    for i, w_box in enumerate(measured_box_widths):
        box_lefts.append(x_cursor)
        x_cursor += w_box + (cluster_gap_px if i < n_blocks - 1 else 0.0)
    packed_w_px = x_cursor + side_margin_px
    if packed_w_px > 1.0:
        fig.set_size_inches(_px_to_in(packed_w_px), _px_to_in(total_h_px), forward=True)
        total_w_px = packed_w_px

    for i, ax_c in enumerate(axes_list):
        w_legend = measured_legend_sizes[i][0]
        left_px = box_lefts[i] + box_pad_px
        ax_c.set_position(
            [
                left_px / total_w_px,
                y0_px / total_h_px,
                max(w_legend, 1.0) / total_w_px,
                content_h_px / total_h_px,
            ]
        )

    fig.canvas.draw()
    renderer = fig.canvas.get_renderer()
    inv_fig = fig.transFigure.inverted()

    # Stored box extents (display px): content bbox + pad. Slots were packed so
    # next_left ≈ prev_right + cluster_gap_px; draw around the live content bbox.
    for i, (ax_c, (leg, p, pcol)) in enumerate(zip(axes_list, legends_data)):
        if not draw_box[i]:
            continue
        bb = leg.get_window_extent(renderer=renderer)
        x0 = float(bb.x0) - box_pad_px
        x1 = float(bb.x1) + box_pad_px
        y0 = float(bb.y0) - box_pad_px
        y1 = float(bb.y1) + box_pad_px

        (fx0, fy0) = inv_fig.transform((x0, y0))
        (fx1, fy1) = inv_fig.transform((x1, y1))

        edge = "none" if flat_single_column else coerce_color(pcol)
        rect = Rectangle(
            (fx0, fy0),
            (fx1 - fx0),
            (fy1 - fy0),
            transform=fig.transFigure,
            fill=False,
            linestyle=box_ls if not flat_single_column else "-",
            linewidth=box_lw if not flat_single_column else 0.0,
            edgecolor=edge,
        )
        rect.set_clip_on(False)
        fig.patches.append(rect)

        if show_title[i] and str(p).strip():
            (tx, ty) = inv_fig.transform((x0, y1 + title_gap_px))
            fig.text(
                tx,
                ty,
                str(p),
                ha="left",
                va="bottom",
                color=coerce_color(pcol),
                fontproperties=legend_fp,
            )

    fig.savefig(out_png, dpi=dpi, bbox_inches="tight", pad_inches=pad, transparent=True)
    plt.close(fig)
    return out_png


def build_output_paths(cfg: Dict[str, Any], period_label: str) -> Tuple[str, str]:
    explicit_xlsx = str(cfg.get("output_path") or "").strip()
    if explicit_xlsx:
        xlsx_path = explicit_xlsx
        out_dir = os.path.dirname(xlsx_path) or str(cfg.get("output_file_path") or ".")
        os.makedirs(out_dir, exist_ok=True)
    else:
        out_dir = str(cfg.get("output_file_path") or cfg.get("output_dir") or os.getcwd())
        os.makedirs(out_dir, exist_ok=True)
        case_id = str(cfg.get("case_id") or "bubble").strip()
        xlsx_path = os.path.join(out_dir, f"{case_id}_Bubble_Output.xlsx")
    prefix = cfg.get("output_prefix", "bubble")
    safe = str(period_label).replace(" ", "_").replace("/", "-")
    base = os.path.join(out_dir, f"{prefix}_{safe}")
    return base + ".png", xlsx_path


def export_excel_with_images(
    cfg: Dict[str, Any],
    plot_png_path: str,
    legend_png_path: Optional[str],
    xlsx_path: str,
    period_label: str,
    *,
    wb=None,
    save: bool = True,
    helper_table: Optional[pd.DataFrame] = None,
    total_net_sales: float = 0.0,
) -> None:
    excel_cfg = cfg.get("excel") or {}
    if excel_cfg.get("enabled") is False:
        return

    col_offset = int(excel_cfg.get("col_offset", 3))
    sheet_name = str(excel_cfg.get("sheet_name", "Margin analyses"))
    anchor_cell = str(excel_cfg.get("image_anchor_cell", "D6"))
    set_widths = bool(excel_cfg.get("set_column_widths", True))
    legend_anchor_cell = str(excel_cfg.get("legend_anchor_cell", "D35"))
    legend_scale = float(excel_cfg.get("legend_scale", 0.35))
    display_dpi = int(excel_cfg.get("display_dpi", 180))
    image_scale = float(excel_cfg.get("image_scale", 1.0))

    if wb is None:
        wb = Workbook()
        ws = wb.active
        ws.title = sheet_name
    else:
        if sheet_name in wb.sheetnames:
            del wb[sheet_name]
        ws = wb.create_sheet(sheet_name)

    font_name = THEME.font_name
    title_font = Font(name=font_name, size=24, color="FF4F2D7F")
    sub_font = Font(name=font_name, size=12, color="FF4F2D7F")
    info_font = Font(name=font_name, size=9, bold=True, color="FF4F2D7F")

    table_base = str(cfg.get("table", ""))
    sub_suffix = str(cfg.get("subtitle_suffix", "")).strip()
    table_name = f"{table_base} — {sub_suffix}".strip(" —") if sub_suffix else table_base

    ws["D1"] = str(cfg.get("title", ""))
    ws["D1"].font = title_font
    ws["D2"] = f"{table_name} | {period_label}"
    ws["D2"].font = sub_font
    ws["D4"] = f"{cfg.get('company', '')} | {table_name}"
    ws["D4"].font = info_font

    fill_grey = PatternFill(fill_type="solid", fgColor="FFF3F1EF")
    for r in range(1, 200):
        for c in range(1, 1 + col_offset):
            ws.cell(row=r, column=c).fill = fill_grey

    fill_white = PatternFill(fill_type="solid", fgColor="FFFFFFFF")
    for r in range(1, 200):
        for c in range(1 + col_offset, 30):
            ws.cell(row=r, column=c).fill = fill_white

    if os.path.exists(plot_png_path):
        img = XLImage(BytesIO(Path(plot_png_path).read_bytes()))
        plot_cfg = cfg.get("plot", {}) or {}
        fig_w, fig_h = plot_cfg.get("figsize", (6.0, 3.2))
        img.width = int(float(fig_w) * display_dpi * image_scale)
        img.height = int(float(fig_h) * display_dpi * image_scale)
        ws.add_image(img, anchor_cell)
        plot_img_h = img.height
    else:
        ws[anchor_cell] = f"(Plot-Bild nicht gefunden: {plot_png_path})"
        plot_img_h = 0

    if legend_png_path and os.path.exists(legend_png_path):
        leg = XLImage(BytesIO(Path(legend_png_path).read_bytes()))
        leg.width = int(leg.width * legend_scale)
        leg.height = int(leg.height * legend_scale)
        # Place legend just under the chart (Excel row ≈ 20px), unless an
        # explicit far-below anchor is configured.
        plot_row = int("".join(ch for ch in anchor_cell if ch.isdigit()) or "6")
        plot_col = "".join(ch for ch in anchor_cell if ch.isalpha()) or "D"
        auto_legend_row = plot_row + max(1, int(np.ceil(float(plot_img_h) / 20.0))) + 2
        configured_row = int(
            "".join(ch for ch in str(legend_anchor_cell) if ch.isdigit()) or "0"
        )
        legend_row = max(auto_legend_row, configured_row) if configured_row else auto_legend_row
        ws.add_image(leg, f"{plot_col}{legend_row}")

    if helper_table is not None:
        labels = sales_basis_labels(cfg.get("sales_basis"))
        table_row = int(excel_cfg.get("helper_table_row", 55))
        table_col = int(excel_cfg.get("helper_table_col", 4))
        headers = [
            "Group",
            f"{labels.metric_label} ({period_label})",
            "Gross profit",
            "GM",
            "Bubble size",
        ]
        for idx, value in enumerate(headers):
            cell = ws.cell(table_row, table_col + idx, value)
            cell.font = Font(name=font_name, size=9, bold=True, color="FFFFFFFF")
            cell.fill = PatternFill(fill_type="solid", fgColor="FF4F2D7F")
        for offset, record in enumerate(helper_table.to_dict("records"), start=1):
            values = [
                record.get("label") or record.get("__group__"),
                float(record.get("revenue") or 0.0) / 1000.0,
                float(record.get("gp") or 0.0) / 1000.0,
                record.get("gm"),
                float(record.get("bubblesize") or 0.0) / 1000.0,
            ]
            for idx, value in enumerate(values):
                ws.cell(table_row + offset, table_col + idx, value)
        total_row = table_row + len(helper_table) + 2
        recon_row = total_row + 1
        reported_row = total_row + 2
        ws.cell(total_row, table_col, labels.total_label)
        ws.cell(total_row, table_col + 1, float(total_net_sales) / 1000.0)
        ws.cell(recon_row, table_col, RECON_DIFFERENCE_LABEL)
        ws.cell(reported_row, table_col, labels.reported_label)
        write_reconciliation_cells(
            wb,
            ws,
            period_columns={period_label: table_col + 1},
            total_row=total_row,
            recon_row=recon_row,
            reported_row=reported_row,
            sales_basis=labels.sales_basis,
        )

    if set_widths:
        for c in range(1, 30):
            ws.column_dimensions[get_column_letter(c)].width = 14
        ws.column_dimensions["A"].width = 3
        ws.column_dimensions["B"].width = 3
        ws.column_dimensions["C"].width = 3
        ws.column_dimensions["D"].width = 20

    ws.sheet_properties.outlinePr.summaryRight = True
    for col_letter in ["A", "B", "C"]:
        ws.column_dimensions[col_letter].outlineLevel = 1
        ws.column_dimensions[col_letter].hidden = True
    ws.column_dimensions["D"].collapsed = True

    if save:
        wb.save(xlsx_path)


def run_bubble_chart(
    df: pd.DataFrame,
    cfg: Dict[str, Any],
    wb=None,
    save: bool = True,
) -> Tuple[None, Optional[str]]:
    cfg = normalize_config(cfg)
    df = apply_filters(df, cfg)
    d = preprocess_input(df, cfg)

    start, end, period_label = get_period(cfg)

    d = d.copy()
    d["revenue_amt"] = compute_amount(d, start, end, cfg, cfg["revenue_col"])
    d["bubble_amt"] = compute_amount(d, start, end, cfg, cfg["bubblesize_col"])
    profit_mode = str(cfg.get("profit_mode", "cost")).lower()
    if profit_mode == "profit":
        d["profit_amt"] = compute_amount(d, start, end, cfg, cfg["profit_col"])
        d["cogs_amt"] = 0.0
    else:
        d["cogs_amt"] = compute_amount(d, start, end, cfg, cfg["cogs_col"])
        d["profit_amt"] = d["revenue_amt"] - d["cogs_amt"]

    g = build_bubble_table(d, cfg)

    fig, _ax, legend_entries = plot_bubble_chart(g, cfg, period_label=period_label)

    png_path, xlsx_path = build_output_paths(cfg, period_label)
    plot_png = png_path.replace(".png", "_plot.png")
    legend_png = png_path.replace(".png", "_legend.png")

    plot_cfg = cfg.get("plot", {}) or {}
    save_dpi = int(plot_cfg.get("save_dpi", plot_cfg.get("dpi", 300)))
    leg_cfg = plot_cfg.get("legend", {}) or {}

    legend_path_out: Optional[str] = None
    if legend_entries:
        legend_path_out = save_legend_png(
            legend_entries=legend_entries,
            out_png=legend_png,
            dpi=save_dpi,
            ncol=int(leg_cfg.get("ncol", 4)),
            font_size=int(leg_cfg.get("font_size", 4)),
            marker_size=float(leg_cfg.get("marker_size", 4.5)),
            cfg=cfg,
            pad=0.10,
        )

    fig.canvas.draw()
    fig.savefig(plot_png, dpi=save_dpi, bbox_inches=None)
    plt.close(fig)

    excel_path_out: Optional[str] = None
    if cfg.get("excel", {}).get("enabled", True):
        export_excel_with_images(
            cfg,
            plot_png,
            legend_path_out,
            xlsx_path,
            period_label,
            wb=wb,
            save=save,
            helper_table=g,
            total_net_sales=float(d["revenue_amt"].sum()),
        )
        excel_path_out = xlsx_path

    if excel_path_out:
        for p in [plot_png, legend_path_out]:
            if p and os.path.exists(p):
                try:
                    os.remove(p)
                except Exception as e:
                    print("Warnung: Konnte temporäres PNG nicht löschen:", e)

    return None, excel_path_out


def main(config_path: str) -> None:
    path = Path(config_path)
    with path.open(encoding="utf-8-sig") as f:
        raw = json.load(f)
    cfg = normalize_config(raw)
    file_path = str(cfg.get("file_path") or "").strip()
    if not file_path:
        raise SystemExit("CONFIG['file_path'] muss gesetzt sein.")
    sheet_name = str(cfg.get("sheet_name") or "Sheet1")
    df = pd.read_excel(file_path, sheet_name=sheet_name, engine="openpyxl")
    _, excel_path = run_bubble_chart(df, cfg)
    print("Finished:", excel_path or "(no excel output)")


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Usage: python bubblescatterplot.py <config.json>", file=sys.stderr)
        sys.exit(2)
    main(sys.argv[1])
