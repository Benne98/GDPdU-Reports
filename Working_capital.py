"""Working capital monthly report — TWC/OWC hierarchy from Master_BS (formula mode)."""
from __future__ import annotations

import json
import sys
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import Alignment, Border, Font, Side

PROJECT_ROOT = Path(__file__).resolve().parent
SCRIPTS_DIR = PROJECT_ROOT / "scripts"
BACKEND_DIR = PROJECT_ROOT / "backend"
if str(SCRIPTS_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPTS_DIR))
if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))

from gst_excel_theme import THEME  # noqa: E402
from databook_excel_layout import (  # noqa: E402
    DIFF_FONT,
    FILL_YELLOW,
    LAYOUT_NA,
    apply_subtotal_row_style,
    check_row_groups_after_table,
    collapse_check_portfolio,
    detect_period_bucket_cols,
    hide_helper_column_group,
    paint_check_source_yellow,
    paint_grey_white_canvas,
    paint_header_band,
    prune_zero_value_rows,
    write_kpi_section_title_row,
    write_mapping_header_row,
)
from databook_periods import (  # noqa: E402
    days_in_month_formula,
    display_bs_period_labels,
    group_month_columns_by_reporting_fy,
    ordered_month_columns_from_df,
    ordered_reporting_columns_from_df,
    yearly_average_period_label,
)
from report_row_layout import (  # noqa: E402
    build_working_capital_row_structure,
    l3_order_from_mapping,
    load_na_l3_orders,
    resolve_wc_kpi_rows,
)

from databook_workbook import (  # noqa: E402
    BS_RECON_MAPPING_FILE,
    CF_NA_L3_ORDER_FILE,
    MASTER_WORKBOOK_STR,
)

DESKTOP_DIR = PROJECT_ROOT / "Desktop"
DEFAULT_INPUT = MASTER_WORKBOOK_STR
DEFAULT_MAPPING_BS = BS_RECON_MAPPING_FILE
DEFAULT_CF_NA_L3_ORDER = CF_NA_L3_ORDER_FILE

SHEET_MASTER_BS = "Master_BS"
SHEET_MASTER_PL = "Master_PL"
SHEET_PL_RECON = "PL_Reconciliation"
SHEET_BS_BUCKET = "BS_Bucket"
SHEET_OUT = "Working_Capital"

SOURCE_COL_CANDIDATES = ("L5", "L6")
REPORTED_FILTER_VALUE = "reported"

_COL_LAYOUT = LAYOUT_NA
MAP_START_COL = _COL_LAYOUT.map_start_col
MAP_END_COL = _COL_LAYOUT.map_end_col
TECH_SPACER_COL = _COL_LAYOUT.spacer_col
POS_COL = _COL_LAYOUT.pos_col

PROJECT_TITLE_ROW = 1
SUBTITLE_ROW = 2
TITLE_ROW = 6
HEADER_ROW = 8
HEADER_ROW_7 = 7
DATA_START_ROW = 9

ROW_HEIGHT = 12
POS_COL_WIDTH = 32
MAP_COL_WIDTH = 12
VALUE_COL_WIDTH = 7.86

FONT_BASE = THEME.font_base
FONT_BASE_BOLD = THEME.font_bold
FONT_HEADER = THEME.font_header
FONT_MAPPING = Font(name=THEME.font_name, size=THEME.font_size, color=THEME.text_header)
FONT_TITLE = Font(
    name=THEME.font_name, size=THEME.font_size, color=THEME.text_brand_title, bold=True
)
FONT_PROJECT = Font(
    name=THEME.font_name, size=THEME.font_size_title, color=THEME.text_brand_title
)
FONT_SUBTITLE = Font(
    name=THEME.font_name, size=THEME.font_size_subtitle, color=THEME.text_brand_title
)
FONT_KPI = Font(name=THEME.font_name, size=THEME.font_size, color=THEME.text_kpi)
FONT_CHECK_RED = DIFF_FONT

FILL_GREY = THEME.fill_tech
FILL_WHITE = THEME.fill_white
FILL_HEADER = THEME.fill_header
FILL_SUBTOTAL = THEME.fill_subtotal
FILL_KPI = THEME.fill_subtotal

ALIGN_LEFT = Alignment(horizontal="left", vertical="center")
ALIGN_RIGHT = Alignment(horizontal="right", vertical="center")

THIN_SIDE = Side(style="thin", color=THEME.border_color)
TOP_BOTTOM_BORDER = Border(top=THIN_SIDE, bottom=THIN_SIDE)
BORDER_SUBTOTAL_TOP = THEME.border_subtotal_top

NUM_FMT_INT = "#,##0;(#,##0);-"
NUM_FMT_KPI = "0.0;(0.0);-"

KPI_SPECS = [
    ("DIO - Days Inventory Outstanding", "dio"),
    ("DSO - Days Sales Outstanding", "dso"),
    ("DPO - Days Payable Outstanding", "dpo"),
    ("CCC - Cash Conversion Cycle", "ccc"),
]

PL_LABEL_FALLBACKS = {
    "cogs": "cost of goods sold",
    "net_sales": "net sales",
}


def col_letter(idx: int) -> str:
    return get_column_letter(idx)


def load_argv_config() -> dict:
    if len(sys.argv) < 2:
        return {}
    path = Path(sys.argv[1])
    if not path.is_file():
        return {}
    return json.loads(path.read_text(encoding="utf-8"))


def path_value(cfg: dict, key: str, default: str) -> str:
    paths = cfg.get("paths") or {}
    val = paths.get(key)
    return str(val) if val else default


def resolve_config() -> dict:
    raw = load_argv_config()
    paths = raw.get("paths") or {}
    return {
        "input_file": path_value(raw, "input_file", path_value(raw, "master_file", DEFAULT_INPUT)),
        "mapping_file": paths.get("mapping_file") or DEFAULT_MAPPING_BS,
        "cf_na_l3_order_file": paths.get("cf_na_l3_order_file") or DEFAULT_CF_NA_L3_ORDER,
        "project_name": str(raw.get("project_name") or "Desktop Test").strip(),
        "company_name": str(raw.get("company_name") or "Group").strip(),
        "fy_end_month": int(raw.get("fy_end_month") or 12),
        "fiscal_start_month": raw.get("fiscal_start_month"),
    }


def detect_source_col(df_: pd.DataFrame) -> str | None:
    candidates = ["Quelle", "Source", "Reported/Adjusted", "Type", "Scenario"]
    sc = next((c for c in candidates if c in df_.columns), None)
    if sc is not None:
        return sc
    for c in df_.columns:
        try:
            s = df_[c].astype(str).str.strip().str.lower()
            if (s == "reported").any():
                return c
        except Exception:
            pass
    return None


def resolve_source_section_column(df_: pd.DataFrame) -> str | None:
    markers = {REPORTED_FILTER_VALUE, "adjusted", "consolidation"}

    def marker_count(col: str) -> int:
        if col not in df_.columns:
            return 0
        series = df_[col].astype(str).str.strip().str.lower()
        return int(series.isin(markers).sum())

    best_col = None
    best_count = 0
    for col in SOURCE_COL_CANDIDATES:
        count = marker_count(col)
        if count > best_count:
            best_col = col
            best_count = count
    if best_col and best_count > 0:
        return best_col
    legacy = detect_source_col(df_)
    if legacy:
        return legacy
    for col in SOURCE_COL_CANDIDATES:
        if col in df_.columns:
            return col
    return None


def detect_na_bucket_col(df_: pd.DataFrame) -> str:
    for cand in ["NA", "L5 - GT NA View", "L5 - GT NA view", "L5 - GT NA", "L5"]:
        if cand in df_.columns:
            return cand
    for c in df_.columns:
        if str(c).strip() == "NA" or str(c).strip().startswith("L5"):
            return c
    raise ValueError("Keine NA-/L5-Bucket-Spalte gefunden.")


def normalize_na_bucket(x) -> str:
    if x is None or (isinstance(x, float) and pd.isna(x)):
        return "Other"
    s = str(x).strip()
    if s in ("TWC", "OWC"):
        return s
    if s.startswith("TWC"):
        return "TWC"
    if s.startswith("OWC"):
        return "OWC"
    return "Other"


def find_row_by_label_contains(ws_, col_idx: int, needle: str) -> int | None:
    needle_l = needle.lower()
    for rr in range(1, ws_.max_row + 1):
        v = ws_.cell(rr, col_idx).value
        if isinstance(v, str) and needle_l in v.lower():
            return rr
    return None


def resolve_pl_l3_label(wb, kind: str) -> str:
    """Resolve PL L3 label from PL_Reconciliation Aggregated block or fallback."""
    fallback_needle = PL_LABEL_FALLBACKS[kind]
    if SHEET_PL_RECON in wb.sheetnames:
        ws_pl = wb[SHEET_PL_RECON]
        for col in range(1, min(ws_pl.max_column or 0, 30) + 1):
            hdr = ws_pl.cell(HEADER_ROW, col).value
            if isinstance(hdr, str) and "aggregated" in hdr.lower():
                row = find_row_by_label_contains(ws_pl, col, fallback_needle)
                if row is not None:
                    val = ws_pl.cell(row, col).value
                    if isinstance(val, str) and val.strip():
                        return val.strip()
        row = find_row_by_label_contains(ws_pl, POS_COL, fallback_needle)
        if row is not None:
            val = ws_pl.cell(row, POS_COL).value
            if isinstance(val, str) and val.strip():
                return val.strip()
    return "Cost of goods sold" if kind == "cogs" else "Net sales"


def subtotal_span_for_l3(struct, idx: int, bucket: str, l3_label: str) -> tuple[int, int]:
    k = idx - 1
    while (
        k >= 0
        and struct[k]["type"] in {"detail", "detail_single"}
        and struct[k].get("NA") == bucket
        and struct[k].get("L3") == l3_label
    ):
        k -= 1
    start = DATA_START_ROW + (k + 1)
    end = DATA_START_ROW + (idx - 1)
    return start, end


def bucket_l4_refs(struct, bucket: str, col_l: str) -> list[str]:
    refs: list[str] = []
    for rr in struct:
        if rr.get("NA") != bucket:
            continue
        if rr["type"] in {"detail", "detail_single"}:
            refs.append(f"{col_l}{rr['_excel_row']}")
    return refs


def _find_row_in_col_any(ws_, col_idx: int, needles: list[str]) -> int | None:
    normalized = [n.strip() for n in needles]
    for rr in range(1, ws_.max_row + 1):
        v = ws_.cell(rr, col_idx).value
        if isinstance(v, str) and v.strip() in normalized:
            return rr
    return None


def _wc_nwc_month_refs(
    nwc_row: int,
    month_cols: list[str],
    period_index: dict[str, int],
    y_cols: list[int],
) -> list[str]:
    refs: list[str] = []
    for mc in month_cols:
        if mc not in period_index:
            continue
        refs.append(f"{col_letter(y_cols[period_index[mc]])}{nwc_row}")
    return refs


def _bs_bucket_bucket_source_formula(
    sheet_name: str,
    bucket_cols: dict[str, int],
    bucket: str,
    assets_row: int,
    el_row: int,
) -> str:
    col = col_letter(bucket_cols[bucket])
    return f"='{sheet_name}'!{col}{assets_row}+'{sheet_name}'!{col}{el_row}"


def _bs_bucket_nwc_source_formula(
    sheet_name: str,
    bucket_cols: dict[str, int],
    assets_row: int,
    el_row: int,
) -> str:
    twc_col = col_letter(bucket_cols["TWC"])
    owc_col = col_letter(bucket_cols["OWC"])
    return (
        f"='{sheet_name}'!{twc_col}{assets_row}"
        f"+'{sheet_name}'!{twc_col}{el_row}"
        f"+'{sheet_name}'!{owc_col}{assets_row}"
        f"+'{sheet_name}'!{owc_col}{el_row}"
    )


def main() -> None:
    cfg = resolve_config()
    input_file = cfg["input_file"]
    mapping_file = cfg["mapping_file"]
    cf_na_l3_order_file = cfg["cf_na_l3_order_file"]
    project_name = cfg["project_name"]
    group_name = cfg["company_name"]
    fy_end_month = cfg["fy_end_month"]

    print("Working_capital —")
    print(f"  input:   {input_file}")
    print(f"  mapping: {mapping_file}")
    print(f"  na_l3:   {cf_na_l3_order_file}")

    wb = load_workbook(input_file)
    if SHEET_MASTER_BS not in wb.sheetnames:
        raise RuntimeError(f"Sheet '{SHEET_MASTER_BS}' nicht gefunden in {input_file}")

    if SHEET_OUT in wb.sheetnames:
        del wb[SHEET_OUT]
    ws = wb.create_sheet(SHEET_OUT)

    df_bs = pd.read_excel(input_file, sheet_name=SHEET_MASTER_BS, engine="openpyxl")
    PERIODS = ordered_month_columns_from_df(df_bs)
    if not PERIODS:
        raise ValueError("Keine Monats-Spalten im Master_BS gefunden.")

    Y_COLS = list(range(POS_COL + 1, POS_COL + 1 + len(PERIODS)))
    source_col = resolve_source_section_column(df_bs)
    na_col = detect_na_bucket_col(df_bs)

    required = {"Entity", "L2", "L3", "L4", na_col} | set(PERIODS)
    missing = required - set(df_bs.columns)
    if missing:
        raise ValueError(f"Master_BS fehlt Spalten: {missing}")

    map_df = pd.read_excel(mapping_file, sheet_name=0, engine="openpyxl")
    if "L3" not in map_df.columns:
        raise ValueError("BS-Mapping muss Spalte 'L3' enthalten.")
    mapping_l3_order = l3_order_from_mapping(map_df)

    na_l3_map_df = pd.read_excel(cf_na_l3_order_file, sheet_name=0, engine="openpyxl")
    na_l3_order = load_na_l3_orders(na_l3_map_df, ["TWC", "OWC"])

    master_bs_ref = SHEET_MASTER_BS
    master_start = 2
    master_end = len(df_bs) + 1

    def master_bs_range(colname: str) -> str:
        loc = df_bs.columns.get_loc(colname) + 1
        letter = get_column_letter(loc)
        return f"{master_bs_ref}!${letter}${master_start}:${letter}${master_end}"

    na_rng = master_bs_range(na_col)
    l3_rng = master_bs_range("L3")
    l4_rng = master_bs_range("L4")
    src_rng = master_bs_range(source_col) if source_col else None
    period_rng = {p: master_bs_range(p) for p in PERIODS}

    pl_l3_rng = None
    pl_src_rng = None
    pl_period_rng: dict[str, str] = {}
    pl_l3_cogs = resolve_pl_l3_label(wb, "cogs")
    pl_l3_ns = resolve_pl_l3_label(wb, "net_sales")

    if SHEET_MASTER_PL in wb.sheetnames:
        df_pl = pd.read_excel(input_file, sheet_name=SHEET_MASTER_PL, engine="openpyxl")
        pl_source = resolve_source_section_column(df_pl)
        if pl_source and "L3" in df_pl.columns:
            pl_master_ref = SHEET_MASTER_PL
            pl_end = len(df_pl) + 1

            def master_pl_range(colname: str) -> str:
                loc = df_pl.columns.get_loc(colname) + 1
                letter = get_column_letter(loc)
                return f"{pl_master_ref}!${letter}${master_start}:${letter}${pl_end}"

            pl_l3_rng = master_pl_range("L3")
            pl_src_rng = master_pl_range(pl_source)
            for p in PERIODS:
                if p in df_pl.columns:
                    pl_period_rng[p] = master_pl_range(p)

    row_structure = build_working_capital_row_structure(
        df_bs,
        mapping_l3_order,
        {"l4_sort_basis": "latest_fy"},
        source_col=source_col or na_col,
        bucket_col=na_col,
        normalize_bucket_fn=normalize_na_bucket,
        na_l3_order=na_l3_order,
    )
    row_structure = prune_zero_value_rows(
        row_structure,
        df_bs,
        PERIODS,
        source_col=source_col or SOURCE_COL_CANDIDATES[0],
        bucket_col=na_col,
    )

    fy_groups = group_month_columns_by_reporting_fy(PERIODS, fy_end_month)
    period_index = {p: i for i, p in enumerate(PERIODS)}
    fy_avg_rows: list[dict] = []
    for fy_end_year, cols in fy_groups.items():
        fy_avg_rows.append(
            {
                "type": "fy_avg",
                "label": f"Yearly average {yearly_average_period_label(fy_end_year, cols, PERIODS, fy_end_month)}",
                "fy_end_year": fy_end_year,
                "month_cols": cols,
            }
        )

    main_row_structure = list(row_structure)
    for i, r in enumerate(main_row_structure):
        r["_idx"] = i
        r["_excel_row"] = DATA_START_ROW + i

    LAST_TABLE_ROW = DATA_START_ROW + len(main_row_structure) - 1
    bucket_total_row: dict[str, int] = {}
    nwc_rownum: int | None = None
    for r in main_row_structure:
        if r["type"] == "total_na":
            bucket_total_row[r.get("bucket", "")] = r["_excel_row"]
        if r["type"] == "total_nwc":
            nwc_rownum = r["_excel_row"]

    kpi_numerator_rows = resolve_wc_kpi_rows(main_row_structure)

    # Header / titles
    pt = ws.cell(PROJECT_TITLE_ROW, POS_COL, f"Project {project_name}")
    pt.font = FONT_PROJECT
    pt.alignment = ALIGN_LEFT
    ws.row_dimensions[PROJECT_TITLE_ROW].height = 36

    st = ws.cell(SUBTITLE_ROW, POS_COL, f"{group_name} Working capital")
    st.font = FONT_SUBTITLE
    st.alignment = ALIGN_LEFT

    tt = ws.cell(TITLE_ROW, POS_COL, f"{group_name} | Working capital {PERIODS[0]} - {PERIODS[-1]}")
    tt.font = FONT_TITLE
    tt.alignment = ALIGN_LEFT

    layout = LAYOUT_NA
    write_mapping_header_row(
        ws,
        layout,
        header_row=HEADER_ROW,
        header_row_7=HEADER_ROW_7,
        labels=["Reported", "", "NA", "L3", "L4"],
    )

    for cc in [POS_COL] + Y_COLS:
        ws.cell(HEADER_ROW_7, cc).fill = FILL_HEADER

    ws.cell(HEADER_ROW, POS_COL, "kEUR").font = FONT_HEADER
    ws.cell(HEADER_ROW, POS_COL).alignment = ALIGN_LEFT
    ws.cell(HEADER_ROW, POS_COL).fill = FILL_HEADER

    for idx, period in enumerate(PERIODS):
        cc = Y_COLS[idx]
        h = ws.cell(HEADER_ROW, cc, period)
        h.font = FONT_HEADER
        h.alignment = ALIGN_RIGHT
        h.fill = FILL_HEADER

    for c in range(MAP_START_COL, MAP_END_COL + 1):
        ws.column_dimensions[col_letter(c)].width = MAP_COL_WIDTH
    ws.column_dimensions[col_letter(POS_COL)].width = POS_COL_WIDTH
    for c in Y_COLS:
        ws.column_dimensions[col_letter(c)].width = VALUE_COL_WIDTH

    map_rep_col = col_letter(MAP_START_COL)
    map_na_col = col_letter(MAP_START_COL + 2)
    map_l3_col = col_letter(MAP_START_COL + 3)
    map_l4_col = col_letter(MAP_START_COL + 4)

    ws.sheet_properties.outlinePr.summaryBelow = True
    ws.sheet_properties.outlinePr.summaryRight = True

    # WC table data rows (main table only — fy averages come after KPI block)
    for i, r in enumerate(main_row_structure):
        excel_row = r["_excel_row"]
        ws.row_dimensions[excel_row].height = ROW_HEIGHT
        rt = r["type"]

        ws.cell(excel_row, MAP_START_COL, "Reported").font = FONT_MAPPING
        ws.cell(excel_row, MAP_START_COL + 1, "").font = FONT_MAPPING
        ws.cell(excel_row, MAP_START_COL + 2, r.get("NA", "")).font = FONT_MAPPING
        ws.cell(excel_row, MAP_START_COL + 3, r.get("L3", "")).font = FONT_MAPPING
        ws.cell(excel_row, MAP_START_COL + 4, r.get("L4", "")).font = FONT_MAPPING
        for c in range(MAP_START_COL, MAP_END_COL + 1):
            ws.cell(excel_row, c).alignment = ALIGN_LEFT

        pc = ws.cell(excel_row, POS_COL, r["label"])
        pc.alignment = ALIGN_LEFT

        is_bold = rt in {"total_na", "total_nwc"}
        pc.font = FONT_BASE_BOLD if is_bold else FONT_BASE

        if rt == "detail":
            pc.alignment = Alignment(horizontal="left", vertical="center", indent=1)
            ws.row_dimensions[excel_row].outlineLevel = 1
            ws.row_dimensions[excel_row].collapsed = True
        elif rt in {"detail_single", "subtotal_l3"}:
            ws.row_dimensions[excel_row].outlineLevel = 0
        elif rt in {"total_na", "total_nwc"}:
            ws.row_dimensions[excel_row].outlineLevel = 0

        rep_crit = f"${map_rep_col}${excel_row}"
        na_crit = f"${map_na_col}${excel_row}"
        l3_crit = f"${map_l3_col}${excel_row}"
        l4_crit = f"${map_l4_col}${excel_row}"

        for p_idx, period in enumerate(PERIODS):
            cell = ws.cell(excel_row, Y_COLS[p_idx])
            cell.alignment = ALIGN_RIGHT
            cell.number_format = NUM_FMT_INT
            cell.font = FONT_BASE_BOLD if is_bold else FONT_BASE
            col_l = col_letter(Y_COLS[p_idx])

            if rt in {"detail", "detail_single"}:
                sum_rng = period_rng[period]
                if src_rng is not None:
                    cell.value = (
                        f"=SUMIFS({sum_rng},"
                        f"{src_rng},{rep_crit},"
                        f"{na_rng},{na_crit},"
                        f"{l3_rng},{l3_crit},"
                        f"{l4_rng},{l4_crit}"
                        f")/1000"
                    )
                else:
                    cell.value = (
                        f"=SUMIFS({sum_rng},"
                        f"{na_rng},{na_crit},"
                        f"{l3_rng},{l3_crit},"
                        f"{l4_rng},{l4_crit}"
                        f")/1000"
                    )
            elif rt == "subtotal_l3":
                start, end = subtotal_span_for_l3(main_row_structure, i, r["NA"], r["label"])
                cell.value = f"=SUM({col_l}{start}:{col_l}{end})" if end >= start else 0
            elif rt == "total_na":
                refs = bucket_l4_refs(main_row_structure, r["bucket"], col_l)
                cell.value = f"=SUM({','.join(refs)})" if refs else 0
            elif rt == "total_nwc":
                twc = bucket_total_row.get("TWC")
                owc = bucket_total_row.get("OWC")
                parts = []
                if twc:
                    parts.append(f"{col_l}{twc}")
                if owc:
                    parts.append(f"{col_l}{owc}")
                cell.value = f"=SUM({','.join(parts)})" if parts else 0

        if rt == "total_nwc":
            apply_subtotal_row_style(
                ws,
                excel_row,
                layout=layout,
                pos_col=POS_COL,
                value_cols=Y_COLS,
                fill=FILL_WHITE,
                pos_border=TOP_BOTTOM_BORDER,
                value_border=TOP_BOTTOM_BORDER,
            )
        else:
            for cc in Y_COLS:
                ws.cell(excel_row, cc).fill = FILL_WHITE

    # KPI block (after Net working capital)
    KPI_TITLE_ROW = LAST_TABLE_ROW + 1
    write_kpi_section_title_row(
        ws,
        KPI_TITLE_ROW,
        title_cols=Y_COLS,
        label_col=POS_COL,
        title_text="KPIs",
        fill_end_col=Y_COLS[-1],
    )
    KPI_START_ROW = LAST_TABLE_ROW + 2
    kpi_rows_by_key: dict[str, int] = {}
    KPI_ROWS: set[int] = set()

    for i_kpi, (label, key) in enumerate(KPI_SPECS):
        rr = KPI_START_ROW + i_kpi
        KPI_ROWS.add(rr)
        kpi_rows_by_key[key] = rr
        ws.row_dimensions[rr].height = ROW_HEIGHT
        ws.cell(rr, POS_COL, label).alignment = ALIGN_LEFT
        ws.cell(rr, POS_COL, label).font = FONT_KPI
        for cc in Y_COLS:
            ws.cell(rr, cc).number_format = NUM_FMT_KPI
            ws.cell(rr, cc).alignment = ALIGN_RIGHT
            ws.cell(rr, cc).font = FONT_KPI

    # Yearly averages after KPI rows (KPI band formatting)
    FY_AVG_START_ROW = KPI_START_ROW + len(KPI_SPECS)
    FY_AVG_ROWS: set[int] = set()
    for i_fy, r in enumerate(fy_avg_rows):
        excel_row = FY_AVG_START_ROW + i_fy
        FY_AVG_ROWS.add(excel_row)
        r["_excel_row"] = excel_row
        ws.row_dimensions[excel_row].height = ROW_HEIGHT
        ws.row_dimensions[excel_row].outlineLevel = 0

        ws.cell(excel_row, MAP_START_COL, "Reported").font = FONT_MAPPING
        ws.cell(excel_row, MAP_START_COL + 1, "").font = FONT_MAPPING
        ws.cell(excel_row, MAP_START_COL + 2, r.get("NA", "")).font = FONT_MAPPING
        ws.cell(excel_row, MAP_START_COL + 3, r.get("L3", "")).font = FONT_MAPPING
        ws.cell(excel_row, MAP_START_COL + 4, r.get("L4", "")).font = FONT_MAPPING
        for c in range(MAP_START_COL, MAP_END_COL + 1):
            ws.cell(excel_row, c).alignment = ALIGN_LEFT

        pc = ws.cell(excel_row, POS_COL, r["label"])
        pc.alignment = ALIGN_LEFT
        pc.font = FONT_KPI

        for p_idx, period in enumerate(PERIODS):
            cell = ws.cell(excel_row, Y_COLS[p_idx])
            cell.alignment = ALIGN_RIGHT
            cell.number_format = NUM_FMT_KPI
            cell.font = FONT_KPI
            month_cols_fy = r.get("month_cols", [])
            if period not in month_cols_fy or not nwc_rownum:
                cell.value = ""
                continue
            refs = [
                f"{col_letter(Y_COLS[period_index[mc]])}{nwc_rownum}"
                for mc in month_cols_fy
                if mc in period_index
            ]
            cell.value = f"=AVERAGE({','.join(refs)})" if refs else ""

        ws.cell(excel_row, POS_COL).fill = FILL_KPI
        for cc in Y_COLS:
            ws.cell(excel_row, cc).fill = FILL_KPI

    GAP_ROW = FY_AVG_START_ROW + len(fy_avg_rows)
    COGS_HELPER_ROW = GAP_ROW + 2
    NS_HELPER_ROW = COGS_HELPER_ROW + 1

    for rr in (COGS_HELPER_ROW, NS_HELPER_ROW):
        ws.row_dimensions[rr].outlineLevel = 2
        ws.row_dimensions[rr].hidden = True

    for rr, label in [(COGS_HELPER_ROW, pl_l3_cogs), (NS_HELPER_ROW, pl_l3_ns)]:
        ws.cell(rr, MAP_START_COL, "Reported").font = FONT_MAPPING
        ws.cell(rr, MAP_START_COL + 1, "").font = FONT_MAPPING
        ws.cell(rr, MAP_START_COL + 3, label).font = FONT_MAPPING
        ws.cell(rr, POS_COL, label).font = FONT_KPI
        ws.cell(rr, POS_COL).alignment = ALIGN_LEFT
        for cc in range(MAP_START_COL, Y_COLS[-1] + 1):
            ws.cell(rr, cc).fill = FILL_WHITE
        for cc in Y_COLS:
            ws.cell(rr, cc).number_format = NUM_FMT_INT
            ws.cell(rr, cc).alignment = ALIGN_RIGHT
            ws.cell(rr, cc).font = FONT_KPI

    rep_cogs = f"${map_rep_col}${COGS_HELPER_ROW}"
    rep_ns = f"${map_rep_col}${NS_HELPER_ROW}"
    l3_cogs = f"${map_l3_col}${COGS_HELPER_ROW}"
    l3_ns = f"${map_l3_col}${NS_HELPER_ROW}"

    for p_idx, period in enumerate(PERIODS):
        cc = Y_COLS[p_idx]
        if pl_period_rng.get(period) and pl_l3_rng and pl_src_rng:
            ws.cell(COGS_HELPER_ROW, cc).value = (
                f"=SUMIFS({pl_period_rng[period]},{pl_src_rng},{rep_cogs},"
                f"{pl_l3_rng},{l3_cogs})/1000"
            )
            ws.cell(NS_HELPER_ROW, cc).value = (
                f"=SUMIFS({pl_period_rng[period]},{pl_src_rng},{rep_ns},"
                f"{pl_l3_rng},{l3_ns})/1000"
            )
        else:
            ws.cell(COGS_HELPER_ROW, cc).value = 0
            ws.cell(NS_HELPER_ROW, cc).value = 0

    row_inventory = kpi_numerator_rows.get("dio")
    row_recv = kpi_numerator_rows.get("dso")
    row_pay = kpi_numerator_rows.get("dpo")

    for p_idx, period in enumerate(PERIODS):
        cc = Y_COLS[p_idx]
        col_l = col_letter(cc)
        days = days_in_month_formula(period, fy_end_month)
        cogs_cell = f"{col_l}{COGS_HELPER_ROW}"
        ns_cell = f"{col_l}{NS_HELPER_ROW}"
        inv_cell = f"{col_l}{row_inventory}" if row_inventory else None
        recv_cell = f"{col_l}{row_recv}" if row_recv else None
        pay_cell = f"{col_l}{row_pay}" if row_pay else None

        if inv_cell:
            ws.cell(kpi_rows_by_key["dio"], cc).value = (
                f'=IFERROR(-{inv_cell}/{cogs_cell}*{days},"n/a")'
            )
        if recv_cell:
            ws.cell(kpi_rows_by_key["dso"], cc).value = (
                f'=IFERROR({recv_cell}/{ns_cell}*{days},"n/a")'
            )
        if pay_cell:
            ws.cell(kpi_rows_by_key["dpo"], cc).value = (
                f'=IFERROR({pay_cell}/{cogs_cell}*{days},"n/a")'
            )

        dio_r = kpi_rows_by_key.get("dio")
        dso_r = kpi_rows_by_key.get("dso")
        dpo_r = kpi_rows_by_key.get("dpo")
        if dio_r and dso_r and dpo_r:
            ws.cell(kpi_rows_by_key["ccc"], cc).value = (
                f'=IFERROR({col_l}{dso_r}+{col_l}{dio_r}-{col_l}{dpo_r},"n/a")'
            )

    check_base_row = NS_HELPER_ROW + 2
    (
        bs_bucket_src_row,
        bs_bucket_sum_row,
        bs_bucket_check_row,
        master_src_row,
        master_sum_row,
        master_check_row,
    ) = check_row_groups_after_table(check_base_row, [3, 3])
    check_rows_to_collapse = [
        bs_bucket_src_row,
        bs_bucket_sum_row,
        bs_bucket_check_row,
        master_src_row,
        master_sum_row,
        master_check_row,
    ]
    master_periods = ordered_reporting_columns_from_df(df_bs)
    display_periods = display_bs_period_labels(master_periods, fy_end_month)
    master_to_display = {
        master_periods[i]: display_periods[i] for i in range(len(master_periods))
    }

    from openpyxl.formatting.rule import CellIsRule

    def _style_check_row(cell, *, red: bool = False) -> None:
        cell.number_format = NUM_FMT_INT
        cell.alignment = ALIGN_RIGHT
        cell.font = FONT_CHECK_RED if red else FONT_BASE

    def _style_check_delta(cell) -> None:
        _style_check_row(cell, red=True)
        ws.conditional_formatting.add(
            cell.coordinate,
            CellIsRule(operator="notEqual", formula=["0"], font=FONT_CHECK_RED),
        )

    if SHEET_BS_BUCKET in wb.sheetnames and nwc_rownum is not None:
        ws_bs = wb[SHEET_BS_BUCKET]
        period_bucket_cols = detect_period_bucket_cols(ws_bs, PERIODS)
        bs_assets_row = _find_row_in_col_any(ws_bs, POS_COL, ["Total assets"])
        bs_el_row = _find_row_in_col_any(
            ws_bs, POS_COL, ["Total equity & liabilities", "Total equity & liabilities "]
        )
        owc_total_row = bucket_total_row.get("OWC")
        twc_total_row = bucket_total_row.get("TWC")
        if (
            not period_bucket_cols
            or bs_assets_row is None
            or bs_el_row is None
            or owc_total_row is None
            or twc_total_row is None
        ):
            print("Warnung: BS_Bucket Bucket-Checks übersprungen (Perioden/Total-Zeilen).")
        else:
            ws.cell(bs_bucket_src_row, POS_COL, "Source - BS_Bucket").font = FONT_BASE
            ws.cell(bs_bucket_src_row, POS_COL).alignment = ALIGN_LEFT
            ws.cell(bs_bucket_sum_row, POS_COL, "Sum WC").font = FONT_BASE
            ws.cell(bs_bucket_sum_row, POS_COL).alignment = ALIGN_LEFT
            ws.cell(bs_bucket_check_row, POS_COL, "Check").font = FONT_CHECK_RED
            ws.cell(bs_bucket_check_row, POS_COL).alignment = ALIGN_LEFT
            bs_yellow_cols: list[int] = []
            for _fy_end_year, month_cols in fy_groups.items():
                if len(month_cols) < 2:
                    continue
                penultimate_period = month_cols[-2]
                last_period = month_cols[-1]
                if penultimate_period not in period_index or last_period not in period_index:
                    continue
                pen_display = master_to_display.get(penultimate_period, penultimate_period)
                last_display = master_to_display.get(last_period, last_period)
                bucket_map_pen = period_bucket_cols.get(pen_display)
                bucket_map_last = period_bucket_cols.get(last_display)
                if bucket_map_pen is None or bucket_map_last is None:
                    continue
                owc_cc = Y_COLS[period_index[penultimate_period]]
                twc_cc = Y_COLS[period_index[last_period]]
                for bucket, cc, bucket_map, total_row in (
                    ("OWC", owc_cc, bucket_map_pen, owc_total_row),
                    ("TWC", twc_cc, bucket_map_last, twc_total_row),
                ):
                    src_cell = ws.cell(bs_bucket_src_row, cc)
                    src_cell.value = _bs_bucket_bucket_source_formula(
                        SHEET_BS_BUCKET, bucket_map, bucket, bs_assets_row, bs_el_row
                    )
                    src_cell.fill = FILL_YELLOW
                    _style_check_row(src_cell)
                    bs_yellow_cols.append(cc)
                    sum_refs = [
                        f"{col_letter(Y_COLS[period_index[mc]])}{total_row}"
                        for mc in month_cols
                        if mc in period_index
                    ]
                    sum_cell = ws.cell(bs_bucket_sum_row, cc)
                    sum_cell.value = f"=SUM({','.join(sum_refs)})" if sum_refs else ""
                    _style_check_row(sum_cell)
                    delta_cell = ws.cell(bs_bucket_check_row, cc)
                    delta_cell.value = f"={sum_cell.coordinate}-{src_cell.coordinate}"
                    _style_check_delta(delta_cell)
            if bs_yellow_cols:
                paint_check_source_yellow(ws, bs_bucket_src_row, [POS_COL, *bs_yellow_cols])
    else:
        print(f"Hinweis: Sheet '{SHEET_BS_BUCKET}' fehlt oder NWC-Zeile nicht gefunden — BS_Bucket-Check übersprungen.")

    if nwc_rownum is not None:
        ws.cell(master_src_row, POS_COL, "Source - Master_BS").font = FONT_BASE
        ws.cell(master_src_row, POS_COL).alignment = ALIGN_LEFT
        ws.cell(master_sum_row, POS_COL, "Sum WC").font = FONT_BASE
        ws.cell(master_sum_row, POS_COL).alignment = ALIGN_LEFT
        ws.cell(master_check_row, POS_COL, "Check").font = FONT_CHECK_RED
        ws.cell(master_check_row, POS_COL).alignment = ALIGN_LEFT
        master_yellow_cols: list[int] = []
        rep_crit = f'"{REPORTED_FILTER_VALUE}"' if src_rng else None
        for _fy_end_year, month_cols in fy_groups.items():
            if not month_cols:
                continue
            fy_master_label = yearly_average_period_label(
                _fy_end_year, month_cols, PERIODS, fy_end_month
            )
            target_period = month_cols[-1]
            if target_period not in period_index:
                continue
            target_cc = Y_COLS[period_index[target_period]]
            nwc_refs = _wc_nwc_month_refs(nwc_rownum, month_cols, period_index, Y_COLS)
            if not nwc_refs:
                continue
            if fy_master_label in df_bs.columns:
                sum_rng = master_bs_range(fy_master_label)
                if rep_crit and src_rng:
                    src_formula = (
                        f"=(SUMIFS({sum_rng},{na_rng},\"TWC\",{src_rng},{rep_crit})"
                        f"+SUMIFS({sum_rng},{na_rng},\"OWC\",{src_rng},{rep_crit}))/1000"
                    )
                else:
                    src_formula = (
                        f"=(SUMIFS({sum_rng},{na_rng},\"TWC\")"
                        f"+SUMIFS({sum_rng},{na_rng},\"OWC\"))/1000"
                    )
            else:
                month_parts: list[str] = []
                for mc in month_cols:
                    if mc not in period_rng:
                        continue
                    pr = period_rng[mc]
                    if rep_crit and src_rng:
                        month_parts.append(
                            f"SUMIFS({pr},{na_rng},\"TWC\",{src_rng},{rep_crit})"
                            f"+SUMIFS({pr},{na_rng},\"OWC\",{src_rng},{rep_crit})"
                        )
                    else:
                        month_parts.append(
                            f"SUMIFS({pr},{na_rng},\"TWC\")+SUMIFS({pr},{na_rng},\"OWC\")"
                        )
                src_formula = (
                    f"=({'+'.join(month_parts)})/1000" if month_parts else ""
                )
            if not src_formula:
                continue
            src_cell = ws.cell(master_src_row, target_cc)
            src_cell.value = src_formula
            src_cell.fill = FILL_YELLOW
            _style_check_row(src_cell)
            master_yellow_cols.append(target_cc)
            sum_cell = ws.cell(master_sum_row, target_cc)
            sum_cell.value = f"=SUM({','.join(nwc_refs)})"
            _style_check_row(sum_cell)
            delta_cell = ws.cell(master_check_row, target_cc)
            delta_cell.value = f"={sum_cell.coordinate}-{src_cell.coordinate}"
            _style_check_delta(delta_cell)
        if master_yellow_cols:
            paint_check_source_yellow(ws, master_src_row, [POS_COL, *master_yellow_cols])

    if check_rows_to_collapse:
        collapse_check_portfolio(ws, check_rows_to_collapse)

    current_row = max(
        NS_HELPER_ROW,
        master_check_row if nwc_rownum else NS_HELPER_ROW,
        bs_bucket_check_row if nwc_rownum else NS_HELPER_ROW,
    )

    ws.sheet_view.showOutlineSymbols = True

    FILL_END_ROW = max(LAST_TABLE_ROW, current_row, FY_AVG_START_ROW + len(fy_avg_rows)) + 100
    FILL_END_COL = Y_COLS[-1] + 40

    for rr in range(1, FILL_END_ROW + 1):
        if rr != PROJECT_TITLE_ROW:
            ws.row_dimensions[rr].height = ROW_HEIGHT

    paint_grey_white_canvas(ws, layout, last_row=FILL_END_ROW, last_col=Y_COLS[-1])

    for rr in range(1, FILL_END_ROW + 1):
        for cc in range(POS_COL, FILL_END_COL + 1):
            if rr in (KPI_ROWS | FY_AVG_ROWS | {KPI_TITLE_ROW}) and cc in ([POS_COL] + Y_COLS):
                continue
            if rr in (HEADER_ROW_7, HEADER_ROW) and cc in ([POS_COL] + Y_COLS):
                continue
            if ws.cell(rr, cc).fill.fill_type is None:
                ws.cell(rr, cc).fill = FILL_WHITE

    paint_header_band(
        ws,
        layout,
        header_rows=[HEADER_ROW_7, HEADER_ROW],
        period_cols=Y_COLS,
    )

    for rr in KPI_ROWS | FY_AVG_ROWS | {KPI_TITLE_ROW}:
        ws.cell(rr, POS_COL).fill = FILL_KPI
        for cc in Y_COLS:
            ws.cell(rr, cc).fill = FILL_KPI

    for r in main_row_structure:
        if r["type"] != "total_nwc":
            continue
        excel_row = r["_excel_row"]
        apply_subtotal_row_style(
            ws,
            excel_row,
            pos_col=POS_COL,
            value_cols=Y_COLS,
            fill=FILL_WHITE,
            pos_border=TOP_BOTTOM_BORDER,
            value_border=TOP_BOTTOM_BORDER,
        )

    hide_helper_column_group(ws, layout)

    wb.calculation.fullCalcOnLoad = True
    wb.save(input_file)
    print(f"Saved sheet '{SHEET_OUT}' to {input_file}")
    print(f"Monatsspalten: {len(PERIODS)} ({PERIODS[0]} … {PERIODS[-1]})")


if __name__ == "__main__":
    main()
