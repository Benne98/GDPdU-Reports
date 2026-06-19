"""
Apply sign multiplication and PL net-result totals on PDF extraction workbook.
JSON config: input_file, output_file
"""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

SHEETS = ["BS", "PL"]
START_ANCHORS = ["Englisch Mapping", "English Mapping"]
END_ANCHORS = ["Level"]
SIGN_HEADER = "Sign"
TOTAL_HEADER = "Total"
FY_TARGETS = ["FY23A", "FY24A", "FY25A"]


def find_header_row_and_cols(ws, max_scan_rows=20):
    for r in range(1, max_scan_rows + 1):
        row_values = [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
        row_values_str = [str(v).strip() if v is not None else "" for v in row_values]
        has_start = any(a in row_values_str for a in START_ANCHORS)
        has_end = any(a in row_values_str for a in END_ANCHORS)
        if has_start and has_end:
            header_to_col = {}
            for c, v in enumerate(row_values_str, start=1):
                if v != "":
                    header_to_col[v] = c
            return r, header_to_col
    raise ValueError(f"Could not find header row (anchors {START_ANCHORS}, {END_ANCHORS}).")


def parse_number(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return v
    s = str(v).strip()
    if s == "" or s.lower() in {"nan", "none"}:
        return None
    neg = False
    if s.startswith("(") and s.endswith(")"):
        neg = True
        s = s[1:-1].strip()
    s = s.replace(" ", "")
    if "." in s and "," in s:
        if s.rfind(",") > s.rfind("."):
            s = s.replace(".", "").replace(",", ".")
        else:
            s = s.replace(",", "")
    else:
        s = s.replace(",", "")
    s = s.replace("'", "")
    try:
        num = float(s)
        if neg:
            num = -num
        if abs(num - int(num)) < 1e-9:
            return int(num)
        return num
    except Exception:
        return None


def parse_sign(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return None if v == 0 else float(v)
    s = str(v).strip()
    if s == "":
        return None
    if s.startswith("(") and s.endswith(")"):
        inner = s[1:-1].strip()
        inner_num = parse_number(inner)
        if inner_num is None:
            return None
        return -abs(float(inner_num))
    num = parse_number(s)
    if num is None or num == 0:
        return None
    return float(num)


def find_last_data_row(ws, header_row):
    for r in range(ws.max_row, header_row, -1):
        for c in range(1, ws.max_column + 1):
            if ws.cell(row=r, column=c).value not in (None, ""):
                return r
    return header_row


def get_anchor_col(header_to_col, anchors):
    for a in anchors:
        if a in header_to_col:
            return header_to_col[a]
    return None


def run(config: dict) -> None:
    input_file = Path(str(config.get("input_file") or ""))
    output_file = Path(str(config.get("output_file") or input_file))
    if not input_file.is_file():
        raise SystemExit(f"input_file not found: {input_file}")

    wb = load_workbook(input_file)

    for sheet in SHEETS:
        if sheet not in wb.sheetnames:
            raise ValueError(f"Sheet '{sheet}' not found. Available: {wb.sheetnames}")

        ws = wb[sheet]
        header_row, header_to_col = find_header_row_and_cols(ws)
        start_col = get_anchor_col(header_to_col, START_ANCHORS)
        end_col = get_anchor_col(header_to_col, END_ANCHORS)
        sign_col = header_to_col.get(SIGN_HEADER)

        if start_col is None or end_col is None or sign_col is None:
            raise ValueError(
                f"Anchors/Sign not found in sheet '{sheet}'. "
                f"start={start_col}, end={end_col}, sign={sign_col}"
            )

        value_cols = list(range(start_col + 1, end_col))

        for r in range(header_row + 1, ws.max_row + 1):
            sgn = parse_sign(ws.cell(row=r, column=sign_col).value)
            if sgn is None:
                continue
            for c in value_cols:
                cell = ws.cell(row=r, column=c)
                num = parse_number(cell.value)
                if num is None:
                    continue
                new_val = num * sgn
                if isinstance(new_val, float) and abs(new_val - int(new_val)) < 1e-9:
                    new_val = int(new_val)
                cell.value = new_val

        if sheet == "PL":
            total_col = header_to_col.get(TOTAL_HEADER)
            if total_col is None:
                raise ValueError(f"Column '{TOTAL_HEADER}' not found in sheet 'PL'.")
            last_row = find_last_data_row(ws, header_row)
            fy_cols = []
            for fy in FY_TARGETS:
                if fy in header_to_col:
                    fy_cols.append(header_to_col[fy])
            if not fy_cols:
                fy_pattern = re.compile(r"^FY\d{2}A$", re.IGNORECASE)

                def fy_key(h):
                    m = re.search(r"FY(\d{2})A", h, re.IGNORECASE)
                    return int(m.group(1)) if m else 0

                fy_headers = sorted(
                    [h for h in header_to_col if fy_pattern.match(h)],
                    key=fy_key,
                )
                fy_cols = [header_to_col[h] for h in fy_headers[-3:]]

            for c in fy_cols:
                total_sum = 0
                for r in range(header_row + 1, last_row):
                    total_flag = ws.cell(row=r, column=total_col).value
                    total_flag_str = str(total_flag).strip().lower() if total_flag is not None else ""
                    if total_flag_str == "yes":
                        continue
                    val = parse_number(ws.cell(row=r, column=c).value)
                    if val is None:
                        continue
                    total_sum += val
                ws.cell(row=last_row, column=c).value = total_sum

    wb.save(output_file)
    print(f"Adjustments saved to {output_file}")


def main() -> None:
    if len(sys.argv) < 2:
        raise SystemExit("Usage: python PDF_Extraction_adjustments.py <config.json>")
    config = json.loads(Path(sys.argv[1]).read_text(encoding="utf-8-sig"))
    run(config)


if __name__ == "__main__":
    main()
