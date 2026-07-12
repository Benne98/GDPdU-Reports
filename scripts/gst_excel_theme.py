"""
Excel theme for General Sales Table — aligned with Financials → Income Statement (frontend).

Reference: frontend/src/components/financials/FinancialStatementTable.tsx
"""

from __future__ import annotations

from dataclasses import dataclass
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


@dataclass(frozen=True)
class GstExcelTheme:
    """ARGB hex without # prefix (openpyxl fgColor)."""

    font_name: str = "Calibri"
    font_size: int = 9
    font_size_title: int = 24
    font_size_subtitle: int = 12
    font_size_company: int = 9

    # Surfaces
    header_bg: str = "FFF8FAFC"
    subtotal_bg: str = "FFF8FAFC"
    gm_column_bg: str = "FFF8FAFC"
    tech_col_bg: str = "FFF1F5F9"
    white_bg: str = "FFFFFFFF"
    period_column_bg: str = "FFF1F5F9"

    zero_row_bg: str = "FFF1F5F9"

    # Text
    text_primary: str = "FF111827"
    text_header: str = "FF475569"
    text_kpi: str = "FF64748B"
    text_account: str = "FF475569"
    text_cagr: str = "FF64748B"
    text_brand_title: str = "FF1E3A5F"

    # Deltas (Income Statement)
    delta_positive: str = "FF10B981"
    delta_negative: str = "FFDC2626"
    delta_zero: str = "FF94A3B8"

    # Borders
    border_color: str = "FFE2E8F0"
    border_strong: str = "FFE2E8F0"

    @property
    def font_base(self) -> Font:
        return Font(name=self.font_name, size=self.font_size, color=self.text_primary)

    @property
    def font_bold(self) -> Font:
        return Font(name=self.font_name, size=self.font_size, bold=True, color=self.text_primary)

    @property
    def font_header(self) -> Font:
        return Font(name=self.font_name, size=self.font_size, bold=True, color=self.text_header)

    @property
    def fill_header(self) -> PatternFill:
        return PatternFill(fill_type="solid", fgColor=self.header_bg)

    @property
    def fill_subtotal(self) -> PatternFill:
        return PatternFill(fill_type="solid", fgColor=self.subtotal_bg)

    @property
    def fill_gm(self) -> PatternFill:
        return PatternFill(fill_type="solid", fgColor=self.gm_column_bg)

    @property
    def fill_white(self) -> PatternFill:
        return PatternFill(fill_type="solid", fgColor=self.white_bg)

    @property
    def fill_period(self) -> PatternFill:
        return PatternFill(fill_type="solid", fgColor=self.period_column_bg)

    @property
    def fill_tech(self) -> PatternFill:
        return PatternFill(fill_type="solid", fgColor=self.tech_col_bg)

    @property
    def fill_zero_row(self) -> PatternFill:
        return PatternFill(fill_type="solid", fgColor=self.zero_row_bg)

    @property
    def font_zero_row(self) -> Font:
        return Font(name=self.font_name, size=self.font_size, color=self.delta_zero)

    @property
    def border_header_bottom(self) -> Border:
        s = Side(style="medium", color=self.border_strong)
        return Border(bottom=s)

    @property
    def border_subtotal_top(self) -> Border:
        s = Side(style="medium", color=self.border_strong)
        return Border(top=s)

    @property
    def border_kpi_section_top(self) -> Border:
        return self.border_subtotal_top


THEME = GstExcelTheme()

SCHEMA_ROWS = [
    ("Header row", "Background", "F8FAFC", THEME.header_bg),
    ("Header row", "Text", "475569 (all period columns)", THEME.text_header),
    ("Header row", "Bottom border", "2px E2E8F0", THEME.border_strong),
    ("Period data columns", "Background", "FFFFFF (no highlight on latest period)", THEME.white_bg),
    ("Data body", "Text", "111827", THEME.text_primary),
    ("Account / detail label", "Text", "475569", THEME.text_account),
    ("Level-1 subtotal row", "Background", "F8FAFC", THEME.subtotal_bg),
    ("Level-1 subtotal row", "Top border", "2px E2E8F0", THEME.border_strong),
    ("Block total (row_type=total)", "Background", "F8FAFC", THEME.subtotal_bg),
    ("Block total", "Top border", "2px E2E8F0", THEME.border_strong),
    ("Detail rows", "Row borders", "none", ""),
    ("Delta positive", "Font", "10B981", THEME.delta_positive),
    ("Delta negative", "Font", "DC2626", THEME.delta_negative),
    ("Delta zero", "Font", "94A3B8", THEME.delta_zero),
    ("GM CAGR column only", "Background", "F8FAFC", THEME.gm_column_bg),
    ("KPI section", "Background", "F8FAFC", THEME.header_bg),
    ("KPI section title", "Text", "1E3A5F", THEME.text_brand_title),
    ("KPI rows", "Text italic", "64748B", THEME.text_kpi),
    ("Hidden key columns", "Background", "F1F5F9", THEME.tech_col_bg),
    ("Title / subtitle", "Text", "1E3A5F", THEME.text_brand_title),
    ("Font family", "All cells", "Calibri", ""),
    ("Font size data", "All cells", "9pt", ""),
    ("All-zero data row", "Background", "F1F5F9", THEME.zero_row_bg),
    ("All-zero data row", "Text", "94A3B8", THEME.delta_zero),
]


def apply_zero_row_conditional_formatting(
    ws,
    *,
    first_row: int,
    last_row: int,
    year_col_indices: list[int],
    style_start_col: int,
    style_end_col: int,
    exclude_cols: set[int] | None = None,
    gray_font: bool = True,
) -> None:
    """
    Gray out rows where every year value column is exactly 0.
    Uses one CF rule; Excel adjusts the row reference per row in the range.
    """
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.utils import get_column_letter

    if first_row > last_row or not year_col_indices:
        return

    year_cols = sorted({c for c in year_col_indices if c > 0})
    if not year_cols:
        return

    skip = exclude_cols or set()
    parts = [f"{get_column_letter(c)}{first_row}=0" for c in year_cols]
    formula = f"AND({','.join(parts)})"

    rule = FormulaRule(
        formula=[formula],
        fill=THEME.fill_zero_row,
        font=THEME.font_zero_row if gray_font else THEME.font_base,
    )

    style_cols = [c for c in range(style_start_col, style_end_col + 1) if c not in skip]
    if not style_cols:
        return

    # Merge contiguous column runs (e.g. skip spacer column in Lead_IS).
    runs: list[tuple[int, int]] = []
    run_start = style_cols[0]
    prev = style_cols[0]
    for c in style_cols[1:]:
        if c == prev + 1:
            prev = c
            continue
        runs.append((run_start, prev))
        run_start = c
        prev = c
    runs.append((run_start, prev))

    for start_c, end_c in runs:
        start_cell = f"{get_column_letter(start_c)}{first_row}"
        end_cell = f"{get_column_letter(end_c)}{last_row}"
        ws.conditional_formatting.add(f"{start_cell}:{end_cell}", rule)


def apply_recon_portfolio_layout(
    ws,
    *,
    blocks: list[dict],
    pos_col: int,
    header_row: int,
    block_title_row: int,
    entity_code_row: int,
    last_used_col: int,
    spacer_cols: set[int],
    visible_block_kinds: frozenset[str] | None = None,
    collapsed_poscol_keys: tuple[str, ...] = ("Difference", "Financial statements"),
    tech_layout=None,
) -> None:
    """
    Column grouping, header bands and entity block titles — aligned with BS_Bucket.py.

    Presentation only: collapsed helper columns (E–I), entity portfolios grouped with
    white spacer separators, Consolidation/Difference/FS blocks always expanded.
    """
    from openpyxl.utils import get_column_letter

    visible = visible_block_kinds or frozenset({"consolidation", "difference", "fs"})

    ws.sheet_view.showOutlineSymbols = True
    ws.sheet_properties.outlinePr.summaryBelow = True
    ws.sheet_properties.outlinePr.applyStyles = True

    if tech_layout is not None:
        from databook_excel_layout import hide_helper_column_group

        hide_helper_column_group(ws, tech_layout)
    else:
        left_group_end = pos_col - 1
        for cc in range(1, left_group_end + 1):
            col = get_column_letter(cc)
            ws.column_dimensions[col].outlineLevel = 2
            ws.column_dimensions[col].hidden = True
        if left_group_end >= 1:
            ws.column_dimensions[get_column_letter(left_group_end)].collapsed = True

    pos_letter = get_column_letter(pos_col)
    ws.column_dimensions[pos_letter].outlineLevel = 0
    ws.column_dimensions[pos_letter].hidden = False

    ws.row_dimensions[entity_code_row].outlineLevel = 2
    ws.row_dimensions[entity_code_row].hidden = True

    for b in blocks:
        c_first = b["startcol"]
        c_last = b["spacer_col"]
        kind = str(b.get("kind") or "")

        if kind in visible:
            for cc in range(c_first, c_last + 1):
                col = get_column_letter(cc)
                ws.column_dimensions[col].outlineLevel = 0
                ws.column_dimensions[col].hidden = False
            continue

        for cc in range(c_first, c_last + 1):
            col = get_column_letter(cc)
            ws.column_dimensions[col].outlineLevel = 1
            ws.column_dimensions[col].hidden = False
        ws.column_dimensions[get_column_letter(c_last)].collapsed = True

    for b in blocks:
        key = str(b.get("key") or "")
        if key not in collapsed_poscol_keys:
            continue
        poscol = b.get("poscol")
        if poscol is None:
            continue
        col_l = get_column_letter(poscol)
        ws.column_dimensions[col_l].outlineLevel = 1
        ws.column_dimensions[col_l].hidden = True
        ws.column_dimensions[col_l].collapsed = True
        for c in range(b["year_startcol"], b["year_endcol"] + 1):
            c_l = get_column_letter(c)
            ws.column_dimensions[c_l].outlineLevel = 0
            ws.column_dimensions[c_l].hidden = False

    from databook_excel_layout import DatabookLayout, paint_header_band, write_entity_block_titles

    period_cols: list[int] = []
    for b in blocks:
        for y_idx in range(b["year_endcol"] - b["year_startcol"] + 1):
            period_cols.append(b["year_startcol"] + y_idx)

    layout = tech_layout if tech_layout is not None else DatabookLayout(pos_col=pos_col)
    paint_header_band(
        ws,
        layout,
        header_rows=[block_title_row, header_row],
        period_cols=period_cols,
        spacer_cols=spacer_cols,
    )
    write_entity_block_titles(ws, blocks, block_title_row=block_title_row)
    for b in blocks:
        key = str(b.get("key") or "")
        if key not in collapsed_poscol_keys:
            continue
        poscol = b.get("poscol")
        if poscol is None:
            continue
        for r in (block_title_row, header_row):
            cell = ws.cell(r, poscol)
            cell.fill = THEME.fill_header
            if r == header_row:
                cell.font = THEME.font_header
                cell.border = THEME.border_header_bottom
