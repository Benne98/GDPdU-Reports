"""Shared Excel layout helpers for databook pipeline sheets."""
from __future__ import annotations

import re
from dataclasses import dataclass

import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from gst_excel_theme import THEME

FILL_YELLOW = PatternFill("solid", fgColor="FFFF00")
DIFF_FONT = Font(name=THEME.font_name, size=THEME.font_size, color=THEME.delta_negative)
ALIGN_LEFT = Alignment(horizontal="left", vertical="center")
ALIGN_RIGHT = Alignment(horizontal="right", vertical="center")
ALIGN_CENTER = Alignment(horizontal="center", vertical="center")
BORDER_SUBTOTAL_TOP_BOTTOM = Border(
    top=THEME.border_subtotal_top.top,
    bottom=THEME.border_subtotal_top.top,
)

FS_CHECK_SOURCE_LABEL = "Source - Financial statements"
FS_CHECK_DELTA_LABEL = "Difference to trial balances"
BS_ALE_CHECK_LABEL = "CHECK: A = L + E"
KPI_SECTION_TITLE_DEFAULT = "KPIs"
FILL_KPI = THEME.fill_subtotal
FONT_MAPPING = Font(name=THEME.font_name, size=THEME.font_size, color=THEME.text_header)


def check_rows_after_table(last_table_row: int, pair_count: int) -> list[int]:
    """
    Allocate check row numbers below the table with one blank row between each
    consecutive check line (still one collapsible check portfolio).
    """
    return check_row_groups_after_table(last_table_row, [1] * pair_count)


def check_row_groups_after_table(last_table_row: int, group_sizes: list[int]) -> list[int]:
    """
    Allocate check rows below the table. Rows within a group are adjacent;
    one blank row is inserted between groups.
    """
    rows: list[int] = []
    current = last_table_row + 2
    for gi, size in enumerate(group_sizes):
        if gi > 0:
            current += 1
        for _ in range(size):
            rows.append(current)
            current += 1
    return rows


def write_kpi_section_title_row(
    ws,
    row: int,
    *,
    title_cols: list[int],
    label_col: int,
    title_text: str = KPI_SECTION_TITLE_DEFAULT,
    spacer_cols: set[int] | None = None,
    fill_end_col: int | None = None,
    title_font: Font | None = None,
) -> None:
    """Lead_IS-style KPI band title (without the '% total output' wording)."""
    skip = spacer_cols or set()
    end_col = fill_end_col if fill_end_col is not None else max(title_cols, default=label_col)
    font = title_font or Font(
        name=THEME.font_name,
        size=THEME.font_size,
        color="00A7B5",
        italic=True,
        bold=True,
    )
    for cc in range(label_col, end_col + 1):
        ws.cell(row, cc).fill = THEME.fill_white if cc in skip else FILL_KPI
        ws.cell(row, cc).border = Border()
    cell = ws.cell(row, label_col, title_text)
    cell.font = font
    cell.alignment = ALIGN_LEFT


def write_bs_ale_check_section(
    ws,
    *,
    blocks: list[dict],
    years: list[str],
    check_row: int,
    total_assets_row: int,
    total_el_row: int,
    pos_col: int,
    block_kinds: frozenset[str],
    num_fmt: str = "#,##0",
    hide_outline: bool = False,
) -> None:
    """Balance check: Total assets + Total equity & liabilities should equal 0."""
    from openpyxl.formatting.rule import CellIsRule

    lbl = ws.cell(check_row, pos_col, BS_ALE_CHECK_LABEL)
    lbl.font = THEME.font_base
    lbl.alignment = ALIGN_LEFT

    for b in blocks:
        if b.get("kind") not in block_kinds:
            continue
        for y_idx, _year in enumerate(years):
            col = b["year_startcol"] + y_idx
            assets_cell = ws.cell(total_assets_row, col)
            el_cell = ws.cell(total_el_row, col)
            bal_cell = ws.cell(check_row, col)
            bal_cell.value = f"={assets_cell.coordinate}+{el_cell.coordinate}"
            bal_cell.number_format = num_fmt
            bal_cell.alignment = ALIGN_RIGHT
            bal_cell.font = DIFF_FONT
            ws.conditional_formatting.add(
                bal_cell.coordinate,
                CellIsRule(operator="notEqual", formula=["0"], font=DIFF_FONT),
            )

    if hide_outline:
        ws.row_dimensions[check_row].outlineLevel = 2
        ws.row_dimensions[check_row].hidden = True


def collapse_check_portfolio(ws, rows: list[int | None]) -> None:
    """Single collapsible check portfolio (FA roll forward pattern), incl. blank rows between checks."""
    valid = [int(rr) for rr in rows if rr is not None]
    if not valid:
        return
    first_row = min(valid)
    last_row = max(valid)
    for rr in range(first_row, last_row + 1):
        ws.row_dimensions[rr].outlineLevel = 1
        ws.row_dimensions[rr].hidden = True
    ws.row_dimensions[last_row].collapsed = True


def collapse_check_outline_rows(ws, rows: list[int | None]) -> None:
    collapse_check_portfolio(ws, rows)


RECON_BLOCK_TITLE_ROW = 7
RECON_HEADER_ROW = 8
DEFAULT_NA_BUCKET_ORDER = ["FA", "TWC", "OWC", "ND", "Other", "Equity"]
BS_RECON_SHEET = "BS_Reconciliation"
AGGREGATED_BLOCK_TITLE = "Aggregated"


def _recon_block_range(
    ws_recon,
    block_title: str,
    *,
    block_title_row: int = RECON_BLOCK_TITLE_ROW,
) -> tuple[int | None, int]:
    titles: list[tuple[int, str]] = []
    for cc in range(1, ws_recon.max_column + 1):
        v = ws_recon.cell(block_title_row, cc).value
        if isinstance(v, str) and v.strip():
            titles.append((cc, v.strip()))

    target = block_title.strip().lower()
    block_start = None
    block_end = ws_recon.max_column + 1
    for idx, (cc, title) in enumerate(titles):
        if title.lower() == target:
            block_start = cc
            if idx + 1 < len(titles):
                block_end = titles[idx + 1][0]
            break
    return block_start, block_end


def find_recon_block_start_col(
    ws_recon,
    block_title: str,
    *,
    block_title_row: int = RECON_BLOCK_TITLE_ROW,
) -> int | None:
    block_start, _ = _recon_block_range(ws_recon, block_title, block_title_row=block_title_row)
    return block_start


def find_recon_block_year_cols(
    ws_recon,
    block_title: str,
    years: list[str],
    *,
    block_title_row: int = RECON_BLOCK_TITLE_ROW,
    header_row: int = RECON_HEADER_ROW,
) -> dict[str, int]:
    """Map period labels to columns within a reconciliation portfolio block."""
    block_start, block_end = _recon_block_range(
        ws_recon, block_title, block_title_row=block_title_row
    )
    if block_start is None:
        return {}

    year_set = set(years)
    year_col: dict[str, int] = {}
    for cc in range(block_start, block_end):
        v = ws_recon.cell(header_row, cc).value
        if isinstance(v, str) and v.strip() in year_set:
            year_col[v.strip()] = cc
    return year_col


def find_recon_block_col_by_header(
    ws_recon,
    block_title: str,
    header_label: str,
    *,
    block_title_row: int = RECON_BLOCK_TITLE_ROW,
    header_row: int = RECON_HEADER_ROW,
) -> int | None:
    block_start, block_end = _recon_block_range(
        ws_recon, block_title, block_title_row=block_title_row
    )
    if block_start is None:
        return None
    target = str(header_label or "").strip()
    for cc in range(block_start, block_end):
        v = ws_recon.cell(header_row, cc).value
        if isinstance(v, str) and v.strip() == target:
            return cc
    return None


_SNAPSHOT_PERIOD_RE = re.compile(r"^[A-Z][a-z]{2}(\d{2})A$", re.IGNORECASE)


def find_recon_block_col_for_period_label(
    ws_recon,
    block_title: str,
    period_label: str,
    *,
    block_title_row: int = RECON_BLOCK_TITLE_ROW,
    header_row: int = RECON_HEADER_ROW,
) -> int | None:
    """Match OPOS snapshot label to recon block column (exact, then FY year suffix e.g. Dec24A→Jul24A)."""
    exact = find_recon_block_col_by_header(
        ws_recon,
        block_title,
        period_label,
        block_title_row=block_title_row,
        header_row=header_row,
    )
    if exact is not None:
        return exact
    target = str(period_label or "").strip()
    m = _SNAPSHOT_PERIOD_RE.match(target)
    if not m:
        return None
    yy_suffix = f"{m.group(1)}A".upper()
    block_start, block_end = _recon_block_range(
        ws_recon, block_title, block_title_row=block_title_row
    )
    if block_start is None:
        return None
    matches: list[int] = []
    for cc in range(block_start, block_end):
        v = ws_recon.cell(header_row, cc).value
        if isinstance(v, str) and v.strip().upper().endswith(yy_suffix):
            matches.append(cc)
    if not matches:
        return None
    return matches[-1]


def find_recon_aggregated_year_cols(
    ws_recon,
    years: list[str],
    *,
    block_title_row: int = RECON_BLOCK_TITLE_ROW,
    header_row: int = RECON_HEADER_ROW,
) -> dict[str, int]:
    return find_recon_block_year_cols(
        ws_recon,
        AGGREGATED_BLOCK_TITLE,
        years,
        block_title_row=block_title_row,
        header_row=header_row,
    )


def find_bs_recon_row_contains(
    ws_recon,
    needle: str,
    *,
    pos_col: int | None = None,
) -> int | None:
    col = pos_col if pos_col is not None else LAYOUT_NA.pos_col
    needle_l = str(needle or "").strip().lower()
    if not needle_l:
        return None
    for rr in range(1, ws_recon.max_row + 1):
        v = ws_recon.cell(rr, col).value
        if isinstance(v, str) and needle_l in v.lower():
            return rr
    return None


def bs_recon_aggregated_ref(
    row: int,
    col: int,
    *,
    sheet_name: str = BS_RECON_SHEET,
) -> str:
    return f"='{sheet_name}'!{get_column_letter(col)}{row}"


def detect_period_bucket_cols(
    ws_,
    period_labels: list[str],
    *,
    period_row: int = RECON_BLOCK_TITLE_ROW,
    bucket_header_row: int = RECON_HEADER_ROW,
    bucket_order: list[str] | None = None,
) -> dict[str, dict[str, int]]:
    """BS_Bucket NA-classification blocks: period label -> {bucket: col}."""
    order = bucket_order or DEFAULT_NA_BUCKET_ORDER
    period_set = {str(p).strip() for p in period_labels}
    period_to_cols: dict[str, dict[str, int]] = {}
    for cc in range(1, ws_.max_column + 1):
        v = ws_.cell(period_row, cc).value
        if not isinstance(v, str):
            continue
        period = v.strip()
        if period not in period_set:
            continue
        bucket_map: dict[str, int] = {}
        steps = 0
        c2 = cc
        while c2 <= ws_.max_column and steps < 60 and len(bucket_map) < len(order):
            hv = ws_.cell(bucket_header_row, c2).value
            if isinstance(hv, str):
                hvs = hv.strip()
                if hvs in order:
                    bucket_map[hvs] = c2
            c2 += 1
            steps += 1
        if all(b in bucket_map for b in order):
            period_to_cols[period] = bucket_map
    return period_to_cols


def assign_recon_block_columns(
    blocks: list[dict],
    *,
    first_col: int,
    n_years: int,
) -> list[dict]:
    current_col = first_col
    for b in blocks:
        b["startcol"] = current_col
        extra = 1 if b.get("has_plpos") else 0
        b["poscol"] = current_col if extra == 1 else None
        b["year_startcol"] = current_col + extra
        b["year_endcol"] = b["year_startcol"] + n_years - 1
        b["spacer_col"] = b["year_endcol"] + 1
        current_col = b["spacer_col"] + 1
    return blocks


def build_entity_recon_blocks(
    entity: str,
    titles: dict,
    *,
    first_col: int,
    n_years: int,
    difference_key: str = "Difference",
    fs_key: str = "Financial statements",
) -> list[dict]:
    """Entity trial balance + Difference + FS (no Aggregated/IC/Consolidation)."""
    blocks = [
        {
            "kind": "entity",
            "key": entity,
            "code": entity,
            "title": entity,
            "has_plpos": False,
        },
        {
            "kind": "difference",
            "key": difference_key,
            "code": difference_key,
            "title": titles.get("difference_title", difference_key),
            "has_plpos": True,
        },
        {
            "kind": "fs",
            "key": fs_key,
            "code": fs_key,
            "title": titles.get("financial_statements_title", fs_key),
            "has_plpos": True,
        },
    ]
    return assign_recon_block_columns(blocks, first_col=first_col, n_years=n_years)

FREE_LEFT_COL = 1  # A — leading placeholder inside helper block
HELPER_COLS_PL = 4  # Reported, blank, L3, L4
HELPER_COLS_BS = 5  # Reported, blank, L2, L3, L4
HELPER_COLS_NA = 5  # Reported, blank, NA/L5, L3, L4
TECH_SPACER_WIDTH = 1.14


@dataclass(frozen=True)
class DatabookLayout:
    """A..?+1 collapsed helper block (grey) | ?+2 POS table start | values."""

    helper_start_col: int = FREE_LEFT_COL
    map_start_col: int = 2
    map_end_col: int = 5
    helper_end_col: int = 6  # trailing placeholder (?+1), end of helper block
    pos_col: int = 7
    first_value_col: int = 8
    fill_padding_rows: int = 100
    fill_padding_cols: int = 40

    @property
    def free_left_col(self) -> int:
        return self.helper_start_col

    @property
    def spacer_col(self) -> int:
        return self.helper_end_col


def databook_column_layout(num_helper_cols: int) -> DatabookLayout:
    """Build column geometry: A..?+1 helper block, ?+2 POS."""
    map_start_col = 2
    map_end_col = map_start_col + num_helper_cols - 1
    helper_end_col = map_end_col + 1
    pos_col = helper_end_col + 1
    return DatabookLayout(
        helper_start_col=FREE_LEFT_COL,
        map_start_col=map_start_col,
        map_end_col=map_end_col,
        helper_end_col=helper_end_col,
        pos_col=pos_col,
        first_value_col=pos_col + 1,
    )


LAYOUT_PL = databook_column_layout(HELPER_COLS_PL)
LAYOUT_BS = databook_column_layout(HELPER_COLS_BS)
LAYOUT_NA = databook_column_layout(HELPER_COLS_NA)


def hide_helper_column_group(ws, layout: DatabookLayout) -> None:
    """Collapse full helper block A..?+1; POS (?+2) is the first visible column."""
    for cc in range(layout.helper_start_col, layout.helper_end_col + 1):
        col = get_column_letter(cc)
        ws.column_dimensions[col].outlineLevel = 2
        ws.column_dimensions[col].hidden = True
    ws.column_dimensions[get_column_letter(layout.helper_end_col)].collapsed = True

    pos_letter = get_column_letter(layout.pos_col)
    ws.column_dimensions[pos_letter].outlineLevel = 0
    ws.column_dimensions[pos_letter].hidden = False


def paint_grey_white_canvas(
    ws,
    layout: DatabookLayout,
    *,
    last_row: int,
    last_col: int | None = None,
) -> None:
    """Grey full helper block A..spacer (incl. padding cols); white from POS onward."""
    fill_end_col = (last_col or 0) + layout.fill_padding_cols
    for rr in range(1, last_row + 1):
        for cc in range(layout.helper_start_col, layout.helper_end_col + 1):
            ws.cell(rr, cc).fill = THEME.fill_tech
        for cc in range(layout.pos_col, fill_end_col + 1):
            ws.cell(rr, cc).fill = THEME.fill_white


def apply_databook_row_border_band(
    ws,
    row: int,
    layout: DatabookLayout,
    border: Border,
    *,
    last_used_col: int,
    skip_cols: set[int] | None = None,
) -> None:
    """Row border through mapping cols + table; skip column A and portfolio spacer cols."""
    skip = skip_cols or set()
    for cc in range(layout.map_start_col, layout.map_end_col + 1):
        ws.cell(row, cc).border = border
    for cc in range(layout.pos_col, last_used_col + 1):
        if cc in skip:
            ws.cell(row, cc).border = Border()
            continue
        ws.cell(row, cc).border = border


def write_mapping_header_row(
    ws,
    layout: DatabookLayout,
    *,
    header_row: int,
    header_row_7: int,
    labels: list[str],
) -> None:
    """Write mapping column headers (Reported, blank, L2/L3/L4/NA) with databook styling."""
    for offset, label in enumerate(labels):
        col = layout.map_start_col + offset
        if col > layout.map_end_col:
            break
        cell = ws.cell(header_row, col, label)
        cell.font = THEME.font_header
        cell.alignment = ALIGN_LEFT
        cell.fill = THEME.fill_header
        ws.cell(header_row_7, col).fill = THEME.fill_header


def apply_subtotal_row_style(
    ws,
    row: int,
    *,
    layout: DatabookLayout | None = None,
    pos_col: int,
    value_cols: list[int],
    fill=None,
    pos_border: Border | None = None,
    value_border: Border | None = None,
    last_used_col: int | None = None,
) -> None:
    """Subtotal/total fill on POS + values; optional borders also on mapping cols (not A/spacer)."""
    row_fill = fill or THEME.fill_subtotal
    ws.cell(row, pos_col).fill = row_fill
    if pos_border is not None:
        ws.cell(row, pos_col).border = pos_border
    for cc in value_cols:
        ws.cell(row, cc).fill = row_fill
        if value_border is not None:
            ws.cell(row, cc).border = value_border
    if layout is not None and pos_border is not None:
        end_col = last_used_col if last_used_col is not None else max(value_cols, default=pos_col)
        for cc in range(layout.map_start_col, layout.map_end_col + 1):
            ws.cell(row, cc).border = pos_border


def header_band_columns(
    layout: DatabookLayout,
    *,
    period_cols: list[int],
    spacer_cols: set[int] | None = None,
) -> list[int]:
    skip = spacer_cols or set()
    cols = list(range(layout.map_start_col, layout.map_end_col + 1))
    cols.append(layout.pos_col)
    cols.extend(c for c in period_cols if c not in skip)
    return sorted({c for c in cols if c > 0 and c not in skip})


def paint_header_band(
    ws,
    layout: DatabookLayout,
    *,
    header_rows: list[int],
    period_cols: list[int],
    spacer_cols: set[int] | None = None,
) -> None:
    """Header fill/border only on mapping columns, POS, and period columns."""
    cols = header_band_columns(layout, period_cols=period_cols, spacer_cols=spacer_cols)
    for r in header_rows:
        for cc in cols:
            ws.cell(r, cc).fill = THEME.fill_header
    for cc in cols:
        ws.cell(header_rows[-1], cc).border = THEME.border_header_bottom


def write_period_headers(
    ws,
    labels: list[str],
    cols: list[int],
    *,
    header_row: int,
    align_right: bool = True,
) -> None:
    alignment = ALIGN_RIGHT if align_right else ALIGN_CENTER
    for label, col in zip(labels, cols):
        cell = ws.cell(header_row, col, label)
        cell.font = THEME.font_header
        cell.alignment = alignment
        cell.fill = THEME.fill_header
        cell.border = THEME.border_header_bottom


def write_entity_block_titles(
    ws,
    blocks: list[dict],
    *,
    block_title_row: int,
    font=None,
) -> None:
    title_font = font or THEME.font_header
    for b in blocks:
        if b.get("has_plpos"):
            title_col = b["year_startcol"]
            if b.get("poscol"):
                ws.cell(block_title_row, b["poscol"]).fill = THEME.fill_white
        else:
            title_col = b["startcol"]
        title_cell = ws.cell(block_title_row, title_col)
        title_cell.font = title_font
        title_cell.alignment = ALIGN_CENTER


def apply_bs_hierarchy_borders(
    ws,
    row_structure: list[dict],
    *,
    layout: DatabookLayout,
    last_used_col: int,
    total_labels: frozenset[str] | None = None,
    skip_cols: set[int] | None = None,
) -> None:
    """L2 subtotals: top border; configured totals: top + bottom."""
    totals = total_labels or frozenset(
        {"Total assets", "Total equity & liabilities"}
    )
    for r in row_structure:
        excel_row = r.get("_excel_row")
        if excel_row is None:
            continue
        row_type = r.get("type")
        label = str(r.get("label") or "")
        if row_type == "subtotal_l2":
            border = THEME.border_subtotal_top
        elif row_type == "total" and label in totals:
            border = BORDER_SUBTOTAL_TOP_BOTTOM
        else:
            continue
        apply_databook_row_border_band(
            ws,
            excel_row,
            layout,
            border,
            last_used_col=last_used_col,
            skip_cols=skip_cols,
        )


def _row_amounts_all_zero(
    df: pd.DataFrame,
    row: dict,
    value_cols: list[str],
    *,
    source_col: str,
    reported_value: str,
    bucket_col: str | None = None,
) -> bool:
    if not value_cols:
        return True
    rep = (
        df[source_col]
        .astype(str)
        .str.strip()
        .str.lower()
        .eq(str(reported_value).strip().lower())
    )
    l2 = str(row.get("L2") or "").strip()
    l3 = str(row.get("L3") or "").strip()
    l4 = str(row.get("L4") or "").strip()
    mask = rep
    if "L2" in df.columns and l2:
        mask = mask & df["L2"].astype(str).str.strip().eq(l2)
    if "L3" in df.columns and l3:
        mask = mask & df["L3"].astype(str).str.strip().eq(l3)
    if "L4" in df.columns:
        l4_series = df["L4"]
        if l4 == "" or l4.lower() in {"", "nan", "none", "null"}:
            l4_mask = l4_series.isna() | l4_series.astype(str).str.strip().str.lower().isin(
                {"", "nan", "none", "null"}
            )
        else:
            l4_mask = l4_series.astype(str).str.strip().eq(l4)
        mask = mask & l4_mask
    bucket_key = bucket_col if bucket_col else "L5"
    if bucket_key in df.columns and row.get("L5") is not None:
        mask = mask & df[bucket_key].astype(str).str.strip().eq(str(row.get("L5")).strip())
    elif "L5" in df.columns and row.get("L5") is not None:
        mask = mask & df["L5"].astype(str).str.strip().eq(str(row.get("L5")).strip())
    subset = df.loc[mask, value_cols]
    if subset.empty:
        return True
    nums = subset.apply(pd.to_numeric, errors="coerce").fillna(0)
    return bool((nums.abs() < 1e-9).all().all())


def prune_zero_value_rows(
    struct: list[dict],
    df: pd.DataFrame,
    value_cols: list[str],
    *,
    source_col: str,
    reported_value: str = "reported",
    bucket_col: str | None = None,
    detail_types: frozenset[str] | None = None,
) -> list[dict]:
    """Remove detail rows whose reported amounts are zero in every period column."""
    if not struct or not value_cols:
        return struct
    dtypes = detail_types or frozenset({"detail", "detail_single"})
    remove_idx: set[int] = set()
    for i, row in enumerate(struct):
        if row.get("type") not in dtypes:
            continue
        if _row_amounts_all_zero(
            df,
            row,
            value_cols,
            source_col=source_col,
            reported_value=reported_value,
            bucket_col=bucket_col,
        ):
            remove_idx.add(i)

    if not remove_idx:
        return struct

    pruned = [r for i, r in enumerate(struct) if i not in remove_idx]

    # Drop subtotal rows that no longer have detail children.
    out: list[dict] = []
    i = 0
    while i < len(pruned):
        row = pruned[i]
        row_type = row.get("type")
        if row_type in {"subtotal", "subtotal_l3"}:
            l3 = str(row.get("L3") or row.get("tech_label") or "").strip()
            has_detail = any(
                r.get("type") in dtypes
                and str(r.get("L3") or "").strip() == l3
                for r in pruned
            )
            if not has_detail:
                i += 1
                continue
        out.append(row)
        i += 1
    return out


def copy_worksheet(src_ws, dest_wb, title: str):
    """Copy a worksheet including cell styles into dest_wb (cross-workbook safe)."""
    from copy import copy

    safe_title = str(title)[:31]
    if safe_title in dest_wb.sheetnames:
        del dest_wb[safe_title]
    dest = dest_wb.create_sheet(safe_title)
    for row in src_ws.iter_rows():
        for cell in row:
            dest_cell = dest[cell.coordinate]
            dest_cell.value = cell.value
            if cell.has_style:
                dest_cell.font = copy(cell.font)
                dest_cell.fill = copy(cell.fill)
                dest_cell.border = copy(cell.border)
                dest_cell.alignment = copy(cell.alignment)
                dest_cell.number_format = cell.number_format
    for col_letter, dim in src_ws.column_dimensions.items():
        dest_dim = dest.column_dimensions[col_letter]
        dest_dim.width = dim.width
        dest_dim.hidden = dim.hidden
        dest_dim.outlineLevel = dim.outlineLevel
        dest_dim.collapsed = dim.collapsed
    for row_idx, dim in src_ws.row_dimensions.items():
        dest_dim = dest.row_dimensions[row_idx]
        dest_dim.height = dim.height
        dest_dim.hidden = dim.hidden
        dest_dim.outlineLevel = dim.outlineLevel
    dest.sheet_format = copy(src_ws.sheet_format)
    dest.sheet_properties = copy(src_ws.sheet_properties)
    dest.freeze_panes = src_ws.freeze_panes
    return dest


def ensure_master_sheet(wb, source_wb, sheet_name: str) -> None:
    """Keep formatted master sheet from workbook; copy from source if missing."""
    if sheet_name in wb.sheetnames:
        return
    if sheet_name in source_wb.sheetnames:
        copy_worksheet(source_wb[sheet_name], wb, sheet_name)
        return
    raise RuntimeError(f"Master sheet '{sheet_name}' not found in target or source workbook.")


def refresh_master_sheet(wb, df: pd.DataFrame, sheet_name: str) -> None:
    """Replace audit master sheet with current dataframe (keeps SUMIFS ranges valid)."""
    if sheet_name in wb.sheetnames:
        del wb[sheet_name]
    ws = wb.create_sheet(sheet_name)
    for c_idx, colname in enumerate(df.columns, start=1):
        ws.cell(1, c_idx, colname)
    for r_idx, row in enumerate(df.itertuples(index=False, name=None), start=2):
        for c_idx, val in enumerate(row, start=1):
            ws.cell(r_idx, c_idx, val)


def recon_check_yellow_columns(
    blocks: list[dict],
    *,
    pos_col: int,
    block_kinds: frozenset[str],
    spacer_cols: set[int] | None = None,
    period_indices: list[int] | None = None,
) -> list[int]:
    """Continuous yellow band from POS through last check block column."""
    skip = spacer_cols or set()
    eligible = [b for b in blocks if b.get("kind") in block_kinds]
    if not eligible:
        return [c for c in [pos_col] if c not in skip]
    last_col = max(b["year_endcol"] for b in eligible)
    cols = {pos_col}
    for b in eligible:
        if b.get("poscol"):
            cols.add(int(b["poscol"]))
        indices = period_indices if period_indices is not None else range(
            b["year_endcol"] - b["year_startcol"] + 1
        )
        for y_idx in indices:
            cols.add(b["year_startcol"] + y_idx)
    return sorted(c for c in cols if c not in skip and c <= last_col)


def paint_check_source_yellow(ws, source_row: int, cols: list[int]) -> None:
    for cc in cols:
        ws.cell(source_row, cc).fill = FILL_YELLOW


def write_fs_check_section(
    ws,
    *,
    blocks: list[dict],
    years: list[str],
    source_row: int,
    delta_row: int,
    anchor_row: int,
    pos_col: int,
    block_kinds: frozenset[str],
    source_values: dict[str, dict[str, float | None]] | None = None,
    source_label: str = FS_CHECK_SOURCE_LABEL,
    delta_label: str = FS_CHECK_DELTA_LABEL,
    num_fmt: str = "#,##0",
    hide_outline: bool = False,
    yellow_block_kinds: frozenset[str] | None = None,
    spacer_cols: set[int] | None = None,
    yellow_period_indices: list[int] | None = None,
) -> list[int]:
    """
    Lead_IS-style FS check: yellow source row, red delta font, optional empty values.
    Returns period column indices touched.
    """
    from openpyxl.formatting.rule import CellIsRule

    src = ws.cell(source_row, pos_col, source_label)
    src.font = THEME.font_base
    src.alignment = ALIGN_LEFT

    delta_lbl = ws.cell(delta_row, pos_col, delta_label)
    delta_lbl.font = THEME.font_base
    delta_lbl.alignment = ALIGN_LEFT

    period_cols: list[int] = []
    values = source_values or {}

    for b in blocks:
        if b.get("kind") not in block_kinds:
            continue
        entity_key = str(b.get("key") or "")
        year_map = values.get(entity_key, {})
        for y_idx, year in enumerate(years):
            col = b["year_startcol"] + y_idx
            period_cols.append(col)
            val = year_map.get(year)
            src_cell = ws.cell(source_row, col, val if val is not None else None)
            src_cell.number_format = num_fmt
            src_cell.alignment = ALIGN_RIGHT
            src_cell.font = THEME.font_base

            net_cell = ws.cell(anchor_row, col)
            chk_cell = ws.cell(source_row, col)
            diff_cell = ws.cell(delta_row, col)
            diff_cell.value = f"={net_cell.coordinate}-{chk_cell.coordinate}"
            diff_cell.number_format = num_fmt
            diff_cell.alignment = ALIGN_RIGHT
            diff_cell.font = THEME.font_base
            ws.conditional_formatting.add(
                diff_cell.coordinate,
                CellIsRule(operator="notEqual", formula=["0"], font=DIFF_FONT),
            )

    yellow_cols = recon_check_yellow_columns(
        blocks,
        pos_col=pos_col,
        block_kinds=yellow_block_kinds or block_kinds,
        spacer_cols=spacer_cols,
        period_indices=yellow_period_indices,
    )
    paint_check_source_yellow(ws, source_row, yellow_cols)

    if hide_outline:
        for rr in (source_row, delta_row):
            ws.row_dimensions[rr].outlineLevel = 2
            ws.row_dimensions[rr].hidden = True

    return period_cols
