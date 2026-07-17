"""Single-process Fast Track workbook batch runner.

Every selected section receives the same openpyxl Workbook. Section adapters must
not save it; the runner performs one atomic save after final sheet ordering.
"""
from __future__ import annotations

import json
import runpy
import sys
import tempfile
from contextlib import contextmanager
from pathlib import Path
from typing import Any, Callable

import pandas as pd

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))
if str(ROOT / "backend") not in sys.path:
    sys.path.insert(0, str(ROOT / "backend"))

from databook_workbook import (  # noqa: E402
    atomic_save_workbook,
    derive_entity_order_by_latest_revenue,
    open_batch_workbook,
    remove_placeholder_fast_track_sheet,
    reorder_workbook_sheets,
)

SECTION_ORDER = ("databook", "opos", "fte", "fa", "revenue")
SectionRunner = Callable[[dict[str, Any], Any], Any]


def _parse_ltm_parts(ltm: str) -> tuple[int, int] | None:
    raw = str(ltm or "").strip()
    if not raw or "-" not in raw:
        return None
    year_s, month_s = raw.split("-", 1)
    try:
        year = int(year_s)
        month = int(month_s)
    except ValueError:
        return None
    if 1 <= month <= 12:
        return year, month
    return None


def _coerce_int(value: Any, default: int | None = None) -> int | None:
    if value in (None, ""):
        return default
    try:
        return int(float(value))
    except (TypeError, ValueError):
        return default


def _apply_manifest_dates(config: dict[str, Any], dates: dict[str, Any] | None = None) -> None:
    """Ensure as-of year/month and fiscal end are present for revenue legacy scripts."""
    dates = dates if isinstance(dates, dict) else {}
    for key in ("ltm_month", "first_fy", "fy_end_month", "fy_end_day", "current_year", "current_month"):
        if config.get(key) in (None, "") and dates.get(key) not in (None, ""):
            config[key] = dates[key]

    ltm = str(config.get("ltm_month") or "").strip()
    if ltm:
        config["ltm_month"] = ltm

    cy = _coerce_int(config.get("current_year"))
    cm = _coerce_int(config.get("current_month"))
    if cy is None or cm is None:
        parsed = _parse_ltm_parts(ltm)
        if parsed:
            cy, cm = parsed
        else:
            from datetime import date

            today = date.today()
            cy = cy if cy is not None else today.year
            cm = cm if cm is not None else today.month
    config["current_year"] = cy
    config["current_month"] = cm

    first_fy = _coerce_int(config.get("first_fy"), default=cy - 3)
    config["first_fy"] = first_fy if first_fy is not None else cy - 3

    fy_end_m = _coerce_int(config.get("fy_end_month"), default=12)
    fy_end_d = _coerce_int(config.get("fy_end_day"), default=31)
    config["fy_end_month"] = fy_end_m if fy_end_m is not None else 12
    config["fy_end_day"] = fy_end_d if fy_end_d is not None else 31
    config["fiscal_year_end_month"] = config["fy_end_month"]
    config["fiscal_year_end_day"] = config["fy_end_day"]
    config["as_of_year"] = config["current_year"]
    config["as_of_month"] = config["current_month"]


def _resolve_sheet_name(file_path: str, sheet_name: str | None) -> str:
    """Resolve worksheet name; never silently assume 'Sheet1'."""
    path = Path(str(file_path or "").strip()).expanduser()
    requested = str(sheet_name or "").strip()
    import openpyxl

    if not path.is_file():
        if requested:
            return requested
        raise ValueError(f"Revenue workbook not found: {file_path}")

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        names = list(wb.sheetnames)
    finally:
        wb.close()
    if not names:
        raise ValueError(f"Revenue workbook has no worksheets: {path}")
    if requested:
        if requested in names:
            return requested
        raise ValueError(
            f"Worksheet named {requested!r} not found in {path.name}. "
            f"Available sheets: {', '.join(names)}"
        )
    return names[0]


def _section_config(manifest: dict[str, Any], name: str) -> tuple[bool, dict[str, Any]]:
    sections = manifest.get("sections")
    if not isinstance(sections, dict):
        raise ValueError("manifest.sections must be an object")
    raw = sections.get(name, False)
    if isinstance(raw, bool):
        return raw, {}
    if not isinstance(raw, dict):
        raise ValueError(f"manifest.sections.{name} must be a boolean or object")
    enabled = bool(
        raw.get(
            "enabled",
            raw.get("selected", raw.get("included", True)),
        )
    )
    # API callers may use the canonical {enabled, config} envelope. Rasa's
    # Fast Track state keeps source metadata and a nested analysis `config`
    # in the same section object, so only unwrap a true envelope.
    envelope_keys = set(raw) - {"enabled", "selected", "included", "config"}
    if not envelope_keys:
        config = dict(raw.get("config") or {})
    else:
        config = dict(raw)
        nested = raw.get("config")
        if isinstance(nested, dict):
            config.update(nested)
    config.pop("enabled", None)
    config.pop("selected", None)
    config.pop("included", None)
    config.pop("config", None)
    return enabled, config


def validate_manifest(manifest: dict[str, Any]) -> dict[str, Any]:
    if not isinstance(manifest, dict):
        raise ValueError("manifest must be an object")
    unknown = set((manifest.get("sections") or {}).keys()) - set(SECTION_ORDER)
    if unknown:
        raise ValueError(f"Unknown Fast Track sections: {', '.join(sorted(unknown))}")
    selected = []
    normalized_sections: dict[str, dict[str, Any]] = {}
    for name in SECTION_ORDER:
        enabled, config = _section_config(manifest, name)
        normalized_sections[name] = {"enabled": enabled, "config": config}
        if enabled:
            selected.append(name)
    output_path = str(manifest.get("output_path") or "").strip()
    if not output_path:
        raise ValueError("manifest.output_path is required")
    source_path = str(
        manifest.get("master_path")
        or manifest.get("input_master_path")
        or normalized_sections["databook"]["config"].get("master_path")
        or ""
    ).strip()
    if normalized_sections["databook"]["enabled"] and not source_path:
        raise ValueError("Databook requires an existing/imported master_path")
    needs_master = any(
        normalized_sections[name]["enabled"] for name in ("databook", "opos", "fte", "fa")
    )
    if needs_master and not source_path:
        raise ValueError(
            "Fast Track requires master_path when databook, opos, fte, or fa is enabled"
        )
    if not selected:
        raise ValueError("Fast Track requires at least one selected section")
    return {
        **manifest,
        "output_path": output_path,
        "master_path": source_path,
        "sections": normalized_sections,
        "selected_sections": selected,
    }


def _run_opos(config: dict[str, Any], wb) -> None:
    from opos import normalize_config, run_opos

    cfg = dict(config)
    cfg.setdefault("company", str(cfg.get("company_name") or cfg.get("group_name") or ""))
    cfg.setdefault("company_name", str(cfg.get("company") or ""))
    cfg.setdefault("title", str(cfg.get("project_name") or "Project"))
    sid = str(cfg.get("case_id") or cfg.get("session_id") or "").strip()
    snaps = cfg.get("snapshots")
    if (
        sid
        and not cfg.get("sides")
        and isinstance(snaps, list)
        and snaps
        and isinstance(snaps[0], dict)
        and "debitor" in snaps[0]
        and "kreditor" in snaps[0]
    ):
        try:
            import sys
            from pathlib import Path

            backend_root = Path(__file__).resolve().parents[1] / "backend"
            if str(backend_root) not in sys.path:
                sys.path.insert(0, str(backend_root))
            from app.routers.fdd_bot import _normalize_opos_config

            cfg = _normalize_opos_config(cfg, sid)
        except Exception:
            pass
    run_opos(normalize_config(cfg), wb=wb, save=False)


def _run_fte(config: dict[str, Any], wb) -> None:
    from FTE_payroll import run_fte_payroll

    cfg = dict(config)
    cfg.setdefault("company", str(cfg.get("company_name") or cfg.get("group_name") or ""))
    cfg.setdefault("title", str(cfg.get("project_name") or "Project"))
    # Fast Track passes master_path; FTE reads Master_PL via master_pl_path
    cfg.setdefault("master_pl_path", str(cfg.get("master_path") or "").strip())
    run_fte_payroll(cfg, wb=wb, save=False)


def _run_fa(config: dict[str, Any], wb) -> None:
    from fixed_assets_rollf import run_fixed_assets_rollf

    cfg = dict(config)
    cfg.setdefault("company", str(cfg.get("company_name") or cfg.get("group_name") or ""))
    cfg.setdefault("title", str(cfg.get("project_name") or "Project"))
    run_fixed_assets_rollf(cfg, wb=wb, save=False)


def _revenue_script_configs(config: dict[str, Any]) -> list[tuple[str, dict[str, Any]]]:
    _apply_manifest_dates(config, {})
    selected = {
        str(value).strip().lower()
        for value in config.get("selected_analyses", [])
        if str(value).strip()
    }
    file_path = str(config.get("file_path") or "").strip()
    if not file_path:
        raise ValueError("Revenue section requires file_path")
    if not selected:
        raise ValueError("Revenue section requires selected_analyses")

    sheet_name = _resolve_sheet_name(file_path, config.get("sheet_name"))
    calc_mode = str(config.get("calc_mode") or "invoice")
    mapping_mode = "year" if str(config.get("date_mapping") or "date") == "period" else "date"
    invoice_col = str(config.get("invoice_col") or config.get("period_col") or "")
    revenue_col = str(config.get("revenue_col") or "")
    profit_mode = str(config.get("margin_mode") or config.get("profit_mode") or "cost")
    cost_col = str(config.get("cost_col") or "")
    profit_col = str(config.get("profit_col") or "")
    start_col = str(config.get("contract_start_col") or config.get("start_col") or "")
    end_col = str(config.get("contract_end_col") or config.get("end_col") or "")
    entity_col = str(config.get("entity_col") or "")
    product_col = str(config.get("product_col") or "")
    customer_col = str(config.get("customer_col") or "")
    segment_col = str(config.get("segment_col") or "")
    region_col = str(config.get("region_col") or "")
    quantity_col = str(config.get("quantity_col") or "")
    apply_fx = bool(config.get("fx_enabled", config.get("apply_fx", False)))
    fx_col = str(config.get("fx_col") or "")
    from revenue_reconciliation import sales_basis_labels

    labels = sales_basis_labels(config.get("sales_basis"))
    sales_basis = labels.sales_basis

    common = {
        "file_path": file_path,
        "sheet_name": sheet_name,
        "output_file_path": str(config["output_file_path"]),
        "output_workbook_path": str(config["output_path"]),
        "use_session_workbook": True,
        "case_id": str(config.get("case_id") or "fast_track"),
        "title": str(config.get("project_name") or "Fast Track"),
        "company": str(config.get("company_name") or ""),
        "first_fy": config.get("first_fy"),
        "current_year": config.get("current_year"),
        "current_month": config.get("current_month"),
        "as_of_year": config.get("current_year"),
        "as_of_month": config.get("current_month"),
        "ltm_month": config.get("ltm_month", ""),
        "fy_end_month": config.get("fy_end_month", 12),
        "fy_end_day": config.get("fy_end_day", 31),
        "fiscal_year_end_month": config.get("fy_end_month", 12),
        "fiscal_year_end_day": config.get("fy_end_day", 31),
        "apply_fx": apply_fx,
        "fx_col": fx_col,
        "filters": config.get("filters") or {},
        "formula_mode": mapping_mode == "date",
        "sales_basis": sales_basis,
    }
    scripts: list[tuple[str, dict[str, Any]]] = []
    if "gst" in selected:
        scripts.append(
            (
                "general_sales_table_MM_verformelt.py",
                {
                    **common,
                    "base_sheet_name": "General sales table",
                    "table": "General sales table",
                    "calc_mode": calc_mode,
                    "invoice_mapping_mode": mapping_mode,
                    "invoice_col": invoice_col,
                    "start_col": start_col,
                    "end_col": end_col,
                    "group_cols": [entity_col, product_col, customer_col],
                    "value_cols": {
                        "revenue": revenue_col,
                        profit_mode: profit_col if profit_mode == "profit" else cost_col,
                    },
                    "profit_mode": profit_mode,
                    "show_revenue": True,
                    "show_gp": True,
                    "show_gm": True,
                    # YTD needs date-based invoice mapping (not period/year columns).
                    "show_ytd": mapping_mode == "date",
                    "show_ytd_delta": mapping_mode == "date",
                    "sort_mode": "YTD_LAST" if mapping_mode == "date" else "FY_LAST",
                    "total_label": labels.total_label,
                    "hierarchy_limits": {
                        "level_2": {"max_items": 4, "other_label": "Other"},
                        "level_3": {"max_items": 3, "other_label": "Other"},
                    },
                    "excel_formatting": {"column_widths": {"fy": 6.5}},
                },
            )
        )
    if "pvm" in selected:
        scripts.append(
            (
                "pvm_verformelt.py",
                {
                    **common,
                    "base_sheet_name": "pvm",
                    "table": "PVM",
                    "period_mode": "FY",
                    "calc_mode": calc_mode,
                    "invoice_mapping_mode": mapping_mode,
                    "invoice_col": invoice_col,
                    "start_col": start_col,
                    "end_col": end_col,
                    "group_col": product_col,
                    "quantity_col": quantity_col,
                    "revenue_col": revenue_col,
                    "profit_mode": profit_mode,
                    "cost_col": cost_col,
                    "profit_col": profit_col,
                    "pvm_method": str(config.get("pvm_method") or "three_components"),
                    "total_label_revenue": labels.total_label,
                },
            )
        )
    if "top" in selected:
        scripts.append(
            (
                "top_report.py",
                {
                    **common,
                    "base_sheet_name": "TOP",
                    "calc_mode": calc_mode,
                    "invoice_mapping_mode": mapping_mode,
                    "invoice_col": invoice_col,
                    "start_col": start_col,
                    "end_col": end_col,
                    "top_col": customer_col,
                    "value_col": revenue_col,
                    "total_label": labels.total_label,
                    "top_bucket": {
                        "enabled": True,
                        "thresholds": [0.2, 0.5, 0.8],
                        "include_other": True,
                    },
                },
            )
        )
    if "churn" in selected:
        if not region_col:
            raise ValueError("Churn analysis requires region_col")
        if not entity_col:
            raise ValueError("Churn analysis requires entity_col")
        scripts.append(
            (
                "churn_verformelt.py",
                {
                    **common,
                    "base_sheet_name": "Churn",
                    "table": "Churn",
                    "value_col": revenue_col,
                    "customer_col": customer_col,
                    "product_col": product_col,
                    "start_col": start_col,
                    "end_col": end_col,
                    "invoice_col": invoice_col,
                    "invoice_mapping_mode": mapping_mode,
                    "group_cols": [entity_col, region_col],
                    "period_mode": "FY",
                    "formula_mode": True,
                    "total_label": labels.total_label,
                },
            )
        )
    if "bars" in selected:
        from vertical_bars import default_bars_config

        bars = default_bars_config(
            entity_col=entity_col,
            segment_col=segment_col,
            customer_col=customer_col,
            product_col=product_col,
            region_col=region_col,
        )
        if not bars:
            raise ValueError(
                "Bars analysis requires at least one of entity/segment/customer/product/region columns"
            )
        scripts.append(
            (
                "vertical_bars.py",
                {
                    **common,
                    "base_sheet_name": f"{labels.metric_label} breakdown",
                    "period_mode": "FY",
                    "calc_mode": calc_mode,
                    "invoice_mapping_mode": mapping_mode,
                    "invoice_col": invoice_col,
                    "start_col": start_col,
                    "end_col": end_col,
                    "value_col": revenue_col,
                    "table": f"{labels.metric_label} breakdown",
                    "subtitle_suffix": "Revenue share by dimension",
                    "bars": bars,
                    "rest_label": "Misc.",
                    "plot": {
                        "show_plot_title": False,
                        "figsize": (5.0, 5.48),
                        "dpi": 300,
                        "show_values": False,
                        "label_min_share": 0.05,
                        "segment_fontsize": 7,
                        "min_share_to_rest": 0.05,
                        "x_step": 0.32,
                        "xtick_rotation": 22,
                    },
                    "excel": {
                        "enabled": True,
                        "col_offset": 3,
                        "sheet_name": f"{labels.metric_label} breakdown",
                        "image_anchor_cell": "D6",
                        "set_column_widths": True,
                    "legend_anchor_cell": "D18",
                    "legend_scale": 0.35,
                    },
                },
            )
        )
    if "gpbridge" in selected:
        if not product_col:
            raise ValueError("GP Bridge analysis requires product_col")
        if profit_mode == "profit" and not profit_col:
            raise ValueError("GP Bridge analysis requires profit_col when margin_mode=profit")
        if profit_mode != "profit" and not cost_col:
            raise ValueError("GP Bridge analysis requires cost_col when margin_mode=cost")
        sheet_bridge = labels.bridge_sheet_title
        table_bridge = labels.bridge_table_title
        period_mode = str(config.get("period_mode") or "FY").strip().upper() or "FY"
        if period_mode not in {"FY", "YTD", "LTM"}:
            period_mode = "FY"
        scripts.append(
            (
                "GP_Bridge.py",
                {
                    **common,
                    "base_sheet_name": sheet_bridge,
                    "period_mode": period_mode,
                    "calc_mode": calc_mode,
                    "invoice_mapping_mode": mapping_mode,
                    "invoice_col": invoice_col,
                    "start_col": start_col,
                    "end_col": end_col,
                    "revenue_col": revenue_col,
                    "profit_mode": profit_mode,
                    "cost_col": cost_col,
                    "profit_col": profit_col,
                    "table": table_bridge,
                    "subtitle_suffix": (
                        "by accrued amounts" if calc_mode == "accrual" else "by invoiced amounts"
                    ),
                    "plot": {
                        "xtick_rotation": 45,
                    },
                    "bridge": {
                        "dim_col": product_col,
                        "top_n": 3,
                        "other_label": "Other",
                        "missing_token": "__MISSING__",
                        "combine_input_other": True,
                        "input_other_tokens": {"Other"},
                        "other_bucket_label": "Other",
                        "sort_mode": "value",
                    },
                    "excel": {
                        "enabled": True,
                        "col_offset": 3,
                        "sheet_name": sheet_bridge,
                        "image_anchor_cell": "D6",
                        "set_column_widths": True,
                        "display_dpi": 180,
                        "image_scale": 0.5,
                    },
                },
            )
        )
    if "hbar" in selected:
        if not entity_col:
            raise ValueError("Horizontal bars analysis requires entity_col")
        if not segment_col:
            raise ValueError("Horizontal bars analysis requires segment_col")
        if not product_col:
            raise ValueError("Horizontal bars analysis requires product_col")
        sheet_hbar = f"{labels.metric_label} hierarchy"
        subtitle_hbar = (
            "by accrued amounts" if calc_mode == "accrual" else "by invoiced amounts"
        )
        scripts.append(
            (
                "horizontal_bars_hierachy.py",
                {
                    **common,
                    "base_sheet_name": sheet_hbar,
                    "period_mode": "FY",
                    "calc_mode": calc_mode,
                    "invoice_mapping_mode": mapping_mode,
                    "invoice_col": invoice_col,
                    "start_col": start_col,
                    "end_col": end_col,
                    "value_col": revenue_col,
                    "table": sheet_hbar,
                    "subtitle_suffix": subtitle_hbar,
                    "dimensions": {
                        "bar_col": entity_col,
                        "parent_col": segment_col,
                        "child_col": product_col,
                    },
                    "bars": {
                        "top_n": 5,
                        "create_other_bar": True,
                        "other_label": "Rest",
                        "bar_order": "value_desc",
                    },
                    "excel": {
                        "enabled": True,
                        "col_offset": 3,
                        "sheet_name": sheet_hbar,
                        "set_column_widths": True,
                        "display_dpi": 180,
                        "image_scale": 0.5,
                        "legend_scale": 0.45,
                    },
                },
            )
        )
    if "arrbridge" in selected:
        if not entity_col:
            raise ValueError("ARR Bridge analysis requires entity_col")
        if not product_col:
            raise ValueError("ARR Bridge analysis requires product_col")
        if not customer_col:
            raise ValueError("ARR Bridge analysis requires customer_col")
        if not start_col or not end_col:
            raise ValueError("ARR Bridge analysis requires contract start/end columns")
        use_arr = bool(config.get("use_arr_calculation", False))
        if use_arr and mapping_mode == "year":
            raise ValueError(
                "ARR calculation requires date mapping (invoice dates), not a period column"
            )
        arr_calc_mode = "ARR" if use_arr else calc_mode
        sheet_arr = "ARR Bridge"
        if use_arr:
            table_arr = "ARR Component Bridge"
            subtitle_arr = "ARR run-rate"
        else:
            table_arr = f"{labels.metric_label} component bridge"
            subtitle_arr = (
                "by accrued amounts" if calc_mode == "accrual" else "by invoiced amounts"
            )
        scripts.append(
            (
                "ARR_Bridge.py",
                {
                    **common,
                    "base_sheet_name": sheet_arr,
                    "period_mode": "FY",
                    "calc_mode": arr_calc_mode,
                    "invoice_mapping_mode": mapping_mode,
                    "invoice_col": invoice_col,
                    "start_col": start_col,
                    "end_col": end_col,
                    "value_col": revenue_col,
                    "customer_col": customer_col,
                    "product_col": product_col,
                    "group_col": entity_col,
                    "group_value_selection": "Total",
                    "table": table_arr,
                    "subtitle_suffix": subtitle_arr,
                    "recurring_filter": {"enabled": False},
                    "plot": {
                        "xtick_rotation": 45,
                        "bars_per_inch": 1.925,
                        "fixed_height_in": 1.5,
                    },
                    "bridge": {
                        "order": ["Upsell", "Downsell", "Cross-sell", "Lost", "New"],
                        "top_n_per_step": 3,
                        "missing_token": "__MISSING__",
                        "combine_input_other": True,
                        "input_other_tokens": {"Other"},
                        "other_bucket_label": "Other",
                        "sort_mode": "value",
                    },
                    "excel": {
                        "enabled": True,
                        "col_offset": 3,
                        "sheet_name": sheet_arr,
                        "image_anchor_cell": "D6",
                        "set_column_widths": True,
                        "display_dpi": 180,
                        "image_scale": 0.5,
                    },
                },
            )
        )
    return scripts


def _ensure_revenue_source_sheet(wb, config: dict[str, Any]) -> None:
    """Populate the hidden source sheet in the shared workbook before legacy revenue scripts."""
    from funktionssammlung import ensure_source_sheet_in_workbook

    file_path = str(config.get("file_path") or "").strip()
    sheet_name = str(config.get("sheet_name") or "").strip()
    if not file_path or not sheet_name:
        return
    ensure_source_sheet_in_workbook(wb, config)


def _run_revenue(config: dict[str, Any], wb) -> None:
    _register_revenue_runners()
    scripts = _revenue_script_configs(config)
    for script_name, script_config in scripts:
        runner = REVENUE_IN_MEMORY_RUNNERS.get(script_name)
        if runner is not None:
            runner(script_config, wb)
        else:
            _run_legacy_script(script_name, script_config)

    if "bubble" not in {
        str(value).strip().lower() for value in config.get("selected_analyses", [])
    }:
        return
    bubble_sheet = _resolve_sheet_name(config["file_path"], config.get("sheet_name"))
    df = pd.read_excel(config["file_path"], sheet_name=bubble_sheet, engine="openpyxl")
    from bubblescatterplot import run_bubble_chart

    from revenue_reconciliation import sales_basis_labels

    bubble_labels = sales_basis_labels(config.get("sales_basis"))
    bubble = {
        **config,
        "title": str(config.get("project_name") or config.get("title") or "Fast Track"),
        "company": str(config.get("company_name") or config.get("company") or ""),
        "sales_basis": bubble_labels.sales_basis,
        "revenue_col": config["revenue_col"],
        "invoice_col": config.get("invoice_col") or config.get("period_col"),
        "invoice_mapping_mode": (
            "year" if str(config.get("date_mapping") or "date") == "period" else "date"
        ),
        "start_col": config.get("contract_start_col"),
        "end_col": config.get("contract_end_col"),
        "profit_mode": config.get("margin_mode", "cost"),
        "cogs_col": config.get("cost_col"),
        "profit_col": config.get("profit_col"),
        "bubblesize_col": config["revenue_col"],
        "group_cols": [
            value
            for value in (config.get("segment_col"), config.get("product_col"))
            if value
        ],
        "period_mode": "FY",
        "table": "Margin analyses",
        "subtitle_suffix": f"{bubble_labels.profit_abbrev} by segment",
        "plot": {
            "x_label": f"{bubble_labels.margin_abbrev} (in%)",
            "y_label": bubble_labels.profit_abbrev,
            "legend": {
                "cluster_box_pad_px": 2,
                "cluster_gap_px": 8,
                "cluster_top_pad_px": 2,
            },
        },
        "excel": {
            "enabled": True,
            "sheet_name": "Margin analyses",
            "image_anchor_cell": "D6",
            "legend_anchor_cell": "D28",
            "legend_scale": 0.5,
        },
    }
    run_bubble_chart(df, bubble, wb=wb, save=False)


@contextmanager
def _legacy_workbook_adapter(wb, *redirect_paths: str):
    """Route legacy script workbook opens/saves/closes to the shared in-memory workbook."""
    import openpyxl
    from openpyxl.workbook.workbook import Workbook as OpenpyxlWorkbook

    original_save = OpenpyxlWorkbook.save
    original_close = OpenpyxlWorkbook.close
    targets: set[Path] = set()
    for raw in redirect_paths:
        if not raw:
            continue
        try:
            targets.add(Path(raw).expanduser().resolve())
        except (TypeError, ValueError, OSError):
            continue

    original_load = openpyxl.load_workbook

    def load_workbook(filename, *args, **kwargs):
        try:
            if Path(filename).expanduser().resolve() in targets:
                return wb
        except TypeError:
            pass
        return original_load(filename, *args, **kwargs)

    def save(workbook, filename):
        try:
            resolved = Path(filename).expanduser().resolve()
            if resolved in targets:
                if workbook is wb:
                    return None
                raise RuntimeError(
                    "Fast Track blocked save to the shared session workbook path from a "
                    f"secondary workbook ({filename}). All sections must use the batch workbook."
                )
        except TypeError:
            pass
        return original_save(workbook, filename)

    def close(workbook):
        if workbook is wb:
            return None
        return original_close(workbook)

    load_patches: list[tuple[Any, str, Any]] = [("openpyxl", "load_workbook", original_load)]
    openpyxl.load_workbook = load_workbook
    for name, mod in list(sys.modules.items()):
        if mod is None or not hasattr(mod, "load_workbook"):
            continue
        current = getattr(mod, "load_workbook")
        if current is load_workbook or current is original_load:
            continue
        load_patches.append((name, "load_workbook", current))
        setattr(mod, "load_workbook", load_workbook)

    OpenpyxlWorkbook.save = save
    OpenpyxlWorkbook.close = close
    try:
        yield
    finally:
        for name, attr, original in load_patches:
            mod = sys.modules.get(name)
            if mod is not None and hasattr(mod, attr):
                setattr(mod, attr, original)
        OpenpyxlWorkbook.close = original_close
        OpenpyxlWorkbook.save = original_save


def _default_databook_steps(
    config: dict[str, Any],
    master_path: str,
    entity_order: list[str],
    *,
    session_output_path: str | None = None,
):
    from databook_workbook import BS_RECON_MAPPING_FILE, CF_NA_L3_ORDER_FILE, PL_RECON_MAPPING_FILE

    project = str(config.get("project_name") or "Project")
    company = str(config.get("company_name") or "Group")
    basis = str(config.get("l4_sort_basis") or "latest_fy")
    session_target = str(session_output_path or master_path)
    shared = {
        "project_name": project,
        "company_name": company,
        "entity_order": entity_order,
        "l4_sort_basis": basis,
        "preserve_master_sheets": True,
        "use_session_workbook": True,
    }
    paths = {
        "mapping_source": "db",
        "master_file": master_path,
        "input_file": master_path,
        "target_file": session_target,
    }
    steps: list[tuple[str, dict[str, Any]]] = [
        (
            "Recon_tables_PL_MM.py",
            {
                **shared,
                "sort_by": "custom",
                "show_fs_check": False,
                "fs_check_values": {"entities": [], "consolidation": []},
                "display": {"titles": {"ic_display_name": "IC eliminations"}},
                "pl_config": {"display_label_map": {}},
                "paths": {
                    **paths,
                    "source_file": master_path,
                    "source_sheet": "Master_PL",
                    "target_file": session_target,
                    "report_sheet": "PL_Reconciliation",
                    "audit_master_sheet": "Master_PL",
                    "mapping_file": PL_RECON_MAPPING_FILE,
                    "append_to_master": False,
                },
            },
        ),
        (
            "Recon_tables_BS.py",
            {
                **shared,
                "paths": {
                    **paths,
                    "source_file": master_path,
                    "target_file": session_target,
                    "report_sheet": "BS_Reconciliation",
                    "mapping_file": BS_RECON_MAPPING_FILE,
                    "append_to_master": False,
                },
            },
        ),
    ]
    if bool(config.get("extended_tables", True)):
        for script in ("BS_Bucket.py", "Lead_IS.py", "Lead_BS.py", "Working_capital.py", "cashflow.py"):
            steps.append(
                (
                    script,
                    {
                        **shared,
                        "paths": {
                            **paths,
                            "mapping_file": BS_RECON_MAPPING_FILE,
                            "pl_mapping_file": PL_RECON_MAPPING_FILE,
                            "cf_na_l3_order_file": CF_NA_L3_ORDER_FILE,
                            "pl_master_file": master_path,
                        },
                    },
                )
            )
    return steps


def _run_legacy_script(script_name: str, config: dict[str, Any]) -> None:
    script = ROOT / script_name
    if not script.is_file():
        raise FileNotFoundError(f"Databook script not found: {script}")
    with tempfile.TemporaryDirectory(prefix="fast-track-config-") as tmp:
        config_path = Path(tmp) / "config.json"
        config_path.write_text(json.dumps(config, ensure_ascii=False), encoding="utf-8")
        old_argv = sys.argv
        try:
            sys.argv = [str(script), str(config_path)]
            runpy.run_path(str(script), run_name="__main__")
        finally:
            sys.argv = old_argv


def _run_databook(config: dict[str, Any], wb) -> None:
    master_path = str(config["master_path"])
    session_output_path = str(config.get("output_path") or master_path)
    entity_order = [str(v).strip() for v in config.get("entity_order", []) if str(v).strip()]
    if not entity_order:
        entity_order = derive_entity_order_by_latest_revenue(wb)
    config["entity_order"] = entity_order
    steps = config.get("steps")
    if steps is None:
        steps = _default_databook_steps(
            config,
            master_path,
            entity_order,
            session_output_path=session_output_path,
        )
    for item in steps:
        if isinstance(item, dict):
            script_name = str(item.get("script") or "")
            step_config = dict(item.get("config") or {})
        else:
            script_name, step_config = item
        step_config.setdefault("use_session_workbook", True)
        _run_legacy_script(script_name, step_config)


SECTION_RUNNERS: dict[str, SectionRunner] = {
    "databook": _run_databook,
    "opos": _run_opos,
    "fte": _run_fte,
    "fa": _run_fa,
    "revenue": _run_revenue,
}

REVENUE_IN_MEMORY_RUNNERS: dict[str, Callable[[dict[str, Any], Any], str]] = {}


def _register_revenue_runners() -> None:
    if REVENUE_IN_MEMORY_RUNNERS:
        return
    from general_sales_table_MM_verformelt import run_hierarchy_report_in_workbook
    from pvm_verformelt import run_pvm_in_workbook
    from top_report import run_top_report_in_workbook
    from churn_verformelt import run_churn_in_workbook
    from vertical_bars import run_vertical_bars_in_workbook
    from GP_Bridge import run_gp_bridge_in_workbook
    from horizontal_bars_hierachy import run_hbar_breakdown_in_workbook
    from ARR_Bridge import run_arr_bridge_in_workbook

    REVENUE_IN_MEMORY_RUNNERS.update(
        {
            "general_sales_table_MM_verformelt.py": run_hierarchy_report_in_workbook,
            "pvm_verformelt.py": run_pvm_in_workbook,
            "top_report.py": run_top_report_in_workbook,
            "churn_verformelt.py": run_churn_in_workbook,
            "vertical_bars.py": run_vertical_bars_in_workbook,
            "GP_Bridge.py": run_gp_bridge_in_workbook,
            "horizontal_bars_hierachy.py": run_hbar_breakdown_in_workbook,
            "ARR_Bridge.py": run_arr_bridge_in_workbook,
        }
    )


def _remove_invalid_output_placeholder(path: Path) -> None:
    """Delete empty or corrupt placeholder files left from earlier Fast Track runs."""
    if not path.is_file():
        return
    try:
        if path.stat().st_size == 0:
            path.unlink()
            return
    except OSError:
        return
    import zipfile

    if not zipfile.is_zipfile(path):
        try:
            path.unlink()
        except OSError:
            pass


def _session_redirect_paths(normalized: dict[str, Any], output_path: Path) -> list[str]:
    paths: list[str] = [str(output_path)]
    master = str(normalized.get("master_path") or "").strip()
    if master:
        paths.append(master)
    output_dir = output_path.parent
    if output_dir.is_dir():
        for pattern in ("*_Master.xlsx", "*_Workbook.xlsx", "*_FastTrack.xlsx", "*_SuSa_Master.xlsx"):
            paths.extend(str(p) for p in output_dir.glob(pattern))
    # De-dupe while preserving order
    seen: set[str] = set()
    unique: list[str] = []
    for raw in paths:
        try:
            key = str(Path(raw).expanduser().resolve())
        except (TypeError, ValueError, OSError):
            key = raw
        if key in seen:
            continue
        seen.add(key)
        unique.append(raw)
    return unique


def run_fast_track(
    manifest: dict[str, Any],
    *,
    section_runners: dict[str, SectionRunner] | None = None,
) -> dict[str, Any]:
    normalized = validate_manifest(manifest)
    wb = open_batch_workbook(normalized["master_path"] or None)
    runners = section_runners or SECTION_RUNNERS
    completed: list[str] = []
    output_path = Path(normalized["output_path"])
    output_path.parent.mkdir(parents=True, exist_ok=True)
    _remove_invalid_output_placeholder(output_path)
    redirect_paths = _session_redirect_paths(normalized, output_path)
    from funktionssammlung import session_workbook_context

    try:
        with _legacy_workbook_adapter(wb, *redirect_paths), session_workbook_context(wb):
            for name in SECTION_ORDER:
                section = normalized["sections"][name]
                if not section["enabled"]:
                    continue
                config = dict(section["config"])
                dates = normalized.get("dates") if isinstance(normalized.get("dates"), dict) else {}
                project = normalized.get("project") if isinstance(normalized.get("project"), dict) else {}
                config.setdefault("output_path", str(output_path))
                config.setdefault("output_file_path", str(output_path.parent))
                config.setdefault("output_workbook_path", str(output_path))
                config.setdefault("use_session_workbook", True)
                config.setdefault("case_id", str(normalized.get("case_id") or "fast_track"))
                config.setdefault("project_name", str(project.get("project_name") or "Fast Track"))
                config.setdefault("company_name", str(project.get("group_name") or ""))
                config.setdefault("first_fy", dates.get("first_fy"))
                config.setdefault("fy_end_month", dates.get("fy_end_month", 12))
                config.setdefault("fy_end_day", dates.get("fy_end_day", 31))
                _apply_manifest_dates(config, dates)
                if name == "databook":
                    config["master_path"] = normalized["master_path"]
                    config.setdefault("l4_sort_basis", "latest_fy")
                    config.setdefault(
                        "extended_tables",
                        bool(config.get("create_extended_tables", True)),
                    )
                else:
                    config.setdefault("master_path", normalized.get("master_path") or "")
                runner = runners.get(name)
                if runner is None:
                    raise ValueError(f"No in-memory Fast Track adapter registered for {name}")
                runner(config, wb)
                completed.append(name)
            remove_placeholder_fast_track_sheet(wb)
            reorder_workbook_sheets(wb)
            output = atomic_save_workbook(wb, normalized["output_path"])
    finally:
        wb.close()
    summary = {
        "success": True,
        "output_file": str(output),
        "output_filename": output.name,
        "completed_sections": completed,
        "sheet_names": list(wb.sheetnames),
        "master_path": normalized.get("master_path") or "",
    }
    print(json.dumps(summary, ensure_ascii=False))
    return summary


def main() -> None:
    if len(sys.argv) != 2:
        raise SystemExit("Usage: python scripts/fast_track_batch.py <manifest.json>")
    manifest = json.loads(Path(sys.argv[1]).read_text(encoding="utf-8-sig"))
    result = run_fast_track(manifest)
    print(json.dumps(result, ensure_ascii=False))


if __name__ == "__main__":
    main()
