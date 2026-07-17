"""Assemble Fast Track PDF from an ordered workbook (reportlab)."""

from __future__ import annotations

import os
import re
import tempfile
from datetime import date
from pathlib import Path
from typing import Any, Optional

from openpyxl import load_workbook
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas

from pdf_report.chart_extract import extract_sheet_images_meta
from pdf_report.sheet_render import RENDER_SCALE, render_sheet_screenshots_paged
from pdf_report.theme import THEME, hex_to_rgb

try:
    from databook_workbook import (
        GROUP_BS_RECON_SHEET,
        GROUP_PL_RECON_SHEET,
        SHEET_GROUP_SPECS,
        _classify_sheet_group,
        _earnings_sort_key,
        _financial_sort_key,
        collect_entity_recon_sheets_ordered,
        entity_order_from_group_recon,
        is_entity_recon_sheet,
        is_fte_output_sheet,
        is_revenue_graph_sheet,
        is_revenue_table_sheet,
        is_section_sheet,
        is_vertical_bars_sheet,
        is_workbook_source_sheet,
    )
except ImportError:  # pragma: no cover
    import sys

    root = Path(__file__).resolve().parents[2]
    if str(root) not in sys.path:
        sys.path.insert(0, str(root))
    from databook_workbook import (
        GROUP_BS_RECON_SHEET,
        GROUP_PL_RECON_SHEET,
        SHEET_GROUP_SPECS,
        _classify_sheet_group,
        _earnings_sort_key,
        _financial_sort_key,
        collect_entity_recon_sheets_ordered,
        entity_order_from_group_recon,
        is_entity_recon_sheet,
        is_fte_output_sheet,
        is_revenue_graph_sheet,
        is_revenue_table_sheet,
        is_section_sheet,
        is_vertical_bars_sheet,
        is_workbook_source_sheet,
    )


PDF_GROUP_SPECS = tuple(s for s in SHEET_GROUP_SPECS if s["key"] != "source")

SECTION_NUMBERS = {
    "executive": 1,
    "earnings": 2,
    "financial": 3,
    "liquidity": 4,
    "appendix": 5,
}

_SHEET_DISPLAY_NAMES = {
    "lead_is": "Lead Income statement",
    "lead_bs": "Lead Balance sheet",
    "bs_bucket": "BS Bucket",
    "bucket_bs": "BS Bucket",
    "general sales table": "General sales table",
    "pvm": "Price-Volume-Mix",
    "top": "Top customers",
    "churn": "Churn",
    "margin analyses": "Margin analyses",
    "bubble": "Margin analyses",
    "gp bridge": "GP Bridge",
    "np bridge": "NP Bridge",
    "arr bridge": "ARR Bridge",
    "fte development": "Payroll accounting",
    "cashflow": "Cashflow",
    "working_capital": "Working Capital",
    "working capital": "Working Capital",
    "bs_reconciliation": "Group BS Reconciliation",
    "pl_reconciliation": "Group PL Reconciliation",
    "net sales breakdown": "Net sales breakdown",
    "gross sales breakdown": "Gross sales breakdown",
    "net sales hierarchy": "Net sales hierarchy",
    "gross sales hierarchy": "Gross sales hierarchy",
    "trade debtors aging": "Trade debtors aging",
    "trade creditors aging": "Trade creditors aging",
    "trade debtors aging detail": "Trade debtors aging detail",
    "trade creditors aging detail": "Trade creditors aging detail",
    "fa roll forward": "FA roll forward",
}

_PDF_SKIP_SHEETS = frozenset(
    {
        "working_capital",
        "working capital",
    }
)


def _content_box() -> tuple[float, float, float, float]:
    """Full-width content box: x, width, top_y, bottom_y."""
    m = THEME.margin
    return m, THEME.page_w - 2 * m, THEME.page_h - m - 18, m + 22


def _to_roman(n: int) -> str:
    if n <= 0:
        return str(n)
    vals = (
        (100, "C"),
        (90, "XC"),
        (50, "L"),
        (40, "XL"),
        (10, "X"),
        (9, "IX"),
        (5, "V"),
        (4, "IV"),
        (1, "I"),
    )
    out: list[str] = []
    remaining = int(n)
    for value, glyph in vals:
        while remaining >= value:
            out.append(glyph)
            remaining -= value
    return "".join(out)


def display_sheet_name(sheet_name: str) -> str:
    """Human-readable sheet title for PDF chrome / outline."""
    raw = str(sheet_name or "").strip()
    if not raw:
        return "Sheet"
    key = raw.lower()
    if key in _SHEET_DISPLAY_NAMES:
        return _SHEET_DISPLAY_NAMES[key]
    if is_entity_recon_sheet(raw):
        return raw.replace("_", " ")
    if raw in {GROUP_BS_RECON_SHEET, GROUP_PL_RECON_SHEET}:
        return _SHEET_DISPLAY_NAMES.get(key, raw.replace("_", " "))
    token = re.sub(r"_\d+$", "", raw).strip()
    mapped = _SHEET_DISPLAY_NAMES.get(token.lower())
    if mapped:
        return mapped
    return raw.replace("_", " ")


def _sort_appendix_sheets(names: list[str], wb) -> list[str]:
    """Group recons first, then entity recons, then other appendix sheets."""
    group = [n for n in (GROUP_BS_RECON_SHEET, GROUP_PL_RECON_SHEET) if n in names]
    try:
        entity_order = entity_order_from_group_recon(wb)
    except Exception:
        entity_order = []
    entities = collect_entity_recon_sheets_ordered(names, entity_order=entity_order or None)
    used = set(group) | set(entities)
    other = sorted([n for n in names if n not in used], key=str.lower)
    return group + entities + other


def _build_section_outline(section_key: str, section_no: int, sheets: list[str]) -> list[dict[str, Any]]:
    """
    Build numbered outline entries for a section divider / sheet mapping.
    Returns list of {sheet_name?, code, title, indent, kind}.
    kind: group (subsection divider only) | item (subsection + sheet) | roman (sheet under Sales)
    """
    entries: list[dict[str, Any]] = []
    if section_key == "earnings":
        lead = [n for n in sheets if n.lower() == "lead_is"]
        sales = [n for n in sheets if is_revenue_table_sheet(n)]
        graphs = [n for n in sheets if is_revenue_graph_sheet(n)]
        fte = [n for n in sheets if is_fte_output_sheet(n)]
        rest = [
            n
            for n in sheets
            if n not in lead and n not in sales and n not in graphs and n not in fte
        ]

        sub_i = 1
        if lead:
            for n in lead:
                entries.append(
                    {
                        "sheet_name": n,
                        "code": f"{section_no}.{sub_i}",
                        "title": display_sheet_name(n),
                        "indent": 0,
                        "kind": "item",
                    }
                )
            sub_i += 1

        if sales or graphs:
            entries.append(
                {
                    "sheet_name": None,
                    "code": f"{section_no}.{sub_i}",
                    "title": "Sales",
                    "indent": 0,
                    "kind": "group",
                }
            )
            sales_group_code = f"{section_no}.{sub_i}"
            for ri, n in enumerate(sales + graphs, start=1):
                entries.append(
                    {
                        "sheet_name": n,
                        "code": _to_roman(ri),
                        "parent_code": sales_group_code,
                        "title": display_sheet_name(n),
                        "indent": 1,
                        "kind": "roman",
                    }
                )
            sub_i += 1

        for n in fte + rest:
            entries.append(
                {
                    "sheet_name": n,
                    "code": f"{section_no}.{sub_i}",
                    "title": display_sheet_name(n),
                    "indent": 0,
                    "kind": "item",
                }
            )
            sub_i += 1
        return entries

    if section_key == "appendix":
        group_sheets = [n for n in sheets if n in {GROUP_BS_RECON_SHEET, GROUP_PL_RECON_SHEET}]
        entity_sheets = [n for n in sheets if is_entity_recon_sheet(n)]
        other = [n for n in sheets if n not in group_sheets and n not in entity_sheets]
        sub_i = 1
        if group_sheets:
            entries.append(
                {
                    "sheet_name": None,
                    "code": f"{section_no}.{sub_i}",
                    "title": "Group reconciliations",
                    "indent": 0,
                    "kind": "group",
                }
            )
            for n in group_sheets:
                entries.append(
                    {
                        "sheet_name": n,
                        "code": "",
                        "title": display_sheet_name(n),
                        "indent": 1,
                        "kind": "roman",
                    }
                )
            sub_i += 1
        if entity_sheets:
            entries.append(
                {
                    "sheet_name": None,
                    "code": f"{section_no}.{sub_i}",
                    "title": "Entity reconciliations",
                    "indent": 0,
                    "kind": "group",
                }
            )
            for n in entity_sheets:
                entries.append(
                    {
                        "sheet_name": n,
                        "code": "",
                        "title": display_sheet_name(n),
                        "indent": 1,
                        "kind": "roman",
                    }
                )
            sub_i += 1
        for n in other:
            entries.append(
                {
                    "sheet_name": n,
                    "code": f"{section_no}.{sub_i}",
                    "title": display_sheet_name(n),
                    "indent": 0,
                    "kind": "item",
                }
            )
            sub_i += 1
        return entries

    for i, n in enumerate(sheets, start=1):
        entries.append(
            {
                "sheet_name": n,
                "code": f"{section_no}.{i}",
                "title": display_sheet_name(n),
                "indent": 0,
                "kind": "item",
            }
        )
    return entries


def _is_skipped_pdf_sheet(name: str) -> bool:
    return str(name or "").strip().lower() in _PDF_SKIP_SHEETS


def build_workbook_toc(wb) -> list[dict[str, Any]]:
    buckets: dict[str, list[str]] = {spec["key"]: [] for spec in PDF_GROUP_SPECS}
    for name in wb.sheetnames:
        if is_section_sheet(name):
            continue
        if is_workbook_source_sheet(name):
            continue
        if _is_skipped_pdf_sheet(name):
            continue
        group = _classify_sheet_group(name)
        if group == "source" or group is None:
            continue
        if group in buckets:
            buckets[group].append(name)

    for key, names in buckets.items():
        if key == "earnings":
            buckets[key] = sorted(names, key=_earnings_sort_key)
        elif key == "financial":
            buckets[key] = sorted(names, key=_financial_sort_key)
        elif key == "appendix":
            buckets[key] = _sort_appendix_sheets(names, wb)
        elif key == "liquidity":
            buckets[key] = sorted(
                names, key=lambda n: (0 if n.lower() == "cashflow" else 1, n.lower())
            )
        else:
            buckets[key] = sorted(names, key=lambda n: n.lower())

    toc: list[dict[str, Any]] = []
    for spec in PDF_GROUP_SPECS:
        sheets = buckets.get(spec["key"]) or []
        if not sheets:
            continue
        section_no = SECTION_NUMBERS.get(spec["key"], len(toc) + 1)
        outline = _build_section_outline(spec["key"], section_no, sheets)
        toc.append(
            {
                "type": "section",
                "key": spec["key"],
                "label": spec["label"],
                "color": spec["color"],
                "section_no": section_no,
                "section_code": str(section_no),
                "sheet_names": list(sheets),
                "outline": outline,
            }
        )
        for entry in outline:
            kind = entry.get("kind") or "item"
            if kind == "group":
                toc.append(
                    {
                        "type": "subsection",
                        "key": spec["key"],
                        "label": spec["label"],
                        "color": spec["color"],
                        "section_no": section_no,
                        "outline_code": entry.get("code") or "",
                        "display_name": entry.get("title") or "",
                        "title": f"{entry.get('code') or ''}  {entry.get('title') or ''}".strip(),
                    }
                )
                continue
            sn = entry.get("sheet_name")
            if not sn:
                continue
            if kind == "item":
                toc.append(
                    {
                        "type": "subsection",
                        "key": spec["key"],
                        "label": spec["label"],
                        "color": spec["color"],
                        "section_no": section_no,
                        "outline_code": entry.get("code") or "",
                        "display_name": entry.get("title") or display_sheet_name(sn),
                        "title": f"{entry.get('code') or ''}  {entry.get('title') or display_sheet_name(sn)}".strip(),
                    }
                )
            toc.append(
                {
                    "type": "sheet",
                    "key": spec["key"],
                    "label": spec["label"],
                    "color": spec["color"],
                    "sheet_name": sn,
                    "display_name": entry.get("title") or display_sheet_name(sn),
                    "outline_code": entry.get("code") or "",
                    "parent_code": entry.get("parent_code") or "",
                    "section_no": section_no,
                }
            )
    return toc


def _clean_project_name(name: str) -> str:
    s = str(name or "").strip()
    if not s or s.lower() in {"project", "finssentials", "fast track"}:
        return "Project"
    return s


def build_fast_track_pdf_report(
    workbook_path: str | Path,
    output_path: str | Path,
    *,
    project_name: str = "Project",
    company_name: str = "",
    report_date: Optional[date] = None,
    recalc_excel: bool = True,
) -> str:
    workbook_path = Path(workbook_path).expanduser().resolve()
    output_path = Path(output_path).expanduser().resolve()
    output_path.parent.mkdir(parents=True, exist_ok=True)

    project_name = _clean_project_name(project_name)
    company_name = str(company_name or "").strip()
    report_date = report_date or date.today()

    source_wb = workbook_path
    if recalc_excel:
        try:
            from pdf_report.excel_recalc import materialize_calculated_copy

            source_wb = materialize_calculated_copy(workbook_path)
        except Exception as exc:
            print(f"WARN: could not materialize calculated workbook: {exc}")
            source_wb = workbook_path

    wb = load_workbook(source_wb, data_only=True)
    toc = build_workbook_toc(wb)

    img_dir = tempfile.mkdtemp(prefix="ft_pdf_assets_")
    c = canvas.Canvas(str(output_path), pagesize=(THEME.page_w, THEME.page_h))

    _draw_cover(
        c,
        project_name=project_name,
        company_name=company_name,
        report_date=report_date,
    )
    c.showPage()

    for item in toc:
        if item["type"] == "section":
            _draw_section_divider(
                c,
                label=item["label"],
                color=item["color"],
                section_no=item.get("section_no"),
                outline=item.get("outline") or [],
            )
            c.showPage()
            continue

        if item["type"] == "subsection":
            _draw_section_divider(
                c,
                label=item.get("title") or item.get("display_name") or item["label"],
                color=item["color"],
                section_no=None,
                outline=[],
                subtitle=item.get("label"),
            )
            c.showPage()
            continue

        sheet_name = item["sheet_name"]
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        display_name = item.get("display_name") or display_sheet_name(sheet_name)
        chart_meta = extract_sheet_images_meta(
            ws, out_dir=os.path.join(img_dir, "charts", _safe(sheet_name))
        )
        is_chart_sheet = bool(chart_meta) or is_revenue_graph_sheet(sheet_name) or is_vertical_bars_sheet(
            sheet_name
        )

        # Build slide heading: "{parent_code} {outline_code}  {display_name}"
        outline_code = item.get("outline_code") or ""
        parent_code = item.get("parent_code") or ""
        if parent_code and outline_code:
            slide_heading = f"{parent_code} {outline_code}  {display_name}"
        elif outline_code:
            slide_heading = f"{outline_code}  {display_name}"
        else:
            slide_heading = display_name
        # The small Excel table title is part of the rendered table image.
        # Page chrome contains only the slide heading.
        titles = [slide_heading]

        page_kwargs = dict(
            sheet_name=sheet_name,
            display_name=display_name,
            section_label=item["label"],
            section_color=item["color"],
            titles=titles,
        )

        if is_chart_sheet and chart_meta:
            _draw_chart_pages(c, chart_meta=chart_meta, **page_kwargs)
            continue

        shots = render_sheet_screenshots_paged(
            ws,
            out_dir=os.path.join(img_dir, "tables", _safe(sheet_name)),
            prefix=_safe(sheet_name),
            rows_per_page=200,
            max_cols=80,
            scale=RENDER_SCALE,
            prefer_single_page=True,
        )
        if not shots and chart_meta:
            _draw_chart_pages(c, chart_meta=chart_meta, **page_kwargs)
            continue
        if not shots:
            continue

        for idx, shot in enumerate(shots):
            page_titles = titles if idx == 0 else [f"{display_name} (continued)"]
            _draw_image_page(
                c,
                image_paths=[shot],
                stack_vertical=True,
                allow_upscale=False,
                sheet_name=sheet_name,
                display_name=display_name,
                section_label=item["label"],
                section_color=item["color"],
                titles=page_titles,
            )

    c.save()
    try:
        wb.close()
    except Exception:
        pass
    return str(output_path)


def _safe(name: str) -> str:
    return re.sub(r"[^\w\-]+", "_", str(name))[:50] or "sheet"


def _set_fill(c: canvas.Canvas, color: str) -> None:
    r, g, b = hex_to_rgb(color)
    c.setFillColorRGB(r, g, b)


def _set_stroke(c: canvas.Canvas, color: str) -> None:
    r, g, b = hex_to_rgb(color)
    c.setStrokeColorRGB(r, g, b)


def _draw_cover(
    c: canvas.Canvas,
    *,
    project_name: str,
    company_name: str,
    report_date: date,
) -> None:
    band_w = THEME.page_w * 0.38
    _set_fill(c, THEME.brand_navy)
    c.rect(0, 0, band_w, THEME.page_h, fill=1, stroke=0)

    c.setFillColorRGB(1, 1, 1)
    c.setFont(THEME.font_name_bold, 28)
    c.drawString(THEME.margin + 8, THEME.page_h - 90, "Finssentials")
    c.setFont(THEME.font_name, 11)
    c.drawString(THEME.margin + 8, THEME.page_h - 112, "Financial Due Diligence")

    x = band_w + 40
    y = THEME.page_h - 160
    _set_fill(c, THEME.brand_navy)
    c.setFont(THEME.font_name_bold, THEME.font_size_title)
    c.drawString(x, y, "Fast Track Report")
    _set_stroke(c, THEME.brand_navy)
    c.setLineWidth(1.2)
    c.line(x, y - 14, x + 220, y - 14)

    y -= 50
    _set_fill(c, THEME.text_muted)
    c.setFont(THEME.font_name, 10)
    c.drawString(x, y, "Project")
    y -= 18
    _set_fill(c, THEME.text_primary)
    c.setFont(THEME.font_name_bold, 16)
    c.drawString(x, y, project_name[:80])

    if company_name:
        y -= 36
        _set_fill(c, THEME.text_muted)
        c.setFont(THEME.font_name, 10)
        c.drawString(x, y, "Company")
        y -= 18
        _set_fill(c, THEME.text_primary)
        c.setFont(THEME.font_name_bold, 14)
        c.drawString(x, y, company_name[:80])

    y -= 48
    _set_fill(c, THEME.text_light)
    c.setFont(THEME.font_name, 9)
    c.drawString(x, y, report_date.strftime("%d %B %Y"))
    c.drawString(x, THEME.margin, "Confidential — for discussion purposes")


def _draw_section_divider(
    c: canvas.Canvas,
    *,
    label: str,
    color: str,
    section_no: int | None = None,
    outline: list[dict[str, Any]] | None = None,
    subtitle: str | None = None,
) -> None:
    _set_fill(c, f"#{str(color).lstrip('#')[-6:]}")
    c.rect(0, 0, THEME.page_w, THEME.page_h, fill=1, stroke=0)
    c.setFillColorRGB(1, 1, 1)
    c.setFont(THEME.font_name_bold, 18)
    c.drawString(THEME.margin + 10, THEME.page_h / 2 + 70, "Finssentials")
    if subtitle:
        c.setFont(THEME.font_name, 12)
        c.drawString(THEME.margin + 10, THEME.page_h / 2 + 48, str(subtitle))
    c.setFont(THEME.font_name_bold, THEME.font_size_section if section_no is not None else 24)
    heading = f"{section_no}  {label}" if section_no is not None else label
    c.drawString(THEME.margin + 10, THEME.page_h / 2 + 18, heading[:100])
    if outline:
        c.setFont(THEME.font_name, 11)
        c.drawString(THEME.margin + 10, THEME.page_h / 2 - 8, "Contents")
        y = THEME.page_h / 2 - 32
        c.setFont(THEME.font_name, 10)
        for entry in (outline or [])[:22]:
            indent = 18 + int(entry.get("indent") or 0) * 18
            code = str(entry.get("code") or "").strip()
            title = str(entry.get("title") or "").strip()
            line = f"{code}  {title}".strip() if code else title
            c.drawString(THEME.margin + indent, y, line[:100])
            y -= 14
            if y < THEME.margin + 40:
                break
        rest = max(0, len(outline or []) - 22)
        if rest:
            c.drawString(THEME.margin + 18, y, f"… and {rest} more")


def _draw_page_chrome(
    c: canvas.Canvas,
    *,
    section_label: str,
    section_color: str,
    display_name: str,
    titles: list[str],
) -> float:
    """Top color bar + heading; footer with section path and Finssentials © year | page."""
    x, w, top_y, _bottom = _content_box()
    _set_fill(c, f"#{str(section_color).lstrip('#')[-6:]}")
    c.rect(0, THEME.page_h - 8, THEME.page_w, 8, fill=1, stroke=0)

    y = top_y
    for i, line in enumerate(titles[:3]):
        _set_fill(c, THEME.brand_navy if i == 0 else THEME.text_muted)
        font = THEME.font_name_bold if i == 0 else THEME.font_name
        size = 14 if i == 0 else 9
        # Shrink font until the full title fits — never truncate with ellipsis
        while size >= 8 and c.stringWidth(line, font, size) > w:
            size -= 1
        c.setFont(font, size)
        c.drawString(x, y, line)
        y -= size + 6

    _set_fill(c, THEME.text_light)
    c.setFont(THEME.font_name, 8)
    left = f"{section_label}  ·  {display_name}"
    c.drawString(THEME.margin, THEME.margin - 4, left[:110])
    year = date.today().year
    page_no = c.getPageNumber()
    right = f"Finssentials © {year}  |  {page_no}"
    c.drawRightString(THEME.page_w - THEME.margin, THEME.margin - 4, right)
    return y - 4


def _native_pt_size(width_px: float, height_px: float, *, dpi: float = 144.0) -> tuple[float, float]:
    """Convert raster px to PDF points at a fixed DPI so fonts stay consistent."""
    return width_px * 72.0 / dpi, height_px * 72.0 / dpi


def _fit_images_on_page(
    image_paths: list[str],
    *,
    max_w: float,
    max_h: float,
    stack_vertical: bool,
    allow_upscale: bool = False,
) -> list[tuple[str, float, float]]:
    """Fit images without blowing up small legends/tables."""
    sized: list[tuple[str, float, float]] = []
    for path in image_paths:
        if not path or not os.path.isfile(path):
            continue
        try:
            img = ImageReader(path)
            iw, ih = img.getSize()
        except Exception:
            continue
        if iw <= 0 or ih <= 0:
            continue
        nw, nh = _native_pt_size(float(iw), float(ih))
        sized.append((path, nw, nh))

    if not sized:
        return []

    gap = 8.0
    if len(sized) == 1 or stack_vertical:
        out: list[tuple[str, float, float]] = []
        n = len(sized)
        avail_each = (max_h - gap * (n - 1)) / max(n, 1)
        for path, nw, nh in sized:
            scale = min(max_w / nw, avail_each / nh)
            if not allow_upscale:
                scale = min(scale, 1.0)
            out.append((path, nw * scale, nh * scale))
        total_h = sum(h for _, _, h in out) + gap * (len(out) - 1)
        if total_h > max_h and total_h > 0:
            f = max_h / total_h
            out = [(p, w * f, h * f) for p, w, h in out]
        return out

    cols = 2 if len(sized) <= 4 else 3
    rows = (len(sized) + cols - 1) // cols
    cell_w = (max_w - gap * (cols - 1)) / cols
    cell_h = (max_h - gap * (rows - 1)) / rows
    out = []
    for path, nw, nh in sized:
        scale = min(cell_w / nw, cell_h / nh)
        if not allow_upscale:
            scale = min(scale, 1.0)
        out.append((path, nw * scale, nh * scale))
    return out


def _draw_image_page(
    c: canvas.Canvas,
    *,
    sheet_name: str,
    section_label: str,
    section_color: str,
    titles: list[str],
    image_paths: list[str],
    display_name: str | None = None,
    stack_vertical: bool = False,
    allow_upscale: bool = False,
) -> None:
    y = _draw_page_chrome(
        c,
        section_label=section_label,
        section_color=section_color,
        display_name=display_name or display_sheet_name(sheet_name),
        titles=titles,
    )
    x, w, _top, bottom = _content_box()
    avail_h = y - bottom
    fitted = _fit_images_on_page(
        image_paths,
        max_w=w,
        max_h=avail_h,
        stack_vertical=stack_vertical or len(image_paths) <= 2,
        allow_upscale=allow_upscale,
    )
    if not fitted:
        c.showPage()
        return

    if stack_vertical or len(fitted) <= 2:
        cy = y
        gap = 8.0
        for path, dw, dh in fitted:
            try:
                c.drawImage(
                    ImageReader(path),
                    x,
                    cy - dh,
                    width=dw,
                    height=dh,
                    preserveAspectRatio=True,
                    mask="auto",
                )
            except Exception:
                pass
            cy -= dh + gap
    else:
        cols = 2 if len(fitted) <= 4 else 3
        gap = 8.0
        cell_w = (w - gap * (cols - 1)) / cols
        cx0 = x
        cy = y
        row_h = 0.0
        col_i = 0
        for path, dw, dh in fitted:
            if col_i == cols:
                cy -= row_h + gap
                row_h = 0.0
                col_i = 0
            px = cx0 + col_i * (cell_w + gap) + max(0, (cell_w - dw) / 2)
            try:
                c.drawImage(
                    ImageReader(path),
                    px,
                    cy - dh,
                    width=dw,
                    height=dh,
                    preserveAspectRatio=True,
                    mask="auto",
                )
            except Exception:
                pass
            row_h = max(row_h, dh)
            col_i += 1
    c.showPage()


def _draw_chart_pages(
    c: canvas.Canvas,
    *,
    sheet_name: str,
    section_label: str,
    section_color: str,
    titles: list[str],
    chart_meta: list,
    display_name: str | None = None,
) -> None:
    """
    Place charts in Excel period order (left→right by anchor col).
    Multi-period sheets: plots on one row, legends on the row below (same order).
    """
    if not chart_meta:
        return

    disp = display_name or display_sheet_name(sheet_name)
    rows = sorted({im.row for im in chart_meta})
    if len(rows) >= 2:
        plot_row = rows[0]
        plots = sorted([im for im in chart_meta if im.row == plot_row], key=lambda im: im.col)
        legends = sorted([im for im in chart_meta if im.row != plot_row], key=lambda im: im.col)
    else:
        by_size = sorted(chart_meta, key=lambda im: im.width_px * im.height_px, reverse=True)
        if len(by_size) >= 2:
            median = by_size[len(by_size) // 2].width_px * by_size[len(by_size) // 2].height_px
            plots = sorted(
                [im for im in chart_meta if im.width_px * im.height_px >= median * 0.55],
                key=lambda im: im.col,
            )
            plot_ids = {id(im) for im in plots}
            legends = sorted(
                [im for im in chart_meta if id(im) not in plot_ids],
                key=lambda im: im.col,
            )
        else:
            plots = sorted(chart_meta, key=lambda im: im.col)
            legends = []

    if len(plots) >= 2:
        y = _draw_page_chrome(
            c,
            section_label=section_label,
            section_color=section_color,
            display_name=disp,
            titles=titles,
        )
        x, w, _top, bottom = _content_box()
        avail_h = y - bottom
        gap = 8.0
        col_w = (w - gap * (len(plots) - 1)) / len(plots)
        plot_budget = avail_h * (0.72 if legends else 0.95)
        leg_budget = avail_h - plot_budget - gap if legends else 0.0

        plot_fitted = []
        for im in plots:
            try:
                iw, ih = ImageReader(im.path).getSize()
                nw, nh = _native_pt_size(float(iw), float(ih))
            except Exception:
                nw, nh = _native_pt_size(float(im.width_px), float(im.height_px))
            scale = min(col_w / nw, plot_budget / nh, 1.0)
            plot_fitted.append((im.path, nw * scale, nh * scale))

        cy = y
        max_plot_h = max((h for _, _, h in plot_fitted), default=0.0)
        for i, (path, dw, dh) in enumerate(plot_fitted):
            px = x + i * (col_w + gap) + max(0.0, (col_w - dw) / 2)
            try:
                c.drawImage(
                    ImageReader(path),
                    px,
                    cy - dh,
                    width=dw,
                    height=dh,
                    preserveAspectRatio=True,
                    mask="auto",
                )
            except Exception:
                pass
        cy -= max_plot_h + gap

        if legends and leg_budget > 20:
            for i, im in enumerate(legends[: len(plots)]):
                try:
                    iw, ih = ImageReader(im.path).getSize()
                    nw, nh = _native_pt_size(float(iw), float(ih))
                except Exception:
                    nw, nh = _native_pt_size(float(im.width_px), float(im.height_px))
                scale = min(col_w / nw, leg_budget / nh, 1.0)
                dw, dh = nw * scale, nh * scale
                px = x + i * (col_w + gap) + max(0.0, (col_w - dw) / 2)
                try:
                    c.drawImage(
                        ImageReader(im.path),
                        px,
                        cy - dh,
                        width=dw,
                        height=dh,
                        preserveAspectRatio=True,
                        mask="auto",
                    )
                except Exception:
                    pass
        c.showPage()
        return

    paths = [im.path for im in plots] + [im.path for im in legends]
    _draw_image_page(
        c,
        sheet_name=sheet_name,
        display_name=disp,
        section_label=section_label,
        section_color=section_color,
        titles=titles,
        image_paths=paths,
        stack_vertical=True,
        allow_upscale=False,
    )
