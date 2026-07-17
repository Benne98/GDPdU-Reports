"""Recalculate an .xlsx via Microsoft Excel on macOS so data_only values exist."""

from __future__ import annotations

import shutil
import subprocess
import tempfile
from pathlib import Path


def excel_available() -> bool:
    return Path("/Applications/Microsoft Excel.app").is_dir()


def _cached_copy_path(src: Path) -> Path:
    return src.parent / f"{src.stem}_calculated{src.suffix}"


def cached_values_score(workbook_path: str | Path) -> int:
    """
    Count non-empty cells (data_only) on key report sheets.
    Uncalculated formula workbooks score low; Excel-recalculated copies score much higher.
    """
    from openpyxl import load_workbook

    path = Path(workbook_path).expanduser().resolve()
    if not path.is_file():
        return 0
    try:
        wb = load_workbook(path, data_only=True, read_only=True)
    except Exception:
        return 0
    score = 0
    preferred = ("Lead_IS", "Lead_BS", "General sales table MM", "Cashflow")
    sheets = [name for name in preferred if name in wb.sheetnames]
    if not sheets:
        sheets = list(wb.sheetnames[:3])
    for name in sheets:
        ws = wb[name]
        for row in ws.iter_rows(min_row=1, max_row=80, max_col=30):
            for cell in row:
                if cell.value not in (None, ""):
                    score += 1
    wb.close()
    return score


def _cached_copy_usable(src: Path, dest_path: Path) -> bool:
    """True when dest_path is a fresh, Excel-recalculated sidecar (not a stale plain copy)."""
    if not dest_path.is_file():
        return False
    if dest_path.stat().st_mtime < src.stat().st_mtime:
        return False
    src_score = cached_values_score(src)
    dest_score = cached_values_score(dest_path)
    if dest_score <= src_score:
        return False
    # Recalculated workbooks typically gain many cached formula values.
    return dest_score >= max(300, int(src_score * 1.25))


def _recalc_via_excel(dest_path: Path) -> bool:
    """Open dest_path in Excel, full rebuild, save. Returns True on success."""
    posix = str(dest_path)
    calc_script = f'''tell application "Microsoft Excel"
    activate
    try
        open workbook workbook file name "{posix}"
        delay 10
        set wb to active workbook
        calculate full rebuild
        delay 15
        save wb
        delay 5
        close wb saving yes
    on error errMsg number errNum
        error errMsg number errNum
    end try
end tell'''
    with tempfile.NamedTemporaryFile("w", suffix=".applescript", delete=False) as fh:
        fh.write(calc_script)
        script_path = fh.name
    try:
        subprocess.run(
            ["osascript", script_path],
            check=True,
            capture_output=True,
            text=True,
            timeout=300,
        )
        return True
    except subprocess.TimeoutExpired:
        print(f"WARN: Excel recalc timed out for {dest_path.name}")
    except subprocess.CalledProcessError as exc:
        err = (exc.stderr or exc.stdout or str(exc))[:500]
        print(f"WARN: Excel recalc failed ({err})")
    finally:
        Path(script_path).unlink(missing_ok=True)
    return False


def materialize_calculated_copy(workbook_path: str | Path, dest: str | Path | None = None) -> Path:
    """
    Copy workbook, open in Excel, full calculate, save.
    Returns path to the calculated copy (with cached formula values).

    By default writes a sidecar ``*_calculated.xlsx`` next to the source and
    reuses it only when it contains materially more cached values than the source.
    """
    src = Path(workbook_path).expanduser().resolve()
    if not src.is_file():
        raise FileNotFoundError(src)

    dest_path = Path(dest).expanduser().resolve() if dest else _cached_copy_path(src)

    if dest_path.resolve() == src.resolve():
        dest_path = _cached_copy_path(src)

    if not excel_available():
        return src

    if _cached_copy_usable(src, dest_path):
        return dest_path

    if dest_path.is_file():
        # Drop stale sidecars (plain copy from a failed recalc, or outdated cache).
        dest_path.unlink(missing_ok=True)

    dest_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(src, dest_path)
    if not _recalc_via_excel(dest_path):
        dest_path.unlink(missing_ok=True)
        print(f"WARN: using uncalculated workbook for {src.name}")
        return src

    if not _cached_copy_usable(src, dest_path):
        dest_path.unlink(missing_ok=True)
        print(f"WARN: Excel recalc did not materialize cached values for {src.name}")
        return src

    return dest_path
