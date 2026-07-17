import re
import sys
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.styles import Border, Side
from openpyxl.formatting.rule import CellIsRule

PROJECT_ROOT = Path(__file__).resolve().parent
SCRIPTS_DIR = PROJECT_ROOT / "scripts"
BACKEND_DIR = PROJECT_ROOT / "backend"
if str(SCRIPTS_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPTS_DIR))
if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))

from gst_excel_theme import THEME, apply_recon_portfolio_layout  # noqa: E402
from databook_excel_layout import (  # noqa: E402
    LAYOUT_BS,
    apply_bs_hierarchy_borders,
    check_row_groups_after_table,
    collapse_bs_l3_groups_on_open,
    collapse_check_outline_rows,
    hide_helper_column_group,
    paint_check_source_yellow,
    paint_grey_white_canvas,
    paint_header_band,
    prune_zero_value_rows,
    recon_check_yellow_columns,
    write_bs_ale_check_section,
    write_entity_block_titles,
)
from databook_periods import (  # noqa: E402
    display_bs_period_labels,
    display_reporting_columns_from_df,
    is_fy_period_column,
    is_ytd_period_column,
    ordered_reporting_columns_from_df,
    split_fy_and_ytd,
)
from report_row_layout import (  # noqa: E402
    NA_BUCKET_ROLLUP,
    build_bs_row_structure,
    l2_l3_order_from_mapping,
)
from databook_workbook import BS_RECON_MAPPING_FILE, MASTER_WORKBOOK_STR  # noqa: E402
from databook_runtime import load_argv_config, path_value  # noqa: E402

_argv_cfg = load_argv_config()

# ================================================
# CONFIG (Desktop work defaults)
# ================================================
DESKTOP_DIR = PROJECT_ROOT / "Desktop"

PROJECT_NAME = str(_argv_cfg.get("project_name") or "Desktop Test")
GROUP_NAME = str(_argv_cfg.get("company_name") or _argv_cfg.get("group_name") or "Group")
UNIT_LABEL = "kEUR"

INPUT_FILE = path_value(_argv_cfg, "input_file", MASTER_WORKBOOK_STR)
MAPPING_FILE_BS = BS_RECON_MAPPING_FILE
SHEET_MASTER = "Master_BS"
REPORT_SHEET = "BS_Bucket"

PERIODS = []  # detected from Master_BS
NA_PERIODS = []  # derived from PERIODS
BS_PERIOD_FROM = ""
BS_PERIOD_TO = ""

ENTITY_SORT_ORDER = {}
ENTITY_RENAME_MAP = {}

IC_MASTER_ENTITY = "Consolidation"
IC_DISPLAY_NAME = "IC eliminations"

SOURCE_COL_CANDIDATES = ("L5", "L6")
REPORTED_FILTER_VALUE = "reported"
L4_SORT_BASIS = "latest_fy"

BS_TOTALS_CONFIG = [
    {"label": "Total assets", "components": ["Fixed assets", "Current assets"], "insert_after": "Current assets"},
    {"label": "Total equity & liabilities", "components": ["Equity", "Liabilities"], "insert_after": "Liabilities"},
]

SHOW_BS_CHECKS = True

FY_END_MONTH = int(_argv_cfg.get("fy_end_month") or 12)
LTM_MONTH = _argv_cfg.get("ltm_month") or "2023-7"
if "show_bs_checks" in _argv_cfg:
    SHOW_BS_CHECKS = bool(_argv_cfg["show_bs_checks"])

NA_BUCKETS = ["FA", "TWC", "OWC", "Other", "ND", "Equity"]

# ================================================
# LAYOUT
# ================================================
# LAYOUT — A free | B..F helpers | G spacer | H POS
_COL_LAYOUT = LAYOUT_BS
MAP_START_COL = _COL_LAYOUT.map_start_col
TECH_SPACER_COL = _COL_LAYOUT.spacer_col
POS_COL = _COL_LAYOUT.pos_col
FIRST_ENTITY_COL = _COL_LAYOUT.first_value_col

PROJECT_TITLE_ROW = 1
SUBTITLE_ROW = 2
TITLE_ROW = 6
ENTITY_CODE_ROW = 3
HEADER_ROW = 8
BLOCK_TITLE_ROW = 7
DATA_START_ROW = 9

BS_MAP_START_COL = MAP_START_COL
BS_MAP_END_COL = _COL_LAYOUT.map_end_col
BS_L2_COL = MAP_START_COL + 2
BS_L3_COL = MAP_START_COL + 3
BS_L4_COL = MAP_START_COL + 4

SPACER_WIDTH = 1.14
VALUE_COL_WIDTH = 7.86
ROW_HEIGHT = 12
REPEAT_POS_WIDTH = 32

# ================================================
# STYLES (gst_excel_theme)
# ================================================
FONT_BASE = THEME.font_base
FONT_BASE_BOLD = THEME.font_bold
FONT_HEADER = THEME.font_header
FONT_MAPPING = Font(name=THEME.font_name, size=THEME.font_size, color=THEME.text_header)
FONT_TITLE = Font(
    name=THEME.font_name,
    size=THEME.font_size,
    bold=True,
    color=THEME.text_brand_title,
)
FONT_PROJECT_TITLE = Font(
    name=THEME.font_name,
    size=THEME.font_size_title,
    color=THEME.text_brand_title,
    bold=False,
)
FONT_SUBTITLE = Font(
    name=THEME.font_name,
    size=THEME.font_size_subtitle,
    color=THEME.text_brand_title,
    bold=False,
)
FONT_CHECK_RED = Font(name=THEME.font_name, size=THEME.font_size, color=THEME.delta_negative)

FILL_GREY = THEME.fill_tech
FILL_WHITE = THEME.fill_white
FILL_HEADER = THEME.fill_header
FILL_SUBTOTAL = THEME.fill_subtotal

ALIGN_LEFT = Alignment(horizontal="left", vertical="center")
ALIGN_CENTER = Alignment(horizontal="center", vertical="center")
ALIGN_RIGHT = Alignment(horizontal="right", vertical="center")

THIN_SIDE = Side(style="thin", color=THEME.border_color)
TOP_BORDER = THEME.border_subtotal_top
BOTTOM_BORDER = THEME.border_header_bottom
TOP_BOTTOM_BORDER = Border(
    top=THEME.border_subtotal_top.top,
    bottom=THEME.border_subtotal_top.top,
)

NUM_FMT_INT = '#,##0;(#,##0);"-"'


# ================================================
# HELPERS
# ================================================
def col_letter(idx: int) -> str:
    return get_column_letter(idx)


def is_blank_entity(x):
    if x is None or pd.isna(x):
        return True
    s = str(x).strip()
    return (s == "") or (s.lower() in {"nan", "none", "null"})


def detect_reporting_period_columns(df_: pd.DataFrame) -> list[str]:
    periods = display_reporting_columns_from_df(df_)
    if not periods:
        raise ValueError("Keine FY/YTD-Spalten im Master_BS gefunden.")
    return periods


def detect_year_columns(df_: pd.DataFrame) -> list[str]:
    fy, _ = split_fy_and_ytd(ordered_reporting_columns_from_df(df_))
    if not fy:
        raise ValueError("Keine FY-Spalten im Master_BS gefunden (erwartet z. B. FY22A, FY23A).")
    return fy


def detect_source_col(df_):
    candidates = ["Quelle", "Source", "Reported/Adjusted", "Type", "Scenario"]
    sc = next((c for c in candidates if c in df_.columns), None)
    if sc is not None:
        return sc
    for c in df_.columns:
        try:
            s = df_[c].astype(str).str.strip().str.lower()
            if (s == "reported").any():
                return c
        except Exception:
            pass
    return None


def resolve_source_section_column(df_: pd.DataFrame) -> str | None:
    markers = {REPORTED_FILTER_VALUE, "adjusted", "consolidation"}

    def marker_count(col: str) -> int:
        if col not in df_.columns:
            return 0
        series = df_[col].astype(str).str.strip().str.lower()
        return int(series.isin(markers).sum())

    best_col = None
    best_count = 0
    for col in SOURCE_COL_CANDIDATES:
        count = marker_count(col)
        if count > best_count:
            best_col = col
            best_count = count
    if best_col and best_count > 0:
        return best_col

    legacy = detect_source_col(df_)
    if legacy:
        return legacy

    for col in SOURCE_COL_CANDIDATES:
        if col in df_.columns:
            return col
    return None


def ic_sign(block):
    return 1


def detect_na_bucket_col(df_):
    """Net-assets bucket column: prefer explicit NA, else legacy L5 view columns."""
    for cand in [
        "NA",
        "L5 - GT NA View",
        "L5 - GT NA view",
        "L5 - GT NA",
        "L5",
        "L5 - Management defined Net Asset View",
    ]:
        if cand in df_.columns:
            return cand
    for c in df_.columns:
        if str(c).strip() == "NA":
            return c
        if str(c).strip().startswith("L5"):
            return c
    raise ValueError("Keine NA-/L5-Bucket-Spalte gefunden (z.B. 'NA' oder 'L5 - GT NA View').")


detect_l5_col = detect_na_bucket_col


def subtotal_l2_refs(bs_row_structure, r, colL):
    refs = []
    for rr2 in bs_row_structure:
        if rr2["L2"] != r["L2"]:
            continue
        if rr2["type"] == "subtotal_l3":
            refs.append(f"{colL}{rr2['_excel_row']}")
        elif (
            rr2["type"] == "detail_single"
            and str(rr2["L3"]).strip() == str(rr2["L4"]).strip()
        ):
            refs.append(f"{colL}{rr2['_excel_row']}")
    return refs


def detail_indent_alignment(r):
    if r["type"] == "detail" and str(r["L3"]).strip() != str(r["L4"]).strip():
        return Alignment(horizontal="left", vertical="center", indent=1)
    return ALIGN_LEFT


def label_font(r):
    return FONT_BASE_BOLD if r["type"] in {"subtotal_l2", "total"} else FONT_BASE


# ================================================
# OPEN WORKBOOK + READ MASTER
# ================================================
wb = load_workbook(INPUT_FILE)

if SHEET_MASTER not in wb.sheetnames:
    raise RuntimeError(f"Sheet '{SHEET_MASTER}' nicht gefunden in {INPUT_FILE}")

if REPORT_SHEET in wb.sheetnames:
    del wb[REPORT_SHEET]
ws = wb.create_sheet(REPORT_SHEET)

df_bs = pd.read_excel(INPUT_FILE, sheet_name=SHEET_MASTER, engine="openpyxl")

PERIODS = detect_reporting_period_columns(df_bs)
FY_COLS, YTD_COLS = split_fy_and_ytd(PERIODS)
BS_PERIOD_FROM = PERIODS[0]
BS_PERIOD_TO = PERIODS[-1]

source_col = resolve_source_section_column(df_bs)
l5_col = detect_l5_col(df_bs)

required = {"Entity", "L2", "L3", "L4", l5_col} | set(PERIODS)
missing = required - set(df_bs.columns)
if missing:
    raise ValueError(f"Master_BS fehlt Spalten: {missing}")


def normalize_entity_keep_ic(x):
    if x is None or pd.isna(x):
        return x
    s = str(x).strip()
    if s == IC_MASTER_ENTITY:
        return s
    return ENTITY_RENAME_MAP.get(s, s) if ENTITY_RENAME_MAP else s


df_bs = df_bs.copy()
if ENTITY_RENAME_MAP:
    df_bs["Entity"] = df_bs["Entity"].apply(normalize_entity_keep_ic)
    ENTITY_SORT_ORDER = {ENTITY_RENAME_MAP.get(k, k): v for k, v in ENTITY_SORT_ORDER.items()}

# ================================================
# LOAD MAPPING (L2/L3/L4)
# ================================================
map_df = pd.read_excel(MAPPING_FILE_BS, sheet_name=0, engine="openpyxl")
map_df = map_df.loc[:, [c for c in map_df.columns if not str(c).startswith("Unnamed")]]
if not {"L2", "L3"}.issubset(map_df.columns):
    raise ValueError("Sortierung_BS_Datenbank muss Spalten 'L2' und 'L3' enthalten.")

map_df = map_df.copy()
map_df["L2"] = map_df["L2"].astype(str)
map_df["L3"] = map_df["L3"].astype(str)
bs_l2_l3_order = l2_l3_order_from_mapping(map_df)

bs_row_structure = build_bs_row_structure(
    df_bs,
    bs_l2_l3_order,
    {"l4_sort_basis": L4_SORT_BASIS},
    source_col=source_col or l5_col,
)

# ================================================
# INSERT TOTALS
# ================================================
def insert_after_anchor(struct, anchor_label, new_row):
    for i, row in enumerate(struct):
        if row["type"] == "subtotal_l2" and row["label"] == anchor_label:
            struct.insert(i + 1, new_row)
            return True
    return False


existing_labels = {r["label"] for r in bs_row_structure}
for t in BS_TOTALS_CONFIG:
    comps = [c for c in t["components"] if c in existing_labels]
    if not comps:
        continue
    total_row = {"type": "total", "label": t["label"], "components": comps, "L2": "", "L3": "", "L4": ""}
    if not insert_after_anchor(bs_row_structure, t["insert_after"], total_row):
        raise RuntimeError(f"Anchor '{t['insert_after']}' für Total '{t['label']}' nicht gefunden.")

bs_row_structure = prune_zero_value_rows(
    bs_row_structure,
    df_bs,
    PERIODS,
    source_col=source_col or SOURCE_COL_CANDIDATES[0],
)

DISPLAY_PERIODS = display_bs_period_labels(PERIODS, FY_END_MONTH, LTM_MONTH)
FY_PERIOD_INDICES = [i for i, p in enumerate(PERIODS) if is_fy_period_column(p)]
NA_PERIODS = [{"label": DISPLAY_PERIODS[i], "year_col": p} for i, p in enumerate(PERIODS)]

# ================================================
# ENTITY LIST
# ================================================
raw_entities = []
for v in pd.unique(df_bs["Entity"]):
    if not is_blank_entity(v):
        raw_entities.append(str(v).strip())

has_ic = IC_MASTER_ENTITY in raw_entities
individual_entities = [e for e in raw_entities if e != IC_MASTER_ENTITY]


def entity_sort_key(e):
    if ENTITY_SORT_ORDER:
        return ENTITY_SORT_ORDER.get(e, 9999)
    return raw_entities.index(e) if e in raw_entities else 9999


individual_entities = sorted(individual_entities, key=entity_sort_key)

# ================================================
# MASTER RANGE HELPERS
# ================================================
MASTER_START = 2
MASTER_END = len(df_bs) + 1


def master_range(colname):
    L = get_column_letter(df_bs.columns.get_loc(colname) + 1)
    return f"'{SHEET_MASTER}'!${L}${MASTER_START}:${L}${MASTER_END}"


src_rng = master_range(source_col) if source_col is not None else None
ent_rng = master_range("Entity")
l2_rng = master_range("L2")
l3_rng = master_range("L3")
l4_rng = master_range("L4")
l5_rng = master_range(l5_col)


def _na_bucket_sumifs_body(
    sum_rng: str,
    *,
    src_rng: str | None,
    rep_crit: str,
    l2_rng: str,
    l2_crit: str,
    l3_rng: str,
    l3_crit: str,
    l4_rng: str,
    l4_crit: str,
    l5_rng: str,
    cat: str,
) -> str:
    """SUMIFS for one NA bucket column; ND column rolls up DL + ND."""
    na_values = NA_BUCKET_ROLLUP.get(cat, (cat,))

    def one(na_val: str) -> str:
        na_crit = f'"{na_val}"'
        if src_rng is not None:
            return (
                f"SUMIFS({sum_rng},{src_rng},{rep_crit},"
                f"{l2_rng},{l2_crit},{l3_rng},{l3_crit},"
                f"{l4_rng},{l4_crit},{l5_rng},{na_crit})"
            )
        return (
            f"SUMIFS({sum_rng},{l2_rng},{l2_crit},{l3_rng},{l3_crit},"
            f"{l4_rng},{l4_crit},{l5_rng},{na_crit})"
        )

    return "+".join(one(v) for v in na_values)


# ================================================
# TITLES
# ================================================
pt = ws.cell(PROJECT_TITLE_ROW, POS_COL, f"Project {PROJECT_NAME}")
pt.font = FONT_PROJECT_TITLE
pt.alignment = ALIGN_LEFT
ws.row_dimensions[PROJECT_TITLE_ROW].height = 36

st = ws.cell(SUBTITLE_ROW, POS_COL, f"{GROUP_NAME} BS Bucket")
st.font = FONT_SUBTITLE
st.alignment = ALIGN_LEFT

bs_title_text = (
    f"{GROUP_NAME} | Balance sheet (reported) "
    f"{BS_PERIOD_FROM} - {BS_PERIOD_TO}"
)
bs_title = ws.cell(TITLE_ROW, POS_COL, bs_title_text)
bs_title.font = FONT_TITLE
bs_title.alignment = ALIGN_LEFT

# ================================================
# HEADERS (mapping + positions)
# ================================================
ws.cell(HEADER_ROW, BS_MAP_START_COL + 0, "Reported").font = FONT_HEADER
ws.cell(HEADER_ROW, BS_MAP_START_COL + 1, "").font = FONT_HEADER
ws.cell(HEADER_ROW, BS_L2_COL, "L2").font = FONT_HEADER
ws.cell(HEADER_ROW, BS_L3_COL, "L3").font = FONT_HEADER
ws.cell(HEADER_ROW, BS_L4_COL, "L4").font = FONT_HEADER
for c in range(BS_MAP_START_COL, BS_MAP_END_COL + 1):
    ws.cell(HEADER_ROW, c).alignment = ALIGN_LEFT

ws.cell(HEADER_ROW, POS_COL, UNIT_LABEL).font = FONT_HEADER
ws.cell(HEADER_ROW, POS_COL).alignment = ALIGN_LEFT

# ================================================
# BLOCK PLAN
# ================================================
blocks = []
for e in individual_entities:
    blocks.append({"kind": "entity", "key": e, "code": e, "title": e})
if len(individual_entities) > 0:
    blocks.append({"kind": "aggregated", "key": "Aggregated", "code": "Aggregated", "title": "Aggregated"})
if has_ic:
    blocks.append({"kind": "ic", "key": IC_DISPLAY_NAME, "code": IC_MASTER_ENTITY, "title": IC_DISPLAY_NAME})
if len(individual_entities) > 0 and has_ic:
    blocks.append({"kind": "consolidation", "key": "Consolidation", "code": "Consolidation", "title": "Consolidation"})

if not blocks:
    raise RuntimeError("Keine Entities im Master_BS gefunden.")

current_col = FIRST_ENTITY_COL
for b in blocks:
    b["startcol"] = current_col
    b["year_startcol"] = current_col
    b["year_endcol"] = current_col + len(PERIODS) - 1
    b["spacer_col"] = b["year_endcol"] + 1
    current_col = b["spacer_col"] + 1

for b in blocks:
    ws.merge_cells(
        start_row=BLOCK_TITLE_ROW,
        start_column=b["startcol"],
        end_row=BLOCK_TITLE_ROW,
        end_column=b["year_endcol"],
    )
    t = ws.cell(BLOCK_TITLE_ROW, b["startcol"], b["title"])
    t.font = FONT_HEADER
    t.alignment = ALIGN_CENTER

    ws.merge_cells(
        start_row=ENTITY_CODE_ROW,
        start_column=b["startcol"],
        end_row=ENTITY_CODE_ROW,
        end_column=b["year_endcol"],
    )
    ec = ws.cell(ENTITY_CODE_ROW, b["startcol"], b["code"])
    ec.font = FONT_BASE
    ec.alignment = ALIGN_CENTER

    for y_idx, year in enumerate(PERIODS):
        h = ws.cell(HEADER_ROW, b["year_startcol"] + y_idx, DISPLAY_PERIODS[y_idx])
        h.font = FONT_HEADER
        h.alignment = ALIGN_RIGHT

    ws.column_dimensions[col_letter(b["spacer_col"])].width = SPACER_WIDTH

ws.row_dimensions[ENTITY_CODE_ROW].outlineLevel = 2
ws.row_dimensions[ENTITY_CODE_ROW].hidden = True
ws.sheet_view.showOutlineSymbols = True


def block_lookup(key):
    return next(bb for bb in blocks if bb["key"] == key)


try:
    con_block = block_lookup("Consolidation")
except StopIteration:
    con_block = blocks[-1]

LAST_USED_COL = max(b["spacer_col"] for b in blocks)
main_spacer_cols = {b["spacer_col"] for b in blocks}

# ================================================
# WRITE MAIN TABLE ROWS
# ================================================
for i, r in enumerate(bs_row_structure):
    r["_idx"] = i
    excel_row = DATA_START_ROW + i
    r["_excel_row"] = excel_row

    ws.cell(excel_row, BS_MAP_START_COL + 0, "Reported").font = FONT_MAPPING
    ws.cell(excel_row, BS_MAP_START_COL + 1, "").font = FONT_MAPPING
    ws.cell(excel_row, BS_L2_COL, r["L2"]).font = FONT_MAPPING
    ws.cell(excel_row, BS_L3_COL, r["L3"]).font = FONT_MAPPING
    ws.cell(excel_row, BS_L4_COL, r["L4"]).font = FONT_MAPPING

    for cc in range(BS_MAP_START_COL, BS_MAP_END_COL + 1):
        ws.cell(excel_row, cc).alignment = ALIGN_LEFT

    pos_cell = ws.cell(excel_row, POS_COL, r["label"])
    pos_cell.alignment = detail_indent_alignment(r)
    pos_cell.font = label_font(r)

    ws.row_dimensions[excel_row].height = ROW_HEIGHT
    if r["type"] == "detail":
        ws.row_dimensions[excel_row].outlineLevel = 1
        ws.row_dimensions[excel_row].hidden = False
    elif r["type"] == "detail_single" and str(r["L3"]).strip() != str(r["L4"]).strip():
        ws.row_dimensions[excel_row].outlineLevel = 1
        ws.row_dimensions[excel_row].hidden = False
    else:
        ws.row_dimensions[excel_row].outlineLevel = 0
        ws.row_dimensions[excel_row].hidden = False

LAST_TABLE_ROW = DATA_START_ROW + len(bs_row_structure) - 1
collapse_bs_l3_groups_on_open(ws, bs_row_structure)

# ================================================
# VALUES for main blocks (SUMIFS / SUM)
# ================================================
map_rep_col = col_letter(BS_MAP_START_COL + 0)
map_l2_col = col_letter(BS_L2_COL)
map_l3_col = col_letter(BS_L3_COL)
map_l4_col = col_letter(BS_L4_COL)


def master_range_year(year_col):
    return master_range(year_col)


for b in blocks:
    if b["kind"] in {"entity", "ic"}:
        entity_ref = f"${col_letter(b['startcol'])}${ENTITY_CODE_ROW}"
    else:
        entity_ref = None

    sign = ic_sign(b)

    for r in bs_row_structure:
        excel_row = r["_excel_row"]
        for y_idx, year in enumerate(PERIODS):
            cell = ws.cell(excel_row, b["year_startcol"] + y_idx)
            cell.alignment = ALIGN_RIGHT
            cell.number_format = NUM_FMT_INT
            cell.font = label_font(r)

            if b["kind"] in {"entity", "ic"}:
                if r["type"] in {"detail", "detail_single"}:
                    sum_rng = master_range_year(year)
                    l2_crit = f"${map_l2_col}${excel_row}"
                    l3_crit = f"${map_l3_col}${excel_row}"
                    l4_crit = f"${map_l4_col}${excel_row}"
                    rep_crit = f"${map_rep_col}${excel_row}"

                    if src_rng is not None:
                        cell.value = (
                            f"={sign}*SUMIFS({sum_rng},"
                            f"{ent_rng},{entity_ref},"
                            f"{src_rng},{rep_crit},"
                            f"{l2_rng},{l2_crit},"
                            f"{l3_rng},{l3_crit},"
                            f"{l4_rng},{l4_crit}"
                            f")/1000"
                        )
                    else:
                        cell.value = (
                            f"={sign}*SUMIFS({sum_rng},"
                            f"{ent_rng},{entity_ref},"
                            f"{l2_rng},{l2_crit},"
                            f"{l3_rng},{l3_crit},"
                            f"{l4_rng},{l4_crit}"
                            f")/1000"
                        )

                elif r["type"] == "subtotal_l3":
                    idx = r["_idx"]
                    k = idx - 1
                    while (
                        k >= 0
                        and bs_row_structure[k]["type"] in {"detail", "detail_single"}
                        and bs_row_structure[k]["L3"] == r["L3"]
                        and bs_row_structure[k]["L2"] == r["L2"]
                    ):
                        k -= 1
                    ds = k + 1
                    de = idx - 1
                    colL = col_letter(b["year_startcol"] + y_idx)
                    cell.value = (
                        f"=SUM({colL}{DATA_START_ROW + ds}:{colL}{DATA_START_ROW + de})"
                        if de >= ds
                        else 0
                    )

                elif r["type"] == "subtotal_l2":
                    colL = col_letter(b["year_startcol"] + y_idx)
                    refs = subtotal_l2_refs(bs_row_structure, r, colL)
                    cell.value = f"=SUM({','.join(refs)})" if refs else 0

                elif r["type"] == "total":
                    colL = col_letter(b["year_startcol"] + y_idx)
                    refs = []
                    for comp in r.get("components", []):
                        for rr2 in bs_row_structure:
                            if rr2["type"] == "subtotal_l2" and rr2["label"] == comp:
                                refs.append(f"{colL}{rr2['_excel_row']}")
                                break
                    cell.value = f"=SUM({','.join(refs)})" if refs else 0

            elif b["kind"] == "aggregated":
                refs = []
                for e in individual_entities:
                    bb = block_lookup(e)
                    refs.append(f"{col_letter(bb['year_startcol'] + y_idx)}{excel_row}")
                cell.value = f"=SUM({','.join(refs)})" if refs else 0

            elif b["kind"] == "consolidation":
                agg = block_lookup("Aggregated")
                icb = block_lookup(IC_DISPLAY_NAME)
                agg_ref = f"{col_letter(agg['year_startcol'] + y_idx)}{excel_row}"
                ic_ref = f"{col_letter(icb['year_startcol'] + y_idx)}{excel_row}"
                cell.value = f"={agg_ref}+{ic_ref}"

# Borders + row fills for subtotals/totals
layout = LAYOUT_BS
apply_bs_hierarchy_borders(
    ws,
    bs_row_structure,
    layout=layout,
    last_used_col=LAST_USED_COL,
    skip_cols=main_spacer_cols,
)

for r in bs_row_structure:
    excel_row = r["_excel_row"]
    fill = FILL_SUBTOTAL if r["type"] in {"subtotal_l2", "total"} else FILL_WHITE
    if r["type"] in {"subtotal_l2", "total"}:
        ws.cell(excel_row, POS_COL).fill = fill
        for c in range(POS_COL + 1, LAST_USED_COL + 1):
            if BS_MAP_START_COL <= c <= BS_MAP_END_COL:
                continue
            if c not in main_spacer_cols:
                ws.cell(excel_row, c).fill = fill
    else:
        for c in range(POS_COL, LAST_USED_COL + 1):
            if c not in main_spacer_cols:
                ws.cell(excel_row, c).fill = fill

# ================================================
# CHECKS (FS + BS_Reconciliation)
# ================================================
CHECK_FS_ROW = None
CHECK_DELTA_ROW = None
CHECK_BSREC_ROW = None
CHECK_BSREC_DELTA_ROW = None
CHECK_ALE_ROW = None
CHECK_BAL_ROW = None
total_assets_row = None
total_el_row = None

if SHOW_BS_CHECKS:
    _check_rows = check_row_groups_after_table(LAST_TABLE_ROW, [1, 2, 2])
    (
        CHECK_ALE_ROW,
        CHECK_FS_ROW,
        CHECK_DELTA_ROW,
        CHECK_BSREC_ROW,
        CHECK_BSREC_DELTA_ROW,
    ) = _check_rows
    CHECK_BAL_ROW = CHECK_BSREC_DELTA_ROW + 2

    total_row_map = {r["label"]: r["_excel_row"] for r in bs_row_structure if r["type"] == "total"}
    total_assets_row = total_row_map["Total assets"]
    total_el_row = total_row_map["Total equity & liabilities"]

    write_bs_ale_check_section(
        ws,
        blocks=blocks,
        years=[PERIODS[i] for i in FY_PERIOD_INDICES],
        check_row=CHECK_ALE_ROW,
        total_assets_row=total_assets_row,
        total_el_row=total_el_row,
        pos_col=POS_COL,
        block_kinds=frozenset({"entity", "ic", "aggregated", "consolidation"}),
        num_fmt=NUM_FMT_INT,
        hide_outline=False,
    )

    ws.cell(CHECK_FS_ROW, POS_COL, "Source - Financial statements").alignment = ALIGN_LEFT
    ws.cell(CHECK_DELTA_ROW, POS_COL, "Difference to trial balances").alignment = ALIGN_LEFT
    ws.cell(CHECK_BSREC_ROW, POS_COL, "Source - BS Reconciliation").alignment = ALIGN_LEFT
    ws.cell(CHECK_BSREC_DELTA_ROW, POS_COL, "Difference to trial balances").alignment = ALIGN_LEFT

    ws.cell(CHECK_FS_ROW, POS_COL).font = FONT_BASE
    ws.cell(CHECK_DELTA_ROW, POS_COL).font = FONT_BASE
    ws.cell(CHECK_BSREC_ROW, POS_COL).font = FONT_BASE
    ws.cell(CHECK_BSREC_DELTA_ROW, POS_COL).font = FONT_BASE

    def setup_value_cell(c, red=False):
        c.number_format = NUM_FMT_INT
        c.alignment = ALIGN_RIGHT
        c.font = FONT_CHECK_RED if red else FONT_BASE

    check_kinds = {"entity", "ic", "aggregated", "consolidation"}

    for b in blocks:
        if b["kind"] not in check_kinds:
            continue

        for y_idx in FY_PERIOD_INDICES:
            c_fs = ws.cell(CHECK_FS_ROW, b["year_startcol"] + y_idx)
            setup_value_cell(c_fs, red=False)

            assets_cell = ws.cell(total_assets_row, b["year_startcol"] + y_idx)
            delta_cell = ws.cell(CHECK_DELTA_ROW, b["year_startcol"] + y_idx)
            delta_cell.value = f"={assets_cell.coordinate}-{c_fs.coordinate}"
            setup_value_cell(delta_cell, red=False)

            ws.conditional_formatting.add(
                delta_cell.coordinate,
                CellIsRule(
                    operator="notEqual",
                    formula=["0"],
                    font=FONT_CHECK_RED,
                ),
            )

    if "BS_Reconciliation" not in wb.sheetnames:
        raise RuntimeError("Sheet 'BS_Reconciliation' wurde nicht gefunden (für Source - BS Reconciliation).")

    ws_rec = wb["BS_Reconciliation"]

    rec_total_assets_row = None
    for rr in range(1, ws_rec.max_row + 1):
        v = ws_rec.cell(rr, POS_COL).value
        if isinstance(v, str) and v.strip() == "Total assets":
            rec_total_assets_row = rr
            break
    if rec_total_assets_row is None:
        raise RuntimeError(
            f"Konnte 'Total assets' in BS_Reconciliation (Spalte {col_letter(POS_COL)}) nicht finden."
        )

    rec_title_to_startcol = {}
    for cc in range(1, ws_rec.max_column + 1):
        v = ws_rec.cell(BLOCK_TITLE_ROW, cc).value
        if isinstance(v, str) and v.strip():
            rec_title_to_startcol[v.strip()] = cc

    def rec_startcol_for(title: str) -> int:
        if title in rec_title_to_startcol:
            return rec_title_to_startcol[title]
        if title == IC_DISPLAY_NAME and "IC eliminations" in rec_title_to_startcol:
            return rec_title_to_startcol["IC eliminations"]
        raise RuntimeError(f"Block '{title}' nicht in BS_Reconciliation (Zeile {BLOCK_TITLE_ROW}) gefunden.")

    for b in blocks:
        if b["kind"] not in check_kinds:
            continue

        rec_start = rec_startcol_for(b["title"])

        for y_idx in FY_PERIOD_INDICES:
            rec_col = rec_start + y_idx
            src_cell = ws.cell(CHECK_BSREC_ROW, b["year_startcol"] + y_idx)
            src_cell.value = f"=BS_Reconciliation!{col_letter(rec_col)}{rec_total_assets_row}"
            setup_value_cell(src_cell, red=False)

            assets_cell = ws.cell(total_assets_row, b["year_startcol"] + y_idx)
            d2 = ws.cell(CHECK_BSREC_DELTA_ROW, b["year_startcol"] + y_idx)
            d2.value = f"={assets_cell.coordinate}-{src_cell.coordinate}"
            setup_value_cell(d2, red=False)

            ws.conditional_formatting.add(
                d2.coordinate,
                CellIsRule(
                    operator="notEqual",
                    formula=["0"],
                    font=FONT_CHECK_RED,
                ),
            )

# ================================================
# NET ASSETS CLASSIFICATION TABLES
# ================================================
NA_START_COL = con_block["spacer_col"] + 1
na_blocks = []
cur = NA_START_COL

for p in NA_PERIODS:
    pos_col = cur
    cat_start = cur + 1
    cat_cols = {k: cat_start + i for i, k in enumerate(NA_BUCKETS)}
    spacer = cat_start + len(NA_BUCKETS)
    endcol = spacer - 1

    na_blocks.append({
        "label": p["label"],
        "year_col": p["year_col"],
        "pos_col": pos_col,
        "cat_cols": cat_cols,
        "startcol": pos_col,
        "endcol": endcol,
        "spacer_col": spacer,
    })

    ws.column_dimensions[col_letter(pos_col)].width = REPEAT_POS_WIDTH
    for c in cat_cols.values():
        ws.column_dimensions[col_letter(c)].width = VALUE_COL_WIDTH
    ws.column_dimensions[col_letter(spacer)].width = SPACER_WIDTH

    cur = spacer + 1

for nb in na_blocks:
    ws.merge_cells(
        start_row=TITLE_ROW,
        start_column=nb["startcol"],
        end_row=TITLE_ROW,
        end_column=nb["endcol"],
    )
    t = ws.cell(TITLE_ROW, nb["startcol"], f"{GROUP_NAME} | Net assets classification {nb['label']}")
    t.font = FONT_TITLE
    t.alignment = ALIGN_LEFT

for nb in na_blocks:
    first_bucket_col = list(nb["cat_cols"].values())[0]
    ws.merge_cells(
        start_row=BLOCK_TITLE_ROW,
        start_column=first_bucket_col,
        end_row=BLOCK_TITLE_ROW,
        end_column=nb["endcol"],
    )
    c = ws.cell(BLOCK_TITLE_ROW, first_bucket_col, nb["label"])
    c.font = FONT_BASE_BOLD
    c.alignment = ALIGN_CENTER

for nb in na_blocks:
    ws.cell(HEADER_ROW, nb["pos_col"], UNIT_LABEL).font = FONT_HEADER
    ws.cell(HEADER_ROW, nb["pos_col"]).alignment = ALIGN_LEFT
    for k, col in nb["cat_cols"].items():
        h = ws.cell(HEADER_ROW, col, k)
        h.font = FONT_HEADER
        h.alignment = ALIGN_RIGHT

for r in bs_row_structure:
    excel_row = r["_excel_row"]

    l2_crit = f"${col_letter(BS_L2_COL)}${excel_row}"
    l3_crit = f"${col_letter(BS_L3_COL)}${excel_row}"
    l4_crit = f"${col_letter(BS_L4_COL)}${excel_row}"
    rep_crit = f"${col_letter(BS_MAP_START_COL)}${excel_row}"

    for nb in na_blocks:
        pcell = ws.cell(excel_row, nb["pos_col"])
        pcell.value = f"={col_letter(POS_COL)}{excel_row}"
        pcell.alignment = detail_indent_alignment(r)
        pcell.font = label_font(r)

        for cat, col in nb["cat_cols"].items():
            c = ws.cell(excel_row, col)
            c.alignment = ALIGN_RIGHT
            c.number_format = NUM_FMT_INT
            c.font = label_font(r)

            l5_crit = f"${col_letter(col)}${HEADER_ROW}"
            sum_rng = master_range(nb["year_col"])

            if r["type"] in {"detail", "detail_single"}:
                body = _na_bucket_sumifs_body(
                    sum_rng,
                    src_rng=src_rng,
                    rep_crit=rep_crit,
                    l2_rng=l2_rng,
                    l2_crit=l2_crit,
                    l3_rng=l3_rng,
                    l3_crit=l3_crit,
                    l4_rng=l4_rng,
                    l4_crit=l4_crit,
                    l5_rng=l5_rng,
                    cat=cat,
                )
                c.value = f"=({body})/1000" if "+" in body else f"={body}/1000"

            elif r["type"] == "subtotal_l3":
                idx = r["_idx"]
                k = idx - 1
                while (
                    k >= 0
                    and bs_row_structure[k]["type"] in {"detail", "detail_single"}
                    and bs_row_structure[k]["L3"] == r["L3"]
                    and bs_row_structure[k]["L2"] == r["L2"]
                ):
                    k -= 1
                ds = k + 1
                de = idx - 1
                colL = col_letter(col)
                c.value = (
                    f"=SUM({colL}{DATA_START_ROW + ds}:{colL}{DATA_START_ROW + de})"
                    if de >= ds
                    else 0
                )

            elif r["type"] == "subtotal_l2":
                colL = col_letter(col)
                refs = subtotal_l2_refs(bs_row_structure, r, colL)
                c.value = f"=SUM({','.join(refs)})" if refs else 0

            elif r["type"] == "total":
                colL = col_letter(col)
                refs = []
                for comp in r.get("components", []):
                    for rr2 in bs_row_structure:
                        if rr2["type"] == "subtotal_l2" and rr2["label"] == comp:
                            refs.append(f"{colL}{rr2['_excel_row']}")
                            break
                c.value = f"=SUM({','.join(refs)})" if refs else 0

for rr in bs_row_structure:
    if rr["type"] not in {"subtotal_l2", "total"}:
        continue
    tr = rr["_excel_row"]
    border = TOP_BOTTOM_BORDER if rr["type"] == "total" else TOP_BORDER
    for nb in na_blocks:
        ws.cell(tr, nb["pos_col"]).border = border
        for col in nb["cat_cols"].values():
            ws.cell(tr, col).border = border

# Extend subtotal fills into NA area
for r in bs_row_structure:
    if r["type"] not in {"subtotal_l2", "total"}:
        continue
    excel_row = r["_excel_row"]
    for nb in na_blocks:
        for c in range(nb["startcol"], nb["spacer_col"] + 1):
            ws.cell(excel_row, c).fill = FILL_SUBTOTAL

if SHOW_BS_CHECKS and na_blocks:
    for nb in na_blocks:
        bucket_cols = [nb["cat_cols"][k] for k in NA_BUCKETS]
        first_col = bucket_cols[0]
        last_col = bucket_cols[-1]

        assets_rng = f"{col_letter(first_col)}{total_assets_row}:{col_letter(last_col)}{total_assets_row}"
        el_rng = f"{col_letter(first_col)}{total_el_row}:{col_letter(last_col)}{total_el_row}"

        equity_col = nb["cat_cols"]["Equity"]

        eq_cell = ws.cell(CHECK_BAL_ROW, equity_col)
        eq_cell.value = f"=SUM({assets_rng})+SUM({el_rng})"
        eq_cell.number_format = NUM_FMT_INT
        eq_cell.alignment = ALIGN_RIGHT
        eq_cell.font = FONT_CHECK_RED

        for c in bucket_cols:
            if c == equity_col:
                continue
            ws.cell(CHECK_BAL_ROW, c).value = None

# ================================================
# FORMATTING
# ================================================
LAST_USED_COL = max(LAST_USED_COL, na_blocks[-1]["spacer_col"] if na_blocks else LAST_USED_COL)

ws.column_dimensions[col_letter(POS_COL)].width = 32

for b in blocks:
    for y_idx in range(len(PERIODS)):
        ws.column_dimensions[col_letter(b["year_startcol"] + y_idx)].width = VALUE_COL_WIDTH
    ws.column_dimensions[col_letter(b["spacer_col"])].width = SPACER_WIDTH

FILL_END_ROW = LAST_TABLE_ROW + 100
for rr in range(1, FILL_END_ROW + 1):
    if rr == PROJECT_TITLE_ROW:
        continue
    ws.row_dimensions[rr].height = ROW_HEIGHT

paint_grey_white_canvas(ws, layout, last_row=FILL_END_ROW, last_col=LAST_USED_COL)

spacer_cols = {b["spacer_col"] for b in blocks} | {nb["spacer_col"] for nb in na_blocks}
for sc in spacer_cols:
    for rr in range(1, FILL_END_ROW + 1):
        ws.cell(rr, sc).border = Border()
        ws.cell(rr, sc).fill = FILL_WHITE
main_period_cols: list[int] = []
for b in blocks:
    for y_idx in range(len(PERIODS)):
        main_period_cols.append(b["year_startcol"] + y_idx)
na_value_cols: list[int] = []
for nb in na_blocks:
    na_value_cols.extend(nb["cat_cols"].values())

paint_header_band(
    ws,
    layout,
    header_rows=[BLOCK_TITLE_ROW, HEADER_ROW],
    period_cols=main_period_cols + na_value_cols,
    spacer_cols=spacer_cols,
)
write_entity_block_titles(ws, blocks, block_title_row=BLOCK_TITLE_ROW)

hide_helper_column_group(ws, layout)

if SHOW_BS_CHECKS:
    yellow_cols = recon_check_yellow_columns(
        blocks,
        pos_col=POS_COL,
        block_kinds=frozenset({"entity", "ic", "aggregated", "consolidation"}),
        spacer_cols=spacer_cols,
        period_indices=FY_PERIOD_INDICES,
    )
    for rr in (CHECK_FS_ROW, CHECK_BSREC_ROW):
        paint_check_source_yellow(ws, rr, yellow_cols)

    collapse_check_outline_rows(
        ws,
        [
            CHECK_ALE_ROW,
            CHECK_FS_ROW,
            CHECK_DELTA_ROW,
            CHECK_BSREC_ROW,
            CHECK_BSREC_DELTA_ROW,
            CHECK_BAL_ROW,
        ],
    )

for b in blocks:
    c_first = b["startcol"]
    c_last = b["spacer_col"]
    if b["kind"] in {"entity", "ic", "aggregated"}:
        for cc in range(c_first, c_last + 1):
            ws.column_dimensions[col_letter(cc)].outlineLevel = 1
            ws.column_dimensions[col_letter(cc)].hidden = False
        ws.column_dimensions[col_letter(c_last)].collapsed = True
    elif b["kind"] == "consolidation":
        for cc in range(c_first, c_last + 1):
            ws.column_dimensions[col_letter(cc)].outlineLevel = 0
            ws.column_dimensions[col_letter(cc)].hidden = False

for nb in na_blocks:
    for cc in range(nb["startcol"], nb["spacer_col"] + 1):
        ws.column_dimensions[col_letter(cc)].outlineLevel = 0
        ws.column_dimensions[col_letter(cc)].hidden = False

for nb in na_blocks:
    posL = col_letter(nb["pos_col"])
    ws.column_dimensions[posL].outlineLevel = 1
    ws.column_dimensions[posL].hidden = True
    ws.column_dimensions[posL].collapsed = True

    for c in nb["cat_cols"].values():
        cL = col_letter(c)
        ws.column_dimensions[cL].outlineLevel = 0
        ws.column_dimensions[cL].hidden = False

# ================================================
# SAVE
# ================================================
wb.save(INPUT_FILE)
print(f"Fertig. Reiter '{REPORT_SHEET}' wurde aktualisiert in:\n{INPUT_FILE}")
print(f"Jahre: {PERIODS}")
print(f"NA-Perioden: {[p['label'] for p in NA_PERIODS]}")
print(f"Entities: {individual_entities}")
print(f"NA-Bucket-Spalte verwendet: {l5_col}")
if source_col is None:
    print("Hinweis: Keine Quelle/Reported-Spalte erkannt; SUMIFS filtert daher nicht auf 'Reported'.")
else:
    print(f"Quelle-Spalte erkannt: '{source_col}' (SUMIFS filtert auf 'Reported').")
