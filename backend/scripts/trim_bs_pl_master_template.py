"""Trim a BS/PL Master workbook to standard columns (L1–L4 + optional L5+).

Usage (local only — do not commit customer source paths):
  python backend/scripts/trim_bs_pl_master_template.py \\
      --input "C:/Users/.../BS_PL_Master Test input.xlsx" \\
      --output docs/fixtures/synthetic/BS_PL_Master_template.xlsx
"""
from __future__ import annotations

import argparse
import re
from pathlib import Path

import openpyxl

BS_SHEET = "Master_BS"
PL_SHEET = "Master_PL"

ID_HEADERS = ["Entity", "Account", "Account description"]
REQUIRED_LEVEL_HEADERS = ["L1", "L2", "L3", "L4"]
_LEVEL_RE = re.compile(r"^L(\d+)$", re.IGNORECASE)


def _headers_for_sheet(src_headers: list, *, pl_sheet: bool) -> list[str]:
    normalized = list(src_headers)
    if pl_sheet and "L1 - BS/PL" in normalized and "L1" not in normalized:
        normalized = ["L1" if h == "L1 - BS/PL" else h for h in normalized]
    level_headers = [h for h in normalized if _LEVEL_RE.match(str(h or ""))]
    level_headers = sorted(
        level_headers,
        key=lambda h: int(_LEVEL_RE.match(str(h)).group(1)),  # type: ignore[union-attr]
    )
    missing = [h for h in REQUIRED_LEVEL_HEADERS if h not in level_headers]
    if missing:
        raise KeyError(f"Missing required level columns: {missing}")
    return ID_HEADERS + level_headers


def trim_sheet(ws_in, ws_out, headers: list[str]) -> int:
    for col_idx, name in enumerate(headers, start=1):
        ws_out.cell(row=1, column=col_idx, value=name)
    src_headers = [c.value for c in ws_in[1]]
    col_map: dict[str, int] = {}
    for idx, h in enumerate(src_headers):
        if h in headers:
            col_map[h] = idx
        if h == "L1 - BS/PL" and "L1" in headers:
            col_map["L1"] = idx
    missing = [h for h in headers if h not in col_map]
    if missing:
        raise KeyError(f"Sheet {ws_in.title!r} missing columns: {missing}")
    out_row = 2
    for row in ws_in.iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        for out_col, name in enumerate(headers, start=1):
            src_idx = col_map[name]
            val = row[src_idx] if src_idx < len(row) else None
            ws_out.cell(row=out_row, column=out_col, value=val)
        out_row += 1
    return out_row - 2


def main() -> None:
    parser = argparse.ArgumentParser(description="Trim BS/PL Master to L1–L4 (+ optional L5+)")
    parser.add_argument("--input", required=True, help="Source xlsx path")
    parser.add_argument("--output", required=True, help="Output xlsx path")
    args = parser.parse_args()

    src = Path(args.input)
    out = Path(args.output)
    out.parent.mkdir(parents=True, exist_ok=True)

    wb_in = openpyxl.load_workbook(src, read_only=True, data_only=True)
    wb_out = openpyxl.Workbook()
    wb_out.remove(wb_out.active)

    bs_in = wb_in[BS_SHEET]
    pl_in = wb_in[PL_SHEET]
    bs_headers = _headers_for_sheet([c.value for c in bs_in[1]], pl_sheet=False)
    pl_headers = _headers_for_sheet([c.value for c in pl_in[1]], pl_sheet=True)
    ws_bs = wb_out.create_sheet(BS_SHEET)
    ws_pl = wb_out.create_sheet(PL_SHEET)

    n_bs = trim_sheet(bs_in, ws_bs, bs_headers)
    n_pl = trim_sheet(pl_in, ws_pl, pl_headers)
    wb_out.save(out)
    print(f"Wrote {out} — {BS_SHEET}: {n_bs} rows, {PL_SHEET}: {n_pl} rows")
    print(f"Columns: {bs_headers}")


if __name__ == "__main__":
    main()
