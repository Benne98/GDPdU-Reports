"""Helpers for post-SuSa databook pipeline steps."""
from __future__ import annotations

import sys
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter

from databook_periods import (
    filter_master_period_columns,
    is_fy_period_column,
    is_month_period_column,
    read_master_bs_headers,
    resolve_adjustments_period_columns,
    resolve_consolidation_period_columns,
    fallback_yearly_columns as fallback_yearly_period_columns,
)

PROJECT_ROOT = Path(__file__).resolve().parent.parent
SCRIPTS_DIR = PROJECT_ROOT / "scripts"
if str(SCRIPTS_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPTS_DIR))

from gst_excel_theme import THEME  # noqa: E402

TEMPLATES_DIR = Path(__file__).resolve().parent / "templates" / "fdd"
PL_MAPPING_TEMPLATE = TEMPLATES_DIR / "Sortierung_PL_Datenbank.xlsx"
FS_EXTRACTION_TEMPLATE = TEMPLATES_DIR / "fs_extraction_template.xlsx"

CONSOLIDATION_META_COLS: list[str] = [
    "Entity",
    "Account",
    "Account description",
    "L1 - BS/PL",
    "L2",
    "L3",
    "L4",
]

CONSOLIDATION_TEXT_COLS = len(CONSOLIDATION_META_COLS)

TEMPLATE_CANVAS_COLS = 100
TEMPLATE_CANVAS_ROWS_BELOW_HEADER = 100

_DEFAULT_META_WIDTHS: dict[str, float] = {
    "Entity": 10.0,
    "Account": 9.0,
    "Account description": 53.0,
    "L1 - BS/PL": 12.0,
    "L2": 16.0,
    "L3": 36.0,
    "L4": 49.0,
    "L5": 8.0,
    "L6": 12.0,
    "NA": 10.0,
}
_DEFAULT_AMOUNT_COL_WIDTH = 13.0

if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from databook_paths import (  # noqa: E402
    legacy_master_workbook_filename,
    master_workbook_filename,
    master_workbook_path,
    resolve_existing_master_path,
    sanitize_project_basename,
)

def consolidation_template_path(session_id: str, output_folder: str | None = None) -> Path:
    folder = Path(output_folder) if output_folder else PROJECT_ROOT / "uploads" / session_id / "output"
    return folder / f"{session_id}_consolidation_template.xlsx"


def adjustments_template_path(session_id: str, output_folder: str | None = None) -> Path:
    folder = Path(output_folder) if output_folder else PROJECT_ROOT / "uploads" / session_id / "output"
    return folder / f"{session_id}_adjustments_template.xlsx"


def _read_master_column_widths(master_path: Path) -> tuple[dict[str, float], float]:
    """Header label -> width from Master_BS; default width for amount/period columns."""
    widths: dict[str, float] = {}
    amount_width = _DEFAULT_AMOUNT_COL_WIDTH
    mwb = load_workbook(master_path)
    try:
        if "Master_BS" not in mwb.sheetnames:
            return widths, amount_width
        mws = mwb["Master_BS"]
        for c in range(1, mws.max_column + 1):
            header = str(mws.cell(1, c).value or "").strip()
            if not header:
                continue
            letter = get_column_letter(c)
            w = mws.column_dimensions[letter].width
            if not w:
                continue
            widths[header] = float(w)
            if is_fy_period_column(header) or is_month_period_column(header):
                amount_width = float(w)
    finally:
        mwb.close()
    return widths, amount_width


def _resolve_column_width(
    header: str,
    master_widths: dict[str, float],
    amount_width: float,
) -> float:
    label = str(header).strip()
    master_w = master_widths.get(label)
    meta_default = _DEFAULT_META_WIDTHS.get(label)
    if meta_default is not None:
        if master_w is None or master_w <= _DEFAULT_AMOUNT_COL_WIDTH:
            return meta_default
        return max(master_w, meta_default)
    if master_w is not None:
        return master_w
    return amount_width


def _apply_template_column_widths(
    ws,
    headers: list[str],
    canvas_cols: int,
    master_path: Path | None,
) -> None:
    master_widths: dict[str, float] = {}
    amount_width = _DEFAULT_AMOUNT_COL_WIDTH
    if master_path and master_path.is_file():
        try:
            master_widths, amount_width = _read_master_column_widths(master_path)
        except Exception:
            pass

    for idx, header in enumerate(headers, start=1):
        w = _resolve_column_width(header, master_widths, amount_width)
        ws.column_dimensions[get_column_letter(idx)].width = w

    for c in range(len(headers) + 1, canvas_cols + 1):
        ws.column_dimensions[get_column_letter(c)].width = amount_width


def _style_header_row(ws, canvas_cols: int, headers: list[str]) -> None:
    """Full-width header band (styled cells across the canvas, not only named columns)."""
    n_cols = len(headers)
    for c in range(1, canvas_cols + 1):
        cell = ws.cell(row=1, column=c)
        if c <= n_cols:
            cell.value = headers[c - 1]
        cell.font = THEME.font_header
        cell.fill = THEME.fill_header
        cell.border = THEME.border_header_bottom
        cell.alignment = Alignment(
            horizontal="left" if c <= CONSOLIDATION_TEXT_COLS else "right",
            vertical="center",
        )
    ws.row_dimensions[1].height = 12


def _paint_template_canvas_white(ws, max_col: int, max_row: int) -> None:
    """White editable area below the header (like master SuSa canvas)."""
    for r in range(2, max_row + 1):
        ws.row_dimensions[r].height = 12
        for c in range(1, max_col + 1):
            cell = ws.cell(row=r, column=c)
            cell.fill = THEME.fill_white
            cell.font = THEME.font_base
            cell.alignment = Alignment(
                horizontal="left" if c <= CONSOLIDATION_TEXT_COLS else "right",
                vertical="center",
            )


def format_databook_template_sheet(ws, headers: list[str], master_path: Path | None = None) -> None:
    """Header row styled like SuSabyYear master output (GST theme)."""
    canvas_cols = max(len(headers), TEMPLATE_CANVAS_COLS)
    canvas_last_row = 1 + TEMPLATE_CANVAS_ROWS_BELOW_HEADER

    _style_header_row(ws, canvas_cols, headers)
    _paint_template_canvas_white(ws, canvas_cols, canvas_last_row)
    _apply_template_column_widths(ws, headers, canvas_cols, master_path)


def format_consolidation_template_sheet(ws, headers: list[str], master_path: Path | None = None) -> None:
    format_databook_template_sheet(ws, headers, master_path=master_path)


def _build_upload_template_workbook(
    config: dict,
    *,
    sheet_name: str,
    period_cols_resolver,
    default_out_name: str,
    path_resolver,
) -> Path:
    session_id = str(config.get("session_id") or config.get("case_id") or "").strip()
    output_folder = str(config.get("output_folder") or config.get("output_file_path") or "").strip()
    out_raw = str(config.get("output_path") or "").strip()
    if out_raw:
        out_path = Path(out_raw)
    elif session_id:
        out_path = path_resolver(session_id, output_folder or None)
    else:
        out_path = TEMPLATES_DIR / default_out_name

    master_raw = str(config.get("master_path") or "").strip()
    master_path = Path(master_raw).expanduser().resolve() if master_raw else None

    period_cols = period_cols_resolver(config)
    template_cols = list(CONSOLIDATION_META_COLS) + period_cols

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]
    ws = wb.create_sheet(sheet_name)
    format_databook_template_sheet(ws, template_cols, master_path=master_path)
    wb.save(out_path)
    return out_path


def build_consolidation_template_from_config(config: dict) -> Path:
    """Write empty Consolidation sheet with meta + filtered period columns."""
    return _build_upload_template_workbook(
        config,
        sheet_name="Consolidation",
        period_cols_resolver=resolve_consolidation_period_columns,
        default_out_name="consolidation_generated.xlsx",
        path_resolver=consolidation_template_path,
    )


def build_adjustments_template_from_config(config: dict) -> Path:
    """Write empty Adjustments sheet with meta + FY columns only."""
    return _build_upload_template_workbook(
        config,
        sheet_name="Adjustments",
        period_cols_resolver=resolve_adjustments_period_columns,
        default_out_name="adjustments_generated.xlsx",
        path_resolver=adjustments_template_path,
    )


def append_sheet_to_master(
    master_path: Path,
    source_path: Path,
    sheet_name: str,
) -> None:
    """Copy a sheet from source workbook into master (replace if exists)."""
    from copy import copy

    master_wb = load_workbook(master_path)
    source_wb = load_workbook(source_path)
    if sheet_name in master_wb.sheetnames:
        del master_wb[sheet_name]
    src = source_wb[sheet_name]
    dest = master_wb.create_sheet(sheet_name)
    for row in src.iter_rows():
        for cell in row:
            dest[cell.coordinate].value = cell.value
            if cell.has_style:
                dest[cell.coordinate].font = copy(cell.font)
                dest[cell.coordinate].fill = copy(cell.fill)
                dest[cell.coordinate].border = copy(cell.border)
                dest[cell.coordinate].alignment = copy(cell.alignment)
                dest[cell.coordinate].number_format = cell.number_format
    for col, dim in src.column_dimensions.items():
        dest.column_dimensions[col].width = dim.width
    master_wb.save(master_path)


def load_pl_recon_mapping_df_from_db() -> "pd.DataFrame":
    """Backend helper: PL recon mapping as DataFrame."""
    import pandas as pd
    from recon_mapping_loader import load_pl_recon_mapping_df

    return load_pl_recon_mapping_df({"paths": {"mapping_source": "db"}})


def load_bs_recon_mapping_df_from_db() -> "pd.DataFrame":
    """Backend helper: BS recon mapping as DataFrame."""
    import pandas as pd
    from recon_mapping_loader import load_bs_recon_mapping_df

    return load_bs_recon_mapping_df({"paths": {"mapping_source": "db"}})


def ensure_pl_mapping_file() -> Path:
    """Deprecated file fallback; recon pipeline uses DB. Kept for API compatibility."""
    try:
        df = load_pl_recon_mapping_df_from_db()
        if "L3" not in df.columns:
            raise ValueError("PL mapping from DB must contain column 'L3'.")
        df = df[["L3"]].copy()
        df["L3"] = df["L3"].astype(str).str.strip()
        df = df[df["L3"].ne("")].drop_duplicates(subset=["L3"], keep="first")
        TEMPLATES_DIR.mkdir(parents=True, exist_ok=True)
        out = TEMPLATES_DIR / "pl_recon_mapping_from_db.xlsx"
        df.to_excel(out, index=False, sheet_name="Mapping")
        return out
    except Exception:
        TEMPLATES_DIR.mkdir(parents=True, exist_ok=True)
        if PL_MAPPING_TEMPLATE.is_file():
            return PL_MAPPING_TEMPLATE
    rows = [
        "Net sales",
        "Own work capitalised",
        "Total output",
        "Cost of goods sold",
        "Gross profit",
        "Personnel expenses",
        "Wages & salaries",
        "Social security",
        "Other operating income",
        "Other operating expenses",
        "Net operating expenses",
        "EBITDA",
        "Depreciation",
        "EBIT",
        "Financial result",
        "EBT",
        "Taxes on income",
        "Net result",
    ]
    wb = Workbook()
    ws = wb.active
    ws.title = "Mapping"
    ws.append(["L3"])
    for l3 in rows:
        ws.append([l3])
    wb.save(PL_MAPPING_TEMPLATE)
    return PL_MAPPING_TEMPLATE


def build_consolidation_template(
    master_path: Path,
    *,
    period_level: str = "yearly",
    output_path: Path | None = None,
    config: dict | None = None,
) -> Path:
    """Empty consolidation template; period columns filtered by yearly/monthly."""
    cfg: dict[str, Any] = dict(config or {})
    cfg["master_path"] = str(master_path)
    cfg["period_level"] = period_level
    if output_path is not None:
        cfg["output_path"] = str(output_path)
    elif "output_path" not in cfg:
        TEMPLATES_DIR.mkdir(parents=True, exist_ok=True)
        cfg["output_path"] = str(TEMPLATES_DIR / "consolidation_generated.xlsx")
    return build_consolidation_template_from_config(cfg)


def prepare_master_pl_for_recon(master_path: Path) -> None:
    """Ensure Master_PL has L6=Reported for rows without section headers."""
    wb = load_workbook(master_path)
    if "Master_PL" not in wb.sheetnames:
        return
    ws = wb["Master_PL"]
    headers = {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1)}
    if "L6" not in headers:
        l6_col = ws.max_column + 1
        ws.cell(1, l6_col, "L6")
    else:
        l6_col = headers["L6"]
    skip_labels = {"Consolidation", "Total PL", "Adjustments", "Check"}
    for r in range(2, ws.max_row + 1):
        label = ws.cell(r, 1).value
        if label in skip_labels:
            continue
        if ws.cell(r, l6_col).value in (None, ""):
            ws.cell(r, l6_col, "Reported")
    wb.save(master_path)


def append_recon_to_master(recon_path: Path, master_path: Path, report_sheet: str = "PL_Reconciliation") -> None:
    """Copy PL_Reconciliation sheet from recon workbook into master workbook."""
    recon_wb = load_workbook(recon_path)
    master_wb = load_workbook(master_path)
    if report_sheet in master_wb.sheetnames:
        del master_wb[report_sheet]
    src = recon_wb[report_sheet]
    dest = master_wb.create_sheet(report_sheet)
    for row in src.iter_rows():
        for cell in row:
            dest[cell.coordinate].value = cell.value
            if cell.has_style:
                dest[cell.coordinate].font = cell.font.copy()
                dest[cell.coordinate].fill = cell.fill.copy()
                dest[cell.coordinate].border = cell.border.copy()
                dest[cell.coordinate].alignment = cell.alignment.copy()
                dest[cell.coordinate].number_format = cell.number_format
    for col, dim in src.column_dimensions.items():
        dest.column_dimensions[col].width = dim.width
    master_wb.save(master_path)


def extract_fs_check_values(checked_excel: Path) -> dict[str, Any]:
    """Best-effort net result values per FY from checked FS workbook PL sheet."""
    wb = load_workbook(checked_excel, read_only=True, data_only=True)
    if "PL" not in wb.sheetnames:
        return {"entities": [], "consolidation": []}
    ws = wb["PL"]
    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    fy_cols = [h for h in headers if isinstance(h, str) and h.upper().startswith("FY")]
    net_row = None
    for r in range(2, ws.max_row + 1):
        g = str(ws.cell(r, 1).value or "").lower()
        if "net result" in g or "jahresüberschuss" in g or "jahresfehlbetrag" in g:
            net_row = r
    if net_row is None:
        return {"entities": [], "consolidation": [0] * len(fy_cols)}
    values = []
    for h in fy_cols:
        idx = headers.index(h) + 1
        v = ws.cell(net_row, idx).value
        try:
            values.append(float(v))
        except (TypeError, ValueError):
            values.append(0.0)
    return {"entities": [], "consolidation": values}


def extract_master_entities(master_path: Path | str) -> list[str]:
    """Unique entity names from Master_BS Entity column (fallback Master_PL), first-seen order."""
    path = Path(master_path)
    if not path.is_file():
        return []

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if "Master_BS" in wb.sheetnames:
            sheet_name = "Master_BS"
        elif "Master_PL" in wb.sheetnames:
            sheet_name = "Master_PL"
        else:
            return []

        ws = wb[sheet_name]
        entity_col = 1
        for c in range(1, min(ws.max_column or 1, 50) + 1):
            header = str(ws.cell(1, c).value or "").strip()
            if header == "Entity":
                entity_col = c
                break

        seen: list[str] = []
        seen_lower: set[str] = set()
        for r in range(2, (ws.max_row or 1) + 1):
            raw = ws.cell(r, entity_col).value
            if raw is None or str(raw).strip() == "":
                continue
            name = str(raw).strip()
            key = name.lower()
            if key in ("entity", "consolidation"):
                continue
            if key not in seen_lower:
                seen.append(name)
                seen_lower.add(key)
        return seen
    finally:
        wb.close()
