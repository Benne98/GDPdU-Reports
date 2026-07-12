"""Tests for FTE_payroll.py — FTE math, layout, integration."""
import json
import tempfile
import unittest
from pathlib import Path

import openpyxl

from FTE_payroll import (
    DATA_START_ROW,
    HEADER_ROW,
    fte_column_layout,
    fte_period_header,
    fte_row_weight,
    fte_row_weight_decimal,
    fte_source_sheet_name,
    load_pl_personnel_lines,
    normalize_config,
    normalize_employment_rate,
    resolve_period_header,
    run_fte_payroll,
)

PERSONNEL_DIR = Path(__file__).resolve().parent.parent.parent / "Desktop" / "FTE Payroll"
MASTER_PL = Path(__file__).resolve().parent.parent.parent / "Desktop" / "BS_PL_Master.xlsx"


def _base_cfg(**overrides):
    cfg = {
        "title": "Test",
        "company": "Co",
        "case_id": "fte_test",
        "periods": [
            {
                "label": "FY2024",
                "file_path": str(PERSONNEL_DIR / "personaltable_2024.xlsx"),
                "sheet_name": "Personal 2024",
            }
        ],
        "columns": {
            "employment": "Beschäftigungsgrad",
            "months_sum": "Summe",
            "payroll_cols": ["Gesamtsumme"],
        },
        "group_cols": ["Bereich"],
        "amount_scale": 1000,
        "header_row": 0,
        "formula_mode": True,
        "use_session_workbook": False,
    }
    cfg.update(overrides)
    return cfg


def _multi_cfg(**overrides):
    cfg = {
        "title": "Desktop Test",
        "company": "Group",
        "case_id": "fte_multiyear",
        "periods": [
            {
                "label": "FY2023",
                "file_path": str(PERSONNEL_DIR / "personaltable_2023.xlsx"),
                "sheet_name": "Personal 2023",
            },
            {
                "label": "FY2024",
                "file_path": str(PERSONNEL_DIR / "personaltable_2024.xlsx"),
                "sheet_name": "Personal 2024",
            },
            {
                "label": "FY2025",
                "header": "YTD25A",
                "file_path": str(PERSONNEL_DIR / "personaltable_2025.xlsx"),
                "sheet_name": "Personal 2025",
            },
        ],
        "columns": {
            "employment": "Beschäftigungsgrad",
            "months_sum": "Summe",
            "payroll_cols": ["Gesamtsumme"],
        },
        "group_cols": ["Bereich"],
        "amount_scale": 1000,
        "header_row": 0,
        "formula_mode": True,
        "use_session_workbook": False,
        "master_pl_path": str(MASTER_PL) if MASTER_PL.is_file() else "",
    }
    cfg.update(overrides)
    return cfg


class TestFteHelpers(unittest.TestCase):
    def test_normalize_employment_rate(self):
        self.assertAlmostEqual(normalize_employment_rate(100), 1.0)
        self.assertAlmostEqual(normalize_employment_rate(1.0), 1.0)

    def test_fte_row_weight(self):
        self.assertEqual(fte_row_weight(100, 12), 1)
        self.assertEqual(fte_row_weight(1.0, 12), 1)
        self.assertEqual(fte_row_weight(100, 10), 1)
        self.assertAlmostEqual(fte_row_weight_decimal(100, 10), 10 / 12)

    def test_fte_period_header(self):
        self.assertEqual(fte_period_header("FY2023"), "FY23A")
        self.assertEqual(fte_period_header("FY2024"), "FY24A")
        self.assertEqual(fte_period_header("YTD2025"), "YTD25A")
        self.assertEqual(
            resolve_period_header({"label": "FY2025", "header": "YTD25A"}, {}),
            "YTD25A",
        )
        self.assertEqual(fte_source_sheet_name("FY2023"), "__SOURCE__FY23A")

    def test_column_layout_one_group(self):
        layout = fte_column_layout(1)
        self.assertEqual(layout.pos_col, 8)
        self.assertEqual(layout.group_crit_cols, (2,))

    def test_column_layout_two_groups(self):
        layout = fte_column_layout(2)
        self.assertEqual(layout.pos_col, 10)
        self.assertEqual(layout.group_crit_cols, (2, 4))

    def test_normalize_config(self):
        cfg = normalize_config(_base_cfg())
        self.assertEqual(cfg["columns"]["payroll_cols"], ["Gesamtsumme"])


class TestFteIntegration(unittest.TestCase):
    @unittest.skipUnless(
        (PERSONNEL_DIR / "personaltable_2024.xlsx").is_file(),
        "Desktop personnel file missing",
    )
    def test_single_period_smoke(self):
        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp) / "out.xlsx"
            path = run_fte_payroll(_base_cfg(case_id="x", output_file_path=str(tmp)))
            self.assertTrue(Path(path).is_file())
            wb = openpyxl.load_workbook(path, data_only=False)
            ws = wb["FTE development"]
            layout = fte_column_layout(1)
            self.assertEqual(ws.cell(HEADER_ROW, layout.first_data_col).value, "FY24A")
            val = ws.cell(DATA_START_ROW, layout.first_data_col).value
            self.assertIsInstance(val, str)
            self.assertIn("SUMIFS", str(val).upper())
            self.assertIn("ROUND", str(val).upper())

    @unittest.skipUnless(
        all((PERSONNEL_DIR / f"personaltable_{y}.xlsx").is_file() for y in (2023, 2024, 2025)),
        "Multi-year personnel files missing",
    )
    def test_three_period_blocks(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = run_fte_payroll(_multi_cfg(case_id="x", output_file_path=str(tmp)))
            wb = openpyxl.load_workbook(path, data_only=False)
            ws = wb["FTE development"]
            layout = fte_column_layout(1)
            labels = [
                ws.cell(r, layout.pos_col).value
                for r in range(1, ws.max_row + 1)
                if ws.cell(r, layout.pos_col).value
            ]
            joined = " ".join(str(x) for x in labels if x)
            self.assertIn("# Average FTEs", joined)
            self.assertIn("Payroll accounting", joined)
            self.assertIn("Personnel expenses", joined)
            self.assertEqual(len([s for s in wb.sheetnames if s.startswith("__SOURCE__")]), 3)
            self.assertIn("__SOURCE__FY23A", wb.sheetnames)
            self.assertIn("__SOURCE__YTD25A", wb.sheetnames)

    @unittest.skipUnless(
        all((PERSONNEL_DIR / f"personaltable_{y}.xlsx").is_file() for y in (2023, 2024, 2025)),
        "Multi-year personnel files missing",
    )
    def test_two_level_hierarchy(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = run_fte_payroll(
                _multi_cfg(
                    case_id="x",
                    output_file_path=str(tmp),
                    group_cols=["Bereich", "Bereichuntergruppe"],
                )
            )
            wb = openpyxl.load_workbook(path, data_only=False)
            ws = wb["FTE development"]
            layout = fte_column_layout(2)
            labels = [ws.cell(r, layout.pos_col).value for r in range(DATA_START_ROW, DATA_START_ROW + 25)]
            indented = [l for l in labels if l and str(l).startswith("    ")]
            self.assertTrue(len(indented) >= 1)
            formula = ws.cell(DATA_START_ROW + 2, layout.first_data_col).value
            self.assertIn("$B", str(formula))
            self.assertIn("$D", str(formula))

    @unittest.skipUnless(MASTER_PL.is_file(), "BS_PL_Master missing")
    def test_pl_bridge_rows(self):
        lines = load_pl_personnel_lines(str(MASTER_PL))
        self.assertTrue(len(lines) >= 1)
        l4s = [x[0] for x in lines]
        self.assertNotIn("Wages & salaries", l4s)
        self.assertEqual(l4s[-1], "Other")
        with tempfile.TemporaryDirectory() as tmp:
            path = run_fte_payroll(_multi_cfg(case_id="x", output_file_path=str(tmp)))
            wb = openpyxl.load_workbook(path, data_only=False)
            ws = wb["FTE development"]
            layout = fte_column_layout(1)
            total_labels = [
                ws.cell(r, layout.pos_col).value
                for r in range(1, ws.max_row + 1)
                if ws.cell(r, layout.pos_col).value
            ]
            self.assertIn("Personnel expenses", total_labels)
            pl_formula = None
            for r in range(1, ws.max_row + 1):
                val = ws.cell(r, layout.first_data_col).value
                if isinstance(val, str) and "SUMIFS" in val.upper() and "Master_PL" in val:
                    pl_formula = val
                    break
            self.assertIsNotNone(pl_formula)

    def test_session_workbook_append(self):
        with tempfile.TemporaryDirectory() as tmp:
            cfg = _base_cfg(
                case_id="sess",
                output_file_path=str(tmp),
                use_session_workbook=True,
            )
            if not (PERSONNEL_DIR / "personaltable_2024.xlsx").is_file():
                self.skipTest("personnel file missing")
            path1 = run_fte_payroll(cfg)
            from openpyxl import Workbook

            wb = openpyxl.load_workbook(path1)
            if "KeepMe" not in wb.sheetnames:
                wb.create_sheet("KeepMe")
                wb.save(path1)
                wb.close()
            path2 = run_fte_payroll(cfg)
            self.assertEqual(path1, path2)
            wb2 = openpyxl.load_workbook(path2)
            self.assertIn("KeepMe", wb2.sheetnames)
            self.assertIn("FTE development", wb2.sheetnames)


if __name__ == "__main__":
    unittest.main()
