"""Tests for OPOS BS_Reconciliation reported refs and check portfolio."""
from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import Workbook

ROOT = Path(__file__).resolve().parent.parent.parent
SCRIPTS = ROOT / "scripts"
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))
if str(SCRIPTS) not in sys.path:
    sys.path.insert(0, str(SCRIPTS))

from databook_excel_layout import (  # noqa: E402
    AGGREGATED_BLOCK_TITLE,
    FILL_YELLOW,
    LAYOUT_BS,
    RECON_BLOCK_TITLE_ROW,
    RECON_HEADER_ROW,
    find_recon_block_col_for_period_label,
)
from opos import (  # noqa: E402
    POS_COL,
    _opos_bs_recon_reported_formula,
    apply_opos_check_sections,
    apply_opos_bs_recon_reported_refs,
    AsOfPeriod,
)


def _seed_bs_recon(wb: Workbook, *, header_label: str = "Dec24A") -> None:
    ws = wb.create_sheet("BS_Reconciliation")
    ws.cell(RECON_BLOCK_TITLE_ROW, 10, AGGREGATED_BLOCK_TITLE)
    ws.cell(RECON_HEADER_ROW, 10, header_label)
    ws.cell(12, LAYOUT_BS.pos_col, "Trade receivables")
    ws.cell(12, 10, 1234)


def test_opos_bs_recon_reported_formula_debitor():
    wb = Workbook()
    _seed_bs_recon(wb)
    formula = _opos_bs_recon_reported_formula(wb, "debitor", "Dec24A")
    assert formula == "='BS_Reconciliation'!J12"


def test_opos_bs_recon_reported_formula_missing_sheet():
    wb = Workbook()
    assert _opos_bs_recon_reported_formula(wb, "debitor", "Dec24A") is None


def test_opos_bs_recon_reported_formula_same_year_suffix():
    wb = Workbook()
    _seed_bs_recon(wb, header_label="Jul24A")
    formula = _opos_bs_recon_reported_formula(wb, "debitor", "Dec24A")
    assert formula == "='BS_Reconciliation'!J12"
    ws = wb["BS_Reconciliation"]
    assert find_recon_block_col_for_period_label(ws, AGGREGATED_BLOCK_TITLE, "Dec24A") == 10


def test_opos_bs_recon_reported_formula_no_cross_year_fallback():
    """Jul25A must not invent Dec24A — leave n/a when the snapshot year is missing."""
    wb = Workbook()
    _seed_bs_recon(wb, header_label="Dec24A")
    assert _opos_bs_recon_reported_formula(wb, "debitor", "Jul25A") is None


def test_apply_opos_summary_check_section_writes_formulas():
    wb = Workbook()
    _seed_bs_recon(wb)
    ws = wb.create_sheet("Trade debtors aging")
    periods = [AsOfPeriod("Dec24A", __import__("pandas").Timestamp("2024-12-31"))]
    footer_rows = {"sum": 20, "recon": 21, "reported": 22}
    summary_layout = {
        "footer_rows": footer_rows,
        "last_col": POS_COL + 1,
    }
    col = POS_COL + 1
    ws.cell(footer_rows["sum"], col, 100)
    ws.cell(footer_rows["reported"], col, None)

    apply_opos_bs_recon_reported_refs(
        wb,
        {"side": "debitor"},
        periods,
        {"footer_rows": footer_rows, "blocks": []},
        summary_layout,
    )
    assert ws.cell(footer_rows["reported"], col).value == "='BS_Reconciliation'!J12"

    apply_opos_check_sections(
        wb,
        {"side": "debitor"},
        periods,
        {"footer_rows": footer_rows, "blocks": [], "last_col": col, "bucket_cols": []},
        summary_layout,
    )
    source_row = footer_rows["reported"] + 2
    check_row = footer_rows["reported"] + 3
    assert ws.cell(source_row, POS_COL).value == "Source - BS Reconciliation"
    assert ws.cell(check_row, POS_COL).value == "Check"
    assert ws.cell(source_row, col).value == "='BS_Reconciliation'!J12"
    assert str(ws.cell(check_row, col).value).startswith("=")
    assert ws.cell(source_row, POS_COL).fill == FILL_YELLOW
    assert ws.cell(source_row, col).fill == FILL_YELLOW
    assert ws.row_dimensions[source_row].outlineLevel == 1
    assert ws.row_dimensions[source_row].hidden is True
