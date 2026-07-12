"""Tests: each FDD strand only clears its own workbook tabs."""

from __future__ import annotations

from openpyxl import Workbook

from databook_workbook import (
    clear_fa_rollf_workbook_sheets,
    clear_fte_workbook_sheets,
    clear_opos_workbook_sheets,
)


def _wb_with_strand_tabs():
    wb = Workbook()
    wb.remove(wb.active)
    wb.create_sheet("Master_BS")
    wb.create_sheet("Trade debtors aging")
    wb.create_sheet("Trade debtors aging detail")
    wb.create_sheet("__SOURCE__AR_Jul25A")
    wb.create_sheet("__SOURCE__FA_Dec25A")
    wb.create_sheet("FA roll forward")
    wb.create_sheet("FTE Development")
    wb.create_sheet("__SOURCE__FY25A")
    wb.create_sheet("PL_Reconciliation")
    return wb


def test_clear_opos_workbook_sheets_preserves_other_strands():
    wb = _wb_with_strand_tabs()
    clear_opos_workbook_sheets(wb)
    assert "Trade debtors aging" not in wb.sheetnames
    assert "__SOURCE__AR_Jul25A" not in wb.sheetnames
    assert "__SOURCE__FA_Dec25A" in wb.sheetnames
    assert "FA roll forward" in wb.sheetnames
    assert "FTE Development" in wb.sheetnames
    assert "Master_BS" in wb.sheetnames


def test_clear_fa_rollf_workbook_sheets_preserves_other_strands():
    wb = _wb_with_strand_tabs()
    clear_fa_rollf_workbook_sheets(wb)
    assert "FA roll forward" not in wb.sheetnames
    assert "__SOURCE__FA_Dec25A" not in wb.sheetnames
    assert "Trade debtors aging" in wb.sheetnames
    assert "__SOURCE__AR_Jul25A" in wb.sheetnames
    assert "FTE Development" in wb.sheetnames


def test_clear_fte_workbook_sheets_preserves_other_strands():
    wb = _wb_with_strand_tabs()
    clear_fte_workbook_sheets(wb)
    assert "FTE Development" not in wb.sheetnames
    assert "__SOURCE__FY25A" not in wb.sheetnames
    assert "Trade debtors aging" in wb.sheetnames
    assert "FA roll forward" in wb.sheetnames
    assert "__SOURCE__FA_Dec25A" in wb.sheetnames


def test_build_output_file_path_reuses_existing_master_by_mtime(tmp_path):
    from funktionssammlung import build_output_file_path

    older = tmp_path / "SessionA_Master.xlsx"
    newer = tmp_path / "BS_PL_Master.xlsx"
    older.write_bytes(b"old")
    newer.write_bytes(b"new")
    cfg = {
        "title": "Other Project",
        "output_file_path": str(tmp_path),
        "case_id": "sess-9",
        "use_session_workbook": True,
    }
    assert build_output_file_path(cfg) == str(newer.resolve())
