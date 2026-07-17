"""Excel recalc sidecar validation (no Microsoft Excel required)."""
from __future__ import annotations

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
SCRIPTS = ROOT / "scripts"
if str(SCRIPTS) not in sys.path:
    sys.path.insert(0, str(SCRIPTS))

from pdf_report.excel_recalc import _cached_copy_usable, cached_values_score


def test_cached_copy_usable_rejects_stale_plain_copy(tmp_path: Path):
    from openpyxl import Workbook

    src = tmp_path / "Demo_FastTrack.xlsx"
    dest = tmp_path / "Demo_FastTrack_calculated.xlsx"

    wb = Workbook()
    ws = wb.active
    ws.title = "Lead_IS"
    ws["A1"] = "Label"
    ws["B1"] = "=1+1"
    wb.save(src)
    wb.save(dest)

    assert cached_values_score(src) == cached_values_score(dest)
    assert not _cached_copy_usable(src, dest)


def test_cached_copy_usable_accepts_richer_sidecar(tmp_path: Path):
    from openpyxl import Workbook

    src = tmp_path / "Demo_FastTrack.xlsx"
    dest = tmp_path / "Demo_FastTrack_calculated.xlsx"

    wb = Workbook()
    ws = wb.active
    ws.title = "Lead_IS"
    ws["A1"] = "Label"
    ws["B1"] = "=1+1"
    wb.save(src)

    wb2 = Workbook()
    ws2 = wb2.active
    ws2.title = "Lead_IS"
    for r in range(1, 41):
        for c in range(1, 21):
            ws2.cell(row=r, column=c, value=float(r * c))
    wb2.save(dest)

    assert cached_values_score(dest) > cached_values_score(src)
    assert _cached_copy_usable(src, dest)
