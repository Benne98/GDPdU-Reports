"""Tests for deriving Master_BS / Master_PL from GL trial-balance services.

Covers (synthetic fixtures only — no DB, no real client data):
  (1) reshape unit test — TB dicts → Master frame: columns, signs NOT double-flipped,
      L-mapping (level_2→L2, level_3→L3, level_4→L4, L5="", L6="Reported"), bare
      gl_account_id in Account, NA=level_1 for BS.
  (2) reconciliation — per-period Master column sums equal the source TB sums.
  (3) net income — exactly one BS equity row per entity, verbatim recon labels.
  (4) workbook integration — written xlsx opens with Master_BS / Master_PL and the
      full META columns.

=== WORKED EXAMPLE (anchor year=2024, fy_end_month=12, fiscal_years=[2023, 2024]) ===
  PL rows (presented: revenue +, expense −):
    REV  level_2='Income'  level_3='Net sales'       level_4='Net sales'  FY2023=+1000 YTD2024=+1200
    COGS level_2='Expense' level_3='Cost of sales'   level_4='Materials'  FY2023=-400  YTD2024=-500
  Master_PL (no re-flip):  REV row → FY23A=+1000 FY24A=+1200 ; COGS row → FY23A=-400 FY24A=-500
  Net income (BS equity, sum of presented PL per column):
    FY23A = 1000 + (-400) = +600 ; FY24A = 1200 + (-500) = +700
  BS rows (raw cumulative closing balance):
    AR  level_1='Assets' level_2='Current assets' level_3='Trade receivables' DEC2023=+600 CM2024-12=+650
  Master_BS: AR row FY23A=+600 FY24A=+650 ; NA='Assets' ; plus one Net income row (+600/+700).

=== EDGE CASES ===
  - zero amount column → 0.0 carried through (no NaN).
  - negative PL value (expense) → stays negative (no double flip).
  - missing TB key for a requested period → filled with 0.0.
  - empty PL → no net-income row injected.
  - fiscal year outside TB span → silently skipped (not derivable from anchor).
"""
from __future__ import annotations

from pathlib import Path

import pytest

from app.services import fin_compat_master_from_gl as mfg


# ---------------------------------------------------------------------------
# Synthetic trial-balance dicts (shape of build_pl/bs_trial_balance output)
# ---------------------------------------------------------------------------
def _pl_row(entity, level_2, level_3, level_4, gid, name, amounts):
    return {
        "entity": entity,
        "level_2": level_2,
        "level_3": level_3,
        "level_4": level_4,
        "account": f"{gid} | {name}",  # the "gid | name" field that must NOT be used
        "gl_account_id": gid,
        "account_name": name,
        "amounts": amounts,
    }


def _bs_row(entity, level_1, level_2, level_3, level_4, gid, name, amounts):
    return {
        "entity": entity,
        "level_1": level_1,
        "level_2": level_2,
        "level_3": level_3,
        "level_4": level_4,
        "account": f"{gid} | {name}",
        "gl_account_id": gid,
        "account_name": name,
        "amounts": amounts,
    }


# anchor year=2024, fy_end_month=12 → period keys 2021-01 .. 2024-12
_FY = [2023, 2024]


def _pl_tb():
    return {
        "statement_type": "PL",
        "year": 2024,
        "month": 12,
        "rows": [
            _pl_row(
                "AT", "Income", "Net sales", "Net sales", "8000", "Umsatzerlöse",
                {"FY2023": 1000.0, "YTD2024": 1200.0, "2024-12": 100.0},
            ),
            _pl_row(
                "AT", "Expense", "Cost of sales", "Materials", "5000", "Materialaufwand",
                {"FY2023": -400.0, "YTD2024": -500.0, "2024-12": -40.0},
            ),
        ],
        "row_count": 2,
    }


def _bs_tb():
    return {
        "statement_type": "BS",
        "year": 2024,
        "month": 12,
        "rows": [
            _bs_row(
                "AT", "Assets", "Current assets", "Trade receivables", "Receivables",
                "1200", "Forderungen",
                {"DEC2023": 600.0, "CM2024-12": 650.0, "2024-12": 650.0},
            ),
            _bs_row(
                "AT", "Equity & liabilities", "Liabilities", "Trade payables", "Payables",
                "1600", "Verbindlichkeiten",
                {"DEC2023": -300.0, "CM2024-12": -350.0, "2024-12": -350.0},
            ),
        ],
        "row_count": 2,
    }


@pytest.fixture
def frames(monkeypatch):
    monkeypatch.setattr(mfg, "build_pl_trial_balance", lambda *a, **k: _pl_tb())
    monkeypatch.setattr(mfg, "build_bs_trial_balance", lambda *a, **k: _bs_tb())
    df_bs, df_pl = mfg.build_master_frames_from_gl(
        session=object(), entity_codes=["AT"], fiscal_years=_FY, fy_end_month=12
    )
    return df_bs, df_pl


# ---------------------------------------------------------------------------
# (1) Reshape unit test
# ---------------------------------------------------------------------------
class TestReshape:
    def test_pl_columns_and_meta(self, frames):
        _, df_pl = frames
        for col in mfg.META_COLS_PL:
            assert col in df_pl.columns
        assert "FY23A" in df_pl.columns and "FY24A" in df_pl.columns
        assert "Dec-2024" in df_pl.columns  # monthly label format Mon-YYYY
        assert "NA" not in df_pl.columns  # NA is BS-only

    def test_bs_columns_and_meta(self, frames):
        df_bs, _ = frames
        for col in mfg.META_COLS_BS:
            assert col in df_bs.columns
        assert "NA" in df_bs.columns

    def test_account_is_bare_gl_id(self, frames):
        _, df_pl = frames
        accounts = set(df_pl["Account"])
        assert "8000" in accounts and "5000" in accounts
        assert all("|" not in a for a in accounts)  # not the "gid | name" field

    def test_l_mapping_pl(self, frames):
        _, df_pl = frames
        rev = df_pl[df_pl["Account"] == "8000"].iloc[0]
        assert rev["L1 - BS/PL"] == "PL"
        assert rev["L2"] == "Income"
        assert rev["L3"] == "Net sales"
        assert rev["L4"] == "Net sales"
        assert rev["L5"] == ""
        assert rev["L6"] == "Reported"

    def test_l_mapping_bs_na_is_level_1(self, frames):
        df_bs, _ = frames
        ar = df_bs[df_bs["Account"] == "1200"].iloc[0]
        assert ar["L1 - BS/PL"] == "BS"
        assert ar["L2"] == "Current assets"
        assert ar["L3"] == "Trade receivables"
        assert ar["NA"] == "Assets"  # DB level_1 lands in NA

    def test_signs_not_double_flipped(self, frames):
        _, df_pl = frames
        rev = df_pl[df_pl["Account"] == "8000"].iloc[0]
        cogs = df_pl[df_pl["Account"] == "5000"].iloc[0]
        assert rev["FY23A"] == 1000.0 and rev["FY24A"] == 1200.0  # revenue stays +
        assert cogs["FY23A"] == -400.0 and cogs["FY24A"] == -500.0  # expense stays −

    def test_missing_period_key_is_zero(self, frames):
        _, df_pl = frames
        rev = df_pl[df_pl["Account"] == "8000"].iloc[0]
        # 2021-01 is a requested period column but not in the TB amounts → 0.0
        assert rev["Jan-2021"] == 0.0


# ---------------------------------------------------------------------------
# (2) Reconciliation — Master column sums == source TB sums
# ---------------------------------------------------------------------------
class TestReconciliation:
    def test_pl_fy_column_sums_match_tb(self, frames):
        _, df_pl = frames
        # FY23A ← TB FY2023 ; FY24A ← TB YTD2024
        tb_fy23 = sum(r["amounts"].get("FY2023", 0.0) for r in _pl_tb()["rows"])
        tb_fy24 = sum(r["amounts"].get("YTD2024", 0.0) for r in _pl_tb()["rows"])
        assert df_pl["FY23A"].sum() == pytest.approx(tb_fy23)
        assert df_pl["FY24A"].sum() == pytest.approx(tb_fy24)

    def test_pl_month_column_sum_matches_tb(self, frames):
        _, df_pl = frames
        tb = sum(r["amounts"].get("2024-12", 0.0) for r in _pl_tb()["rows"])
        assert df_pl["Dec-2024"].sum() == pytest.approx(tb)

    def test_bs_fy_column_sums_match_tb_excluding_net_income(self, frames):
        df_bs, _ = frames
        src = df_bs[df_bs["Account"] != mfg.NET_INCOME_ACCOUNT]
        tb_fy23 = sum(r["amounts"].get("DEC2023", 0.0) for r in _bs_tb()["rows"])
        tb_cm = sum(r["amounts"].get("CM2024-12", 0.0) for r in _bs_tb()["rows"])
        assert src["FY23A"].sum() == pytest.approx(tb_fy23)
        assert src["FY24A"].sum() == pytest.approx(tb_cm)


# ---------------------------------------------------------------------------
# (3) Net income row
# ---------------------------------------------------------------------------
class TestNetIncome:
    def test_exactly_one_row_per_entity(self, frames):
        df_bs, _ = frames
        ni = df_bs[df_bs["Account"] == mfg.NET_INCOME_ACCOUNT]
        assert len(ni) == 1
        assert ni.iloc[0]["Entity"] == "AT"

    def test_labels_match_recon_taxonomy(self, frames):
        df_bs, _ = frames
        ni = df_bs[df_bs["Account"] == mfg.NET_INCOME_ACCOUNT].iloc[0]
        assert ni["L1 - BS/PL"] == "BS"
        assert ni["L2"] == "Equity"
        assert ni["L3"] == "Net retained profits"
        assert ni["L4"] == "Net income"
        assert ni["NA"] == "Equity"
        assert ni["Account description"] == mfg.NET_INCOME_DESCRIPTION

    def test_net_income_equals_pl_column_sums(self, frames):
        df_bs, df_pl = frames
        ni = df_bs[df_bs["Account"] == mfg.NET_INCOME_ACCOUNT].iloc[0]
        # FY23A: 1000 + (-400) = 600 ; FY24A: 1200 + (-500) = 700
        assert ni["FY23A"] == pytest.approx(df_pl["FY23A"].sum())
        assert ni["FY24A"] == pytest.approx(df_pl["FY24A"].sum())
        assert ni["FY23A"] == pytest.approx(600.0)
        assert ni["FY24A"] == pytest.approx(700.0)

    def test_no_net_income_when_pl_empty(self, monkeypatch):
        monkeypatch.setattr(mfg, "build_pl_trial_balance", lambda *a, **k: {"rows": []})
        monkeypatch.setattr(mfg, "build_bs_trial_balance", lambda *a, **k: _bs_tb())
        df_bs, _ = mfg.build_master_frames_from_gl(
            session=object(), entity_codes=None, fiscal_years=_FY, fy_end_month=12
        )
        assert (df_bs["Account"] == mfg.NET_INCOME_ACCOUNT).sum() == 0


# ---------------------------------------------------------------------------
# (4) Workbook integration
# ---------------------------------------------------------------------------
class TestWorkbook:
    def test_written_xlsx_has_both_sheets_and_meta(self, frames, tmp_path):
        from openpyxl import load_workbook

        df_bs, df_pl = frames
        out = tmp_path / "case_SuSa_Master.xlsx"
        mfg.write_master_workbook(df_bs, df_pl, str(out))
        assert out.is_file()

        wb = load_workbook(out)
        try:
            assert "Master_BS" in wb.sheetnames
            assert "Master_PL" in wb.sheetnames
            bs_headers = [c.value for c in wb["Master_BS"][1]]
            pl_headers = [c.value for c in wb["Master_PL"][1]]
            for col in mfg.META_COLS_BS:
                assert col in bs_headers, f"Master_BS missing {col}"
            for col in mfg.META_COLS_PL:
                assert col in pl_headers, f"Master_PL missing {col}"
            assert "FY24A" in bs_headers and "FY24A" in pl_headers
        finally:
            wb.close()


# ---------------------------------------------------------------------------
# (5) Multi-entity isolation (C1 regression)
#
# The TB services take a SINGLE legal_entity_code. A comma-joined string used to
# resolve to "no filter" → every tenant leaked into the Master. build_master_*
# must now call the TB services once per allowed entity_code, so the result is
# strictly limited to the requested+allowed set.
# ---------------------------------------------------------------------------
class TestMultiEntityIsolation:
    # Synthetic 3-entity dataset; only AT + DE are in the allow-list, CH must NOT leak.
    _PL_BY_ENTITY = {
        "AT": [_pl_row("AT", "Income", "Net sales", "Net sales", "8000", "AT rev",
                       {"FY2023": 100.0, "YTD2024": 110.0})],
        "DE": [_pl_row("DE", "Income", "Net sales", "Net sales", "8000", "DE rev",
                       {"FY2023": 200.0, "YTD2024": 210.0})],
        "CH": [_pl_row("CH", "Income", "Net sales", "Net sales", "8000", "CH rev",
                       {"FY2023": 999.0, "YTD2024": 999.0})],
    }
    _BS_BY_ENTITY = {
        "AT": [_bs_row("AT", "Assets", "Current assets", "Trade receivables", "Recv",
                       "1200", "AT AR", {"DEC2023": 50.0, "CM2024-12": 55.0})],
        "DE": [_bs_row("DE", "Assets", "Current assets", "Trade receivables", "Recv",
                       "1200", "DE AR", {"DEC2023": 60.0, "CM2024-12": 65.0})],
        "CH": [_bs_row("CH", "Assets", "Current assets", "Trade receivables", "Recv",
                       "1200", "CH AR", {"DEC2023": 777.0, "CM2024-12": 777.0})],
    }

    @staticmethod
    def _fake_pl(session, year, month, entity=None):
        # Mirrors build_pl_trial_balance: a single legal_entity_code filters to it;
        # None/"all"/comma-string returns ALL rows (the leak the fix must avoid).
        rows = TestMultiEntityIsolation._PL_BY_ENTITY.get(str(entity))
        if rows is None:
            rows = [r for rs in TestMultiEntityIsolation._PL_BY_ENTITY.values() for r in rs]
        return {"rows": rows}

    @staticmethod
    def _fake_bs(session, year, month, entity=None):
        rows = TestMultiEntityIsolation._BS_BY_ENTITY.get(str(entity))
        if rows is None:
            rows = [r for rs in TestMultiEntityIsolation._BS_BY_ENTITY.values() for r in rs]
        return {"rows": rows}

    def test_only_allowed_entities_appear(self, monkeypatch):
        monkeypatch.setattr(mfg, "build_pl_trial_balance", self._fake_pl)
        monkeypatch.setattr(mfg, "build_bs_trial_balance", self._fake_bs)

        allowed = ["AT", "DE"]
        df_bs, df_pl = mfg.build_master_frames_from_gl(
            session=object(), entity_codes=allowed, fiscal_years=_FY, fy_end_month=12
        )

        pl_entities = set(df_pl["Entity"].unique())
        # BS includes injected net-income rows (one per entity) — those entities
        # still derive only from the allowed PL set, so the subset check holds.
        bs_entities = set(df_bs["Entity"].unique())

        assert pl_entities <= set(allowed), f"PL leaked entities: {pl_entities - set(allowed)}"
        assert bs_entities <= set(allowed), f"BS leaked entities: {bs_entities - set(allowed)}"
        assert "CH" not in pl_entities and "CH" not in bs_entities
        # Positive: both allowed entities are present.
        assert {"AT", "DE"} <= pl_entities

    def test_comma_string_is_not_passed_to_tb(self, monkeypatch):
        seen_entities: list = []

        def _spy_pl(session, year, month, entity=None):
            seen_entities.append(entity)
            return self._fake_pl(session, year, month, entity)

        monkeypatch.setattr(mfg, "build_pl_trial_balance", _spy_pl)
        monkeypatch.setattr(mfg, "build_bs_trial_balance", self._fake_bs)

        mfg.build_master_frames_from_gl(
            session=object(), entity_codes=["AT", "DE"], fiscal_years=_FY, fy_end_month=12
        )
        # Each call must receive a SINGLE code, never a comma-joined string / None.
        assert seen_entities == ["AT", "DE"]
        assert all("," not in str(e) for e in seen_entities)


# ---------------------------------------------------------------------------
# Period / FY label helpers (pure)
# ---------------------------------------------------------------------------
class TestLabelHelpers:
    def test_month_label_format(self):
        assert mfg._month_master_label(2024, 1) == "Jan-2024"
        assert mfg._month_master_label(2024, 12) == "Dec-2024"

    def test_fy_label_format(self):
        assert mfg._fy_master_label(2024) == "FY24A"

    def test_anchor_from_fiscal_years(self):
        assert mfg._anchor_from_fiscal_years([2022, 2024, 2023], 12) == (2024, 12)
        assert mfg._anchor_from_fiscal_years([2024], 6) == (2024, 6)

    def test_anchor_rejects_empty(self):
        with pytest.raises(ValueError):
            mfg._anchor_from_fiscal_years([], 12)
