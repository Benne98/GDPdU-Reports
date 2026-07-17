from openpyxl import Workbook
import pandas as pd

from bubblescatterplot import export_excel_with_images
from revenue_reconciliation import (
    NA_VALUE,
    resolve_pl_net_sales_references,
    resolve_pl_sales_references,
    sales_basis_labels,
    write_reconciliation_cells,
)


def _workbook_with_pl_reconciliation():
    wb = Workbook()
    ws = wb.active
    ws.title = "PL_Reconciliation"
    ws["B1"] = "Aggregated"
    ws["C3"] = "FY24A"
    ws["D3"] = "FY25A"
    ws["B10"] = "Net sales"
    ws["C10"] = 100
    ws["D10"] = 120
    return wb


def _workbook_with_aggregated_and_fs_blocks():
    """Aggregated left, Financial statements right — same FY headers (the FS overwrite bug)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "PL_Reconciliation"
    ws["B7"] = "Aggregated"
    ws["F7"] = "Financial statements"
    ws["B8"] = "FY24A"
    ws["C8"] = "FY25A"
    ws["F8"] = "FY24A"
    ws["G8"] = "FY25A"
    ws["A12"] = "Net sales"
    ws["B12"] = 111  # Aggregated FY24
    ws["C12"] = 122  # Aggregated FY25
    ws["F12"] = 999  # FS FY24 (must NOT be used)
    ws["G12"] = 888  # FS FY25
    ws["A13"] = "Gross sales"
    ws["B13"] = 211
    ws["C13"] = 222
    ws["F13"] = 777
    ws["G13"] = 666
    return wb


def test_sales_basis_labels():
    net = sales_basis_labels("net")
    assert net.metric_label == "Net sales"
    assert net.total_label == "Total Net sales"
    assert net.reported_label == "Reported Net sales"
    gross = sales_basis_labels("gross")
    assert gross.metric_label == "Gross sales"
    assert gross.total_label == "Total Gross sales"
    assert gross.reported_label == "Reported Gross sales"
    assert sales_basis_labels("unknown").sales_basis == "net"


def test_resolves_aggregated_net_sales_period_cells():
    refs = resolve_pl_net_sales_references(
        _workbook_with_pl_reconciliation(),
        ["FY24A", "FY25A"],
    )
    assert refs["FY24A"].coordinate == "C10"
    assert refs["FY25A"].formula == "='PL_Reconciliation'!D10"


def test_prefers_aggregated_columns_over_financial_statements():
    wb = _workbook_with_aggregated_and_fs_blocks()
    refs = resolve_pl_sales_references(wb, ["FY24A", "FY25A"], sales_basis="net")
    assert refs["FY24A"].coordinate == "B12"
    assert refs["FY25A"].coordinate == "C12"
    assert refs["FY24A"].formula == "='PL_Reconciliation'!B12"


def test_resolves_gross_sales_row_in_aggregated():
    wb = _workbook_with_aggregated_and_fs_blocks()
    refs = resolve_pl_sales_references(wb, ["FY24A", "FY25A"], sales_basis="gross")
    assert refs["FY24A"].coordinate == "B13"
    assert refs["FY25A"].coordinate == "C13"


def test_missing_gross_sales_row_returns_empty():
    wb = _workbook_with_pl_reconciliation()
    refs = resolve_pl_sales_references(wb, ["FY24A"], sales_basis="gross")
    assert refs == {}


def test_writes_reported_minus_total_and_na_fallback():
    wb = _workbook_with_pl_reconciliation()
    report = wb.create_sheet("TOP")
    report["E20"] = 95
    write_reconciliation_cells(
        wb,
        report,
        period_columns={"FY24A": 5},
        total_row=20,
        recon_row=21,
        reported_row=22,
    )
    assert report["E22"].value == "='PL_Reconciliation'!C10"
    assert report["E21"].value == "=E22-E20"

    no_pl = Workbook()
    target = no_pl.active
    write_reconciliation_cells(
        no_pl,
        target,
        period_columns={"FY24A": 2},
        total_row=2,
        recon_row=3,
        reported_row=4,
    )
    assert target["B3"].value == NA_VALUE
    assert target["B4"].value == NA_VALUE


def test_write_uses_gross_sales_basis():
    wb = _workbook_with_aggregated_and_fs_blocks()
    report = wb.create_sheet("GST")
    report["E20"] = 95
    write_reconciliation_cells(
        wb,
        report,
        period_columns={"FY24A": 5},
        total_row=20,
        recon_row=21,
        reported_row=22,
        sales_basis="gross",
    )
    assert report["E22"].value == "='PL_Reconciliation'!B13"
    assert report["E21"].value == "=E22-E20"


def test_bubble_embeds_helper_table_and_net_sales_reconciliation(tmp_path):
    wb = _workbook_with_pl_reconciliation()
    helper = pd.DataFrame(
        [{"label": "Product A", "revenue": 95_000, "gp": 40_000, "gm": 0.42, "bubblesize": 95_000}]
    )
    export_excel_with_images(
        {"excel": {"enabled": True, "sheet_name": "Margin analyses"}, "sales_basis": "net"},
        str(tmp_path / "missing-plot.png"),
        None,
        str(tmp_path / "output.xlsx"),
        "FY24A",
        wb=wb,
        save=False,
        helper_table=helper,
        total_net_sales=95_000,
    )
    ws = wb["Margin analyses"]
    assert ws["D55"].value == "Group"
    assert ws["E55"].value == "Net sales (FY24A)"
    assert ws["D58"].value == "Total Net sales"
    assert ws["E60"].value == "='PL_Reconciliation'!C10"
    assert ws["E59"].value == "=E60-E58"


def test_bubble_gross_sales_labels(tmp_path):
    wb = _workbook_with_aggregated_and_fs_blocks()
    helper = pd.DataFrame(
        [{"label": "Product A", "revenue": 95_000, "gp": 40_000, "gm": 0.42, "bubblesize": 95_000}]
    )
    export_excel_with_images(
        {"excel": {"enabled": True, "sheet_name": "Margin analyses"}, "sales_basis": "gross"},
        str(tmp_path / "missing-plot.png"),
        None,
        str(tmp_path / "output.xlsx"),
        "FY24A",
        wb=wb,
        save=False,
        helper_table=helper,
        total_net_sales=95_000,
    )
    ws = wb["Margin analyses"]
    assert ws["E55"].value == "Gross sales (FY24A)"
    assert ws["D58"].value == "Total Gross sales"
    assert ws["D60"].value == "Reported Gross sales"
    assert ws["E60"].value == "='PL_Reconciliation'!B13"
