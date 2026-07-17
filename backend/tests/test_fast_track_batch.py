from __future__ import annotations

import sys
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
import pandas as pd
import pytest

ROOT = Path(__file__).resolve().parents[2]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from scripts.fast_track_batch import (  # noqa: E402
    _apply_manifest_dates,
    _revenue_script_configs,
    run_fast_track,
    validate_manifest,
)


def test_manifest_honours_rasa_included_flags(tmp_path):
    master = tmp_path / "Master.xlsx"
    master.touch()
    manifest = {
        "output_path": str(tmp_path / "Fast_Track.xlsx"),
        "master_path": str(master),
        "sections": {
            "databook": {"included": False},
            "opos": {"included": False},
            "fte": {"included": True, "periods": [{"file_path": "fte.xlsx"}]},
            "fa": {"included": False},
            "revenue": {"included": False},
        },
    }
    normalized = validate_manifest(manifest)
    assert normalized["selected_sections"] == ["fte"]


def test_batch_uses_shared_workbook_and_writes_one_output(tmp_path):
    output = tmp_path / "Fast_Track.xlsx"
    master = tmp_path / "Master.xlsx"
    from openpyxl import Workbook

    mwb = Workbook()
    mwb.save(master)
    mwb.close()
    seen_ids: list[int] = []

    def run_fte(config, wb):
        seen_ids.append(id(wb))
        wb.create_sheet("FTE Development")

    def run_fa(config, wb):
        seen_ids.append(id(wb))
        wb.create_sheet("FA roll forward")

    result = run_fast_track(
        {
            "output_path": str(output),
            "master_path": str(master),
            "sections": {
                "databook": False,
                "opos": False,
                "fte": True,
                "fa": True,
                "revenue": False,
            },
        },
        section_runners={"fte": run_fte, "fa": run_fa},
    )

    assert len(set(seen_ids)) == 1
    assert result["completed_sections"] == ["fte", "fa"]
    wb = load_workbook(output)
    assert "FTE Development" in wb.sheetnames
    assert "FA roll forward" in wb.sheetnames
    wb.close()


def test_revenue_fast_track_builds_all_selected_analysis_configs():
    configs = _revenue_script_configs(
        {
            "selected_analyses": ["gst", "pvm", "top", "churn"],
            "file_path": "sales.xlsx",
            "sheet_name": "Data",
            "output_path": "/tmp/Fast_Track.xlsx",
            "output_file_path": "/tmp",
            "revenue_col": "Revenue",
            "invoice_col": "Invoice date",
            "entity_col": "Entity",
            "product_col": "Product",
            "customer_col": "Customer",
            "quantity_col": "Quantity",
            "cost_col": "Cost",
            "margin_mode": "cost",
            "calc_mode": "accrual",
            "contract_start_col": "Contract start",
            "contract_end_col": "Contract end",
        }
    )
    assert [script for script, _config in configs] == [
        "general_sales_table_MM_verformelt.py",
        "pvm_verformelt.py",
        "top_report.py",
        "churn_verformelt.py",
    ]
    pvm = configs[1][1]
    assert pvm["pvm_method"] == "three_components"
    assert pvm["calc_mode"] == "accrual"
    assert pvm["start_col"] == "Contract start"
    assert pvm["end_col"] == "Contract end"
    assert pvm["sales_basis"] == "net"
    assert pvm["total_label_revenue"] == "Total Net sales"
    assert configs[0][1]["total_label"] == "Total Net sales"
    assert configs[2][1]["total_label"] == "Total Net sales"
    assert configs[3][1]["total_label"] == "Total Net sales"


def test_revenue_fast_track_gross_sales_basis_labels():
    configs = _revenue_script_configs(
        {
            "selected_analyses": ["gst", "pvm", "top", "churn"],
            "file_path": "sales.xlsx",
            "sheet_name": "Data",
            "output_path": "/tmp/Fast_Track.xlsx",
            "output_file_path": "/tmp",
            "sales_basis": "gross",
            "revenue_col": "Revenue",
            "invoice_col": "Invoice date",
            "entity_col": "Entity",
            "product_col": "Product",
            "customer_col": "Customer",
            "quantity_col": "Quantity",
            "cost_col": "Cost",
            "margin_mode": "cost",
            "calc_mode": "invoice",
        }
    )
    assert all(cfg["sales_basis"] == "gross" for _, cfg in configs)
    assert configs[0][1]["total_label"] == "Total Gross sales"
    assert configs[1][1]["total_label_revenue"] == "Total Gross sales"
    assert configs[2][1]["total_label"] == "Total Gross sales"
    assert configs[3][1]["total_label"] == "Total Gross sales"


def test_revenue_fast_track_gpbridge_defaults():
    configs = _revenue_script_configs(
        {
            "selected_analyses": ["gpbridge"],
            "file_path": "sales.xlsx",
            "sheet_name": "Data",
            "output_path": "/tmp/Fast_Track.xlsx",
            "output_file_path": "/tmp",
            "sales_basis": "gross",
            "revenue_col": "Revenue",
            "invoice_col": "Invoice date",
            "product_col": "Product Name",
            "cost_col": "Cost",
            "margin_mode": "cost",
            "calc_mode": "accrual",
            "contract_start_col": "Start",
            "contract_end_col": "End",
            "filters": {"enabled": True, "rules": [{"col": "Region", "op": "eq", "value": "EMEA"}]},
            "fx_enabled": True,
            "fx_col": "FX",
            "period_mode": "FY",
        }
    )
    assert len(configs) == 1
    name, cfg = configs[0]
    assert name == "GP_Bridge.py"
    assert cfg["bridge"]["dim_col"] == "Product Name"
    assert cfg["profit_mode"] == "cost"
    assert cfg["cost_col"] == "Cost"
    assert cfg["calc_mode"] == "accrual"
    assert cfg["invoice_mapping_mode"] == "date"
    assert cfg["sales_basis"] == "gross"
    assert cfg["base_sheet_name"] == "GP Bridge"
    assert cfg["excel"]["sheet_name"] == "GP Bridge"
    assert cfg["subtitle_suffix"] == "by accrued amounts"
    assert cfg["filters"]["enabled"] is True
    assert cfg["apply_fx"] is True
    assert cfg["fx_col"] == "FX"
    assert cfg["period_mode"] == "FY"
    assert cfg["start_col"] == "Start"
    assert cfg["end_col"] == "End"


def test_revenue_fast_track_gpbridge_profit_mode_np_sheet():
    configs = _revenue_script_configs(
        {
            "selected_analyses": ["gpbridge"],
            "file_path": "sales.xlsx",
            "sheet_name": "Data",
            "output_path": "/tmp/Fast_Track.xlsx",
            "output_file_path": "/tmp",
            "revenue_col": "Revenue",
            "period_col": "FY",
            "date_mapping": "period",
            "product_col": "Product",
            "profit_col": "GP",
            "margin_mode": "profit",
            "calc_mode": "invoice",
        }
    )
    _, cfg = configs[0]
    assert cfg["profit_mode"] == "profit"
    assert cfg["profit_col"] == "GP"
    assert cfg["base_sheet_name"] == "NP Bridge"
    assert cfg["invoice_mapping_mode"] == "year"
    assert cfg["subtitle_suffix"] == "by invoiced amounts"


def test_revenue_fast_track_hbar_defaults():
    configs = _revenue_script_configs(
        {
            "selected_analyses": ["hbar"],
            "file_path": "sales.xlsx",
            "sheet_name": "Data",
            "output_path": "/tmp/Fast_Track.xlsx",
            "output_file_path": "/tmp",
            "sales_basis": "net",
            "revenue_col": "Revenue",
            "invoice_col": "Invoice date",
            "entity_col": "Entity",
            "segment_col": "Segment",
            "product_col": "Product Name",
            "calc_mode": "accrual",
            "contract_start_col": "Start",
            "contract_end_col": "End",
            "filters": {"enabled": True, "rules": [{"col": "Region", "op": "eq", "value": "EMEA"}]},
            "fx_enabled": True,
            "fx_col": "FX",
        }
    )
    assert len(configs) == 1
    name, cfg = configs[0]
    assert name == "horizontal_bars_hierachy.py"
    assert cfg["dimensions"] == {
        "bar_col": "Entity",
        "parent_col": "Segment",
        "child_col": "Product Name",
    }
    assert cfg["bars"]["top_n"] == 5
    assert cfg["bars"]["create_other_bar"] is True
    assert cfg["bars"]["other_label"] == "Rest"
    assert cfg["calc_mode"] == "accrual"
    assert cfg["invoice_mapping_mode"] == "date"
    assert cfg["base_sheet_name"] == "Net sales hierarchy"
    assert cfg["excel"]["sheet_name"] == "Net sales hierarchy"
    assert cfg["subtitle_suffix"] == "by accrued amounts"
    assert cfg["filters"]["enabled"] is True
    assert cfg["apply_fx"] is True
    assert cfg["fx_col"] == "FX"
    assert cfg["start_col"] == "Start"
    assert cfg["end_col"] == "End"
    assert cfg["value_col"] == "Revenue"


def test_revenue_fast_track_arrbridge_uses_global_calc_mode():
    configs = _revenue_script_configs(
        {
            "selected_analyses": ["arrbridge"],
            "file_path": "sales.xlsx",
            "sheet_name": "Data",
            "output_path": "/tmp/Fast_Track.xlsx",
            "output_file_path": "/tmp",
            "sales_basis": "gross",
            "revenue_col": "Revenue",
            "invoice_col": "Invoice date",
            "entity_col": "Entity",
            "product_col": "Product Name",
            "customer_col": "Customer",
            "calc_mode": "accrual",
            "use_arr_calculation": False,
            "contract_start_col": "Start",
            "contract_end_col": "End",
            "filters": {"enabled": True, "rules": [{"col": "Region", "op": "eq", "value": "EMEA"}]},
            "fx_enabled": True,
            "fx_col": "FX",
        }
    )
    assert len(configs) == 1
    name, cfg = configs[0]
    assert name == "ARR_Bridge.py"
    assert cfg["calc_mode"] == "accrual"
    assert cfg["group_col"] == "Entity"
    assert cfg["customer_col"] == "Customer"
    assert cfg["product_col"] == "Product Name"
    assert cfg["value_col"] == "Revenue"
    assert cfg["base_sheet_name"] == "ARR Bridge"
    assert cfg["table"] == "Gross sales component bridge"
    assert cfg["subtitle_suffix"] == "by accrued amounts"
    assert cfg["filters"]["enabled"] is True
    assert cfg["apply_fx"] is True
    assert cfg["start_col"] == "Start"
    assert cfg["end_col"] == "End"


def test_revenue_fast_track_arrbridge_use_arr_overrides_calc_mode():
    configs = _revenue_script_configs(
        {
            "selected_analyses": ["arrbridge"],
            "file_path": "sales.xlsx",
            "sheet_name": "Data",
            "output_path": "/tmp/Fast_Track.xlsx",
            "output_file_path": "/tmp",
            "revenue_col": "Revenue",
            "invoice_col": "Invoice date",
            "entity_col": "Entity",
            "product_col": "Product",
            "customer_col": "Customer",
            "calc_mode": "invoice",
            "use_arr_calculation": True,
            "contract_start_col": "Start",
            "contract_end_col": "End",
        }
    )
    _, cfg = configs[0]
    assert cfg["calc_mode"] == "ARR"
    assert cfg["table"] == "ARR Component Bridge"
    assert cfg["subtitle_suffix"] == "ARR run-rate"


def test_revenue_fast_track_arrbridge_rejects_arr_with_period_mapping():
    import pytest

    with pytest.raises(ValueError, match="date mapping"):
        _revenue_script_configs(
            {
                "selected_analyses": ["arrbridge"],
                "file_path": "sales.xlsx",
                "sheet_name": "Data",
                "output_path": "/tmp/Fast_Track.xlsx",
                "output_file_path": "/tmp",
                "revenue_col": "Revenue",
                "period_col": "FY",
                "date_mapping": "period",
                "entity_col": "Entity",
                "product_col": "Product",
                "customer_col": "Customer",
                "calc_mode": "invoice",
                "use_arr_calculation": True,
                "contract_start_col": "Start",
                "contract_end_col": "End",
            }
        )


def test_revenue_fast_track_applies_current_year_from_ltm_month():
    configs = _revenue_script_configs(
        {
            "selected_analyses": ["gst"],
            "file_path": "sales.xlsx",
            "sheet_name": "Data",
            "output_path": "/tmp/Fast_Track.xlsx",
            "output_file_path": "/tmp",
            "ltm_month": "2025-06",
            "revenue_col": "Revenue",
            "invoice_col": "Invoice date",
            "entity_col": "Entity",
            "product_col": "Product",
            "customer_col": "Customer",
            "cost_col": "Cost",
            "margin_mode": "cost",
            "calc_mode": "invoice",
        }
    )
    gst = configs[0][1]
    assert gst["output_workbook_path"] == "/tmp/Fast_Track.xlsx"
    assert gst["current_year"] == 2025
    assert gst["current_month"] == 6


def test_revenue_fast_track_falls_back_current_year_when_ltm_missing():
    configs = _revenue_script_configs(
        {
            "selected_analyses": ["gst"],
            "file_path": "sales.xlsx",
            "sheet_name": "Data",
            "output_path": "/tmp/Fast_Track.xlsx",
            "output_file_path": "/tmp",
            "revenue_col": "Revenue",
            "invoice_col": "Invoice date",
            "entity_col": "Entity",
            "product_col": "Product",
            "customer_col": "Customer",
            "cost_col": "Cost",
            "margin_mode": "cost",
            "calc_mode": "invoice",
        }
    )
    gst = configs[0][1]
    assert gst["current_year"]
    assert gst["current_month"]


def test_apply_manifest_dates_coerces_string_numbers():
    cfg: dict[str, Any] = {}
    _apply_manifest_dates(
        cfg,
        {
            "ltm_month": "2025-06",
            "first_fy": "2020",
            "fy_end_month": "12",
            "fy_end_day": "31",
        },
    )
    assert cfg["current_year"] == 2025
    assert cfg["current_month"] == 6
    assert cfg["first_fy"] == 2020
    assert cfg["fy_end_month"] == 12
    assert cfg["fy_end_day"] == 31
    assert isinstance(cfg["first_fy"], int)

    configs = _revenue_script_configs(
        {
            **cfg,
            "selected_analyses": ["churn"],
            "file_path": "/tmp/x.xlsx",
            "sheet_name": "Data",
            "output_path": "/tmp/out.xlsx",
            "output_file_path": "/tmp",
            "revenue_col": "R",
            "invoice_col": "I",
            "entity_col": "E",
            "product_col": "P",
            "customer_col": "C",
            "start_col": "S",
            "end_col": "E",
            "calc_mode": "invoice",
        }
    )
    churn_cfg = configs[0][1]
    assert isinstance(churn_cfg["first_fy"], int)
    assert isinstance(churn_cfg["fy_end_month"], int)


def test_legacy_adapter_keeps_shared_workbook_open(tmp_path):
    from openpyxl import Workbook

    from scripts.fast_track_batch import _legacy_workbook_adapter

    output = tmp_path / "FastTrack.xlsx"
    output.touch()
    wb = Workbook()
    wb.active.title = "Master_PL"
    wb.create_sheet("OPOS")

    with _legacy_workbook_adapter(wb, str(output)):
        import openpyxl

        loaded = openpyxl.load_workbook(str(output))
        loaded.create_sheet("GST")
        loaded.close()

    assert "GST" in wb.sheetnames
    assert "OPOS" in wb.sheetnames


def test_batch_blocks_secondary_workbook_save_to_session_path(tmp_path):
    from openpyxl import Workbook

    from scripts.fast_track_batch import _legacy_workbook_adapter

    output = tmp_path / "FastTrack.xlsx"
    shared = Workbook()
    shared.active.title = "Master_PL"
    rogue = Workbook()
    rogue.create_sheet("Only Revenue")

    with _legacy_workbook_adapter(shared, str(output)):
        import openpyxl

        loaded = openpyxl.load_workbook(str(output))
        assert loaded is shared
        with pytest.raises(RuntimeError, match="secondary workbook"):
            rogue.save(str(output))


def test_batch_preserves_master_when_revenue_runs(tmp_path):
    from openpyxl import Workbook, load_workbook

    master = tmp_path / "Master.xlsx"
    output = tmp_path / "FastTrack.xlsx"
    mwb = Workbook()
    mwb.active.title = "Master_PL"
    mwb.create_sheet("Master_BS")
    mwb.save(master)
    mwb.close()

    def run_revenue(config, wb):
        wb.create_sheet("General sales table")
        wb.create_sheet("__SOURCE__Sales")

    result = run_fast_track(
        {
            "output_path": str(output),
            "master_path": str(master),
            "sections": {
                "databook": False,
                "opos": False,
                "fte": False,
                "fa": False,
                "revenue": {"enabled": True, "config": {"selected_analyses": ["gst"]}},
            },
        },
        section_runners={"revenue": run_revenue},
    )
    assert result["completed_sections"] == ["revenue"]
    wb = load_workbook(output)
    assert "Master_PL" in wb.sheetnames
    assert "Master_BS" in wb.sheetnames
    assert "General sales table" in wb.sheetnames
    assert "__SOURCE__Sales" in wb.sheetnames
    wb.close()


def test_remove_invalid_output_placeholder_deletes_empty_file(tmp_path):
    from scripts.fast_track_batch import _remove_invalid_output_placeholder

    path = tmp_path / "FastTrack.xlsx"
    path.touch()
    _remove_invalid_output_placeholder(path)
    assert not path.exists()


def test_ensure_source_session_mode_skips_disk(tmp_path):
    from funktionssammlung import ensure_source_sheet_in_output

    output = tmp_path / "FastTrack.xlsx"
    output.touch()
    name = ensure_source_sheet_in_output(
        {"use_session_workbook": True, "sheet_name": "Data", "file_path": "x.xlsx"},
        str(output),
    )
    assert name == "__SOURCE__Sales"


def test_ensure_source_session_mode_populates_shared_workbook(tmp_path):
    from openpyxl import Workbook

    from funktionssammlung import ensure_source_sheet_in_output, session_workbook_context

    sales = tmp_path / "sales.xlsx"
    df = pd.DataFrame({"Revenue": [100], "Entity": ["A"]})
    df.to_excel(sales, sheet_name="Data", index=False)

    shared = Workbook()
    shared.active.title = "Master_PL"
    with session_workbook_context(shared):
        name = ensure_source_sheet_in_output(
            {
                "use_session_workbook": True,
                "sheet_name": "Data",
                "file_path": str(sales),
            },
            str(tmp_path / "FastTrack.xlsx"),
        )
    assert name == "__SOURCE__Sales"
    assert "__SOURCE__Sales" in shared.sheetnames


def test_remove_placeholder_fast_track_sheet():
    from openpyxl import Workbook

    from databook_workbook import remove_placeholder_fast_track_sheet

    wb = Workbook()
    wb.active.title = "Fast Track"
    wb.create_sheet("GST")
    remove_placeholder_fast_track_sheet(wb)
    assert "Fast Track" not in wb.sheetnames
    assert "GST" in wb.sheetnames


def test_revenue_report_col_offset_includes_spacers():
    from funktionssammlung import revenue_report_col_offset

    assert revenue_report_col_offset(3) == 5
    assert revenue_report_col_offset(1) == 3
    assert revenue_report_col_offset(3, configured=6) == 6


def test_revenue_report_helper_col_letters():
    from funktionssammlung import (
        REVENUE_REPORT_HELPER_START_COL,
        revenue_report_helper_col,
        revenue_report_helper_letter,
    )

    assert REVENUE_REPORT_HELPER_START_COL == 2
    assert revenue_report_helper_col(0) == 2
    assert revenue_report_helper_col(2) == 4
    assert revenue_report_helper_letter(0) == "B"
    assert revenue_report_helper_letter(1) == "C"
    assert revenue_report_helper_letter(2) == "D"


def test_resolve_sheet_name_uses_first_sheet_when_missing(tmp_path):
    from openpyxl import Workbook

    from scripts.fast_track_batch import _resolve_sheet_name

    path = tmp_path / "sales.xlsx"
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Daten"
    wb.save(path)
    wb.close()
    assert _resolve_sheet_name(str(path), "") == "Daten"
    assert _resolve_sheet_name(str(path), None) == "Daten"
