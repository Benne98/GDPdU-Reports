"""Fast Track workbook sheet grouping (section dividers, colors, order)."""
from __future__ import annotations

from pathlib import Path
import sys

ROOT = Path(__file__).resolve().parents[2]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from openpyxl import Workbook

from databook_workbook import (
    SHEET_GROUP_SPECS,
    apply_group_tab_colors,
    ensure_group_section_sheets,
    entity_order_from_group_recon,
    reorder_workbook_sheets,
    section_sheet_name,
)


def test_section_sheet_names_and_colors():
    wb = Workbook()
    wb.remove(wb.active)
    wb.create_sheet("Lead_IS")
    names = ensure_group_section_sheets(wb)
    assert names["executive"] == "==> Executive Summary"
    assert names["earnings"] == "==> Earnings"
    assert names["financial"] == "==> Assets"
    assert names["liquidity"] == "==> Liquidity"
    assert names["source"] == "==> Source"
    assert len(names) == 6

    ws = wb["==> Earnings"]
    assert ws["B2"].value == "Finssentials"
    assert ws["B3"].value == "Earnings"
    assert str(ws["B2"].font.color.rgb).upper().endswith("FFFFFF")
    assert str(ws["A1"].fill.fgColor.rgb).upper().endswith("2E5A8A")
    assert str(ws.sheet_properties.tabColor.rgb).upper().endswith("2E5A8A")


def test_full_group_order_earnings_financial_liquidity_appendix():
    wb = Workbook()
    wb.remove(wb.active)
    for name in (
        "Master_PL",
        "Master_BS",
        "Cashflow",
        "Margin analyses",
        "pvm",
        "Lead_IS",
        "General sales table",
        "TOP",
        "Churn",
        "Net sales breakdown",
        "FTE Development",
        "__SOURCE__Sales",
        "__SOURCE__FTE_FY24A",
        "Lead_BS",
        "BS_Bucket",
        "Working_Capital",
        "FA roll forward",
        "Trade creditors aging",
        "Trade debtors aging",
        "__SOURCE__FA_Dec23A",
        "__SOURCE__AR_Jul25A",
        "PL_Reconciliation",
        "BS_Reconciliation",
        "Beta_BS_Reconciliation",
        "Beta_PL_Reconciliation",
        "Alpha_BS_Reconciliation",
        "Alpha_PL_Reconciliation",
    ):
        wb.create_sheet(name)

    # Group PL recon entity block order: Alpha then Beta
    pl = wb["PL_Reconciliation"]
    pl.cell(7, 2, "Alpha")
    pl.cell(7, 8, "Beta")
    pl.cell(7, 14, "Aggregated")

    reorder_workbook_sheets(wb)
    names = wb.sheetnames

    assert names[0] == "==> Executive Summary"
    assert names[1] == "Net sales breakdown"
    assert names[2] == "==> Earnings"
    earnings = names[names.index("==> Earnings") + 1 : names.index("==> Assets")]
    assert earnings[0] == "Lead_IS"
    assert set(earnings[1:5]) == {"General sales table", "pvm", "TOP", "Churn"}
    assert earnings[5] == "Margin analyses"
    assert earnings[6] == "FTE Development"
    assert "__SOURCE__Sales" not in earnings

    financial = names[names.index("==> Assets") + 1 : names.index("==> Liquidity")]
    assert financial[:5] == [
        "Lead_BS",
        "BS_Bucket",
        "Working_Capital",
        "FA roll forward",
        "Trade debtors aging",
    ]
    assert financial[5] == "Trade creditors aging"
    assert "__SOURCE__AR_Jul25A" not in financial

    liquidity = names[names.index("==> Liquidity") + 1 : names.index("==> Appendix")]
    assert liquidity == ["Cashflow"]

    appendix = names[names.index("==> Appendix") + 1 : names.index("==> Source")]
    assert appendix[:2] == ["BS_Reconciliation", "PL_Reconciliation"]
    assert appendix[2:6] == [
        "Alpha_BS_Reconciliation",
        "Alpha_PL_Reconciliation",
        "Beta_BS_Reconciliation",
        "Beta_PL_Reconciliation",
    ]
    assert "Master_BS" not in appendix

    source = names[names.index("==> Source") + 1 :]
    assert source[0:2] == ["Master_BS", "Master_PL"]
    assert "__SOURCE__Sales" in source
    assert "__SOURCE__FTE_FY24A" in source
    assert "__SOURCE__AR_Jul25A" in source
    assert "__SOURCE__FA_Dec23A" in source
    assert wb["__SOURCE__Sales"].sheet_state == "visible"
    assert wb["Master_BS"].sheet_state == "visible"


def test_entity_order_from_group_recon():
    wb = Workbook()
    wb.remove(wb.active)
    wb.create_sheet("PL_Reconciliation")
    wb.create_sheet("Alpha_BS_Reconciliation")
    wb.create_sheet("Beta_PL_Reconciliation")
    wb["PL_Reconciliation"].cell(7, 1, "Beta")
    wb["PL_Reconciliation"].cell(7, 5, "Alpha")
    wb["PL_Reconciliation"].cell(7, 9, "Aggregated")
    assert entity_order_from_group_recon(wb) == ["Beta", "Alpha"]


def test_tab_colors_match_group():
    wb = Workbook()
    wb.remove(wb.active)
    wb.create_sheet("Lead_IS")
    wb.create_sheet("Cashflow")
    ensure_group_section_sheets(wb)
    apply_group_tab_colors(wb)
    color = {spec["key"]: spec["color"] for spec in SHEET_GROUP_SPECS}
    assert wb["Lead_IS"].sheet_properties.tabColor.rgb.endswith(color["earnings"])
    assert wb["Cashflow"].sheet_properties.tabColor.rgb.endswith(color["liquidity"])
    assert wb[section_sheet_name("Executive Summary")].sheet_properties.tabColor.rgb.endswith(
        color["executive"]
    )
