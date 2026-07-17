"""Tests for databook Excel layout helpers."""
from __future__ import annotations

import sys
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import Workbook

SCRIPTS = Path(__file__).resolve().parent.parent.parent / "scripts"
if str(SCRIPTS) not in sys.path:
    sys.path.insert(0, str(SCRIPTS))

from databook_excel_layout import (  # noqa: E402
    AGGREGATED_BLOCK_TITLE,
    HELPER_COLS_BS,
    HELPER_COLS_PL,
    LAYOUT_BS,
    LAYOUT_NA,
    LAYOUT_PL,
    DatabookLayout,
    RECON_BLOCK_TITLE_ROW,
    RECON_HEADER_ROW,
    bs_recon_aggregated_ref,
    check_row_groups_after_table,
    collapse_check_portfolio,
    databook_column_layout,
    find_bs_recon_row_contains,
    find_recon_block_col_by_header,
    find_recon_block_start_col,
    find_recon_block_year_cols,
    header_band_columns,
    paint_grey_white_canvas,
    paint_header_band,
    prune_zero_value_rows,
)


def test_databook_column_layout_pl_geometry():
    layout = databook_column_layout(HELPER_COLS_PL)
    assert layout.helper_start_col == 1
    assert layout.map_start_col == 2
    assert layout.map_end_col == 5
    assert layout.helper_end_col == 6
    assert layout.spacer_col == 6
    assert layout.pos_col == 7
    assert layout.first_value_col == 8


def test_databook_column_layout_bs_geometry():
    layout = databook_column_layout(HELPER_COLS_BS)
    assert layout.map_start_col == 2
    assert layout.map_end_col == 6
    assert layout.helper_end_col == 7
    assert layout.pos_col == 8


def test_header_band_columns_skip_helper_block():
    layout = LAYOUT_PL
    cols = header_band_columns(layout, period_cols=[8, 9, 10])
    assert 1 not in cols
    assert layout.helper_end_col not in cols
    assert cols == [2, 3, 4, 5, 7, 8, 9, 10]


def test_paint_header_band_does_not_touch_helper_block():
    wb = Workbook()
    ws = wb.active
    layout = LAYOUT_PL
    paint_header_band(
        ws,
        layout,
        header_rows=[7, 8],
        period_cols=[8, 9],
    )
    assert ws.cell(8, 1).fill.fill_type is None or ws.cell(8, 1).fill.fgColor.rgb in (
        None,
        "00000000",
    )
    assert ws.cell(8, 2).fill.fgColor.rgb.endswith("F8FAFC")
    assert ws.cell(8, layout.helper_end_col).fill.fill_type is None or (
        ws.cell(8, layout.helper_end_col).fill.fgColor.rgb in (None, "00000000")
    )


def test_paint_grey_white_canvas_greys_full_helper_block():
    wb = Workbook()
    ws = wb.active
    layout = LAYOUT_BS
    paint_grey_white_canvas(ws, layout, last_row=5, last_col=12)
    # A through spacer G grey
    assert ws.cell(3, 1).fill.fgColor.rgb.endswith("F1F5F9")
    assert ws.cell(3, layout.helper_end_col).fill.fgColor.rgb.endswith("F1F5F9")
    assert ws.cell(3, 2).fill.fgColor.rgb.endswith("F1F5F9")
    assert ws.cell(3, layout.pos_col).fill.fgColor.rgb.endswith("FFFFFF")


def test_prune_zero_value_rows_removes_all_zero_details():
    df = pd.DataFrame(
        {
            "L3": ["Revenue", "Revenue", "Costs", "Costs"],
            "L4": ["A", "B", "C", "D"],
            "L5": ["reported", "reported", "reported", "reported"],
            "FY23A": [100, 0, 0, 50],
            "FY24A": [0, 0, 0, 0],
        }
    )
    struct = [
        {"type": "detail", "L3": "Revenue", "L4": "A"},
        {"type": "detail", "L3": "Revenue", "L4": "B"},
        {"type": "subtotal", "L3": "Revenue", "tech_label": "Revenue"},
        {"type": "detail", "L3": "Costs", "L4": "C"},
        {"type": "detail", "L3": "Costs", "L4": "D"},
        {"type": "subtotal", "L3": "Costs", "tech_label": "Costs"},
    ]
    pruned = prune_zero_value_rows(
        struct,
        df,
        ["FY23A", "FY24A"],
        source_col="L5",
        reported_value="reported",
    )
    labels = {(r.get("L3"), r.get("L4"), r.get("type")) for r in pruned}
    assert ("Revenue", "B", "detail") not in labels
    assert ("Costs", "C", "detail") not in labels
    assert ("Revenue", "A", "detail") in labels
    assert ("Costs", "D", "detail") in labels


def test_check_row_groups_after_table_adjacent_within_group():
    rows = check_row_groups_after_table(100, [1, 2, 2])
    assert rows == [102, 104, 105, 107, 108]


def test_collapse_check_portfolio_single_group_with_blank_rows():
    wb = Workbook()
    ws = wb.active
    check_rows = [20, 22, 24]
    collapse_check_portfolio(ws, check_rows)
    for rr in range(20, 25):
        assert ws.row_dimensions[rr].outlineLevel == 1
        assert ws.row_dimensions[rr].hidden is True
    assert ws.row_dimensions[24].collapsed is True


def _build_minimal_bs_recon_ws(ws):
    ws.cell(RECON_BLOCK_TITLE_ROW, 10, "Aggregated")
    ws.cell(RECON_BLOCK_TITLE_ROW, 14, "Financial statements")
    ws.cell(RECON_HEADER_ROW, 10, "Dec24A")
    ws.cell(RECON_HEADER_ROW, 11, "Dec23A")
    ws.cell(12, LAYOUT_NA.pos_col, "Trade receivables")
    ws.cell(12, 10, 500)
    ws.cell(12, 14, 999)


def test_find_recon_block_col_by_header_aggregated_only():
    wb = Workbook()
    ws = wb.active
    _build_minimal_bs_recon_ws(ws)
    col = find_recon_block_col_by_header(ws, AGGREGATED_BLOCK_TITLE, "Dec24A")
    assert col == 10
    assert find_recon_block_col_by_header(ws, AGGREGATED_BLOCK_TITLE, "Dec23A") == 11


def test_find_bs_recon_row_and_ref():
    wb = Workbook()
    ws = wb.active
    _build_minimal_bs_recon_ws(ws)
    row = find_bs_recon_row_contains(ws, "trade receivables")
    assert row == 12
    assert bs_recon_aggregated_ref(row, 10) == "='BS_Reconciliation'!J12"


def test_find_recon_block_year_cols_and_start():
    wb = Workbook()
    ws = wb.active
    _build_minimal_bs_recon_ws(ws)
    assert find_recon_block_start_col(ws, AGGREGATED_BLOCK_TITLE) == 10
    cols = find_recon_block_year_cols(ws, AGGREGATED_BLOCK_TITLE, ["Dec24A", "Dec23A"])
    assert cols == {"Dec24A": 10, "Dec23A": 11}


def test_find_recon_aggregated_year_cols():
    from databook_excel_layout import find_recon_aggregated_year_cols

    wb = Workbook()
    ws = wb.active
    _build_minimal_bs_recon_ws(ws)
    cols = find_recon_aggregated_year_cols(ws, ["Dec24A", "Dec23A"])
    assert cols == {"Dec24A": 10, "Dec23A": 11}


def test_find_recon_block_col_for_period_label_year_suffix_fallback():
    from databook_excel_layout import find_recon_block_col_for_period_label

    wb = Workbook()
    ws = wb.active
    _build_minimal_bs_recon_ws(ws)
    ws.cell(RECON_HEADER_ROW, 10, "Jul24A")
    assert find_recon_block_col_for_period_label(ws, AGGREGATED_BLOCK_TITLE, "Dec24A") == 10
    assert find_recon_block_col_for_period_label(ws, AGGREGATED_BLOCK_TITLE, "Jul24A") == 10


def test_find_recon_block_col_for_period_label_no_cross_year_fallback():
    from databook_excel_layout import find_recon_block_col_for_period_label

    wb = Workbook()
    ws = wb.active
    _build_minimal_bs_recon_ws(ws)
    ws.cell(RECON_HEADER_ROW, 10, "Dec23A")
    ws.cell(RECON_HEADER_ROW, 11, "Dec24A")
    assert find_recon_block_col_for_period_label(ws, AGGREGATED_BLOCK_TITLE, "Jul25A") is None


def test_apply_databook_row_border_band_skips_spacer_cols():
    from openpyxl.styles import Border, Side

    from databook_excel_layout import LAYOUT_PL, apply_databook_row_border_band

    wb = Workbook()
    ws = wb.active
    border = Border(top=Side(style="thin"))
    spacer_col = 20
    apply_databook_row_border_band(
        ws,
        10,
        LAYOUT_PL,
        border,
        last_used_col=25,
        skip_cols={spacer_col},
    )
    assert ws.cell(10, LAYOUT_PL.pos_col).border.top.style == "thin"
    assert ws.cell(10, spacer_col).border.top is None
