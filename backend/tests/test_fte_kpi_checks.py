"""Tests for FTE KPI section and PL_Reconciliation checks."""
from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import Workbook

ROOT = Path(__file__).resolve().parent.parent.parent
SCRIPTS = ROOT / "scripts"
BACKEND = ROOT / "backend"
for p in (ROOT, SCRIPTS, BACKEND):
    if str(p) not in sys.path:
        sys.path.insert(0, str(p))

from databook_excel_layout import (  # noqa: E402
    AGGREGATED_BLOCK_TITLE,
    LAYOUT_PL,
    RECON_BLOCK_TITLE_ROW,
    RECON_HEADER_ROW,
)
from FTE_payroll import (  # noqa: E402
    PeriodColumn,
    _write_fte_kpi_and_checks,
    fte_column_layout,
)


def _seed_pl_recon(wb: Workbook) -> None:
    ws = wb.create_sheet("PL_Reconciliation")
    ws.cell(RECON_BLOCK_TITLE_ROW, 10, AGGREGATED_BLOCK_TITLE)
    ws.cell(RECON_HEADER_ROW, 10, "FY24A")
    ws.cell(20, LAYOUT_PL.pos_col, "Total output")
    ws.cell(21, LAYOUT_PL.pos_col, "Personnel expenses")
    ws.cell(20, 10, 5000)
    ws.cell(21, 10, 800)


def test_fte_kpi_section_writes_pct_and_pl_recon_refs():
    wb = Workbook()
    _seed_pl_recon(wb)
    ws = wb.active
    layout = fte_column_layout(1)
    period_cols = [PeriodColumn(col_idx=layout.first_data_col, header="FY24A", period_idx=0)]
    total_pe_row = 30

    ws.cell(total_pe_row, layout.pos_col, "Personnel expenses")
    ws.cell(total_pe_row, layout.first_data_col, 750)

    meta = _write_fte_kpi_and_checks(
        wb,
        ws,
        layout,
        period_cols,
        total_pe_row=total_pe_row,
        formula_mode=True,
    )

    assert ws.cell(meta["kpi_title_row"], layout.pos_col).value == "KPIs in % of total output"
    assert ws.cell(meta["kpi_pe_row"], layout.pos_col).value == "Personnel expenses"
    pct_formula = ws.cell(meta["kpi_pe_row"], layout.first_data_col).value
    assert str(pct_formula).startswith("=IFERROR(")
    assert "*100" in str(pct_formula)

    helper_formula = ws.cell(meta["total_output_helper_row"], layout.first_data_col).value
    assert helper_formula == "='PL_Reconciliation'!J20"
    assert ws.row_dimensions[meta["total_output_helper_row"]].hidden is True

    src_formula = ws.cell(meta["check_src_row"], layout.first_data_col).value
    assert src_formula == "='PL_Reconciliation'!J21"
    delta_formula = ws.cell(meta["check_delta_row"], layout.first_data_col).value
    assert str(delta_formula).startswith("=")
    assert str(delta_formula).endswith(f"-{ws.cell(meta['check_src_row'], layout.first_data_col).coordinate}")


def test_fte_kpi_row_layout_after_personnel_total():
    from databook_excel_layout import check_row_groups_after_table

    total_pe_row = 50
    kpi_title = total_pe_row + 1
    kpi_pe = total_pe_row + 2
    helper = kpi_pe + 3
    check_rows = check_row_groups_after_table(helper, [2])
    assert kpi_title == 51
    assert kpi_pe == 52
    assert helper == 55
    assert check_rows == [57, 58]
