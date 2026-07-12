"""Unit tests for extract_master_entities."""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook

from databook_helpers import extract_master_entities


def _write_master(path: Path, entities: list[str], sheet: str = "Master_BS") -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    ws.append(["Entity", "L2", "L3", "L4", "FY2022"])
    for ent in entities:
        ws.append([ent, "Fixed assets", "FA", "Detail", 100])
    wb.save(path)


def test_extract_master_entities_from_bs_sheet(tmp_path):
    master = tmp_path / "master.xlsx"
    _write_master(master, ["Atlas", "Beta", "Atlas", "Consolidation", "Gamma"])
    assert extract_master_entities(master) == ["Atlas", "Beta", "Gamma"]


def test_extract_master_entities_fallback_pl_sheet(tmp_path):
    master = tmp_path / "master.xlsx"
    wb = Workbook()
    ws_bs = wb.active
    ws_bs.title = "Other"
    ws_pl = wb.create_sheet("Master_PL")
    ws_pl.append(["Entity", "L2", "L3", "L4", "FY2022"])
    ws_pl.append(["EntityA", "Revenue", "R", "D", 50])
    wb.save(master)
    assert extract_master_entities(master) == ["EntityA"]


def test_extract_master_entities_empty_when_missing_sheets(tmp_path):
    master = tmp_path / "empty.xlsx"
    Workbook().save(master)
    assert extract_master_entities(master) == []
