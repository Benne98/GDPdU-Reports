"""SUMIFS partner keys with *, ?, ~ must match source text literally."""

from __future__ import annotations

import sys
from pathlib import Path

import pytest
from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[2]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from funktionssammlung import escape_excel_sumifs_wildcards  # noqa: E402
from opos import (  # noqa: E402
    KEY_COL,
    MISSING_GROUP_LABEL,
    write_opos_formula_key_column,
)


def test_escape_excel_sumifs_wildcards():
    assert escape_excel_sumifs_wildcards("**Schlussrechnung**") == "~*~*Schlussrechnung~*~*"
    assert escape_excel_sumifs_wildcards("*Schlussrechnung") == "~*Schlussrechnung"
    assert escape_excel_sumifs_wildcards("a?b") == "a~?b"
    assert escape_excel_sumifs_wildcards("a~b") == "a~~b"
    assert escape_excel_sumifs_wildcards("plain text") == "plain text"


def test_write_opos_formula_key_column_escapes_wildcards():
    wb = Workbook()
    ws = wb.active
    layout = {
        "key_col": KEY_COL,
        "header_row": 8,
        "helper_right_col": 4,
    }
    cfg = {"columns": {"partner_id": "Text"}}
    partner_row_names = {
        88: "**Schlussrechnung**",
        89: "*Schlussrechnung",
        9: MISSING_GROUP_LABEL,
    }
    write_opos_formula_key_column(ws, layout, cfg, partner_row_names, footer_rows={"total": 167})

    assert ws.cell(88, KEY_COL).value == "~*~*Schlussrechnung~*~*"
    assert ws.cell(89, KEY_COL).value == "~*Schlussrechnung"
    assert ws.cell(9, KEY_COL).value == MISSING_GROUP_LABEL


def test_escaped_schlussrechnung_keys_sum_without_overlap():
    rows = {
        "**Schlussrechnung**": 34858.51,
        "*Schlussrechnung": 18434.61,
    }

    def literal_sum(partner_key: str) -> float:
        crit = escape_excel_sumifs_wildcards(partner_key)
        for source_text, amount in rows.items():
            if escape_excel_sumifs_wildcards(source_text) == crit:
                return amount
        return 0.0

    footer = literal_sum("**Schlussrechnung**") + literal_sum("*Schlussrechnung")
    assert footer == pytest.approx(53293.12)
