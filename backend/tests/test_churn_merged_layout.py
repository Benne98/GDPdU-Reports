"""Merged churn bridge: total row order, latest-FY sort, bucket formulas."""
import unittest

import pandas as pd
from openpyxl import Workbook

from Churn import (
    apply_top_buckets,
    build_merged_horizontal_bridge,
    finalize_merged_bridge_layout,
    normalize_config,
    write_churn_export_to_workbook,
)
from churn_verformelt import apply_churn_formulas


def _base_cfg(**overrides):
    cfg = {
        "fy_end_year": 2025,
        "fy_end_month": 12,
        "fy_end_day": 31,
        "first_fy": 2022,
        "current_year": 2025,
        "current_month": 12,
        "value_col": "Amount",
        "customer_col": "Customer",
        "product_col": "Product",
        "start_col": "Start",
        "end_col": "End",
        "invoice_col": "Invoice",
        "group_cols": ["Entity"],
        "total_label": "Total",
        "sort": {"top_level_by": "ARR2", "top_level_desc": True},
    }
    cfg.update(overrides)
    return normalize_config(cfg)


def _mini_df() -> pd.DataFrame:
    return pd.DataFrame({
        "Entity": ["A", "B", "C"],
        "Customer": ["C1", "C2", "C3"],
        "Product": ["P1", "P2", "P3"],
        "Amount": [12000.0, 24000.0, 36000.0],
        "Start": pd.to_datetime(["2022-01-01", "2023-01-01", "2024-01-01"]),
        "End": pd.to_datetime(["2025-12-31", "2025-12-31", "2025-12-31"]),
        "Invoice": pd.to_datetime(["2022-01-01", "2023-01-01", "2024-01-01"]),
    })


class TestChurnMergedLayout(unittest.TestCase):
    def test_merged_total_is_last_row(self):
        merged = build_merged_horizontal_bridge(_mini_df(), _base_cfg())
        self.assertFalse(merged.empty)
        self.assertEqual(merged.iloc[-1]["row_type"], "total")
        total_idx = merged.index[merged["row_type"] == "total"]
        self.assertEqual(len(total_idx), 1)
        self.assertEqual(total_idx[0], len(merged) - 1)

    def test_merged_sort_by_latest_fy(self):
        df = pd.DataFrame({
            "Entity": ["Low", "High", "Mid"],
            "Customer": ["C1", "C2", "C3"],
            "Product": ["P1", "P2", "P3"],
            "Amount": [1000.0, 9000.0, 5000.0],
            "Start": pd.to_datetime(["2022-01-01"] * 3),
            "End": pd.to_datetime(["2025-12-31"] * 3),
            "Invoice": pd.to_datetime(["2022-01-01"] * 3),
        })
        merged = build_merged_horizontal_bridge(df, _base_cfg())
        leaves = merged[merged["row_type"] == "leaf"]["Entity"].tolist()
        fy_latest = sorted(c for c in merged.columns if str(c).startswith("FY"))[-1]
        fy_pos = [i for i, c in enumerate(merged.columns) if c == fy_latest][-1]
        leaf_vals = merged[merged["row_type"] == "leaf"].set_index("Entity").iloc[:, fy_pos]
        self.assertEqual(leaves, leaf_vals.sort_values(ascending=False).index.tolist())

    def test_finalize_merged_sort_uses_rightmost_fy(self):
        cfg = _base_cfg()
        raw = pd.DataFrame({
            "Entity": ["A", "B"],
            "FY22A": [1, 100],
            "Upsell": [0, 0],
            "Downsell": [0, 0],
            "Cross-sell": [0, 0],
            "Lost": [0, 0],
            "NRR": [0, 0],
            "New": [0, 0],
            "FY23A": [50, 5],
            "FY24A": [200, 10],
        })
        out = finalize_merged_bridge_layout(raw, cfg)
        leaves = out[out["row_type"] == "leaf"]["Entity"].tolist()
        self.assertEqual(leaves[0], "A")

    def test_bucket_formulas_fy(self):
        cfg = _base_cfg(
            formula_mode=True,
            top_bucket={
                "enabled": True,
                "bucket_mode": "number",
                "numbers": (2,),
                "create_other_bucket": False,
            },
        )
        merged = build_merged_horizontal_bridge(_mini_df(), cfg)
        wb = Workbook()
        ws = wb.active
        ws.title = "Churn"
        write_churn_export_to_workbook(wb, merged, "Churn", cfg, formula_mode=True)
        from Churn import format_churn_excel

        layout = format_churn_excel(wb, "Churn", merged, cfg, formula_mode=True)
        layout["report_sheet"] = "Churn"
        wb.create_sheet("Source")
        wb["Source"].cell(row=1, column=1).value = "Entity"

        bucket_rows = [
            r
            for r in range(layout["data_start_row"], layout["table_bottom_row"] + 1)
            if str(wb["Churn"].cell(row=r, column=layout["row_type_col"]).value) == "bucket"
        ]
        self.assertGreater(len(bucket_rows), 0)

        ws = wb["Churn"]
        fy_col = layout["fy_col_indices"][0][1]
        br = bucket_rows[0]
        apply_churn_formulas(wb, layout, cfg, "Source", {2022: "Z", 2023: "AA", 2024: "AB", 2025: "AC"})
        val = ws.cell(row=br, column=fy_col).value
        self.assertIsInstance(val, str)
        self.assertTrue(str(val).startswith("=SUM(") or str(val).startswith('=IF(COUNT('))


if __name__ == "__main__":
    unittest.main()
