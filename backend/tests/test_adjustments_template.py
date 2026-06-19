"""Adjustments template: yearly FY columns only, formatting, and CLI smoke."""
from __future__ import annotations

import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))
if str(ROOT / "backend") not in sys.path:
    sys.path.insert(0, str(ROOT / "backend"))
if str(ROOT / "scripts") not in sys.path:
    sys.path.insert(0, str(ROOT / "scripts"))

from databook_helpers import (  # noqa: E402
    CONSOLIDATION_META_COLS,
    build_adjustments_template_from_config,
)
from databook_periods import (  # noqa: E402
    fallback_yearly_columns,
    filter_master_period_columns,
    read_master_bs_headers,
    resolve_adjustments_period_columns,
)
from gst_excel_theme import THEME  # noqa: E402


def _write_sample_master(path: Path) -> None:
    cols = list(CONSOLIDATION_META_COLS) + [
        "L5",
        "FY24A",
        "FY25A",
        "Jan-2024",
        "Feb-2024",
        "Dec-2025",
        "Comments",
    ]
    df = pd.DataFrame(columns=cols)
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="Master_BS", index=False)


class TestAdjustmentsTemplate(unittest.TestCase):
    def test_resolve_adjustments_yearly_from_master(self):
        with tempfile.TemporaryDirectory() as tmp:
            master = Path(tmp) / "master.xlsx"
            _write_sample_master(master)
            cols = resolve_adjustments_period_columns({"master_path": str(master)})
        self.assertEqual(cols, ["FY24A", "FY25A"])
        self.assertTrue(all(str(c).startswith("FY") for c in cols))

    def test_fallback_yearly_from_date_settings(self):
        cols = fallback_yearly_columns(
            {"first_fy": 2024, "ltm_month": "2025-12", "fy_end_month": 12}
        )
        self.assertEqual(cols, ["FY24A", "FY25A"])
        self.assertNotIn("FY22A", cols)

    def test_build_yearly_from_master_no_month_columns(self):
        with tempfile.TemporaryDirectory() as tmp:
            master = Path(tmp) / "master.xlsx"
            _write_sample_master(master)
            headers = read_master_bs_headers(master)
            out = Path(tmp) / "adj.xlsx"
            build_adjustments_template_from_config(
                {
                    "master_path": str(master),
                    "output_path": str(out),
                }
            )
            cols = list(pd.read_excel(out, sheet_name="Adjustments", nrows=0).columns)
            self.assertEqual(
                cols,
                list(CONSOLIDATION_META_COLS) + filter_master_period_columns(headers, "yearly"),
            )
            period_cols = [c for c in cols if c not in CONSOLIDATION_META_COLS]
            self.assertTrue(all(str(c).startswith("FY") for c in period_cols))
            self.assertNotIn("Jan-2024", cols)

    def test_fallback_without_master(self):
        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp) / "adj.xlsx"
            build_adjustments_template_from_config(
                {
                    "output_path": str(out),
                    "first_fy": 2024,
                    "ltm_month": "2025-12",
                    "fy_end_month": 12,
                    "fiscal_start_month": 1,
                }
            )
            cols = list(pd.read_excel(out, sheet_name="Adjustments", nrows=0).columns)
            period_cols = [c for c in cols if c not in CONSOLIDATION_META_COLS]
            self.assertEqual(period_cols, ["FY24A", "FY25A"])
            self.assertNotIn("FY22A", period_cols)

    def test_header_formatting_and_canvas(self):
        with tempfile.TemporaryDirectory() as tmp:
            master = Path(tmp) / "master.xlsx"
            _write_sample_master(master)
            out = Path(tmp) / "fmt.xlsx"
            build_adjustments_template_from_config(
                {
                    "master_path": str(master),
                    "output_path": str(out),
                    "first_fy": 2024,
                    "ltm_month": "2025-12",
                }
            )
            wb = load_workbook(out)
            ws = wb["Adjustments"]
            cell = ws.cell(1, 1)
            self.assertEqual(cell.fill.start_color.rgb, THEME.fill_header.start_color.rgb)
            self.assertEqual(cell.font.name, THEME.font_header.name)
            empty_header = ws.cell(1, 50)
            self.assertEqual(empty_header.fill.start_color.rgb, THEME.fill_header.start_color.rgb)
            body = ws.cell(2, 1)
            self.assertEqual(body.fill.start_color.rgb, THEME.fill_white.start_color.rgb)
            self.assertEqual(ws.cell(101, 100).fill.start_color.rgb, THEME.fill_white.start_color.rgb)
            self.assertGreater(ws.column_dimensions["C"].width, 20)

    def test_cli_script(self):
        with tempfile.TemporaryDirectory() as tmp:
            master = Path(tmp) / "master.xlsx"
            out = Path(tmp) / "tpl.xlsx"
            _write_sample_master(master)
            cfg = {
                "session_id": "smoke",
                "output_folder": tmp,
                "master_path": str(master),
                "output_path": str(out),
                "first_fy": 2024,
                "ltm_month": "2025-12",
            }
            cfg_path = Path(tmp) / "config.json"
            cfg_path.write_text(json.dumps(cfg), encoding="utf-8")
            proc = subprocess.run(
                [sys.executable, str(ROOT / "adjustments_template.py"), str(cfg_path)],
                capture_output=True,
                text=True,
                cwd=str(ROOT),
            )
            self.assertEqual(proc.returncode, 0, proc.stderr or proc.stdout)
            self.assertTrue(out.is_file())
            cols = list(pd.read_excel(out, sheet_name="Adjustments", nrows=0).columns)
            self.assertIn("FY24A", cols)
            self.assertNotIn("FY22A", cols)
            self.assertNotIn("Jan-2024", cols)


if __name__ == "__main__":
    unittest.main()
