"""Smoke tests for Fast Track PDF report extraction and ordering."""

from __future__ import annotations

import sys
from pathlib import Path

import pytest
from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[2]
SCRIPTS = ROOT / "scripts"
for p in (str(ROOT), str(SCRIPTS)):
    if p not in sys.path:
        sys.path.insert(0, p)


def test_build_workbook_toc_skips_source_and_orders_sections():
    from pdf_report.builder import PDF_GROUP_SPECS, build_workbook_toc, display_sheet_name

    wb = Workbook()
    # remove default
    default = wb.active
    wb.remove(default)

    wb.create_sheet("==> Source")
    wb.create_sheet("Master_PL")
    wb.create_sheet("__SOURCE__Sales")
    wb.create_sheet("==> Earnings")
    wb.create_sheet("Lead_IS")
    wb.create_sheet("General sales table")
    wb.create_sheet("Churn")
    wb.create_sheet("Bubble")
    wb.create_sheet("FTE Development")
    wb.create_sheet("==> Assets")
    wb.create_sheet("Lead_BS")
    wb.create_sheet("Working_Capital")
    wb.create_sheet("Cashflow")
    wb.create_sheet("Net sales breakdown")
    wb.create_sheet("Alpha_BS_Reconciliation")
    wb.create_sheet("BS_Reconciliation")
    wb.create_sheet("PL_Reconciliation")
    wb.create_sheet("Alpha_PL_Reconciliation")

    toc = build_workbook_toc(wb)
    assert all(item.get("key") != "source" for item in toc)
    assert not any(item.get("sheet_name") in {"Master_PL", "__SOURCE__Sales"} for item in toc)
    assert not any(item.get("sheet_name") == "Working_Capital" for item in toc)

    section_labels = [i["label"] for i in toc if i["type"] == "section"]
    assert "Executive Summary" in section_labels
    assert "Earnings" in section_labels
    assert "Assets" in section_labels
    assert "Liquidity" in section_labels
    assert "Appendix" in section_labels
    assert "Source" not in section_labels

    idx = {lab: section_labels.index(lab) for lab in section_labels}
    assert idx["Executive Summary"] < idx["Earnings"] < idx["Assets"] < idx["Liquidity"]
    assert idx["Liquidity"] < idx["Appendix"]

    keys = [s["key"] for s in PDF_GROUP_SPECS]
    assert "source" not in keys

    exec_sec = next(i for i in toc if i["type"] == "section" and i["key"] == "executive")
    assert exec_sec["section_no"] == 1
    assert any(e["code"] == "1.1" for e in exec_sec["outline"])

    earn_sec = next(i for i in toc if i["type"] == "section" and i["key"] == "earnings")
    assert earn_sec["section_no"] == 2
    codes = [(e.get("code"), e.get("title"), e.get("indent")) for e in earn_sec["outline"]]
    assert ("2.1", "Lead Income statement", 0) in codes
    assert ("2.2", "Sales", 0) in codes
    assert any(c[0] == "I" and c[2] == 1 for c in codes)
    assert display_sheet_name("Lead_IS") == "Lead Income statement"
    assert display_sheet_name("FTE Development") == "Payroll accounting"

    # Subsection dividers before Lead IS / Sales / Payroll
    sub_titles = [i.get("title") for i in toc if i["type"] == "subsection"]
    assert any("Lead Income statement" in (t or "") for t in sub_titles)
    assert any(t and t.strip().endswith("Sales") or (t or "").endswith("Sales") for t in sub_titles)
    assert any("Payroll accounting" in (t or "") for t in sub_titles)

    lead_item = next(i for i in toc if i.get("sheet_name") == "Lead_IS")
    assert lead_item["display_name"] == "Lead Income statement"
    assert lead_item["outline_code"] == "2.1"

    appendix_sheets = [i["sheet_name"] for i in toc if i["type"] == "sheet" and i["key"] == "appendix"]
    assert appendix_sheets[:2] == ["BS_Reconciliation", "PL_Reconciliation"]
    assert appendix_sheets.index("BS_Reconciliation") < appendix_sheets.index("Alpha_BS_Reconciliation")


def test_extract_sheet_table_skips_hidden_rows_and_cols():
    from pdf_report.table_extract import extract_sheet_table

    wb = Workbook()
    ws = wb.active
    ws.title = "Demo"
    ws["D1"] = "Finssentials"
    ws["D2"] = "Demo table"
    ws["D4"] = "Company | Demo table"
    ws["A6"] = "hidden_helper"
    ws["D6"] = "Line"
    ws["E6"] = "FY24"
    ws["D7"] = "Revenue"
    ws["E7"] = 1000
    ws["D8"] = "Hidden detail"
    ws["E8"] = 50
    ws["D9"] = "Reported"
    ws["E9"] = 1050

    ws.column_dimensions["A"].hidden = True
    ws.row_dimensions[8].hidden = True
    ws.row_dimensions[8].outlineLevel = 1

    table = extract_sheet_table(ws)
    assert table.title_lines
    assert any("Demo" in t for t in table.title_lines)

    # Flatten visible text values
    texts = []
    for row in table.rows:
        texts.append([str(c.value) if c.value is not None else "" for c in row])
    flat = " | ".join(" / ".join(r) for r in texts)
    assert "Hidden detail" not in flat
    assert "hidden_helper" not in flat
    assert "Revenue" in flat
    assert "Reported" in flat


def test_build_fast_track_pdf_report_writes_file(tmp_path: Path):
    pytest.importorskip("reportlab")
    from pdf_report.builder import build_fast_track_pdf_report

    wb = Workbook()
    ws = wb.active
    ws.title = "Lead_IS"
    ws["D1"] = "Finssentials"
    ws["D2"] = "Income Statement"
    ws["D6"] = "Account"
    ws["E6"] = "FY24A"
    ws["D7"] = "Revenue"
    ws["E7"] = 12345
    xlsx = tmp_path / "Demo_FastTrack.xlsx"
    wb.save(xlsx)

    pdf = tmp_path / "Demo_FastTrack_Report.pdf"
    out = build_fast_track_pdf_report(
        xlsx,
        pdf,
        project_name="Demo Project",
        company_name="Demo Co",
        recalc_excel=False,
    )
    assert Path(out).is_file()
    assert Path(out).stat().st_size > 500
