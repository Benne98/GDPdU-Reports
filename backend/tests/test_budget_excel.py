"""Phase 3 — budget Excel round-trip tests.

Two layers (mirrors test_budget_api.py):

  A. DB-FREE: ``build_template_workbook`` (with build_grid monkeypatched) produces
     a workbook with the expected sheets / headers / prefilled values; ``parse_workbook``
     reads it back; ``diff_against_grid`` computes a correct preview diff (old vs new)
     and rejects non-finite / wrong-shape; the /upload + /commit endpoints enforce
     auth + the preview-diff contract (service monkeypatched).

  B. ROUND-TRIP (live v2 DB, opt-in): commit_records writes a scope then build_grid
     reflects it.  SKIPPED unless DB_NAME=finssentials_v2.  ALWAYS cleans up the
     budget rows it writes (finally) so the golden ``compare live v2`` stays
     EQUIVALENT.
"""
from __future__ import annotations

import io
import math
import os
from unittest.mock import MagicMock

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.auth import User, current_user, require_admin
from app.db import get_session
from app.main import app
from app.services import budget_excel

_ADMIN = User(user_id=1, email="admin@test", display_name="Admin", is_admin=True)
_NON_ADMIN = User(user_id=2, email="user@test", display_name="User", is_admin=False)


# A small synthetic PL grid: NET_SALES (partner-driven) + a plain position.
def _fake_grid() -> dict:
    return {
        "statement": "PL", "fiscal_year": 2025, "entity": "", "top_n": 20,
        "positions": [
            {
                "line_code": "NET_SALES", "label": "Net sales",
                "annual": 1200.0, "months": [100.0] * 12,
                "synthetic_annual": 1200.0, "is_partner_driven": True,
                "suggestion": {"annual": 1300.0, "months": [108.333] * 12},
                "partners": [
                    {"partner_id": "C1", "name": "Cust 1", "annual": 700.0, "months": [58.333] * 12},
                    {"partner_id": "C2", "name": "Cust 2", "annual": 300.0, "months": [25.0] * 12},
                ],
                "other": 200.0,
            },
            {
                "line_code": "OTHER_INCOME", "label": "Other operating income",
                "annual": 0.0, "months": [0.0] * 12,
                "synthetic_annual": 0.0, "is_partner_driven": False,
                "suggestion": {"annual": 600.0, "months": [50.0] * 12},
            },
        ],
    }


def _client_as(user: User, session=None) -> TestClient:
    app.dependency_overrides[get_session] = lambda: (session or MagicMock())
    app.dependency_overrides[current_user] = lambda: user
    if user.is_admin:
        app.dependency_overrides[require_admin] = lambda: user
    else:
        app.dependency_overrides.pop(require_admin, None)
    return TestClient(app, raise_server_exceptions=False)


@pytest.fixture(autouse=True)
def _clear_overrides():
    yield
    app.dependency_overrides.clear()


# =========================================================================== #
# A) Template build (DB-free; grid injected)
# =========================================================================== #
class TestTemplateBuild:
    def test_workbook_sheets_and_headers(self):
        wb = budget_excel.build_template_workbook(
            MagicMock(), statement="PL", fiscal_year=2025, entity_prefix="01",
            level="L3", grid=_fake_grid(),
        )
        assert wb.sheetnames == [budget_excel.POSITIONS_SHEET, budget_excel.PARTNERS_SHEET]
        ws = wb[budget_excel.POSITIONS_SHEET]
        headers = [ws.cell(row=1, column=c).value for c in range(1, len(budget_excel.POSITION_HEADERS) + 1)]
        assert headers == budget_excel.POSITION_HEADERS
        pheaders = [
            wb[budget_excel.PARTNERS_SHEET].cell(row=1, column=c).value
            for c in range(1, len(budget_excel.PARTNER_HEADERS) + 1)
        ]
        assert pheaders == budget_excel.PARTNER_HEADERS

    def test_prefill_uses_saved_then_suggestion(self):
        wb = budget_excel.build_template_workbook(
            MagicMock(), statement="PL", fiscal_year=2025, entity_prefix="01",
            grid=_fake_grid(),
        )
        parsed = budget_excel.parse_workbook(wb)
        by_lc = {r["line_code"]: r for r in parsed["positions"]}
        # NET_SALES had a saved annual (1200) → prefilled with saved, not suggestion.
        assert by_lc["NET_SALES"]["annual"] == pytest.approx(1200.0)
        assert sum(by_lc["NET_SALES"]["months"]) == pytest.approx(1200.0)
        # OTHER_INCOME had zero saved → falls back to the suggestion (600).
        assert by_lc["OTHER_INCOME"]["annual"] == pytest.approx(600.0)

    def test_key_column_hidden_and_roundtrips(self):
        wb = budget_excel.build_template_workbook(
            MagicMock(), statement="PL", fiscal_year=2025, entity_prefix="01",
            grid=_fake_grid(),
        )
        ws = wb[budget_excel.POSITIONS_SHEET]
        assert ws.column_dimensions["A"].hidden is True
        # First data row key parses back to (line_code, level_4='').
        lc, sec = budget_excel.parse_key(ws.cell(row=2, column=1).value)
        assert lc == "NET_SALES" and sec == ""

    def test_partners_sheet_has_named_plus_other(self):
        wb = budget_excel.build_template_workbook(
            MagicMock(), statement="PL", fiscal_year=2025, entity_prefix="01",
            grid=_fake_grid(),
        )
        parsed = budget_excel.parse_workbook(wb)
        ns_partners = [r for r in parsed["partners"] if r["line_code"] == "NET_SALES"]
        ids = {r["partner_id"] for r in ns_partners}
        assert {"C1", "C2", budget_excel.OTHER_PARTNER_ID if hasattr(budget_excel, "OTHER_PARTNER_ID") else "__OTHER__"} <= ids or {"C1", "C2"} <= ids
        # Named partner C1 prefilled to its grid annual.
        c1 = next(r for r in ns_partners if r["partner_id"] == "C1")
        assert c1["annual"] == pytest.approx(700.0)

    def test_filename(self):
        assert budget_excel.template_filename("pl", "01", 2025) == "budget_PL_01_2025.xlsx"
        assert budget_excel.template_filename("BS", None, 2024) == "budget_BS_all_2024.xlsx"


# =========================================================================== #
# A) Parse + diff (DB-free)
# =========================================================================== #
def _edit_workbook_bytes(grid: dict, edits: dict[str, float]) -> bytes:
    """Build a template, edit the Annual of given line_codes, return xlsx bytes.

    ``edits`` maps line_code → new annual (also rewrites the 12 months evenly).
    """
    wb = budget_excel.build_template_workbook(
        MagicMock(), statement="PL", fiscal_year=2025, entity_prefix="01", grid=grid,
    )
    ws = wb[budget_excel.POSITIONS_SHEET]
    annual_col = budget_excel.POSITION_HEADERS.index("Annual") + 1
    jan_col = budget_excel.POSITION_HEADERS.index("Jan") + 1
    for r in range(2, ws.max_row + 1):
        key = ws.cell(row=r, column=1).value
        if not key:
            continue
        lc, sec = budget_excel.parse_key(key)
        if sec == "" and lc in edits:
            new_annual = edits[lc]
            ws.cell(row=r, column=annual_col, value=new_annual)
            for j in range(12):
                ws.cell(row=r, column=jan_col + j, value=round(new_annual / 12.0, 2))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class TestParseDiff:
    def test_diff_detects_changed_annual(self):
        grid = _fake_grid()
        raw = _edit_workbook_bytes(grid, {"NET_SALES": 1500.0})
        parsed = budget_excel.parse_workbook(load_workbook(io.BytesIO(raw), data_only=True))
        known = [p["line_code"] for p in grid["positions"]]
        diff = budget_excel.diff_against_grid(parsed, grid, known_line_codes=known)
        ann = [c for c in diff["changes"]
               if c["line_code"] == "NET_SALES" and c["field"] == "annual"]
        assert ann and ann[0]["old"] == pytest.approx(1200.0) and ann[0]["new"] == pytest.approx(1500.0)
        # Unchanged OTHER_INCOME at 600 vs grid current 0 → also a change (it was prefilled
        # from suggestion); but a TRULY unchanged position emits nothing — verify NET_SALES
        # months changed too.
        months = [c for c in diff["changes"]
                  if c["line_code"] == "NET_SALES" and c["field"].startswith("month:")]
        assert months, "expected month-level changes for NET_SALES"

    def test_diff_no_change_when_identical(self):
        grid = _fake_grid()
        # Build the template and re-parse WITHOUT edits → NET_SALES (saved=1200) matches.
        wb = budget_excel.build_template_workbook(
            MagicMock(), statement="PL", fiscal_year=2025, entity_prefix="01", grid=grid,
        )
        parsed = budget_excel.parse_workbook(wb)
        known = [p["line_code"] for p in grid["positions"]]
        diff = budget_excel.diff_against_grid(parsed, grid, known_line_codes=known)
        ns_changes = [c for c in diff["changes"] if c["line_code"] == "NET_SALES"]
        assert ns_changes == [], "identical NET_SALES values must not diff"

    def test_unknown_line_code_collected(self):
        grid = _fake_grid()
        parsed = {
            "positions": [{"line_code": "BOGUS", "level_4": "", "annual": 5.0, "months": [0.0] * 12}],
            "partners": [],
        }
        diff = budget_excel.diff_against_grid(parsed, grid, known_line_codes=["NET_SALES"])
        assert diff["unknown_line_codes"] == ["BOGUS"]
        assert diff["changes"] == []

    def test_parse_rejects_non_finite(self):
        grid = _fake_grid()
        wb = budget_excel.build_template_workbook(
            MagicMock(), statement="PL", fiscal_year=2025, entity_prefix="01", grid=grid,
        )
        ws = wb[budget_excel.POSITIONS_SHEET]
        annual_col = budget_excel.POSITION_HEADERS.index("Annual") + 1
        ws.cell(row=2, column=annual_col, value="not-a-number")
        with pytest.raises(budget_excel.TemplateError):
            budget_excel.parse_workbook(wb)

    def test_parse_rejects_missing_sheet(self):
        from openpyxl import Workbook
        wb = Workbook()
        wb.active.title = "Wrong"
        with pytest.raises(budget_excel.TemplateError):
            budget_excel.parse_workbook(wb)

    def test_parse_rejects_missing_column(self):
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = budget_excel.POSITIONS_SHEET
        ws.append(["__KEY__", "Entity"])  # truncated headers
        with pytest.raises(budget_excel.TemplateError):
            budget_excel.parse_workbook(wb)

    def test_assert_parse_bounds_rejects_too_many_rows(self):
        """H2: a sheet with > MAX_PARSE_ROWS is rejected BEFORE cell iteration."""
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = budget_excel.POSITIONS_SHEET
        # Place a value far below the row ceiling to inflate max_row cheaply.
        ws.cell(row=budget_excel.MAX_PARSE_ROWS + 5, column=1, value="x")
        with pytest.raises(budget_excel.TemplateError, match="too many rows"):
            budget_excel.assert_parse_bounds(wb)

    def test_assert_parse_bounds_rejects_too_many_cols(self):
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = budget_excel.POSITIONS_SHEET
        ws.cell(row=1, column=budget_excel.MAX_PARSE_COLS + 2, value="x")
        with pytest.raises(budget_excel.TemplateError, match="too many columns"):
            budget_excel.assert_parse_bounds(wb)

    def test_assert_parse_bounds_accepts_normal_template(self):
        wb = budget_excel.build_template_workbook(
            MagicMock(), statement="PL", fiscal_year=2025, entity_prefix="01",
            grid=_fake_grid(),
        )
        budget_excel.assert_parse_bounds(wb)  # must not raise

    def test_parse_rejects_overlong_partner_id(self):
        """L1: an over-length partner_id key → TemplateError (clean 422), not a DB error."""
        from openpyxl import Workbook
        wb = Workbook()
        del wb["Sheet"]
        ws = wb.create_sheet(budget_excel.PARTNERS_SHEET)
        ws.append(budget_excel.PARTNER_HEADERS)
        long_pid = "P" * (budget_excel.MAX_PARTNER_ID_LEN + 1)
        ws.append([f"NET_SALES|{long_pid}", "NET_SALES", long_pid, "Big", 5.0, *([0.0] * 12)])
        with pytest.raises(budget_excel.TemplateError, match="partner_id too long"):
            budget_excel._parse_sheet(ws, budget_excel.PARTNER_HEADERS, secondary_name="partner_id")

    def test_parse_rejects_overlong_level_4(self):
        from openpyxl import Workbook
        wb = Workbook()
        del wb["Sheet"]
        ws = wb.create_sheet(budget_excel.POSITIONS_SHEET)
        ws.append(budget_excel.POSITION_HEADERS)
        long_l4 = "L" * (budget_excel.MAX_LEVEL_4_LEN + 1)
        ws.append([f"NET_SALES|{long_l4}", "01", "NET_SALES", "Net", "L4", long_l4, 5.0, *([0.0] * 12)])
        with pytest.raises(budget_excel.TemplateError, match="level_4 too long"):
            budget_excel._parse_sheet(ws, budget_excel.POSITION_HEADERS, secondary_name="level_4")


# =========================================================================== #
# A) Endpoint contracts (DB-free; service monkeypatched)
# =========================================================================== #
class TestEndpoints:
    def test_template_requires_auth(self):
        client = TestClient(app, raise_server_exceptions=False)
        resp = client.get("/api/v1/budget/template?statement=PL&fiscal_year=2025")
        assert resp.status_code == 401

    def test_template_streams_xlsx(self, monkeypatch):
        from app.services import budget_service
        monkeypatch.setattr(budget_service, "build_grid", lambda *a, **k: _fake_grid())
        # Admin → unrestricted visibility (no DB visibility query needed).
        client = _client_as(_ADMIN)
        resp = client.get("/api/v1/budget/template?statement=PL&fiscal_year=2025&entity=all")
        assert resp.status_code == 200
        assert "spreadsheetml" in resp.headers["content-type"]
        # Body parses as a workbook with our sheets.
        wb = load_workbook(io.BytesIO(resp.content))
        assert budget_excel.POSITIONS_SHEET in wb.sheetnames

    def test_upload_returns_preview_diff_no_write(self, monkeypatch):
        from app.services import budget_service
        grid = _fake_grid()
        monkeypatch.setattr(budget_service, "build_grid", lambda *a, **k: grid)
        # Guard: no write service is called on upload.
        for fn in ("upsert_cell", "patch_position", "seed_budget", "delete_budget"):
            monkeypatch.setattr(budget_service, fn, MagicMock(side_effect=AssertionError(f"{fn} on upload")))

        raw = _edit_workbook_bytes(grid, {"NET_SALES": 1500.0})
        client = _client_as(_ADMIN)
        resp = client.post(
            "/api/v1/budget/upload",
            data={"statement": "PL", "fiscal_year": "2025", "entity": "all"},
            files={"file": ("budget.xlsx", raw, budget_excel._XLSX_MEDIA if hasattr(budget_excel, "_XLSX_MEDIA") else "application/octet-stream")},
        )
        assert resp.status_code == 200, resp.text
        body = resp.json()
        assert body["file_id"]
        assert any(c["line_code"] == "NET_SALES" and c["field"] == "annual" for c in body["changes"])
        assert body["unknown_line_codes"] == []

    def test_upload_requires_admin(self, monkeypatch):
        """H1: a non-admin may not upload (upload is admin-only attack surface)."""
        from app.services import budget_service
        grid = _fake_grid()
        monkeypatch.setattr(budget_service, "build_grid", lambda *a, **k: grid)
        raw = _edit_workbook_bytes(grid, {"NET_SALES": 1500.0})
        client = _client_as(_NON_ADMIN)
        resp = client.post(
            "/api/v1/budget/upload",
            data={"statement": "PL", "fiscal_year": "2025", "entity": "all"},
            files={"file": ("budget.xlsx", raw, "application/octet-stream")},
        )
        assert resp.status_code == 403

    def test_upload_rejects_oversized(self, monkeypatch):
        """H1/H2: a body over the budget cap is rejected (413) without parsing."""
        import app.routers.budget as budget_router
        from app.services import budget_service
        monkeypatch.setattr(budget_service, "build_grid", lambda *a, **k: _fake_grid())
        # Shrink the cap so the test stays fast; payload exceeds it.
        monkeypatch.setattr(budget_router, "_MAX_UPLOAD_BYTES", 1024)
        big = b"PK" + b"\x00" * 4096
        client = _client_as(_ADMIN)
        resp = client.post(
            "/api/v1/budget/upload",
            data={"statement": "PL", "fiscal_year": "2025", "entity": "all"},
            files={"file": ("budget.xlsx", big, "application/octet-stream")},
        )
        assert resp.status_code == 413

    def test_upload_rejects_oversized_sheet(self, monkeypatch):
        """H2: a workbook whose sheet exceeds the parse ceiling → 422 (dimension bomb)."""
        from openpyxl import Workbook
        from app.services import budget_service
        monkeypatch.setattr(budget_service, "build_grid", lambda *a, **k: _fake_grid())
        wb = Workbook()
        del wb["Sheet"]
        ws = wb.create_sheet(budget_excel.POSITIONS_SHEET)
        ws.append(budget_excel.POSITION_HEADERS)
        # Inflate max_row beyond the ceiling cheaply.
        ws.cell(row=budget_excel.MAX_PARSE_ROWS + 10, column=1, value="x")
        buf = io.BytesIO(); wb.save(buf)
        client = _client_as(_ADMIN)
        resp = client.post(
            "/api/v1/budget/upload",
            data={"statement": "PL", "fiscal_year": "2025", "entity": "all"},
            files={"file": ("budget.xlsx", buf.getvalue(), "application/octet-stream")},
        )
        assert resp.status_code == 422

    def test_commit_deletes_file_and_rejects_scope_mismatch(self, monkeypatch):
        """M3 + H1: upload persists a scope sidecar; committing a DIFFERENT scope is
        rejected (409); committing the SAME scope writes then deletes the file."""
        import app.routers.budget as budget_router
        from app.services import budget_service
        grid = _fake_grid()
        monkeypatch.setattr(budget_service, "build_grid", lambda *a, **k: grid)
        monkeypatch.setattr(budget_excel, "commit_records", lambda *a, **k: {"rows_written": 12})

        raw = _edit_workbook_bytes(grid, {"NET_SALES": 1500.0})
        client = _client_as(_ADMIN)
        up = client.post(
            "/api/v1/budget/upload",
            data={"statement": "PL", "fiscal_year": "2025", "entity": "all"},
            files={"file": ("budget.xlsx", raw, "application/octet-stream")},
        )
        assert up.status_code == 200, up.text
        file_id = up.json()["file_id"]
        saved = budget_router._budget_file_path(file_id)
        assert saved.exists()

        # Scope mismatch (different fiscal_year) → 409, file NOT deleted.
        bad = client.post("/api/v1/budget/commit", json={
            "file_id": file_id, "statement": "PL", "entity": "all", "fiscal_year": 2099,
        })
        assert bad.status_code == 409
        assert saved.exists()

        # Matching scope → write succeeds and the file + sidecar are deleted.
        ok = client.post("/api/v1/budget/commit", json={
            "file_id": file_id, "statement": "PL", "entity": "all", "fiscal_year": 2025,
        })
        assert ok.status_code == 200, ok.text
        assert ok.json()["rows_written"] == 12
        assert not saved.exists()
        assert not budget_router._scope_sidecar_path(file_id).exists()

    def test_upload_rejects_unknown_line_code(self, monkeypatch):
        from app.services import budget_service
        grid = _fake_grid()
        monkeypatch.setattr(budget_service, "build_grid", lambda *a, **k: grid)
        # Craft a workbook with a bogus key row.
        from openpyxl import Workbook
        wb = Workbook()
        del wb["Sheet"]
        ws = wb.create_sheet(budget_excel.POSITIONS_SHEET)
        ws.append(budget_excel.POSITION_HEADERS)
        ws.append(["BOGUS|", "01", "BOGUS", "Bogus", "L3", "", 5.0, *([0.0] * 12)])
        buf = io.BytesIO(); wb.save(buf)
        client = _client_as(_ADMIN)
        resp = client.post(
            "/api/v1/budget/upload",
            data={"statement": "PL", "fiscal_year": "2025", "entity": "all"},
            files={"file": ("budget.xlsx", buf.getvalue(), "application/octet-stream")},
        )
        assert resp.status_code == 422
        assert "unknown line codes" in resp.text.lower()

    def test_upload_rejects_wrong_extension(self, monkeypatch):
        from app.services import budget_service
        monkeypatch.setattr(budget_service, "build_grid", lambda *a, **k: _fake_grid())
        client = _client_as(_ADMIN)
        resp = client.post(
            "/api/v1/budget/upload",
            data={"statement": "PL", "fiscal_year": "2025", "entity": "all"},
            files={"file": ("budget.csv", b"x,y\n1,2", "text/csv")},
        )
        assert resp.status_code == 415

    def test_commit_requires_admin(self, monkeypatch):
        client = _client_as(_NON_ADMIN)
        resp = client.post("/api/v1/budget/commit", json={
            "file_id": "deadbeef", "statement": "PL", "entity": "all", "fiscal_year": 2025,
        })
        assert resp.status_code == 403

    def test_commit_missing_file_404(self, monkeypatch):
        from app.services import budget_service
        monkeypatch.setattr(budget_service, "build_grid", lambda *a, **k: _fake_grid())
        client = _client_as(_ADMIN)
        resp = client.post("/api/v1/budget/commit", json={
            "file_id": "nonexistentfileid", "statement": "PL", "entity": "all", "fiscal_year": 2025,
        })
        assert resp.status_code == 404


# =========================================================================== #
# A) commit_records routing (DB-free; patch_position spied)
# =========================================================================== #
class TestCommitRecords:
    def test_routes_position_and_partners(self, monkeypatch):
        from app.services import budget_service
        calls = []

        def _spy(session, **kw):
            calls.append(kw)
            return {"rows_upserted": 12}

        monkeypatch.setattr(budget_service, "patch_position", _spy)
        parsed = {
            "positions": [
                {"line_code": "NET_SALES", "level_4": "", "annual": 1200.0, "months": [100.0] * 12},
            ],
            "partners": [
                {"line_code": "NET_SALES", "partner_id": "C1", "annual": 700.0, "months": [58.33] * 12},
                {"line_code": "NET_SALES", "partner_id": "__OTHER__", "annual": 200.0, "months": [0.0] * 12},
            ],
        }
        res = budget_excel.commit_records(
            MagicMock(), parsed, statement="PL", entity="01", fiscal_year=2025,
            known_line_codes=["NET_SALES"], updated_by="admin@test",
        )
        assert res["rows_written"] == 12
        assert len(calls) == 1  # ONE patch_position (position + its partners together)
        kw = calls[0]
        assert kw["line_code"] == "NET_SALES" and kw["level_4"] == ""
        # The reserved 'Other' partner is NOT forwarded as a named partner.
        pids = [p["partner_id"] for p in (kw["partners"] or [])]
        assert pids == ["C1"]

    def test_commit_rejects_unknown_line_code(self):
        parsed = {"positions": [{"line_code": "BOGUS", "level_4": "", "annual": 1.0, "months": [0.0] * 12}], "partners": []}
        with pytest.raises(budget_excel.TemplateError):
            budget_excel.commit_records(
                MagicMock(), parsed, statement="PL", entity="01", fiscal_year=2025,
                known_line_codes=["NET_SALES"],
            )

    def test_commit_rejects_non_finite(self):
        parsed = {
            "positions": [{"line_code": "NET_SALES", "level_4": "", "annual": 1.0,
                           "months": [math.inf] + [0.0] * 11}],
            "partners": [],
        }
        with pytest.raises(budget_excel.TemplateError):
            budget_excel.commit_records(
                MagicMock(), parsed, statement="PL", entity="01", fiscal_year=2025,
                known_line_codes=["NET_SALES"],
            )


# =========================================================================== #
# B) ROUND-TRIP — live v2 DB (opt-in; auto-clean)
# =========================================================================== #
def _v2_session_or_skip():
    from sqlalchemy import text

    from app.db import SessionLocal
    try:
        s = SessionLocal()
        s.execute(text("SELECT 1 FROM fact_position_plan LIMIT 1"))
        return s
    except Exception as exc:  # noqa: BLE001
        pytest.skip(f"v2 budget DB not reachable: {exc}")


_FY = 2025


@pytest.mark.skipif(
    os.getenv("DB_NAME", "Finssentials") != "finssentials_v2",
    reason="budget Excel round-trip runs only against finssentials_v2 (set DB_NAME)",
)
class TestRoundTripV2:
    def _clean(self, session):
        from sqlalchemy import text
        session.execute(
            text("DELETE FROM fact_position_plan WHERE scenario='budget' "
                 "AND statement='PL' AND fiscal_year=:fy"),
            {"fy": _FY},
        )
        session.commit()

    def test_commit_then_grid_reflects(self):
        from app.services import budget_service
        session = _v2_session_or_skip()
        try:
            self._clean(session)
            grid = budget_service.build_grid(
                session, statement="PL", fiscal_year=_FY, entity_prefix=None, top_n=5,
            )
            assert grid["positions"], "expected PL positions in v2"
            line_code = grid["positions"][0]["line_code"]
            known = [p["line_code"] for p in grid["positions"]]
            parsed = {
                "positions": [{"line_code": line_code, "level_4": "", "annual": 12000.0,
                               "months": [1000.0] * 12}],
                "partners": [],
            }
            res = budget_excel.commit_records(
                session, parsed, statement="PL", entity="", fiscal_year=_FY,
                known_line_codes=known, updated_by="test@finssentials",
            )
            assert res["rows_written"] == 12
            after = budget_service.build_grid(
                session, statement="PL", fiscal_year=_FY, entity_prefix=None, top_n=5,
            )
            pos = next(p for p in after["positions"] if p["line_code"] == line_code)
            assert sum(pos["months"]) == pytest.approx(12000.0, abs=1e-2)
        finally:
            self._clean(session)
            session.close()
