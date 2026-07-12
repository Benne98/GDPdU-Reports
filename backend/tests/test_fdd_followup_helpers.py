"""Tests for FDD bot follow-up fixes (OPOS layout and BS reported row)."""

import os
import re
from pathlib import Path
import sys

import pandas as pd
import pytest
from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[2]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from opos import (
    KEY_COL,
    POS_COL,
    AgingBucket,
    AsOfPeriod,
    _excel_sheet_ref,
    _opos_sumifs_criteria,
    col_letter,
    financial_statements_cell_ref,
    load_reported_from_financial_statements,
)
from funktionssammlung import (
    build_output_file_path,
    find_session_output_file,
    prune_init_placeholder_sheet,
    touch_session_workbook,
)
from FTE_payroll import build_pl_personnel_rows


def _normalize_header_key(name: str) -> str:
    s = str(name).replace("\u00a0", " ").strip().lower()
    return re.sub(r"\s+", " ", s)


def test_opos_layout_constants():
    assert POS_COL == 4
    assert KEY_COL == 2


def test_opos_sumifs_partner_criterion_uses_key_col():
    bucket = AgingBucket(key="overdue_1_30", label="1-30", min_days=1, max_days=30)
    period = AsOfPeriod(label="Dec25A", date=pd.Timestamp("2025-12-31"))
    body = _opos_sumifs_criteria(
        "__SOURCE__Dec25A",
        {"partner id": 1, "amount": 2, "due date": 3},
        {"columns": {"partner_id": "partner id", "amount": "amount", "due_date": "due date"}},
        period,
        bucket,
        col_idx=5,
        partner_row=12,
    )
    assert f"${col_letter(KEY_COL)}12" in body


def test_excel_sheet_ref_escapes_spaces():
    assert _excel_sheet_ref("Trade debtors detail") == "'Trade debtors detail'"


def test_header_key_normalization_nbsp_tolerant():
    assert _normalize_header_key("Bilanz\u00a0position") == _normalize_header_key("Bilanz position")


def test_load_reported_from_financial_statements():
    wb = Workbook()
    ws = wb.active
    ws.title = "Financial statements"
    ws.cell(8, 11, "Dec25A")
    ws.cell(9, 10, "Trade receivables")
    ws.cell(9, 11, 123.45)

    periods = [AsOfPeriod(label="Dec25A", date=pd.Timestamp("2025-12-31"))]
    cfg = {"side": "debitor"}
    reported = load_reported_from_financial_statements(wb, cfg, periods)
    assert reported["Dec25A"] == pytest.approx(123.45)


def test_financial_statements_cell_ref():
    wb = Workbook()
    ws = wb.active
    ws.title = "BS_Recon"
    ws.cell(8, 11, "Dec25A")
    ws.cell(9, 10, "Fixed assets")
    ws.cell(9, 11, 999.0)

    ref = financial_statements_cell_ref(wb, "Fixed assets", "Dec25A")
    assert ref == "'BS_Recon'!$K$9"


def test_build_output_file_path_uses_project_name(tmp_path):
    cfg = {
        "title": "My Project",
        "output_file_path": str(tmp_path),
        "case_id": "sess-123",
        "use_session_workbook": True,
    }
    path = build_output_file_path(cfg)
    assert path.endswith("My_Project_Workbook.xlsx")


def test_build_pl_personnel_rows_default_without_master():
    rows = build_pl_personnel_rows("", "Master_PL", {})
    assert len(rows) >= 1
    assert all(r["L3"] == "Personnel expenses" for r in rows)
    assert all(r["L2"] == "Expense" for r in rows)
    assert all(r["L4"] for r in rows)


def test_find_session_output_file_project_workbook(tmp_path):
    cfg = {
        "title": "Acme Deal",
        "output_file_path": str(tmp_path),
        "case_id": "sess-1",
        "use_session_workbook": True,
    }
    path = tmp_path / "Acme_Deal_Workbook.xlsx"
    path.write_bytes(b"fake")
    found = find_session_output_file(cfg)
    assert found == str(path)


def test_touch_session_workbook_returns_path(tmp_path):
    cfg = {
        "title": "New Project",
        "output_file_path": str(tmp_path),
        "case_id": "sess-2",
        "use_session_workbook": True,
    }
    path = touch_session_workbook(cfg)
    assert path.endswith("New_Project_Workbook.xlsx")


def test_prune_init_placeholder_sheet():
    from openpyxl import Workbook

    wb = Workbook()
    wb.active.title = "__init__"
    wb.create_sheet("FTE development")
    prune_init_placeholder_sheet(wb)
    assert "__init__" not in wb.sheetnames
    assert "FTE development" in wb.sheetnames
