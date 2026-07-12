"""Smoke tests for recon pipeline script registration and config shape."""
from __future__ import annotations

from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parents[2]


def test_recon_pipeline_scripts_exist():
    from app.routers import fdd_bot

    for key in ("recon_pl", "recon_bs", "bs_bucket", "lead_is", "lead_bs", "working_capital", "cashflow"):
        path = fdd_bot.SCRIPTS[key]
        assert path.is_file(), f"Missing script for {key}: {path}"


def test_recon_core_config_shape():
    """recon-core uses custom entity order and default English section titles."""
    from app.routers.fdd_bot import _default_recon_display_titles

    titles = _default_recon_display_titles()
    assert titles["aggregated_title"] == "Aggregated"
    assert titles["ic_display_name"] == "IC eliminations"

    entity_order = ["Beta", "Atlas"]
    pl_config = {
        "sort_by": "custom",
        "entity_order": entity_order,
        "l4_sort_basis": "latest_fy",
        "paths": {"append_to_master": True},
    }
    assert pl_config["sort_by"] == "custom"
    assert pl_config["entity_order"] == entity_order


def test_recon_pipeline_step_sheet_names():
    """Expected output sheets after full pipeline."""
    expected = {
        "recon_pl": "PL_Reconciliation",
        "recon_bs": "BS_Reconciliation",
        "bs_bucket": "BS_Bucket",
        "lead_is": "Lead_IS",
        "lead_bs": "Lead_BS",
    }
    assert set(expected.values()) == {
        "PL_Reconciliation",
        "BS_Reconciliation",
        "BS_Bucket",
        "Lead_IS",
        "Lead_BS",
    }


def test_databook_runtime_load_empty_without_argv(monkeypatch):
    import sys

    from databook_runtime import load_argv_config

    monkeypatch.setattr(sys, "argv", ["script.py"])
    assert load_argv_config() == {}


def test_apply_custom_entity_order():
    from databook_runtime import apply_custom_entity_order

    entities = ["Gamma", "Alpha", "Beta"]
    order = ["Beta", "Alpha"]
    assert apply_custom_entity_order(entities, order) == ["Beta", "Alpha", "Gamma"]


def test_sanitize_entity_recon_sheet_name():
    from databook_workbook import sanitize_entity_recon_sheet_name

    assert sanitize_entity_recon_sheet_name("WoSH", "bs") == "WoSH_BS_Reconciliation"
    assert sanitize_entity_recon_sheet_name("WoSH", "pl") == "WoSH_PL_Reconciliation"
    assert sanitize_entity_recon_sheet_name("Bad/Name*", "bs") == "BadName_BS_Reconciliation"


def test_entity_recon_sheet_names_bs_before_pl():
    from databook_workbook import entity_recon_sheet_names

    names = entity_recon_sheet_names(["Alpha", "Beta"])
    assert names == [
        "Alpha_BS_Reconciliation",
        "Alpha_PL_Reconciliation",
        "Beta_BS_Reconciliation",
        "Beta_PL_Reconciliation",
    ]


def test_build_entity_recon_blocks_three_portfolio_blocks():
    import sys
    from pathlib import Path

    scripts = Path(__file__).resolve().parents[2] / "scripts"
    if str(scripts) not in sys.path:
        sys.path.insert(0, str(scripts))
    from databook_excel_layout import build_entity_recon_blocks

    blocks = build_entity_recon_blocks(
        "WoSH",
        {"difference_title": "Difference", "financial_statements_title": "Financial statements"},
        first_col=8,
        n_years=3,
    )
    assert [b["kind"] for b in blocks] == ["entity", "difference", "fs"]
    assert blocks[0]["key"] == "WoSH"
    assert blocks[-1]["spacer_col"] == blocks[-1]["year_endcol"] + 1


def test_is_entity_recon_sheet_excludes_group_tabs():
    from databook_workbook import is_entity_recon_sheet

    assert is_entity_recon_sheet("WoSH_BS_Reconciliation") is True
    assert is_entity_recon_sheet("BS_Reconciliation") is False
    assert is_entity_recon_sheet("PL_Reconciliation") is False


def test_reorder_workbook_sheets_entity_recon_left_of_masters():
    from openpyxl import Workbook

    from databook_workbook import reorder_workbook_sheets

    wb = Workbook()
    wb.remove(wb.active)
    wb.create_sheet("Master_BS")
    wb.create_sheet("BS_Reconciliation")
    wb.create_sheet("Alpha_BS_Reconciliation")
    wb.create_sheet("Alpha_PL_Reconciliation")
    wb.create_sheet("Beta_BS_Reconciliation")
    wb.create_sheet("Beta_PL_Reconciliation")

    reorder_workbook_sheets(wb)
    assert wb.sheetnames[:4] == [
        "Alpha_BS_Reconciliation",
        "Alpha_PL_Reconciliation",
        "Beta_BS_Reconciliation",
        "Beta_PL_Reconciliation",
    ]
    assert wb.sheetnames[4] == "Master_BS"
    assert wb.sheetnames[5] == "BS_Reconciliation"


def test_reorder_workbook_sheets_order():
    from openpyxl import Workbook

    from databook_workbook import reorder_workbook_sheets

    wb = Workbook()
    wb.remove(wb.active)
    wb.create_sheet("Cashflow")
    wb.create_sheet("Master_BS")
    wb.create_sheet("__SOURCE__FA_Dec23A")
    wb.create_sheet("__SOURCE__FY25A")
    wb.create_sheet("FTE Development")
    wb.create_sheet("__SOURCE__AR_Jul25A")
    wb.create_sheet("BS_Reconciliation")
    wb.create_sheet("Trade debtors aging")
    wb.create_sheet("FA roll forward")

    reorder_workbook_sheets(wb)
    assert wb.sheetnames == [
        "Master_BS",
        "BS_Reconciliation",
        "Cashflow",
        "FA roll forward",
        "FTE Development",
        "Trade debtors aging",
        "__SOURCE__AR_Jul25A",
        "__SOURCE__FA_Dec23A",
        "__SOURCE__FY25A",
    ]
