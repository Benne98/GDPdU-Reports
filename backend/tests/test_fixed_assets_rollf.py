"""Tests for fixed_assets_rollf.py — bridge layout, formulas, multi-period."""
import tempfile
import unittest
from pathlib import Path

import openpyxl

from fixed_assets_rollf import (
    TABLE_TITLE_ROW,
    build_bridge_columns,
    compute_closing,
    fa_column_layout,
    grid_label_to_display,
    normalize_config,
    run_fixed_assets_rollf,
    source_sheet_name,
    visible_movements_for_agg,
    HEADER_ROW,
    HEADER_ROW_7,
    POS_COL,
)

ANLAGEN_DIR = Path(__file__).resolve().parent.parent.parent / "Desktop" / "Anlagengitter"
EXAMPLE_2024 = ANLAGEN_DIR / "anlagengitter_2024.xlsx"

DEP_COLS = ["Afa des Jahres", "Afa Abgang"]
CHECK_RED = "DC2626"


def _cell_font_rgb(cell) -> str:
    color = cell.font.color
    if color is None:
        return ""
    rgb = color.rgb
    if rgb is None:
        return ""
    return str(rgb).upper()


def _find_total_row(ws, pos_col: int, label: str = "Fixed assets") -> int | None:
    for r in range(9, ws.max_row + 1):
        if ws.cell(r, pos_col).value == label:
            return r
    return None


def _example_cfg(**overrides):
    cfg = {
        "title": "Test Project",
        "company": "Co",
        "case_id": "fa_test",
        "periods": [
            {
                "label": "FY2024",
                "file_path": str(EXAMPLE_2024),
                "sheet_name": "Anlagengitter 2024",
            }
        ],
        "columns": {
            "opening": "Buchwert GJ-Beg",
            "additions": "Zugang",
            "disposals": "Abgang",
            "depreciation_cols": DEP_COLS,
        },
        "group_cols": ["Bilanzposition"],
        "amount_scale": 1000,
        "header_row": 0,
        "formula_mode": True,
    }
    cfg.update(overrides)
    return cfg


def _multi_period_cfg(**overrides):
    cfg = {
        "title": "Desktop Test",
        "company": "Group",
        "case_id": "fa_multiyear",
        "periods": [
            {
                "label": "FY2023",
                "file_path": str(ANLAGEN_DIR / "anlagengitter_2023.xlsx"),
                "sheet_name": "Anlagengitter 2023",
            },
            {
                "label": "FY2024",
                "file_path": str(ANLAGEN_DIR / "anlagengitter_2024.xlsx"),
                "sheet_name": "Anlagengitter 2024",
            },
            {
                "label": "YTD2025",
                "file_path": str(ANLAGEN_DIR / "anlagengitter_2025.xlsx"),
                "sheet_name": "Anlagengitter 2025",
            },
        ],
        "columns": {
            "opening": "Buchwert GJ-Beg",
            "additions": "Zugang",
            "disposals": "Abgang",
            "depreciation_cols": DEP_COLS,
        },
        "group_cols": ["Bilanzposition"],
        "amount_scale": 1000,
        "header_row": 0,
        "formula_mode": True,
        "balance_check_col": "Lfd Buchwert",
    }
    cfg.update(overrides)
    return cfg


class TestFaRollfHelpers(unittest.TestCase):
    def test_grid_label_to_display(self):
        self.assertEqual(grid_label_to_display("FY2024"), "Dec23A")
        self.assertEqual(grid_label_to_display("FY2023"), "Dec22A")
        self.assertEqual(grid_label_to_display("YTD2025"), "Dec24A")

    def test_source_sheet_name_includes_fa_prefix(self):
        self.assertEqual(source_sheet_name("FY2024"), "__SOURCE__FA_Dec23A")

    def test_compute_closing(self):
        row = {
            "opening": 100.0,
            "additions": 20.0,
            "disposals": -5.0,
            "depreciation": -10.0,
        }
        self.assertEqual(compute_closing(row), 105.0)

    def test_bridge_columns_single_period(self):
        agg = {("A",): {"opening": 1, "additions": 0, "disposals": 0, "depreciation": 0, "closing": 1}}
        cols, spacers = build_bridge_columns(["FY2024"], [agg])
        self.assertEqual(len(cols), 1)
        self.assertEqual(cols[0].kind, "period")
        self.assertEqual(cols[0].header, "Dec23A")
        self.assertEqual(spacers, [])

    def test_bridge_columns_multi_period(self):
        agg = {
            ("A",): {
                "opening": 1,
                "additions": 2,
                "disposals": -1,
                "depreciation": -1,
                "closing": 1,
            }
        }
        cols, _ = build_bridge_columns(["FY2023", "FY2024", "YTD2025"], [agg, agg, agg])
        headers = [c.header for c in cols]
        self.assertEqual(
            headers,
            ["Dec22A", "Add.", "Disp.", "D&A", "Dec23A", "Add.", "Disp.", "D&A", "Dec24A"],
        )

    def test_visible_movements_hides_zero_block(self):
        agg = {
            (): {
                "opening": 10,
                "additions": 0,
                "disposals": 0,
                "depreciation": -1,
                "closing": 9,
            }
        }
        visible = visible_movements_for_agg(agg)
        self.assertIn("depreciation", visible)
        self.assertNotIn("additions", visible)
        self.assertNotIn("disposals", visible)

    def test_bs_recon_aggregated_ref_for_fa_check(self):
        import sys
        from pathlib import Path

        scripts = Path(__file__).resolve().parent.parent.parent / "scripts"
        if str(scripts) not in sys.path:
            sys.path.insert(0, str(scripts))
        from databook_excel_layout import (
            AGGREGATED_BLOCK_TITLE,
            RECON_BLOCK_TITLE_ROW,
            RECON_HEADER_ROW,
            bs_recon_aggregated_ref,
            find_bs_recon_row_contains,
            find_recon_block_start_col,
        )

        from databook_excel_layout import LAYOUT_NA

        wb = openpyxl.Workbook()
        ws = wb.create_sheet("BS_Reconciliation")
        ws.cell(RECON_BLOCK_TITLE_ROW, 10, AGGREGATED_BLOCK_TITLE)
        ws.cell(RECON_HEADER_ROW, 10, "Dec23A")
        ws.cell(15, LAYOUT_NA.pos_col, "Fixed assets")
        agg_start = find_recon_block_start_col(ws, AGGREGATED_BLOCK_TITLE)
        fa_row = find_bs_recon_row_contains(ws, "Fixed assets")
        self.assertEqual(agg_start, 10)
        self.assertEqual(fa_row, 15)
        self.assertEqual(bs_recon_aggregated_ref(fa_row, agg_start), "='BS_Reconciliation'!J15")


class TestFaRollfIntegration(unittest.TestCase):
    @unittest.skipUnless(EXAMPLE_2024.is_file(), "Desktop Anlagengitter example missing")
    def test_single_period_formulas(self):
        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp) / "out.xlsx"
            path = run_fixed_assets_rollf(_example_cfg(output_file_path=str(out)))
            wb = openpyxl.load_workbook(path, data_only=False)
            ws = wb.active
            layout = fa_column_layout(1)
            self.assertEqual(ws.cell(HEADER_ROW, layout.first_data_col).value, "Dec23A")
            self.assertEqual(ws.cell(TABLE_TITLE_ROW, layout.pos_col).value, "Co | FA roll forward Dec23A - Dec23A")
            self.assertIsNotNone(ws.cell(HEADER_ROW_7, layout.first_data_col).fill.fgColor.rgb)
            detail_row = 9
            val = ws.cell(detail_row, layout.first_data_col).value
            self.assertIsInstance(val, str)
            self.assertIn("SUMIFS", val.upper())
            self.assertNotIn("__GROUP__", wb["__SOURCE__FA_Dec23A"].cell(1, 1).value or "")

    @unittest.skipUnless(
        all((ANLAGEN_DIR / f"anlagengitter_{y}.xlsx").is_file() for y in (2023, 2024, 2025)),
        "Multi-year Anlagengitter files missing",
    )
    def test_multi_period_bridge(self):
        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp) / "out.xlsx"
            path = run_fixed_assets_rollf(_multi_period_cfg(output_file_path=str(out)))
            wb = openpyxl.load_workbook(path, data_only=False)
            ws = wb.active
            layout = fa_column_layout(1)

            headers = []
            for c in range(layout.first_data_col, ws.max_column + 1):
                v = ws.cell(HEADER_ROW, c).value
                if v:
                    headers.append(v)
            self.assertEqual(headers[0], "Dec22A")
            self.assertEqual(headers[-1], "Dec24A")
            self.assertIn("Dec23A", headers)
            self.assertNotIn("Opening", headers)
            self.assertIsNone(ws.cell(HEADER_ROW_7, layout.first_data_col).value)

            period_fill = ws.cell(HEADER_ROW, layout.first_data_col).fill.fgColor.rgb
            self.assertIn("F1F5F9", str(period_fill).upper())

            dec23_col = None
            for c in range(layout.first_data_col, ws.max_column + 1):
                if ws.cell(HEADER_ROW, c).value == "Dec23A":
                    dec23_col = c
                    break
            self.assertIsNotNone(dec23_col)
            bridge_formula = ws.cell(9, dec23_col).value
            self.assertIsInstance(bridge_formula, str)
            self.assertIn("SUMIFS", str(bridge_formula).upper())
            self.assertIn("$B9", str(bridge_formula))

            total_row = _find_total_row(ws, layout.pos_col)
            self.assertIsNotNone(total_row)
            total_formula = ws.cell(total_row, layout.first_data_col).value
            self.assertIn("SUM", str(total_formula).upper())

            check_delta_row = total_row + 5
            check_bsrec_delta_row = total_row + 8
            self.assertEqual(ws.cell(check_delta_row, layout.pos_col).value, "Check")
            self.assertNotIn(CHECK_RED, _cell_font_rgb(ws.cell(check_delta_row, layout.pos_col)))
            self.assertIn(CHECK_RED, _cell_font_rgb(ws.cell(check_delta_row, layout.first_data_col)))
            self.assertIn(CHECK_RED, _cell_font_rgb(ws.cell(check_bsrec_delta_row, layout.first_data_col)))

            check_hidden = ws.row_dimensions[total_row + 4].hidden
            self.assertTrue(check_hidden)

            bridge_row = total_row + 10
            self.assertEqual(ws.cell(bridge_row, layout.pos_col).value, "Bridge check")
            self.assertTrue(ws.row_dimensions[total_row + 6].hidden)
            dec23_col = next(
                c
                for c in range(layout.first_data_col, ws.max_column + 1)
                if ws.cell(HEADER_ROW, c).value == "Dec23A"
            )
            bridge_delta = ws.cell(bridge_row, dec23_col).value
            self.assertIn("-(", str(bridge_delta))

            src_sheet_names = [s for s in wb.sheetnames if s.startswith("__SOURCE__")]
            self.assertEqual(len(src_sheet_names), 3)

    @unittest.skipUnless(
        all((ANLAGEN_DIR / f"anlagengitter_{y}.xlsx").is_file() for y in (2023, 2024, 2025)),
        "Multi-year Anlagengitter files missing",
    )
    def test_two_level_hierarchy(self):
        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp) / "out_2level.xlsx"
            path = run_fixed_assets_rollf(
                _multi_period_cfg(
                    output_file_path=str(out),
                    group_cols=["Geschäftsbereich", "Anlagenklasse"],
                )
            )
            wb = openpyxl.load_workbook(path, data_only=False)
            ws = wb.active
            layout = fa_column_layout(2)
            labels = [ws.cell(r, layout.pos_col).value for r in range(9, 9 + 20)]
            indented = [l for l in labels if l and str(l).startswith("    ")]
            self.assertTrue(len(indented) >= 1)
            formula = ws.cell(9, layout.first_data_col).value
            self.assertIn("$B9", str(formula))
            self.assertIn("$D9", str(formula))

            total_row = _find_total_row(ws, layout.pos_col)
            self.assertIsNotNone(total_row)
            check_delta_row = total_row + 5
            bridge_row = total_row + 10
            dec23_col = next(
                c
                for c in range(layout.first_data_col, ws.max_column + 1)
                if ws.cell(HEADER_ROW, c).value == "Dec23A"
            )
            self.assertEqual(ws.cell(check_delta_row, layout.pos_col).value, "Check")
            self.assertNotIn(CHECK_RED, _cell_font_rgb(ws.cell(check_delta_row, layout.pos_col)))
            self.assertIn(CHECK_RED, _cell_font_rgb(ws.cell(check_delta_row, layout.first_data_col)))
            self.assertIn(CHECK_RED, _cell_font_rgb(ws.cell(bridge_row, dec23_col)))


class TestFaRollfNormalizeConfig(unittest.TestCase):
    def test_normalize_requires_columns(self):
        with self.assertRaises(ValueError):
            normalize_config({"periods": [{"label": "FY2024", "file_path": "/x"}]})

    def test_normalize_depreciation_cols(self):
        cfg = normalize_config(
            {
                "periods": [{"label": "FY2024", "file_path": "/x"}],
                "columns": {
                    "opening": "A",
                    "additions": "B",
                    "disposals": "C",
                    "depreciation_cols": ["D1", "D2"],
                },
            }
        )
        self.assertEqual(cfg["columns"]["depreciation_cols"], ["D1", "D2"])


if __name__ == "__main__":
    unittest.main()
