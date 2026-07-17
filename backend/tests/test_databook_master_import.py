"""Master import (quick path) replaces session workbook 1:1."""

from __future__ import annotations

import sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[2]
BACKEND = ROOT / "backend"
for p in (ROOT, BACKEND):
    if str(p) not in sys.path:
        sys.path.insert(0, str(p))

from databook_helpers import import_master_workbook_sheets, mark_master_workbook_imported  # noqa: E402


def test_import_master_workbook_sheets_replaces_workbook(tmp_path):
    src = tmp_path / "upload.xlsx"
    df = pd.DataFrame(
        {
            "Entity": ["E1"],
            "Account": ["1000"],
            "Account description": ["Cash"],
            "L1 - BS/PL": ["BS"],
            "L2": ["CA"],
            "L3": ["Cash"],
            "L4": ["Cash"],
            "L5": [""],
            "L6": ["Reported"],
            "NA": ["CA"],
            "FY24A": [100.0],
            "YTD24A": [250.0],
            "YTD25A": [300.0],
        }
    )
    with pd.ExcelWriter(src, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="Master_BS", index=False)
        df.to_excel(writer, sheet_name="Master_PL", index=False)

    target = tmp_path / "session_master.xlsx"
    wb = Workbook()
    wb.create_sheet("PL_Reconciliation")
    wb.remove(wb.active)
    wb.save(target)

    import_master_workbook_sheets(src, target)

    out = load_workbook(target, data_only=True)
    assert out.sheetnames == ["Master_BS", "Master_PL"]
    ws = out["Master_BS"]
    headers = {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1)}
    ytd24 = ws.cell(2, headers["YTD24A"]).value
    ytd25 = ws.cell(2, headers["YTD25A"]).value
    assert float(ytd24) == 250.0
    assert float(ytd25) == 300.0
    out.close()


def test_mark_master_workbook_imported_flag(tmp_path, monkeypatch):
    from databook_helpers import (
        clear_master_workbook_imported,
        mark_master_workbook_imported,
        master_workbook_imported,
        should_preserve_master_sheets,
    )

    session_id = "test_session_flag"
    uploads = tmp_path / "uploads" / session_id
    monkeypatch.setattr(
        "databook_helpers.PROJECT_ROOT",
        tmp_path,
    )
    master = tmp_path / "output" / "Project_Master.xlsx"

    assert not master_workbook_imported(session_id)
    mark_master_workbook_imported(session_id, master)
    assert master_workbook_imported(session_id)
    assert should_preserve_master_sheets(session_id)
    clear_master_workbook_imported(session_id)
    assert not master_workbook_imported(session_id)
