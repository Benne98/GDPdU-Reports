"""FTE Development v2 — formula spec + golden aggregation (months_col, sum_components).

Financial proof (months_col tenure):
  FTE per row = active_months_in_FY × Beschäftigungsgrad / 100 / 12
  Display = ROUND(SUM(FTE per row), 0) per dimension × FY.

Payroll (sum_components):
  payroll per row = SUM(component_cols) + optional social_col
  Display = SUM(payroll) / 1000 (EURk, negative in output).

Worked example (FY22, synthetic fixture):
  Admin: 6 months × 50% / 12 = 0.25 FTE → 0; payroll 30k+6k = 36 EURk
  Sales: 12 months × 100% / 12 = 1 FTE → 1; payroll 60k+12k = 72 EURk
"""
from __future__ import annotations

import os
import sys
import tempfile
import unittest
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

import fte_development_verformelt as fte  # noqa: E402

FIXTURE_ROWS_FY22 = [
    {
        "Entity": "E1",
        "Bereich": "Sales",
        "Beschäftigungsgrad": 100,
        "Summe": 12,
        "Grundgehalt": 60_000,
        "Sozialversicherung": 12_000,
    },
    {
        "Entity": "E1",
        "Bereich": "Admin",
        "Beschäftigungsgrad": 50,
        "Summe": 6,
        "Grundgehalt": 30_000,
        "Sozialversicherung": 6_000,
    },
]
FIXTURE_ROWS_FY23 = [
    {
        "Entity": "E1",
        "Bereich": "Sales",
        "Beschäftigungsgrad": 100,
        "Summe": 12,
        "Grundgehalt": 66_000,
        "Sozialversicherung": 13_200,
    },
]

GOLDEN_FTE = {
    (2022, "Admin"): 0,
    (2022, "Sales"): 1,
    (2023, "Sales"): 1,
}
GOLDEN_PAYROLL_EURK = {
    (2022, "Admin"): 36.0,
    (2022, "Sales"): 72.0,
    (2023, "Sales"): 79.2,
}


def _write_fixture(path: Path, rows: list[dict]) -> None:
    pd.DataFrame(rows).to_excel(path, index=False)


def _base_cfg(tmpdir: str, p22: str, p23: str) -> dict:
    return {
        "case_id": "golden_fte",
        "output_file_path": tmpdir,
        "file_path": p22,
        "sheet_name": "Sheet1",
        "first_fy": 2022,
        "last_fy": 2023,
        "upload_mode": "per_fy_grid",
        "resolved_entity_year_files": [
            {"fy_label": "FY2022", "paths": [p22]},
            {"fy_label": "FY2023", "paths": [p23]},
        ],
        "fte_mapping": {
            "employment_pct_col": "Beschäftigungsgrad",
            "months_col": "Summe",
            "tenure_mode": "months_col",
        },
        "fte_tenure_mode": "months_col",
        "payroll_mapping": {
            "component_cols": ["Grundgehalt"],
            "social_col": "Sozialversicherung",
        },
        "fte_payroll_mode": "sum_components",
        "dimensions": [{"source_col": "Bereich", "output_label": "Department"}],
        "preset_metrics": ["fte", "payroll"],
        "formula_mode": True,
    }


class TestFteRowCalculations(unittest.TestCase):
    def test_compute_row_fte_months_col(self):
        cfg = _base_cfg("", "", "")
        row = pd.Series(FIXTURE_ROWS_FY22[1])
        self.assertAlmostEqual(fte.compute_row_fte(row, cfg), 0.25)

    def test_compute_row_payroll_sum_components(self):
        cfg = _base_cfg("", "", "")
        row = pd.Series(FIXTURE_ROWS_FY22[0])
        self.assertAlmostEqual(fte.compute_row_payroll(row, cfg), 72_000.0)

    def test_normalize_config_requires_dimensions(self):
        with self.assertRaises(ValueError):
            fte.normalize_config({"dimensions": []})


class TestFteGoldenWorkbook(unittest.TestCase):
    def setUp(self) -> None:
        self.tmp = tempfile.mkdtemp()
        self.p22 = os.path.join(self.tmp, "fy22.xlsx")
        self.p23 = os.path.join(self.tmp, "fy23.xlsx")
        _write_fixture(Path(self.p22), FIXTURE_ROWS_FY22)
        _write_fixture(Path(self.p23), FIXTURE_ROWS_FY23)

    def test_merge_and_aggregate_matches_golden(self):
        cfg = _base_cfg(self.tmp, self.p22, self.p23)
        df = fte.merge_uploads(cfg)
        df["_fte"] = df.apply(lambda r: fte.compute_row_fte(r, cfg), axis=1)
        df["_pay"] = df.apply(lambda r: fte.compute_row_payroll(r, cfg), axis=1)
        for (year, dept), exp_fte in GOLDEN_FTE.items():
            sub = df[(df[fte.HELPER_FY] == year) & (df["Bereich"] == dept)]
            got = round(float(sub["_fte"].sum()), 0)
            self.assertEqual(got, exp_fte, f"FTE {year}/{dept}")
        for (year, dept), exp_pay in GOLDEN_PAYROLL_EURK.items():
            sub = df[(df[fte.HELPER_FY] == year) & (df["Bereich"] == dept)]
            got = round(float(sub["_pay"].sum()) / 1000, 1)
            self.assertAlmostEqual(got, exp_pay, places=1, msg=f"Payroll {year}/{dept}")

    def test_build_workbook_structure_and_formulas(self):
        cfg = _base_cfg(self.tmp, self.p22, self.p23)
        out = fte.build_fte_workbook(cfg)
        self.assertTrue(os.path.isfile(out))
        wb = load_workbook(out)
        self.assertIn("__SOURCE__", wb.sheetnames)
        self.assertIn("FTE Development", wb.sheetnames)
        ws_src = wb["__SOURCE__"]
        headers = [ws_src.cell(1, c).value for c in range(1, (ws_src.max_column or 0) + 1)]
        self.assertIn("_fte_avg", headers)
        self.assertIn("_payroll", headers)
        self.assertIn("_fte_fy", headers)
        ws = wb["FTE Development"]
        self.assertEqual(ws.cell(2, 2).value, "FY22A")
        self.assertEqual(ws.cell(2, 3).value, "FY23A")
        sales_fte_22 = str(ws.cell(5, 2).value or "")
        self.assertIn("ROUND", sales_fte_22)
        self.assertIn("SUMIFS", sales_fte_22)
        sales_pay_22 = str(ws.cell(9, 2).value or "")
        self.assertIn("/1000", sales_pay_22)
        wb.close()


if __name__ == "__main__":
    unittest.main()
