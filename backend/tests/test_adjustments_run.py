"""Smoke test: Adjustments.py run(config) against template + minimal master."""
from __future__ import annotations

import json
import sys
import tempfile
import unittest
from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[2]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))
if str(ROOT / "backend") not in sys.path:
    sys.path.insert(0, str(ROOT / "backend"))

from databook_helpers import CONSOLIDATION_META_COLS, build_adjustments_template_from_config  # noqa: E402

# Import after path setup
sys.path.insert(0, str(ROOT))
from Adjustments import run  # noqa: E402


def _write_minimal_master(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Master_PL"
    headers = list(CONSOLIDATION_META_COLS) + ["L5", "FY24A", "FY25A"]
    for c, h in enumerate(headers, start=1):
        ws.cell(1, c, h)
    ws.cell(2, 1, "Entity A")
    ws.cell(2, 5, "PL")
    ws.cell(2, 8, "Reported")
    ws.cell(2, 9, 100)
    ws.cell(2, 10, 200)
    ws.cell(3, 1, "Total PL")
    wb.save(path)


class TestAdjustmentsRun(unittest.TestCase):
    def test_run_updates_master_from_template(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            master = tmp_path / "master.xlsx"
            _write_minimal_master(master)

            adj_out = tmp_path / "adj_template.xlsx"
            build_adjustments_template_from_config(
                {
                    "output_path": str(adj_out),
                    "first_fy": 2024,
                    "ltm_month": "2025-12",
                    "fy_end_month": 12,
                    "fiscal_start_month": 1,
                }
            )

            adj_df = pd.read_excel(adj_out, sheet_name="Adjustments")
            adj_df.loc[0] = [
                "AdjCo",
                "9000",
                "One-off",
                "PL",
                "Adj",
                "Detail",
                "Line",
                50,
                75,
            ]
            filled = tmp_path / "adj_filled.xlsx"
            with pd.ExcelWriter(filled, engine="openpyxl") as writer:
                adj_df.to_excel(writer, sheet_name="Adjustments", index=False)

            run({"master_path": str(master), "adjustments_path": str(filled)})

            wb = load_workbook(master)
            ws = wb["Master_PL"]
            texts = [ws.cell(r, 1).value for r in range(1, ws.max_row + 1)]
            self.assertIn("Adjustments", texts)
            self.assertIn("Total PL", texts)
            self.assertIn("AdjCo", texts)

            fy24_col = next(
                c for c in range(1, ws.max_column + 1) if ws.cell(1, c).value == "FY24A"
            )
            adj_value = None
            for r in range(1, ws.max_row + 1):
                if ws.cell(r, 1).value == "AdjCo":
                    adj_value = ws.cell(r, fy24_col).value
                    break
            self.assertEqual(adj_value, 0.05)


if __name__ == "__main__":
    unittest.main()
