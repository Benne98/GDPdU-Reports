"""Churn GST-style single-column hierarchy."""
import unittest

import pandas as pd
from openpyxl import Workbook

from Churn import (
    build_churn_display_order,
    build_churn_hierarchy_rows,
    finalize_merged_bridge_layout,
    format_churn_excel,
    normalize_config,
    write_churn_export_to_workbook,
)
from churn_verformelt import apply_churn_formulas


def _three_level_leaves() -> pd.DataFrame:
    return pd.DataFrame({
        "Region": ["Americas", "Americas", "Europe"],
        "Country": ["USA", "USA", "Germany"],
        "City": ["NYC", "LA", "Berlin"],
        "FY23A": [100, 200, 300],
        "Upsell": [0, 0, 0],
        "Downsell": [0, 0, 0],
        "Cross-sell": [0, 0, 0],
        "Lost": [0, 0, 0],
        "NRR": [0, 0, 0],
        "New": [0, 0, 0],
        "FY24A": [1000, 2000, 3000],
    })


def _base_cfg(**overrides):
    cfg = {
        "fy_end_year": 2024,
        "fy_end_month": 12,
        "fy_end_day": 31,
        "value_col": "Amount",
        "customer_col": "Customer",
        "product_col": "Product",
        "start_col": "Start",
        "end_col": "End",
        "invoice_col": "Invoice",
        "total_label": "Total",
    }
    cfg.update(overrides)
    return normalize_config(cfg)


class TestChurnHierarchy(unittest.TestCase):
    def test_single_label_column_three_levels(self):
        cfg = _base_cfg(group_cols=["Region", "Country", "City"])
        hier = build_churn_hierarchy_rows(_three_level_leaves(), cfg)
        self.assertIn("label", hier.columns)
        self.assertNotIn("Region", hier.columns)
        self.assertNotIn("Country", hier.columns)
        self.assertNotIn("City", hier.columns)
        self.assertIn("key", hier.columns)
        self.assertEqual(len(hier[hier["row_type"] == "leaf"]), 3)
        self.assertEqual(len(hier[hier["row_type"] == "hierarchy"]), 4)

    def test_display_order_post_order(self):
        cfg = _base_cfg(
            group_cols=["Region", "Country", "City"],
            sort={"top_level_by": "ARR2", "top_level_desc": True},
        )
        hier = build_churn_hierarchy_rows(_three_level_leaves(), cfg)
        ordered = build_churn_display_order(hier, cfg)
        labels = ordered["label"].tolist()
        nyc_idx = labels.index("NYC")
        la_idx = labels.index("LA")
        usa_idx = labels.index("USA")
        americas_idx = labels.index("Americas")
        self.assertLess(nyc_idx, usa_idx)
        self.assertLess(la_idx, usa_idx)
        self.assertLess(usa_idx, americas_idx)

    def test_level1_row_exists(self):
        cfg = _base_cfg(group_cols=["Region", "Country", "City"])
        out = finalize_merged_bridge_layout(_three_level_leaves(), cfg)
        level1 = out[out["level"] == 1]
        self.assertGreaterEqual(len(level1), 1)
        self.assertEqual(out.iloc[-1]["row_type"], "total")

    def test_finalize_sort_uses_rightmost_fy(self):
        cfg = _base_cfg(
            group_cols=["Entity"],
            sort={"top_level_by": "ARR2", "top_level_desc": True},
        )
        raw = pd.DataFrame({
            "Entity": ["Low", "High"],
            "FY23A": [1, 100],
            "Upsell": [0, 0],
            "Downsell": [0, 0],
            "Cross-sell": [0, 0],
            "Lost": [0, 0],
            "NRR": [0, 0],
            "New": [0, 0],
            "FY24A": [10, 1000],
        })
        out = finalize_merged_bridge_layout(raw, cfg)
        leaves = out[out["row_type"] == "leaf"]["Entity"].tolist()
        self.assertEqual(leaves[0], "High")

    def test_parent_formula_sums_children(self):
        cfg = _base_cfg(
            group_cols=["Region", "Country", "City"],
            formula_mode=True,
            first_fy=2023,
            current_year=2024,
            current_month=12,
        )
        merged = finalize_merged_bridge_layout(_three_level_leaves(), cfg)
        wb = Workbook()
        wb.active.title = "Churn"
        write_churn_export_to_workbook(wb, merged, "Churn", cfg, formula_mode=True)
        layout = format_churn_excel(wb, "Churn", merged, cfg, formula_mode=True)
        layout["report_sheet"] = "Churn"

        src = wb.create_sheet("Source")
        for i, h in enumerate(["Region", "Country", "City", "FY23", "FY24"], start=1):
            src.cell(row=1, column=i).value = h

        ws = wb["Churn"]
        label_col = layout["label_col"]
        row_type_col = layout["row_type_col"]
        fy_col = layout["fy_col_indices"][0][1]

        def row_for_label(label: str) -> int:
            for r in range(layout["data_start_row"], layout["table_bottom_row"] + 1):
                if str(ws.cell(row=r, column=label_col).value or "").strip().lower() == label.lower():
                    return r
            raise AssertionError(f"Row for label {label!r} not found")

        usa_row = row_for_label("USA")
        nyc_row = row_for_label("NYC")
        self.assertEqual(str(ws.cell(row=usa_row, column=row_type_col).value).lower(), "hierarchy")
        self.assertEqual(str(ws.cell(row=nyc_row, column=row_type_col).value).lower(), "leaf")

        apply_churn_formulas(wb, layout, cfg, "Source", {2023: "D", 2024: "E"})
        usa_fy = ws.cell(row=usa_row, column=fy_col).value
        nyc_fy = ws.cell(row=nyc_row, column=fy_col).value
        self.assertIsInstance(usa_fy, str)
        self.assertIn("SUM(", str(usa_fy))
        self.assertIsInstance(nyc_fy, str)
        self.assertIn("SUMIFS", str(nyc_fy).upper())


if __name__ == "__main__":
    unittest.main()
