"""
Insert adjustment block into Master_PL of the SuSa master workbook.
Driven by JSON config (first CLI argument) from the FDD backend.

Expects:
- master_path: existing SuSa master after consolidation (in-place update)
- adjustments_path: uploaded template (sheet 'Adjustments', meta + FY columns)

The master must already contain L5; this script does not add columns.
"""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

PROJECT_ROOT = Path(__file__).resolve().parent
SCRIPTS_DIR = PROJECT_ROOT / "scripts"
if str(SCRIPTS_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPTS_DIR))

from databook_amounts import FMT_KEUR, scale_to_keur  # noqa: E402

# --------------------------------------------------
# FORMAT STANDARDS
# --------------------------------------------------
BASE_FONT = Font(name="GT Walsheim LC Light", size=8)
HEADER_FONT = Font(name="GT Walsheim LC Light", size=8, bold=True)

FMT_AMOUNT = FMT_KEUR

WHITE_FILL = PatternFill("solid", fgColor="FFFFFF")
PURPLE_FILL = PatternFill("solid", fgColor="C9A0DC")

HEADER_FILL = PatternFill("solid", fgColor="FFF2F2F2")  # #F2F2F2
BOTTOM_BORDER = Border(bottom=Side(style="thin"))

TEXT_COLS = {
    "Entity",
    "Account",
    "Account description",
    "L1 - BS/PL",
    "L2",
    "L3",
    "L4",
    "L5",
}

MONTH_RE = re.compile(
    r"^(Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)-\d{4}$"
)
FY_RE = re.compile(r"^FY\d{2,4}A?$")

# SuSa master canvas extends beyond the last period header (see SuSabyYear.py).
CANVAS_EXTRA_COLS = 20


# --------------------------------------------------
# Helper functions
# --------------------------------------------------
def find_row(ws, text):
    for r in range(1, ws.max_row + 1):
        if ws.cell(r, 1).value == text:
            return r
    raise ValueError(f"'{text}' not found in sheet '{ws.title}'.")


def find_row_optional(ws, text):
    for r in range(1, ws.max_row + 1):
        if ws.cell(r, 1).value == text:
            return r
    return None


def get_table_col_limit(ws) -> int:
    """
    Last column with a non-empty header in row 1.

    This is more robust than using '-' as the boundary because headers such as
    'L1 - BS/PL' also contain a hyphen but are not period columns.
    """
    last = 1
    for c in range(1, ws.max_column + 1):
        val = ws.cell(1, c).value
        if val is not None and str(val).strip():
            last = c
    return last


def get_required_col(ws, header_name: str) -> int:
    """
    Return the column index of an existing header.
    This script no longer inserts L5 itself. The master creation logic must
    already provide the L5 column.
    """
    for c in range(1, ws.max_column + 1):
        if ws.cell(1, c).value == header_name:
            return c
    raise ValueError(
        f"Required column '{header_name}' not found in sheet '{ws.title}'. "
        f"Please ensure the master output already contains this column."
    )


def get_format_limit_col(ws, l5_col=None) -> int:
    """
    Rightmost column to paint white / purple backgrounds.
    Matches Consolidation.py and the SuSa master canvas extension.
    """
    table = get_table_col_limit(ws)
    limit = min(max(ws.max_column, table), table + CANVAS_EXTRA_COLS)
    if l5_col is not None:
        limit = max(limit, l5_col)
    return limit


def is_amount_header(header) -> bool:
    """
    Amount columns are FY columns and monthly period columns such as Jan-2024.
    Non-period text columns are excluded.
    """
    if header is None:
        return False

    h = str(header).strip()
    if not h or h in TEXT_COLS:
        return False

    if FY_RE.fullmatch(h):
        return True

    return bool(MONTH_RE.fullmatch(h))


def fy_columns_from_df(df: pd.DataFrame) -> list[str]:
    return [c for c in df.columns if isinstance(c, str) and FY_RE.fullmatch(c.strip())]


def apply_header_style(ws, l5_col=None):
    """
    Header format: #F2F2F2, bold, bottom border, GT Walsheim LC Light 8.
    """
    limit = get_format_limit_col(ws, l5_col=l5_col)
    for c in range(1, limit + 1):
        cell = ws.cell(1, c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.border = BOTTOM_BORDER
        cell.alignment = Alignment(
            horizontal="left" if cell.value in TEXT_COLS else "right",
            vertical="center",
        )


def _style_data_cell(cell, header):
    cell.font = BASE_FONT
    cell.fill = WHITE_FILL
    if header in TEXT_COLS:
        cell.alignment = Alignment(horizontal="left", vertical="center")
    else:
        cell.alignment = Alignment(horizontal="right", vertical="center")
        if is_amount_header(header):
            cell.number_format = FMT_AMOUNT


def write_df(ws, start_row, df, l5_col=None):
    header_map = {
        ws.cell(1, c).value: c
        for c in range(1, ws.max_column + 1)
        if ws.cell(1, c).value is not None
    }
    limit = get_format_limit_col(ws, l5_col=l5_col)

    for i, (_, row) in enumerate(df.iterrows()):
        r = start_row + i

        for col, val in row.items():
            if col not in header_map:
                continue

            c = header_map[col]
            cell = ws.cell(r, c, val)
            _style_data_cell(cell, col)

        for c in range(1, limit + 1):
            header = ws.cell(1, c).value
            cell = ws.cell(r, c)
            if cell.value is None:
                _style_data_cell(cell, header)


def fill_single_white_row_to_limit(ws, row, l5_col=None):
    """
    Fill one row white across the relevant table area, including L5.
    """
    limit = get_format_limit_col(ws, l5_col=l5_col)
    for c in range(1, limit + 1):
        cell = ws.cell(row, c)
        cell.fill = WHITE_FILL
        cell.font = BASE_FONT


def fill_block_white_to_limit(ws, start_row, end_row, l5_col=None):
    """
    Fill the complete adjustment block white across the relevant table area,
    not only cells with values.
    """
    if end_row < start_row:
        return

    limit = get_format_limit_col(ws, l5_col=l5_col)
    for r in range(start_row, end_row + 1):
        for c in range(1, limit + 1):
            cell = ws.cell(r, c)
            cell.fill = WHITE_FILL
            cell.font = BASE_FONT


def add_adjustment_header(ws, row, l5_col=None):
    limit = get_format_limit_col(ws, l5_col=l5_col)
    for c in range(1, limit + 1):
        cell = ws.cell(row, c)
        cell.fill = PURPLE_FILL
        cell.font = BASE_FONT
        cell.alignment = Alignment(horizontal="left", vertical="center")
        if c == 1:
            cell.value = "Adjustments"


def find_consolidation_block(ws, total_row):
    """
    Find the Consolidation block in Master_PL:
    - Header in column A is 'Consolidation'
    - Start = first non-empty row after the header
    - End = last non-empty row before Total PL

    Returns (cons_header_row, cons_start, cons_end) or (None, None, None).
    """
    header_row = find_row_optional(ws, "Consolidation")
    if header_row is None or header_row >= total_row:
        return None, None, None

    r = header_row + 1
    while r < total_row and (ws.cell(r, 1).value in (None, "")):
        r += 1

    cons_start = r if r < total_row else None
    if cons_start is None:
        return header_row, None, None

    cons_end = None
    for rr in range(cons_start, total_row):
        if ws.cell(rr, 1).value not in (None, ""):
            cons_end = rr

    return header_row, cons_start, cons_end


def validate_adjustment_input(adj_df: pd.DataFrame, fy_cols: list[str]) -> None:
    required_cols = [
        "Entity",
        "Account",
        "Account description",
        "L2",
        "L3",
        "L4",
    ]
    if "L1" not in adj_df.columns and "L1 - BS/PL" not in adj_df.columns:
        required_cols.append("L1 or L1 - BS/PL")

    missing = [c for c in required_cols if c not in adj_df.columns]
    if missing:
        raise ValueError(f"Adjustment input is missing required columns: {missing}")

    if not fy_cols:
        raise ValueError(
            "Adjustment input does not contain any FY columns. "
            "Expected columns such as FY23A, FY24A, FY25A."
        )


def _read_adjustments_excel(path: str) -> pd.DataFrame:
    try:
        return pd.read_excel(path, sheet_name="Adjustments", engine="openpyxl")
    except ValueError:
        return pd.read_excel(path, engine="openpyxl")


def _prepare_adj_pl(
    adj: pd.DataFrame,
    fy_cols: list[str],
    *,
    scale_to_keur_enabled: bool = True,
) -> pd.DataFrame:
    if "L1" in adj.columns:
        l1_norm = adj["L1"].astype(str).str.strip().str.upper()
        adj_pl = adj[l1_norm == "PL"].copy()
    elif "L1 - BS/PL" in adj.columns:
        l1_norm = adj["L1 - BS/PL"].astype(str).str.strip().str.upper()
        adj_pl = adj[l1_norm == "PL"].copy()
    else:
        raise ValueError("Adjustment input must contain 'L1' or 'L1 - BS/PL'.")

    adj_pl["L1 - BS/PL"] = "PL"
    adj_pl["L5"] = "Adjusted"

    result = adj_pl.reindex(columns=[
        "Entity",
        "Account",
        "Account description",
        "L1 - BS/PL",
        "L2",
        "L3",
        "L4",
        "L5",
    ] + fy_cols)
    return scale_to_keur(result, fy_cols, enabled=scale_to_keur_enabled)


def run(config: dict) -> None:
    master_path = str(config.get("master_path") or config.get("output_path") or "").strip()
    adj_path = str(config.get("adjustments_path") or config.get("adj_path") or "").strip()
    if not master_path or not adj_path:
        raise SystemExit("Config must define master_path and adjustments_path.")

    master_file = Path(master_path).expanduser().resolve()
    adj_file = Path(adj_path).expanduser().resolve()
    if not master_file.is_file():
        raise SystemExit(f"Master workbook not found: {master_file}")
    if not adj_file.is_file():
        raise SystemExit(f"Adjustments file not found: {adj_file}")

    adj = _read_adjustments_excel(str(adj_file))
    fy_cols = fy_columns_from_df(adj)
    validate_adjustment_input(adj, fy_cols)

    scale_to_keur_enabled = bool(config.get("scale_to_keur", True))

    wb = load_workbook(master_file)
    ws = wb["Master_PL"]

    old_total = find_row(ws, "Total PL")
    pl_start = 2
    pl_end = old_total - 1

    cons_header_row, cons_start, cons_end = find_consolidation_block(ws, old_total)

    l5_col = get_required_col(ws, "L5")
    apply_header_style(ws, l5_col=l5_col)

    reported_end = pl_end
    if cons_header_row is not None:
        reported_end = cons_header_row - 2
        if reported_end < pl_start:
            reported_end = pl_end

    for r in range(pl_start, reported_end + 1):
        cell = ws.cell(r, l5_col)
        if cell.value in (None, ""):
            cell.value = "Reported"
        cell.font = BASE_FONT
        cell.fill = WHITE_FILL
        cell.alignment = Alignment(horizontal="left", vertical="center")

    ws.delete_rows(old_total, 1)
    insert_row = old_total

    fill_single_white_row_to_limit(ws, insert_row - 1, l5_col=l5_col)

    adj_pl = _prepare_adj_pl(adj, fy_cols, scale_to_keur_enabled=scale_to_keur_enabled)
    if adj_pl.empty:
        raise ValueError("No PL rows found in adjustment input.")

    ws.insert_rows(insert_row, amount=len(adj_pl) + 4)

    add_adjustment_header(ws, insert_row, l5_col=l5_col)
    fill_single_white_row_to_limit(ws, insert_row + 1, l5_col=l5_col)

    adj_start = insert_row + 2
    write_df(ws, adj_start, adj_pl, l5_col=l5_col)
    adj_end = adj_start + len(adj_pl) - 1

    fill_single_white_row_to_limit(ws, adj_end + 1, l5_col=l5_col)
    fill_block_white_to_limit(ws, adj_start, adj_end, l5_col=l5_col)

    for r in range(adj_start, adj_end + 1):
        cell = ws.cell(r, l5_col, "Adjusted")
        cell.font = BASE_FONT
        cell.fill = WHITE_FILL
        cell.alignment = Alignment(horizontal="left", vertical="center")

    total_row = adj_end + 2
    ws.cell(total_row, 1, "Total PL").font = BASE_FONT

    limit = get_format_limit_col(ws, l5_col=l5_col)

    for c in range(1, limit + 1):
        cell = ws.cell(total_row, c)
        cell.fill = WHITE_FILL
        cell.font = BASE_FONT

        header = ws.cell(1, c).value

        if header in TEXT_COLS:
            cell.alignment = Alignment(horizontal="left", vertical="center")
            continue

        if not is_amount_header(header):
            continue

        col = get_column_letter(c)

        expr = (
            f"SUBTOTAL(9,{col}{pl_start}:{col}{reported_end})"
            f"+SUBTOTAL(9,{col}{adj_start}:{col}{adj_end})"
        )

        if cons_start is not None and cons_end is not None:
            expr += f"-SUBTOTAL(9,{col}{cons_start}:{col}{cons_end})"

        cell.value = f"={expr}"
        cell.number_format = FMT_AMOUNT
        cell.alignment = Alignment(horizontal="right", vertical="center")

    wb.save(master_file)
    print(
        f"Adjustments applied to {master_file}: Total PL = Reported (without Consolidation) "
        "+ Adjustments - Consolidation."
    )


def main() -> None:
    if len(sys.argv) < 2:
        raise SystemExit("Usage: python Adjustments.py <config.json>")
    config_path = Path(sys.argv[1])
    config = json.loads(config_path.read_text(encoding="utf-8-sig"))
    run(config)


if __name__ == "__main__":
    main()
