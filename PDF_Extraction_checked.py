"""
Add check rows and formatting to PDF extraction workbook.
JSON config: input_file, output_file
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

SHEETS = ["BS", "PL"]
LEVEL_HEADER = "Level"
GERMAN_HEADER = "German Mapping"
ENGLISH_HEADERS = ["Englisch Mapping", "English Mapping"]
NET_RESULT_KEYWORDS = [
    "net result", "net income", "net income / loss", "net income/loss",
    "net profit", "net loss", "konzernjahresfehlbetrag", "jahresfehlbetrag",
    "jahresüberschuss", "jahresergebnis", "period result",
]
BG_COLOR = None
NUMBER_FORMAT = '#,##0;(#,##0);-'
MANUAL_LABEL = "Please insert net result manually:"


def norm(s):
    return str(s).strip().lower() if s is not None else ""


def find_col(ws, header_name):
    for c in range(1, ws.max_column + 1):
        if str(ws.cell(1, c).value).strip() == header_name:
            return c
    raise ValueError(f"Column '{header_name}' not found")


def find_col_any(ws, header_names):
    for c in range(1, ws.max_column + 1):
        v = str(ws.cell(1, c).value).strip()
        if v in header_names:
            return c
    return None


def find_last_table_row(ws, german_col):
    last = 1
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, german_col).value not in (None, ""):
            last = r
    return last


def find_table_end_col(ws):
    for c in range(ws.max_column, 0, -1):
        v = ws.cell(1, c).value
        if v not in (None, ""):
            return c
    return ws.max_column


def find_first_value_col(ws, col_level):
    start_col = None
    for c in range(1, col_level):
        hdr = str(ws.cell(1, c).value)
        if hdr and "mapping" in hdr.lower():
            start_col = c + 1
    if start_col is None:
        raise ValueError("Start of value columns not found")
    return start_col


def find_net_result_row(ws, last_table_row, german_col, english_col=None):
    found = None
    for r in range(2, last_table_row + 1):
        g = norm(ws.cell(r, german_col).value)
        e = norm(ws.cell(r, english_col).value) if english_col else ""
        if any(k in g for k in NET_RESULT_KEYWORDS) or any(k in e for k in NET_RESULT_KEYWORDS):
            found = r
    return found if found else last_table_row


def infer_bg_color(ws, sample_row, sample_col, fallback="FFFFFF"):
    cell = ws.cell(sample_row, sample_col)
    fill = cell.fill
    if fill and fill.fill_type == "solid":
        fc = fill.fgColor
        if fc is not None and getattr(fc, "type", None) == "rgb" and fc.rgb:
            return fc.rgb[-6:]
    return fallback


def run(config: dict) -> None:
    input_file = Path(str(config.get("input_file") or ""))
    output_file = Path(str(config.get("output_file") or input_file))
    if not input_file.is_file():
        raise SystemExit(f"input_file not found: {input_file}")

    wb = load_workbook(input_file)
    table_last_row = {}
    check_rows_by_sheet = {}
    value_cols_by_sheet = {}
    table_end_col_by_sheet = {}
    format_last_row_by_sheet = {}
    manual_row_by_sheet = {}

    for sheet in SHEETS:
        ws = wb[sheet]
        col_german = find_col(ws, GERMAN_HEADER)
        col_level = find_col(ws, LEVEL_HEADER)
        col_english = find_col_any(ws, ENGLISH_HEADERS)
        last_table_row = find_last_table_row(ws, col_german)
        table_last_row[sheet] = last_table_row
        first_value_col = find_first_value_col(ws, col_level)
        last_value_col = col_level - 1
        value_cols_by_sheet[sheet] = (first_value_col, last_value_col)
        table_end_col_by_sheet[sheet] = find_table_end_col(ws)
        check_start_row = last_table_row + 2

        if sheet == "PL":
            checks = [("L3 - Check", "*L3*"), ("L4 - Check", "*L4*"), ("L1 - Check", None)]
            net_result_row = find_net_result_row(ws, last_table_row, col_german, col_english)
            manual_row = check_start_row + len(checks)
            manual_row_by_sheet[sheet] = manual_row
        else:
            checks = [("L2 - Check", "*L2*"), ("L3 - Check", "*L3*"), ("L4 - Check", "*L4*")]
            net_result_row = None
            manual_row = None

        check_rows = []
        for i, (label, level_pattern) in enumerate(checks):
            r = check_start_row + i
            check_rows.append(r)
            ws.cell(r, col_german, label)
            for c in range(first_value_col, last_value_col + 1):
                col_letter = get_column_letter(c)
                level_letter = get_column_letter(col_level)
                if sheet == "PL" and label == "L1 - Check":
                    net_ref = f"{col_letter}{net_result_row}"
                    man_ref = f"{col_letter}{manual_row}"
                    formula = f"=N({net_ref})-N({man_ref})"
                else:
                    base_sumifs = (
                        f"SUMIFS("
                        f"{col_letter}$2:{col_letter}${last_table_row},"
                        f"{level_letter}$2:{level_letter}${last_table_row},\"{level_pattern}\")"
                    )
                    if sheet == "PL":
                        net_ref = f"{col_letter}{net_result_row}"
                        formula = f"={base_sumifs}-N({net_ref})"
                    else:
                        formula = f"={base_sumifs}"
                cell = ws.cell(r, c, formula)
                cell.number_format = NUMBER_FORMAT

        check_rows_by_sheet[sheet] = check_rows
        if sheet == "PL":
            ws.cell(manual_row, col_german, MANUAL_LABEL)
            for c in range(first_value_col, last_value_col + 1):
                ws.cell(manual_row, c).number_format = NUMBER_FORMAT
            format_last_row_by_sheet[sheet] = manual_row
        else:
            format_last_row_by_sheet[sheet] = max(check_rows)

    font_default = Font(name="Inter", size=8, color="000000")
    font_red = Font(name="Inter", size=8, color="9C0006")

    for sheet in SHEETS:
        ws = wb[sheet]
        col_german = find_col(ws, GERMAN_HEADER)
        last_format_row = format_last_row_by_sheet[sheet]
        check_rows = check_rows_by_sheet[sheet]
        table_end_col = table_end_col_by_sheet[sheet]
        first_value_col, last_value_col = value_cols_by_sheet[sheet]
        bg_hex = BG_COLOR if BG_COLOR else infer_bg_color(ws, 2, col_german, "FFFFFF")
        bg_fill = PatternFill(fill_type="solid", fgColor=bg_hex)

        for r in range(1, last_format_row + 1):
            for c in range(col_german, table_end_col + 1):
                cell = ws.cell(r, c)
                cell.fill = bg_fill
                cell.font = font_default
                if first_value_col <= c <= last_value_col:
                    cell.number_format = NUMBER_FORMAT

        for r in check_rows:
            for c in range(first_value_col, last_value_col + 1):
                col_letter = get_column_letter(c)
                addr = f"{col_letter}{r}"
                ws.conditional_formatting.add(
                    addr,
                    FormulaRule(formula=[f"{addr}<>0"], font=font_red),
                )

    wb.save(output_file)
    print(f"Checked workbook saved to {output_file}")


def main() -> None:
    if len(sys.argv) < 2:
        raise SystemExit("Usage: python PDF_Extraction_checked.py <config.json>")
    config = json.loads(Path(sys.argv[1]).read_text(encoding="utf-8-sig"))
    run(config)


if __name__ == "__main__":
    main()
