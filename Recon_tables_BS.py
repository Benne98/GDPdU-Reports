import re
import sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.styles import Border, Side
from openpyxl.formatting.rule import CellIsRule

PROJECT_ROOT = Path(__file__).resolve().parent
SCRIPTS_DIR = PROJECT_ROOT / "scripts"
if str(SCRIPTS_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPTS_DIR))

from gst_excel_theme import THEME, apply_recon_portfolio_layout  # noqa: E402
from databook_excel_layout import (  # noqa: E402
    LAYOUT_BS,
    apply_bs_hierarchy_borders,
    check_row_groups_after_table,
    collapse_check_portfolio,
    paint_grey_white_canvas,
    refresh_master_sheet,
    write_bs_ale_check_section,
    write_fs_check_section,
    assign_recon_block_columns,
    build_entity_recon_blocks,
)
from report_row_layout import build_bs_row_structure, l2_l3_order_from_mapping  # noqa: E402
from databook_workbook import (  # noqa: E402
    BS_RECON_MAPPING_FILE,
    GROUP_BS_RECON_SHEET,
    MASTER_WORKBOOK_STR,
    remove_stale_entity_recon_sheets,
    sanitize_entity_recon_sheet_name,
)
from databook_runtime import apply_custom_entity_order, load_argv_config, path_value  # noqa: E402

BACKEND_DIR = PROJECT_ROOT / "backend"
if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))

from databook_periods import display_bs_period_labels  # noqa: E402

_argv_cfg = load_argv_config()

# =============================================
# CONFIG (BS) — Desktop work defaults
# =============================================
DESKTOP_DIR = PROJECT_ROOT / "Desktop"

PROJECT_NAME = "Desktop Test"
GROUP_NAME   = "Group"

YEARS = []
source_col_bs = None
DISPLAY_YEARS: list[str] = []  # detected from Master_BS FY columns
UNIT_LABEL = "kEUR"

REC_PERIOD_FROM = ""
REC_PERIOD_TO   = ""

ENTITY_SORT_ORDER = {}

ENTITY_RENAME_MAP = {}

IC_MASTER_ENTITY = "Consolidation"
IC_DISPLAY_NAME  = "IC eliminations"

DIFF_KEY = "Difference"
FS_KEY   = "Financial statements"

# =============================================
# BS TOTALS CONFIG (components refer to L2 labels)
# =============================================
BS_TOTALS_CONFIG = [
    {"label": "Total assets",             "components": ["Fixed assets", "Current assets"], "insert_after": "Current assets"},
    {"label": "Total equity & liabilities","components": ["Equity", "Liabilities"],          "insert_after": "Liabilities"},
]

# =============================================
# BS CHECK VALUES (Financial statements) - per year
# =============================================
SHOW_BS_CHECKS = True

BS_FS_CHECK_TOTAL_ASSETS = {}

FY_END_MONTH = int(_argv_cfg.get("fy_end_month") or 12)
LTM_MONTH = _argv_cfg.get("ltm_month") or "2023-7"
if "show_bs_checks" in _argv_cfg:
    SHOW_BS_CHECKS = bool(_argv_cfg["show_bs_checks"])

# =============================================
# PATHS (BS)
# =============================================
SOURCE_FILE     = path_value(_argv_cfg, "source_file", MASTER_WORKBOOK_STR)
MAPPING_FILE_BS = BS_RECON_MAPPING_FILE
TARGET_FILE     = path_value(_argv_cfg, "target_file", MASTER_WORKBOOK_STR)

REPORT_SHEET_BS   = GROUP_BS_RECON_SHEET
MASTER_SHEET_BS_OUT = "Master_BS"

SOURCE_SHEET_BS = "Master_BS"
SOURCE_HEADER_ROW = 0
SOURCE_COL_CANDIDATES = ("L5", "L6")
REPORTED_FILTER_VALUE = "reported"
L4_SORT_BASIS = "latest_fy"

# =============================================
# LAYOUT — A free | B..F helpers | G spacer | H POS
_COL_LAYOUT = LAYOUT_BS
MAP_START_COL = _COL_LAYOUT.map_start_col
TECH_SPACER_COL = _COL_LAYOUT.spacer_col
POS_COL = _COL_LAYOUT.pos_col
FIRST_ENTITY_COL = _COL_LAYOUT.first_value_col

PROJECT_TITLE_ROW = 1
SUBTITLE_ROW      = 2

ENTITY_CODE_ROW = 3
HEADER_ROW      = 8
BLOCK_TITLE_ROW = HEADER_ROW - 1
DATA_START_ROW  = HEADER_ROW + 1

# BS mapping columns: Reported, blank, L2, L3, L4
BS_MAP_START_COL = MAP_START_COL
BS_MAP_END_COL   = _COL_LAYOUT.map_end_col
BS_L2_COL        = MAP_START_COL + 2
BS_L3_COL        = MAP_START_COL + 3
BS_L4_COL        = MAP_START_COL + 4

SPACER_WIDTH     = 1.14
VALUE_COL_WIDTH  = 7.86
ROW_HEIGHT       = 12
REPEAT_POS_WIDTH = 32

# =============================================
# STYLES — gst_excel_theme / Finssentials output
# =============================================
FONT_BASE = THEME.font_base
FONT_BASE_BOLD = THEME.font_bold
FONT_HEADER = THEME.font_header
FONT_MAPPING = Font(name=THEME.font_name, size=THEME.font_size, color=THEME.text_header)
FONT_TITLE = Font(
    name=THEME.font_name,
    size=THEME.font_size,
    bold=True,
    color=THEME.text_brand_title,
)
FONT_PROJECT_TITLE = Font(
    name=THEME.font_name,
    size=THEME.font_size_title,
    color=THEME.text_brand_title,
    bold=False,
)
FONT_SUBTITLE = Font(
    name=THEME.font_name,
    size=THEME.font_size_subtitle,
    color=THEME.text_brand_title,
    bold=False,
)

FILL_TECH = THEME.fill_tech
FILL_WHITE = THEME.fill_white
FILL_HEADER = THEME.fill_header
FILL_SUBTOTAL = THEME.fill_subtotal
CHECK_FILL = PatternFill("solid", fgColor="FFFEF3C7")

ALIGN_LEFT = Alignment(horizontal="left", vertical="center")
ALIGN_CENTER = Alignment(horizontal="center", vertical="center")
ALIGN_RIGHT = Alignment(horizontal="right", vertical="center")

BORDER_HEADER_BOTTOM = THEME.border_header_bottom
BORDER_SUBTOTAL_TOP = THEME.border_subtotal_top
DIFF_FONT = Font(name=THEME.font_name, size=THEME.font_size, color=THEME.delta_negative)

NUM_FMT_INT = '#,##0;(#,##0);"-"'

# =============================================
# HELPERS
# =============================================
def col_letter(idx: int) -> str:
    return get_column_letter(idx)

def is_blank_entity(x):
    if x is None or pd.isna(x):
        return True
    s = str(x).strip()
    return (s == "") or (s.lower() in {"nan", "none", "null"})

def detect_year_columns(df_: pd.DataFrame) -> list[str]:
    pattern = re.compile(r"^FY(\d{2}|\d{4})A$", re.IGNORECASE)
    years = [str(c).strip() for c in df_.columns if pattern.match(str(c).strip())]
    if not years:
        raise ValueError("Keine FY-Spalten im Master_BS gefunden (erwartet z. B. FY22A, FY23A).")

    def year_sort_key(x: str):
        m = pattern.match(x)
        return int(m.group(1))

    return sorted(dict.fromkeys(years), key=year_sort_key)


def resolve_source_section_column(df_: pd.DataFrame) -> str | None:
    markers = {REPORTED_FILTER_VALUE, "adjusted", "consolidation"}

    def marker_count(col: str) -> int:
        if col not in df_.columns:
            return 0
        series = df_[col].astype(str).str.strip().str.lower()
        return int(series.isin(markers).sum())

    best_col = None
    best_count = 0
    for col in SOURCE_COL_CANDIDATES:
        count = marker_count(col)
        if count > best_count:
            best_col = col
            best_count = count
    if best_col and best_count > 0:
        return best_col

    legacy = detect_source_col(df_)
    if legacy:
        return legacy

    for col in SOURCE_COL_CANDIDATES:
        if col in df_.columns:
            return col
    return None


def detect_source_col(df_):
    candidates = ["Quelle", "Source", "Reported/Adjusted", "Type", "Scenario"]
    sc = next((c for c in candidates if c in df_.columns), None)
    if sc is not None:
        return sc
    for c in df_.columns:
        try:
            s = df_[c].astype(str).str.strip().str.lower()
            if (s == "reported").any():
                return c
        except Exception:
            pass
    return None

def ic_sign(block):
    return -1 if block["kind"] == "ic" else 1

# =============================================
# LOAD BS MAPPING (L2/L3 hierarchy order; L4 ignored)
# =============================================
bs_map_df = pd.read_excel(MAPPING_FILE_BS, sheet_name=0, engine="openpyxl")
bs_map_df = bs_map_df.loc[:, [c for c in bs_map_df.columns if not str(c).startswith("Unnamed")]]
if not {"L2", "L3"}.issubset(bs_map_df.columns):
    raise ValueError("Sortierung_BS_Datenbank muss Spalten 'L2' und 'L3' enthalten.")

bs_map_df = bs_map_df.copy()
bs_map_df["L2"] = bs_map_df["L2"].astype(str)
bs_map_df["L3"] = bs_map_df["L3"].astype(str)
bs_l2_l3_order = l2_l3_order_from_mapping(bs_map_df)



def build_bs_group_blocks(individual_entities: list[str], has_ic: bool) -> list[dict]:
    blocks = []
    for e in individual_entities:
        blocks.append({"kind": "entity", "key": e, "code": e, "title": e, "has_plpos": False})
    if len(individual_entities) > 0:
        blocks.append({"kind": "aggregated", "key": "Aggregated", "code": "Aggregated", "title": "Aggregated", "has_plpos": False})
    if has_ic:
        blocks.append({"kind": "ic", "key": IC_DISPLAY_NAME, "code": IC_MASTER_ENTITY, "title": IC_DISPLAY_NAME, "has_plpos": False})
    if len(individual_entities) > 0 and has_ic:
        blocks.append({"kind": "consolidation", "key": "Consolidation", "code": "Consolidation", "title": "Consolidation", "has_plpos": False})
    blocks.append({"kind": "difference", "key": DIFF_KEY, "code": DIFF_KEY, "title": DIFF_KEY, "has_plpos": True})
    blocks.append({"kind": "fs", "key": FS_KEY, "code": FS_KEY, "title": FS_KEY, "has_plpos": True})
    return assign_recon_block_columns(blocks, first_col=FIRST_ENTITY_COL, n_years=len(YEARS))


def build_bs_entity_blocks(entity: str) -> list[dict]:
    titles = {"difference_title": DIFF_KEY, "financial_statements_title": FS_KEY}
    return build_entity_recon_blocks(
        entity,
        titles,
        first_col=FIRST_ENTITY_COL,
        n_years=len(YEARS),
        difference_key=DIFF_KEY,
        fs_key=FS_KEY,
    )


def build_sorted_bs_entities(df_bs: pd.DataFrame, cfg: dict) -> tuple[list[str], bool]:
    raw_entities = []
    for v in pd.unique(df_bs["Entity"]):
        if not is_blank_entity(v):
            raw_entities.append(str(v).strip())
    has_ic = IC_MASTER_ENTITY in raw_entities
    individual_entities = [e for e in raw_entities if e != IC_MASTER_ENTITY]
    entity_order = list(cfg.get("entity_order") or [])
    if entity_order:
        individual_entities = apply_custom_entity_order(individual_entities, entity_order)
    elif ENTITY_SORT_ORDER:
        def entity_sort_key(e):
            return ENTITY_SORT_ORDER.get(e, raw_entities.index(e) if e in raw_entities else 9999)
        individual_entities = sorted(individual_entities, key=entity_sort_key)
    return individual_entities, has_ic


def insert_after_anchor_bs(struct, anchor_label, new_row):
    for i, row in enumerate(struct):
        if row["type"] == "subtotal_l2" and row["label"] == anchor_label:
            struct.insert(i + 1, new_row)
            return True
    return False


def prepare_bs_row_structure(df_bs: pd.DataFrame) -> list:
    global YEARS, DISPLAY_YEARS, REC_PERIOD_FROM, REC_PERIOD_TO
    YEARS = detect_year_columns(df_bs)
    REC_PERIOD_FROM = YEARS[0]
    REC_PERIOD_TO = YEARS[-1]
    DISPLAY_YEARS = display_bs_period_labels(YEARS, FY_END_MONTH, LTM_MONTH)
    bs_row_structure = build_bs_row_structure(
        df_bs,
        bs_l2_l3_order,
        {"l4_sort_basis": L4_SORT_BASIS},
        source_col=source_col_bs or SOURCE_COL_CANDIDATES[0],
    )
    existing_labels = {r["label"] for r in bs_row_structure}
    for t in BS_TOTALS_CONFIG:
        comps = [c for c in t["components"] if c in existing_labels]
        if not comps:
            continue
        total_row = {
            "type": "total",
            "label": t["label"],
            "components": comps,
            "L2": "",
            "L3": "",
            "L4": "",
        }
        if not insert_after_anchor_bs(bs_row_structure, t["insert_after"], total_row):
            raise RuntimeError(f"Anchor '{t['insert_after']}' for BS total '{t['label']}' not found.")
    return bs_row_structure


def write_bs_reconciliation_sheet(
    wb,
    sheet_name: str,
    *,
    df_bs: pd.DataFrame,
    bs_row_structure: list,
    blocks: list[dict],
    individual_entities: list[str],
    is_entity_sheet: bool = False,
) -> None:
    if sheet_name in wb.sheetnames:
        del wb[sheet_name]
    ws_bs = wb.create_sheet(sheet_name)

    MASTER_START = 2
    MASTER_END = len(df_bs) + 1

    def master_range_bs(colname):
        L = get_column_letter(df_bs.columns.get_loc(colname) + 1)
        return f"{MASTER_SHEET_BS_OUT}!${L}${MASTER_START}:${L}${MASTER_END}"

    src_rng = master_range_bs(source_col_bs) if source_col_bs is not None else None
    ent_rng = master_range_bs("Entity")
    l2_rng = master_range_bs("L2")
    l3_rng = master_range_bs("L3")
    l4_rng = master_range_bs("L4")

    LAST_USED_COL = blocks[-1]["spacer_col"] if blocks else FIRST_ENTITY_COL

    pt = ws_bs.cell(PROJECT_TITLE_ROW, POS_COL, f"Project {PROJECT_NAME}")
    pt.font = FONT_PROJECT_TITLE
    pt.alignment = ALIGN_LEFT
    ws_bs.row_dimensions[PROJECT_TITLE_ROW].height = 36

    title_name = (
        individual_entities[0]
        if is_entity_sheet and individual_entities
        else GROUP_NAME
    )

    st = ws_bs.cell(SUBTITLE_ROW, POS_COL, f"{title_name} Reconciliation - Balance sheet")
    st.font = FONT_SUBTITLE
    st.alignment = ALIGN_LEFT

    rec_title_text = (
        f"{title_name} | Reconciliation (Trial Balances) "
        f"{REC_PERIOD_FROM} - {REC_PERIOD_TO}"
    )
    rec_title = ws_bs.cell(6, POS_COL, rec_title_text)
    rec_title.font = FONT_TITLE
    rec_title.alignment = ALIGN_LEFT

    ws_bs.cell(HEADER_ROW, BS_MAP_START_COL + 0, "Reported").font = FONT_HEADER
    ws_bs.cell(HEADER_ROW, BS_MAP_START_COL + 1, "").font = FONT_HEADER
    ws_bs.cell(HEADER_ROW, BS_L2_COL, "L2").font = FONT_HEADER
    ws_bs.cell(HEADER_ROW, BS_L3_COL, "L3").font = FONT_HEADER
    ws_bs.cell(HEADER_ROW, BS_L4_COL, "L4").font = FONT_HEADER
    for c in range(BS_MAP_START_COL, BS_MAP_END_COL + 1):
        ws_bs.cell(HEADER_ROW, c).alignment = ALIGN_LEFT

    ph = ws_bs.cell(HEADER_ROW, POS_COL, UNIT_LABEL)
    ph.font = FONT_HEADER
    ph.alignment = ALIGN_LEFT

    for b in blocks:
        if b.get("has_plpos"):
            ws_bs.merge_cells(start_row=BLOCK_TITLE_ROW, start_column=b["year_startcol"],
                              end_row=BLOCK_TITLE_ROW, end_column=b["year_endcol"])
            t = ws_bs.cell(BLOCK_TITLE_ROW, b["year_startcol"], b["title"])
            t.font = FONT_BASE_BOLD
            t.alignment = ALIGN_CENTER
            if b.get("kind") not in {"difference", "fs"}:
                ws_bs.cell(BLOCK_TITLE_ROW, b["poscol"]).value = None
                ws_bs.cell(BLOCK_TITLE_ROW, b["poscol"]).fill = FILL_WHITE
        else:
            ws_bs.merge_cells(start_row=BLOCK_TITLE_ROW, start_column=b["startcol"],
                              end_row=BLOCK_TITLE_ROW, end_column=b["year_endcol"])
            t = ws_bs.cell(BLOCK_TITLE_ROW, b["startcol"], b["title"])
            t.font = FONT_BASE_BOLD
            t.alignment = ALIGN_CENTER

        if b.get("has_plpos"):
            ws_bs.merge_cells(start_row=ENTITY_CODE_ROW, start_column=b["year_startcol"],
                              end_row=ENTITY_CODE_ROW, end_column=b["year_endcol"])
            ec = ws_bs.cell(ENTITY_CODE_ROW, b["year_startcol"], b["code"])
        else:
            ws_bs.merge_cells(start_row=ENTITY_CODE_ROW, start_column=b["startcol"],
                              end_row=ENTITY_CODE_ROW, end_column=b["year_endcol"])
            ec = ws_bs.cell(ENTITY_CODE_ROW, b["startcol"], b["code"])
        ec.font = FONT_BASE
        ec.alignment = ALIGN_CENTER

        if b.get("has_plpos"):
            hpos = ws_bs.cell(HEADER_ROW, b["poscol"], UNIT_LABEL)
            hpos.font = FONT_HEADER
            hpos.alignment = ALIGN_LEFT

        for y_idx, year in enumerate(YEARS):
            h = ws_bs.cell(HEADER_ROW, b["year_startcol"] + y_idx, DISPLAY_YEARS[y_idx])
            h.font = FONT_HEADER
            h.alignment = ALIGN_RIGHT
        ws_bs.column_dimensions[col_letter(b["spacer_col"])].width = SPACER_WIDTH

    for i, r in enumerate(bs_row_structure):
        r["_idx"] = i
        excel_row = DATA_START_ROW + i
        r["_excel_row"] = excel_row
        ws_bs.cell(excel_row, BS_MAP_START_COL + 0, "Reported").font = FONT_MAPPING
        ws_bs.cell(excel_row, BS_MAP_START_COL + 1, "").font = FONT_MAPPING
        ws_bs.cell(excel_row, BS_L2_COL, r["L2"]).font = FONT_MAPPING
        ws_bs.cell(excel_row, BS_L3_COL, r["L3"]).font = FONT_MAPPING
        ws_bs.cell(excel_row, BS_L4_COL, r["L4"]).font = FONT_MAPPING
        for c in range(BS_MAP_START_COL, BS_MAP_END_COL + 1):
            ws_bs.cell(excel_row, c).alignment = ALIGN_LEFT
        pc = ws_bs.cell(excel_row, POS_COL, r["label"])
        if r["type"] == "detail" and str(r["L3"]).strip() != str(r["L4"]).strip():
            pc.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        else:
            pc.alignment = ALIGN_LEFT
        pc.font = FONT_BASE_BOLD if r["type"] in {"subtotal_l2", "total"} else FONT_BASE
        for b in blocks:
            if not b.get("has_plpos"):
                continue
            pcell = ws_bs.cell(excel_row, b["poscol"])
            pcell.value = f"={col_letter(POS_COL)}{excel_row}"
            if r["type"] == "detail" and str(r["L3"]).strip() != str(r["L4"]).strip():
                pcell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
            else:
                pcell.alignment = ALIGN_LEFT
            pcell.font = FONT_BASE_BOLD if r["type"] in {"subtotal_l2", "total"} else FONT_BASE
        ws_bs.row_dimensions[excel_row].height = ROW_HEIGHT
        if r["type"] == "detail":
            ws_bs.row_dimensions[excel_row].outlineLevel = 1
            ws_bs.row_dimensions[excel_row].hidden = False
        elif r["type"] == "detail_single" and str(r["L3"]).strip() != str(r["L4"]).strip():
            ws_bs.row_dimensions[excel_row].outlineLevel = 1
            ws_bs.row_dimensions[excel_row].hidden = False
        elif r["type"] in {"detail_single", "subtotal_l3", "subtotal_l2", "total"}:
            ws_bs.row_dimensions[excel_row].outlineLevel = 0
            ws_bs.row_dimensions[excel_row].hidden = False

    bs_l2_col = col_letter(BS_L2_COL)
    bs_l3_col = col_letter(BS_L3_COL)
    bs_l4_col = col_letter(BS_L4_COL)

    def block_lookup(key):
        return next(bb for bb in blocks if bb["key"] == key)

    for b in blocks:
        if b["kind"] in {"entity", "ic"}:
            entity_ref = f"${col_letter(b['startcol'])}${ENTITY_CODE_ROW}"
        else:
            entity_ref = None
        sign = ic_sign(b)
        for r in bs_row_structure:
            excel_row = r["_excel_row"]
            for y_idx, year in enumerate(YEARS):
                cell = ws_bs.cell(excel_row, b["year_startcol"] + y_idx)
                cell.alignment = ALIGN_RIGHT
                cell.number_format = NUM_FMT_INT
                cell.font = FONT_BASE_BOLD if r["type"] in {"subtotal_l2", "total"} else FONT_BASE
                if b["kind"] in {"entity", "ic"}:
                    if r["type"] in {"detail", "detail_single"}:
                        sum_rng = master_range_bs(year)
                        l2_crit = f"${bs_l2_col}${excel_row}"
                        l3_crit = f"${bs_l3_col}${excel_row}"
                        l4_crit = f"${bs_l4_col}${excel_row}"
                        bs_rep_col = col_letter(BS_MAP_START_COL)
                        rep_crit = f"${bs_rep_col}${excel_row}"
                        if src_rng is not None:
                            cell.value = (
                                f"={sign}*SUMIFS({sum_rng},{ent_rng},{entity_ref},"
                                f"{src_rng},{rep_crit},{l2_rng},{l2_crit},{l3_rng},{l3_crit},"
                                f"{l4_rng},{l4_crit})/1000"
                            )
                        else:
                            cell.value = (
                                f"={sign}*SUMIFS({sum_rng},{ent_rng},{entity_ref},"
                                f"{l2_rng},{l2_crit},{l3_rng},{l3_crit},{l4_rng},{l4_crit})/1000"
                            )
                    elif r["type"] == "subtotal_l3":
                        idx = r["_idx"]
                        k = idx - 1
                        while k >= 0 and bs_row_structure[k]["type"] == "detail" and bs_row_structure[k]["L3"] == r["L3"] and bs_row_structure[k]["L2"] == r["L2"]:
                            k -= 1
                        detail_start_idx = k + 1
                        detail_end_idx = idx - 1
                        colL = col_letter(b["year_startcol"] + y_idx)
                        if detail_end_idx >= detail_start_idx:
                            start_row = DATA_START_ROW + detail_start_idx
                            end_row = DATA_START_ROW + detail_end_idx
                            cell.value = f"=SUM({colL}{start_row}:{colL}{end_row})"
                        else:
                            cell.value = 0
                    elif r["type"] == "subtotal_l2":
                        colL = col_letter(b["year_startcol"] + y_idx)
                        refs = []
                        for rr in bs_row_structure:
                            if rr["L2"] != r["L2"]:
                                continue
                            if rr["type"] == "subtotal_l3":
                                refs.append(f"{colL}{rr['_excel_row']}")
                            elif rr["type"] == "detail_single" and str(rr["L3"]).strip() == str(rr["L4"]).strip():
                                refs.append(f"{colL}{rr['_excel_row']}")
                        cell.value = f"=SUM({','.join(refs)})" if refs else 0
                    elif r["type"] == "total":
                        colL = col_letter(b["year_startcol"] + y_idx)
                        refs = []
                        for comp in r.get("components", []):
                            for rr in bs_row_structure:
                                if rr["type"] == "subtotal_l2" and rr["label"] == comp:
                                    refs.append(f"{colL}{rr['_excel_row']}")
                                    break
                        cell.value = f"=SUM({','.join(refs)})" if refs else 0
                elif b["kind"] == "aggregated":
                    refs = []
                    for e in individual_entities:
                        bb = block_lookup(e)
                        refs.append(f"{col_letter(bb['year_startcol'] + y_idx)}{excel_row}")
                    cell.value = f"=SUM({','.join(refs)})" if refs else 0
                elif b["kind"] == "consolidation":
                    agg_block = block_lookup("Aggregated")
                    ic_block = block_lookup(IC_DISPLAY_NAME)
                    agg_ref = f"{col_letter(agg_block['year_startcol'] + y_idx)}{excel_row}"
                    ic_ref = f"{col_letter(ic_block['year_startcol'] + y_idx)}{excel_row}"
                    cell.value = f"={agg_ref}+{ic_ref}"
                elif b["kind"] == "fs":
                    cell.value = None
                elif b["kind"] == "difference":
                    fs_block = block_lookup(FS_KEY)
                    fs_ref = f"{col_letter(fs_block['year_startcol'] + y_idx)}{excel_row}"
                    if any(bb["kind"] == "consolidation" for bb in blocks):
                        con_block = block_lookup("Consolidation")
                        con_ref = f"{col_letter(con_block['year_startcol'] + y_idx)}{excel_row}"
                        cell.value = f"={fs_ref}-{con_ref}"
                    elif is_entity_sheet:
                        ent_block = next(bb for bb in blocks if bb["kind"] == "entity")
                        ent_ref = f"{col_letter(ent_block['year_startcol'] + y_idx)}{excel_row}"
                        cell.value = f"={fs_ref}-{ent_ref}"
                    else:
                        cell.value = None

    LAST_TABLE_ROW = DATA_START_ROW + len(bs_row_structure) - 1
    spacer_cols = {b["spacer_col"] for b in blocks}
    for r in bs_row_structure:
        excel_row = r["_excel_row"]
        if r["type"] in {"subtotal_l2", "total"}:
            fill = FILL_SUBTOTAL
            ws_bs.cell(excel_row, POS_COL).fill = fill
            for c in range(POS_COL + 1, LAST_USED_COL + 1):
                if c not in spacer_cols:
                    ws_bs.cell(excel_row, c).fill = fill
        else:
            for c in range(POS_COL, LAST_USED_COL + 1):
                if c not in spacer_cols:
                    ws_bs.cell(excel_row, c).fill = FILL_WHITE

    layout = LAYOUT_BS
    apply_bs_hierarchy_borders(
        ws_bs,
        bs_row_structure,
        layout=layout,
        last_used_col=LAST_USED_COL,
        skip_cols=spacer_cols,
    )
    LAST_CONTENT_ROW = LAST_TABLE_ROW

    total_assets_row = next((r["_excel_row"] for r in bs_row_structure if r["type"] == "total" and r["label"] == "Total assets"), None)
    total_el_row = next((r["_excel_row"] for r in bs_row_structure if r["type"] == "total" and r["label"] == "Total equity & liabilities"), None)

    if SHOW_BS_CHECKS and total_assets_row is not None and total_el_row is not None:
        _check_rows = check_row_groups_after_table(LAST_TABLE_ROW, [1, 2])
        CHECK_ALE_ROW, CHECK_FS_ROW, CHECK_DELTA_ROW = _check_rows
        LAST_CONTENT_ROW = CHECK_DELTA_ROW
    else:
        CHECK_ALE_ROW = CHECK_FS_ROW = CHECK_DELTA_ROW = None

    ws_bs.column_dimensions[col_letter(POS_COL)].width = 32
    for b in blocks:
        if b.get("has_plpos"):
            ws_bs.column_dimensions[col_letter(b["poscol"])].width = REPEAT_POS_WIDTH
        for y_idx in range(len(YEARS)):
            ws_bs.column_dimensions[col_letter(b["year_startcol"] + y_idx)].width = VALUE_COL_WIDTH
        ws_bs.column_dimensions[col_letter(b["spacer_col"])].width = SPACER_WIDTH

    FILL_END_ROW = LAST_CONTENT_ROW + 100
    FILL_END_COL = LAST_USED_COL + 40
    paint_grey_white_canvas(ws_bs, layout, last_row=FILL_END_ROW, last_col=LAST_USED_COL)
    for rr in range(DATA_START_ROW, LAST_TABLE_ROW + 1):
        for cc in range(BS_MAP_START_COL, BS_MAP_END_COL + 1):
            ws_bs.cell(rr, cc).font = FONT_MAPPING
    for sc in spacer_cols:
        for rr in range(1, FILL_END_ROW + 1):
            ws_bs.cell(rr, sc).border = Border()
            ws_bs.cell(rr, sc).fill = FILL_WHITE
    apply_recon_portfolio_layout(
        ws_bs,
        blocks=blocks,
        pos_col=POS_COL,
        header_row=HEADER_ROW,
        block_title_row=BLOCK_TITLE_ROW,
        entity_code_row=ENTITY_CODE_ROW,
        last_used_col=LAST_USED_COL,
        spacer_cols=spacer_cols,
        visible_block_kinds=frozenset({"entity", "difference", "fs"}) if is_entity_sheet else frozenset({"consolidation", "difference", "fs"}),
        tech_layout=layout,
    )
    ws_bs.row_dimensions[HEADER_ROW].height = ROW_HEIGHT
    ws_bs.row_dimensions[BLOCK_TITLE_ROW].height = ROW_HEIGHT

    if SHOW_BS_CHECKS and total_assets_row is not None and total_el_row is not None:
        check_block_kinds = frozenset({"entity"}) if is_entity_sheet else frozenset({"entity", "ic", "aggregated", "consolidation"})
        fs_values: dict[str, dict[str, float | None]] = {}
        for b in blocks:
            if b["kind"] not in check_block_kinds:
                continue
            fy_vals = BS_FS_CHECK_TOTAL_ASSETS.get(b["key"], {})
            fs_values[b["key"]] = {
                year: (fy_vals.get(year) / 1000 if fy_vals.get(year) is not None else None)
                for year in YEARS
            }
        write_bs_ale_check_section(
            ws_bs, blocks=blocks, years=YEARS, check_row=CHECK_ALE_ROW,
            total_assets_row=total_assets_row, total_el_row=total_el_row, pos_col=POS_COL,
            block_kinds=check_block_kinds, num_fmt=NUM_FMT_INT,
        )
        write_fs_check_section(
            ws_bs, blocks=blocks, years=YEARS, source_row=CHECK_FS_ROW, delta_row=CHECK_DELTA_ROW,
            anchor_row=total_assets_row, pos_col=POS_COL, block_kinds=check_block_kinds,
            yellow_block_kinds=check_block_kinds, spacer_cols=spacer_cols,
            source_values=fs_values if BS_FS_CHECK_TOTAL_ASSETS else None,
        )
        collapse_check_portfolio(ws_bs, [CHECK_ALE_ROW, CHECK_FS_ROW, CHECK_DELTA_ROW])


def main():
    global df_bs, source_col_bs, YEARS, DISPLAY_YEARS, REC_PERIOD_FROM, REC_PERIOD_TO
    global PROJECT_NAME, GROUP_NAME

    cfg = dict(_argv_cfg)
    PROJECT_NAME = str(cfg.get("project_name") or PROJECT_NAME).strip()
    GROUP_NAME = str(cfg.get("company_name") or GROUP_NAME).strip()

    df_bs = pd.read_excel(SOURCE_FILE, sheet_name=SOURCE_SHEET_BS, engine="openpyxl", header=SOURCE_HEADER_ROW)
    source_col_bs = resolve_source_section_column(df_bs)
    print(f"Verwende Source-Spalte für Reported/Adjusted: {source_col_bs}")

    required_bs = {"Entity", "L2", "L3", "L4"} | set(detect_year_columns(df_bs))
    missing_bs = required_bs - set(df_bs.columns)
    if missing_bs:
        raise ValueError(f"Master_BS fehlt Spalten: {missing_bs}")

    df_bs = df_bs.copy()

    def normalize_entity_keep_ic(x):
        if x is None or pd.isna(x):
            return x
        s = str(x).strip()
        if s == IC_MASTER_ENTITY:
            return s
        return ENTITY_RENAME_MAP.get(s, s) if ENTITY_RENAME_MAP else s

    if ENTITY_RENAME_MAP:
        df_bs["Entity"] = df_bs["Entity"].apply(normalize_entity_keep_ic)

    bs_row_structure = prepare_bs_row_structure(df_bs)
    individual_entities, has_ic = build_sorted_bs_entities(df_bs, cfg)

    from funktionssammlung import open_report_workbook

    wb = open_report_workbook(cfg, TARGET_FILE)

    remove_stale_entity_recon_sheets(wb, individual_entities)

    if not bool(cfg.get("preserve_master_sheets")):
        refresh_master_sheet(wb, df_bs, MASTER_SHEET_BS_OUT)

    group_blocks = build_bs_group_blocks(individual_entities, has_ic)
    write_bs_reconciliation_sheet(
        wb, REPORT_SHEET_BS, df_bs=df_bs, bs_row_structure=bs_row_structure,
        blocks=group_blocks, individual_entities=individual_entities, is_entity_sheet=False,
    )

    for entity in individual_entities:
        entity_sheet = sanitize_entity_recon_sheet_name(entity, "bs")
        entity_blocks = build_bs_entity_blocks(entity)
        write_bs_reconciliation_sheet(
            wb, entity_sheet, df_bs=df_bs, bs_row_structure=bs_row_structure,
            blocks=entity_blocks, individual_entities=[entity], is_entity_sheet=True,
        )

    wb.save(TARGET_FILE)
    print(f"Saved: {TARGET_FILE}")


if __name__ == "__main__":
    main()
