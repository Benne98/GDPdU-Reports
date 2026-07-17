"""Shared PL reconciliation references and revenue bottom-row helpers."""
from __future__ import annotations

import re
from dataclasses import dataclass
from typing import Iterable

from openpyxl.utils import get_column_letter


PL_RECON_SHEET = "PL_Reconciliation"
AGGREGATED_BLOCK_TITLE = "Aggregated"
TOTAL_NET_SALES_LABEL = "Total Net sales"
RECON_DIFFERENCE_LABEL = "Recon. difference"
REPORTED_NET_SALES_LABEL = "Reported Net sales"
NA_VALUE = "n/a"

VALID_SALES_BASES = frozenset({"net", "gross"})


def _norm(value) -> str:
    text = str(value or "").replace("\u00a0", " ").strip().casefold()
    return re.sub(r"[^a-z0-9]+", "", text)


def normalize_sales_basis(value) -> str:
    basis = str(value or "net").strip().lower()
    return basis if basis in VALID_SALES_BASES else "net"


@dataclass(frozen=True)
class SalesBasisLabels:
    sales_basis: str
    metric_label: str
    total_label: str
    reported_label: str
    profit_label: str
    margin_label: str
    profit_abbrev: str
    margin_abbrev: str
    bridge_sheet_title: str
    bridge_table_title: str
    pl_row_needles: tuple[str, ...]


def sales_basis_labels(sales_basis: str | None = "net") -> SalesBasisLabels:
    basis = normalize_sales_basis(sales_basis)
    if basis == "gross":
        return SalesBasisLabels(
            sales_basis="gross",
            metric_label="Gross sales",
            total_label="Total Gross sales",
            reported_label="Reported Gross sales",
            profit_label="Gross profit",
            margin_label="Gross margin",
            profit_abbrev="GP",
            margin_abbrev="GM",
            bridge_sheet_title="GP Bridge",
            bridge_table_title="Gross Profit Bridge",
            pl_row_needles=("grosssales",),
        )
    return SalesBasisLabels(
        sales_basis="net",
        metric_label="Net sales",
        total_label=TOTAL_NET_SALES_LABEL,
        reported_label=REPORTED_NET_SALES_LABEL,
        profit_label="Net profit",
        margin_label="Net margin",
        profit_abbrev="NP",
        margin_abbrev="NM",
        bridge_sheet_title="NP Bridge",
        bridge_table_title="Net Profit Bridge",
        pl_row_needles=("netsales",),
    )


def canonical_period(value) -> str | None:
    """Return a common FY/YTD/LTM key for Excel header variants."""
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        year = int(value)
        if 1900 <= year <= 2200:
            return f"FY{year % 100:02d}A"

    raw = str(value).replace("\u00a0", " ").strip().upper()
    compact = re.sub(r"[^A-Z0-9]+", "", raw)
    match = re.fullmatch(r"(FY|YTD|LTM)?(\d{2}|\d{4})(A)?", compact)
    if not match:
        return None
    prefix = match.group(1) or "FY"
    year = int(match.group(2))
    if year >= 100:
        year %= 100
    return f"{prefix}{year:02d}A"


@dataclass(frozen=True)
class PLNetSalesReference:
    period: str
    row: int
    column: int
    sheet_name: str = PL_RECON_SHEET

    @property
    def coordinate(self) -> str:
        return f"{get_column_letter(self.column)}{self.row}"

    @property
    def formula(self) -> str:
        escaped = self.sheet_name.replace("'", "''")
        return f"='{escaped}'!{self.coordinate}"


# Backward-compatible alias.
PLSalesReference = PLNetSalesReference


def _find_sales_metric_row(ws, needles: tuple[str, ...]) -> int | None:
    needle_set = {_norm(n) for n in needles if _norm(n)}
    if not needle_set:
        return None
    exact: list[int] = []
    contains: list[int] = []
    for row in ws.iter_rows():
        for cell in row:
            token = _norm(cell.value)
            if not token:
                continue
            if token in needle_set:
                exact.append(cell.row)
            elif any(n in token for n in needle_set):
                contains.append(cell.row)
    if exact:
        return min(exact)
    return min(contains) if contains else None


def _find_net_sales_row(ws) -> int | None:
    return _find_sales_metric_row(ws, ("netsales",))


def _block_title_row(ws) -> int | None:
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row or 1, 20)):
        for cell in row:
            if _norm(cell.value) == "aggregated":
                return cell.row
    return None


def _aggregated_block_range(ws) -> tuple[int | None, int]:
    """Return (start_col, end_col_exclusive) for the Aggregated portfolio block."""
    title_row = _block_title_row(ws)
    if title_row is None:
        return None, (ws.max_column or 0) + 1

    titles: list[tuple[int, str]] = []
    for cc in range(1, (ws.max_column or 0) + 1):
        raw = ws.cell(title_row, cc).value
        if isinstance(raw, str) and raw.strip():
            titles.append((cc, raw.strip()))

    block_start = None
    block_end = (ws.max_column or 0) + 1
    for idx, (cc, title) in enumerate(titles):
        if _norm(title) == "aggregated":
            block_start = cc
            if idx + 1 < len(titles):
                block_end = titles[idx + 1][0]
            break
    return block_start, block_end


def _period_columns_in_aggregated_block(ws) -> dict[str, int]:
    """Map canonical periods → first (leftmost) column inside Aggregated only."""
    block_start, block_end = _aggregated_block_range(ws)
    if block_start is None:
        return {}

    title_row = _block_title_row(ws) or 1
    candidates: list[tuple[int, int, str]] = []
    row_start = max(1, title_row - 1)
    row_end = min(ws.max_row or 1, title_row + 6)
    for row_idx in range(row_start, row_end + 1):
        for col_idx in range(block_start, block_end):
            period = canonical_period(ws.cell(row_idx, col_idx).value)
            if period:
                candidates.append((row_idx, col_idx, period))
    if not candidates:
        return {}

    row_counts: dict[int, int] = {}
    for row_idx, _col_idx, _period in candidates:
        row_counts[row_idx] = row_counts.get(row_idx, 0) + 1
    header_row = max(row_counts, key=lambda r: (row_counts[r], -abs(r - title_row)))

    # Keep the first (leftmost) column per period — never later FS/entity blocks.
    period_cols: dict[str, int] = {}
    for row_idx, col_idx, period in sorted(candidates, key=lambda t: t[1]):
        if row_idx != header_row:
            continue
        period_cols.setdefault(period, col_idx)
    return period_cols


def resolve_pl_sales_references(
    workbook,
    periods: Iterable[str] | None = None,
    *,
    sales_basis: str | None = "net",
    sheet_name: str = PL_RECON_SHEET,
) -> dict[str, PLSalesReference]:
    """Resolve Aggregated Net/Gross sales cells without fixed row or column numbers."""
    if workbook is None or sheet_name not in workbook.sheetnames:
        return {}
    ws = workbook[sheet_name]
    labels = sales_basis_labels(sales_basis)
    sales_row = _find_sales_metric_row(ws, labels.pl_row_needles)
    if sales_row is None:
        return {}

    period_cols = _period_columns_in_aggregated_block(ws)
    if not period_cols:
        return {}

    requested = None
    if periods is not None:
        requested = {canonical_period(period) for period in periods}
        requested.discard(None)

    return {
        period: PLSalesReference(period, sales_row, col_idx, sheet_name)
        for period, col_idx in period_cols.items()
        if requested is None or period in requested
    }


def resolve_pl_net_sales_references(
    workbook,
    periods: Iterable[str] | None = None,
    *,
    sheet_name: str = PL_RECON_SHEET,
) -> dict[str, PLNetSalesReference]:
    """Backward-compatible Net-sales resolver."""
    return resolve_pl_sales_references(
        workbook,
        periods,
        sales_basis="net",
        sheet_name=sheet_name,
    )


def write_reconciliation_cells(
    workbook,
    worksheet,
    *,
    period_columns: dict[str, int],
    total_row: int,
    recon_row: int,
    reported_row: int,
    sales_basis: str | None = "net",
) -> dict[str, PLSalesReference]:
    """Write PL references and Reported-minus-Total formulas, else n/a."""
    references = resolve_pl_sales_references(
        workbook,
        period_columns,
        sales_basis=sales_basis,
    )
    for period, col_idx in period_columns.items():
        key = canonical_period(period)
        ref = references.get(key or "")
        if ref is None:
            worksheet.cell(reported_row, col_idx).value = NA_VALUE
            worksheet.cell(recon_row, col_idx).value = NA_VALUE
            continue
        worksheet.cell(reported_row, col_idx).value = ref.formula
        col_letter = get_column_letter(col_idx)
        worksheet.cell(recon_row, col_idx).value = (
            f"={col_letter}{reported_row}-{col_letter}{total_row}"
        )
    return references
